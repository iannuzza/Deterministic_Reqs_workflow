import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from validate_downstream_coherence import (
    _approved_upstream_coverage,
    _check_srs_artifact_coherence,
    _check_sysml_requirement_coverage,
    _primary_to_supplementary_coverage,
    _source_spec_node_metrics,
    _source_spec_target_metrics,
    downstream_contract_fingerprint,
    _extract_srs_requirement_ids,
)
from generate_snapshot_reports import (
    _approved_supplementary_imports,
    _classify_integrated_source_rows,
    _generated_requirement_statement,
    _hierarchy_level_label,
    _snapshot_statement_for,
    _snapshot_statement_index,
)
from run_srs_gen_spec_agent import _filter_mode_rows_to_srs_allocations, _write_stage_report
from report_generation_service import generate_snapshot_report
from workflow_gui import WorkflowGUI, tk


class EmptySrsContract:
    snapshot_id = "snap-test"
    expected = {}

    def source_ids_for_target(self, target):
        return set()


class SrsDownstreamCoherenceTests(unittest.TestCase):
    def test_traceability_boxes_contain_wrapped_text_at_every_zoom(self):
        root = tk.Tk()
        root.withdraw()
        self.addCleanup(root.destroy)
        canvas = tk.Canvas(root)
        gui = object.__new__(WorkflowGUI)
        gui.traceability_selected_edge = None
        gui.traceability_selected_block = None
        graph = {
            "nodes": [
                {
                    "id": "supplementary:long",
                    "kind": "supplementary_spec",
                    "label": "Supplementary specification: " + "LongSpecificationName" * 8,
                    "upstream_coverage": ["Primary coverage (" + "Extended name " * 20 + "): 123 / 123 (100.0%)"],
                },
                {"id": "supplementary:short", "kind": "supplementary_spec", "label": "Short specification"},
            ],
            "edges": [], "source_spec_edges": [], "source_spec_relationship_edges": [],
        }
        with patch("workflow_gui.approved_generated_document_dependency_graph", return_value=graph):
            for zoom in (0.65, 1.0, 1.5, 2.25):
                with self.subTest(zoom=zoom):
                    gui.traceability_zoom = zoom
                    gui._draw_generated_dependency_graph(canvas)
                    rectangles = []
                    for node in graph["nodes"]:
                        items = canvas.find_withtag("traceability_node:" + node["id"])
                        rectangle = next(item for item in items if canvas.type(item) == "rectangle")
                        text = next(item for item in items if canvas.type(item) == "text")
                        box = canvas.coords(rectangle)
                        bounds = canvas.bbox(text)
                        self.assertLessEqual(box[0], bounds[0])
                        self.assertLessEqual(box[1], bounds[1])
                        self.assertGreaterEqual(box[2], bounds[2])
                        self.assertGreaterEqual(box[3], bounds[3])
                        rectangles.append(box)
                    self.assertLess(rectangles[0][3], rectangles[1][1])

    def test_traceability_table_and_diagram_share_node_selection(self):
        graph = {
            "nodes": [
                {"id": "source:main", "kind": "source_spec", "label": "Main"},
                {"id": "document:child", "kind": "document", "label": "Child"},
            ],
            "edges": [{"source": "source:main", "target": "document:child", "coverage": {"covered": 1, "total": 2}}],
            "source_spec_edges": [],
            "source_spec_relationship_edges": [],
        }
        gui = object.__new__(WorkflowGUI)
        gui.traceability_tree = MagicMock()
        gui.traceability_tree.exists.return_value = True
        gui.traceability_window = MagicMock()
        gui.traceability_graph_canvas = MagicMock()
        gui.traceability_graph_canvas.bbox.return_value = (0, 0, 200, 54)
        gui.traceability_window_status_var = MagicMock()
        gui.traceability_selected_edge = None
        gui.traceability_selected_block = None
        gui.traceability_zoom = 1.0
        font = MagicMock()
        font.measure.side_effect = lambda text: len(text) * 7
        with patch("workflow_gui.approved_generated_document_dependency_graph", return_value=graph), \
             patch("workflow_gui.tkfont.Font", return_value=font):
            gui.traceability_tree.selection.return_value = ("document:child",)
            gui._on_traceability_tree_select()
            self.assertEqual(gui.traceability_selected_block, "document:child")
            self.assertEqual(gui.traceability_graph_canvas.create_line.call_args.kwargs["fill"], "#005fbd")
            self.assertEqual(gui.traceability_graph_canvas.create_rectangle.call_args.kwargs["fill"], "#cfe8ff")

            gui.traceability_graph_canvas.find_overlapping.return_value = (42,)
            gui.traceability_graph_canvas.gettags.return_value = ("traceability_node:source:main",)
            gui._on_traceability_canvas_click(SimpleNamespace(x=1, y=1))
            self.assertEqual(gui.traceability_selected_block, "source:main")
            gui.traceability_tree.selection_set.assert_called_with("source:main")
            gui.traceability_tree.insert.assert_any_call(
                "", "end", iid="source:main", text="Main",
                values=("Source ledger: Primary spec", "", "", ""),
                tags=("traceability_even",), open=True,
            )
            self.assertEqual(gui.traceability_graph_canvas.create_line.call_args.kwargs["fill"], "#005fbd")

            gui.traceability_graph_canvas.gettags.return_value = (
                "traceability_edge:source:main->document:child:0",
                "traceability_target:document:child",
            )
            gui._on_traceability_canvas_click(SimpleNamespace(x=1, y=1))
            self.assertEqual(gui.traceability_selected_block, "document:child")
            gui.traceability_tree.selection_set.assert_called_with("document:child")
            self.assertEqual(gui.traceability_graph_canvas.create_line.call_args.kwargs["fill"], "#005fbd")
            self.assertIn("traceability_target:document:child", gui.traceability_graph_canvas.create_line.call_args.kwargs["tags"])

            gui.traceability_tree.selection.return_value = ()
            gui._on_traceability_tree_select()
            self.assertEqual(gui.traceability_selected_block, "document:child")

    def test_supplementary_primary_coverage_appears_in_box_and_table(self):
        node = {
            "id": "supplementary:extra.pdf",
            "kind": "supplementary_spec",
            "label": "Supplementary specification: extra.pdf",
            "upstream_coverage_metrics": [{"upstream": "main.pdf", "covered": 3, "total": 4, "percentage": 75.0}],
            "upstream_coverage": ["Primary coverage (main.pdf): 3 / 4 (75.0%)"],
        }
        graph = {"nodes": [node], "edges": [], "source_spec_edges": [], "source_spec_relationship_edges": []}
        gui = object.__new__(WorkflowGUI)
        gui.traceability_tree = MagicMock()
        gui.traceability_window = MagicMock()
        gui.traceability_graph_canvas = MagicMock()
        gui.traceability_graph_canvas.bbox.return_value = (0, 0, 200, 54)
        gui.traceability_window_status_var = MagicMock()
        gui.traceability_selected_edge = None
        gui.traceability_selected_block = None
        gui.traceability_zoom = 1.0
        font = MagicMock()
        font.measure.side_effect = lambda text: len(text) * 7
        with patch("workflow_gui.approved_generated_document_dependency_graph", return_value=graph), \
             patch("workflow_gui.tkfont.Font", return_value=font):
            gui._refresh_traceability_window()

        values = gui.traceability_tree.insert.call_args.kwargs["values"]
        self.assertEqual(values[:3], ("main.pdf", "3 / 4", "75.0%"))
        self.assertTrue(any(
            "Primary coverage (main.pdf): 3 / 4 (75.0%)" in call.kwargs.get("text", "")
            for call in gui.traceability_graph_canvas.create_text.call_args_list
        ))

    def test_supplementary_primary_coverage_uses_target_denominator_only_when_linked(self):
        edges = [{
            "source": "source:main.pdf",
            "target": "supplementary:extra.pdf",
            "source_coverage": {"covered": 2, "total": 10, "percentage": 20.0},
            "target_coverage": {"covered": 3, "total": 4, "percentage": 75.0},
        }]
        self.assertEqual(
            _primary_to_supplementary_coverage("source:main.pdf", "supplementary:extra.pdf", "main.pdf", edges),
            {"upstream": "main.pdf", "covered": 3, "total": 4, "percentage": 75.0},
        )
        self.assertIsNone(
            _primary_to_supplementary_coverage("source:main.pdf", "supplementary:other.pdf", "main.pdf", edges)
        )
        self.assertIsNone(
            _primary_to_supplementary_coverage("source:other.pdf", "supplementary:extra.pdf", "other.pdf", edges)
        )

    def test_stage1_catalog_export_preserves_csv_layout_and_labels_authority(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "catalog.xlsx"
            catalog = Path(directory) / "source.csv"
            rows = _classify_integrated_source_rows(
                [
                    {"id": "REQ-1", "notes": "Kept", "source_spec": "", "category": "System"},
                    {"id": "SUP-1", "notes": "Included", "source_spec": "Extra.pdf", "category": "System"},
                ],
                {"REQ-1"},
                "Main.pdf",
            )
            resolved = SimpleNamespace(snapshot_id="snap-test", snapshot_hash="hash", source_path="snapshot")
            connection = MagicMock()
            with patch("report_generation_service.resolve_complete_authoritative_input", return_value=(resolved, "selected")), \
                 patch("report_generation_service.connect", return_value=connection), \
                 patch("report_generation_service.record_event"):
                generate_snapshot_report(
                    Path(directory), project_id="test", snapshot_id="snap-test",
                    use_latest_approved=False, report_kind="Stage 1 catalog",
                    output_path=output, rows=rows, source_catalog_path=catalog,
                )

            workbook = load_workbook(output)
            self.assertEqual(
                [tuple(value or "" for value in row) for row in workbook.active.values],
                [tuple(rows[0]), *(tuple(row.values()) for row in rows)],
            )
            self.assertEqual(workbook.active["D1"].value, "source_type")
            self.assertEqual(workbook.active["E1"].value, "category")
            self.assertEqual([row["source_type"] for row in rows], ["Primary", "Supplementary"])
            self.assertEqual(
                _classify_integrated_source_rows([{"id": "REQ-1", "source_spec": "Extra.pdf"}], {"REQ-1"}, "Main.pdf")[0]["source_type"],
                "Supplementary",
            )
            self.assertEqual(
                _classify_integrated_source_rows([{"id": "REQ-1", "source_spec": "sources/Main.pdf"}], {"REQ-1"}, "Main.pdf")[0]["source_type"],
                "Primary",
            )
            self.assertIn("ReportData", workbook.active.tables)
            metadata = dict(workbook["Report Metadata"].values)
            self.assertIn("not canonical snapshot authority", metadata["Authority"])
            self.assertEqual(metadata["Catalog source"], str(catalog))
            with self.assertRaises(RuntimeError):
                _classify_integrated_source_rows([{"id": "UNKNOWN", "source_spec": ""}], {"REQ-1"}, "Main.pdf")

    def test_supplementary_source_classification_requires_approved_snapshot_match(self):
        source_rows = [
            {
                "canonical_id": "canonical-primary",
                "source_req_id": "PRIMARY-REQ-1",
                "requirement_statement": "Primary source statement",
            },
            {
                "canonical_id": "canonical-supplementary",
                "source_req_id": "SUPPLEMENTARY-REQ-1",
                "requirement_statement": "Supplementary source statement",
            },
        ]
        mapping_rows = [
            {
                "requirement_id": "SUPPLEMENTARY-REQ-1",
                "canonical_requirement_id": "mapping-canonical-1",
                "review_decision": "approved",
                "supplementary_source_spec": "Supplementary Spec.pdf",
            },
            {
                "requirement_id": "NOT-IN-SNAPSHOT",
                "canonical_requirement_id": "mapping-canonical-2",
                "review_decision": "approved",
                "supplementary_source_spec": "Stale Spec.pdf",
            },
            {
                "requirement_id": "PRIMARY-REQ-1",
                "canonical_requirement_id": "canonical-primary",
                "review_decision": "pending_review",
                "supplementary_source_spec": "Unapproved Spec.pdf",
            },
        ]

        imports, source_specs_by_id = _approved_supplementary_imports(
            source_rows,
            mapping_rows,
            "Primary Spec.pdf",
        )

        self.assertEqual([row["source_req_id"] for row in imports], ["SUPPLEMENTARY-REQ-1"])
        self.assertEqual(source_specs_by_id["SUPPLEMENTARY-REQ-1"], {"Supplementary Spec.pdf"})
        self.assertEqual(source_specs_by_id["canonical-supplementary"], {"Supplementary Spec.pdf"})
        self.assertNotIn("PRIMARY-REQ-1", source_specs_by_id)
        self.assertNotIn("NOT-IN-SNAPSHOT", source_specs_by_id)

    def test_generated_report_statement_comes_only_from_approved_snapshot(self):
        snapshot_rows = [{
            "canonical_id": "canonical-17",
            "source_req_id": "SRC-REQ-17",
            "requirement_statement": "Approved source statement",
        }]
        matrix_row = {
            "canonical_id": "canonical-17",
            "source_req_id": "SRC-REQ-17",
            "requirement_statement": "Stale derived statement",
        }

        statement_index = _snapshot_statement_index(snapshot_rows)

        self.assertEqual(
            _snapshot_statement_for(matrix_row, statement_index),
            "Approved source statement",
        )
        self.assertEqual(_snapshot_statement_for({"source_req_id": "unknown"}, statement_index), "")

    def test_generated_local_id_resolves_matrix_statement_through_approved_source_id(self):
        snapshot_rows = [{
            "canonical_id": "canonical-17",
            "source_req_id": "SRC-REQ-17",
            "requirement_statement": "ADSP",
        }]
        generated_row = {
            "requirement_id": "DOC-REQ-004",
            "source_req_id": "SRC-REQ-17",
            "requirement_statement": "Stale derived statement",
        }
        interaction_rows = [{
            "From block": "Source Block",
            "To block": "Destination Block",
            "Requirement IDs": "SRC-REQ-17",
        }]
        statement_index = _snapshot_statement_index(snapshot_rows)

        statement, statement_source = _generated_requirement_statement(
            generated_row,
            "DRS",
            statement_index,
            {"SRC-REQ-17"},
            interaction_rows,
        )
        unapproved_statement, unapproved_source = _generated_requirement_statement(
            generated_row,
            "DRS",
            statement_index,
            set(),
            interaction_rows,
        )

        self.assertEqual(
            statement,
            "The Source Block block shall implement: connection to the Destination Block block.",
        )
        self.assertEqual(statement_source, "Source-backed interaction matrix")
        self.assertEqual(unapproved_statement, "ADSP")
        self.assertEqual(unapproved_source, "Approved canonical snapshot")

    def test_inclusive_report_uses_architecture_allocation_for_hierarchy_level(self):
        self.assertEqual(
            _hierarchy_level_label("DRS", "top_digital_architecture", "DRS"),
            "top_digital",
        )
        self.assertEqual(
            _hierarchy_level_label("ARS", "top_analog_architecture", "ARS"),
            "top_analog",
        )
        self.assertEqual(_hierarchy_level_label("SRS", "", "SRS"), "System")

    def test_reciprocal_source_coverage_deduplicates_ids_and_uses_exact_references(self):
        source_coverage, target_coverage = _source_spec_target_metrics(
            {"SRC-1", "SRC-2", "SRC-3"},
            [
                {"source_req_id": "SRC-1", "ipos_req_id": "IPOS-1"},
                {"source_origin_req_ids": "SRC-2; SRC-2", "ipos_req_id": "IPOS-2"},
                {"source_req_id": "SRC-1", "ipos_req_id": "IPOS-1"},
                {"source_req_id": "SRC-20", "ipos_req_id": "IPOS-3"},
            ],
        )

        self.assertEqual((source_coverage["covered"], source_coverage["total"]), (2, 3))
        self.assertEqual(source_coverage["covered_unique_source_req_ids"], ["SRC-1", "SRC-2"])
        self.assertEqual((target_coverage["covered"], target_coverage["total"]), (2, 3))
        self.assertEqual(target_coverage["covered_unique_target_req_ids"], ["IPOS-1", "IPOS-2"])

    def test_source_spec_metrics_use_each_sources_own_denominator(self):
        primary = _source_spec_node_metrics(
            {"P-1", "P-2", "P-3"},
            [{
                "target": "DRS",
                "coverage": {
                    "covered_unique_upstream_req_ids": ["P-1", "P-1", "P-3"],
                    "total_unique_upstream_req_ids": ["P-1", "P-2", "P-3"],
                },
            }],
        )
        supplementary = _source_spec_node_metrics(
            {"SUP-1", "SUP-2"},
            [{
                "target": "Digital_IPOS:block-a",
                "coverage": {
                    "covered_unique_upstream_req_ids": ["SUP-1", "SUP-2"],
                    "total_unique_upstream_req_ids": ["SUP-1", "SUP-2"],
                },
            }],
        )

        self.assertEqual((primary["covered"], primary["total"]), (2, 3))
        self.assertAlmostEqual(primary["percentage"], 200 / 3)
        self.assertEqual((supplementary["covered"], supplementary["total"], supplementary["percentage"]), (2, 2, 100.0))
        self.assertEqual(supplementary["downstream_coverage"][0]["target"], "Digital_IPOS:block-a")

    def test_upstream_coverage_is_unique_and_limited_to_covering_specifications(self):
        nodes = {
            "spec-upstream": {"label": "Upstream spec"},
            "primary-source": {"label": "Primary source"},
            "supplementary-source": {"label": "Supplementary source"},
        }
        coverage = {"covered": 123, "total": 123, "percentage": 100.0}
        edges = [
            {"source": "spec-upstream", "target": "spec-downstream", "kind": "generated_document_covers", "coverage": coverage},
            {"source": "spec-upstream", "target": "spec-downstream", "kind": "generated_document_covers", "coverage": coverage},
            {"source": "primary-source", "target": "spec-downstream", "kind": "approved_source_covers", "coverage": coverage},
            {"source": "supplementary-source", "target": "spec-downstream", "kind": "approved_supplementary_covers", "coverage": coverage},
        ]

        upstream_labels, upstream_statistics, upstream_metrics = _approved_upstream_coverage(
            "spec-downstream", nodes, edges
        )

        self.assertEqual(upstream_labels, ["Upstream spec", "Primary source", "Supplementary source"])
        self.assertEqual(
            upstream_statistics,
            [
                "Upstream spec: 123 / 123 (100.0%)",
                "Primary source: 123 / 123 (100.0%)",
                "Supplementary source: 123 / 123 (100.0%)",
            ],
        )
        self.assertEqual(len(upstream_metrics), 3)
        self.assertEqual(upstream_metrics[0]["covered"], 123)
        self.assertEqual(upstream_metrics[0]["percentage"], 100.0)

    def test_explicit_edge_precedes_source_metric_for_same_upstream_document(self):
        nodes = {"source:primary.pdf": {"label": "Primary source specification: primary.pdf"}}
        edges = [
            {
                "source": "source:primary.pdf", "target": "DRS", "kind": "approved_source_covers",
                "coverage": {"covered": 5, "total": 20, "percentage": 25.0},
            },
            {
                "source": "source:primary.pdf", "target": "DRS", "kind": "approved_source_covers",
                "coverage": {"covered": 8, "total": 20, "percentage": 40.0},
            },
        ]

        labels, details, metrics = _approved_upstream_coverage("DRS", nodes, edges)

        self.assertEqual(labels, ["Primary source specification: primary.pdf"])
        self.assertEqual(details, ["Primary source specification: primary.pdf: 5 / 20 (25.0%)"])
        self.assertEqual(len(metrics), 1)

    def test_requirement_id_extraction_ignores_template_examples(self):
        rendered = (
            "- `SRS-REQ-001` (for newly authored requirements)\n"
            "| Block | SRS-REQ-002 | Source |\n"
            "[SRS-REQ-003] Requirement:\n"
        )

        self.assertEqual(
            _extract_srs_requirement_ids(rendered),
            {"SRS-REQ-002", "SRS-REQ-003"},
        )

    def test_empty_approved_partition_is_a_passing_srs_report(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_path = Path(temporary_directory) / "stage_srs_report.md"
            status = _write_stage_report(
                report_path,
                [],
                [],
                [],
                [],
                0,
                0,
                "go",
                valid_empty=True,
            )

            report = report_path.read_text(encoding="utf-8")

        self.assertEqual(status, "pass")
        self.assertIn("SRS generation status: pass", report)
        self.assertNotIn("No requirements were loaded", report)

    def test_mode_rows_require_approved_srs_allocation(self):
        mode_rows = [
            {"source_req_id": "SRC-SYS-001", "mode": "boot"},
            {"source_req_id": "SRC-DIG-002", "mode": "boot"},
            {"source_req_id": "MODE-derived-003", "mode": "boot"},
        ]
        allocations = {
            "SRC-SYS-001": {"owning_target": "SRS"},
            "SRC-DIG-002": {"owning_target": "DRS"},
        }

        filtered = _filter_mode_rows_to_srs_allocations(mode_rows, allocations)

        self.assertEqual([row["source_req_id"] for row in filtered], ["SRC-SYS-001"])

    def test_stale_docx_requirement_is_rejected_for_empty_approved_partition(self):
        contract = EmptySrsContract()
        with tempfile.TemporaryDirectory() as temporary_directory:
            repo_root = Path(temporary_directory)
            stage3_root = repo_root / "artifacts/stage3_srs"
            stage3_root.mkdir(parents=True)
            (stage3_root / "srs_traceability_matrix.csv").write_text(
                "srs_req_id,source_req_id\n", encoding="utf-8"
            )
            (stage3_root / "system_requirements_specification.md").write_text(
                "Snapshot ID: snap-test\n"
                f"Downstream contract fingerprint: {downstream_contract_fingerprint(contract)}\n",
                encoding="utf-8",
            )
            docx_path = stage3_root / "system_requirements_specification.docx"
            with zipfile.ZipFile(docx_path, "w") as archive:
                archive.writestr(
                    "word/document.xml",
                    '<w:document xmlns:w="urn:test"><w:body><w:tbl><w:tr>'
                    "<w:tc><w:p><w:r><w:t>PMU</w:t></w:r></w:p></w:tc>"
                    "<w:tc><w:p><w:r><w:t>SRS-REQ-022</w:t></w:r></w:p></w:tc>"
                    "</w:tr></w:tbl></w:body></w:document>",
                )

            findings = _check_srs_artifact_coherence(repo_root, contract, {})

        self.assertIn("SRS_RENDERED_REQUIREMENT_ID_SET_MISMATCH: missing=0, extra=1", findings)


class OperationalFingerprintTests(unittest.TestCase):
    def setUp(self):
        temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(temporary_directory.cleanup)
        self.repo_root = Path(temporary_directory.name)
        self.contract = EmptySrsContract()
        self.contract.expected = {
            "SOURCE_001": {
                "canonical_id": "canonical-test",
                "source_req_id": "SOURCE_001",
                "allocation_class": "top_digital_architecture",
                "owning_target": "DRS",
                "approved_block": "Digital",
                "owning_domain": "Digital",
                "allocation_rationale": "Approved digital architecture requirement.",
                "lineage_mode": "normal_hierarchical",
                "source_origin_req_ids": "SOURCE_001",
                "hierarchy_parent_req_ids": "",
                "lineage_candidate_parent_req_ids": "",
            },
        }
        self.fingerprint = downstream_contract_fingerprint(self.contract)
        stage3_root = self.repo_root / "artifacts/stage3_srs"
        stage3_root.mkdir(parents=True)
        (stage3_root / "srs_traceability_matrix.csv").write_text(
            "srs_req_id,source_req_id\n", encoding="utf-8",
        )
        (stage3_root / "system_requirements_specification.md").write_text(
            f"Snapshot ID: {self.contract.snapshot_id}\n"
            f"Downstream contract fingerprint: {self.fingerprint}\n",
            encoding="utf-8",
        )
        sysml_root = self.repo_root / "artifacts/stage2_mirco_arc/sysml"
        sysml_root.mkdir(parents=True)
        (sysml_root / "DigitalSubsystem.sysml").write_text(
            f"// Snapshot ID: {self.contract.snapshot_id}\n"
            f"// Downstream contract fingerprint: {self.fingerprint}\n"
            'attribute sourceReqId = "SOURCE_001";\n',
            encoding="utf-8",
        )

    def findings(self):
        with patch("validate_downstream_coherence.created_ipos_block_directories", return_value={}):
            return (
                _check_srs_artifact_coherence(self.repo_root, self.contract, self.contract.expected),
                _check_sysml_requirement_coverage(self.repo_root, self.contract, self.contract.expected),
            )

    def test_candidate_only_drift_preserves_hash_and_srs_sysml_freshness(self):
        self.assertEqual(self.findings(), ([], []))
        metadata = self.contract.expected["SOURCE_001"]
        for candidates in ("PARENT_001", "PARENT_001;PARENT_002"):
            with self.subTest(candidates=candidates):
                metadata["lineage_candidate_parent_req_ids"] = candidates
                self.assertEqual(downstream_contract_fingerprint(self.contract), self.fingerprint)
                self.assertEqual(self.findings(), ([], []))
                self.assertEqual(metadata["lineage_candidate_parent_req_ids"], candidates)
        del metadata["lineage_candidate_parent_req_ids"]
        self.assertEqual(downstream_contract_fingerprint(self.contract), self.fingerprint)
        self.assertEqual(self.findings(), ([], []))

    def test_approved_contract_changes_still_reject_stale_srs_and_sysml(self):
        metadata = self.contract.expected["SOURCE_001"]
        for field, changed_value in {
            "hierarchy_parent_req_ids": "PARENT_001",
            "lineage_mode": "split_lineage",
            "source_origin_req_ids": "SOURCE_002",
            "approved_block": "Other Block",
            "owning_domain": "Analog",
            "allocation_class": "top_analog_architecture",
            "owning_target": "ARS",
            "allocation_rationale": "Changed approved allocation.",
        }.items():
            with self.subTest(field=field):
                original_value = metadata[field]
                metadata[field] = changed_value
                self.assertNotEqual(downstream_contract_fingerprint(self.contract), self.fingerprint)
                srs_findings, sysml_findings = self.findings()
                self.assertIn("SRS_STALE_CONTRACT_FINGERPRINT", srs_findings)
                self.assertIn("SYSML_STALE_CONTRACT_FINGERPRINT", sysml_findings)
                metadata[field] = original_value
        self.contract.expected["SOURCE_002"] = dict(metadata)
        self.assertNotEqual(downstream_contract_fingerprint(self.contract), self.fingerprint)
        srs_findings, sysml_findings = self.findings()
        self.assertIn("SRS_STALE_CONTRACT_FINGERPRINT", srs_findings)
        self.assertIn("SYSML_STALE_CONTRACT_FINGERPRINT", sysml_findings)
        del self.contract.expected["SOURCE_002"]
        self.contract.snapshot_id = "snap-other"
        srs_findings, sysml_findings = self.findings()
        self.assertTrue(any(finding.startswith("SRS_STALE_SNAPSHOT_MARKERS:") for finding in srs_findings))
        self.assertTrue(any(finding.startswith("SYSML_STALE_SNAPSHOT_MARKERS:") for finding in sysml_findings))


if __name__ == "__main__":
    unittest.main()