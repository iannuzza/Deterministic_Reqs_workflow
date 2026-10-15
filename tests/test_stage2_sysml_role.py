import unittest
import sys
import tempfile
import csv
import json
import hashlib
import tkinter as tk
from tkinter import ttk
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch, MagicMock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from validate_downstream_coherence import resolve_sysml_materialization_scope, sysml_materialization_findings, _check_sysml_requirement_coverage, downstream_contract_fingerprint
from generate_architecture_sysml import build_block_model, build_model, build_digital_subsystem_model, build_shared_model, sysml_header
from review_architecture_sysml import _review
from workflow_gui import WorkflowGUI
from approved_snapshot_resolver import ApprovedResolverError


class Stage2SysmlRoleTests(unittest.TestCase):
    def test_relation_contract_never_promotes_unresolved_clock_ports(self):
        from sysml_relations import RelationEndpoint, RelationInventory, documented_link, serialize_relations

        source = RelationEndpoint("Timing Unit", "topDigital.timingUnit")
        target = RelationEndpoint("Compute Unit", "topDigital.computeUnit")
        records, gaps = documented_link("clock", source, target, ("approved:clock-1",), "selected")
        inventory = RelationInventory("selected", tuple(records), tuple(gaps))
        self.assertEqual(len(inventory.view("block_relation")), 1)
        self.assertEqual(inventory.view("block_relation")[0].kind, "clock")
        self.assertEqual(inventory.view("signal_connection"), ())
        self.assertEqual(gaps[0]["status"], "unresolved")
        serialized = "\n".join(serialize_relations(inventory))
        self.assertIn(": BlockRelation connect", serialized)
        self.assertNotIn(": SignalConnection connect", serialized)
        source = RelationEndpoint("Timing Unit", "topDigital.timingUnit", "clockOut", "out")
        target = RelationEndpoint("Compute Unit", "topDigital.computeUnit", "clockIn", "in")
        records, gaps = documented_link("clock", source, target, ("approved:clock-1",), "selected")
        self.assertEqual([row.relation_class for row in records], ["block_relation", "signal_connection"])
        self.assertEqual(gaps, [])

    def test_relation_serialization_checks_real_ports_and_mutated_connectors(self):
        from dataclasses import replace
        from sysml_relations import RelationEndpoint, RelationInventory, documented_link, serialize_relations, validate_serialized_relations

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            records, gaps = documented_link("reset", RelationEndpoint(source.name, source.node, "resetOut", "out"),
                RelationEndpoint(target.name, target.node, "resetIn", "in"), ("approved-reset",), scope.snapshot_id)
            inventory = RelationInventory(scope.snapshot_id, tuple(records), tuple(gaps))
            original = scope.system_path.read_text(encoding="utf-8")
            text = original + "\n" + "\n".join(serialize_relations(inventory))
            scope.system_path.write_text(text, encoding="utf-8")
            report = validate_serialized_relations(scope, inventory)
            self.assertTrue(any("PORT_NOT_MATERIALIZED" in item for item in report["findings"]))
            for block, direction, port in ((source, "out", "resetOut"), (target, "in", "resetIn")):
                block.path.write_text(block.path.read_text(encoding="utf-8") + f"\n{direction} port {port} : ArchitecturePort {{}}", encoding="utf-8")
            self.assertEqual(validate_serialized_relations(scope, inventory)["findings"], [])
            signal = records[1]
            declaration = f"connection {signal.identity} : SignalConnection connect {signal.source.address} to {signal.target.address};"
            for mutated, code in ((text.replace(declaration, ""), "MISSING_OR_DUPLICATE"),
                                  (text.replace(declaration, declaration.replace(signal.source.address, "topDigital.invented.port")), "CHANGED"),
                                  (text + "\nconnection guessed : SignalConnection connect a.port to b.port;", "UNSUPPORTED")):
                scope.system_path.write_text(mutated, encoding="utf-8")
                self.assertTrue(any(code in item for item in validate_serialized_relations(scope, inventory)["findings"]))

    def test_relation_inventory_uses_snapshot_evidence_without_path_guessing(self):
        from sysml_relations import build_relation_inventory

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            identity = next(iter(contract.expected))
            contract.requirement_input = SimpleNamespace(rows=[{"source_req_id": identity,
                "requirement_statement": f"The {source.name}.clockOut signal shall be connected to the {target.name}.clockIn signal, with frequency 12 MHz."}])
            ports = [{"Owner": block.name, "Port name": name, "Direction": direction, "Ownership status": "approved"}
                     for block, name, direction in ((source, "clockOut", "output"), (target, "clockIn", "input"))]
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(len(inventory.view("signal_connection")), 1)
            self.assertEqual(inventory.view("block_relation")[0].kind, "clock")
            self.assertEqual(inventory.view("signal_connection")[0].qualifiers, (("frequency", "12 MHz"),))
            contract.requirement_input.rows[0]["requirement_statement"] = "The top.u_unknown.clock signal shall be connected to the top.u_other.clock signal."
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertFalse(inventory.view("signal_connection"))
            self.assertTrue(inventory.gaps)

    def test_relation_bindings_retain_block_clock_when_port_is_unknown(self):
        from sysml_relations import build_relation_inventory

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            identity = next(iter(contract.expected))
            contract.requirement_input = SimpleNamespace(rows=[{"source_req_id": identity,
                "requirement_statement": "The chip.u_clock.clock_out signal shall be connected to the chip.u_compute.clock_in signal."}])
            bindings = ({"endpoint": "chip.u_clock", "block": source.name}, {"endpoint": "chip.u_compute", "block": target.name})
            inventory = build_relation_inventory(scope, contract, [], [], endpoint_bindings=bindings)
            self.assertEqual(inventory.view("block_relation")[0].status, "resolved")
            self.assertEqual(inventory.view("block_relation")[0].kind, "clock")
            self.assertFalse(inventory.view("signal_connection"))
            ports = [{"Owner": block.name, "Port name": name, "Direction": direction, "Ownership status": "approved"}
                     for block, name, direction in ((source, "clock_out", "output"), (target, "clock_in", "input"))]
            inventory = build_relation_inventory(scope, contract, ports, [], endpoint_bindings=bindings)
            self.assertEqual(len(inventory.view("signal_connection")), 1)
            inventory = build_relation_inventory(scope, contract, ports, [], endpoint_bindings=(*bindings, bindings[0]))
            self.assertFalse(inventory.view("signal_connection"))
            self.assertTrue(inventory.gaps)

    def test_relation_hierarchy_projection_contains_only_materialized_parts(self):
        from sysml_relations import containment_projection

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, scope, _ = self.materialized_fixture(root)
            block = next(block for block in scope.blocks if block.target == "Analog IPOS")
            edges = containment_projection(scope)
            self.assertIn(("topAnalog.analogSubsystem", block.node), edges)
            self.assertNotIn(("topDigital", block.node), edges)
            text = scope.system_path.read_text(encoding="utf-8")
            scope.system_path.write_text(text.replace(f"part {block.node.rsplit('.', 1)[-1]} : {block.definition};", ""), encoding="utf-8")
            self.assertNotIn(("topAnalog.analogSubsystem", block.node), containment_projection(scope))

    def test_relation_matrix_cannot_invent_type_or_promote_block_evidence(self):
        from sysml_relations import build_relation_inventory

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            identity = next(iter(contract.expected))
            contract.requirement_input = SimpleNamespace(rows=[{"source_req_id": identity,
                "requirement_statement": f"The {source.name} shall provide clock signals to the {target.name}."}])
            matrix = [{"From block": source.name, "To block": target.name, "Requirement IDs": identity, "Relation type": "reset"}]
            inventory = build_relation_inventory(scope, contract, [], matrix)
            self.assertEqual(inventory.view("block_relation")[0].kind, "clock")
            self.assertFalse(inventory.view("signal_connection"))
            self.assertTrue(any(gap["status"] == "conflicting" for gap in inventory.gaps))
            matrix[0]["Requirement IDs"] = "unapproved"
            inventory = build_relation_inventory(scope, contract, [], matrix)
            self.assertTrue(any("selected-snapshot" in gap["reason"] for gap in inventory.gaps))

    def test_relation_governed_matrix_cells_reach_gui_without_signal_inference(self):
        from dataclasses import replace
        from sysml_relations import build_relation_inventory, relation_projection, serialize_relations

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            identity = next(iter(contract.expected))
            contract.requirement_input = SimpleNamespace(rows=[{"source_req_id": identity,
                "requirement_statement": "Interconnect connection matrix destination source"}])
            cell = {"From block": source.name, "To block": target.name, "Signal/control": "crossbar connection",
                    "Requirement IDs": identity, "Trigger": "matrix cell"}
            inventory = build_relation_inventory(scope, contract, [], [cell])
            self.assertEqual(inventory.view("block_relation"), ())
            inventory = build_relation_inventory(scope, contract, [], [cell], governed_interaction_rows=[cell])
            scope = replace(scope, relation_inventory=inventory)
            bus = inventory.view("block_relation")[0]
            self.assertEqual((bus.kind, bus.source.node, bus.target.node), ("bus", source.node, target.node))
            self.assertIn(identity, bus.evidence)
            self.assertEqual(inventory.view("signal_connection"), ())
            projection = relation_projection(scope, "block_relation", selected_node=source.node, kind="bus")
            self.assertEqual(projection["serialized"], 0)
            self.assertEqual(projection["gaps"], 1)
            self.assertEqual(projection["rows"][0]["status"], "missing")
            scope.system_path.write_text(scope.system_path.read_text(encoding="utf-8") + "\n" + "\n".join(serialize_relations(inventory)), encoding="utf-8")
            projection = relation_projection(scope, "block_relation", kind="bus")
            self.assertEqual(projection["serialized"], 1)
            self.assertEqual(projection["rows"][0]["status"], "serialized")
            window = tk.Tk()
            try:
                gui = object.__new__(WorkflowGUI)
                gui.sysml_hierarchy_selected_path = source.path
                gui._get_sysml_materialization_scope = lambda: scope
                notebook = ttk.Notebook(window)
                notebook.pack(fill=tk.BOTH, expand=True)
                gui._build_sysml_relation_tabs(notebook)
                for size in ("640x700", "1100x700"):
                    window.geometry(size)
                    tree, fields, kind, focus, status = gui.sysml_relation_tables["block_relation"]
                    kind.set("bus")
                    focus.set(True)
                    gui._refresh_sysml_relation_tabs()
                    window.update()
                    self.assertEqual(len(tree.get_children()), 1)
                    self.assertIn(identity, tree.item(tree.get_children()[0], "values")[fields.index("evidence")])
                    self.assertIn("1 serialized", status.get())
                    signal_tree = gui.sysml_relation_tables["signal_connection"][0]
                    self.assertTrue(all(signal_tree.item(item, "values")[3] != "serialized" for item in signal_tree.get_children()))
            finally:
                window.destroy()
            forged = {**cell, "To block": source.name}
            inventory = build_relation_inventory(scope, contract, [], [forged], governed_interaction_rows=[cell])
            self.assertEqual(inventory.view("block_relation")[0].target.node, target.node)
            self.assertTrue(any("not independently supported" in gap["reason"] for gap in inventory.gaps))
            boundary_id = "approved-boundary-cell"
            contract.expected[boundary_id] = {"approved_block": "Digital"}
            contract.requirement_input.rows.append({"source_req_id": boundary_id, "requirement_statement": "Interconnect Connection Matrix"})
            external = {**cell, "To block": "External Compute", "Requirement IDs": boundary_id}
            inventory = build_relation_inventory(scope, contract, [], [cell, external], governed_interaction_rows=[cell, external])
            scope = replace(scope, relation_inventory=inventory)
            projection = relation_projection(scope, "block_relation", kind="bus")
            boundary = next(row for row in projection["rows"] if boundary_id in row["evidence"])
            self.assertEqual(boundary["status"], "outside_scope")
            self.assertEqual(boundary["target"], "External Compute")
            self.assertIn("crossbar connection", boundary["qualifiers"])
            self.assertFalse(any(gap["status"] == "non_materializable" for gap in inventory.gaps))
            self.assertEqual(inventory.view("signal_connection"), ())
            self.assertNotIn(("protocol", "AHB"), inventory.view("block_relation")[0].qualifiers)
            unknown = {**cell, "Requirement IDs": "unapproved-cell"}
            inventory = build_relation_inventory(scope, contract, [], [], governed_interaction_rows=[unknown])
            self.assertEqual(inventory.view("block_relation"), ())
            self.assertTrue(inventory.gaps)

    def test_relation_endpoint_diagnostics_and_normalized_logical_identity(self):
        from dataclasses import replace
        from sysml_relations import build_relation_inventory, relation_projection, RelationEndpoint, documented_link

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            ports = [{"Owner": source.name, "Port name": "clockOut", "Direction": "output", "Ownership status": "approved"},
                     {"Owner": target.name, "Port name": "HCLK", "Direction": "input", "Ownership status": "approved",
                      "Type / details": source.name.swapcase().replace(" ", "_") + ".clockOut"}]
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(len(inventory.view("signal_connection")), 1)
            ports[1]["Type / details"] = "chip.unmapped.clockOut"
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(inventory.view("signal_connection"), ())
            projection = relation_projection(replace(scope, relation_inventory=inventory), "signal_connection")
            self.assertEqual(projection["rows"][0]["status"], "unmapped_instance")
            self.assertIn("chip.unmapped.clockOut", projection["rows"][0]["reason"])
            for details, expected in ((source.name.replace(" ", "_") + ".missingPort", "port_not_found"),
                                      ("chip.unmapped.", "incomplete_reference"), ("chip.unmapped. clockOut", "unsupported_reference")):
                ports[1]["Type / details"] = details
                inventory = build_relation_inventory(scope, contract, ports, [])
                scoped = replace(scope, relation_inventory=inventory)
                projection = relation_projection(scoped, "signal_connection", kind=expected)
                self.assertTrue(projection["rows"], expected)
                self.assertTrue(all(row["status"] == expected for row in projection["rows"]))
                self.assertEqual(inventory.view("signal_connection"), ())
            ports[1]["Type / details"] = source.node + ".clockOut"
            ports[0]["Direction"] = "input"
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(inventory.signal_checks[1]["connection_status"], "unresolved")
            self.assertEqual(inventory.signal_checks[1]["resolution_status"], "direction_conflict")
            _, gaps = documented_link("bus", RelationEndpoint(source.name, source.node, token=source.name),
                                      RelationEndpoint(target.name, target.node, token=target.name), ("source-evidence",), scope.snapshot_id)
            self.assertEqual(gaps[0]["resolution_status"], "block_only")
            self.assertEqual(gaps[0]["status"], "unresolved")

    def test_relation_instance_names_resolve_hierarchical_connections_without_substrings(self):
        from dataclasses import replace
        from sysml_relations import build_relation_inventory
        from validate_downstream_coherence import connection_instance_block_keys

        with tempfile.TemporaryDirectory() as temporary:
            contract, scope, _ = self.materialized_fixture(Path(temporary))
            source, target = scope.blocks[:2]
            for prefix in ("", "chip.digital."):
                details = prefix + "u_" + source.name.swapcase().replace(" ", "_") + "_0.clockOut"
                ports = [{"Owner": source.name, "Port name": "clockOut", "Direction": "output", "Ownership status": "approved"},
                         {"Owner": target.name, "Port name": "HCLK", "Direction": "input", "Ownership status": "approved",
                          "Type / details": details}]
                inventory = build_relation_inventory(scope, contract, ports, [])
                self.assertEqual(len(inventory.view("signal_connection")), 1, details)
                self.assertEqual(inventory.signal_checks[1]["resolution_status"], "resolved")
                self.assertEqual(inventory.view("signal_connection")[0].source.token, details)
            for block in scope.blocks:
                instance = "chip.digital.u_" + block.name.swapcase().replace(" ", "_").replace("-", "_")
                self.assertEqual(connection_instance_block_keys(instance, [item.name for item in scope.blocks]), (block.key,))
            identity = next(iter(contract.expected))
            contract.requirement_input = SimpleNamespace(rows=[{"source_req_id": identity,
                "requirement_statement": "The " + details + " signal shall be connected to the chip.digital.u_"
                + target.name.swapcase().replace(" ", "_") + ".HCLK signal."}])
            ports[1]["Type / details"] = ""
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(len(inventory.view("signal_connection")), 1)
            self.assertEqual(inventory.view("signal_connection")[0].target.node, target.node)
            contract.requirement_input = SimpleNamespace(rows=[])
            ports[1]["Type / details"] = details
            self.assertEqual(connection_instance_block_keys("chip.u_pMu", ["PMU"]), ("pmu",))
            self.assertEqual(connection_instance_block_keys("chip.u_fifo", ["Smart FIFO"]), ("smart fifo",))
            self.assertEqual(connection_instance_block_keys("chip.u_notpmu", ["PMU"]), ())
            self.assertEqual(connection_instance_block_keys("chip.u_pmu.unknown", ["PMU"]), ())
            ports[1]["Type / details"] = details.replace("clockOut", "CLOCKOUT")
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(inventory.signal_checks[1]["resolution_status"], "port_not_found")
            ports[1]["Type / details"] = details
            ports[0]["Direction"] = "input"
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(inventory.signal_checks[1]["resolution_status"], "direction_conflict")
            ports[0]["Direction"] = "output"
            collision = replace(source, name="Aux " + source.name, key="aux " + source.key, node=source.node + "Aux")
            ambiguous_scope = replace(scope, blocks=(*scope.blocks, collision))
            ports[1]["Type / details"] = "chip.u_aux_" + source.name.replace(" ", "_") + ".clockOut"
            inventory = build_relation_inventory(ambiguous_scope, contract, ports, [])
            self.assertEqual(inventory.signal_checks[1]["resolution_status"], "identity_conflict")
            self.assertFalse(inventory.view("signal_connection"))

    def test_relation_table_headers_sort_naturally_and_preserve_selection(self):
        window = tk.Tk()
        self.addCleanup(window.destroy)
        gui = object.__new__(WorkflowGUI)
        notebook = ttk.Notebook(window)
        gui._build_sysml_relation_tabs(notebook)
        for tree, fields, _kind, _focus, _status in gui.sysml_relation_tables.values():
            items = [tree.insert("", tk.END, values=tuple(value if field == "source" else "" for field in fields))
                     for value in ("port10", "port2", "port1", "")]
            tree.selection_set(items[0])
            tree.focus(items[0])
            for field in fields:
                self.assertTrue(tree.heading(field, "command"))
            window.tk.call(tree.heading("source", "command"))
            self.assertEqual([tree.set(item, "source") for item in tree.get_children()], ["port1", "port2", "port10", ""])
            window.tk.call(tree.heading("source", "command"))
            self.assertEqual([tree.set(item, "source") for item in tree.get_children()], ["port10", "port2", "port1", ""])
            self.assertEqual(tree.selection(), (items[0],))
            self.assertEqual(tree.focus(), items[0])
            gui._sort_sysml_table(tree)
            self.assertIn("[desc]", tree.heading("source", "text"))
            tree.configure(columns=("source_page",))
            gui._configure_sysml_table_headings(tree, ("source_page",), details=True)
            for item, value in zip(items, ("100", "20", "3", "")):
                tree.set(item, "source_page", value)
            window.tk.call(tree.heading("source_page", "command"))
            self.assertEqual([tree.set(item, "source_page") for item in tree.get_children()], ["3", "20", "100", ""])

    def test_relation_io_block_descriptions_preserve_peers_without_inventing_ports(self):
        from dataclasses import replace
        from sysml_relations import build_relation_inventory, signal_description_projection, relation_projection, serialize_relations

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            ports = []
            descriptions = ("wire " + target.name + " clock", "wire clock for " + target.name,
                            "wire clock slow 32k to " + target.name.replace(" ", "_"),
                            "wire " + target.name + " master clock", "wire reset " + target.name,
                            "wire 16MHz clock for " + target.name, "wire AHB bus to " + target.name,
                            "wire data to " + target.name, "wire interaction with " + target.name)
            for index, details in enumerate(descriptions):
                ports.append({"Owner": source.name, "Port name": "out_" + str(index), "Direction": "output",
                              "Type / details": details, "Ownership status": "approved", "Source page": "40", "Table": "I/O List"})
            ports.append({**ports[0], "Port name": "inputClock", "Direction": "input", "Type / details": "wire clock from " + target.name})
            inventory = build_relation_inventory(scope, contract, ports, [])
            scoped = replace(scope, relation_inventory=inventory)
            self.assertEqual(len(inventory.view("block_relation")), len(ports))
            self.assertTrue({"clock", "reset", "bus", "data", "interaction"} <= {row.kind for row in inventory.view("block_relation")})
            self.assertEqual(inventory.view("signal_connection"), ())
            details = signal_description_projection(scoped, selected_node=source.node)["rows"]
            self.assertTrue(all(row["target"] == target.name for row in details))
            self.assertTrue(all(row["status"] == "block_only" and row["peer_port_status"] == "not_established" for row in details))
            incoming = next(row for row in inventory.view("block_relation") if dict(row.qualifiers)["local_port"].endswith(".inputClock"))
            self.assertEqual((incoming.source.node, incoming.target.node), (target.node, source.node))
            self.assertTrue(all(row["status"] == "missing" for row in relation_projection(scoped, "block_relation")["rows"]))
            self.assertNotIn(": SignalConnection connect", "\n".join(serialize_relations(inventory)))
            association = next(row for row in inventory.view("block_relation") if row.kind == "interaction")
            self.assertEqual(dict(association.qualifiers)["directionality"], "unspecified")
            with patch.dict("validate_downstream_coherence.APPROVED_LABEL_CANONICALS", {"consumer alias": target.name}):
                candidate = build_relation_inventory(scope, contract, [{**ports[0], "Type / details": "wire consumer alias clock"}], [])
                self.assertEqual(candidate.view("block_relation")[0].target.node, target.node)
                self.assertEqual(candidate.gaps[0]["resolution_status"], "block_only")
            containment = inventory.view("containment")
            for rejected in ("wire clock for " + target.name + " or " + source.name,
                             "wire clock for " + target.name + " if enabled", "wire not " + target.name + " clock"):
                candidate = build_relation_inventory(scope, contract, [{**ports[0], "Type / details": rejected}], [])
                self.assertEqual(candidate.view("block_relation"), ())
                self.assertEqual(candidate.view("signal_connection"), ())
                self.assertEqual(candidate.view("containment"), containment)

    def test_relation_io_table_preserves_clock_reset_bus_and_unresolved_ports(self):
        from dataclasses import replace
        from sysml_relations import build_relation_inventory, relation_projection

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]

            def port(block, name, direction, details=""):
                return {"Owner": block.name, "Port name": name, "Direction": direction, "Type / details": details,
                        "Ownership status": "approved", "Table": "Controller I/O List", "Source page": "12"}

            ports = [port(source, "clockOut", "output"), port(source, "resetOut", "output"),
                     port(source, "addressOut", "output"), port(source, "dataIn", "input"),
                     port(target, "HCLK", "input", "AHB Handler Clock 16MHz clock " + source.node + ".clockOut"),
                     port(target, "HRESETn", "input", source.node + ".resetOut"),
                     port(target, "HADDR", "input", "[AW-1:0] AHB Slave Port " + source.node + ".addressOut"),
                     port(target, "HRDATA", "output", "[31:0] AHB Slave Port " + source.node + ".dataIn"),
                     port(target, "HWRITE", "input", "AHB Slave Port " + source.node + "."),
                     port(target, "asyncClock", "input", "Clock reference chip.u_unknown.clockOut")]
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(len(inventory.view("signal_connection")), 4)
            self.assertEqual({row.kind for row in inventory.view("signal_connection")}, {"clock", "reset", "bus"})
            data = next(row for row in inventory.view("signal_connection") if row.source.port == "HRDATA")
            self.assertEqual(data.target.port, "dataIn")
            scope = replace(scope, relation_inventory=inventory)
            projection = relation_projection(scope, "signal_connection", selected_node=target.node, kind="clock")
            self.assertTrue(any(row["target"].endswith(".HCLK") for row in projection["rows"]))
            self.assertTrue(any(row["source"] == "chip.u_unknown.clockOut" for row in projection["rows"]))
            self.assertTrue(any("incomplete" in gap["reason"] for gap in inventory.gaps))
            self.assertFalse(any(row.target.port == "HWRITE" for row in inventory.view("signal_connection")))
            ambiguous = port(target, "ambiguousClock", "input", "Clock from a.clock or b.clock")
            inventory = build_relation_inventory(scope, contract, [ambiguous], [])
            self.assertEqual(inventory.view("signal_connection"), ())
            self.assertTrue(any("ambiguous" in gap["reason"] for gap in inventory.gaps))

    def test_relation_signal_checks_cover_every_scoped_block_including_analog(self):
        from dataclasses import replace
        from sysml_relations import build_relation_inventory, signal_description_projection, validate_serialized_relations

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            contract, scope, _ = self.materialized_fixture(root)
            ports = []
            for block in scope.blocks:
                for name, details in (("HCLK", "Clock 16MHz chip.unmapped.clk_out"), ("HRESETn", "chip.unmapped.reset_out"),
                                      ("HADDR", "[AW-1:0] AHB Slave Port chip.fabric.haddr"), ("status", "[3:0]")):
                    if name == "status" and block.target == "Analog IPOS":
                        details = "16 bits"
                    ports.append({"Owner": block.name, "Port name": name, "Direction": "" if name == "status" else "input", "Type / details": details,
                                  "Ownership status": "approved", "Table": "Interface List", "Source page": "7"})
            ports.append({**ports[0], "Owner": "Unmaterialized Block"})
            inventory = build_relation_inventory(scope, contract, ports, [])
            self.assertEqual(len(inventory.signal_checks), 4 * len(scope.blocks))
            scope = replace(scope, relation_inventory=inventory)
            self.assertTrue(any(block.target == "Analog IPOS" for block in scope.blocks))
            for block in scope.blocks:
                projection = signal_description_projection(scope, selected_node=block.node)
                self.assertEqual(len(projection["rows"]), 4)
                self.assertEqual(len(signal_description_projection(scope, selected_node=block.node, kind="clock")["rows"]), 1)
                self.assertEqual(next(row for row in projection["rows"] if row["port"] == "status")["description_status"], "not_established")
                self.assertEqual(next(row for row in projection["rows"] if row["port"] == "status")["direction_status"], "invalid_or_missing")
                self.assertEqual(next(row for row in projection["rows"] if row["port"] == "HADDR")["width"], "[AW-1:0]")
                self.assertTrue(all(row["node"] == block.node for row in projection["rows"]))
            self.assertEqual(inventory.view("signal_connection"), ())
            report = validate_serialized_relations(scope, inventory)
            self.assertTrue(all(row["source_signals"] == 4 for row in report["blocks"]))
            self.assertTrue(all(row["descriptions_not_established"] == 1 for row in report["blocks"]))
            window = tk.Tk()
            try:
                gui = object.__new__(WorkflowGUI)
                gui._get_sysml_materialization_scope = lambda: scope
                notebook = ttk.Notebook(window)
                notebook.pack(fill=tk.BOTH, expand=True)
                gui._build_sysml_relation_tabs(notebook)
                gui.sysml_signal_view.set("details")
                for size in ("640x700", "1100x700"):
                    window.geometry(size)
                    for block in scope.blocks:
                        gui.sysml_hierarchy_selected_path = block.path
                        gui.sysml_relation_tables["signal_connection"][3].set(True)
                        gui._refresh_sysml_relation_tabs()
                        window.update()
                        tree, fields, kind, focus, status = gui.sysml_relation_tables["signal_connection"]
                        self.assertEqual(len(tree.get_children()), 4)
                        self.assertEqual(tree.heading("status", "text"), "Endpoint Status")
                        self.assertEqual(tree.heading("model_status", "text"), "Model Status")
                        self.assertIn("4 source signals", status.get())
                        self.assertTrue(all(tree.item(item, "values")[fields.index("source")].startswith(block.node + ".")
                                            for item in tree.get_children()))
                        kind.set("unmapped_instance")
                        gui._refresh_sysml_relation_tabs()
                        self.assertEqual(len(tree.get_children()), 3)
                        self.assertTrue(all(tree.item(item, "values")[fields.index("status")] == "unmapped_instance"
                                            for item in tree.get_children()))
                        kind.set("all")
                gui.sysml_hierarchy_selected_path = scope.digital_path
                gui._refresh_sysml_relation_tabs()
                self.assertEqual(len(tree.get_children()), 4 * sum(block.target == "Digital IPOS" for block in scope.blocks))
                gui.sysml_signal_view.set("connections")
                gui._refresh_sysml_relation_tabs()
                self.assertNotIn("description", gui.sysml_relation_tables["signal_connection"][1])
                tree = gui.sysml_relation_tables["signal_connection"][0]
                self.assertEqual(tree.heading("resolution_status", "text"), "Endpoint Status")
                self.assertEqual(tree.heading("model_status", "text"), "Model Status")
                self.assertNotIn("status", tree.cget("displaycolumns"))
            finally:
                window.destroy()

    def test_relation_legacy_crosscheck_delegates_to_selected_central_scope(self):
        import crosscheck_sysml_connections as command

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, scope, _ = self.materialized_fixture(root)
            output = root / "relation-crosscheck.md"
            with patch.object(sys, "argv", ["crosscheck_sysml_connections.py", "--snapshot-id", scope.snapshot_id, "--output", str(output)]), \
                    patch.object(command, "resolve_sysml_materialization_scope", return_value=scope) as resolve:
                self.assertEqual(command.main(), 0)
                self.assertEqual(resolve.call_args.args[1], scope.snapshot_id)
                self.assertIn("not a complete-connectivity claim", output.read_text(encoding="utf-8"))

    def test_relation_three_tabs_show_actual_and_missing_separately(self):
        from dataclasses import replace
        from sysml_relations import RelationEndpoint, RelationInventory, documented_link, relation_projection

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, scope, _ = self.materialized_fixture(root)
            source, target = scope.blocks[:2]
            records, gaps = documented_link("clock", RelationEndpoint(source.name, source.node), RelationEndpoint(target.name, target.node), ("clock-evidence",), scope.snapshot_id)
            scope = replace(scope, relation_inventory=RelationInventory(scope.snapshot_id, tuple(records), tuple(gaps)))
            self.assertEqual(relation_projection(scope, "block_relation")["rows"][0]["status"], "missing")
            self.assertEqual(relation_projection(scope, "signal_connection")["serialized"], 0)
            window = tk.Tk()
            try:
                gui = object.__new__(WorkflowGUI)
                gui.sysml_hierarchy_selected_path = source.path
                gui._get_sysml_materialization_scope = lambda: scope
                notebook = ttk.Notebook(window)
                notebook.pack(fill=tk.BOTH, expand=True)
                notebook.add(ttk.Frame(notebook), text="Hierarchy")
                gui._build_sysml_relation_tabs(notebook)
                for size in ("640x700", "1100x700"):
                    window.geometry(size)
                    gui._refresh_sysml_relation_tabs()
                    window.update()
                    self.assertEqual([notebook.tab(tab, "text") for tab in notebook.tabs()], ["Hierarchy", "Block Relations", "Signal Connections"])
                    tree, fields, kind, focus, status = gui.sysml_relation_tables["block_relation"]
                    self.assertEqual(len(tree.get_children()), 1)
                    self.assertIn("1 gaps", status.get())
                    kind.set("reset")
                    gui._refresh_sysml_relation_tabs()
                    self.assertEqual(len(tree.get_children()), 0)
                    self.assertIn("1 hidden", status.get())
                    kind.set("all")
            finally:
                window.destroy()

    def test_relation_clock_reset_details_share_endpoint_and_model_status_for_all_blocks(self):
        from dataclasses import replace
        from sysml_relations import build_relation_inventory, relation_projection, signal_description_projection, serialize_relations

        with tempfile.TemporaryDirectory() as temporary:
            _, scope, _ = self.materialized_fixture(Path(temporary))
            original_models = {block.path: block.path.read_text(encoding="utf-8") for block in scope.blocks}
            original_system = scope.system_path.read_text(encoding="utf-8")
            for source in scope.blocks:
                target = next(block for block in scope.blocks if block != source)
                for kind, output, input_name in (("clock", "clockOut", "HCLK"), ("reset", "resetOut", "HRESETn")):
                    ports = [{"Owner": source.name, "Port name": output, "Direction": "output", "Ownership status": "approved"},
                             {"Owner": target.name, "Port name": input_name, "Direction": "input", "Ownership status": "approved",
                              "Type / details": source.node + "." + output}]
                    inventory = build_relation_inventory(scope, SimpleNamespace(expected={}), ports, [])
                    scoped = replace(scope, relation_inventory=inventory)
                    connection = next(row for row in relation_projection(scoped, "signal_connection")["rows"] if row["status"] == "missing")
                    detail = next(row for row in signal_description_projection(scoped)["rows"] if row["port"] == input_name)
                    self.assertEqual(connection["kind"], kind)
                    self.assertEqual((detail["resolution_status"], detail["model_status"]), ("resolved", "missing"))
                    self.assertEqual((connection["resolution_status"], connection["model_status"]), ("resolved", "missing"))
                    self.assertEqual(detail["relation_ids"], [connection["identity"]])
                    self.assertEqual(len(signal_description_projection(scoped, kind="missing")["rows"]), 1)
                    for block, direction, port in ((source, "out", output), (target, "in", input_name)):
                        block.path.write_text(original_models[block.path] + f"\n{direction} port {port} : ArchitecturePort {{}}", encoding="utf-8")
                    scope.system_path.write_text(original_system + "\n" + "\n".join(serialize_relations(inventory)), encoding="utf-8")
                    detail = next(row for row in signal_description_projection(scoped)["rows"] if row["port"] == input_name)
                    self.assertEqual((detail["resolution_status"], detail["model_status"]), ("resolved", "serialized"))
                    self.assertEqual(len(signal_description_projection(scoped, kind="serialized")["rows"]), 1)
                    source.path.write_text(original_models[source.path], encoding="utf-8")
                    detail = next(row for row in signal_description_projection(scoped)["rows"] if row["port"] == input_name)
                    self.assertEqual((detail["resolution_status"], detail["model_status"]), ("resolved", "invalid"))
                    scope.system_path.write_text(original_system, encoding="utf-8")
                    for path, text in original_models.items():
                        path.write_text(text, encoding="utf-8")

    def test_relation_reports_are_separate_by_view_and_non_authoritative(self):
        from report_generation_service import write_relation_reports
        from sysml_relations import validate_serialized_relations
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, scope, _ = self.materialized_fixture(root)
            report = validate_serialized_relations(scope, scope.relation_inventory)
            write_relation_reports(root / "reports", report, scope.relation_inventory.payload())
            payload = json.loads((root / "reports/relation_summary.json").read_text(encoding="utf-8"))
            self.assertFalse(payload["authority"])
            self.assertEqual(set(payload["views"]), {"containment", "block_relation", "signal_connection"})
            workbook = load_workbook(root / "reports/relation_summary.xlsx")
            try:
                self.assertIn("Gaps", workbook.sheetnames)
                self.assertIn("Signal Checks", workbook.sheetnames)
                self.assertIn("Signal Status Crosscheck", workbook.sheetnames)
                self.assertEqual(workbook["View Coverage"].max_row, 4)
            finally:
                workbook.close()

    def test_readiness_excel_receipt_binds_snapshot_and_summary_bytes(self):
        import run_final_sysml_phase as phase
        import io
        from contextlib import redirect_stdout

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            folder = root / "artifacts/validation/readiness_checklists"
            folder.mkdir(parents=True)
            summary = folder / "readiness_summary.json"
            summary.write_text(json.dumps({"binding": {"snapshot_id": "selected"}}), encoding="utf-8")
            workbook = folder / "readiness_summary.xlsx"
            workbook.write_bytes(b"Excel receipt fixture")
            assessment = {"binding": {"snapshot_id": "selected"}, "status": "NOT_READY", "publication_status": "BLOCKED", "models": {}}
            output = io.StringIO()
            with redirect_stdout(output):
                phase._log_readiness_reports(root, assessment)
            line = next(line for line in output.getvalue().splitlines() if line.startswith("[SYSML_READINESS] "))
            receipt = json.loads(line[len("[SYSML_READINESS] "):])
            self.assertEqual(receipt["snapshot_id"], "selected")
            self.assertEqual(receipt["summary_json_sha256"], hashlib.sha256(summary.read_bytes()).hexdigest())
            self.assertEqual(receipt["summary_excel_sha256"], hashlib.sha256(workbook.read_bytes()).hexdigest())

    def test_final_sysml_readiness_runs_immediately_after_generation_and_logs_excel(self):
        import run_final_sysml_phase as phase
        import readiness_checklist as policy
        import validate_downstream_coherence as central
        import io
        from contextlib import redirect_stdout

        events = []
        assessment = {"binding": {"snapshot_id": "selected"}, "status": "NOT_READY", "publication_status": "BLOCKED", "models": {}}
        with patch.object(sys, "argv", ["run_final_sysml_phase.py", "--snapshot-id", "selected"]), \
                patch.object(phase, "resolve_complete_snapshot_id", return_value="selected"), \
            patch.object(policy, "sync_checklist_from_excel", side_effect=lambda: events.append("sync") or True), \
                patch.object(phase, "_run", side_effect=lambda root, label, command: events.append(label) or (1 if label.endswith("structural review") else 0)), \
                patch.object(central, "assess_final_readiness", side_effect=lambda *args: events.append("checklist") or assessment), \
                patch.object(central, "write_readiness_reports", side_effect=lambda *args: events.append("reports")), \
                patch.object(central, "stamp_final_readiness") as stamp:
            output = io.StringIO()
            with redirect_stdout(output):
                self.assertEqual(phase.main(), 1)
            self.assertEqual(events[:5], ["sync", "Final SysML generation", "checklist", "reports", "Final SysML structural review"])
            self.assertIn("[REPORT] Open SysML readiness Excel: ", output.getvalue())
            self.assertIn("readiness_summary.xlsx", output.getvalue())
            stamp.assert_not_called()

    def test_final_sysml_refreshes_readiness_before_coherence_without_waiving_failures(self):
        import run_final_sysml_phase as phase
        import readiness_checklist as policy
        import validate_downstream_coherence as central
        import io
        from contextlib import redirect_stdout

        for coherence_exit in (0, 1):
            with self.subTest(coherence_exit=coherence_exit):
                events = []
                assessment = {"binding": {"snapshot_id": "selected"}, "status": "NOT_READY", "publication_status": "BLOCKED", "models": {}}

                def run(root, label, command):
                    events.append(label)
                    self.assertIn("selected", command)
                    return coherence_exit if label == "Final SysML central coherence validation" else 0

                with patch.object(sys, "argv", ["run_final_sysml_phase.py", "--snapshot-id", "selected"]), \
                        patch.object(phase, "resolve_complete_snapshot_id", return_value="selected"), \
                        patch.object(policy, "sync_checklist_from_excel", return_value=False), \
                        patch.object(phase, "_run", side_effect=run), \
                        patch.object(phase, "_log_readiness_reports"), \
                        patch.object(central, "assess_final_readiness", side_effect=lambda *args: events.append("assess") or assessment), \
                        patch.object(central, "write_readiness_reports"), \
                        patch.object(central, "stamp_final_readiness", side_effect=lambda *args: events.append("stamp")) as stamp:
                    output = io.StringIO()
                    with redirect_stdout(output):
                        self.assertEqual(phase.main(), coherence_exit)
                    stamp.assert_called_once()
                    self.assertLess(events.index("Final SysML structural review"), events.index("stamp"))
                    self.assertLess(events.index("stamp"), events.index("Final SysML central coherence validation"))
                    self.assertEqual("Final SysML readiness declaration validation" in events, coherence_exit == 0)
                    self.assertEqual("Final SysML phase: PASS" in output.getvalue(), coherence_exit == 0)
                    if coherence_exit == 0:
                        self.assertIn("implementation readiness: NOT_READY", output.getvalue())

    def test_readiness_rule_workbook_matches_json_without_changing_policy(self):
        from readiness_checklist import CHECKLIST_PATH, checklist_definition, checklist_fingerprint
        from report_generation_service import write_readiness_rule_workbook
        from openpyxl import load_workbook

        original = CHECKLIST_PATH.read_bytes()
        definition = checklist_definition()
        expected = [(category, identity, text) for category, items in definition["categories"].items()
                    for identity, text in items.items()]
        with tempfile.TemporaryDirectory() as directory:
            output = write_readiness_rule_workbook(Path(directory) / "readiness_checklist.xlsx")
            workbook = load_workbook(output)
            try:
                sheet = workbook["Checklist Rules"]
                self.assertEqual(list(sheet.iter_rows(min_row=2, values_only=True)), expected)
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.tables["ReadinessRules"].ref, "A1:C" + str(len(expected) + 1))
                self.assertTrue(all(cell.alignment.wrap_text for row in sheet.iter_rows(min_row=2) for cell in row))
                metadata = dict(workbook["Policy Metadata"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(metadata["Source SHA256"], checklist_fingerprint())
                self.assertEqual(metadata["Policy version"], definition["version"])
                self.assertEqual(metadata["Rule count"], len(expected))
                self.assertIn("synchronized", metadata["Policy source"])
            finally:
                workbook.close()
        self.assertEqual(CHECKLIST_PATH.read_bytes(), original)

    def test_readiness_rule_workbook_refresh_preserves_concurrent_user_edits(self):
        import readiness_checklist as policy
        from report_generation_service import write_readiness_rule_workbook
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "readiness_checklist.json"
            path.write_bytes(policy.CHECKLIST_PATH.read_bytes())
            with patch.object(policy, "CHECKLIST_PATH", path):
                output = write_readiness_rule_workbook()
                expected = hashlib.sha256(output.read_bytes()).hexdigest()
                definition = policy.checklist_definition()
                definition["categories"]["Scope and intent"]["scope.boundary"] = "Revised master boundary rule"
                path.write_text(json.dumps(definition), encoding="utf-8")
                write_readiness_rule_workbook(expected_workbook_sha256=expected)
                self.assertTrue(policy.checklist_fingerprint())
                expected = hashlib.sha256(output.read_bytes()).hexdigest()
                workbook = load_workbook(output)
                workbook["Checklist Rules"]["C2"] = "Unsynchronized user edit"
                workbook.save(output)
                workbook.close()
                original = output.read_bytes()
                with self.assertRaisesRegex(ValueError, "refusing to overwrite"):
                    write_readiness_rule_workbook(expected_workbook_sha256=expected)
                self.assertEqual(output.read_bytes(), original)

    def test_readiness_rule_workbook_edits_sync_to_json_and_invalidate_old_results(self):
        import readiness_checklist as policy
        from report_generation_service import write_readiness_rule_workbook
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "readiness_checklist.json"
            path.write_bytes(policy.CHECKLIST_PATH.read_bytes())
            with patch.object(policy, "CHECKLIST_PATH", path):
                workbook_path = write_readiness_rule_workbook()
                before = policy.checklist_fingerprint()
                workbook = load_workbook(workbook_path)
                workbook["Checklist Rules"]["C2"] = "The model boundary and every documented external dependency are clear"
                workbook["Checklist Rules"].append(["User extensions", "custom.integration", "Integration assumptions are explicitly documented"])
                workbook.save(workbook_path)
                workbook.close()
                with self.assertRaisesRegex(ValueError, "differ from JSON"):
                    policy.checklist_fingerprint()
                self.assertTrue(policy.sync_checklist_from_excel())
                self.assertNotEqual(policy.checklist_fingerprint(), before)
                self.assertEqual(policy.checklist_definition()["categories"]["Scope and intent"]["scope.boundary"], "The model boundary and every documented external dependency are clear")
                rows = policy.build_checklist({"scope.boundary": {"Assessment result": "pass", "Evidence status": "supported", "Evidence references": ["old-proof"]}})
                self.assertEqual(len(rows), len(policy.REQUIRED_IDS) + 1)
                self.assertEqual(rows[0]["Assessment result"], "uncertain")
                self.assertFalse(policy.sync_checklist_from_excel())
                original = path.read_bytes()
                workbook = load_workbook(workbook_path)
                workbook["Checklist Rules"].append(["User extensions", "custom.integration", "Duplicate"])
                workbook.save(workbook_path)
                workbook.close()
                with self.assertRaisesRegex(ValueError, "Duplicate"):
                    policy.sync_checklist_from_excel()
                self.assertEqual(path.read_bytes(), original)

    def test_readiness_rule_workbook_invalid_edits_do_not_change_json(self):
        import readiness_checklist as policy
        from report_generation_service import write_readiness_rule_workbook
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "readiness_checklist.json"
            original = policy.CHECKLIST_PATH.read_bytes()
            path.write_bytes(original)
            with patch.object(policy, "CHECKLIST_PATH", path):
                workbook_path = write_readiness_rule_workbook()
                original_workbook = workbook_path.read_bytes()
                for case in ("missing", "formula", "conflict", "lock", "replace"):
                    with self.subTest(case=case):
                        path.write_bytes(original)
                        workbook_path.write_bytes(original_workbook)
                        workbook = load_workbook(workbook_path)
                        workbook["Checklist Rules"]["C2"] = "Edited boundary obligation"
                        if case == "missing":
                            workbook["Checklist Rules"].delete_rows(2)
                        elif case == "formula":
                            workbook["Checklist Rules"]["C2"] = '=CONCAT("Not", " a literal rule")'
                        workbook.save(workbook_path)
                        workbook.close()
                        if case == "conflict":
                            path.write_bytes(original + b"\n")
                        lock = workbook_path.with_name("~$" + workbook_path.name)
                        if case == "lock":
                            lock.touch()
                        before = path.read_bytes()
                        before_workbook = workbook_path.read_bytes()
                        replace = Path.replace

                        def replace_policy(source, target):
                            if case == "replace" and source.suffix == ".xlsx":
                                raise PermissionError("Workbook replacement denied")
                            return replace(source, target)

                        try:
                            with patch.object(Path, "replace", replace_policy), self.assertRaises((ValueError, OSError)):
                                policy.sync_checklist_from_excel()
                            self.assertEqual(path.read_bytes(), before)
                            self.assertEqual(workbook_path.read_bytes(), before_workbook)
                        finally:
                            lock.unlink(missing_ok=True)

    def test_final_sysml_readiness_rejects_invalid_rule_edits_before_generation(self):
        import run_final_sysml_phase as phase
        import readiness_checklist as policy
        import io
        from contextlib import redirect_stdout

        with patch.object(sys, "argv", ["run_final_sysml_phase.py", "--snapshot-id", "selected"]), \
                patch.object(phase, "resolve_complete_snapshot_id", return_value="selected"), \
                patch.object(policy, "sync_checklist_from_excel", side_effect=ValueError("Invalid saved rules")), \
                patch.object(phase, "_run") as run:
            output = io.StringIO()
            with redirect_stdout(output):
                self.assertEqual(phase.main(), 1)
            run.assert_not_called()
            self.assertIn("synchronization blocked", output.getvalue())

    def test_readiness_checklist_fails_closed_without_semantic_coverage(self):
        from readiness_checklist import build_checklist, checklist_decision

        rows = build_checklist(candidate_evidence_refs=["approved:source-1"])
        self.assertGreaterEqual(len(rows), 45)
        self.assertEqual(checklist_decision(rows), "NOT_READY")
        self.assertTrue(all(row["Assessment result"] == "uncertain" for row in rows))
        self.assertTrue(all(row["Evidence references"] == [] for row in rows))

    def test_readiness_integrated_catalog_and_optional_diagnostics(self):
        import readiness_checklist as policy

        definition = policy.checklist_definition()
        self.assertEqual(list(definition["categories"]), [
            "Scope and intent", "Interfaces", "Clocks and resets", "FSMs and control behavior", "Functional behavior",
            "Datapath and data semantics", "Configuration, CSR, memory and registers",
            "Timing, ordering and performance constraints", "Internal structure and composition",
            "Consistency and conflict checks", "Traceability and governed evidence support",
            "Residual gaps and readiness blockers", "Optional DFT, DFX and diagnostics", "Final READY or NOT READY decision"])
        checks = {identity: {"Assessment result": "pass", "Evidence status": "supported", "Reason": "Synthetic governed adapter proof",
                             "Evidence references": ["fixture:approved:scope"]}
                  for identity in policy.REQUIRED_IDS - policy.OPTIONAL_IDS - {"readiness.decision"}}
        rows = policy.build_checklist(checks)
        optional = next(row for row in rows if row["Checklist ID"] == "diagnostics.behavior")
        self.assertEqual(optional["Assessment result"], "uncertain")
        self.assertEqual(optional["Readiness impact"], "informational")
        self.assertEqual(policy.checklist_decision(rows), "READY")
        self.assertNotIn(optional, policy.blocking_items(rows))
        import validate_downstream_coherence as central
        blocked = central._readiness_model_result(checks, ["snapshot_or_dependency_gate_failed"], [], ["fixture:approved:scope"])
        self.assertEqual(blocked["status"], "NOT_READY")
        self.assertEqual(next(row for row in blocked["checklist"] if row["Checklist ID"] == "readiness.decision")["Assessment result"], "fail")
        with patch.object(policy, "checklist_definition", return_value={**definition, "customized_rule_ids": ["readiness.decision"]}):
            self.assertEqual(policy.checklist_decision(policy.build_checklist(checks)), "NOT_READY")
        with patch.object(policy, "checklist_definition", return_value={**definition, "required_optional_rule_ids": ["diagnostics.behavior"]}):
            self.assertEqual(policy.checklist_decision(rows), "NOT_READY")
            required = policy.build_checklist(checks)
            self.assertEqual(policy.checklist_decision(required), "NOT_READY")
            self.assertIn("diagnostics.behavior", [row["Checklist ID"] for row in policy.blocking_items(required)])
        checks["readiness.decision"] = checks["io.names"]
        del checks["data.overflow"]
        self.assertEqual(policy.checklist_decision(policy.build_checklist(checks)), "NOT_READY")

    def test_readiness_integrated_standard_coverage_requires_all_local_checks(self):
        import readiness_checklist as policy

        checks = {identity: {"Assessment result": "pass", "Evidence status": "supported", "Reason": "Synthetic governed adapter proof",
                             "Evidence references": ["fixture:approved:scope"]}
                  for identity in policy.REQUIRED_IDS - policy.OPTIONAL_IDS - {"readiness.decision"}}
        for identity in policy.STANDARD_COVERED_IDS:
            checks[identity].update({"Coverage basis": "standard", "Standard-defined scope": "fixture:interface-A:standard-defined behavior"})
        self.assertEqual(policy.checklist_decision(policy.build_checklist(checks)), "READY")
        for identity in policy.STANDARD_PREREQUISITES:
            with self.subTest(missing=identity):
                incomplete = {key: dict(value) for key, value in checks.items() if key != identity}
                incomplete["readiness.decision"] = checks["io.names"]
                rows = policy.build_checklist(incomplete)
                self.assertEqual(policy.checklist_decision(rows), "NOT_READY")
                self.assertTrue(all(row["Assessment result"] == "uncertain" for row in rows
                                    if row["Checklist ID"] in policy.STANDARD_COVERED_IDS))
                forged = policy.build_checklist(checks)
                next(row for row in forged if row["Checklist ID"] == identity).update({"Assessment result": "not_applicable", "Evidence references": []})
                self.assertEqual(policy.checklist_decision(forged), "NOT_READY")
            checks["io.sideband"]["Assessment result"] = "not_applicable"
            self.assertEqual(policy.checklist_decision(policy.build_checklist(checks)), "READY")
            checks["io.standard"]["Assessment result"] = "not_applicable"
            self.assertEqual(policy.checklist_decision(policy.build_checklist(checks)), "NOT_READY")
            checks["io.standard"]["Assessment result"] = "pass"
        checks["io.meaning"]["Standard-defined scope"] = ""
        self.assertEqual(policy.checklist_decision(policy.build_checklist(checks)), "NOT_READY")

    def test_readiness_checklist_requires_complete_supported_rows(self):
        from readiness_checklist import build_checklist, checklist_decision

        rows = build_checklist()
        for row in rows:
            row.update({"Assessment result": "pass", "Evidence status": "supported",
                        "Reason": "Synthetic policy-mechanics fixture, not production source admission",
                        "Evidence references": ["fixture:complete-semantic-evidence"]})
        self.assertEqual(checklist_decision(rows), "READY")
        self.assertEqual(checklist_decision(rows[:-1]), "NOT_READY")
        rows[0]["Assessment result"] = "uncertain"
        self.assertEqual(checklist_decision(rows), "NOT_READY")
        rows[0]["Readiness impact"] = "non_blocking"
        self.assertEqual(checklist_decision(rows), "NOT_READY")
        rows[0].update({"Assessment result": "not_applicable", "Evidence status": "missing"})
        self.assertEqual(checklist_decision(rows), "NOT_READY")
        rows[0].update({"Assessment result": "not_applicable", "Evidence status": "supported"})
        self.assertEqual(checklist_decision(rows), "READY")
        self.assertEqual(checklist_decision([rows[0], *rows[1:-1], rows[0]]), "NOT_READY")
        rows[0]["Evidence references"] = []
        self.assertEqual(checklist_decision(rows), "NOT_READY")

    def test_readiness_checklist_rejects_truncated_policy_and_forged_ready(self):
        import readiness_checklist as checklist
        import validate_downstream_coherence as central

        with patch.object(checklist.json, "loads", return_value={"version": "governed-ai-readiness-checklist-v2", "authority": False, "categories": {}}):
            with self.assertRaises(ValueError):
                checklist.build_checklist()
        self.assertEqual(central.readiness_declaration(None, None, "model"), central.READINESS_NOT_READY)
        binding = {"checklist_fingerprint": checklist.checklist_fingerprint(), "model_bodies": {"model": "body"}}
        assessment = {"binding": binding, "policy_version": central.READINESS_POLICY_VERSION, "authority": False,
                      "publication_status": "ELIGIBLE", "models": {"model": {"status": "READY", "findings": [], "checklist": []}}}
        self.assertEqual(central.readiness_declaration(assessment, binding, "model"), central.READINESS_NOT_READY)
        assessment["models"] = []
        self.assertEqual(central.readiness_declaration(assessment, binding, "model"), central.READINESS_NOT_READY)

    def test_readiness_checklist_reports_blocked_scope_without_allowing_stamping(self):
        from dataclasses import replace
        import validate_downstream_coherence as central
        from readiness_checklist import REQUIRED_IDS

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, model = self.readiness_fixture(root)
            scope = replace(scope, findings=("SYSML_SCOPE_STALE_IPOS: core",))
            original = scope.system_path.read_bytes()
            with patch.object(central, "resolve_downstream_contract", return_value=contract), \
                    patch.object(central, "resolve_sysml_materialization_scope", return_value=scope), \
                    patch.object(central, "validate", side_effect=central.DownstreamCoherenceError("governed domain conflict")), \
                    patch("sysml_generation_service.resolve_sysml_descriptive_model", side_effect=central.DownstreamCoherenceError("stale description scope")):
                assessment = central.assess_final_readiness(root, "selected")
                self.assertEqual(assessment["status"], "NOT_READY")
                self.assertEqual(assessment["publication_status"], "BLOCKED")
                central.write_readiness_reports(root, assessment)
                summary_path = root / "artifacts/validation/readiness_checklists/readiness_summary.json"
                summary = json.loads(summary_path.read_text(encoding="utf-8"))
                self.assertEqual(summary["checklist_reference"]["excel"], "config/readiness_checklist.xlsx")
                self.assertIn("synchronized to JSON", summary["checklist_reference"]["editing"])
                self.assertEqual(len(summary["blocks"]), len(scope.blocks))
                self.assertEqual(len(summary["structural_models"]), 3)
                self.assertTrue(all(row["status"] == "NOT_READY" for row in summary["blocks"]))
                self.assertTrue(all(sum(row["checklist_results"].values()) == len(REQUIRED_IDS) for row in summary["blocks"]))
                summary_md = summary_path.with_suffix(".md").read_text(encoding="utf-8")
                self.assertIn("| Block/model | Readiness |", summary_md)
                self.assertIn("Checklist reference to be satisfied", summary_md)
                self.assertIn("../../../config/readiness_checklist.xlsx", summary_md)
                self.assertTrue(all(block.name in summary_md for block in scope.blocks))
                from openpyxl import load_workbook
                summary_excel = load_workbook(summary_path.with_suffix(".xlsx"))
                try:
                    self.assertEqual(summary_excel["Block Readiness"].max_row, len(scope.blocks) + 1)
                    self.assertEqual(summary_excel["System and Shared"].max_row, 4)
                    self.assertTrue(summary_excel["Block Readiness"].cell(2, 10).hyperlink.target.endswith(".xlsx"))
                    reference = next(cells[1] for cells in summary_excel["Metadata"].iter_rows(min_row=2)
                                     if cells[0].value == "Checklist reference")
                    self.assertEqual(reference.hyperlink.target, "../../../config/readiness_checklist.xlsx")
                finally:
                    summary_excel.close()
                for key, result in assessment["models"].items():
                    self.assertEqual(len(result["checklist"]), len(REQUIRED_IDS))
                    self.assertTrue(result["blocking_items_summary"])
                    self.assertEqual(central.readiness_declaration(assessment, assessment["binding"], key), central.READINESS_NOT_READY)
                    self.assertTrue(central.readiness_model_report_path(root, key).exists())
                    workbook = load_workbook(central.readiness_model_report_path(root, key).with_suffix(".xlsx"))
                    try:
                        self.assertEqual(workbook["Checklist"].max_row, len(REQUIRED_IDS) + 1)
                        self.assertEqual([cell.value for cell in workbook["Checklist"][1]][:6],
                                         ["Category", "Checklist item", "Assessment result", "Evidence status", "Reason", "Readiness impact"])
                        self.assertEqual(workbook["Checklist"]["C2"].data_type, "s")
                        self.assertIn("Residual Gaps", workbook.sheetnames)
                    finally:
                        workbook.close()
                with self.assertRaises(central.DownstreamCoherenceError):
                    central.stamp_final_readiness(root, assessment)
            self.assertEqual(scope.system_path.read_bytes(), original)

    def test_central_sysml_checker_rejects_changed_governed_descriptions(self):
        import validate_downstream_coherence as central
        from workflow_routing import build_ipos_functional_input
        from sysml_generation_service import resolve_sysml_descriptive_model
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, _model_dir = self.materialized_fixture(root)
            models = {}
            inputs = {}
            for block in scope.blocks:
                domain = "digital" if block.target == "Digital IPOS" else "analog"
                record = {"record_id": "context", "statement": "The Acquisition domain performs measurement.",
                          "mapped_block": block.name, "domain": domain, "snapshot_id": "selected", "admission_state": "admitted"}
                inputs[block.name] = build_ipos_functional_input(block.name, {"Function": "perform measurement", "Inputs": "samples", "Outputs": "results"}, domain=domain, descriptive_source_records=[record])
            contract.ipos_descriptive_input = lambda root, *, block, domain: inputs[block]
            for block in scope.blocks:
                domain = "digital" if block.target == "Digital IPOS" else "analog"
                models[block.name] = resolve_sysml_descriptive_model(root, contract, block=block.name, domain=domain)
                block.path.write_text(sysml_header(scope.snapshot_id, scope.fingerprint) + build_block_model(block.name, {}, [], [], [], scope=scope, descriptive_sections=models[block.name]["sections"]), encoding="utf-8")
            scope.manifest_path.write_text(json.dumps({"descriptive_model_version": "governed-ipos-shared-v1", "descriptive_models": models}), encoding="utf-8")
            self.assertEqual(central.check_sysml_governed_overviews(root, contract, scope), [])
            path = scope.blocks[0].path
            path.write_text(path.read_text(encoding="utf-8").replace("Acquisition domain", "Routing domain"), encoding="utf-8")
            self.assertIn("SYSML_GOVERNED_DESCRIPTION_CHANGED:" + scope.blocks[0].name, central.check_sysml_governed_overviews(root, contract, scope))

    def test_missing_assessment_and_duplicate_declarations_fail_closed(self):
        import validate_downstream_coherence as central
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, _model = self.readiness_fixture(root)
            self.assertEqual(central.check_sysml_readiness_declarations(root, contract, scope), [])
            text = scope.system_path.read_text(encoding="utf-8")
            scope.system_path.write_text(text.replace(central.READINESS_NOT_READY, central.READINESS_READY), encoding="utf-8")
            self.assertTrue(central.check_sysml_readiness_declarations(root, contract, scope))
            scope.system_path.write_text("// " + central.READINESS_NOT_READY + "\n" + text, encoding="utf-8")
            self.assertTrue(central.check_sysml_readiness_declarations(root, contract, scope))

    def readiness_fixture(self, root):
        contract, scope, _model_dir = self.materialized_fixture(root)
        statement = 'data carries measurement samples interpreted as unsigned 16-bit words using standard protocol specification "AMBA AHB5 specification revision 1.0" without local timing and handshake deviations.'
        contract.requirement_input = SimpleNamespace(rows=[{"source_req_id": identity, "statement": statement} for identity in contract.expected])
        contract.allocation = lambda identity: contract.expected.get(identity, {})
        contract.descriptive_context = lambda *args, **kwargs: ()
        port_path = root / "artifacts/stage2_mirco_arc/source_port_catalog.csv"
        with port_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Owner", "Port name", "Direction", "Type / details", "Ownership status"])
            writer.writeheader()
            for block in scope.blocks:
                writer.writerow({"Owner": block.name, "Port name": "data", "Direction": "input", "Type / details": "16 bits", "Ownership status": "approved"})
        files = [scope.system_path, scope.digital_path, scope.shared_path, *(block.path for block in scope.blocks)]
        scope.manifest_path.write_text(json.dumps({"snapshot_id": "selected", "model_path": scope.system_path.as_posix(), "descriptive_model_version": "governed-ipos-shared-v1",
                                      "model_files": {path.as_posix(): hashlib.sha256(path.read_bytes()).hexdigest() for path in files}}), encoding="utf-8")
        scope = resolve_sysml_materialization_scope(root, contract=contract)
        model = {"snapshot_id": "selected", "sections": {"functionality": ["The core performs measurement."], "scope": ["The acquisition domain performs measurement."]}, "audit": [], "residuals": [], "authority": False}
        return contract, scope, model

    def test_final_readiness_is_bound_stamped_and_rejects_forged_or_stale_outputs(self):
        import validate_downstream_coherence as central
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, model = self.readiness_fixture(root)
            with patch.object(central, "resolve_downstream_contract", return_value=contract), patch.object(central, "validate", return_value={"decision": "PASS", "findings": []}), patch("sysml_generation_service.resolve_sysml_descriptive_model", return_value=model):
                assessment = central.assess_final_readiness(root, "selected")
                self.assertEqual(assessment["status"], "NOT_READY")
                from readiness_checklist import REQUIRED_IDS
                self.assertEqual(len(assessment["models"][scope.blocks[0].path.relative_to(root).as_posix()]["checklist"]), len(REQUIRED_IDS))
                central.stamp_final_readiness(root, assessment)
                self.assertEqual(central.check_sysml_readiness_declarations(root, contract, scope), [])
                text = scope.system_path.read_text(encoding="utf-8")
                self.assertEqual(text.splitlines()[0], "// " + central.READINESS_NOT_READY)
                self.assertEqual(text.count("// " + central.READINESS_NOT_READY), 1)
                key = scope.blocks[0].path.relative_to(root).as_posix()
                checklist_path = central.readiness_model_report_path(root, key)
                checklist_before = checklist_path.read_text(encoding="utf-8")
                checklist_path.write_text("{}", encoding="utf-8")
                self.assertIn("SYSML_READINESS_CHECKLIST_REPORT_CHANGED:" + key, central.check_sysml_readiness_declarations(root, contract, scope))
                checklist_path.write_text(checklist_before, encoding="utf-8")
                with patch("readiness_checklist.checklist_fingerprint", return_value="changed-policy"):
                    self.assertIn("SYSML_READINESS_ASSESSMENT_STALE", central.check_sysml_readiness_declarations(root, contract, scope))
                scope.system_path.write_text(text + "// Changed model body\n", encoding="utf-8")
                self.assertTrue(central.check_sysml_readiness_declarations(root, contract, scope))
                with self.assertRaises(ValueError):
                    central.stamp_final_readiness(root, assessment)
                scope.system_path.write_text(text, encoding="utf-8")
                report_path = root / "artifacts/validation/final_readiness.json"
                forged = json.loads(report_path.read_text(encoding="utf-8"))
                forged["publication_status"] = "FORGED"
                report_path.write_text(json.dumps(forged), encoding="utf-8")
                self.assertIn("SYSML_READINESS_ASSESSMENT_CHANGED", central.check_sysml_readiness_declarations(root, contract, scope))
                report_path.write_text(json.dumps(assessment), encoding="utf-8")
                review_path = root / "artifacts/orchestrator/ipos_blocker_reviews/core/blocker_review.xlsx"
                review_path.parent.mkdir(parents=True)
                review_path.write_bytes(b"changed residual review")
                self.assertTrue(central.check_sysml_readiness_declarations(root, contract, scope))

    def test_readiness_json_round_trip_preserves_qualifiers_and_rejects_real_changes(self):
        import validate_downstream_coherence as central

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, model = self.readiness_fixture(root)
            with patch.object(central, "resolve_downstream_contract", return_value=contract), \
                    patch.object(central, "validate", return_value={"decision": "PASS", "findings": []}), \
                    patch("sysml_generation_service.resolve_sysml_descriptive_model", return_value=model):
                assessment = central.assess_final_readiness(root, "selected")
            assessment["relation_report"] = {"gaps": [{"qualifiers": (("condition", "enabled"), ("direction", "input")), "evidence": ["source-1"]}]}
            self.assertEqual(assessment["status"], "NOT_READY")
            self.assertEqual(assessment["publication_status"], "ELIGIBLE")
            report_path = root / "artifacts/validation/final_readiness.json"
            report_path.parent.mkdir(parents=True, exist_ok=True)
            report_path.write_text(json.dumps(assessment), encoding="utf-8")
            for key, result in assessment["models"].items():
                path = central.readiness_model_report_path(root, key)
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(json.dumps({"validator": "validate_downstream_coherence", "authority": False,
                                           "policy_version": assessment["policy_version"], "model": key,
                                           "binding": assessment["binding"], **result}), encoding="utf-8")
            with patch.object(central, "assess_final_readiness", return_value=assessment):
                self.assertEqual(central.check_sysml_readiness_declarations(root, contract, scope), [])
                for change in ("qualifier", "evidence", "status", "publication", "snapshot", "binding"):
                    with self.subTest(change=change):
                        modified = json.loads(json.dumps(assessment))
                        if change == "qualifier":
                            modified["relation_report"]["gaps"][0]["qualifiers"][0][1] = "disabled"
                        elif change == "evidence":
                            modified["relation_report"]["gaps"][0]["evidence"] = ["source-2"]
                        elif change == "status":
                            modified["status"] = "READY"
                        elif change == "publication":
                            modified["publication_status"] = "BLOCKED"
                        elif change == "snapshot":
                            modified["binding"]["snapshot_id"] = "other-snapshot"
                        else:
                            modified["binding"]["contract_fingerprint"] = "changed-contract"
                        report_path.write_text(json.dumps(modified), encoding="utf-8")
                        expected = "SYSML_READINESS_ASSESSMENT_STALE" if change in {"snapshot", "binding"} else "SYSML_READINESS_ASSESSMENT_CHANGED"
                        self.assertIn(expected, central.check_sysml_readiness_declarations(root, contract, scope))

    def test_nonstandard_io_and_standard_local_deviations_deny_readiness_not_publication(self):
        import validate_downstream_coherence as central
        port = {"Port name": "data", "Direction": "input", "Type / details": "16 bits"}
        self.assertIn("interface_timing_or_handshake_undefined", central.interface_readiness_findings(port, [{"statement": "data carries AHB measurement samples."}]))
        custom = 'data carries samples interpreted as unsigned words using protocol specification "Custom interface protocol revision 1.0" without local timing deviations.'
        self.assertIn("interface_timing_or_handshake_undefined", central.interface_readiness_findings(port, [{"statement": custom}]))
        deviating = 'data carries samples using protocol specification "AMBA AHB5 specification revision 1.0" without local timing deviations except an additional wait state.'
        self.assertIn("interface_timing_or_handshake_undefined", central.interface_readiness_findings(port, [{"statement": deviating}]))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, model = self.readiness_fixture(root)
            for row in contract.requirement_input.rows:
                row["statement"] = "data carries measurement samples."
            with patch.object(central, "resolve_downstream_contract", return_value=contract), patch.object(central, "validate", return_value={"decision": "PASS", "findings": []}), patch("sysml_generation_service.resolve_sysml_descriptive_model", return_value=model):
                assessment = central.assess_final_readiness(root, "selected")
                self.assertEqual(assessment["status"], "NOT_READY")
                self.assertEqual(assessment["generation_status"], "GENERATED")
                self.assertEqual(assessment["publication_status"], "ELIGIBLE")
                central.stamp_final_readiness(root, assessment)
                self.assertEqual(scope.system_path.read_text(encoding="utf-8").splitlines()[0], "// " + central.READINESS_NOT_READY)
                self.assertEqual(central.check_sysml_readiness_declarations(root, contract, scope), [])

    def test_sysml_uses_governed_ipos_descriptions_without_inventing_parts(self):
        from workflow_routing import build_ipos_functional_input
        from sysml_generation_service import resolve_sysml_descriptive_model
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, _model_dir = self.materialized_fixture(root)
            for block in scope.blocks:
                domain = "digital" if block.target == "Digital IPOS" else "analog"
                record = {"record_id": "context", "statement": "The Acquisition domain performs measurement.",
                          "mapped_block": block.name, "domain": domain, "snapshot_id": "selected", "admission_state": "admitted"}
                functional = build_ipos_functional_input(block.name, {"Function": "perform measurement", "Inputs": "samples", "Outputs": "results"}, domain=domain, descriptive_source_records=[record])
                contract.ipos_descriptive_input = MagicMock(return_value=functional)
                model = resolve_sysml_descriptive_model(root, contract, block=block.name, domain=domain)
                text = build_block_model(block.name, {}, [], [], [], scope=scope, descriptive_sections=model["sections"])
                self.assertIn("governedOverview:", text)
                self.assertIn("Acquisition domain", text)
                self.assertEqual(text.count("part def "), 1)
                self.assertNotIn("AcquisitionFSM", text)
                contract.ipos_descriptive_input.assert_called_once_with(root, block=block.name, domain=domain)

    def test_gui_snapshot_choice_preserves_workflow_default_and_manual_override(self):
        for bound, choice, requested, resolved in (
            (None, "Latest approved", None, "latest"),
            ("from-2b", "Current workflow snapshot", "from-2b", "from-2b"),
            ("from-2b", "snap-10 | date | reviewer | approved", "snap-10", "snap-10"),
        ):
            with self.subTest(choice=choice):
                gui = object.__new__(WorkflowGUI)
                gui.root = MagicMock()
                gui._activate_popup = MagicMock()
                gui.workflow_snapshot_id = bound
                connection = MagicMock()
                connection.execute.return_value.fetchall.return_value = [
                    {"snapshot_id": f"snap-{number}", "approved_at": "date", "approved_by": "reviewer", "status": "approved"}
                    for number in range(1, 11)
                ]
                with patch("workflow_gui.connect", return_value=connection), \
                        patch("workflow_gui.resolve_complete_snapshot_id", return_value=resolved) as resolve, \
                        patch("workflow_gui.tk.Toplevel") as dialog, patch("workflow_gui.tk.StringVar") as selected, \
                        patch("workflow_gui.ttk.Frame"), patch("workflow_gui.ttk.Label"), \
                        patch("workflow_gui.ttk.Combobox") as combo, patch("workflow_gui.ttk.Button") as button:
                    selected.return_value.get.return_value = choice

                    def confirm():
                        command = next(call.kwargs["command"] for call in button.call_args_list if call.kwargs.get("text") == "Use Snapshot")
                        command()

                    dialog.return_value.wait_window.side_effect = confirm
                    self.assertEqual(gui._snapshot_selector_args(), ["--snapshot-id", resolved])
                    self.assertNotIn("LIMIT", connection.execute.call_args.args[0])
                    self.assertNotIn("status = 'approved'", connection.execute.call_args.args[0])
                    choices = combo.call_args.kwargs["values"]
                    self.assertEqual(len(choices), 12 if bound else 11)
                    self.assertIn("snap-10 | date | reviewer | approved", choices)
                    selected.assert_called_once_with(value="Current workflow snapshot" if bound else "Latest approved")
                    self.assertEqual(resolve.call_args.kwargs["snapshot_id"], requested)
                    self.assertEqual(gui.workflow_snapshot_id, resolved)

    def test_snapshot_selector_shows_impacted_history_without_authorizing_generation(self):
        gui = object.__new__(WorkflowGUI)
        gui.root = MagicMock()
        gui._activate_popup = MagicMock()
        gui.workflow_snapshot_id = "current-approved"
        connection = MagicMock()
        connection.execute.return_value.fetchall.return_value = [
            {"snapshot_id": "older", "approved_at": "date", "approved_by": "reviewer", "status": "impacted"},
        ]
        label = "older | date | reviewer | impacted (history only)"
        with patch("workflow_gui.connect", return_value=connection), \
                patch("workflow_gui.resolve_complete_snapshot_id", side_effect=ApprovedResolverError("impacted snapshot rejected")) as resolve, \
                patch("workflow_gui.tk.Toplevel") as dialog, patch("workflow_gui.tk.StringVar") as selected, \
                patch("workflow_gui.ttk.Frame"), patch("workflow_gui.ttk.Label"), \
                patch("workflow_gui.ttk.Combobox") as combo, patch("workflow_gui.ttk.Button") as button, \
                patch("workflow_gui.messagebox.showerror") as error:
            selected.return_value.get.return_value = label
            dialog.return_value.wait_window.side_effect = lambda: next(
                call.kwargs["command"] for call in button.call_args_list if call.kwargs.get("text") == "Use Snapshot"
            )()
            self.assertIsNone(gui._snapshot_selector_args())
            self.assertIn(label, combo.call_args.kwargs["values"])
            self.assertEqual(resolve.call_args.kwargs["snapshot_id"], "older")
            error.assert_called_once()
            self.assertEqual(gui.workflow_snapshot_id, "current-approved")

    def test_viewer_keeps_2b_binding_and_refreshes_for_next_approval(self):
        gui = object.__new__(WorkflowGUI)
        first = MagicMock()
        first.require_valid.return_value = first
        second = MagicMock()
        second.require_valid.return_value = second
        with patch.object(WorkflowGUI, "_latest_approved_snapshot_id", return_value="unrelated-latest") as latest, \
                patch("workflow_gui.resolve_sysml_materialization_scope", side_effect=[first, second]) as resolve:
            gui.workflow_snapshot_id = "from-2b"
            self.assertIs(gui._get_sysml_materialization_scope(), first)
            self.assertIs(gui._get_sysml_materialization_scope(), first)
            gui.workflow_snapshot_id = "next-2b"
            self.assertIs(gui._get_sysml_materialization_scope(), second)
            self.assertEqual([call.args[1] for call in resolve.call_args_list], ["from-2b", "next-2b"])
            latest.assert_not_called()

    def test_successful_freeze_binds_snapshot_and_refreshes_views(self):
        gui = object.__new__(WorkflowGUI)
        gui._refresh_review_context = MagicMock()
        gui._refresh_work_area_and_traceability = MagicMock()
        gui._refresh_traceability_after_generation = MagicMock()
        gui.append_chat = MagicMock()
        on_ready = MagicMock()
        with patch.object(WorkflowGUI, "_latest_approved_snapshot_id", return_value="new-2b"):
            gui._after_snapshot_freeze(on_ready)
        self.assertEqual(gui.workflow_snapshot_id, "new-2b")
        self.assertEqual(gui._refresh_traceability_after_generation.call_args.kwargs["snapshot_id"], "new-2b")
        gui._refresh_traceability_after_generation.call_args.args[0]()
        self.assertEqual(gui._refresh_work_area_and_traceability.call_count, 2)
        on_ready.assert_called_once()

    def test_busy_is_set_before_worker_starts(self):
        gui = object.__new__(WorkflowGUI)
        gui.is_running = False
        with patch("workflow_gui.threading.Thread") as thread, patch("workflow_gui.messagebox.showwarning") as warning:
            gui._run_command_async(["command"], "first")
            self.assertTrue(gui.is_running)
            gui._run_command_async(["command"], "second")
            thread.assert_called_once()
            warning.assert_called_once()

    def test_generic_downstream_launchers_keep_workflow_snapshot(self):
        gui = object.__new__(WorkflowGUI)
        gui.workflow_snapshot_id = "from-2b"
        gui._run_command_async = MagicMock()
        gui._latest_approved_snapshot_id = MagicMock(return_value="unrelated-latest")
        gui.run_stage_script("6")
        self.assertEqual(gui._run_command_async.call_args.args[0][-2:], ["--snapshot-id", "from-2b"])
        gui.run_stage_script_after("4", MagicMock())
        self.assertEqual(gui._run_command_async.call_args.args[0][-2:], ["--snapshot-id", "from-2b"])
        gui._latest_approved_snapshot_id.assert_not_called()

    def test_missing_snapshot_does_not_crash_gui_sysml_scope(self):
        gui = object.__new__(WorkflowGUI)
        error = ApprovedResolverError("No complete approved snapshot is available. Checked: none")
        with patch("workflow_gui.resolve_sysml_materialization_scope", side_effect=error) as resolve:
            self.assertIsNone(gui._get_sysml_materialization_scope())
            self.assertEqual(gui._concrete_sysml_block_paths(), [])
            self.assertFalse(gui._sysml_block_file_is_concrete(Path("historical.sysml")))
            self.assertEqual(gui.sysml_scope_diagnostic, str(error))
            resolve.assert_called_once()

    def test_missing_snapshot_is_not_a_current_sysml_review(self):
        with patch("workflow_gui.resolve_sysml_materialization_scope", side_effect=ApprovedResolverError("No approved snapshot")):
            self.assertFalse(WorkflowGUI._sysml_review_is_current("unavailable-snapshot"))

    def exercise_gui_widgets(self, root, scope):
        window = tk.Tk()
        window.withdraw()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.sysml_materialization_scope = scope
            gui.sysml_scope_diagnostic = ""
            gui.sysml_selected_block = None
            gui.sysml_selected_connection = None
            gui.sysml_tree = ttk.Treeview(window)
            gui.sysml_hierarchy_tree = ttk.Treeview(window)
            gui.sysml_graph_canvas = tk.Canvas(window)
            gui.sysml_file_choice_var = tk.StringVar(window)
            gui.sysml_file_choice = ttk.Combobox(window, textvariable=gui.sysml_file_choice_var)
            gui.sysml_file_var = tk.StringVar(window)
            gui.sysml_viewer = tk.Text(window)
            gui.work_tabs = MagicMock()
            gui.sysml_architecture_tab = MagicMock()
            with patch("workflow_gui.REPO_ROOT", root), patch("workflow_gui.resolve_sysml_materialization_scope", return_value=scope):
                gui._refresh_sysml_architecture()
                self.assertEqual(set(gui.sysml_file_choice_paths.values()), {block.path for block in scope.blocks})
                for block in scope.blocks:
                    gui.sysml_file_choice_var.set(block.path.relative_to(root).as_posix())
                    gui._on_sysml_file_choice()
                    self.assertEqual(gui.sysml_viewer.get("1.0", "end-1c"), block.path.read_text(encoding="utf-8"))
                    parent = next(item for item in gui.sysml_hierarchy_tree.get_children(gui.sysml_hierarchy_tree.get_children()[0])
                                  if gui.sysml_hierarchy_tree.item(item, "text") == block.target)
                    item = next(item for item in gui.sysml_hierarchy_tree.get_children(parent)
                                if gui.sysml_hierarchy_tree.item(item, "text") == block.name)
                    gui.sysml_hierarchy_tree.selection_set(item)
                    gui._on_sysml_hierarchy_select()
                    self.assertEqual(gui.sysml_file_var.get(), block.path.relative_to(root).as_posix())
                    self.assertEqual(gui.sysml_viewer.get("1.0", "end-1c"), block.path.read_text(encoding="utf-8"))
                window.update_idletasks()
        finally:
            window.destroy()

    def test_real_tk_dropdown_tree_and_source_opening(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _, scope, _ = self.materialized_fixture(root)
            self.exercise_gui_widgets(root, scope)

    def materialized_fixture(self, root):
        expected = {}
        for target, slug, name, source_id in (
            ("Digital IPOS", "control-engine", "CONTROL Engine", "SOURCE_001"),
            ("Digital IPOS", "packet-router", "Packet Router", "SOURCE_002"),
            ("Analog IPOS", "signal-front-end", "Signal Front End", "req-3"),
        ):
            stage = "stage6_digital_ipos" if target == "Digital IPOS" else "stage7_analog_ipos"
            output = root / "artifacts" / stage
            directory = output / "blocks" / slug
            directory.mkdir(parents=True)
            (directory / "document.md").write_text("Generated block document", encoding="utf-8")
            matrix = output / "ipos_traceability_matrix.csv"
            fields = ["mapped_block", "source_req_id", "snapshot_id", "ipos_req_id", "requirement_statement"]
            exists = matrix.exists()
            with matrix.open("a", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                if not exists:
                    writer.writeheader()
                writer.writerow(dict(zip(fields, (name, source_id, "selected", source_id, "Complete approved behavior."))))
            expected[source_id] = {"approved_block": name, "owning_target": target}
        contract = SimpleNamespace(snapshot_id="selected", expected=expected)
        scope = resolve_sysml_materialization_scope(root, contract=contract).require_valid()
        inventory = {block.name: {"Function": "Perform approved function."} for block in scope.blocks}
        rows = scope.ipos_rows(root)
        model_dir = root / "artifacts/stage2_mirco_arc/sysml"
        (model_dir / "blocks").mkdir(parents=True)
        header = sysml_header(scope.snapshot_id, scope.fingerprint)
        (model_dir / "shared_definitions.sysml").write_text(header + build_shared_model(), encoding="utf-8")
        (model_dir / "DigitalSubsystem.sysml").write_text(header + build_digital_subsystem_model(inventory, [], [], scope=scope), encoding="utf-8")
        scope.system_path.write_text(header + build_model([], inventory, [], scope=scope), encoding="utf-8")
        for block in scope.blocks:
            block.path.write_text(header + build_block_model(block.name, inventory[block.name], [], [], rows[block.name], scope=scope), encoding="utf-8")
        return contract, scope, model_dir

    def test_scope_parity_generation_validation_review_and_gui(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, model_dir = self.materialized_fixture(root)
            self.assertEqual(sysml_materialization_findings(root, scope), [])
            self.assertEqual(_check_sysml_requirement_coverage(root, contract, contract.expected, scope), [])
            with patch("review_architecture_sysml.read_csv", return_value=[]):
                review = _review(scope.system_path, model_dir, root, root, root, root, root, scope=scope)
            self.assertTrue(all(row["status"] == "PASS" for row in review), review)
            gui = object.__new__(WorkflowGUI)
            gui.sysml_materialization_scope = scope
            gui.sysml_scope_diagnostic = ""
            with patch("workflow_gui.REPO_ROOT", root):
                self.assertEqual(set(gui._concrete_sysml_block_paths()), {block.path for block in scope.blocks})
                self.assertTrue(all(gui._sysml_block_file_is_concrete(block.path) for block in scope.blocks))
            for block in scope.blocks:
                self.assertEqual(scope.block_for_node(block.node).path, block.path)
                self.assertEqual(scope.block_for_label(block.name.swapcase()), block)
            self.assertEqual(scope.snapshot_id, contract.snapshot_id)
            self.assertEqual(scope.fingerprint, downstream_contract_fingerprint(contract))

    def test_central_crosscheck_rejects_extra_models_and_unresolved_references(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract, scope, model_dir = self.materialized_fixture(root)
            extra = model_dir / "blocks/inventory_only.sysml"
            extra.write_text("package Unsupported {}", encoding="utf-8")
            digital = model_dir / "DigitalSubsystem.sysml"
            original = digital.read_text(encoding="utf-8")
            digital.write_text(original + "\nprivate import UnknownPackage::*;\npart bad : UndefinedBlock;\n", encoding="utf-8")
            findings = _check_sysml_requirement_coverage(root, contract, contract.expected, scope)
            for code in ("SYSML_STALE_IPOS_BLOCK_FILE", "SYSML_UNRESOLVED_IMPORT", "SYSML_UNRESOLVED_PART"):
                self.assertTrue(any(finding.startswith(code) for finding in findings), findings)
            scope.blocks[0].path.unlink()
            self.assertTrue(any("SYSML_CREATED_IPOS_BLOCK_MISSING" in finding for finding in sysml_materialization_findings(root, scope)))

    def test_gui_canvas_callbacks_open_every_scope_model(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _, scope, _ = self.materialized_fixture(root)
            gui = object.__new__(WorkflowGUI)
            gui.sysml_materialization_scope = scope
            gui.sysml_scope_diagnostic = ""
            gui.sysml_selected_block = None
            gui.sysml_selected_connection = None
            gui.work_tabs = MagicMock()
            gui.sysml_architecture_tab = MagicMock()
            gui._select_sysml_file = MagicMock()
            gui._show_sysml_block_details = MagicMock()
            for attribute, draw, tag in (
                ("sysml_graph_canvas", gui._draw_sysml_architecture, "sysml_node"),
                ("sysml_hierarchy_window_canvas", gui._draw_sysml_hierarchy_window, "sysml_hierarchy_node"),
            ):
                canvas = MagicMock()
                setattr(gui, attribute, canvas)
                canvas.find_withtag.return_value = (1,)
                with patch("workflow_gui.REPO_ROOT", root):
                    draw()
                    callback = next(call.args[2] for call in canvas.tag_bind.call_args_list
                                    if call.args[:2] == (tag, "<Button-1>"))
                    for block in scope.blocks:
                        canvas.gettags.return_value = (tag, block.node if tag == "sysml_node" else block.path.as_posix())
                        callback(SimpleNamespace())
                        gui._select_sysml_file.assert_called_with(block.path)
                    emitted_tags = [call.kwargs.get("tags", ()) for call in canvas.create_rectangle.call_args_list]
                    for block in scope.blocks:
                        self.assertIn((tag, block.node if tag == "sysml_node" else block.path.as_posix()), emitted_tags)

    def test_resolved_extra_and_duplicate_parts_are_not_hidden(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _, scope, _ = self.materialized_fixture(root)
            text = scope.digital_path.read_text(encoding="utf-8")
            block = next(block for block in scope.blocks if block.target == "Digital IPOS")
            scope.digital_path.write_text(text + f"\npart extra : {block.definition};\npart def InventoryOnly {{}}\n", encoding="utf-8")
            findings = sysml_materialization_findings(root, scope)
            self.assertTrue(any(finding.startswith("SYSML_OUT_OF_SCOPE_PART:") for finding in findings), findings)
            self.assertTrue(any(finding.startswith("SYSML_OUT_OF_SCOPE_PART_DEFINITION:") for finding in findings), findings)
            analog = next(block for block in scope.blocks if block.target == "Analog IPOS")
            text = scope.system_path.read_text(encoding="utf-8")
            declaration = f"part {analog.node.rsplit('.', 1)[-1]} : {analog.definition};"
            text = text.replace(declaration, "").replace("part def TopAnalogArchitecture {", "part def TopAnalogArchitecture {\n" + declaration)
            scope.system_path.write_text(text, encoding="utf-8")
            findings = sysml_materialization_findings(root, scope)
            self.assertIn(f"SYSML_BLOCK_PART_MISSING_OR_DUPLICATE: {analog.name}", findings)

    def test_materialization_scope_reports_missing_and_unapproved_documents(self):
        contract = SimpleNamespace(snapshot_id="selected", expected={
            "SOURCE_001": {"owning_target": "Digital IPOS", "approved_block": "Control Engine"},
        })
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch("validate_downstream_coherence.created_ipos_block_directories", return_value={"unknown": root}):
                scope = resolve_sysml_materialization_scope(root, contract=contract)
            self.assertIn("SYSML_SCOPE_IPOS_DOCUMENT_MISSING: Control Engine", scope.findings)
            self.assertTrue(any("OUT_OF_SCOPE_IPOS" in finding for finding in scope.findings))
            self.assertEqual(scope.blocks[0].name, "Control Engine")

    def test_valid_empty_scope_and_ambiguous_manifest_are_explicit(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            contract = SimpleNamespace(snapshot_id="selected", expected={})
            scope = resolve_sysml_materialization_scope(root, contract=contract).require_valid()
            self.assertEqual(scope.blocks, ())
            self.assertEqual(sysml_materialization_findings(root, scope), [])
            scope.system_path.parent.mkdir(parents=True)
            header = sysml_header(scope.snapshot_id, scope.fingerprint)
            scope.shared_path.write_text(header + build_shared_model(), encoding="utf-8")
            scope.digital_path.write_text(header + build_digital_subsystem_model({}, [], [], scope=scope), encoding="utf-8")
            scope.system_path.write_text(header + build_model([], {}, [], scope=scope), encoding="utf-8")
            self.assertEqual(sysml_materialization_findings(root, scope), [])
            self.assertEqual(list((scope.system_path.parent / "blocks").glob("*.sysml")), [])
            for prefix in ("alpha", "beta"):
                (scope.manifest_path.parent / f"{prefix}_architecture_model_manifest.json").write_text("{}", encoding="utf-8")
            ambiguous = resolve_sysml_materialization_scope(root, contract=contract)
            self.assertTrue(any(finding.startswith("SYSML_MANIFEST_ADDRESS_AMBIGUOUS") for finding in ambiguous.findings))

    def test_stage3_gate_treats_stage2_sysml_review_as_non_blocking(self):
        repo_root = Path(__file__).resolve().parents[1]
        gate = (repo_root / "scripts/run_stage3_srs_gate.py").read_text(encoding="utf-8")
        review_section = gate[gate.index("review_path"):gate.index("print(\"Stage 3 SRS full flow")]
        self.assertIn("non-blocking derived-artifact diagnostic", review_section)
        self.assertNotIn("return 2", review_section)

    def test_rules_define_mapping_and_snapshot_as_authority(self):
        repo_root = Path(__file__).resolve().parents[1]
        rules = (repo_root / ".github/copilot-instructions.md").read_text(encoding="utf-8")
        skill = (repo_root / ".github/skills/workflow-stage-gate/SKILL.md").read_text(encoding="utf-8")
        for text in (rules, skill):
            self.assertIn("sole architectural authority", text)
            self.assertIn("Stage 2 SysML is derived", text)


if __name__ == "__main__":
    unittest.main()