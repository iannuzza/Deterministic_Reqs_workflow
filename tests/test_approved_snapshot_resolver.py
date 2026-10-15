import unittest
import hashlib
import json
import tempfile
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from approved_snapshot_resolver import (
    ApprovedResolverError, resolve_complete_authoritative_input,
    resolve_governed_descriptive_context, validate_governed_source_context,
)
from approved_profile_resolver import ProfileResolutionError, resolve_snapshot_profile
from canonical_store import fingerprint
from run_allocation_crosscheck import validate_materialization_manifest
from requirement_corpus import RequirementInput
from snapshot_manager import create_approved_snapshot


class _Connection:
    def __init__(self, rows):
        self.rows = rows

    def execute(self, _query, _params):
        return self

    def fetchall(self):
        return self.rows

    def close(self):
        pass


class ApprovedSnapshotResolverTests(unittest.TestCase):
    def test_governed_descriptions_require_explicit_membership_and_provenance(self):
        text = "The Control domain coordinates acquisition."
        record = {"record_id": "context-1", "statement": text,
                  "content_sha256": hashlib.sha256(text.encode()).hexdigest(),
                  "source_file": "source.pdf", "source_revision": "1",
                  "source_spans_json": json.dumps([{"page": 2, "line": 3,
                        "text_file": "extract.txt", "artifact_sha256": "a" * 64}]),
                  "evidence_kind": "descriptive", "mapped_block": "Control",
                  "domain": "digital", "review_decision": "pending",
                  "review_rationale": "Reviewed source section ownership"}
        profile = {"block_defs": {"Control": {}}, "governed_source_context": {"version": 1, "records": [record]}}
        snapshot = {"project_id": "P"}
        with patch("approved_snapshot_resolver.resolve_approved_snapshot", return_value=snapshot), \
                patch("approved_snapshot_resolver.resolve_snapshot_profile", return_value=profile):
            self.assertEqual(resolve_governed_descriptive_context(Path("."), snapshot_id="S", project_id="P"), ())
            record["review_decision"] = "approved"
            admitted = resolve_governed_descriptive_context(Path("."), snapshot_id="S", project_id="P")
            self.assertEqual(admitted[0]["snapshot_id"], "S")
            self.assertEqual(admitted[0]["admission_state"], "admitted")
            record["statement"] += " Altered."
            with self.assertRaisesRegex(ApprovedResolverError, "text_hash_mismatch"):
                resolve_governed_descriptive_context(Path("."), snapshot_id="S", project_id="P")
        self.assertEqual(validate_governed_source_context({}), [])

    def test_profile_resolution_is_hash_pinned_and_rejects_mutated_payload(self):
        payload = {"block_defs": {"Control": {}}}
        digest = fingerprint(payload)
        snapshot = {"project_id": "P", "metadata_json": json.dumps({"profiles": {"architecture": digest}})}
        connection = _Connection([{"profile_payload_json": json.dumps(payload)}])
        with patch("approved_profile_resolver.connect", return_value=connection):
            self.assertEqual(resolve_snapshot_profile(Path("."), snapshot), payload)
            connection.rows = [{"profile_payload_json": "{}"}]
            with self.assertRaisesRegex(ProfileResolutionError, "fingerprint mismatch"):
                resolve_snapshot_profile(Path("."), snapshot)

    def test_pending_context_is_not_approval_and_admitted_owner_requires_review(self):
        text = "The Acquisition domain performs measurement."
        record = {"record_id": "candidate", "statement": text,
                  "content_sha256": hashlib.sha256(text.encode()).hexdigest(),
                  "source_file": "source.pdf", "source_revision": "",
                  "source_spans_json": json.dumps([{"page": 1, "line": 2,
                      "text_file": "extract.txt", "artifact_sha256": "a" * 64}]),
                  "evidence_kind": "descriptive", "review_decision": "pending"}
        profile = {"block_defs": {"Core": {}},
                   "governed_source_context": {"version": 1, "records": [record]}}
        self.assertEqual(validate_governed_source_context(profile), [])
        record.update(review_decision="approved", mapped_block="Other", domain="digital")
        findings = validate_governed_source_context(profile)
        self.assertIn("governed_source_context_unknown_owner:candidate", findings)
        self.assertIn("governed_source_context_missing_source_revision:candidate", findings)
        self.assertIn("governed_source_context_missing_review_rationale:candidate", findings)

    def test_existing_review_gate_binds_admitted_context_to_original_candidate(self):
        from validate_stage2_profile import validate_approval_evidence, _draft_review_sha256
        text = "The Core domain coordinates acquisition."
        record = {"record_id": "candidate", "statement": text,
                  "content_sha256": hashlib.sha256(text.encode()).hexdigest(),
                  "source_file": "source.pdf", "source_revision": "1",
                  "source_spans_json": json.dumps([{"page": 1, "line": 2,
                      "text_file": "extract.txt", "artifact_sha256": "a" * 64}]),
                  "evidence_kind": "descriptive", "review_decision": "pending",
                  "mapped_block": "", "domain": "", "review_rationale": ""}
        draft = {"block_defs": {"Core": {}}, "approval": {},
                 "governed_source_context": {"version": 1, "records": [record]}}
        profile = json.loads(json.dumps(draft))
        profile["governed_source_context"]["records"][0].update(
            review_decision="approved", mapped_block="Core", domain="digital", review_rationale="Explicit source ownership")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            mapping = root / "mapping.csv"
            mapping.write_text("requirement_id\nREQ_1\n", encoding="utf-8")
            profile["approval"] = {"status": "approved", "approved_by": "Reviewer",
                "reviewed_draft_profile_sha256": _draft_review_sha256(draft),
                "reviewed_mapping_preview_csv_sha256": hashlib.sha256(mapping.read_bytes()).hexdigest()}
            draft_path, profile_path = root / "draft.json", root / "profile.json"
            draft_path.write_text(json.dumps(draft), encoding="utf-8")
            profile_path.write_text(json.dumps(profile), encoding="utf-8")
            from openpyxl import Workbook, load_workbook
            from mapping_review_sync import write_descriptive_review_sheet, DESCRIPTIVE_SHEET
            workbook_path = mapping.with_suffix(".xlsx")
            workbook = Workbook()
            workbook.save(workbook_path)
            workbook.close()
            write_descriptive_review_sheet(workbook_path, draft)
            workbook = load_workbook(workbook_path)
            for column, field in enumerate(("review_decision", "mapped_block", "domain", "review_rationale"), 1):
                workbook[DESCRIPTIVE_SHEET].cell(2, column).value = profile["governed_source_context"]["records"][0][field]
            workbook.save(workbook_path)
            workbook.close()
            self.assertEqual(validate_approval_evidence(profile_path, draft_path, mapping), [])
            changed = profile["governed_source_context"]["records"][0]
            changed["statement"] += " Invented behavior."
            changed["content_sha256"] = hashlib.sha256(changed["statement"].encode()).hexdigest()
            profile_path.write_text(json.dumps(profile), encoding="utf-8")
            self.assertIn("Approved source context differs from the reviewed descriptive candidate",
                          validate_approval_evidence(profile_path, draft_path, mapping))

    def test_new_approval_has_new_snapshot_identity_for_same_material(self):
        connection = MagicMock()
        connection.execute.return_value.fetchone.return_value = {
            "allocation_class": "system_level", "owning_target": "SRS",
            "lineage_mode": "normal_hierarchical", "source_origin_req_ids": "SOURCE-1",
            "hierarchy_parent_req_ids": "",
        }
        profiles = {"vocabulary": {"content_fingerprint": "vocabulary-hash", "revision_number": "1"}}
        with patch("snapshot_manager.connect", return_value=connection), \
                patch("snapshot_manager.resolve_approved_profiles", return_value=profiles), \
                patch("snapshot_manager.record_event"), \
                patch("snapshot_manager.datetime") as clock:
            clock.now.return_value.isoformat.return_value = "2026-01-01T00:00:00.000001+00:00"
            options = {"project_id": "P", "canonical_revision_ids": ["revision-current"],
                       "mapping_ids": ["mapping-approved"], "approved_by": "reviewer",
                       "retrieval": {"version": "1", "hash": "retrieval-hash"}, "semantic_enabled": False}
            first = create_approved_snapshot(Path("."), **options)
            clock.now.return_value.isoformat.return_value = "2026-01-01T00:00:00.000002+00:00"
            second = create_approved_snapshot(Path("."), **options)
        self.assertNotEqual(first, second)
        inserts = [call.args for call in connection.execute.call_args_list if "INSERT INTO snapshots" in call.args[0]]
        self.assertEqual(len(inserts), 2)
        self.assertEqual(inserts[0][1][11], inserts[1][1][11])
        self.assertEqual(inserts[0][1][0], first)
        self.assertEqual(inserts[1][1][0], second)

    def _input(self, snapshot_id, mapping):
        return RequirementInput(
            stage="test",
            snapshot_id=snapshot_id,
            corpus_hash="corpus",
            snapshot_hash="snapshot",
            source_path=Path("canonical.sqlite"),
            rows=({"canonical_id": "req-1"},),
            source_links={"req-1": ("SRC-1",)},
            mapping=mapping,
        )

    def test_explicit_incomplete_snapshot_is_rejected(self):
        incomplete = self._input("bad", {})
        with patch("approved_snapshot_resolver.resolve_authoritative_input", return_value=incomplete):
            with self.assertRaisesRegex(ApprovedResolverError, "bad.*approved mapping set is empty"):
                resolve_complete_authoritative_input(Path("."), "test", project_id="P", snapshot_id="bad")

    def test_no_complete_snapshot_is_rejected(self):
        with patch("canonical_store.connect", return_value=_Connection([])):
            with self.assertRaisesRegex(ApprovedResolverError, "No complete approved snapshot"):
                resolve_complete_authoritative_input(Path("."), "test", project_id="P")

    def test_multiple_complete_snapshots_choose_newest_then_id(self):
        candidates = [
            {"snapshot_id": "older", "approved_at": "2026-09-14T10:00:00Z"},
            {"snapshot_id": "newer-b", "approved_at": "2026-09-15T10:00:00Z"},
            {"snapshot_id": "newer-a", "approved_at": "2026-09-15T10:00:00Z"},
        ]
        inputs = {row["snapshot_id"]: self._input(row["snapshot_id"], {"req-1": "block"}) for row in candidates}
        with patch("canonical_store.connect", return_value=_Connection(candidates)):
            with patch("approved_snapshot_resolver.resolve_authoritative_input", side_effect=lambda *args, **kwargs: inputs[kwargs["snapshot_id"]]):
                selected, record = resolve_complete_authoritative_input(Path("."), "test", project_id="P")
        self.assertEqual(selected.snapshot_id, "newer-b")
        self.assertEqual(record["selected_snapshot_id"], "newer-b")
        self.assertEqual(record["selection_mode"], "deterministic_complete_candidate")

    def test_valid_explicit_snapshot_is_preserved(self):
        valid = self._input("explicit-valid", {"req-1": "block"})
        with patch("approved_snapshot_resolver.resolve_authoritative_input", return_value=valid):
            selected, record = resolve_complete_authoritative_input(
                Path("."), "test", project_id="P", snapshot_id="explicit-valid"
            )
        self.assertEqual(selected.snapshot_id, "explicit-valid")
        self.assertEqual(record["selection_mode"], "explicit")
        self.assertEqual(record["selected_snapshot_id"], "explicit-valid")

    def test_incomplete_latest_snapshot_never_falls_back_to_history(self):
        candidates = [{"snapshot_id": "new", "approved_at": "2026-01-02"},
                      {"snapshot_id": "old", "approved_at": "2026-01-01"}]
        with patch("canonical_store.connect", return_value=_Connection(candidates)), \
                patch("approved_snapshot_resolver.resolve_authoritative_input", return_value=self._input("new", {})) as resolve:
            with self.assertRaisesRegex(ApprovedResolverError, "new.*approved mapping set is empty"):
                resolve_complete_authoritative_input(Path("."), "test", project_id="P")
            resolve.assert_called_once_with(Path("."), "test", project_id="P", snapshot_id="new")

    def test_manifest_records_selection_and_rejections(self):
        manifest = {
            "materialized": True,
            "row_count": 1,
            "selection": {
                "selected_snapshot_id": "good",
                "attempted_snapshots": [
                    {"snapshot_id": "bad", "accepted": False, "reason": "mapping set is empty"},
                    {"snapshot_id": "good", "accepted": True, "reason": "complete"},
                ],
            },
        }
        validate_materialization_manifest(manifest, [{"snapshot_id": "good"}])
        self.assertEqual(manifest["selection"]["selected_snapshot_id"], "good")
        self.assertEqual(manifest["selection"]["attempted_snapshots"][0]["reason"], "mapping set is empty")

    def test_stale_manifest_is_rejected(self):
        manifest = {
            "materialized": True,
            "row_count": 1,
            "selection": {
                "selected_snapshot_id": "old",
                "attempted_snapshots": [{"snapshot_id": "old", "accepted": True}],
            },
        }
        with self.assertRaisesRegex(RuntimeError, "do not match"):
            validate_materialization_manifest(manifest, [{"snapshot_id": "new"}])


if __name__ == "__main__":
    unittest.main()
