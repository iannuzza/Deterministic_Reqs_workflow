import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import ipos_blocker_review as review


class IPOSBlockerReviewTests(unittest.TestCase):
    def test_file_operation_failures_are_not_reported_as_content_gaps(self):
        from workflow_routing import build_ipos_functional_input
        functional = build_ipos_functional_input("Module", {},
            [{"source_req_id": "one", "statement": "The user shall select MODE."}], domain="digital")
        for code in ("ipos_generation_failure", "ipos_publication_failure"):
            with self.subTest(code=code):
                finding = code + ":[WinError 5] Access is denied: 'temporary.docx' -> 'output.docx'"
                rows = review.blocker_rows("Module", [finding], functional)
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]["Likely failing layer"], "generation/publication tooling")
                self.assertEqual(rows[0]["Observed problem"], finding)
                self.assertNotIn("Supported claim/section fails", rows[0]["Observed problem"])
                self.assertIn("No missing source meaning", rows[0]["Expected supported meaning"])

    def test_missing_descriptions_generate_and_retain_postgeneration_review(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            writer = Mock()
            with patch("generate_ipos_specs._block_inventory_record", return_value={}):
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module",
                         rows=[{"source_req_id": "one", "statement": "The user shall select MODE."}], candidate_rows=[], writer=writer)
            writer.assert_called_once()
            self.assertEqual(result["generation_status"], "GENERATABLE")
            self.assertEqual(result["publication_status"], "NOT_ATTEMPTED")
            self.assertEqual(result["readiness_status"], "NOT_ASSESSED")
            self.assertTrue(result["residuals"])
            self.assertTrue(review.review_paths(root, "Module")[0].exists())

    def test_source_update_readiness_action_remains_valid_review_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            row = self.blocker()
            row["Resulting action"] = "generate supported content; retain source/spec update gap for readiness"
            path = review.write_blocker_review(root, "snapshot", "Module", [row])
            self.assertTrue(path.exists())
            self.assertEqual(review.read_blocker_review(root, snapshot_id="snapshot", block="Module", require_saved=False)["snapshot_id"], "snapshot")

    def test_unedited_residual_review_does_not_veto_materialization(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            functional = SimpleNamespace(block="Module")
            self.assertIs(review.apply_descriptive_omissions(root, snapshot_id="snapshot", functional_input=functional), functional)

    def test_archived_section_omit_preserves_legacy_scope_and_new_writing_rule(self):
        from workflow_routing import build_ipos_functional_input, build_ipos_claim_plan, _ipos_claim_presentation_role, IPOS_MAIN_FUNCTION_WRITING_RULE
        functional = build_ipos_functional_input("Module",
            {"Function": "Control sampling", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "configuration", "statement": "The user shall select CHANNEL_MODE."},
             {"source_req_id": "transfer", "statement": "The controller shall transfer samples to memory."}],
            domain="digital")
        plan = build_ipos_claim_plan(functional)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook, _manifest, _companion = review.review_paths(root, "Module")
            workbook.parent.mkdir(parents=True)
            receipt_path = workbook.with_name("descriptive_omissions.json")
            for modern in (False, True):
                identities = sorted(claim.claim_id for claim in plan.claims
                                    if _ipos_claim_presentation_role(claim, main_function_summary=modern) == "functions")
                receipt = {"snapshot_id": "snapshot", "block": "Module", "scope": "descriptive_only",
                           "authority": False, "gate_override": False, "workbook_sha256": "verified-archive",
                           "omitted_claim_ids": identities}
                if modern:
                    receipt["writing_rule"] = IPOS_MAIN_FUNCTION_WRITING_RULE
                receipt_path.write_text(json.dumps(receipt), encoding="utf-8")
                context = {"decisions": [{"Section or claim": "1.4 Supported block functions",
                                         "User decision": "omit (as non-blocking)"}]}
                with patch.object(review, "read_blocker_review", return_value=context) as read:
                    applied = review.apply_descriptive_omissions(root, snapshot_id="snapshot", functional_input=functional)
                    self.assertEqual(list(applied.descriptive_omitted_claim_ids), identities)
                    self.assertEqual(read.call_args.kwargs["history_sha256"], "verified-archive")
                self.assertEqual(applied.local_requirement_evidence, functional.local_requirement_evidence)
                self.assertEqual(len(identities), 2 if modern else 1)

    def test_docx_requirement_body_preserves_literal_multiplication_and_ids(self):
        from docx import Document
        from workflow_routing import apply_docx_ipos_requirement_style_guard

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ipos.docx"
            document = Document()
            document.add_paragraph("Description stays unchanged.")
            document.add_paragraph("IPOS-MODULE-001")
            document.add_paragraph("Acc_P16 + (Acc_P - Acc_N) 4")
            document.add_paragraph("Covers: REQ-A")
            document.add_paragraph("[End]")
            document.save(path)
            statement = "Acc_P*16 + (Acc_P - Acc_N) * 4"
            apply_docx_ipos_requirement_style_guard(path, authored_requirements={"IPOS-MODULE-001": statement})
            self.assertEqual([paragraph.text for paragraph in Document(path).paragraphs],
                             ["Description stays unchanged.", "IPOS-MODULE-001", statement, "Covers: REQ-A", "[End]"])
            with self.assertRaises(ValueError):
                apply_docx_ipos_requirement_style_guard(path, authored_requirements={"IPOS-MODULE-002": statement})

    def test_saved_omit_binding_preserves_all_local_requirement_ids_and_statements(self):
        from workflow_routing import build_ipos_functional_input

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            blocker = self.blocker()
            blocker["Section or claim"] = "drop responsibility"
            path = review.write_blocker_review(root, "snapshot", "Module", [blocker])
            workbook = load_workbook(path)
            workbook.active.cell(2, 10).value = "omit (as non-blocking)"
            workbook.save(path)
            workbook.close()
            rows = [{"mapped_block": "Module", "source_req_id": "REQ-A", "requirement_statement": "Exact normative A."},
                    {"mapped_block": "Module", "source_req_id": "REQ-B", "requirement_statement": "Exact normative B."}]
            functional = build_ipos_functional_input("Module", {"Function": "Control local operation."}, rows, domain="digital")
            plan = SimpleNamespace(claims=(SimpleNamespace(claim_id="drop", responsibilities=("drop responsibility",)),
                                           SimpleNamespace(claim_id="keep", responsibilities=("keep responsibility",))))
            with patch("workflow_routing.build_ipos_claim_plan", return_value=plan):
                applied = review.apply_descriptive_omissions(root, snapshot_id="snapshot", functional_input=functional)
            self.assertEqual(applied.descriptive_omitted_claim_ids, ("drop",))
            self.assertEqual(applied.local_requirement_ids, ("REQ-A", "REQ-B"))
            self.assertEqual(applied.local_requirement_evidence, ("Exact normative A.", "Exact normative B."))
            self.assertEqual(applied.candidate_requirement_records, functional.candidate_requirement_records)

    def test_archived_omit_receipt_remains_bound_and_rejects_tampered_workbook(self):
        from workflow_routing import build_ipos_functional_input

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            blocker = self.blocker()
            blocker["Section or claim"] = "drop responsibility"
            path = review.write_blocker_review(root, "snapshot", "Module", [blocker])
            workbook = load_workbook(path)
            workbook.active.cell(2, 10).value = "omit (as non-blocking)"
            workbook.save(path)
            workbook.close()
            functional = build_ipos_functional_input("Module", {"Function": "Control local operation."}, domain="digital")
            plan = SimpleNamespace(claims=(SimpleNamespace(claim_id="drop", responsibilities=("drop responsibility",)),))
            with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", return_value=functional), \
                    patch("workflow_routing.build_ipos_claim_plan", return_value=plan), \
                    patch("workflow_routing.ipos_claim_preflight_findings", return_value=[]):
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module", rows=[],
                                                 candidate_rows=[], writer=lambda *args, **kwargs: None)
                self.assertEqual(result["status"], "READY")
                self.assertFalse(path.exists())
                receipt = json.loads(path.with_name("descriptive_omissions.json").read_text(encoding="utf-8"))
                self.assertEqual(receipt["writing_rule"], "block_local_main_functions_v1")
                applied = review.apply_descriptive_omissions(root, snapshot_id="snapshot", functional_input=functional)
                self.assertEqual(applied.descriptive_omitted_claim_ids, ("drop",))
                receipt_bytes = path.with_name("descriptive_omissions.json").read_bytes()
                for decision in ("user clarified", "need source/spec update - clarification"):
                    review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
                    workbook = load_workbook(path)
                    workbook.active.cell(2, 10).value = decision
                    workbook.active.cell(2, 11).value = "Existing governed evidence only; no new fact."
                    workbook.save(path)
                    workbook.close()
                    result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module",
                                                     rows=[], candidate_rows=[], writer=lambda *args, **kwargs: None)
                    self.assertEqual(result["generation_status"], "GENERATABLE")
                    self.assertTrue(path.exists())
                    self.assertEqual(path.with_name("descriptive_omissions.json").read_bytes(), receipt_bytes)
                    applied = review.apply_descriptive_omissions(root, snapshot_id="snapshot", functional_input=functional)
                    self.assertEqual(applied.descriptive_omitted_claim_ids, ("drop",))
                with patch("workflow_routing.build_ipos_claim_plan", return_value=SimpleNamespace(claims=())):
                    with self.assertRaises(ValueError):
                        review.apply_descriptive_omissions(root, snapshot_id="snapshot", functional_input=functional)
                archived = next(path.parent.glob("history/*/blocker_review.xlsx"))
                workbook = load_workbook(archived)
                workbook.active.cell(2, 11).value = "Changed after publication."
                workbook.save(archived)
                workbook.close()
                with self.assertRaises(ValueError):
                    review.apply_descriptive_omissions(root, snapshot_id="snapshot", functional_input=functional)

    def test_omit_excludes_only_descriptive_contribution(self):
        from workflow_routing import DescriptiveSectionContribution, ipos_claim_section_plan

        plan = SimpleNamespace(claims=(SimpleNamespace(claim_id="omit-item"), SimpleNamespace(claim_id="keep-item")))
        contribution = DescriptiveSectionContribution("omit-item", "functions", "retained", (), (), "supported")
        kept = DescriptiveSectionContribution("keep-item", "functions", "retained", (), (), "supported")
        with patch("workflow_routing.descriptive_section_contributions", return_value=(contribution, kept)):
            result = ipos_claim_section_plan(plan, omitted_claim_ids=("omit-item",))
            self.assertEqual(result[0].decision, "excluded")
            self.assertEqual(result[0].reason, "user_omit_descriptive_only")
            self.assertEqual(result[1], kept)
            self.assertEqual(len(plan.claims), 2)
            with self.assertRaises(ValueError):
                ipos_claim_section_plan(plan, omitted_claim_ids=("unknown-item",))

    def test_explicit_omit_accepts_saved_decision_without_clarification_note(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            workbook = load_workbook(path)
            workbook.active.cell(2, 10).value = "omit (as non-blocking)"
            workbook.save(path)
            workbook.close()
            context = review.read_blocker_review(root, snapshot_id="snapshot", block="Module")
            self.assertEqual(context["decisions"][0]["User decision"], "omit (as non-blocking)")
            self.assertEqual(context["notes"], [])

    def test_user_decision_dropdown_has_exact_flow_choices(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            workbook = load_workbook(path)
            sheet = workbook.active
            validation = sheet.data_validations.dataValidation[0]
            self.assertEqual(validation.type, "list")
            self.assertEqual(validation.formula1, '"' + ",".join(review.REVIEW_CHOICES) + '"')
            self.assertEqual(str(validation.sqref), "J2")
            self.assertFalse(validation.showDropDown)
            self.assertTrue(validation.showErrorMessage)
            self.assertEqual(validation.errorStyle, "stop")
            self.assertFalse(sheet.cell(2, 10).protection.locked)
            self.assertFalse(sheet.cell(2, 11).protection.locked)
            self.assertTrue(sheet.cell(2, 8).protection.locked)
            self.assertTrue(sheet.cell(2, 12).protection.locked)
            for choice in review.REVIEW_CHOICES:
                sheet.cell(2, 10).value = choice
                sheet.cell(2, 11).value = "Review context, not approval."
                workbook.save(path)
                context = review.read_blocker_review(root, snapshot_id="snapshot", block="Module")
                self.assertEqual(context["decisions"][0]["User decision"], choice)
                self.assertEqual(bool(context["future_source_updates"]), choice == "need source/spec update - clarification")
            workbook.close()

    def test_selected_source_update_records_gap_without_generation_veto(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            workbook = load_workbook(path)
            workbook.active.cell(2, 10).value = "need source/spec update - clarification"
            workbook.save(path)
            workbook.close()
            with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", return_value=SimpleNamespace(block="Module")), \
                    patch("workflow_routing.ipos_claim_preflight_findings", return_value=[]) as check:
                writer = Mock()
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module", rows=[], candidate_rows=[], writer=writer)
            self.assertEqual(result["status"], "READY")
            self.assertEqual(result["review_context"]["future_source_updates"][0]["Resulting action"],
                             "generate supported content; retain source/spec update gap for readiness")
            check.assert_called_once()
            writer.assert_called_once()

    def test_dropdown_upgrade_preserves_notes_and_does_not_count_as_user_save(self):
        for reviewed in (False, True):
            with self.subTest(reviewed=reviewed), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
                workbook = load_workbook(path)
                workbook.active.data_validations.dataValidation = []
                workbook.active.cell(2, 10).protection = review.Protection(locked=True)
                workbook.save(path)
                workbook.close()
                manifest_path = review.review_paths(root, "Module")[1]
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                manifest["initial_sha256"] = review._digest(path)
                manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
                if reviewed:
                    workbook = load_workbook(path)
                    workbook.active.cell(2, 11).value = "Keep this existing user note."
                    workbook.save(path)
                    workbook.close()
                review.upgrade_blocker_review_dropdown(root, snapshot_id="snapshot", block="Module")
                workbook = load_workbook(path)
                self.assertEqual(len(workbook.active.data_validations.dataValidation), 1)
                self.assertFalse(workbook.active.cell(2, 10).protection.locked)
                self.assertEqual(workbook.active.cell(2, 11).value, "Keep this existing user note." if reviewed else None)
                workbook.close()
                if reviewed:
                    self.assertEqual(len(review.read_blocker_review(root, snapshot_id="snapshot", block="Module")["notes"]), 1)
                else:
                    with self.assertRaises(ValueError):
                        review.read_blocker_review(root, snapshot_id="snapshot", block="Module")

    def test_unsaved_or_open_review_holds_only_its_block(self):
        for opened in (False, True):
            with self.subTest(opened=opened), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                path = review.write_blocker_review(root, "snapshot", "First", [self.blocker("First")])
                if opened:
                    path.with_name("~$" + path.name).touch()
                with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", side_effect=lambda block, *args, **kwargs: SimpleNamespace(block=block)), \
                        patch("workflow_routing.ipos_claim_preflight_findings", return_value=[]) as check, \
                        patch("generate_ipos_specs._write_ipos_document") as writer, \
                        patch.object(review, "emit_block_review") as emit:
                    results = review.assess_ipos_blocks(root, snapshot_id="snapshot", kind="digital",
                                                       blocks=["First", "Second"], rows=[], candidate_rows=[])
                self.assertEqual([value["status"] for value in results], ["BLOCKED" if opened else "READY", "READY"])
                if opened:
                    self.assertIn("waiting_reason", results[0])
                self.assertEqual(check.call_count, 1 if opened else 2)
                self.assertEqual(writer.call_count, 1 if opened else 2)
                self.assertEqual(writer.call_args.args[5], "Second")
                self.assertEqual(emit.call_count, 2)

    def test_omit_does_not_waive_unrelated_checker_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            blocker = self.blocker()
            path = review.write_blocker_review(root, "snapshot", "Module", [blocker])
            workbook = load_workbook(path)
            workbook.active.cell(2, 10).value = "omit (as non-blocking)"
            workbook.active.cell(2, 11).value = "Omit selected description only; other failures remain blocking."
            workbook.save(path)
            workbook.close()
            with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", return_value=SimpleNamespace()), \
                    patch("workflow_routing.ipos_claim_preflight_findings", return_value=["still_active"]) as check, \
                    patch.object(review, "apply_descriptive_omissions", side_effect=lambda root, **kwargs: kwargs["functional_input"]), \
                    patch.object(review, "blocker_rows", return_value=[blocker]):
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module", rows=[], candidate_rows=[])
            self.assertEqual(result["status"], "BLOCKED")
            check.assert_called_once()

    def test_future_source_update_keeps_residual_review_and_runs_supported_generation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            blocker = self.blocker()
            blocker.update({"Flow recommendation": "need source/spec update - clarification",
                            "Resulting action": "remain blocked pending future source/spec update"})
            path = review.write_blocker_review(root, "snapshot", "Module", [blocker])
            workbook = load_workbook(path)
            workbook.active.cell(2, 11).value = "Detail for a future source update."
            workbook.save(path)
            workbook.close()
            with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", return_value=SimpleNamespace(block="Module")), \
                    patch("workflow_routing.ipos_claim_preflight_findings", return_value=[]) as check:
                writer = Mock()
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module", rows=[], candidate_rows=[], writer=writer)
            self.assertEqual(result["status"], "READY")
            self.assertEqual(len(result["review_context"]["future_source_updates"]), 1)
            check.assert_called_once()
            writer.assert_called_once()
            self.assertTrue(path.exists())

    def test_cleared_checker_retires_review_without_note_approval_authority(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            workbook = load_workbook(path)
            workbook.active.cell(2, 11).value = "Inspect existing supported meaning."
            workbook.save(path)
            workbook.close()
            with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", return_value=SimpleNamespace(block="Module")), \
                    patch("workflow_routing.ipos_claim_preflight_findings", return_value=[]):
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module",
                                                 rows=[], candidate_rows=[], writer=lambda *args, **kwargs: None)
            self.assertEqual(result["status"], "READY")
            self.assertFalse(path.exists())
            self.assertTrue(list(path.parent.glob("history/*/blocker_review.xlsx")))

    def test_generator_publishes_ready_block_and_keeps_blocked_directory(self):
        import generate_ipos_specs as ipos

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "config").mkdir()
            (root / "config/project_context.json").write_text('{"project_name":"test"}', encoding="utf-8")
            source_rows = [{"source_req_id": "REQ-A", "canonical_id": "A", "statement": "Approved A"},
                           {"source_req_id": "REQ-B", "canonical_id": "B", "statement": "Approved B"}]
            allocation = {"REQ-A": {"allocation_class": "block_local_digital", "approved_block": "Module A"},
                          "REQ-B": {"allocation_class": "block_local_digital", "approved_block": "Module B"}}
            snapshot = SimpleNamespace(snapshot_id="snapshot", rows=source_rows, source_path=root / "canonical.sqlite")
            contract = SimpleNamespace(allocation=lambda identity: allocation[identity], descriptive_context=lambda *args, **kwargs: ())
            ledger = [{"req_id": identity, "hierarchy_or_coverage_class": "block_local_digital", "owning_block": metadata["approved_block"]}
                      for identity, metadata in allocation.items()]
            block_root = root / ipos.CONFIG["digital"]["output_dir"] / "blocks"
            blocked = block_root / "module-a"
            blocked.mkdir(parents=True)
            published = blocked / "digital_ipos_module-a.md"
            published.write_text("Historical blocked document", encoding="utf-8")
            states = [{"block": "Module A", "status": "BLOCKED", "review_path": "a.xlsx", "blocker_count": 1},
                      {"block": "Module B", "status": "READY", "blocker_count": 0}]
            with patch.object(ipos, "resolve_complete_authoritative_input", return_value=(snapshot, {})), \
                    patch.object(ipos, "resolve_downstream_contract", return_value=contract), \
                    patch.object(ipos, "validate_generation_inputs", return_value={"decision": "PASS"}), \
                    patch.object(ipos, "refresh_ledger"), patch.object(ipos, "_read_csv", return_value=[]), \
                    patch.object(ipos, "read_csv", return_value=ledger), \
                    patch.object(ipos, "_inventory_blocks", return_value=["Module A", "Module B"]), \
                    patch.object(ipos, "approved_snapshot_candidate_rows", return_value=[]), \
                    patch.object(ipos, "_require_ipos_claim_preflights"), \
                    patch.object(ipos, "_write_ipos_document") as publish, \
                    patch.object(review, "assess_ipos_blocks", return_value=states):
                trace = ipos.generate(root, "digital", snapshot_id="snapshot")
            publish.assert_called_once()
            self.assertEqual(publish.call_args.args[5], "Module B")
            self.assertEqual(published.read_text(encoding="utf-8"), "Historical blocked document")
            self.assertIn("REQ-A", trace.read_text(encoding="utf-8"))
            self.assertIn("REQ-B", trace.read_text(encoding="utf-8"))
            result = json.loads((root / "artifacts/orchestrator/ipos_generation_results.json").read_text(encoding="utf-8"))
            self.assertEqual(result["status"], "BLOCKED")
            self.assertEqual([value["status"] for value in result["blocks"]], ["BLOCKED", "PASS"])

    def test_gui_auto_opens_dedicated_excel_and_checks_save_close_before_recheck(self):
        import tkinter as tk
        from tkinter import ttk
        from unittest.mock import MagicMock
        from workflow_gui import WorkflowGUI

        window = tk.Tk()
        window.withdraw()
        try:
            with tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
                gui = object.__new__(WorkflowGUI)
                gui.root = window
                gui.is_running = False
                for name in ("append_log", "refresh_work_area", "_on_work_area_select_path", "_run_command_async"):
                    setattr(gui, name, MagicMock())
                with patch("workflow_gui.REPO_ROOT", root), patch("workflow_gui.messagebox.showwarning") as warning:
                    gui._show_ipos_blocker_review({"block": "Module", "snapshot_id": "snapshot", "path": str(path), "blocker_count": 1})
                    gui._on_work_area_select_path.assert_called_once_with(path.resolve())
                    dialog = gui._ipos_review_dialogs["Module"]
                    self.assertIsNone(window.grab_current())
                    def descendants(widget):
                        for child in widget.winfo_children():
                            yield child
                            yield from descendants(child)
                    button = next(child for child in descendants(dialog) if isinstance(child, ttk.Button)
                                  and child.cget("text") == "Re-evaluate (saved + closed)")
                    button.invoke()
                    warning.assert_called_once()
                    gui._run_command_async.assert_not_called()
                    workbook = load_workbook(path)
                    workbook.active.cell(2, 11).value = "Recheck existing source support."
                    workbook.save(path)
                    workbook.close()
                    button.invoke()
                    command = gui._run_command_async.call_args.args[0]
                    self.assertEqual(command[command.index("--block") + 1], "Module")
                    self.assertEqual(command[command.index("--snapshot-id") + 1], "snapshot")
                    review.write_blocker_review(root, "next-snapshot", "Module", [self.blocker()])
                    gui._show_ipos_blocker_review({"block": "Module", "snapshot_id": "next-snapshot",
                                                  "path": str(path), "blocker_count": 1})
                    current_dialog = gui._ipos_review_dialogs["Module"]
                    self.assertIsNot(current_dialog, dialog)
                    self.assertFalse(dialog.winfo_exists())
                    self.assertIsNone(window.grab_current())
                    workbook = load_workbook(path)
                    workbook.active.cell(2, 11).value = "Recheck only the current snapshot."
                    workbook.save(path)
                    workbook.close()
                    current_button = next(child for child in descendants(current_dialog) if isinstance(child, ttk.Button)
                                          and child.cget("text") == "Re-evaluate (saved + closed)")
                    current_button.invoke()
                    command = gui._run_command_async.call_args.args[0]
                    self.assertEqual(command[command.index("--block") + 1], "Module")
                    self.assertEqual(command[command.index("--snapshot-id") + 1], "next-snapshot")
        finally:
            window.destroy()

    def test_failed_publication_restores_existing_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "artifacts/orchestrator").mkdir(parents=True)
            output = root / "output"
            output.mkdir()
            markdown = output / "digital_ipos_module.md"
            markdown.write_text("Existing published document", encoding="utf-8")
            def failing_writer(*args, **kwargs):
                markdown.write_text("Rejected candidate", encoding="utf-8")
                (output / "document_contract_report.json").write_text("{}", encoding="utf-8")
                raise RuntimeError("Publication failure")
            with self.assertRaises(RuntimeError):
                review.publish_ipos_document(failing_writer, root, {"docx": "module.docx"}, "digital", "Title", output, "Module")
            self.assertEqual(markdown.read_text(encoding="utf-8"), "Existing published document")
            self.assertFalse((output / "document_contract_report.json").exists())

    def test_blocked_block_does_not_prevent_other_block_assessment(self):
        states = [{"block": "First", "status": "BLOCKED", "review_path": "first.xlsx", "blocker_count": 1},
                  {"block": "Second", "status": "READY", "blocker_count": 0}]
        with patch.object(review, "assess_ipos_block", side_effect=states) as check, \
                patch.object(review, "emit_block_review") as emit:
            results = review.assess_ipos_blocks(Path("."), snapshot_id="snapshot", kind="digital", blocks=["First", "Second"],
                                               rows=[{"mapped_block": "First"}, {"mapped_block": "Second"}], candidate_rows=[])
        self.assertEqual([value["status"] for value in results], ["BLOCKED", "READY"])
        self.assertEqual(check.call_count, 2)
        self.assertEqual(len(check.call_args.kwargs["rows"]), 1)
        self.assertEqual(emit.call_count, 2)

    def test_saved_note_rechecks_but_never_bypasses_current_blocker(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            functional = SimpleNamespace(block="Module", local_requirement_ids=())
            with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", return_value=functional), \
                    patch("workflow_routing.ipos_claim_preflight_findings", return_value=["current_failure"]) as check, \
                    patch.object(review, "blocker_rows", return_value=[self.blocker()]) as project:
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module", rows=[], candidate_rows=[])
                self.assertEqual(result["status"], "BLOCKED")
                check.assert_called_once()
                workbook = load_workbook(path)
                workbook.active.cell(2, 11).value = "Check this approved relationship again."
                workbook.save(path)
                workbook.close()
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module", rows=[], candidate_rows=[])
                self.assertEqual(result["status"], "BLOCKED")
                self.assertEqual(check.call_count, 2)
                self.assertEqual(project.call_args.kwargs["user_notes"][0]["Note"], "Check this approved relationship again.")

    def test_ready_block_has_no_excel_or_review_interruption(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch("generate_ipos_specs._block_inventory_record", return_value={}), \
                    patch("workflow_routing.build_ipos_functional_input", return_value=SimpleNamespace(block="Module")), \
                    patch("workflow_routing.ipos_claim_preflight_findings", return_value=[]):
                result = review.assess_ipos_block(root, snapshot_id="snapshot", kind="digital", block="Module",
                                                 rows=[], candidate_rows=[], writer=lambda *args, **kwargs: None)
            self.assertEqual(result["status"], "READY")
            self.assertFalse(review.review_paths(root, "Module")[0].exists())

    def blocker(self, block="Module"):
        return {"Block": block, "Section or claim": "1.4 Functions", "Failure type": "Readback mismatch",
                "Expected supported meaning": "Supported relationship", "Observed problem": "Not localized",
                "Likely failing layer": "realization/readback; not localized",
                "Flow recommendation": "user clarified", "Why": "Review existing evidence; no new authority",
                "Resulting action": "re-evaluate after user note"}

    def test_exact_columns_one_block_notes_and_save_close(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            with self.assertRaises(ValueError):
                review.read_blocker_review(root, snapshot_id="snapshot", block="Module")
            workbook = load_workbook(path)
            self.assertEqual(tuple(cell.value for cell in workbook.active[1]), review.REVIEW_COLUMNS)
            self.assertIsNone(workbook.active.cell(2, 1).value)
            workbook.active.cell(2, 11).value = "Check the existing supported relationship; do not infer text."
            workbook.save(path)
            workbook.close()
            result = review.read_blocker_review(root, snapshot_id="snapshot", block="Module")
            self.assertEqual(len(result["notes"]), 1)
            self.assertFalse(result["authority"])
            path.with_name("~$" + path.name).touch()
            with self.assertRaises(ValueError):
                review.read_blocker_review(root, snapshot_id="snapshot", block="Module")

    def test_no_mixed_blocks_no_omit_approval_and_future_source_update(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaises(ValueError):
                review.write_blocker_review(root, "snapshot", "Module", [self.blocker("Other")])
            item = self.blocker()
            item.update({"Flow recommendation": "need source/spec update - clarification",
                         "Resulting action": "remain blocked pending future source/spec update"})
            path = review.write_blocker_review(root, "snapshot", "Module", [item])
            workbook = load_workbook(path)
            workbook.active.cell(2, 10).value = "omit"
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError):
                review.read_blocker_review(root, snapshot_id="snapshot", block="Module")

    def test_identical_blockers_preserve_user_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = review.write_blocker_review(root, "snapshot", "Module", [self.blocker()])
            workbook = load_workbook(path)
            workbook.active.cell(2, 11).value = "Existing review note"
            workbook.save(path)
            workbook.close()
            before = path.read_bytes()
            self.assertEqual(review.write_blocker_review(root, "snapshot", "Module", [self.blocker()]), path)
            self.assertEqual(path.read_bytes(), before)