import sys
import inspect
import io
import tempfile
import tkinter as tk
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from tkinter import ttk
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from report_generation_service import _ordered_fields, _style_data_sheet
from sysml_generation_service import SysMLCompletenessError, validate_snapshot_requirement_coverage
from workflow_gui import (
    SINGLE_STAGE_LABELS,
    SINGLE_STAGE_LABEL_TO_KEY,
    RANGE_STAGE_LABELS,
    RANGE_STAGE_LABEL_TO_KEY,
    WorkflowGUI,
    WorkflowDiagramRenderer,
    WORKFLOW_NODE_MEANINGS,
    format_gui_review_context,
    build_gui_review_context,
    WORKFLOW_DIAGRAM_NODES,
    WORKFLOW_DIAGRAM_EDGES,
    compact_workflow_scrollregion,
    compact_workflow_node_y,
    responsive_layout_scale,
    scale_window_geometry,
)
from workflow_cli import STAGE_ORDER, build_parser, cmd_validate
from generate_stage2_specs import (
    _cascade_previous_concrete_owners,
    _ensure_previous_concrete_owner_defs,
    _is_generic_source_paragraph,
    _preview_blocks_for_requirement,
    _read_previous_blocks_by_requirement,
    _read_profile_source_section_aliases,
    _read_source_matrix_edges,
    _write_mapping_preview_markdown,
)
from generate_stage1_requirements import _source_section_owner
from generate_architecture_sysml import is_concrete_block
from run_stage2_micro_arch_and_crosscheck import (
    _ensure_interaction_blocks,
    _entity_kind,
    _is_concrete_block,
    _link_interaction_requirements_to_blocks,
    _map_blocks,
    _explicit_target_block,
)
from run_stage6_digital_ipos_gate import _run_activity
from generate_ipos_specs import _convert_ipos_markdown_to_docx


class FinalPolishTests(unittest.TestCase):
    def test_final_sysml_block_dialog_opens_only_run_bound_excel(self):
        import hashlib
        import json

        root = tk.Tk()
        root.withdraw()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.root = root
            gui._set_responsive_geometry = MagicMock()
            gui._open_xlsx_in_excel = MagicMock()

            def descendants(widget):
                for child in widget.winfo_children():
                    yield child
                    yield from descendants(child)

            with tempfile.TemporaryDirectory() as temporary:
                repo_root = Path(temporary)
                folder = repo_root / "artifacts/validation/readiness_checklists"
                folder.mkdir(parents=True)
                report = folder / "readiness_summary.xlsx"
                report.write_bytes(b"Excel opener fixture")
                summary = folder / "readiness_summary.json"
                summary.write_text(json.dumps({"binding": {"snapshot_id": "selected"}, "status": "NOT_READY", "publication_status": "BLOCKED"}), encoding="utf-8")
                receipt = {"snapshot_id": "selected", "summary_json_sha256": hashlib.sha256(summary.read_bytes()).hexdigest(),
                           "summary_excel_sha256": hashlib.sha256(report.read_bytes()).hexdigest()}
                command = ["python", "run_final_sysml_phase.py", "--snapshot-id", "selected"]
                with patch("workflow_gui.REPO_ROOT", repo_root), patch("validate_downstream_coherence.validate") as validate, \
                        patch("workflow_gui.messagebox.showerror") as showerror:
                    gui._show_downstream_block_dialog(["Final SysML central coherence validation: FAIL", "[SYSML_READINESS] " + json.dumps(receipt)], command)
                    dialog = next(child for child in root.winfo_children() if isinstance(child, tk.Toplevel))
                    text = next(child for child in descendants(dialog) if isinstance(child, tk.Text)).get("1.0", "end-1c")
                    self.assertIn("Final SysML execution was blocked", text)
                    self.assertIn("readiness_summary.xlsx", text)
                    self.assertNotIn("cannot continue to Stage 3", text)
                    self.assertNotIn("downstream_coherence_report.json", text)
                    button = next(child for child in descendants(dialog) if isinstance(child, ttk.Button) and child.cget("text") == "Open Excel summary")
                    button.invoke()
                    gui._open_xlsx_in_excel.assert_called_once_with(report)
                    report.write_bytes(b"changed workbook")
                    button.invoke()
                    showerror.assert_called_once()
                    self.assertEqual(gui._open_xlsx_in_excel.call_count, 1)
                    dialog.destroy()
                    gui._show_downstream_block_dialog(["BLOCK_DOWNSTREAM"], command)
                    dialog = next(child for child in root.winfo_children() if isinstance(child, tk.Toplevel))
                    button = next(child for child in descendants(dialog) if isinstance(child, ttk.Button) and child.cget("text") == "Open Excel summary")
                    self.assertTrue(button.instate(["disabled"]))
                    dialog.destroy()
                    receipt["snapshot_id"] = "other"
                    gui._show_downstream_block_dialog(["[SYSML_READINESS] " + json.dumps(receipt)], command)
                    dialog = next(child for child in root.winfo_children() if isinstance(child, tk.Toplevel))
                    button = next(child for child in descendants(dialog) if isinstance(child, ttk.Button) and child.cget("text") == "Open Excel summary")
                    self.assertTrue(button.instate(["disabled"]))
                    validate.assert_not_called()
        finally:
            root.destroy()

    def test_sysml_generation_button_is_visible_and_uses_final_pipeline(self):
        window = tk.Tk()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.root = window
            gui.sysml_architecture_tab = ttk.Frame(window)
            gui.sysml_architecture_tab.pack(fill=tk.BOTH, expand=True)
            gui._bind_button_help = MagicMock()
            gui._bind_copy_support = MagicMock()
            gui._refresh_sysml_architecture = MagicMock()
            gui._snapshot_selector_args = MagicMock(return_value=["--snapshot-id", "selected"])
            gui._run_command_async = MagicMock()
            gui._build_sysml_architecture_tab()
            for geometry in ("1100x700", "640x700"):
                window.geometry(geometry)
                window.update()
                button = gui.sysml_generate_btn
                self.assertTrue(button.winfo_ismapped())
                self.assertGreaterEqual(button.winfo_width(), button.winfo_reqwidth())
            gui.sysml_generate_btn.invoke()
            command = gui._run_command_async.call_args.args[0]
            self.assertEqual(Path(command[1]).name, "run_final_sysml_phase.py")
            self.assertEqual(command[-2:], ["--snapshot-id", "selected"])
        finally:
            window.destroy()

    def test_sysml_readiness_excel_log_link_opens_excel(self):
        window = tk.Tk()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.log = tk.Text(window)
            gui.log.tag_bind = MagicMock(wraps=gui.log.tag_bind)
            gui._open_xlsx_in_excel = MagicMock()
            gui._open_run_artifact = MagicMock()
            path = Path(__file__).resolve().parent / "readiness summary.xlsx"
            gui._append_log_ui("[REPORT] Open SysML readiness Excel: " + str(path))
            click = next(call.args[2] for call in gui.log.tag_bind.call_args_list if call.args[1] == "<Button-1>")
            click(None)
            gui._open_xlsx_in_excel.assert_called_once_with(path)
            gui._open_run_artifact.assert_not_called()
            self.assertIn(str(path), gui.log.get("1.0", tk.END))
        finally:
            window.destroy()

    def test_sysml_generation_action_uses_bound_snapshot_and_final_pipeline(self):
        gui = object.__new__(WorkflowGUI)
        gui._snapshot_selector_args = MagicMock(return_value=["--snapshot-id", "selected"])
        gui._run_command_async = MagicMock()
        gui.regenerate_sysml()
        command, label = gui._run_command_async.call_args.args
        self.assertEqual(Path(command[1]).name, "run_final_sysml_phase.py")
        self.assertEqual(command[-2:], ["--snapshot-id", "selected"])
        self.assertIn("readiness checklist", label)
        self.assertIsNotNone(gui._run_command_async.call_args.kwargs["on_success"])
        gui._snapshot_selector_args.return_value = None
        gui._run_command_async.reset_mock()
        gui.regenerate_sysml()
        gui._run_command_async.assert_not_called()

    def test_sysml_navigation_retains_stale_scope_without_weakening_validation(self):
        from validate_downstream_coherence import DownstreamCoherenceError, SysMLMaterializationScope

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            model = root / "Block.sysml"
            model.write_text("// Snapshot ID: previous\npart def Block {}", encoding="utf-8")
            scope = SimpleNamespace(findings=("SYSML_SCOPE_STALE_IPOS: Block",),
                                    existing_model_paths=lambda _root: ([model], ["SYSML_SCOPE_STALE_IPOS: Block"]),
                                    structural_paths=lambda: ())
            gui = object.__new__(WorkflowGUI)
            gui.workflow_snapshot_id = "current"
            with patch("workflow_gui.resolve_sysml_materialization_scope", return_value=scope):
                self.assertIs(gui._get_sysml_materialization_scope(), scope)
                self.assertEqual(gui._concrete_sysml_block_paths(), [model])
                self.assertIn("SYSML_SCOPE_STALE_IPOS", gui.sysml_scope_diagnostic)
                self.assertIn("SYSML_NAVIGATION_STALE_MODEL", gui.sysml_scope_diagnostic)
                self.assertIn("historical/unverified", gui._sysml_navigation_label(model, "Block"))
                self.assertIn("Model: previous | Workflow: current", gui._sysml_navigation_label(model, "Block"))
            with self.assertRaises(DownstreamCoherenceError):
                SysMLMaterializationScope.require_valid(scope)

    def test_sysml_navigation_lists_existing_files_when_scope_resolution_fails(self):
        from validate_downstream_coherence import DownstreamCoherenceError

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / "artifacts/stage2_mirco_arc/sysml/blocks"
            folder.mkdir(parents=True)
            model = folder / "block.sysml"
            model.write_text("part def Block {}", encoding="utf-8")
            gui = object.__new__(WorkflowGUI)
            gui.workflow_snapshot_id = "current"
            gui.sysml_hierarchy_tree = MagicMock()
            gui.sysml_hierarchy_tree.get_children.return_value = ()
            with patch("workflow_gui.REPO_ROOT", root), patch("workflow_gui.resolve_sysml_materialization_scope", side_effect=DownstreamCoherenceError("scope unavailable")):
                self.assertEqual(gui._concrete_sysml_block_paths(), [model])
                gui._populate_sysml_hierarchy([model])
            self.assertTrue(any(call.kwargs.get("values") == (model.as_posix(),)
                                for call in gui.sysml_hierarchy_tree.insert.call_args_list))
            self.assertIn("scope unavailable", gui.sysml_scope_diagnostic)

    def test_system_traceability_top_menu_syncs_selection_and_centers(self):
        window = tk.Tk()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.root = window
            gui.layout_scale = 1.0
            gui.traceability_selected_block = None
            gui.traceability_selected_edge = None
            gui._bind_button_help = MagicMock()
            gui._on_traceability_window_resize = MagicMock()
            gui._on_layout_window_resize = MagicMock()
            graph = {"nodes": [{"id": "DRS", "kind": "document", "label": "DRS"},
                               {"id": "Digital_IPOS:main-controller", "kind": "ipos_block", "label": "Main Controller", "document_qualification": "Digital IPOS"}],
                     "edges": [{"source": "DRS", "target": "Digital_IPOS:main-controller", "coverage": {"covered": 2, "total": 3}}]}
            with patch("workflow_gui.approved_generated_document_dependency_graph", return_value=graph):
                gui._open_system_traceability_window()
                window.update()
                choice = next(label for label, node_id in gui.traceability_block_choices.items() if node_id == "Digital_IPOS:main-controller")
                gui.traceability_block_choice.set(choice)
                gui.traceability_block_choice.event_generate("<<ComboboxSelected>>")
                window.update()
                self.assertEqual(gui.traceability_selected_block, "Digital_IPOS:main-controller")
                self.assertEqual(gui.traceability_tree.selection(), ("Digital_IPOS:main-controller",))
                canvas = gui.traceability_graph_canvas
                bounds = canvas.bbox("traceability_node:Digital_IPOS:main-controller")
                self.assertAlmostEqual(canvas.canvasx(canvas.winfo_width() / 2), (bounds[0] + bounds[2]) / 2, delta=2)
                self.assertAlmostEqual(canvas.canvasy(canvas.winfo_height() / 2), (bounds[1] + bounds[3]) / 2, delta=2)
                self.assertTrue(gui.traceability_block_choice.winfo_ismapped())
                gui.traceability_tree.selection_set("DRS")
                gui._on_traceability_tree_select()
                self.assertEqual(gui.traceability_block_choice_var.get(), "DRS")
                gui.traceability_block_choice.set("All blocks")
                gui._on_traceability_block_choice()
                self.assertIsNone(gui.traceability_selected_block)
                gui.traceability_block_choice.set("DRS")
                gui._on_traceability_block_choice()
                rectangle = next(item for item in canvas.find_withtag("traceability_node:DRS") if canvas.type(item) == "rectangle")
                self.assertEqual(canvas.itemcget(rectangle, "fill"), "#cfe8ff")
        finally:
            window.destroy()

    def test_execution_log_horizontal_scrollbar_stays_visible_when_shrunk(self):
        window = tk.Tk()
        try:
            with patch.object(WorkflowGUI, "_load_gui_state", return_value={"layout_saved_by_user": True, "last_user_layout": {}}), \
                    patch.object(WorkflowGUI, "_restore_gui_layout"), patch.object(WorkflowGUI, "_save_gui_state"):
                gui = WorkflowGUI(window)
                window.geometry("1100x700")
                window.update()
                gui.log.insert(tk.END, "wide-log-entry " * 200)
                pane = gui.main_vertical_pane
                for height in (240, 160, 115, 240):
                    pane.sashpos(0, 120)
                    pane.sashpos(1, pane.winfo_height() - height)
                    window.update()
                    scrollbar = gui.log_x_scroll
                    self.assertTrue(scrollbar.winfo_ismapped(), height)
                    self.assertGreaterEqual(scrollbar.winfo_height(), scrollbar.winfo_reqheight(), height)
                    self.assertGreater(scrollbar.winfo_width(), 100)
                    self.assertLessEqual(scrollbar.winfo_y() + scrollbar.winfo_height(), scrollbar.master.winfo_height())
                    gui.log.xview_moveto(1.0)
                    window.update_idletasks()
                    self.assertGreater(gui.log.xview()[0], 0)
        finally:
            window.destroy()

    def test_sysml_hierarchy_spreads_horizontally_and_caps_overview_ports(self):
        window = tk.Tk()
        try:
            with tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                system, digital = root / "System.sysml", root / "Digital.sysml"
                system.write_text("part def System { part topDigital : DigitalSubsystem; connection link : Link connect topDigital.block0 to topDigital.block1; }", encoding="utf-8")
                digital.write_text("part def DigitalSubsystem {" + " ".join(f"part block{index} : Block{index};" for index in range(8)) + "}", encoding="utf-8")
                blocks = []
                for index in range(8):
                    path = root / f"Block{index}.sysml"
                    path.write_text("part def Block {" + " ".join(f"in port data{port_index} : Signal;" for port_index in range(100)) + "}", encoding="utf-8")
                    blocks.append(SimpleNamespace(path=path, node=f"topDigital.block{index}", target="Digital IPOS", definition=f"Block{index}"))
                scope = SimpleNamespace(system_path=system, digital_path=digital, system_definition="System", blocks=blocks,
                                        structural_paths=lambda: (system, digital),
                                        block_for_path=lambda path: next((block for block in blocks if block.path == path), None))
                gui = object.__new__(WorkflowGUI)
                canvas = tk.Canvas(window, highlightthickness=0, borderwidth=0)
                canvas.pack(fill=tk.BOTH, expand=True)
                gui.sysml_hierarchy_window_canvas = canvas
                gui.sysml_materialization_scope = scope
                gui.sysml_scope_diagnostic = ""
                gui._get_sysml_materialization_scope = MagicMock(return_value=scope)
                gui._concrete_sysml_block_paths = MagicMock(return_value=[block.path for block in blocks])
                for geometry in ("1200x360", "420x480"):
                    window.geometry(geometry)
                    window.update()
                    gui._draw_sysml_hierarchy_window()
                    rectangles = []
                    for block in blocks:
                        item = next(item for item in canvas.find_withtag(block.path.as_posix()) if canvas.type(item) == "rectangle")
                        rectangles.append(canvas.coords(item))
                    self.assertGreaterEqual(len({bounds[0] for bounds in rectangles}), 3)
                    self.assertLessEqual(len({bounds[1] for bounds in rectangles}), 3)
                    self.assertTrue(all(bounds[3] - bounds[1] <= 118 for bounds in rectangles))
                    left, right = min(bounds[0] for bounds in rectangles), max(bounds[2] for bounds in rectangles)
                    for path in (system, digital):
                        item = next(item for item in canvas.find_withtag(path.as_posix()) if canvas.type(item) == "rectangle")
                        bounds = canvas.coords(item)
                        self.assertAlmostEqual((bounds[0] + bounds[2]) / 2, (left + right) / 2)
                    self.assertEqual(len(canvas.find_withtag("sysml_hierarchy_architecture_edge")), 0)
                    self.assertEqual(len(canvas.find_withtag("sysml_hierarchy_dependency_edge")), 9)
                    self.assertTrue(any(canvas.itemcget(item, "text") == "+96 in / +0 out" for item in canvas.find_withtag(blocks[0].path.as_posix()) if canvas.type(item) == "text"))
                    self.assertIn("data99", blocks[0].path.read_text(encoding="utf-8"))
        finally:
            window.destroy()

    def test_saved_window_ratios_follow_monitor_resolution_and_negative_origin(self):
        gui = object.__new__(WorkflowGUI)
        window = MagicMock()
        window.geometry.return_value = "1600x800+80+-900"
        window.state.return_value = "normal"
        window.minsize.return_value = (320, 240)
        gui._monitor_work_area = MagicMock(return_value=(0, -1000, 2000, 0))
        entry = gui._capture_layout_window(window)
        self.assertEqual(entry["rect_ratios"], [0.8, 0.8, 0.04, 0.1])
        for area, expected in (((0, 0, 1000, 600), "800x480+40+60"),
                               ((-1500, -900, 0, 0), "1200x720+-1440+-810")):
            gui._monitor_work_area.return_value = area
            gui._apply_layout_window(window, entry)
            window.geometry.assert_called_with(expected)

    def test_autosave_keeps_user_layout_and_hidden_pane_ratios(self):
        import json
        window = tk.Tk()
        try:
            window.geometry("900x600")
            pane = ttk.Panedwindow(window, orient=tk.HORIZONTAL)
            pane.pack(fill=tk.BOTH, expand=True)
            pane.add(ttk.Frame(pane), weight=1)
            pane.add(ttk.Frame(pane), weight=1)
            window.update()
            gui = object.__new__(WorkflowGUI)
            gui.root = window
            gui.work_pane = pane
            gui.layout_scale = 1.0
            gui.verbose_trace_enabled = False
            gui._gui_layout_ready = True
            original = {"geometry": "900x600+0+0", "sashes": {"work_pane": [315]}, "sash_ratios": {"work_pane": [0.35]}}
            gui.gui_state = {"layout_saved_by_user": True, "last_user_layout": original, "previous_user_layout": original}
            gui._append_log_ui = MagicMock()
            with tempfile.TemporaryDirectory() as directory, patch("workflow_gui.GUI_STATE_PATH", Path(directory) / "state.json"):
                pane.sashpos(0, 450)
                gui._save_gui_state()
                self.assertEqual(gui.gui_state["last_user_layout"], original)
                pane.pack_forget()
                window.update()
                gui._save_gui_state(user_initiated=True)
                saved = json.loads((Path(directory) / "state.json").read_text(encoding="utf-8"))
                self.assertEqual(saved["last_user_layout"]["sash_ratios"]["work_pane"], [0.35])
                self.assertEqual(saved["previous_user_layout"], original)
                self.assertIn("rect_ratios", saved["last_user_layout"]["windows"]["root"])
                self.assertFalse(saved["last_user_layout"]["windows"]["sysml_hierarchy_window"]["open"])
        finally:
            window.destroy()

    def test_saved_pane_ratios_survive_resolution_changes_and_hidden_panes(self):
        window = tk.Tk()
        try:
            window.geometry("1000x600")
            pane = ttk.Panedwindow(window, orient=tk.HORIZONTAL)
            pane.pack(fill=tk.BOTH, expand=True)
            pane.add(ttk.Frame(pane), weight=1)
            pane.add(ttk.Frame(pane), weight=1)
            window.update()
            gui = object.__new__(WorkflowGUI)
            layout = {"sash_ratios": {"work_pane": [0.35]}}
            for geometry in ("1000x600", "600x400", "1400x800"):
                window.geometry(geometry)
                window.update()
                self.assertTrue(gui._apply_saved_pane_layout("work_pane", pane, layout))
                self.assertAlmostEqual(pane.sashpos(0) / pane.winfo_width(), 0.35, delta=0.003)
            self.assertFalse(gui._apply_saved_pane_layout("unknown_pane", pane, {}))
            legacy = {"geometry": "2000x1000", "sashes": {"work_pane": [700]}}
            self.assertTrue(gui._apply_saved_pane_layout("work_pane", pane, legacy))
            self.assertAlmostEqual(pane.sashpos(0) / pane.winfo_width(), 0.35, delta=0.003)
            pane.pack_forget()
            window.update()
            self.assertFalse(gui._apply_saved_pane_layout("work_pane", pane, layout))
        finally:
            window.destroy()

    def test_saved_layout_on_negative_monitor_does_not_fall_back_to_default_size(self):
        self.assertEqual(scale_window_geometry("1920x1009+80+-937", 1), "1920x1009+80-937")
        gui = object.__new__(WorkflowGUI)
        gui.root = MagicMock()
        gui.root.geometry.return_value = "1920x1009+80+-937"
        gui.root.winfo_screenwidth.return_value = 1280
        gui.root.winfo_screenheight.return_value = 800
        gui.layout_scale = 0.8
        gui._monitor_work_area = MagicMock(return_value=(0, -1080, 2000, 80))
        gui.ensure_visible()
        gui._monitor_work_area.assert_called_once_with(1040, -433)
        gui.root.geometry.assert_called_with("1920x1009+80+-937")

    def test_sysml_hierarchy_zoom_reaches_near_zero_and_fits_complete_graph(self):
        window = tk.Tk()
        try:
            gui = object.__new__(WorkflowGUI)
            canvas = tk.Canvas(window, width=420, height=240, highlightthickness=0, borderwidth=0)
            canvas.pack()
            window.update()
            gui.sysml_hierarchy_window_canvas = canvas
            gui.sysml_hierarchy_zoom = 1.0
            gui.sysml_hierarchy_zoom_var = tk.StringVar(master=window)

            def draw():
                canvas.delete("all")
                zoom = gui.sysml_hierarchy_zoom
                canvas.create_rectangle(0, 0, 980 * zoom, 6000 * zoom)
                canvas.configure(scrollregion=canvas.bbox("all"))

            gui._draw_sysml_hierarchy_window = MagicMock(side_effect=draw)
            gui._zoom_sysml_hierarchy(0)
            self.assertEqual(gui.sysml_hierarchy_zoom, 0.001)
            self.assertEqual(gui.sysml_hierarchy_zoom_var.get(), "Zoom: 0.1%")
            gui._reset_sysml_hierarchy_zoom()
            gui._fit_sysml_hierarchy_zoom()
            self.assertLess(gui.sysml_hierarchy_zoom, 0.5)
            bounds = canvas.bbox("all")
            self.assertGreaterEqual(bounds[0], canvas.canvasx(0))
            self.assertLessEqual(bounds[2], canvas.canvasx(canvas.winfo_width()))
            self.assertGreaterEqual(bounds[1], canvas.canvasy(0))
            self.assertLessEqual(bounds[3], canvas.canvasy(canvas.winfo_height()))
            window.geometry("300x160")
            window.update()
            gui._fit_sysml_hierarchy_zoom()
            bounds = canvas.bbox("all")
            self.assertLessEqual(bounds[2] - bounds[0], canvas.winfo_width())
            self.assertLessEqual(bounds[3] - bounds[1], canvas.winfo_height())
        finally:
            window.destroy()

    def test_sysml_connections_popup_deduplicates_directions_and_reuses_window(self):
        window = tk.Tk()
        try:
            first, second, third, isolated = (Path(name + ".sysml") for name in ("First", "Second", "Third", "Isolated"))
            scope = SimpleNamespace(block_for_path=lambda path: SimpleNamespace(name=path.stem))
            gui = object.__new__(WorkflowGUI)
            gui.sysml_hierarchy_window_canvas = tk.Canvas(window)
            gui.sysml_hierarchy_window_canvas.pack()
            gui._get_sysml_materialization_scope = MagicMock(return_value=scope)
            gui._sysml_architecture_connections = MagicMock(return_value=[
                {"source": first, "target": second}, {"source": first, "target": second},
                {"source": second, "target": first}, {"source": third, "target": first},
                {"source": third, "target": second},
            ])
            gui._show_sysml_connections_popup(first)
            popup = gui.sysml_connections_window
            window.update()
            tree = gui.sysml_connections_tree
            rows = [tree.item(item, "values") for item in tree.get_children()]
            self.assertEqual([(row[0], row[1]) for row in rows], [("Second", "Bidirectional"), ("Third", "Incoming")])
            self.assertEqual(gui.sysml_connections_heading_var.get(), "Connected to: First")
            self.assertTrue(popup.winfo_viewable())
            self.assertIsNone(window.grab_current())
            gui._select_sysml_file = MagicMock()
            gui._select_sysml_hierarchy_path = MagicMock()
            tree.selection_set(tree.get_children()[0])
            gui._on_sysml_connected_block_open()
            gui._select_sysml_file.assert_called_once_with(second)
            gui._select_sysml_hierarchy_path.assert_called_once_with(second)
            gui._show_sysml_connections_popup(second)
            self.assertIs(gui.sysml_connections_window, popup)
            rows = [tree.item(item, "values") for item in tree.get_children()]
            self.assertEqual([(row[0], row[1]) for row in rows], [("First", "Bidirectional"), ("Third", "Incoming")])
            gui._show_sysml_connections_popup(isolated)
            self.assertFalse(tree.get_children())
            self.assertEqual(gui.sysml_connections_status_var.get(), "No explicit architecture connections.")
            popup.destroy()
            gui._show_sysml_connections_popup(third)
            self.assertIsNot(gui.sysml_connections_window, popup)
            rows = [gui.sysml_connections_tree.item(item, "values") for item in gui.sysml_connections_tree.get_children()]
            self.assertTrue(all(row[1] == "Outgoing" for row in rows))
        finally:
            window.destroy()

    def test_sysml_connections_are_explicit_scoped_and_port_aware(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            system, digital = root / "System.sysml", root / "Digital.sysml"
            first, second = root / "First.sysml", root / "Second.sysml"
            system.write_text("connection forward : Link connect topDigital.first.out to topDigital.second.input;\n"
                              "connection reverse : Link connect topDigital.second to topDigital.first;\n"
                              "connection unknown : Link connect topDigital.first to topDigital.missing;\n"
                              "/* connection comment : Link connect topDigital.first to topDigital.second; */\n", encoding="utf-8")
            digital.write_text("part first : First; part second : Second;", encoding="utf-8")
            scope = SimpleNamespace(system_path=system, digital_path=digital, system_definition="System",
                                    blocks=(SimpleNamespace(node="topDigital.first", path=first), SimpleNamespace(node="topDigital.second", path=second)),
                                    structural_paths=lambda: (system, digital))
            connections = WorkflowGUI._sysml_architecture_connections(scope)
            self.assertEqual([(value["source"], value["target"]) for value in connections], [(first, second), (second, first)])

    def test_sysml_hierarchy_dropdown_highlights_connections_and_centers(self):
        window = tk.Tk()
        try:
            with tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                system, digital = root / "System.sysml", root / "Digital.sysml"
                first, second = root / "First.sysml", root / "Second.sysml"
                system.write_text("part def System { part topDigital : DigitalSubsystem; connection link : Link connect topDigital.first to topDigital.second; }", encoding="utf-8")
                digital.write_text("part def DigitalSubsystem { part first : First; part second : Second; }", encoding="utf-8")
                first.write_text("part def First { in port data : Signal; }", encoding="utf-8")
                second.write_text("part def Second { out port data : Signal; }", encoding="utf-8")
                blocks = (SimpleNamespace(node="topDigital.first", path=first, target="Digital IPOS", definition="First"),
                          SimpleNamespace(node="topDigital.second", path=second, target="Digital IPOS", definition="Second"))
                scope = SimpleNamespace(system_path=system, digital_path=digital, system_definition="System", blocks=blocks,
                                        structural_paths=lambda: (system, digital),
                                        block_for_path=lambda path: next((block for block in blocks if block.path == path), None))
                gui = object.__new__(WorkflowGUI)
                gui.sysml_hierarchy_window_canvas = tk.Canvas(window, width=420, height=240, highlightthickness=0, borderwidth=0)
                gui.sysml_hierarchy_window_canvas.pack()
                window.update_idletasks()
                gui.sysml_materialization_scope = scope
                gui.sysml_scope_diagnostic = ""
                gui._get_sysml_materialization_scope = MagicMock(return_value=scope)
                gui._concrete_sysml_block_paths = MagicMock(return_value=[first, second])
                gui._show_sysml_block_details = MagicMock()
                gui._draw_sysml_architecture = MagicMock()
                gui.sysml_hierarchy_source_var = tk.StringVar(value="selected")
                gui.sysml_hierarchy_source_paths = {"selected": second}
                with patch("workflow_gui.REPO_ROOT", root):
                    gui._on_sysml_hierarchy_source_choice()
                self.assertFalse(hasattr(gui, "sysml_connections_window"))
                self.assertEqual(gui.sysml_selected_block, "topDigital.second")
                canvas = gui.sysml_hierarchy_window_canvas
                edges = canvas.find_withtag("sysml_hierarchy_architecture_edge")
                self.assertEqual(len(edges), 0)
                self.assertEqual(len(canvas.find_withtag("sysml_hierarchy_dependency_edge")), 3)
                selected = canvas.find_withtag("sysml_hierarchy_path:" + second.as_posix())
                rectangle = next(item for item in selected if canvas.type(item) == "rectangle")
                self.assertEqual(canvas.itemcget(rectangle, "outline"), "#1976d2")
                bounds = canvas.bbox("sysml_hierarchy_path:" + second.as_posix())
                self.assertAlmostEqual(canvas.canvasx(canvas.winfo_width() / 2), (bounds[0] + bounds[2]) / 2, delta=2)
                self.assertAlmostEqual(canvas.canvasy(canvas.winfo_height() / 2), (bounds[1] + bounds[3]) / 2, delta=2)
                gui.work_tabs = MagicMock()
                gui.sysml_architecture_tab = MagicMock()
                gui._select_sysml_file = MagicMock()
                canvas.xview_moveto(0)
                canvas.yview_moveto(0)
                window.update()
                first_bounds = canvas.bbox("sysml_hierarchy_path:" + first.as_posix())
                click_x = round((first_bounds[0] + first_bounds[2]) / 2 - canvas.canvasx(0))
                click_y = round(first_bounds[1] + 16 - canvas.canvasy(0))
                with patch("workflow_gui.REPO_ROOT", root):
                    canvas.event_generate("<Motion>", x=click_x, y=click_y)
                    canvas.event_generate("<Button-1>", x=click_x, y=click_y)
                    window.update()
                gui._select_sysml_file.assert_called_with(first)
                self.assertEqual(gui.sysml_hierarchy_selected_path, first)
                bounds = canvas.bbox("sysml_hierarchy_path:" + first.as_posix())
                self.assertAlmostEqual(canvas.canvasx(canvas.winfo_width() / 2), (bounds[0] + bounds[2]) / 2, delta=2)
                self.assertAlmostEqual(canvas.canvasy(canvas.winfo_height() / 2), (bounds[1] + bounds[3]) / 2, delta=2)
                rectangle = next(item for item in canvas.find_withtag("sysml_hierarchy_path:" + second.as_posix()) if canvas.type(item) == "rectangle")
                self.assertEqual(canvas.itemcget(rectangle, "outline"), "#8493a1")
        finally:
            window.destroy()

    def test_canvas_centers_selected_node_at_edges_and_after_zoom(self):
        window = tk.Tk()
        try:
            canvas = tk.Canvas(window, width=420, height=260, highlightthickness=0, borderwidth=0)
            canvas.pack()
            window.update_idletasks()
            canvas.configure(scrollregion=(0, 0, 2400, 1800))
            canvas.create_rectangle(0, 0, 160, 80, tags=("selected",))
            self.assertTrue(WorkflowGUI._center_canvas_node(canvas, "selected"))
            self.assertAlmostEqual(canvas.canvasx(canvas.winfo_width() / 2), 80, delta=2)
            self.assertAlmostEqual(canvas.canvasy(canvas.winfo_height() / 2), 40, delta=2)
            canvas.move("selected", 2000, 1500)
            canvas.scale("selected", 0, 0, 1.5, 1.5)
            self.assertTrue(WorkflowGUI._center_canvas_node(canvas, "selected"))
            bounds = canvas.bbox("selected")
            self.assertAlmostEqual(canvas.canvasx(canvas.winfo_width() / 2), (bounds[0] + bounds[2]) / 2, delta=2)
            self.assertAlmostEqual(canvas.canvasy(canvas.winfo_height() / 2), (bounds[1] + bounds[3]) / 2, delta=2)
            self.assertFalse(WorkflowGUI._center_canvas_node(canvas, "missing"))
        finally:
            window.destroy()

    def test_traceability_table_selection_centers_existing_or_new_selection(self):
        gui = object.__new__(WorkflowGUI)
        gui.traceability_tree = MagicMock()
        gui.traceability_tree.selection.return_value = ("block-id",)
        gui.traceability_graph_canvas = MagicMock()
        gui.traceability_selected_block = None
        gui.traceability_selected_edge = "old-edge"
        gui._draw_generated_dependency_graph = MagicMock()
        gui._center_canvas_node = MagicMock()
        gui._on_traceability_tree_select()
        self.assertEqual(gui.traceability_selected_block, "block-id")
        self.assertIsNone(gui.traceability_selected_edge)
        gui._draw_generated_dependency_graph.assert_called_once()
        gui._center_canvas_node.assert_called_with(gui.traceability_graph_canvas, "traceability_node:block-id")
        gui._on_traceability_tree_select()
        self.assertEqual(gui._center_canvas_node.call_count, 2)
        gui._draw_generated_dependency_graph.assert_called_once()

    def test_ipos_blocker_report_log_link_opens_existing_artifact_viewer(self):
        gui = object.__new__(WorkflowGUI)
        gui.log = MagicMock()
        gui.log.index.return_value = "3.0"
        gui.refresh_work_area = MagicMock()
        gui.work_tabs = MagicMock()
        gui.work_browser_tab = object()
        gui._on_work_area_select_path = MagicMock()
        path = Path(tempfile.gettempdir()) / "remaining_ipos_blockers.md"
        gui._append_log_ui("[REPORT] Open IPOS blocker report: " + str(path))
        click = next(call.args[2] for call in gui.log.tag_bind.call_args_list if call.args[1] == "<Button-1>")
        click(None)
        gui.refresh_work_area.assert_called_once()
        gui.work_tabs.select.assert_called_once_with(gui.work_browser_tab)
        gui._on_work_area_select_path.assert_called_once_with(path)
        self.assertIn(str(path), gui.log.insert.call_args.args[1])

    def test_ipos_blocker_report_evidence_is_visible_and_does_not_change_run_status(self):
        gui = object.__new__(WorkflowGUI)
        gui.root = MagicMock()
        gui.append_log = MagicMock()
        gui.refresh_work_area = MagicMock()
        path = Path(tempfile.gettempdir()) / "remaining_ipos_blockers.md"
        with patch("generate_cross_spec_issue_report.write_remaining_ipos_blocker_report", return_value=path) as write:
            run_context = {"exit_code": 0, "stages": ["6"], "started_at": 1, "finished_at": 2}
            gui._append_ipos_blocker_report_evidence(["python", "--snapshot-id", "snapshot-test"], run_context=run_context)
        self.assertEqual(write.call_args.args[1], "snapshot-test")
        self.assertEqual(write.call_args.kwargs["run_context"], run_context)
        messages = [call.args[0] for call in gui.append_log.call_args_list]
        self.assertTrue(any("report generated" in value for value in messages))
        self.assertTrue(any("historical evidence separated" in value for value in messages))
        self.assertIn("[REPORT] Open IPOS blocker report: " + str(path.resolve()), messages)
        self.assertIn("[REPORT] Open IPOS blocker report: " + str(path.with_suffix(".docx").resolve()), messages)
        with patch("generate_cross_spec_issue_report.write_remaining_ipos_blocker_report", side_effect=ValueError("snapshot mismatch")):
            gui._append_ipos_blocker_report_evidence(["python", "--snapshot-id", "other-snapshot"])
        self.assertIn("not generated: snapshot mismatch", gui.append_log.call_args.args[0])

    def test_ipos_blocker_report_generated_after_failed_range_without_rerunning_flow(self):
        gui = object.__new__(WorkflowGUI)
        gui.is_running = False
        gui.root = MagicMock()
        gui.root.after.side_effect = lambda delay, callback, *args: callback(*args)
        for method in ("_set_status", "clear_stage_failure", "set_active_stage", "append_log", "append_chat", "set_stage_failure", "_append_ipos_blocker_report_evidence"):
            setattr(gui, method, MagicMock())
        process = MagicMock()
        process.stdout = ["[2026-10-12] Stage 6 Digital IPOS run 1/1\n", "STOP at stage 6 (exit=1)\n"]
        process.wait.return_value = 1
        command = ["python", "workflow_cli.py", "run", "--from-stage", "5", "--to-stage", "7", "--snapshot-id", "snapshot-test"]
        with patch("workflow_gui.threading.Thread") as thread, patch("workflow_gui.subprocess.Popen", return_value=process) as popen:
            thread.side_effect = lambda **kwargs: SimpleNamespace(start=kwargs["target"])
            gui._run_command_async(command, "Stages 5-7", stage="5")
        gui._append_ipos_blocker_report_evidence.assert_called_once()
        self.assertEqual(gui._append_ipos_blocker_report_evidence.call_args.args, (command,))
        run_context = gui._append_ipos_blocker_report_evidence.call_args.kwargs["run_context"]
        self.assertEqual(run_context["exit_code"], 1)
        self.assertEqual(run_context["stages"], ["6"])
        self.assertLessEqual(run_context["started_at"], run_context["finished_at"])
        gui.set_stage_failure.assert_called_once()
        self.assertEqual(gui.set_stage_failure.call_args.args[0], "6")
        self.assertEqual(popen.call_count, 1)
        self.assertFalse(gui.is_running)

    def test_ipos_blocker_report_after_success_keeps_run_status_and_scope(self):
        gui = object.__new__(WorkflowGUI)
        gui.is_running = False
        gui.root = MagicMock()
        for method in ("_set_status", "clear_stage_failure", "set_active_stage", "append_log", "append_chat",
                       "set_stage_failure", "_append_ipos_blocker_report_evidence"):
            setattr(gui, method, MagicMock())
        process = MagicMock()
        process.stdout = ["[2026-10-15] Stage 6 Digital IPOS run 1/1\n"]
        process.wait.return_value = 0
        command = ["python", "workflow_cli.py", "run", "--stage", "6", "--snapshot-id", "snapshot-test",
                   "--block", "Module"]
        with patch("workflow_gui.threading.Thread") as thread, patch("workflow_gui.subprocess.Popen", return_value=process) as popen:
            thread.side_effect = lambda **kwargs: SimpleNamespace(start=kwargs["target"])
            gui._run_command_async(command, "Stage 6 Digital IPOS", stage="6")
        run_context = gui._append_ipos_blocker_report_evidence.call_args.kwargs["run_context"]
        self.assertEqual(run_context["exit_code"], 0)
        self.assertEqual(run_context["stages"], ["6"])
        self.assertEqual(run_context["block"], "Module")
        gui.set_stage_failure.assert_not_called()
        self.assertEqual(popen.call_count, 1)
        self.assertFalse(gui.is_running)

    def test_open_validation_report_uses_read_only_viewer_without_file_association(self):
        root = tk.Tk()
        root.withdraw()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.root = root
            gui._set_responsive_geometry = MagicMock()
            result = {"selected_snapshot_id": "snapshot-test", "findings": ["SRS_STALE_CONTRACT_FINGERPRINT"]}

            def descendants(widget):
                for child in widget.winfo_children():
                    yield child
                    yield from descendants(child)

            with tempfile.TemporaryDirectory() as temporary_dir:
                repo_root = Path(temporary_dir)
                report_path = repo_root / "artifacts/validation/downstream_coherence_report.json"
                with patch("workflow_gui.REPO_ROOT", repo_root), \
                        patch("workflow_gui.DOWNSTREAM_VALIDATION_REPORT_PATH", report_path), \
                        patch("validate_downstream_coherence.validate", return_value=result), \
                        patch("workflow_gui.os.startfile", create=True) as startfile, \
                        patch("workflow_gui.messagebox.showerror") as showerror:
                    gui._show_downstream_block_dialog([], ["python", "workflow_cli.py", "run", "--stage", "3"])
                    dialog = next(child for child in root.winfo_children() if isinstance(child, tk.Toplevel))
                    diagnostic_viewer = next(child for child in descendants(dialog) if isinstance(child, tk.Text))
                    diagnostic_text = diagnostic_viewer.get("1.0", "end-1c")
                    self.assertIn("snapshot-test", diagnostic_text)
                    self.assertIn("SRS_STALE_CONTRACT_FINGERPRINT", diagnostic_text)
                    self.assertIn("python workflow_cli.py run --stage 3", diagnostic_text)
                    self.assertEqual(diagnostic_viewer.cget("state"), tk.DISABLED)
                    self.assertTrue(diagnostic_viewer.bind("<Control-c>"))
                    self.assertTrue(diagnostic_viewer.bind("<Button-3>"))
                    diagnostic_viewer.tag_add(tk.SEL, "1.0", "end-1c")
                    gui._copy_text_selection(diagnostic_viewer)
                    self.assertEqual(root.clipboard_get(), diagnostic_text)
                    diagnostic_viewer.insert("1.0", "unwanted edit")
                    self.assertEqual(diagnostic_viewer.get("1.0", "end-1c"), diagnostic_text)
                    open_button = next(child for child in descendants(dialog) if isinstance(child, ttk.Button) and child.cget("text") == "Open validation report")
                    open_button.invoke()
                    report_dialog = next(child for child in dialog.winfo_children() if isinstance(child, tk.Toplevel))
                    viewer = next(child for child in descendants(report_dialog) if isinstance(child, tk.Text))
                    self.assertEqual(viewer.get("1.0", "end-1c"), report_path.read_text(encoding="utf-8"))
                    self.assertEqual(viewer.cget("state"), tk.DISABLED)
                    self.assertTrue(viewer.bind("<Control-c>"))
                    self.assertTrue(viewer.bind("<Button-3>"))
                    viewer.tag_add(tk.SEL, "1.0", "end-1c")
                    gui._copy_text_selection(viewer)
                    self.assertEqual(root.clipboard_get(), viewer.get("1.0", "end-1c"))
                    self.assertEqual(root.grab_current(), report_dialog)
                    viewer.insert("1.0", "unwanted edit")
                    self.assertNotIn("unwanted edit", viewer.get("1.0", "end-1c"))
                    close_button = next(child for child in descendants(report_dialog) if isinstance(child, ttk.Button) and child.cget("text") == "Close")
                    close_button.invoke()
                    self.assertEqual(root.grab_current(), dialog)
                    startfile.assert_not_called()
                    showerror.assert_not_called()
                    report_path.unlink()
                    open_button.invoke()
                    showerror.assert_called_once()
                    self.assertFalse(any(isinstance(child, tk.Toplevel) for child in dialog.winfo_children()))
        finally:
            root.update_idletasks()
            root.destroy()

    def test_gui_stage_run_tracks_all_executable_stages_and_failure(self):
        stages = ("0", "1", "2", "2a", "3", "4", "5", "6", "7")
        for stage in stages:
            with self.subTest(stage=stage):
                gui = object.__new__(WorkflowGUI)
                gui.is_running = False
                gui.root = MagicMock()
                gui.root.after.side_effect = lambda delay, callback, *args: callback(*args)
                for method in ("_set_status", "clear_stage_failure", "set_active_stage", "append_log", "append_chat", "set_stage_failure"):
                    setattr(gui, method, MagicMock())
                process = MagicMock()
                process.stdout = [
                    f"[2026-10-06 13:00:00] Stage {stage.upper()} run 1/1\n",
                    "Stage 1.4 legacy substep; references Stage 3 and Stage 7\n",
                    f"[2026-10-06 13:00:01] STOP at stage {stage} (exit=1)\n",
                ]
                process.wait.return_value = 1
                with patch("workflow_gui.threading.Thread") as thread, \
                        patch("workflow_gui.subprocess.Popen", return_value=process):
                    thread.side_effect = lambda **kwargs: SimpleNamespace(start=kwargs["target"])
                    gui._run_command_async(["python", "workflow_cli.py", "run"], "Stage range")
                self.assertEqual([call.args[0] for call in gui.set_active_stage.call_args_list], [None, stage, None])
                self.assertEqual([call.args[0] for call in gui.clear_stage_failure.call_args_list], [None, stage])
                gui.set_stage_failure.assert_called_once()
                self.assertEqual(gui.set_stage_failure.call_args.args[0], stage)
                self.assertEqual(gui.set_active_stage.call_args.args, (None,))
                self.assertFalse(gui.is_running)

    def test_gui_stage_range_stays_on_stage_that_fails(self):
        gui = object.__new__(WorkflowGUI)
        gui.is_running = False
        gui.root = MagicMock()
        gui.root.after.side_effect = lambda delay, callback, *args: callback(*args)
        for method in ("_set_status", "clear_stage_failure", "set_active_stage", "append_log", "append_chat", "set_stage_failure"):
            setattr(gui, method, MagicMock())
        process = MagicMock()
        process.stdout = [
            "[2026-10-06 13:00:00] Stage 5 run 1/1\n",
            "[2026-10-06 13:00:01] Stage 6 Digital IPOS run 1/1\n",
            "[2026-10-06 13:00:02] Stage 7 Analog IPOS run 1/1\n",
            "ERROR: Stage 7 failed; Stage 3 artifact missing\n",
            "[2026-10-06 13:00:03] STOP at stage 7 (exit=3)\n",
        ]
        process.wait.return_value = 3
        with patch("workflow_gui.threading.Thread") as thread, \
                patch("workflow_gui.subprocess.Popen", return_value=process) as popen:
            thread.side_effect = lambda **kwargs: SimpleNamespace(start=kwargs["target"])
            gui._run_command_async(["python", "workflow_cli.py", "run"], "Stages 5-7", stage="5")
        self.assertEqual([call.args[0] for call in gui.set_active_stage.call_args_list], ["5", "5", "6", "7", None])
        self.assertEqual(gui.set_stage_failure.call_args.args[0], "7")
        self.assertEqual(popen.call_count, 1)
        self.assertEqual(popen.call_args.kwargs["env"]["PYTHONUNBUFFERED"], "1")

    def test_full_workflow_diagram_retains_failure_after_run_and_redraw(self):
        root = tk.Tk()
        try:
            frame = ttk.Frame(root)
            frame.pack(fill=tk.BOTH, expand=True)
            renderer = WorkflowDiagramRenderer(frame)
            for stage, node_id in (("2", "s2"), ("2a", "s2a"), ("7", "stage7")):
                with self.subTest(stage=stage):
                    renderer.set_stage_failures({})
                    renderer.set_active_stage(stage)
                    rect = renderer.node_items[node_id][0]
                    self.assertEqual(renderer.canvas.itemcget(rect, "fill"), renderer.COLORS["active"]["fill"])
                    renderer.set_stage_failures({stage: "Gate failed"})
                    renderer.set_active_stage(None)
                    renderer.refresh()
                    rect = renderer.node_items[node_id][0]
                    self.assertEqual(renderer.canvas.itemcget(rect, "fill"), renderer.COLORS["failed"]["fill"])
                    renderer.set_stage_failures({})
                    renderer.set_active_stage(stage)
                    self.assertEqual(renderer.canvas.itemcget(rect, "fill"), renderer.COLORS["active"]["fill"])
        finally:
            root.destroy()

    def test_standalone_analog_ipos_sets_diagram_stage(self):
        gui = object.__new__(WorkflowGUI)
        gui._snapshot_selector_args = MagicMock(return_value=["--use-latest-approved"])
        gui._run_command_async = MagicMock()
        gui.run_stage_cli("7")
        self.assertEqual(gui._run_command_async.call_args.kwargs["stage"], "7")

    def test_final_sysml_run_sets_diagram_stage(self):
        gui = object.__new__(WorkflowGUI)
        gui._snapshot_selector_args = MagicMock(return_value=["--snapshot-id", "snapshot-test"])
        gui._run_command_async = MagicMock()
        gui.regenerate_sysml()
        self.assertEqual(gui._run_command_async.call_args.kwargs.get("stage"), "sysml")

    def test_final_sysml_compact_boxes_track_run_result_after_redraw(self):
        root = tk.Tk()
        root.withdraw()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.root = MagicMock()
            gui.root.after.side_effect = lambda delay, callback, *args: callback(*args)
            gui.diagram_canvas = tk.Canvas(root)
            gui.compact_workflow_zoom = 1.0
            gui.active_stage = None
            gui.hover_stage = None
            gui.stage_failure_reasons = {}
            gui.diagram_hint_var = tk.StringVar(master=root)
            gui._refresh_generated_sysml = MagicMock()
            for method in ("_set_status", "append_log", "append_chat"):
                setattr(gui, method, MagicMock())
            gui._snapshot_selector_args = MagicMock(return_value=["--snapshot-id", "snapshot-test"])
            gui._draw_compact_workflow_graph()
            for result in (0, 1, "exception", 0):
                with self.subTest(result=result):
                    gui.is_running = False
                    process = MagicMock()
                    process.stdout = ["Final SysML phase: PASS; implementation readiness: NOT_READY\n"] if result == 0 else ["Generation failed\n"]
                    process.wait.return_value = result

                    def start_process(*args, **kwargs):
                        self.assertEqual(gui.active_stage, "sysml")
                        self.assertFalse(getattr(gui, "sysml_generation_succeeded", False))
                        for key in ("sysml", "sysml_readiness"):
                            self.assertEqual(gui.diagram_canvas.itemcget(gui.stage_rects[key], "fill"), WorkflowDiagramRenderer.COLORS["active"]["fill"])
                        if result == "exception":
                            raise OSError("Cannot start SysML generation")
                        return process

                    with patch("workflow_gui.threading.Thread") as thread, \
                            patch("workflow_gui.subprocess.Popen", side_effect=start_process) as popen:
                        thread.side_effect = lambda **kwargs: SimpleNamespace(start=kwargs["target"])
                        gui.regenerate_sysml()
                    popen.assert_called_once()
                    self.assertIsNone(gui.active_stage)
                    self.assertEqual(gui.sysml_generation_succeeded, result == 0)
                    self.assertEqual("sysml" in gui.stage_failure_reasons, result != 0)
                    gui._draw_compact_workflow_graph()
                    expected = "active" if result == 0 else "failed"
                    for key in ("sysml", "sysml_readiness"):
                        self.assertEqual(gui.diagram_canvas.itemcget(gui.stage_rects[key], "fill"), WorkflowDiagramRenderer.COLORS[expected]["fill"])
            gui._snapshot_selector_args.return_value = None
            with patch.object(gui, "_run_command_async") as run:
                gui.regenerate_sysml()
            run.assert_not_called()
        finally:
            root.update_idletasks()
            root.destroy()

    def test_saved_pane_reapplication_does_not_log(self):
        gui = object.__new__(WorkflowGUI)
        gui.root = MagicMock()
        gui.gui_state = {"layout_saved_by_user": True, "last_user_layout": {"sashes": {}}}
        gui.layout_scale = 1.0
        gui._gui_pane_names = MagicMock(return_value=("work_pane",))
        gui.work_pane = MagicMock()
        gui.work_pane.panes.return_value = ("left", "right")
        gui._apply_saved_pane_layout = MagicMock(return_value=True)
        gui._reapply_saved_gui_sashes()
        gui._apply_saved_pane_layout.assert_called_once_with("work_pane", gui.work_pane, gui.gui_state["last_user_layout"])
        gui.root.update_idletasks.assert_called_once_with()
        gui.root.after.assert_not_called()

    def test_gui_layout_scales_to_screen_and_preserves_saved_position(self):
        self.assertEqual(responsive_layout_scale(1024, 600), 0.72)
        self.assertEqual(responsive_layout_scale(1440, 900), 1.0)
        self.assertEqual(responsive_layout_scale(3840, 2160), 1.25)
        self.assertEqual(scale_window_geometry("1200x800+40+60", 0.8), "960x640+40+60")
        self.assertEqual(scale_window_geometry("invalid", 1.2), "invalid")

    def test_secondary_window_geometry_is_scaled_and_screen_bounded(self):
        class WindowStub:
            geometry_value = None
            minimum = None

            def winfo_screenwidth(self):
                return 800

            def winfo_screenheight(self):
                return 600

            def geometry(self, value):
                self.geometry_value = value

            def minsize(self, width, height):
                self.minimum = (width, height)

        gui = WorkflowGUI.__new__(WorkflowGUI)
        gui.layout_scale = 1.25
        window = WindowStub()
        gui._set_responsive_geometry(window, 1100, 800, 640, 480)
        self.assertEqual(window.geometry_value, "752x540")
        self.assertEqual(window.minimum, (752, 540))

    def test_stage6_execution_log_marks_activity_and_run_boundaries(self):
        class LogStub:
            def __init__(self):
                self.entries = []

            def insert(self, index, text, tags):
                self.entries.append((index, text, tags))

            def see(self, _index):
                pass

        gui = WorkflowGUI.__new__(WorkflowGUI)
        gui.log = LogStub()
        gui._append_log_ui("[IPOS_ACTIVITY] START | 2026-09-25T10:00:00+02:00 | Generate IPOS block")
        gui._append_log_ui("=== START Stage 6 Digital IPOS | start=2026-09-25T10:00:00+02:00 ===")
        self.assertEqual(gui.log.entries[0][2], ("ipos_activity",))
        self.assertEqual(gui.log.entries[1][2], ("ipos_run_boundary",))

    def test_stage6_runner_emits_timestamped_activity_start_and_end(self):
        output = io.StringIO()
        with redirect_stdout(output):
            result = _run_activity("Validate Digital IPOS gate", lambda: SimpleNamespace(returncode=0))
        self.assertEqual(result.returncode, 0)
        self.assertRegex(output.getvalue(), r"\[IPOS_ACTIVITY\] START \| \d{4}-\d{2}-\d{2}T")
        self.assertRegex(output.getvalue(), r"\[IPOS_ACTIVITY\] END \| \d{4}-\d{2}-\d{2}T")
        self.assertIn("exit=0", output.getvalue())

    def test_ipos_pandoc_error_includes_stderr_and_office_lock_guidance(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "digital_ipos_main-controller.docx"
            lock = root / "~$gital_ipos_main-controller.docx"
            lock.write_text("lock", encoding="utf-8")
            result = SimpleNamespace(returncode=1, stdout="", stderr="Permission denied")
            with patch("generate_ipos_specs.subprocess.run", return_value=result):
                with self.assertRaisesRegex(RuntimeError, "Permission denied") as error:
                    _convert_ipos_markdown_to_docx(root / "input.md", root / "template.docx", output, root)
            self.assertIn(str(lock), str(error.exception))
            self.assertIn("remove the stale lock", str(error.exception))

    def test_validate_all_runs_diagram_rendering_crosscheck_after_stage_gates(self):
        args = build_parser().parse_args(["validate", "--all", "--snapshot-id", "snapshot-test"])
        invoked = []

        def run_script(_repo_root, script, _dry_run):
            invoked.append(script)
            return 0

        def run_snapshot_script(_repo_root, command, _dry_run):
            invoked.append(command[0])
            return 0

        with tempfile.TemporaryDirectory() as temp_root:
            with patch("workflow_cli._resolve_workflow_snapshot", return_value="snapshot-test"), \
                    patch("workflow_cli._require_downstream_coherence"), \
                    patch("workflow_cli._log_manifest"), \
                    patch("workflow_cli._run_python", side_effect=run_script), \
                    patch("workflow_cli._run_python_cmd", side_effect=run_snapshot_script):
                result = cmd_validate(args, Path(temp_root))

        self.assertEqual(result, 0)
        self.assertEqual(invoked[-1], "scripts/validate_workflow_diagram_rendering.py")
        self.assertEqual(len(invoked), len(STAGE_ORDER) + 1)

    def test_entity_kind_controls_architecture_hierarchy_membership(self):
        block_defs = {
            "Concrete": {"entity_kind": "concrete_block"},
            "Boundary": {"entity_kind": "interface"},
            "Unassigned": {"entity_kind": "unassigned"},
        }
        self.assertEqual(_entity_kind("Boundary", block_defs), "interface")
        self.assertTrue(_is_concrete_block("Concrete", block_defs))
        self.assertFalse(_is_concrete_block("Boundary", block_defs))
        self.assertFalse(is_concrete_block("Boundary", {"Entity kind": "interface"}))

    def test_interaction_endpoint_does_not_synthesize_a_block(self):
        block_defs = {"Concrete": {"entity_kind": "concrete_block"}, "Unassigned": {"entity_kind": "unassigned"}}
        augmented = _ensure_interaction_blocks(
            block_defs,
            {},
            [{"From block": "ISPU debug", "To block": "Concrete"}],
        )
        self.assertNotIn("ISPU debug", augmented)

    def test_report_fields_put_requirement_and_hierarchy_columns_first(self):
        fields = _ordered_fields([
            {"notes": "n", "block": "ADC", "source_req_id": "REQ-1", "statement": "shall"}
        ])
        self.assertEqual(fields[:3], ["source_req_id", "block", "statement"])

    def test_gui_stage3_cli_prompts_snapshot_without_refreshing_mutable_sources(self):
        gui = object.__new__(WorkflowGUI)
        gui._snapshot_selector_args = MagicMock(return_value=["--snapshot-id", "previous-approved"])
        gui._run_command_async = MagicMock()
        gui._prepare_stage3_run = MagicMock()
        gui.run_stage_cli("3")
        gui._snapshot_selector_args.assert_called_once()
        gui._prepare_stage3_run.assert_not_called()
        command = gui._run_command_async.call_args.args[0]
        args = build_parser().parse_args(command[2:])
        self.assertEqual(args.stage, "3")
        self.assertEqual(args.snapshot_id, "previous-approved")
        self.assertTrue(args.prepare_snapshot_inputs)
        self.assertNotIn("--use-latest-approved", command)
        gui._snapshot_selector_args.return_value = None
        gui._run_command_async.reset_mock()
        gui.run_stage_cli("3")
        gui._run_command_async.assert_not_called()

    def test_gui_stage_range_dispatches_one_cli_command(self):
        gui = object.__new__(WorkflowGUI)
        gui._run_command_async = MagicMock()
        gui._snapshot_selector_args = MagicMock(return_value=["--snapshot-id", "approved-test"])
        for start, end in (("0", "2"), ("2a", "5"), ("3", "7"), ("6", "7")):
            with self.subTest(start=start, end=end):
                gui._run_command_async.reset_mock()
                gui._snapshot_selector_args.reset_mock()
                gui.from_stage_var = SimpleNamespace(get=lambda: next(label for label, key in RANGE_STAGE_LABEL_TO_KEY.items() if key == start))
                gui.to_stage_var = SimpleNamespace(get=lambda: next(label for label, key in RANGE_STAGE_LABEL_TO_KEY.items() if key == end))
                gui.run_stage_range_cli()
                gui._run_command_async.assert_called_once()
                command = gui._run_command_async.call_args.args[0]
                self.assertEqual(Path(command[1]).name, "workflow_cli.py")
                args = build_parser().parse_args(command[2:])
                self.assertEqual(args.from_stage, start)
                self.assertEqual(args.to_stage, end)
                self.assertEqual(args.prepare_snapshot_inputs, start in {"3", "4", "5", "6", "7"})
                if end == "2":
                    gui._snapshot_selector_args.assert_not_called()
                    self.assertIsNone(args.snapshot_id)
                else:
                    gui._snapshot_selector_args.assert_called_once()
                    self.assertEqual(args.snapshot_id, "approved-test")
                self.assertEqual(gui._run_command_async.call_args.kwargs["approval_after_success"], end == "2")
    def test_gui_stage_range_cancelled_snapshot_launches_nothing(self):
        gui = object.__new__(WorkflowGUI)
        gui._snapshot_selector_args = MagicMock(return_value=None)
        gui._run_command_async = MagicMock()
        gui.run_stage_range("3", "7")
        gui._snapshot_selector_args.assert_called_once()
        gui._run_command_async.assert_not_called()

    def test_gui_stage_range_preserves_approval_and_order_guards(self):
        gui = object.__new__(WorkflowGUI)
        gui.root = None
        gui._run_command_async = MagicMock()
        with patch("workflow_gui.messagebox.showinfo") as approval, \
                patch("workflow_gui.messagebox.showwarning") as invalid:
            gui.run_stage_range("0", "7")
            approval.assert_called_once()
            gui.run_stage_range("5", "3")
            invalid.assert_called_once()
        gui._run_command_async.assert_not_called()

    def test_report_sheet_has_filter_freeze_pane_and_table(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["source_req_id", "block", "statement"])
        worksheet.append(["REQ-1", "ADC", "shall sample"])
        _style_data_sheet(worksheet, ["source_req_id", "block", "statement"], 1)
        with io.BytesIO() as buffer:
            workbook.save(buffer)
            buffer.seek(0)
            worksheet = load_workbook(buffer).active
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(len(worksheet.tables), 1)
        table = worksheet.tables["ReportData"]
        self.assertEqual(table.ref, "A1:C2")
        self.assertIsNotNone(table.autoFilter)
        self.assertEqual(table.autoFilter.ref, "A1:C2")

    def test_sysml_hierarchy_source_scrollbar_survives_resize(self):
        root = tk.Tk()
        root.withdraw()
        self.addCleanup(root.destroy)
        gui = object.__new__(WorkflowGUI)
        gui.root = root
        with patch.object(gui, "_set_responsive_geometry"), \
                patch.object(gui, "_bind_button_help"), \
                patch.object(gui, "_bind_copy_support"), \
                patch.object(gui, "_refresh_sysml_relation_tabs"), \
                patch.object(gui, "_draw_sysml_hierarchy_window"):
            gui._open_sysml_hierarchy_window()
        text = gui.sysml_hierarchy_details
        scrollbar = next(child for child in text.master.winfo_children() if isinstance(child, ttk.Scrollbar))
        text_horizontal = next(child for child in text.master.winfo_children()
                       if isinstance(child, ttk.Scrollbar) and str(child.cget("orient")) == tk.HORIZONTAL)
        port_canvas = gui.sysml_hierarchy_port_canvas
        port_scrollbar = next(child for child in port_canvas.master.winfo_children() if isinstance(child, ttk.Scrollbar))
        port_horizontal = next(child for child in port_canvas.master.winfo_children()
                       if isinstance(child, ttk.Scrollbar) and str(child.cget("orient")) == tk.HORIZONTAL)
        port_canvas.configure(scrollregion=(0, 0, 3000, 3000))
        text.configure(state=tk.NORMAL)
        text.insert("1.0", ("SysML source line " * 100 + "\n") * 200)
        text.configure(state=tk.DISABLED)
        details_pane = text.master.master
        self.assertIsInstance(details_pane, ttk.Panedwindow)
        self.assertEqual(str(details_pane.cget("orient")), tk.HORIZONTAL)
        self.assertEqual(len(details_pane.panes()), 2)
        for geometry in ("1100x700", "620x420", "900x550"):
            with self.subTest(geometry=geometry):
                gui.sysml_hierarchy_window.geometry(geometry)
                root.update()
                details_pane.sashpos(0, int(details_pane.winfo_width() * 0.3))
                root.update()
                self.assertTrue(port_scrollbar.winfo_ismapped())
                self.assertGreater(port_scrollbar.winfo_width(), 1)
                self.assertLessEqual(port_scrollbar.winfo_x() + port_scrollbar.winfo_width(), port_canvas.master.winfo_width())
                initial_port_width = gui.sysml_hierarchy_port_canvas.winfo_width()
                initial_text_width = text.winfo_width()
                details_pane.sashpos(0, int(details_pane.winfo_width() * 0.7))
                root.update()
                self.assertGreater(gui.sysml_hierarchy_port_canvas.winfo_width(), initial_port_width)
                self.assertLess(text.winfo_width(), initial_text_width)
                self.assertTrue(scrollbar.winfo_ismapped())
                self.assertGreater(scrollbar.winfo_width(), 1)
                self.assertGreater(scrollbar.winfo_height(), 1)
                self.assertGreater(text.winfo_width(), 1)
                self.assertLessEqual(scrollbar.winfo_x() + scrollbar.winfo_width(), text.master.winfo_width())
                text.yview_moveto(0)
                root.update()
                self.assertEqual(scrollbar.get()[0], 0)
                root.tk.call(scrollbar.cget("command"), "moveto", 1.0)
                root.update()
                self.assertGreater(text.yview()[0], 0)
                self.assertGreater(scrollbar.get()[0], 0)
                self.assertTrue(port_scrollbar.winfo_ismapped())
                self.assertGreater(port_scrollbar.winfo_width(), 1)
                self.assertGreater(port_scrollbar.winfo_height(), 1)
                self.assertLessEqual(port_scrollbar.winfo_x() + port_scrollbar.winfo_width(), port_canvas.master.winfo_width())
                port_canvas.yview_moveto(0)
                root.update()
                self.assertEqual(port_scrollbar.get()[0], 0)
                root.tk.call(port_scrollbar.cget("command"), "moveto", 1.0)
                root.update()
                self.assertGreater(port_canvas.yview()[0], 0)
                self.assertGreater(port_scrollbar.get()[0], 0)
                for widget, horizontal in ((text, text_horizontal), (port_canvas, port_horizontal)):
                    self.assertTrue(horizontal.winfo_ismapped())
                    self.assertGreater(horizontal.winfo_width(), 1)
                    self.assertGreater(horizontal.winfo_height(), 1)
                    self.assertLessEqual(horizontal.winfo_y() + horizontal.winfo_height(), horizontal.master.winfo_height())
                    widget.xview_moveto(0)
                    root.tk.call(horizontal.cget("command"), "moveto", 1.0)
                    root.update()
                    self.assertGreater(widget.xview()[0], 0)
                notebook = gui.sysml_relation_notebook
                for tab in notebook.tabs():
                    notebook.select(tab)
                    root.update()
                    frame = notebook.nametowidget(tab)
                    containers = [frame, *frame.winfo_children()]
                    scrollbars = [child for container in containers for child in container.winfo_children()
                                  if isinstance(child, ttk.Scrollbar)]
                    self.assertEqual({str(child.cget("orient")) for child in scrollbars}, {tk.HORIZONTAL, tk.VERTICAL})
                    for child in scrollbars:
                        self.assertTrue(child.winfo_ismapped())
                        self.assertGreater(child.winfo_width(), 1)
                        self.assertGreater(child.winfo_height(), 1)
                        self.assertLessEqual(child.winfo_x() + child.winfo_width(), child.master.winfo_width())
                        self.assertLessEqual(child.winfo_y() + child.winfo_height(), child.master.winfo_height())

    def test_gui_context_is_deterministic_without_display(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            context = build_gui_review_context(root)
            rendered = format_gui_review_context(context)
            self.assertIn("Snapshot: unavailable", rendered)
            self.assertIn("blocked: no complete approved snapshot/materialization", rendered)
            self.assertIn("Authority scanner: not run", rendered)

    def test_architecture_approval_dialog_stays_open_for_stage2b(self):
        source = inspect.getsource(WorkflowGUI._show_architecture_profile_approval_dialog)
        self.assertIn("1. Open Architecture Map in Excel (orange button): review/change, save, and close the Excel file.", source)
        self.assertIn("2. Enable Gate from Workbook (green button).", source)
        self.assertIn("3. Run Stage 2A.", source)
        self.assertIn("4. After Stage 2A passes, Freeze Stage 2B Snapshot.", source)
        self.assertLess(source.index('text="Freeze Stage 2B Snapshot"'), source.index('text="Run Stage 2A"'))
        stage2a_handler = source.split("def run_stage2a_after_approval", 1)[1].split("def freeze_stage2b_snapshot", 1)[0]
        self.assertIn('self.run_stage_cli("2a")', stage2a_handler)
        self.assertNotIn("close_dialog()", stage2a_handler)
        snapshot_handler = source.split("def freeze_stage2b_snapshot", 1)[1].split("def approve_all_mapping_rows", 1)[0]
        self.assertIn("close_dialog()", snapshot_handler)

    def test_workflow_graph_exposes_bootstrap_snapshot_reports_and_sysml(self):
        labels = " ".join(node["label"] for node in WORKFLOW_DIAGRAM_NODES)
        self.assertIn("Bootstrap", labels)
        self.assertIn("Snapshot", labels)
        self.assertIn("XLSX", labels)
        self.assertIn("SysML", labels)

    def test_new_workflow_nodes_have_hover_descriptions(self):
        for node_id in ("s2b", "s2c", "s2d", "s2e", "reports", "sysml"):
            self.assertTrue(WORKFLOW_NODE_MEANINGS[node_id])

    def test_sysml_attribute_parser_reads_typed_escaped_description(self):
        source = 'attribute description : String = "Full \\"quoted\\" description";'
        self.assertEqual(WorkflowGUI._sysml_attribute_value(source, "description"), 'Full "quoted" description')

    def test_button_help_popups_close_when_pointer_leaves(self):
        bind_source = inspect.getsource(WorkflowGUI._bind_button_help)
        show_source = inspect.getsource(WorkflowGUI._show_hover_help)
        self.assertIn('button.bind("<Enter>"', bind_source)
        self.assertIn('button.bind("<Leave>"', bind_source)
        self.assertIn('self._hide_hover_help()', bind_source)
        self.assertNotIn("after(2200", show_source)
        self.assertIn("self._close_popup()", show_source)
        self.assertIn("self._activate_popup(popup, \"hover\")", show_source)
        self.assertIn("self.help_popup_owner_widget = event.widget", show_source)

    def test_concrete_source_owner_wins_over_generic_paragraph_terms(self):
        block_defs = {"Control Block": {}, "Unassigned": {}}
        row = {
            "source": "Section 1.2 Requirements (under Section 1.1 Control Block), paragraph 001",
            "source_section_owner": "Control Block",
        }
        self.assertFalse(_is_generic_source_paragraph(row, block_defs))
        self.assertEqual(_preview_blocks_for_requirement(row, [], block_defs), ["Control Block"])

    def test_concrete_candidates_win_before_generic_source_context(self):
        block_defs = {"Regmap": {"function": "register map"}, "Unassigned": {}}
        ontology_links = {
            f"REQ-{index}": [
                {
                    "role": "Regmap",
                    "related_role": "",
                    "relation_type": "supports",
                    "function_or_property_evidence": "register address range",
                }
            ]
            for index in range(3000, 3007)
        }
        for index in range(3000, 3007):
            row = {
                "source_req_id": f"REQ-{index}",
                "source": "Section 13.4.1 Register Map (under Section 13.4 Register Map), paragraph 001",
                "requirement_statement": "register address range",
                "category": "Digital",
                "requirement_type": "other",
            }
            self.assertEqual(
                _map_blocks(row, {}, block_defs, {}, ontology_links),
                ["Regmap"],
            )
            self.assertEqual(
                _preview_blocks_for_requirement(row, ["Regmap"], block_defs),
                ["Regmap"],
            )

    def test_configured_source_alias_wins_without_ontology_candidate(self):
        row = {
            "source": "Section 13.4.1 STBIO1 Register Map (under Section 13.4 STBIO1 Register Map), paragraph 034",
        }
        self.assertEqual(
            _preview_blocks_for_requirement(
                row,
                [],
                {"Regmap": {}, "Unassigned": {}},
                {"Regmap": ["Register Map"]},
            ),
            ["Regmap"],
        )

    def test_source_paragraph_is_last_resort_context(self):
        block_defs = {"Register Block": {}, "Unassigned": {}}
        source_only = {
            "source": "Section 7.1 Operating mode (under Section 7 Control), paragraph 004",
        }
        self.assertEqual(
            _preview_blocks_for_requirement(source_only, [], block_defs),
            ["Operating mode"],
        )

        concrete = dict(source_only)
        concrete["source_section_owner"] = "Register Block"
        self.assertEqual(
            _preview_blocks_for_requirement(concrete, [], block_defs),
            ["Register Block"],
        )

    def test_adc_dft_test_paragraph_stays_source_context_until_approved(self):
        block_defs = {"ADC": {}, "Regmap": {}, "Unassigned": {}}
        source = (
            "Section 18.3 ADC TEST LOW NOISE "
            "(under Section 18 Digital DFT), paragraph 006 (page 160)"
        )
        self.assertEqual(
            _source_section_owner(source, {"ADC": ["ADC"], "Regmap": ["Regmap"]}),
            "",
        )
        for requirement_id in ("DDS_STBIO1_0701", "DDS_STBIO1_0702"):
            row = {
                "source_req_id": requirement_id,
                "source": source,
                "requirement_statement": "Enable ADC TEST LOW NOISE by register writes.",
                "source_section_owner": "ADC",
            }
            self.assertEqual(
                _preview_blocks_for_requirement(row, ["ADC", "Regmap"], block_defs),
                ["ADC TEST LOW NOISE"],
            )
            self.assertEqual(
                _map_blocks(row, {}, block_defs, {}, {}, []),
                ["Unassigned"],
            )

    def test_configured_interaction_requirement_maps_to_declared_endpoints(self):
        block_to_requirements = {
            "Source Block": [],
            "Destination Block": [],
        }
        block_defs = {"Source Block": {}, "Destination Block": {}, "Unassigned": {}}
        interaction_rows = [
            {
                "From block": "Source Block",
                "To block": "Destination Block",
                "Requirement IDs": "REQ-CONNECTION",
            }
        ]
        _link_interaction_requirements_to_blocks(
            block_to_requirements,
            block_defs,
            interaction_rows,
        )
        self.assertEqual(block_to_requirements, {
            "Source Block": ["REQ-CONNECTION"],
            "Destination Block": ["REQ-CONNECTION"],
        })
        row = {
            "source_req_id": "REQ-CONNECTION",
            "source": "Section 4.1 Connection Matrix (under Section 4 Interconnect), paragraph 002",
            "requirement_statement": "Source Block connects to Destination Block",
        }
        self.assertEqual(
            _map_blocks(row, {}, block_defs, {}, {}, interaction_rows),
            ["Source Block", "Destination Block"],
        )

    def test_source_matrix_rows_do_not_get_single_preview_owner(self):
        row = {
            "source_req_id": "REQ-MATRIX",
            "source": "Section 4.1 XBAR Connection Matrix (under Section 4 Interconnect)",
            "requirement_statement": "Source Block connects to Destination Block",
        }
        self.assertEqual(
            _preview_blocks_for_requirement(
                row,
                ["Destination Block"],
                {"Regmap": {}, "Destination Block": {}, "Unassigned": {}},
                {},
                {"REQ-MATRIX": ["Source Block", "Destination Block"]},
            ),
            ["Unassigned"],
        )

    def test_connection_matrix_context_is_safe_without_edge_ledger(self):
        row = {
            "source_req_id": "REQ-MATRIX",
            "source": "Section 4.1 XBAR Connection Matrix",
            "requirement_statement": "Source Block connects to Destination Block",
        }
        self.assertEqual(
            _preview_blocks_for_requirement(
                row,
                ["Regmap"],
                {"Regmap": {}, "Unassigned": {}},
            ),
            ["Unassigned"],
        )

    def test_explicit_to_target_beats_incidental_adc_term(self):
        row = {
            "requirement_statement": "Configure ADC calibration. [TO: IPOS_MAIN_CTRL]",
        }
        block_defs = {"ADC": {}, "Main Controller": {}, "Unassigned": {}}
        aliases = {"Main Controller": ["IPOS_MAIN_CTRL"]}
        self.assertEqual(_explicit_target_block(row, block_defs, aliases), "Main Controller")
        self.assertEqual(
            _map_blocks(row, {"ADC": ["adc"]}, block_defs, aliases, {}, []),
            ["Main Controller"],
        )

    def test_mapping_preview_markdown_is_derived_from_csv(self):
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "architecture_mapping_preview.csv"
            csv_path.write_text(
                "candidate_block,requirement_id\nADC,REQ-1\nUnassigned,REQ-2\n",
                encoding="utf-8",
            )
            _write_mapping_preview_markdown(csv_path)
            markdown = csv_path.with_suffix(".md").read_text(encoding="utf-8")
            self.assertIn("| ADC | 1 | REQ-1 |", markdown)
            self.assertIn("| Unassigned | 1 | REQ-2 |", markdown)

    def test_ambiguous_ontology_roles_do_not_create_single_block_owner(self):
        block_defs = {
            "Register Block": {"function": "register access"},
            "Source Block": {"function": "source endpoint"},
            "Destination Block": {"function": "destination endpoint"},
            "Unassigned": {},
        }
        row = {
            "source_req_id": "REQ-MATRIX",
            "source": "Section 4.1 Connection Matrix (under Section 4 Interconnect), paragraph 002",
            "requirement_statement": "Source Block connects to Destination Block",
            "category": "Digital",
            "requirement_type": "functional",
        }
        ontology_links = {
            "REQ-MATRIX": [
                {
                    "role": "Register Block",
                    "related_role": "Source Block; Destination Block",
                    "relation_type": "supports",
                    "function_or_property_evidence": "connection matrix participants",
                }
            ]
        }
        self.assertEqual(
            _map_blocks(row, {}, block_defs, {}, ontology_links),
            ["Unassigned"],
        )

    def test_previous_traceability_reader_accepts_current_and_legacy_schemas(self):
        with tempfile.TemporaryDirectory() as directory:
            current_path = Path(directory) / "current.csv"
            current_path.write_text("source_req_id,block\nREQ-1,Existing Block\n", encoding="utf-8")
            legacy_path = Path(directory) / "legacy.csv"
            legacy_path.write_text("Requirement ID,Block(s)\nREQ-2,Legacy Block\n", encoding="utf-8")
            self.assertEqual(_read_previous_blocks_by_requirement(current_path), {"REQ-1": "Existing Block"})
            self.assertEqual(_read_previous_blocks_by_requirement(legacy_path), {"REQ-2": "Legacy Block"})

    def test_previous_concrete_owner_cascades_over_generated_candidate(self):
        preview_owner_ids = {"Generated Candidate": ["REQ-1"], "Existing Block": []}
        cascaded = _cascade_previous_concrete_owners(
            preview_owner_ids,
            {"REQ-1": "Existing Block"},
            {"Generated Candidate": {}, "Existing Block": {}, "Unassigned": {}},
        )
        self.assertEqual(cascaded, 1)
        self.assertEqual(preview_owner_ids["Generated Candidate"], [])
        self.assertEqual(preview_owner_ids["Existing Block"], ["REQ-1"])

    def test_previous_concrete_owner_def_is_added_when_profile_lacks_it(self):
        block_defs = {"Generated Candidate": {}, "Unassigned": {}}
        added = _ensure_previous_concrete_owner_defs(block_defs, {"REQ-1": "Existing Block"})
        self.assertEqual(added, 1)
        self.assertIn("Existing Block", block_defs)
        self.assertIn("approved upstream traceability", block_defs["Existing Block"]["function"])

    def test_gui_stage_menu_exposes_strengthened_workflow(self):
        keys = {SINGLE_STAGE_LABEL_TO_KEY[label] for label in SINGLE_STAGE_LABELS}
        self.assertEqual(
            keys,
            {"0", "1", "2", "2a", "2f", "3", "4", "5", "6", "7", "arch-compare", "regenerate_workbook"},
        )
        self.assertNotIn("S2B - Merge Impact Analysis", SINGLE_STAGE_LABELS)
        self.assertNotIn("S2H - Immutable Approved Snapshot", SINGLE_STAGE_LABELS)
        self.assertIn("Stage 6 - Digital IPOS", SINGLE_STAGE_LABELS)
        self.assertIn("Stage 7 - Analog IPOS", SINGLE_STAGE_LABELS)
        self.assertEqual(SINGLE_STAGE_LABELS[-1], "Optional Architecture Comparison")
        self.assertEqual(SINGLE_STAGE_LABEL_TO_KEY["Stage 6 - Digital IPOS"], "6")
        self.assertEqual(SINGLE_STAGE_LABEL_TO_KEY["Optional Architecture Comparison"], "arch-compare")
        labels_by_id = {node["id"]: node["label"] for node in WORKFLOW_DIAGRAM_NODES}
        self.assertIn("ARS", labels_by_id["stage4"])
        self.assertIn("DRS", labels_by_id["stage5"])
        self.assertIn("Digital IPOS", labels_by_id["stage6"])
        self.assertIn("Optional", labels_by_id["arch_compare"])
        edge_pairs = {(source, target) for source, target, _kind in WORKFLOW_DIAGRAM_EDGES}
        self.assertIn(("stage4", "stage5"), edge_pairs)
        self.assertIn(("stage5", "stage6"), edge_pairs)
        self.assertIn(("stage6", "stage7"), edge_pairs)
        self.assertIn(("stage5", "arch_compare"), edge_pairs)
        self.assertNotIn(("stage4", "arch_compare"), edge_pairs)
        self.assertNotIn(("stage3", "stage5"), edge_pairs)
        range_keys = {RANGE_STAGE_LABEL_TO_KEY[label] for label in RANGE_STAGE_LABELS}
        self.assertEqual(range_keys, set(STAGE_ORDER))
        self.assertNotIn("arch-compare", range_keys)

    def test_compact_workflow_scrollregion_contains_rightmost_content(self):
        bounds = (0, 0, 1576, 255)
        scrollregion = compact_workflow_scrollregion(bounds, margin=16)
        self.assertEqual(scrollregion, (-16, -16, 1592, 271))

    def test_compact_workflow_final_sysml_follows_ipos_then_readiness_reports(self):
        nodes = {node["id"]: node for node in WORKFLOW_DIAGRAM_NODES}
        edges = {(source, target) for source, target, _kind in WORKFLOW_DIAGRAM_EDGES}
        self.assertIn(("stage6", "stage7"), edges)
        self.assertIn(("stage7", "sysml"), edges)
        self.assertIn(("sysml", "sysml_readiness"), edges)
        self.assertNotIn(("s2e", "sysml"), edges)
        self.assertGreaterEqual(nodes["sysml"]["x"], nodes["stage7"]["x"] + 130)
        self.assertGreaterEqual(nodes["sysml_readiness"]["x"], nodes["sysml"]["x"] + 130)
        self.assertIn("Excel", nodes["sysml_readiness"]["label"])
        self.assertIn("NOT READY", WORKFLOW_NODE_MEANINGS["sysml_readiness"])
        self.assertIn("without rerunning", WORKFLOW_NODE_MEANINGS["sysml"])
        self.assertFalse(any(source in {"sysml", "sysml_readiness"} and target.startswith(("s2", "stage"))
                             for source, target in edges))

    def test_compact_workflow_labels_fit_boxes_at_every_zoom(self):
        root = tk.Tk()
        root.withdraw()
        try:
            gui = object.__new__(WorkflowGUI)
            gui.diagram_canvas = tk.Canvas(root)
            gui._bind_stage_click = MagicMock()
            gui._refresh_stage_styles = MagicMock()
            stage_font_sizes = {}
            for display_scale in (1.0, 1.5, 2.0):
                root.tk.call("tk", "scaling", display_scale)
                for zoom_step in range(6, 21):
                    zoom = zoom_step / 10
                    gui.compact_workflow_zoom = zoom
                    gui._draw_compact_workflow_graph()
                    for node in WORKFLOW_DIAGRAM_NODES:
                        with self.subTest(scale=display_scale, zoom=zoom, node=node["id"]):
                            items = gui.diagram_canvas.find_withtag(node["id"])
                            rectangle = next(item for item in items if gui.diagram_canvas.type(item) == "rectangle")
                            label = next(item for item in items if gui.diagram_canvas.type(item) == "text")
                            box = gui.diagram_canvas.coords(rectangle)
                            bounds = gui.diagram_canvas.bbox(label)
                            padding = 5 * zoom
                            self.assertGreaterEqual(bounds[0], box[0] + padding)
                            self.assertGreaterEqual(bounds[1], box[1] + padding)
                            self.assertLessEqual(bounds[2], box[2] - padding)
                            self.assertLessEqual(bounds[3], box[3] - padding)
                            self.assertEqual(gui.diagram_canvas.itemcget(label, "text"), node["label"])
                            self.assertAlmostEqual(box[2] - box[0], (112 if node["kind"] == "main" else 150) * zoom)
                    if display_scale == 1.0 and zoom in (0.6, 2.0):
                        stage_font_sizes[zoom] = int(root.tk.splitlist(gui.diagram_canvas.itemcget(gui.stage_labels["0"], "font"))[1])
            self.assertGreater(stage_font_sizes[2.0], stage_font_sizes[0.6])
        finally:
            root.update_idletasks()
            root.destroy()

    def test_full_workflow_labels_fit_boxes_at_display_scales(self):
        root = tk.Tk()
        root.withdraw()
        try:
            renderer = WorkflowDiagramRenderer(ttk.Frame(root))
            for display_scale in (1.0, 1.5, 2.0):
                root.tk.call("tk", "scaling", display_scale)
                renderer.refresh()
                for node_id, (rectangle, label) in renderer.node_items.items():
                    with self.subTest(scale=display_scale, node=node_id):
                        box = renderer.canvas.coords(rectangle)
                        bounds = renderer.canvas.bbox(label)
                        self.assertGreaterEqual(bounds[0], box[0] + 5)
                        self.assertGreaterEqual(bounds[1], box[1] + 5)
                        self.assertLessEqual(bounds[2], box[2] - 5)
                        self.assertLessEqual(bounds[3], box[3] - 5)
        finally:
            root.update_idletasks()
            root.destroy()

    def test_compact_workflow_keeps_digital_and_analog_ipos_on_distinct_rows(self):
        nodes = {node["id"]: node for node in WORKFLOW_DIAGRAM_NODES}
        self.assertGreaterEqual(
            compact_workflow_node_y(nodes["stage7"]) - compact_workflow_node_y(nodes["stage6"]),
            42,
        )
        self.assertNotEqual(
            compact_workflow_node_y(nodes["stage4"]),
            compact_workflow_node_y(nodes["stage5"]),
        )

    def test_cli_separates_digital_ipos_from_architecture_comparison(self):
        parser = build_parser()
        digital_args = parser.parse_args(["run", "--stage", "6"])
        comparison_args = parser.parse_args([
            "arch-compare", "--project-to-compare", "sibling-project",
        ])
        self.assertEqual(digital_args.command, "run")
        self.assertEqual(digital_args.stage, "6")
        self.assertFalse(hasattr(digital_args, "project_to_compare"))
        self.assertEqual(comparison_args.command, "arch-compare")
        self.assertEqual(comparison_args.project_to_compare, "sibling-project")

    def test_sysml_coverage_fails_for_omitted_block_requirement(self):
        snapshot = SimpleNamespace(
            rows=({"canonical_id": "CAN-1", "source_req_id": "REQ-1"},),
            source_links={"CAN-1": ("REQ-1",)},
            mapping={"CAN-1": "ADC"},
        )
        with self.assertRaises(SysMLCompletenessError):
            validate_snapshot_requirement_coverage(snapshot, [], {"ADC": []})

    def test_sysml_coverage_accepts_full_block_requirement(self):
        snapshot = SimpleNamespace(
            rows=({"canonical_id": "CAN-1", "source_req_id": "REQ-1"},),
            source_links={"CAN-1": ("REQ-1",)},
            mapping={"CAN-1": "ADC"},
        )
        result = validate_snapshot_requirement_coverage(
            snapshot,
            [],
            {"ADC": [{"source_req_id": "REQ-1"}]},
        )
        self.assertEqual(result["missing_by_block"], {})


if __name__ == "__main__":
    unittest.main()
