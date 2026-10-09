import sys
import unittest
import csv
import json
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from ingest_source_spec import discover_tag_pattern, stage_source, _write_workbook, _sync_saved_review_workbook, migrate_review_schema
from merge_engine import register_staged_ingestion, merge_staging_batch
from canonical_store import connect
from supplementary_source_integrity import bounded_source_bodies, validate_staged_bodies, validate_ingestion_batch


class SourceIngestionTests(unittest.TestCase):
    full_statement = "The unit shall hold its output high until completion is reported."
    prefix = "The unit shall hold its output high until"

    def _batch(self, root, statement=None):
        ingestion = root / "artifacts/source_ingestion/supplement_001"
        ingestion.mkdir(parents=True)
        metadata = {"ingestion_id": "supplement_001", "source_spec": "supplement.pdf", "source_revision": "1"}
        (ingestion / "source_metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
        body = bounded_source_bodies(["[REQ_UNIT_1] REQUIREMENT: " + self.full_statement + " [END]"])[0]
        row = {"staged_id": "supplement_001:REQ_UNIT_1", **body,
               "source_spec": "supplement.pdf", "source_revision": "1"}
        if statement is not None:
            row["requirement_statement"] = statement
        review = {**row, "req_class": "new", "classification": "System", "review_decision": "approved"}
        self._csv(ingestion / "normalized_requirements.csv", [row])
        self._csv(ingestion / "comparison_results.csv", [review])
        self._csv(ingestion / "raw_ocr_pages.csv", [{"source_spec": "supplement.pdf", "page": "1",
                                                   "text": "[REQ_UNIT_1] REQUIREMENT: " + self.full_statement + " [END]"}])
        extracts = ingestion / "ocr_extracts"
        extracts.mkdir()
        text_path = extracts / "supplement_p001.txt"
        text_path.write_text("[REQ_UNIT_1] REQUIREMENT: " + self.full_statement + " [END]\n", encoding="utf-8")
        self._csv(ingestion / "ocr_index.csv", [{"page": "1", "text_file": str(text_path)}])
        _write_workbook(ingestion / "comparison_results.csv")
        return ingestion

    def _csv(self, path, rows):
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)

    def test_discovers_bracketed_tagged_requirement_family(self):
        pattern = discover_tag_pattern([
            "[IPOS_STBIO1_MAIN_CONTROLLER_ 7096] REQUIREMENT:\nThe controller shall enter operative mode.",
            "[IPOS_STBIO1_MAIN_CONTROLLER_5015] Requirement:\nThe controller shall support reset.",
            "PPG signal description",
        ])
        self.assertEqual(pattern, r"IPOS_STBIO1_MAIN_CONTROLLER_\d+")

    def test_rejects_sources_without_tagged_requirement_headers(self):
        with self.assertRaises(ValueError):
            discover_tag_pattern(["The device shall expose a status register."])

    def test_rejects_lexically_complete_prefix_of_full_body(self):
        pages = ["[REQ_UNIT_1] REQUIREMENT: The unit shall hold its output high until\ncompletion is reported. [END]"]
        row = {"source_req_id": "REQ_UNIT_1", "source_page": "1",
               "requirement_statement": "The unit shall hold its output high until"}
        with self.assertRaisesRegex(ValueError, "preservation"):
            validate_staged_bodies([row], pages)

    def test_collects_explicit_cross_page_continuation(self):
        pages = ["[REQ_UNIT_1] REQUIREMENT: The unit shall hold its output high until",
                 "completion is reported. [END]\n[REQ_UNIT_2] REQUIREMENT: The unit shall reset. [END]"]
        bodies = bounded_source_bodies(pages)
        self.assertEqual(bodies[0]["requirement_statement"], "The unit shall hold its output high until completion is reported.")
        self.assertEqual(bodies[0]["source_page"], "1")
        self.assertEqual(bodies[1]["source_page"], "2")
        validate_staged_bodies(bodies, pages)

    def test_never_uses_next_requirement_end_to_complete_previous(self):
        with self.assertRaisesRegex(ValueError, "bounded"):
            bounded_source_bodies(["[REQ_UNIT_1] REQUIREMENT: The unit shall wait.",
                                  "[REQ_UNIT_2] REQUIREMENT: The unit shall reset. [END]"])

    def test_preserves_short_and_inline_terminated_bodies(self):
        bodies = bounded_source_bodies(["[REQ_UNIT_1] REQUIREMENT: Reset shall clear flags. [END]"])
        validate_staged_bodies(bodies, ["[REQ_UNIT_1] REQUIREMENT: Reset shall clear flags. [END]"])
        self.assertEqual(bodies[0]["requirement_statement"], "Reset shall clear flags.")

    def test_cached_prefix_rejected_without_extraction_or_registration(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root, self.prefix)
            before = {str(path.relative_to(ingestion)): path.read_bytes() for path in ingestion.rglob("*") if path.is_file()}
            with patch("ingest_source_spec.PdfReader") as reader, patch("ingest_source_spec.register_staged_ingestion") as register:
                with self.assertRaisesRegex(ValueError, "preservation"):
                    stage_source(root, root / "supplement.pdf", "1")
                reader.assert_not_called()
                register.assert_not_called()
            self.assertEqual(before, {str(path.relative_to(ingestion)): path.read_bytes() for path in ingestion.rglob("*") if path.is_file()})

    def test_cached_workbook_prefix_rejected_before_reuse_or_sync(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root)
            workbook = load_workbook(ingestion / "comparison_results.xlsx")
            worksheet = workbook["Source Ingestion Review"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(2, headers.index("requirement_statement") + 1).value = self.prefix
            workbook.save(ingestion / "comparison_results.xlsx")
            workbook.close()
            before = (ingestion / "comparison_results.csv").read_bytes()
            with self.assertRaisesRegex(ValueError, "review changed"):
                stage_source(root, root / "supplement.pdf", "1")
            with self.assertRaisesRegex(ValueError, "review changed"):
                _sync_saved_review_workbook(ingestion)
            with patch("merge_engine.connect") as connection:
                with self.assertRaisesRegex(ValueError, "review changed"):
                    register_staged_ingestion(root, ingestion)
                connection.assert_not_called()
            self.assertEqual(before, (ingestion / "comparison_results.csv").read_bytes())

    def test_merge_rejects_workbook_only_prefix_before_impact_analysis(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root)
            batch_id = register_staged_ingestion(root, ingestion)
            workbook = load_workbook(ingestion / "comparison_results.xlsx")
            worksheet = workbook["Source Ingestion Review"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(2, headers.index("requirement_statement") + 1).value = self.prefix
            workbook.save(ingestion / "comparison_results.xlsx")
            workbook.close()
            with patch("merge_engine.analyze_merge") as impact:
                with self.assertRaisesRegex(ValueError, "review changed"):
                    merge_staging_batch(root, project_id=root.name, staging_batch_id=batch_id, approved_by="reviewer")
                impact.assert_not_called()

    def test_page_gaps_and_duplicate_source_ids_are_not_valid_continuations(self):
        with tempfile.TemporaryDirectory() as directory:
            ingestion = self._batch(Path(directory))
            self._csv(ingestion / "raw_ocr_pages.csv", [{"source_spec": "supplement.pdf", "page": "2",
                                                       "text": "[REQ_UNIT_1] REQUIREMENT: " + self.full_statement + " [END]"}])
            with self.assertRaisesRegex(ValueError, "missing or duplicated"):
                validate_ingestion_batch(ingestion)
        pages = ["[REQ_UNIT_1] REQUIREMENT: " + self.full_statement + " [END]"] * 2
        with self.assertRaisesRegex(ValueError, "preservation"):
            validate_staged_bodies([{"source_req_id": "REQ_UNIT_1", "source_page": "1",
                                    "requirement_statement": self.full_statement}], pages)

    def test_direct_registration_rejects_prefix_before_database_access(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root, self.prefix)
            with patch("merge_engine.connect") as connection:
                with self.assertRaisesRegex(ValueError, "preservation"):
                    register_staged_ingestion(root, ingestion)
                connection.assert_not_called()

    def test_csv_only_review_and_schema_migration_cannot_reuse_prefix(self):
        with tempfile.TemporaryDirectory() as directory:
            ingestion = self._batch(Path(directory), self.prefix)
            (ingestion / "comparison_results.xlsx").unlink()
            before = (ingestion / "comparison_results.csv").read_bytes()
            with self.assertRaisesRegex(ValueError, "preservation"):
                _sync_saved_review_workbook(ingestion)
            with self.assertRaisesRegex(ValueError, "preservation"):
                migrate_review_schema(ingestion)
            self.assertEqual(before, (ingestion / "comparison_results.csv").read_bytes())
            self.assertFalse((ingestion / "comparison_results.xlsx").exists())

    def test_split_header_discovery_and_collection_preserve_start_page(self):
        pages = ["[REQ_UNIT_", "1] REQUIREMENT: The unit shall report its status. [END]"]
        self.assertEqual(discover_tag_pattern(pages), r"REQ_UNIT_\d+")
        body = bounded_source_bodies(pages)[0]
        self.assertEqual(body["source_page"], "1")
        self.assertEqual(body["source_req_id"], "REQ_UNIT_1")
        validate_staged_bodies([body], pages)

    def test_raw_page_csv_cannot_hide_longer_stored_extracted_body(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root, self.prefix)
            self._csv(ingestion / "raw_ocr_pages.csv", [{"source_spec": "supplement.pdf", "page": "1",
                                                       "text": "[REQ_UNIT_1] REQUIREMENT: " + self.prefix + " [END]"}])
            with self.assertRaisesRegex(ValueError, "extracted text/raw-page mismatch"):
                stage_source(root, root / "supplement.pdf", "1")
            with patch("merge_engine.connect") as connection:
                with self.assertRaisesRegex(ValueError, "extracted text/raw-page mismatch"):
                    register_staged_ingestion(root, ingestion)
                connection.assert_not_called()

    def test_valid_cached_batch_is_read_only_and_completed_merge_is_revalidated(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root)
            before = {str(path.relative_to(ingestion)): path.read_bytes() for path in ingestion.rglob("*") if path.is_file()}
            with patch("ingest_source_spec.PdfReader") as reader:
                self.assertEqual(stage_source(root, root / "supplement.pdf", "1"), ingestion)
                reader.assert_not_called()
            self.assertEqual(before, {str(path.relative_to(ingestion)): path.read_bytes() for path in ingestion.rglob("*") if path.is_file()})
            batch_id = register_staged_ingestion(root, ingestion)
            merge_staging_batch(root, project_id=root.name, staging_batch_id=batch_id, approved_by="reviewer")
            connection = connect(root)
            approved_before = [tuple(row) for row in connection.execute("SELECT * FROM canonical_requirement_revisions")]
            connection.execute("UPDATE staged_requirements SET requirement_text = ?", (self.prefix,))
            connection.commit()
            connection.close()
            with patch("merge_engine.analyze_merge") as impact:
                with self.assertRaisesRegex(ValueError, "body mismatch"):
                    merge_staging_batch(root, project_id=root.name, staging_batch_id=batch_id, approved_by="reviewer")
                impact.assert_not_called()
            connection = connect(root)
            try:
                self.assertEqual(approved_before, [tuple(row) for row in connection.execute("SELECT * FROM canonical_requirement_revisions")])
            finally:
                connection.close()

    def test_full_body_reaches_canonical_revision_without_truncation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root)
            batch_id = register_staged_ingestion(root, ingestion)
            merge_staging_batch(root, project_id=root.name, staging_batch_id=batch_id, approved_by="reviewer")
            connection = connect(root)
            try:
                row = connection.execute("SELECT requirement_text FROM canonical_requirement_revisions").fetchone()
                self.assertEqual(row[0], self.full_statement)
            finally:
                connection.close()

    def test_direct_merge_and_reregistration_reject_stored_prefix(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ingestion = self._batch(root)
            batch_id = register_staged_ingestion(root, ingestion)
            connection = connect(root)
            connection.execute("UPDATE staged_requirements SET requirement_text = ?", (self.prefix,))
            connection.commit()
            connection.close()
            with patch("merge_engine.analyze_merge") as impact:
                with self.assertRaisesRegex(ValueError, "body mismatch"):
                    merge_staging_batch(root, project_id=root.name, staging_batch_id=batch_id, approved_by="reviewer")
                impact.assert_not_called()
            with self.assertRaisesRegex(ValueError, "body mismatch"):
                register_staged_ingestion(root, ingestion)
            connection = connect(root)
            try:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM canonical_requirements").fetchone()[0], 0)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM merge_operations").fetchone()[0], 0)
            finally:
                connection.close()

    def test_new_extraction_keeps_changed_content_pending_review(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            previous = self._batch(root, self.prefix)
            metadata = json.loads((previous / "source_metadata.json").read_text(encoding="utf-8"))
            metadata["source_revision"] = "0"
            (previous / "source_metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
            primary = root / "artifacts/stage1_requirements/requirements_summary.csv"
            primary.parent.mkdir(parents=True)
            self._csv(primary, [{"id": "BASE_1", "requirement_statement": self.prefix}])
            pages = ["[REQ_UNIT_1] REQUIREMENT: The unit shall hold its output high until",
                     "completion is reported. [END]"]
            with patch("ingest_source_spec.PdfReader", return_value=SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in pages])):
                ingestion = stage_source(root, root / "supplement.pdf", "1")
            validated = validate_ingestion_batch(ingestion)
            self.assertEqual(validated[0]["requirement_statement"], self.full_statement)
            with (ingestion / "comparison_results.csv").open(encoding="utf-8") as handle:
                review = next(csv.DictReader(handle))
            self.assertEqual(review["review_decision"], "pending")
            with self.assertRaises(PermissionError):
                merge_staging_batch(root, project_id=root.name, staging_batch_id="stage-" + ingestion.name.replace("_", "-"), approved_by="reviewer")


if __name__ == "__main__":
    unittest.main()