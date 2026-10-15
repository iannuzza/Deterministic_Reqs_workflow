import importlib
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
reporting = importlib.import_module("generate_cross_spec_issue_report")
contracts = importlib.import_module("spec_document_contract")


class CrossSpecIssueReportTests(unittest.TestCase):
    def test_unverified_saved_review_is_not_presented_as_current_blockers(self):
        review = {"snapshot_id": "snapshot-test", "current_evidence": {"status": "UNVERIFIED",
            "reason": "No matching completed-run evidence"}, "decision_summary": [
            {"block": "Module", "document_type": "Digital IPOS", "current_status": "BLOCKED"}]}
        text = reporting.remaining_ipos_blocker_markdown(review)
        self.assertIn("UNVERIFIED", text)
        self.assertIn(reporting.markdown_literal("No matching completed-run evidence"), text)
        self.assertNotIn("Listed blocks remain BLOCKED", text)
        self.assertNotIn("## Module", text)

    def test_concrete_blocker_claim_rows_keep_ids_only_in_appendix(self):
        review = {"snapshot_id": "snapshot-test", "decision_summary": [
            {"block": "Module", "document_type": "Digital IPOS", "current_status": "BLOCKED"}],
            "clean_generation_observations": {"blocks": [{"block": "Module", "document_type": "Digital IPOS",
                "contract_report": {"checks": {"semantic_fidelity": {"findings": [
                    "ipos_final_unmodeled_content:audit:private-claim-id",
                    "ipos_final_claim_missing_or_changed:audit:private-claim-id"]}}}}]}}
        details = {"Module": {"claims": {"private-claim-id": {"section": "1.4 Supported block functions",
            "name": "masked contribution", "meaning": "If masked, the contribution is zero."}}}}
        text = reporting.remaining_ipos_blocker_markdown(review, details)
        self.assertIn("masked contribution", text)
        self.assertIn(reporting.markdown_literal("If masked, the contribution is zero."), text)
        self.assertNotIn("private-claim-id", text)
        self.assertIn("does not isolate the lost fact", text)
        self.assertIn("Detailed evidence", text)

    def test_remaining_blockers_are_short_current_and_project_agnostic(self):
        review = {
            "snapshot_id": "snapshot-test",
            "minimum_shared_or_profile_adjustments": {"block_preflight": []},
            "decision_summary": [
                {"block": "Module A", "document_type": "Digital IPOS", "current_status": "BLOCKED"},
                {"block": "Module B", "document_type": "Digital IPOS", "current_status": "PASS"},
            ],
            "issues": [
                {"block": "Module A", "identity": "current-probe:semantic_fidelity", "blocking_status": "blocking",
                 "issue_class": "parser_or_projection_limitation", "user_clarification_needed": "none"},
                {"block": "Module A", "identity": "old", "blocking_status": "blocking", "historical_only": True,
                 "reason": "HISTORICAL DUMP", "issue_class": "missing_source_information", "user_clarification_needed": "yes"},
            ],
            "clean_generation_observations": {"blocks": [
                {"block": "Module A", "document_type": "Digital IPOS", "contract_report": {"checks": {
                    "semantic_fidelity": {"findings": [
                        "ipos_final_claim_missing_or_changed:module.md:1.3",
                        "ipos_final_claim_missing_or_changed:module.docx:1.3",
                    ]}}}},
            ]},
        }
        details = {"Module A": {"sections": {"1.3": {"title": "1.3 Internal blocks or functions",
                                 "meaning": "Process samples into ordered results"}}}}
        text = reporting.remaining_ipos_blocker_markdown(review, details)
        self.assertIn("## Module A", text)
        self.assertNotIn("Module B", text)
        self.assertNotIn("HISTORICAL DUMP", text)
        self.assertEqual(text.count(reporting.markdown_literal("1.3 Internal blocks or functions")), 1)
        self.assertIn("Expected supported meaning", text)
        self.assertIn("Process samples into ordered results", text)
        self.assertIn("section mapping/readback", text)
        self.assertLess(len(text.splitlines()), 20)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / "artifacts/orchestrator"
            folder.mkdir(parents=True)
            import json
            source = folder / "descriptive_obligation_review.json"
            source.write_text(json.dumps(review), encoding="utf-8")
            before = source.read_bytes()
            current = {"status": "BLOCKED", "blocks": [{"block": "Module A", "status": "BLOCKED",
                "reason": "Missing/changed supported claim: Process samples into ordered results"}]}
            with patch.object(reporting, "current_ipos_blocker_evidence", return_value=current):
                output = reporting.write_remaining_ipos_blocker_report(root, "snapshot-test")
            self.assertEqual(output.read_text(encoding="utf-8"),
                reporting.remaining_ipos_blocker_markdown({**review, "current_evidence": current}))
            self.assertTrue(output.with_suffix(".docx").is_file())
            word_text = "\n".join(event.get("text", "") if event["kind"] != "table" else " ".join(cell for row in event["rows"] for cell in row)
                                  for event in contracts.drs_document_events(output.with_suffix(".docx")))
            self.assertIn("Module A", word_text)
            self.assertNotIn("Module B", word_text)
            self.assertIn("Missing/changed supported claim", word_text)
            self.assertIn("Process samples into ordered results", word_text)
            self.assertEqual(source.read_bytes(), before)
            with self.assertRaises(ValueError):
                reporting.write_remaining_ipos_blocker_report(root, "other-snapshot")

    def test_current_evidence_is_bound_to_snapshot_artifacts_and_run_scope(self):
        import json
        import time
        from contextlib import ExitStack

        real_binding = contracts.validate_materialization_chain
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            root = Path(directory)
            (root / "config").mkdir()
            (root / "config/project_context.json").write_text('{"project_name":"Generic"}', encoding="utf-8")
            orchestrator = root / "artifacts/orchestrator"
            orchestrator.mkdir(parents=True)
            results_path = orchestrator / "ipos_generation_results.json"
            result = {"snapshot_id": "snapshot-test", "scope": "all", "blocks": [
                {"block": "Module", "status": "PASS"}, {"block": "Other", "status": "PASS"}]}
            results_path.write_text(json.dumps(result), encoding="utf-8")
            report = {"status": "PASS", "residuals": [{"authority": False}],
                      "checks": {"semantic_fidelity": {"findings": []}}}
            for name in ("Module", "Other"):
                folder = root / reporting.ipos.CONFIG["digital"]["output_dir"] / "blocks" / reporting.ipos._block_slug(name)
                folder.mkdir(parents=True)
                (folder / reporting.ipos._block_markdown_name("digital", name)).write_text(
                    "- Snapshot: `snapshot-test`\n", encoding="utf-8")
                for filename in (reporting.ipos._block_docx_name("digital", name),
                                 "descriptive_summary_audit.csv", "materialization_audit.json"):
                    (folder / filename).write_text("fixture", encoding="utf-8")
                (folder / "document_contract_report.json").write_text(json.dumps(report), encoding="utf-8")
            contract = SimpleNamespace(target=lambda source_id: "Digital IPOS", descriptive_context=lambda *args, **kwargs: ())
            stack.enter_context(patch.object(reporting.ipos, "resolve_complete_authoritative_input",
                return_value=(SimpleNamespace(snapshot_id="snapshot-test"), {})))
            stack.enter_context(patch.object(reporting.ipos, "resolve_downstream_contract", return_value=contract))
            stack.enter_context(patch.object(reporting, "approved_snapshot_candidate_rows", side_effect=
                lambda snapshot, contract, domain: [{"mapped_block": name, "source_req_id": name} for name in ("Module", "Other")]
                    if domain == "digital" else []))
            stack.enter_context(patch.object(reporting, "build_ipos_functional_input", return_value=object()))
            stack.enter_context(patch("ipos_blocker_review.apply_descriptive_omissions", side_effect=lambda root, **kwargs: kwargs["functional_input"]))
            stack.enter_context(patch("workflow_routing.build_ipos_normalized_records", return_value=()))
            binding = stack.enter_context(patch.object(contracts, "validate_materialization_chain", return_value=[]))
            checker = stack.enter_context(patch.object(contracts, "ipos_document_contract_report", return_value=report))
            evidence = reporting.current_ipos_blocker_evidence(root, "snapshot-test")
            self.assertEqual(evidence["status"], "PASS")
            self.assertEqual(len(evidence["blocks"]), 2)
            self.assertTrue(all(row["status"] == "PASS" for row in evidence["blocks"]))
            self.assertIn("residuals: 1", evidence["blocks"][0]["reason"])
            binding.return_value = ["content changed after audited materialization"]
            self.assertEqual(reporting.current_ipos_blocker_evidence(root, "snapshot-test")["status"], "UNVERIFIED")
            binding.return_value = []
            checker.return_value = {"status": "FAIL", "checks": {"semantic_fidelity": {
                "findings": ["ipos_final_claim_missing_or_changed:1.4"]}}}
            self.assertEqual(reporting.current_ipos_blocker_evidence(root, "snapshot-test")["status"], "BLOCKED")
            checker.return_value = report
            result["blocks"][0].update(status="BLOCKED", findings=["ipos_publication_failure:access denied"])
            results_path.write_text(json.dumps(result), encoding="utf-8")
            blocked = reporting.current_ipos_blocker_evidence(root, "snapshot-test")
            self.assertEqual(blocked["status"], "BLOCKED")
            self.assertIn("access denied", blocked["blocks"][0]["reason"])
            result["snapshot_id"] = "another-snapshot"
            results_path.write_text(json.dumps(result), encoding="utf-8")
            self.assertEqual(reporting.current_ipos_blocker_evidence(root, "snapshot-test")["status"], "UNVERIFIED")
            result["snapshot_id"] = "snapshot-test"
            result["blocks"][0]["status"] = "PASS"
            results_path.write_text(json.dumps(result), encoding="utf-8")
            run = {"stages": ["6"], "block": "Module", "started_at": 0,
                   "finished_at": time.time() + 1, "exit_code": 0}
            scoped = reporting.current_ipos_blocker_evidence(root, "snapshot-test", run)
            self.assertEqual([row["status"] for row in scoped["blocks"]], ["PASS", "UNVERIFIED"])
            run.update(block=None, started_at=time.time() + 2, finished_at=time.time() + 3)
            self.assertEqual(reporting.current_ipos_blocker_evidence(root, "snapshot-test", run)["status"], "UNVERIFIED")
            from docx import Document
            for name in ("Module", "Other"):
                folder = root / reporting.ipos.CONFIG["digital"]["output_dir"] / "blocks" / reporting.ipos._block_slug(name)
                markdown = folder / reporting.ipos._block_markdown_name("digital", name)
                Document().save(markdown.with_suffix(".docx"))
                contracts.write_materialization_audit(folder / "materialization_audit.json", "IPOS", (), (),
                    markdown, markdown.with_suffix(".docx"))
            binding.side_effect = real_binding
            self.assertEqual(reporting.current_ipos_blocker_evidence(root, "snapshot-test")["status"], "PASS")
            markdown.write_text("- Snapshot: `snapshot-test`\nChanged after audit\n", encoding="utf-8")
            changed = reporting.current_ipos_blocker_evidence(root, "snapshot-test")
            self.assertEqual(changed["status"], "UNVERIFIED")
            self.assertTrue(any("content changed" in finding for row in changed["blocks"] for finding in row["findings"]))

    def test_new_current_pass_supersedes_saved_blockers_without_editing_history(self):
        import json
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / "artifacts/orchestrator"
            folder.mkdir(parents=True)
            review = {"snapshot_id": "snapshot-test", "decision_summary": [
                {"block": "Module", "document_type": "Digital IPOS", "current_status": "BLOCKED"}]}
            source = folder / "descriptive_obligation_review.json"
            source.write_text(json.dumps(review), encoding="utf-8")
            before = source.read_bytes()
            current = {"status": "PASS", "blocks": [{"block": "Module", "status": "PASS",
                "reason": "Current snapshot-bound document checks PASS; residuals: 0"}]}
            with patch.object(reporting, "current_ipos_blocker_evidence", return_value=current):
                output = reporting.write_remaining_ipos_blocker_report(root, "snapshot-test")
            text = output.read_text(encoding="utf-8")
            word_text = "\n".join(event.get("text", "") if event["kind"] != "table" else
                " ".join(cell for row in event["rows"] for cell in row)
                for event in contracts.drs_document_events(output.with_suffix(".docx")))
            for rendered in (text, word_text):
                self.assertIn("Current document status: PASS", rendered)
                self.assertIn("NOT_ASSESSED", rendered.replace("\\_", "_"))
                self.assertNotIn("remain BLOCKED", rendered)
            self.assertEqual(source.read_bytes(), before)

    def test_actual_current_fidelity_failure_remains_blocking_not_missing_source(self):
        value = {"reason": "ipos_final_claim_missing_or_changed:audit:claim", "current_failure_layer": "semantic_fidelity", "blocking": True}
        reporting.populate_report_issue_notes(value)
        self.assertEqual(value["blocking_status"], "blocking")
        self.assertEqual(value["issue_class"], "parser_or_projection_limitation")
        self.assertEqual(value["user_clarification_needed"], "none")
        self.assertIn("Approved meaning is present", value["missing_or_unresolved_information"])
        self.assertIn("must not be published blindly", value["blocking_reason"])

    def test_tightened_trace_separates_report_only_history_from_effective_blockers(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Process samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "structured", "statement": "0xA000 CONTROL_REG 32 0xA040 reset_n"}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        previous = {"true_descriptive_blockers": [{"document_type": "Digital IPOS", "block": "Module", "identity": "old:1",
                    "reason": "old table blocker", "source_spans": review["issues"][0]["source_spans"], "provenance": ["approved source"]}]}
        reporting.finalize_descriptive_review(review, previous, [functional])
        review["snapshot_id"] = "test-snapshot"
        with tempfile.TemporaryDirectory() as directory:
            reporting.attach_ipos_decision_trace(Path(directory), review)
        self.assertEqual(len(review["removed_from_active_blocking_set"]), 1)
        self.assertEqual(review["removed_from_active_blocking_set"][0]["identity"], "old:1")
        self.assertEqual(review["decision_summary"][0]["effective_descriptive_blocker_ids"], [])
        review["clean_generation_observations"] = {"blocks": [{"document_type": "Digital IPOS", "block": "Module", "probe_status": "BLOCKED"}]}
        self.assertEqual(reporting.ipos_decision_summary(review)[0]["clean_generation_possible_now"], "no")

    def test_decision_main_starts_with_summary_and_keeps_trace_in_appendix(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Process samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "detail", "statement": "This signal governs the operation as previously explained."}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        reporting.finalize_descriptive_review(review, {}, [functional])
        main = reporting.descriptive_obligation_review_markdown(review)
        appendix = reporting.descriptive_obligation_review_appendix_markdown(review)
        self.assertTrue(main.startswith("| IPOS block |"))
        self.assertIn("| partial | no |", main)
        self.assertNotIn("This signal governs", main)
        self.assertIn("This signal governs", appendix)
        self.assertEqual(reporting.ipos_decision_summary(review)[0]["effective_descriptive_blocker_ids"], [])
        review["minimum_shared_or_profile_adjustments"]["block_preflight"][0]["publication_status"] = "PASS"
        self.assertEqual(reporting.ipos_decision_summary(review)[0]["clean_generation_possible_now"], "yes")

    def test_fixed_notes_stale_artifact_is_blocking_without_claiming_missing_meaning(self):
        blocks = [{"document_type": "Digital IPOS", "block": "Module", "publication_status": "FAIL_EXISTING_BLOCK_GATE",
                   "validation_output": "ipos_descriptive_audit_invalid_or_changed:old-claim"}]
        value = reporting.build_artifact_alignment_issues(blocks, "test-snapshot")[0]
        self.assertEqual(value["issue_class"], "stale_alignment")
        self.assertEqual(value["blocking_status"], "blocking")
        self.assertEqual(value["user_clarification_needed"], "none")
        self.assertIn("not shown to be missing", value["missing_or_unresolved_information"])
        self.assertIn("does not regenerate them or waive the gate", value["blocking_reason"])
        reporting.populate_report_issue_notes(value, historical=True)
        self.assertEqual(value["blocking_status"], "report-only")
        self.assertFalse(value["blocking"])

    def test_fixed_notes_distinguish_parser_missing_information_and_structured_content(self):
        parser = {"reason": "unsupported multi-clause or non-normative construction", "blocking": False,
                  "problem_category": reporting.DESCRIPTIVE_ISSUE_CATEGORIES[0],
                  "source_spans": [{"text": "If timer expires the operation is aborted and proper flag is activated."}]}
        missing = {"reason": "catalog description absent", "approved_port": {"Port name": "o_DONE"},
                   "confirmed_missing_information": "The approved interface row has no output-generation description for o_DONE",
                   "confirmed_missing_scope": "the approved interface row",
                   "confirmed_user_clarification": "State when o_DONE is generated and cleared.",
                   "confirmed_clarification_type": "clarify_output_generation"}
        structured = {"reason": "nonprose_output_path:register", "output_path": "register",
                      "problem_category": reporting.DESCRIPTIVE_ISSUE_CATEGORIES[1]}
        for value in (parser, missing, structured):
            reporting.populate_report_issue_notes(value)
            reporting.validate_report_issue_notes(value)
            self.assertTrue(all(field in value for field in reporting.REPORT_NOTE_FIELDS))
        self.assertEqual(parser["issue_class"], "parser_or_projection_limitation")
        self.assertIn("Approved meaning is present", parser["missing_or_unresolved_information"])
        self.assertEqual(parser["user_clarification_needed"], "none")
        self.assertEqual(missing["issue_class"], "missing_source_information")
        self.assertEqual(missing["blocking_status"], "report-only")
        self.assertEqual(missing["clarification_type"], "clarify_output_generation")
        self.assertEqual(structured["issue_class"], "non_prose_misclassification")
        self.assertEqual(structured["blocking_status"], "non-blocking")
        self.assertEqual(structured["user_clarification_needed"], "none")
        with self.assertRaises(ValueError):
            reporting.validate_report_issue_notes(dict(parser, clarification_type="guess_actor"))
        with self.assertRaises(ValueError):
            reporting.validate_report_issue_notes(dict(parser, clarification_type="make_actor_explicit"))
        with self.assertRaises(ValueError):
            reporting.validate_report_issue_notes(dict(parser, issue_class="missing_source_information"))

    def test_v2_review_preserves_previous_evidence_and_keeps_required_gaps_blocking(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "", "Inputs": "samples", "Outputs": "results"}, [], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        self.assertTrue(review["essential_issue_set"])
        self.assertTrue(any(value["blocking"] and value["problem_category"] == reporting.DESCRIPTIVE_ISSUE_CATEGORIES[4]
                            for value in review["issues"]))
        previous = {"true_descriptive_blockers": [{"document_type": "Digital IPOS", "block": "Module", "identity": "disposition:old",
                    "reason": "old projection", "source_spans": [], "provenance": ["approved source"],
                    "diagnostic_notes": {"missing_from_spec": "Not determined."}}]}
        reporting.finalize_descriptive_review(review, previous, [functional])
        self.assertEqual(review["previous_review"], previous)
        self.assertEqual(review["reclassified_issue_set"][0]["previous_identity"], "disposition:old")
        report = self.user_report_fixture()
        old = dict(report["issues"][0])
        reporting.apply_current_descriptive_review(report, review)
        self.assertEqual({key: report["issues"][0][key] for key in old}, old)
        self.assertEqual(report["issues"][0]["blocking_status"], "report-only")
        self.assertTrue(all(field in report["issues"][0] for field in reporting.REPORT_NOTE_FIELDS))
        self.assertEqual(report["historical_issue_count"], 1)
        self.assertIn("Current IPOS Assessment", reporting.report_markdown(report))
        count = report["issue_count"]
        reporting.apply_current_descriptive_review(report, review)
        self.assertEqual(report["issue_count"], count)

    def test_v2_signal_review_is_scoped_and_does_not_equate_mentions_with_behavior(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Process samples", "Inputs": "i_CONFIG: configuration", "Outputs": "results"},
            [{"source_req_id": "description", "statement": "i_CONFIG is the configuration field."}], domain="digital")
        ports = [{"Owner": "MODULE", "Port name": "i_CONFIG", "Direction": "input", "Ownership status": "approved", "Type / details": "Configuration field"},
                 {"Owner": "Elsewhere", "Port name": "other", "Ownership status": "approved"},
                 {"Owner": "Module", "Port name": "unapproved", "Ownership status": "pending"}]
        report = reporting.build_missing_information_report([functional], ports)
        self.assertEqual(len(report["port_assessments"]), 1)
        self.assertEqual(len(report["issues"]), 1)
        point = report["issues"][0]
        self.assertFalse(point["blocking"])
        self.assertEqual(point["problem_category"], reporting.DESCRIPTIVE_ISSUE_CATEGORIES[2])
        self.assertIn("source-wide absence is not established", point["unresolved_or_missing_information"])
        self.assertEqual(point["source_spans"][0]["text"], "i_CONFIG is the configuration field.")
        self.assertIn("not a complete semantic proof", reporting.missing_information_markdown(report))

    def test_v2_review_audits_optional_omissions_in_exactly_six_categories(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Process frames", "Inputs": "frames", "Outputs": "results"},
            [{"source_req_id": "unknown", "statement": "This signal governs the operation as previously explained."}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        self.assertEqual(tuple(review["categories"]), reporting.DESCRIPTIVE_ISSUE_CATEGORIES)
        self.assertEqual(len(review["issues"]), 1)
        point = review["issues"][0]
        self.assertEqual(point["problem_category"], reporting.DESCRIPTIVE_ISSUE_CATEGORIES[5])
        self.assertFalse(point["blocking"])
        self.assertEqual(point["disposition"], "excluded")
        self.assertEqual(point["source_spans"][0]["text"], "This signal governs the operation as previously explained.")
        for field in ("blocking_reason", "unresolved_or_missing_information", "clarification_need", "diagnostic_notes"):
            self.assertTrue(point[field])
        self.assertFalse(review["authority"])
        self.assertFalse(review["gate_override"])

    def test_relative_clause_notes_identify_existing_antecedent_not_missing_signal(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Control sampling", "Inputs": "configuration", "Outputs": "samples"},
            [{"source_req_id": "duration", "statement":
              "The sampling phase shall be configurable by the user with i_CYCLES values from 0 up to 15. This signal indicates the number of clocks that o_SAMPLE_start shall stay high."}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        point = next(value for value in review["issues"] if value["recovery_status"] == "validated_shared_interpretation")
        notes = point["diagnostic_notes"]
        self.assertFalse(point["blocking"])
        self.assertEqual(point["problem_category"], reporting.DESCRIPTIVE_ISSUE_CATEGORIES[0])
        self.assertIn("exact identifiers, quantities and dependency linkage are retained", notes["missing_from_spec"])
        self.assertIn("i_CYCLES", notes["what_spec_already_states"])
        self.assertIn("from 0 up to 15", notes["what_spec_already_states"])
        self.assertIn("o_SAMPLE_start", notes["what_spec_already_states"])
        self.assertIn("No source amendment", notes["user_action"])
        self.assertIn("no unresolved obligation remains", notes["blocking_mechanism"])

    def test_descriptive_notes_distinguish_parser_limits_from_missing_flag_bindings(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Process transactions", "Inputs": "transactions", "Outputs": "status"},
            [{"source_req_id": "timeout", "statement":
              "If timer expires the operation is aborted and proper flag is activated."}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        notes = review["issues"][0]["diagnostic_notes"]
        self.assertEqual(review["issues"][0]["problem_category"], reporting.DESCRIPTIVE_ISSUE_CATEGORIES[0])
        self.assertEqual(review["issues"][0]["recovery_status"], "candidate_only_not_implemented")
        self.assertIn("no shall/must marker", notes["blocking_mechanism"])
        self.assertIn("does not identify the concrete signal", notes["signal_register_information"])
        self.assertIn("Source-wide absence has not been established", notes["signal_register_information"])
        self.assertIn("Approved block-local scope: Module", notes["application_scope"])
        self.assertIn("preserving every condition", notes["required_resolution"])
        self.assertIn("not needed to describe the stated behavior generically", notes["missing_from_spec"])
        self.assertIn("which signal or register bit", notes["user_action"])
        markdown = reporting.descriptive_obligation_review_appendix_markdown(review)
        for label in ("Diagnostic note", "What the specification already states", "What is missing or not identified", "What the user needs to do", "Blocking mechanism", "Signal/register information", "Where the rule applies", "Required resolution"):
            self.assertIn(label, markdown)

    def test_procedural_notes_do_not_promote_list_numbers_to_fixed_values(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Process frames", "Inputs": "frames", "Outputs": "results"},
            [{"source_req_id": "procedure", "statement":
              "The frame shall work as follows: 1. Select how many times a particular frame shall be repeated 2. The result shall be accumulated."}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        point = next(value for value in review["issues"] if value["reason"] == "no_supported_claim")
        self.assertIn("procedural Select how phrase", point["diagnostic_notes"]["blocking_mechanism"])
        self.assertIn("controlling signal/register field", point["diagnostic_notes"]["signal_register_information"])
        self.assertIn("may be a procedure-step marker", point["diagnostic_notes"]["signal_register_information"])
        self.assertIn("instead of bypassing the filter", point["diagnostic_notes"]["required_resolution"])
        self.assertEqual(point["source_spans"][0]["text"], "Select how many times a particular frame shall be repeated 2.")

    def test_descriptive_review_has_exactly_six_categories_and_preserves_optional_gaps(self):
        functional = reporting.build_ipos_functional_input("Module",
            {"Function": "Acquire samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "register", "statement": "0xA000 CONTROL_REG 32 0xA040 reset_n"},
             {"source_req_id": "behavior", "statement": "When enabled, it shall acquire samples."}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        self.assertEqual(tuple(review["categories"]), reporting.DESCRIPTIVE_ISSUE_CATEGORIES)
        structured = [value for value in review["issues"] if value["problem_category"] == reporting.DESCRIPTIVE_ISSUE_CATEGORIES[1]]
        self.assertEqual(len(structured), 1)
        self.assertEqual(structured[0]["output_path"], "register")
        self.assertTrue(any("behavior" in span["record_id"] for value in review["issues"]
                            for span in value["source_spans"]))
        adjustment = review["minimum_shared_or_profile_adjustments"]
        self.assertFalse(adjustment["authority"])
        self.assertFalse(adjustment["gate_override"])
        self.assertFalse(adjustment["specifications_regenerated"])
        self.assertEqual(adjustment["block_preflight"][0]["preflight_status"], "PASS")
        markdown = reporting.descriptive_obligation_review_appendix_markdown(review)
        self.assertEqual(markdown.count("\n## "), 8)
        self.assertIn("0xA000 CONTROL_REG 32 0xA040 reset_n", markdown)

    def test_descriptive_review_reports_supported_normalized_actor_without_certification(self):
        functional = reporting.build_ipos_functional_input("Module Core",
            {"Function": "Acquire samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "guard", "statement": "During startup MODULE_CORE shall acquire samples."}], domain="digital")
        review = reporting.build_descriptive_obligation_review([functional])
        self.assertFalse(review["essential_issue_set"])
        recoveries = [value for value in review["issues"] if value["problem_category"] == reporting.DESCRIPTIVE_ISSUE_CATEGORIES[0]]
        self.assertEqual(len(recoveries), 1)
        self.assertEqual(review["minimum_shared_or_profile_adjustments"]["block_preflight"][0]["preflight_status"], "PASS")
        self.assertTrue(recoveries[0]["retrieval_support"])
        self.assertNotIn("certified", review)

    def test_presentation_omits_review_notes_without_guessing_missing_content(self):
        original = "The interface accepts samples. External placement is not specified."
        self.assertEqual(contracts.clean_descriptive_review_text(original), "The interface accepts samples.")
        self.assertEqual(contracts.clean_descriptive_review_text("Flow control: need clarification"), "")
        self.assertEqual(contracts.clean_descriptive_review_text("- TBD-001: Review the owner."), "")
        self.assertEqual(contracts.clean_descriptive_review_text(
            "Gated; control expression: u_pm u.enable; signal spelling: need clarification"), "")
        supported = "When reset is low, the output remains low."
        self.assertEqual(contracts.clean_descriptive_review_text(supported), supported)

    def user_report_fixture(self):
        finding = reporting.issue("SRS", "Section | 3", ["REQ-1"], "projection_gap",
                                  "Guard not modeled.", ["source.md#L1"])
        finding.update(issue_id="REVIEW-0001", source_excerpt="The unit shall wait if ready.\n```source```")
        finding["source_spans"] = [{"record_id": "REQ-1", "start": 31, "end": 100,
                       "text": "The unit shall wait if ready."}]
        return {"authority": False, "user_review_only": True, "snapshot_id": "test-snapshot", "run_start": "test-run",
                "issue_count": 1, "issue_counts_by_document": {"SRS": 1}, "issues": [finding],
                "documents": [{"document_type": "SRS", "gate_status": "PASS"}]}

    def test_clean_presentation_records_omissions_and_protects_normative_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "spec.md"
            lines = ["Snapshot ID: approved", "## 8. Assumptions and TBD", "Flow control: need clarification",
                     "| Device | An interrupt triggers wake-up. Local function: need clarification. |",
                     "[DRS-REQ-001]", "The clock shall run at 16 MHz.", "[End]"]
            projected = contracts.apply_clean_document_presentation(lines, path, "DRS")
            self.assertEqual(projected[-3:], lines[-3:])
            self.assertIn("An interrupt triggers wake-up.", projected[3])
            self.assertFalse(contracts.PUBLICATION_REVIEW_MARKERS.search("\n".join(projected)))
            audit = importlib.import_module("json").loads(path.with_name("descriptive_publication_audit.json").read_text())
            self.assertEqual(len(audit["omissions"]), 3)
            self.assertFalse(audit["authority"])
            with self.assertRaises(ValueError):
                contracts.apply_clean_document_presentation(["[DRS-REQ-001]", "The clock shall run: TBD."], path, "DRS")

    def test_accounting_cannot_drop_supported_meaning_or_unreported_notes(self):
        original = "The clock runs at 16 MHz. Timing: need clarification."
        audit = {"omissions": [{"original": original}]}
        self.assertTrue(contracts.publication_text_is_accounted(original, "The clock runs at 16 MHz.", "The clock runs at 16 MHz.", audit))
        self.assertFalse(contracts.publication_text_is_accounted(original, "The clock runs.", "The clock runs.", audit))
        self.assertFalse(contracts.publication_text_is_accounted(original, "The clock runs at 16 MHz.", "The clock runs at 16 MHz.", {}))

    def test_mixed_supported_prose_is_arranged_before_tables(self):
        contract = contracts.load_drs_contract(Path(__file__).resolve().parents[1] / contracts.DRS_TEMPLATE)
        mixed = "The clock runs at 16 MHz. Timing: need clarification."
        lines = []
        for title in contract["sections"]:
            lines.append("## " + title)
            if title.startswith("4."):
                lines.extend(["### 4.7 Clock and synchronization across domains", "| Clock | Frequency |",
                              "| --- | --- |", "| ClockA | 16 MHz |", mixed])
        arranged = contracts.arrange_drs_contract_lines(lines, contract)
        self.assertLess(arranged.index(mixed), arranged.index("| Clock | Frequency |"))

    def test_user_markdown_preserves_issue_fields_and_escapes_evidence(self):
        report = self.user_report_fixture()
        report["issues"][0].update(audit_decision="rejected", disposition="unresolved",
                                  rendered_excerpt="Generated evidence.", retained_excerpt="Supported evidence.",
                                  proposed_clarification="Review the guard.")
        markdown = reporting.report_markdown(report)
        self.assertEqual(markdown.count("### REVIEW-0001"), 1)
        self.assertIn("For user review only", markdown)
        for label in ("Consolidated Descriptive Issue Report", "Observed run start", "Total entries",
                      "Document Summary", "IPOS Block Preflight", "Allowed Conclusions", "Document",
                      "Location", "Blocker class", "Explanation", "Conclusion", "Requirement or context references",
                      "Source references", "Audit decision", "Disposition", "Source excerpt",
                      "Generated document excerpt", "Retained supported content", "Source fragment",
                      "Candidate proposal for user review; NOT approved"):
            self.assertIn(label, markdown)
        self.assertNotRegex(markdown, r"\b(?:centralizzato|revisione|Documento|Posizione|Spiegazione|Conclusione|Riferimenti|Estratto|sorgente|Frammento|Proposta)\b")
        self.assertIn("Section \\| 3", markdown)
        self.assertIn("REQ\\-1", markdown)
        self.assertIn("````text\n" + report["issues"][0]["source_excerpt"] + "\n````", markdown)
        with self.assertRaises(ValueError):
            reporting.report_markdown({**report, "authority": True})

    def test_ipos_publication_observation_requires_fresh_files_real_gate_and_no_markers(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / reporting.ipos.CONFIG["digital"]["output_dir"] / "blocks" / "controller"
            folder.mkdir(parents=True)
            markdown = folder / reporting.ipos._block_markdown_name("digital", "Controller")
            markdown.write_text("Supported behavior.", encoding="utf-8")
            markdown.with_suffix(".docx").write_bytes(b"fixture")
            with patch.object(reporting.subprocess, "run") as run_gate, patch.object(reporting, "drs_document_events") as events:
                events.return_value = [{"kind": "paragraph", "text": "Supported behavior."}]
                run_gate.return_value = SimpleNamespace(returncode=0, stdout="block gate PASS", stderr="")
                state = reporting.observe_ipos_block_publication(root, "digital", "Controller", "snapshot", "2000-01-01 00:00:00")
                self.assertEqual(state["publication_status"], "PASS")
                self.assertEqual(len(state["artifact_hashes"]), 2)
                self.assertEqual(run_gate.call_args.args[0][-2:], ["--block", "Controller"])
                events.return_value = [{"kind": "table", "rows": [["Clock", "TBD"]]}]
                self.assertEqual(reporting.observe_ipos_block_publication(root, "digital", "Controller", "snapshot", "2000-01-01 00:00:00")["publication_status"], "FAIL_INLINE_MARKERS")
                run_gate.return_value = SimpleNamespace(returncode=1, stdout="", stderr="materialization failed")
                self.assertEqual(reporting.observe_ipos_block_publication(root, "digital", "Controller", "snapshot", "2000-01-01 00:00:00")["publication_status"], "FAIL_EXISTING_BLOCK_GATE")
                run_gate.reset_mock()
                self.assertEqual(reporting.observe_ipos_block_publication(root, "digital", "Controller", "snapshot", "2099-01-01 00:00:00")["publication_status"], "NOT_REGENERATED_IN_THIS_RUN")
                run_gate.assert_not_called()

    @unittest.skipUnless(shutil.which("pandoc"), "Pandoc is required for DOCX conversion")
    def test_user_docx_contains_issue_evidence_and_review_notice(self):
        with tempfile.TemporaryDirectory() as directory:
            report = self.user_report_fixture()
            markdown, docx = reporting.export_user_report(report, Path(directory) / "report.json")
            self.assertTrue(markdown.is_file())
            document = importlib.import_module("docx").Document(docx)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("REVIEW-0001", text)
            self.assertIn("For user review only", text)
            self.assertIn("REQ-1", text)
            self.assertIn(report["issues"][0]["source_excerpt"], text)
            self.assertIn("offset 31:100:", text)
            namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            languages = document.styles.element.xpath(".//w:lang")
            self.assertTrue(languages)
            self.assertTrue(all(element.get("{" + namespaces["w"] + "}val") == "en-US"
                                for element in languages))

    def test_audit_omissions_and_candidates_are_retained_without_authority(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "audit.csv"
            path.write_text("decision,reason,statement,source\nselected,eligible,supported,source-a\n"
                            "rejected,projection_gap,unmodeled,source-b\n"
                            "candidate_only,not rendered,candidate,source-c\n", encoding="utf-8")
            before = path.read_bytes()
            findings = reporting.collect_csv_issues(path, "SRS", root)
            self.assertEqual(len(findings), 2)
            self.assertEqual([value["audit_decision"] for value in findings], ["rejected", "candidate_only"])
            self.assertTrue(all(value["user_review_only"] for value in findings))
            self.assertEqual(path.read_bytes(), before)

    def test_publication_markers_preserve_location_and_exact_requirement_reference(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "spec.md"
            path.write_text("# Assumptions and TBD\nSupported function.\n"
                            "REQ-10: need clarification.\nState retention is not specified.\n", encoding="utf-8")
            findings = reporting.collect_publication_issues(path, "DRS", root, ["REQ-1", "REQ-10"])
            self.assertEqual(len(findings), 2)
            self.assertEqual(findings[0]["requirement_or_context_references"], ["REQ-10"])
            self.assertIn("line 3", findings[0]["location"])

    def test_conclusions_are_bounded_and_do_not_label_parser_gaps_as_source_damage(self):
        for blocker, expected in (("unsupported predicate grammar", "Need clarification"),
                                  ("ipos_source_relationship_changed", "Source-model relationship needs review"),
                                  ("unsupported structured/table/procedure evidence", "Structured evidence not sufficiently recoverable")):
            finding = reporting.issue("IPOS", "block", ["context"], blocker, "Observed finding", ["source"])
            self.assertEqual(finding["conclusion"], expected)
            self.assertIn(finding["conclusion"], reporting.CONCLUSIONS)

    def test_unclaimed_and_withheld_dispositions_are_reported(self):
        span = reporting.ipos.build_ipos_normalized_records(
            reporting.build_ipos_functional_input("Controller", {"Function": "control measurements"},
                                                  [{"source_req_id": "REQ-1", "requirement_statement": "The Controller shall start sampling."}],
                                                  domain="digital"))
        record = next(value for value in span if value.record_id.startswith("requirement:"))
        source_span = importlib.import_module("workflow_routing").DescriptiveSourceSpan(
            record.record_id, 0, len(record.text), record.text)
        disposition = SimpleNamespace(decision="unresolved", reason="no_supported_claim",
                                      spans=(source_span,), proposition_ids=("proposition-1",))
        plan = SimpleNamespace(evidence=SimpleNamespace(records=(record,), unresolved_contexts=()),
                               dispositions=(disposition,))
        findings = reporting.collect_plan_disposition_issues(plan, "IPOS", "Controller")
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["blocker_class"], "no_supported_claim")
        self.assertEqual(findings[0]["source_spans"][0]["text"], record.text)

    def test_closure_classes_do_not_treat_parser_gaps_as_source_damage_or_safe_omissions(self):
        cases = (("unsupported predicate grammar", "not worth pursuing in this step"),
                 ("unsupported multi-clause or non-normative construction", "not worth pursuing in this step"),
                 ("unbalanced source delimiters", "source clarification needed"),
                 ("unresolved subject reference", "source clarification needed"),
                 ("ipos_source_relationship_changed:REQ-1", "source-model relationship review needed"),
                 ("ambiguous guard boundary", "source-model relationship review needed"),
                 ("unsupported structured/table/procedure evidence", "structured-evidence recovery beyond this step"))
        for reason, expected in cases:
            self.assertEqual(reporting.ipos_closure_classification(reason), expected)
        self.assertEqual(reporting.ipos_closure_classification("equivalent duplicate", "excluded"), "safe to omit")

    def test_closure_routes_exact_points_without_overriding_blocked_partial_generation(self):
        report = self.user_report_fixture()
        report["gate_override"] = False
        functional = SimpleNamespace(block="Controller", domain="digital")
        span = importlib.import_module("workflow_routing").DescriptiveSourceSpan("REQ-1", 0, 8, "Evidence")
        context = SimpleNamespace(context_id="context-1", reason="unsupported predicate grammar", spans=(span,))
        plan = SimpleNamespace(evidence=SimpleNamespace(
            records=(SimpleNamespace(record_id="REQ-1", provenance=("source.md#L1",)),),
            unresolved_contexts=(context,)), claims=("supported-claim",), dispositions=())
        finding = "descriptive_synthesis_unresolved:unsupported predicate grammar"
        with patch.object(reporting, "build_ipos_claim_plan", return_value=plan), patch.object(
                reporting, "ipos_claim_preflight_findings", return_value=[finding]):
            reporting.apply_ipos_closure_review(report, [functional])
            first_ids = [value["issue_id"] for value in report["issues"]]
            reporting.apply_ipos_closure_review(report, [functional])
        closure = report["ipos_closure_review"]
        block = closure["blocks"][0]
        self.assertEqual(block["partial_generation_decision"], "NO-GO")
        self.assertFalse(block["supported_remainder_sufficient_under_current_contract"])
        self.assertEqual(block["supported_claims"], 1)
        self.assertEqual(len(block["points"]), 2)
        self.assertEqual(closure["minimum_clean_ipos_set"], [])
        self.assertEqual([value["issue_id"] for value in report["issues"]], first_ids)
        self.assertEqual(report["issues"][0]["source_excerpt"], self.user_report_fixture()["issues"][0]["source_excerpt"])
        self.assertIn("Bounded IPOS Closure Review", reporting.report_markdown(report))
        for point in block["points"]:
            self.assertTrue(point["review_ids"])
            self.assertIn(point["classification"], reporting.IPOS_CLOSURE_CLASSES)
        with self.assertRaises(ValueError):
            reporting.apply_ipos_closure_review({**report, "gate_override": True}, [functional])

    def test_closure_clean_set_requires_current_preflight_and_existing_certification(self):
        report = self.user_report_fixture()
        report["gate_override"] = False
        report["documents"].append({"document_type": "Digital IPOS", "blocks": [
            {"block": "Controller", "publication_status": "PASS", "publication_not_certified": False}]})
        functional = SimpleNamespace(block="Controller", domain="digital")
        plan = SimpleNamespace(evidence=SimpleNamespace(records=(), unresolved_contexts=()), claims=(), dispositions=())
        with patch.object(reporting, "build_ipos_claim_plan", return_value=plan), patch.object(
                reporting, "ipos_claim_preflight_findings", return_value=[]):
            reporting.apply_ipos_closure_review(report, [functional])
            self.assertEqual(report["ipos_closure_review"]["minimum_clean_ipos_set"],
                             [{"document_type": "Digital IPOS", "block": "Controller"}])
            report["documents"][-1]["blocks"][0]["publication_not_certified"] = True
            reporting.apply_ipos_closure_review(report, [functional])
            self.assertEqual(report["ipos_closure_review"]["minimum_clean_ipos_set"], [])


if __name__ == "__main__":
    unittest.main()