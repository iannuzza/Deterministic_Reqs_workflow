import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from mapping_review_sync import _row_is_valid, sync_mapping_workbook_to_csv
from generate_stage2_specs import _write_mapping_workbook


FIELDS = ["requirement_id", "approved_classification", "review_decision", "approved_block", "reviewer_notes", "allocation_class", "owning_target", "allocation_rationale", "lineage_mode"]


class MappingReviewSyncTests(unittest.TestCase):
    def _write_csv(self, path, rows):
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(rows)

    def _write_workbook(self, path, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Architecture Mapping Review"
        sheet.append(FIELDS)
        for row in rows:
            sheet.append([row.get(field, "") for field in FIELDS])
        workbook.save(path)
        workbook.close()

    def test_row_validity_predicate_matches_complete_and_incomplete_review_rows(self):
        complete = {
            "approved_classification": "Digital",
            "approved_block": "Controller",
            "allocation_class": "block_local_digital",
            "owning_target": "Digital IPOS",
            "allocation_rationale": "Supported by reviewed architecture evidence.",
            "lineage_mode": "normal_hierarchical",
        }
        self.assertTrue(_row_is_valid(complete))
        self.assertFalse(_row_is_valid(dict(complete, owning_target="Analog IPOS")))

    def test_saved_assignments_preserve_workbook_blanks(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            workbook_path = root / "architecture_mapping_preview.xlsx"
            baseline = [
                {"requirement_id": "REQ-1", "approved_classification": "Digital", "review_decision": "pending_review", "approved_block": "", "reviewer_notes": "", "allocation_class": "block_local_digital", "owning_target": "Digital IPOS", "allocation_rationale": "prefilled", "lineage_mode": "normal_hierarchical"},
                {"requirement_id": "REQ-2", "approved_classification": "Analog", "review_decision": "pending_review", "approved_block": "", "reviewer_notes": "", "allocation_class": "", "owning_target": "", "allocation_rationale": "", "lineage_mode": ""},
            ]
            self._write_csv(csv_path, baseline)
            saved = [dict(baseline[0], approved_block="Controller", allocation_class="", owning_target="", allocation_rationale="", lineage_mode=""), dict(baseline[1], approved_block="ADC")]
            self._write_workbook(workbook_path, saved)
            self.assertEqual(sync_mapping_workbook_to_csv(workbook_path, csv_path), 2)
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual([(row["approved_block"], row["review_decision"]) for row in rows], [("Controller", "pending_review"), ("ADC", "pending_review")])
            self.assertEqual(rows[0]["allocation_class"], "")

    def test_missing_workbook_id_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            workbook_path = root / "architecture_mapping_preview.xlsx"
            baseline = [
                {"requirement_id": "REQ-1", "approved_classification": "Digital", "review_decision": "pending_review", "approved_block": "", "reviewer_notes": ""},
                {"requirement_id": "REQ-2", "approved_classification": "Analog", "review_decision": "pending_review", "approved_block": "", "reviewer_notes": ""},
            ]
            self._write_csv(csv_path, baseline)
            self._write_workbook(workbook_path, [baseline[0]])
            with self.assertRaisesRegex(ValueError, "missing requirement IDs: REQ-2"):
                sync_mapping_workbook_to_csv(workbook_path, csv_path)

    def test_workbook_decision_is_authoritative(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            workbook_path = root / "architecture_mapping_preview.xlsx"
            baseline = [{
                "requirement_id": "REQ-1",
                "approved_classification": "Digital",
                "review_decision": "pending_review",
                "approved_block": "",
                "reviewer_notes": "",
            }]
            self._write_csv(csv_path, baseline)
            saved = [dict(baseline[0], review_decision="approved", approved_block="Controller")]
            self._write_workbook(workbook_path, saved)
            sync_mapping_workbook_to_csv(workbook_path, csv_path)
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                row = next(csv.DictReader(handle))
            self.assertEqual(row["review_decision"], "approved")
            self.assertEqual(row["approved_block"], "Controller")

    def test_xlsx_csv_xlsx_roundtrip_preserves_reviewed_value(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            workbook_path = root / "architecture_mapping_preview.xlsx"
            baseline = [{
                "requirement_id": "REQ-1",
                "approved_classification": "Digital",
                "review_decision": "pending_review",
                "approved_block": "Controller",
                "reviewer_notes": "",
            }]
            self._write_csv(csv_path, baseline)
            _write_mapping_workbook(csv_path)

            workbook = load_workbook(workbook_path)
            try:
                sheet = workbook["Architecture Mapping Review"]
                decision_column = FIELDS.index("review_decision") + 1
                sheet.cell(row=2, column=decision_column).value = "approved"
                workbook.save(workbook_path)
            finally:
                workbook.close()

            sync_mapping_workbook_to_csv(workbook_path, csv_path)
            regenerated_path = _write_mapping_workbook(csv_path)
            regenerated = load_workbook(regenerated_path, data_only=False)
            try:
                decision_cell = regenerated["Architecture Mapping Review"].cell(
                    row=2,
                    column=decision_column,
                )
                self.assertEqual(decision_cell.value, "approved")
            finally:
                regenerated.close()

    def test_accepted_decision_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            workbook_path = root / "architecture_mapping_preview.xlsx"
            baseline = [{
                "requirement_id": "REQ-1",
                "approved_classification": "Digital",
                "review_decision": "pending_review",
                "approved_block": "Controller",
                "reviewer_notes": "",
            }]
            self._write_csv(csv_path, baseline)
            self._write_workbook(workbook_path, [dict(baseline[0], review_decision="accepted")])
            with self.assertRaisesRegex(ValueError, "Unsupported review_decision"):
                sync_mapping_workbook_to_csv(workbook_path, csv_path)

    def test_valid_workbook_edits_default_to_approved(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            workbook_path = root / "architecture_mapping_preview.xlsx"
            baseline = [{
                "requirement_id": "REQ-1",
                "approved_classification": "Digital",
                "review_decision": "pending_review",
                "approved_block": "",
                "reviewer_notes": "",
                "allocation_class": "block_local_digital",
                "owning_target": "Digital IPOS",
                "allocation_rationale": "Supported by the reviewed architecture evidence.",
                "lineage_mode": "normal_hierarchical",
            }]
            self._write_csv(csv_path, baseline)
            self._write_workbook(workbook_path, [dict(baseline[0], approved_block="Controller", review_decision="")])
            sync_mapping_workbook_to_csv(workbook_path, csv_path)
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                row = next(csv.DictReader(handle))
            self.assertEqual(row["review_decision"], "approved")

    def test_generated_workbook_has_explicit_decision_dropdown_and_row_validity_formatting(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            self._write_csv(csv_path, [{
                "requirement_id": "REQ-1",
                "approved_classification": "Digital",
                "review_decision": "pending_review",
                "approved_block": "Controller",
                "reviewer_notes": "",
                "allocation_class": "block_local_digital",
                "owning_target": "Digital IPOS",
                "allocation_rationale": "supported",
                "lineage_mode": "normal_hierarchical",
            }])

            workbook_path = _write_mapping_workbook(csv_path)
            workbook = load_workbook(workbook_path, data_only=False)
            try:
                sheet = workbook["Architecture Mapping Review"]
                self.assertFalse(sheet.protection.sheet)
                decision_column = FIELDS.index("review_decision") + 1
                decision_cell = sheet.cell(row=2, column=decision_column)
                self.assertEqual(decision_cell.value, "pending_review")
                self.assertFalse(decision_cell.protection.locked)
                self.assertFalse(str(decision_cell.value).startswith("="))

                validations = list(sheet.data_validations.dataValidation)
                self.assertTrue(any(
                    validation.formula1 == "'Approved Block Choices'!$G$2:$G$6"
                    and decision_cell.coordinate in str(validation.sqref)
                    for validation in validations
                ))
                choices = workbook["Approved Block Choices"]
                self.assertEqual(
                    [choices.cell(row=row, column=7).value for row in range(2, 7)],
                    ["pending_review", "approved", "reassigned", "rejected", "needs_clarification"],
                )

                formulas = [
                    formula
                    for rules in sheet.conditional_formatting._cf_rules.values()
                    for rule in rules
                    for formula in (rule.formula or [])
                ]
                self.assertTrue(any('$C2="approved"' in formula for formula in formulas))
                self.assertFalse(any('$C2="accepted"' in formula for formula in formulas))
                self.assertTrue(any(formula.startswith("NOT(AND(") for formula in formulas))
                self.assertTrue(any("$C2" in formula for formula in formulas))
                self.assertFalse(any("$C$2" in formula for formula in formulas))
                ranges = [str(range_string) for range_string in sheet.conditional_formatting]
                self.assertTrue(any("C2" in range_string for range_string in ranges))
                self.assertFalse(any("A2:I2" in range_string for range_string in ranges))
                self.assertFalse(any(cell.data_type == "f" for cell in sheet[1] + sheet[2]))
            finally:
                workbook.close()

    def test_approved_row_uses_green_conditional_formatting_only_on_editable_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            self._write_csv(csv_path, [{
                "requirement_id": "REQ-1",
                "approved_classification": "Digital",
                "review_decision": "approved",
                "approved_block": "Controller",
                "reviewer_notes": "",
                "allocation_class": "block_local_digital",
                "owning_target": "Digital IPOS",
                "allocation_rationale": "supported",
                "lineage_mode": "normal_hierarchical",
            }])
            workbook_path = _write_mapping_workbook(csv_path)
            workbook = load_workbook(workbook_path, data_only=False)
            try:
                sheet = workbook["Architecture Mapping Review"]
                editable_columns = {
                    index + 1 for index, field in enumerate(FIELDS)
                    if field in {
                        "approved_classification", "review_decision", "approved_block",
                        "reviewer_notes", "allocation_class", "owning_target",
                        "allocation_rationale", "lineage_mode",
                    }
                }
                for column in editable_columns:
                    self.assertEqual(sheet.cell(row=2, column=column).fill.fgColor.rgb, "00C6EFCE")
                self.assertNotEqual(sheet.cell(row=2, column=1).fill.fgColor.rgb, "00C6EFCE")
                self.assertTrue(any(
                    'C2="approved"' in formula
                    for rules in sheet.conditional_formatting._cf_rules.values()
                    for rule in rules
                    for formula in (rule.formula or [])
                ))
            finally:
                workbook.close()

    def test_manual_pending_to_approved_has_red_initial_fill_and_green_rule(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "architecture_mapping_preview.csv"
            row = {
                "requirement_id": "REQ-1",
                "approved_classification": "Digital",
                "review_decision": "pending_review",
                "approved_block": "Controller",
                "reviewer_notes": "",
                "allocation_class": "block_local_digital",
                "owning_target": "Digital IPOS",
                "allocation_rationale": "supported",
                "lineage_mode": "normal_hierarchical",
            }
            self._write_csv(csv_path, [row])
            workbook_path = _write_mapping_workbook(csv_path)
            workbook = load_workbook(workbook_path, data_only=False)
            try:
                sheet = workbook["Architecture Mapping Review"]
                editable_columns = {
                    index + 1 for index, field in enumerate(FIELDS)
                    if field in {
                        "approved_classification", "review_decision", "approved_block",
                        "reviewer_notes", "allocation_class", "owning_target",
                        "allocation_rationale", "lineage_mode",
                    }
                }
                for column in editable_columns:
                    self.assertEqual(sheet.cell(row=2, column=column).fill.fgColor.rgb, "00FFC7CE")

                sheet.cell(row=2, column=FIELDS.index("review_decision") + 1).value = "approved"
                formulas = [
                    formula
                    for rules in sheet.conditional_formatting._cf_rules.values()
                    for rule in rules
                    for formula in (rule.formula or [])
                ]
                decision_column_letter = "C"
                self.assertTrue(any(
                    f'${decision_column_letter}2="approved"' in formula
                    for formula in formulas
                ))
                self.assertTrue(any(
                    f'${decision_column_letter}2="approved"' in formula
                    for formula in formulas
                ))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
