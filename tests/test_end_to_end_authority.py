import json
import tempfile
import unittest
from pathlib import Path

import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from approved_snapshot_resolver import ApprovedResolverError, resolve_authoritative_input
from canonical_store import connect, create_snapshot_id, fingerprint
from report_generation_service import generate_snapshot_report
from s0_source_bootstrap import approve_candidate, assess_source, persist_candidate
from snapshot_manager import create_approved_snapshot
from workflow_gui import WORKFLOW_DIAGRAM_EDGES, WORKFLOW_DIAGRAM_NODES


class EndToEndAuthorityTests(unittest.TestCase):
    def _base_authority(self, root: Path, *, profile_state: str = "approved") -> None:
        connection = connect(root)
        connection.executescript(
            """
            INSERT INTO projects(project_id, project_name, created_at) VALUES ('p1', 'Demo', 'now');
            INSERT INTO source_specs(source_spec_id, project_id, source_name, source_kind, source_scope, target_scope, created_at)
            VALUES ('primary-1', 'p1', 'primary.pdf', 'primary', 'system', 'system', 'now');
            INSERT INTO source_revisions(source_revision_id, source_spec_id, revision_number, content_fingerprint, ingestion_timestamp)
            VALUES ('source-rev-1', 'primary-1', '1', 'source-hash', 'now');
            INSERT INTO canonical_requirements(canonical_requirement_id, project_id, source_req_id, current_revision_id, created_at, updated_at)
            VALUES ('can-1', 'p1', 'REQ-1', 'rev-1', 'now', 'now');
            INSERT INTO canonical_requirement_revisions(revision_id, canonical_requirement_id, revision_number, requirement_text, content_fingerprint, approved_classification, created_at)
            VALUES ('rev-1', 'can-1', 1, 'The controller shall reset.', 'req-hash', 'functional', 'now');
            INSERT INTO requirement_provenance(revision_id, source_spec_id, source_revision_id, source_req_id)
            VALUES ('rev-1', 'primary-1', 'source-rev-1', 'REQ-1');
            """
        )
        for kind in ("vocabulary", "taxonomy", "architecture"):
            payload = {"kind": kind, "version": "1"}
            connection.execute(
                """INSERT INTO profile_revisions(
                    profile_revision_id, project_id, profile_kind, profile_name, revision_number,
                    profile_payload_json, content_fingerprint, lifecycle_state, approval_state,
                    approved_by, approved_at, created_at)
                    VALUES (?, 'p1', ?, ?, '1', ?, ?, ?, ?, 'tester', 'now', 'now')""",
                (
                    f"profile-{kind}", kind, kind, json.dumps(payload), fingerprint(payload),
                    profile_state, profile_state,
                ),
            )
        connection.commit()
        connection.close()

    def test_canonical_store_is_separate_from_retrieval_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            connection = connect(root)
            database_tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
            connection.close()
            self.assertTrue((root / "data/canonical/canonical_store.sqlite").exists())
            self.assertIn("canonical_requirements", database_tables)
            self.assertNotIn("chunks", database_tables)
            self.assertNotIn("chunk_fts", database_tables)

    def test_authoritative_resolver_rejects_unapproved_profile_context(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self._base_authority(root, profile_state="pending")
            material = {"canonical": ["rev-1"], "mapping": [], "profiles": {}, "retrieval": {"version": "1"}}
            snapshot_id = create_snapshot_id(material)
            connection = connect(root)
            connection.execute(
                """INSERT INTO snapshots(snapshot_id, project_id, snapshot_kind, status,
                    canonical_revision_set_json, mapping_revision_set_json, content_fingerprint, approved_at, created_at)
                    VALUES (?, 'p1', 'mapping', 'approved', '[\"rev-1\"]', '[]', ?, 'now', 'now')""",
                (snapshot_id, fingerprint(material)),
            )
            connection.commit()
            connection.close()
            with self.assertRaises(ApprovedResolverError):
                resolve_authoritative_input(root, "3", project_id="p1", snapshot_id=snapshot_id)

    def test_approved_primary_bootstrap_preserves_source_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "primary.html"
            source.write_text("[REQ_MAIN_0001] Requirement: The controller shall reset.", encoding="utf-8")
            assessment = assess_source(source, {"requirement_id_rules": {"table_req_id_patterns": [r"REQ_MAIN_\d+"]}})
            profile_id = persist_candidate(root, project_id="p1", assessment=assessment, actor="tester")
            approve_candidate(root, profile_id=profile_id, actor="reviewer", approved_prefix="REQ_MAIN")
            connection = connect(root)
            payload = json.loads(connection.execute("SELECT profile_payload_json FROM profile_revisions WHERE profile_revision_id = ?", (profile_id,)).fetchone()[0])
            connection.close()
            self.assertEqual(payload["source_path"], str(source))
            self.assertEqual(payload["approved_prefix"], "REQ_MAIN")
            self.assertEqual(payload["status"], "tagged")

    def test_multi_source_multi_scope_provenance_is_retained(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            connection = connect(root)
            connection.execute("INSERT INTO projects(project_id, project_name, created_at) VALUES ('p1', 'Demo', 'now')")
            for source_id, name, scope in (("supp-a", "analog.pdf", "analog"), ("supp-d", "digital.pdf", "digital")):
                connection.execute(
                    """INSERT INTO source_specs(source_spec_id, project_id, source_name, source_kind, source_scope, target_scope, created_at)
                    VALUES (?, 'p1', ?, 'supplementary', ?, 'controller', 'now')""",
                    (source_id, name, scope),
                )
                connection.execute(
                    """INSERT INTO source_revisions(source_revision_id, source_spec_id, revision_number, content_fingerprint, ingestion_timestamp)
                    VALUES (?, ?, '1', ?, 'now')""",
                    (f"{source_id}-rev", source_id, f"{source_id}-hash"),
                )
                connection.execute(
                    """INSERT INTO ingestion_batches(ingestion_batch_id, project_id, batch_kind, started_at)
                    VALUES (?, 'p1', 'supplementary', 'now')""",
                    (f"{source_id}-batch",),
                )
                connection.execute(
                    """INSERT INTO source_integration_records(
                    integration_record_id, source_spec_id, source_revision_id, ingestion_batch_id,
                    integration_scope, created_at)
                    VALUES (?, ?, ?, ?, ?, 'now')""",
                    (f"{source_id}-integration", source_id, f"{source_id}-rev", f"{source_id}-batch", scope),
                )
            connection.commit()
            rows = connection.execute("SELECT source_kind, source_scope, integration_scope FROM source_specs JOIN source_integration_records USING (source_spec_id) ORDER BY source_scope").fetchall()
            connection.close()
            self.assertEqual([(row[0], row[1], row[2]) for row in rows], [
                ("supplementary", "analog", "analog"),
                ("supplementary", "digital", "digital"),
            ])

    def test_snapshot_report_is_derived_and_records_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self._base_authority(root)
            snapshot_id = create_approved_snapshot(
                root,
                project_id="p1",
                canonical_revision_ids=["rev-1"],
                mapping_ids=[],
                approved_by="tester",
                retrieval={"version": "1", "hash": "retrieval-hash"},
                semantic_enabled=False,
            )
            output = root / "artifacts/stage3_srs/SRS_snapshot_report.xlsx"
            generate_snapshot_report(
                root,
                project_id="p1",
                snapshot_id=snapshot_id,
                use_latest_approved=False,
                report_kind="SRS",
                output_path=output,
                rows=[{"source_req_id": "REQ-1", "block": "Controller", "statement": "The controller shall reset."}],
            )
            from openpyxl import load_workbook
            workbook = load_workbook(output, read_only=False)
            self.assertIn("SRS", workbook.sheetnames)
            self.assertIn("Report Metadata", workbook.sheetnames)
            metadata = {row[0]: row[1] for row in workbook["Report Metadata"].iter_rows(values_only=True) if row[0]}
            workbook.close()
            self.assertEqual(metadata["Snapshot ID"], snapshot_id)
            connection = connect(root)
            report = connection.execute("SELECT snapshot_id, output_format, status FROM report_metadata").fetchone()
            connection.close()
            self.assertEqual(tuple(report), (snapshot_id, "xlsx", "generated"))

    def test_gui_graph_exposes_supplementary_review_and_derived_outputs(self):
        node_ids = {node["id"] for node in WORKFLOW_DIAGRAM_NODES}
        edge_pairs = {(source, target) for source, target, _kind in WORKFLOW_DIAGRAM_EDGES}
        self.assertTrue({"supp_import", "supp_ingest", "supp_review", "supp_merge"}.issubset(node_ids))
        self.assertIn(("s2e", "reports"), edge_pairs)
        self.assertIn(("s2e", "sysml"), edge_pairs)


if __name__ == "__main__":
    unittest.main()
