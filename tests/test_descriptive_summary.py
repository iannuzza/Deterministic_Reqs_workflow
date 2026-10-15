import csv
import json
import os
import re
import sys
import tempfile
import unittest
from dataclasses import replace
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from workflow_routing import (
    DESCRIPTIVE_PROFILES,
    SRS_DESCRIPTIVE_PROFILE,
    DRS_DESCRIPTIVE_PROFILE,
    ARS_DESCRIPTIVE_PROFILE,
    IPOS_DESCRIPTIVE_PROFILE,
    descriptive_profile,
    descriptive_fact_findings,
    run_descriptive_flow,
    validate_descriptive_flow,
    DescriptiveFact,
    DescriptiveParagraph,
    DescriptiveWritingProfile,
    DESCRIPTIVE_STRUCTURE_VERSION,
    DescriptiveSourceSpan,
    DescriptiveContext,
    DescriptiveMeaning,
    DescriptiveProposition,
    DescriptiveDependency,
    DescriptiveClaim,
    DescriptiveClaimIntent,
    DescriptiveFunctionalClaim,
    DescriptiveAcceptance,
    DescriptiveEvidenceModel,
    DescriptiveUnresolvedContext,
    build_descriptive_synthesis_plan,
    validate_descriptive_synthesis_plan,
    descriptive_synthesis_plan_audit,
    descriptive_section_contributions,
    ipos_publication_quality_findings,
    DescriptiveDisposition,
    DescriptiveStructure,
    validate_descriptive_structure,
    descriptive_structure_audit,
    validate_descriptive_structure_audit,
    DESCRIPTIVE_APPLICATION_PROFILES,
    DescriptiveAssessment,
    apply_descriptive_profile,
    descriptive_profile_application_audit,
    validate_descriptive_profile_application_audit,
    compose_descriptive_paragraphs,
    validate_descriptive_paragraphs,
    DESCRIPTIVE_MODE_TOPIC_SPECS,
    DESCRIPTIVE_POWER_TOPIC_SPECS,
    DESCRIPTIVE_DIGITAL_TOPIC_SPECS,
    assign_ipos_descriptive_topic,
    assemble_descriptive_summary,
    build_ipos_functional_input,
    compose_ipos_overview as compose_shared_ipos_overview,
    _legacy_compose_ipos_overview as compose_ipos_overview,
    compose_technical_block_purpose,
    compose_drs_interaction_summary,
    compose_drs_role_summary,
    drs_section_prose_findings,
    document_author_name,
    validate_document_author_fields,
    project_architecture_interactions,
    project_architecture_interfaces,
    summarize_architecture_interactions,
    validate_drs_block_descriptions,
    _ipos_aggregate_responsibility,
    _ipos_evidence_summary,
    _legacy_validate_ipos_descriptive_output as validate_ipos_descriptive_output,
    write_ipos_descriptive_audit,
    write_descriptive_audit,
    _ipos_structural_behavior_facets,
    ipos_source_preservation_findings,
    _ipos_claim_preservation_findings,
    _ipos_shared_flow,
    build_ipos_descriptive_structure,
    ipos_descriptive_structural_audit,
    write_ipos_structural_audit,
    validate_ipos_structural_audit,
)
from spec_document_contract import NormalizedSourceRecord, SemanticUnit, stable_payload_hash
from dataclasses import asdict
from low_power_descriptive import (
    architecture_records_from_function_rows,
    assemble_low_power_descriptive,
    is_drs_top_level_record,
    power_domain_records_from_ocr,
    read_selected_low_power_audit,
    write_low_power_audit,
)
from ipos_semantic_normalizer import detail_evidence_facets, compose_detail_groups
from generate_ipos_specs import approved_snapshot_candidate_rows
from run_srs_gen_spec_agent import _render_srs_low_power_section
from run_drs_gen_spec_agent import (
    BlockInfo, _drs_block_relationships, _drs_descriptive_blocks, _drs_intro_rows,
    _drs_function_summary, _drs_power_domain_rows, _drs_topic_narrative, _drs_shared_facet_rows,
    _drs_domain_function_sentence, _drs_clock_path_rows, _drs_clock_path_markdown_row, _drs_sync_connection_rows,
    _drs_power_clock_routes, _drs_power_reset_outputs, _drs_section_interaction_summary,
    _drs_clock_section_summary, _render_approved_port_table, validate_drs_descriptive_artifacts,
    _drs_rendered_prose_findings, _drs_top_level_overview,
    _drs_processing_summary,
)


class DescriptiveSummaryTests(unittest.TestCase):
    def assert_nonessential_ipos_residual(self, functional):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings, ipos_derivative_residuals
        plan = build_ipos_claim_plan(functional)
        omitted = [value for value in plan.dispositions if value.reason.startswith("omitted_nonessential:")]
        self.assertTrue(omitted)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        records = {record.record_id: record for record in plan.evidence.records}
        residuals = ipos_derivative_residuals(functional)
        for disposition in omitted:
            self.assertEqual(disposition.decision, "excluded")
            self.assertTrue(disposition.spans)
            for span in disposition.spans:
                self.assertEqual(records[span.record_id].text[span.start:span.end], span.text)
                self.assertTrue(any(span.record_id in value.get("source_record_ids", ())
                                    and value["readiness_impact"] == "requires_assessment"
                                    and value["authority"] is False for value in residuals))
        self.assertTrue(validate_descriptive_synthesis_plan(replace(plan, dispositions=()),
                        expected_evidence=plan.evidence, expected_intents=plan.intents,
                        profile=IPOS_DESCRIPTIVE_PROFILE, owner=functional.block))
        return omitted

    def test_drs_dedicated_connection_tables_preserve_approved_fields(self):
        from run_drs_gen_spec_agent import _drs_reset_connection_rows, _drs_connection_table_findings
        from workflow_routing import render_connection_table_sections
        from spec_document_contract import drs_document_events
        row = {"source_req_id": "RESET", "approved_classification": "Digital", "lifecycle_state": "approved",
               "requirement_statement": "The top.reset_n signal shall be connected to the top.unit.reset_n signal."}
        incomplete = {**row, "source_req_id": "INCOMPLETE",
                      "requirement_statement": "The top.controller. signal shall be connected to the top.unit.other_reset_n signal."}
        self.assertEqual(_drs_reset_connection_rows([row, incomplete]), [
            ("RESET", "top.reset_n", "top.unit.reset_n"), ("INCOMPLETE", "", "top.unit.other_reset_n")])
        self.assertEqual(_drs_reset_connection_rows([{**row, "lifecycle_state": "staged"},
                                                   {**row, "approved_classification": "Analog"},
                                                   {**row, "source_req_id": ""}]), [])
        conditional = {**row, "requirement_statement": "When enabled, " + row["requirement_statement"]}
        self.assertEqual(_drs_reset_connection_rows([conditional]), [])
        clock = {**row, "source_req_id": "CLOCK", "requirement_statement":
                 "The top.clk signal shall be connected to the top.unit.clk signal, with frequency 8 MHz and with clock gating top.enable | top.test."}
        self.assertEqual(_drs_clock_path_rows([clock], preserve_source_fields=True), [
            ("CLOCK", "top.clk", "top.unit.clk", "8 MHz", "top.enable | top.test")])
        sections = render_connection_table_sections("4.7", {
            "clock": ["| CLOCK | top.clk | top.unit.clk |  |  |"],
            "reset": ["| RESET | top.reset_n | top.unit.reset_n |"]}, first_table_number=8)
        text = "\n".join(sections)
        self.assertIn("#### 4.7.1 Clock connections", text)
        self.assertIn("##### Table 8. Clock connections", text)
        self.assertIn("#### 4.7.2 Reset connections", text)
        self.assertIn("##### Table 9. Reset connections", text)
        self.assertIn("| Source requirement | Source | Destination | Clock frequency | Clock gate condition |", text)
        self.assertIn("| Source requirement | Source | Destination |", text)
        rows = {"clock": [("CLOCK", "top.clk", "top.unit.clk", "", "")],
            "reset": [("RESET", "top.reset_n", "top.unit.reset_n")]}
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "tables.md"
            path.write_text(text, encoding="utf-8")
            self.assertEqual(_drs_connection_table_findings(drs_document_events(path), rows), [])
            for altered in (text.replace("top.unit.reset_n", "top.other.reset_n"),
                    text.replace("| CLOCK |", "| UNKNOWN |"), text.replace("Table 8.", "Table 99."),
                    text.replace("4.7.2 Reset connections", "4.8.2 Reset connections"),
                    text + text.splitlines()[-1] + "\n"):
                path.write_text(altered, encoding="utf-8")
                self.assertTrue(_drs_connection_table_findings(drs_document_events(path), rows))

    def test_docx_navigation_keeps_toc_entries_before_navigation(self):
        from workflow_routing import _reorder_docx_navigation
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        entries = [("Heading1", "Specification"), ("Heading2", "0. Document Navigation"),
                   ("Heading3", "0.1 Table of contents"), ("", "1. Block overview"),
                   ("", "1.5.1 Clock connections"), ("", "1.5.2 Reset connections"),
                   ("Heading3", "0.2 Document control"), ("Heading2", "1. Block overview")]
        paragraphs = ["<w:p>" + ("<w:pPr><w:pStyle w:val='" + style + "'/></w:pPr>" if style else "")
                      + "<w:r><w:t>" + text + "</w:t></w:r></w:p>" for style, text in entries]
        document = "<w:document xmlns:w='" + namespace[1:-1] + "'><w:body>" + "".join(paragraphs) + "</w:body></w:document>"
        root = ET.fromstring(_reorder_docx_navigation(document))
        texts = ["".join(node.text or "" for node in paragraph.iter(namespace + "t"))
                 for paragraph in root.iter(namespace + "p")]
        navigation = texts.index("0. Document Navigation")
        self.assertLess(texts.index("1.5.1 Clock connections"), navigation)
        self.assertLess(texts.index("1.5.2 Reset connections"), navigation)
        self.assertLess(navigation, texts.index("0.2 Document control"))
        self.assertGreater(len(texts) - 1, navigation)

    def test_ipos_numbered_connection_navigation(self):
        from workflow_routing import apply_ipos_document_navigation
        lines = ["# Specification", "", "## 0. Document Navigation", "### 0.1 Table of contents",
                 "- [old](#old)", "", "### 0.2 Document control", "#### Table 1. Version history",
                 "| Version |", "|---|", "| 0.1 |", "", "### 0.3 Requirement navigation",
                 "#### Table. Requirement navigation", "| Requirement |", "|---|", "| ITEM |", "",
                 "## 1. Block overview", "### 1.5 General architecture",
                 "#### 1.5.1 Clock connections", "##### Table. Clock connections",
                 "| Source |", "|---|", "| clk |", "",
                 "#### 1.5.2 Reset connections", "##### Table. Reset connections",
                 "| Source |", "|---|", "| rst |", "", "## 2. Source I/O",
                 "#### Table. Source I/O", "| Port |", "|---|", "| input |", "", "## 3. Block requirements"]
        output = apply_ipos_document_navigation(lines)
        self.assertEqual(output[:3], ["# Specification", "", "## 0. Document Navigation"])
        text = "\n".join(output)
        self.assertIn("##### Table 3. Clock connections", text)
        self.assertIn("##### Table 4. Reset connections", text)
        toc = text.split("### 0.1 Table of contents", 1)[1].split("### 0.2 Document control", 1)[0]
        self.assertIn("[1.5.1 Clock connections](#151-clock-connections)", toc)
        self.assertIn("[1.5.2 Reset connections](#152-reset-connections)", toc)
        index = text.split("### 0.4 Table index", 1)[1].split("## 1. Block overview", 1)[0]
        self.assertIn("[Table 3. Clock connections](#table-3-clock-connections)", index)
        self.assertIn("[Table 4. Reset connections](#table-4-reset-connections)", index)
        self.assertEqual(index.count("- [Table "), 5)
        self.assertEqual(apply_ipos_document_navigation(output), output)

    def test_ipos_connection_tables_keep_functional_prose_and_supported_fields(self):
        functional = build_ipos_functional_input("Module",
            {"Function": "Control clock and reset sequencing", "Inputs": "requests", "Outputs": "status"},
            [{"source_req_id": "CLOCK", "statement": "The top.source_clk signal shall be connected to the top.dest_clk signal, with frequency 8 MHz and with clock gating top.enable | top.test."},
             {"source_req_id": "RESET", "statement": "The top.reset_n signal shall be connected to the top.unit.reset_n signal."},
             {"source_req_id": "BEHAVIOR", "statement": "When request is high, status shall be high."}], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        prose = " ".join(sections["internal_functions"])
        self.assertIn("When request is high, status is high.", prose)
        self.assertNotIn("connected to", prose)
        tables = "\n".join(sections["general_architecture"])
        self.assertIn("| Source requirement | Source | Destination | Clock frequency | Clock gate condition |", tables)
        self.assertIn("| CLOCK | top.source_clk | top.dest_clk | 8 MHz | top.enable \\| top.test |", tables)
        self.assertIn("| Source requirement | Source | Destination |", tables)
        self.assertIn("| RESET | top.reset_n | top.unit.reset_n |", tables)
        self.assertEqual(sum(row["decision"] == "selected" for row in audit), 5)

    def test_ipos_connection_tables_leave_unattested_fields_empty(self):
        functional = build_ipos_functional_input("Module",
            {"Function": "Control clocks", "Inputs": "requests", "Outputs": "status"},
            [{"source_req_id": "CLOCK", "statement": "The top.source_clk signal shall be connected to the top.dest_clk signal."},
             {"source_req_id": "RESET", "statement": "signal shall be connected to the top.unit.reset_n signal."},
             {"source_req_id": "INCOMPLETE", "statement": "The top.controller. signal shall be connected to the top.unit.other_reset_n signal."}], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        tables = "\n".join(sections["general_architecture"])
        self.assertIn("| CLOCK | top.source_clk | top.dest_clk |  |  |", tables)
        self.assertIn("| RESET |  | top.unit.reset_n |", tables)
        self.assertIn("| INCOMPLETE |  | top.unit.other_reset_n |", tables)
        omitted = next(row for row in audit if row.get("candidate_evidence_requirement_id") == "INCOMPLETE"
                   and row["decision"] == "omitted")
        self.assertEqual(json.loads(omitted["span_json"])["text"], "The top.controller.")
        self.assertFalse(sections["internal_functions"])

    def test_ipos_connection_fidelity_rejects_missing_changed_duplicate_or_misplaced_rows(self):
        from workflow_routing import build_ipos_claim_plan, _ipos_connection_table_findings, _ipos_connection_audit_findings, apply_ipos_document_navigation
        from spec_document_contract import drs_document_events
        functional = build_ipos_functional_input("Module",
            {"Function": "Control clocks", "Inputs": "requests", "Outputs": "status"},
            [{"source_req_id": "CLOCK", "statement": "The top.source_clk signal shall be connected to the top.dest_clk signal, with frequency 8 MHz and without clock gating."}], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        plan = build_ipos_claim_plan(functional)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.md"
            text = "\n".join(apply_ipos_document_navigation([
                "### 1.5 General architecture", "", *sections["general_architecture"]])) + "\n"
            path.write_text(text, encoding="utf-8")
            events = drs_document_events(path)
            self.assertEqual(_ipos_connection_table_findings(events, plan.claims), [])
            for altered in (text.replace("8 MHz", "9 MHz"), text.replace("top.dest_clk", "top.other_clk"),
                            text.replace("Not gated", "top.enable"), text.replace("| CLOCK |", "| UNKNOWN |"),
                            text.replace("Table 1. Clock connections", "Table 9. Clock connections"),
                            text.replace("1.5.1 Clock connections", "1.5.9 Clock connections"),
                            text.replace("Clock frequency", "Frequency"), text.replace("1.5 General architecture", "1.4 Supported block functions"),
                            text + text.splitlines()[-1] + "\n", text.rsplit("\n", 2)[0] + "\n"):
                path.write_text(altered, encoding="utf-8")
                self.assertTrue(_ipos_connection_table_findings(drs_document_events(path), plan.claims))
        selected = next(row for row in audit if row.get("candidate_evidence_requirement_id") == "CLOCK")
        claim = next(value for value in plan.claims if value.claim_id == selected["unit_id"])
        unit = json.loads(selected["semantic_unit_json"])
        self.assertEqual(_ipos_connection_audit_findings(unit["text"], claim), [])
        self.assertTrue(_ipos_connection_audit_findings(unit["text"].replace("8 MHz", "9 MHz"), claim))

    def test_ipos_materialization_dispositions_keep_exact_omitted_inputs_without_false_retention(self):
        from workflow_routing import build_ipos_normalized_records, build_ipos_semantic_units
        functional = build_ipos_functional_input("Module",
            {"Function": "Process samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "structured", "statement": "0xA000 CONTROL_REG 32 0xA040 reset_n"},
             {"source_req_id": "optional", "statement": "This signal governs the operation as previously explained."}], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        records = build_ipos_normalized_records(functional)
        units = build_ipos_semantic_units(functional, sections, audit)
        represented = {identity for unit in units for identity in unit.source_record_ids}
        self.assertTrue({record.record_id for record in records if record.decision == "retained"}.issubset(represented))
        excluded = [record for record in records if record.record_id.startswith("requirement:")]
        self.assertEqual(len(excluded), 2)
        self.assertTrue(all(record.decision == "excluded" and record.provenance for record in excluded))
        self.assertEqual(excluded[0].text, "0xA000 CONTROL_REG 32 0xA040 reset_n")
        self.assertEqual(excluded[1].text, "This signal governs the operation as previously explained.")
        self.assertIn("nonprose_output_path:register", excluded[0].rationale)
        self.assertIn("omitted_nonessential:", excluded[1].rationale)

    def test_ipos_required_section_policy_omits_unsupported_independent_detail(self):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings
        inventory = {"Function": "Control sampling", "Inputs": "configuration", "Outputs": "samples"}
        functional = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "detail", "statement": "This signal governs the operation as previously explained."}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        omitted = [value for value in plan.dispositions if value.reason.startswith("omitted_nonessential:")]
        self.assertEqual(len(omitted), 1)
        sections, audit = compose_shared_ipos_overview(functional)
        self.assertNotIn("previously explained", str(sections))
        self.assertTrue(any(value["decision"] == "omitted" for value in audit))
        missing = replace(functional, function="")
        self.assertTrue(ipos_claim_preflight_findings(missing))

    def test_ipos_same_record_signal_reference_preserves_range_output_and_dependency(self):
        from workflow_routing import build_ipos_claim_plan, _ipos_evidence_semantic_findings
        inventory = {"Function": "Control sampling", "Inputs": "configuration", "Outputs": "samples"}
        text = "The sampling phase shall be configurable by the user with i_CYCLES values from 0 up to 15. This signal indicates the number of clock that o_START shall stay high."
        functional = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "duration", "statement": text}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertFalse(plan.evidence.unresolved_contexts)
        self.assertEqual(_ipos_evidence_semantic_findings(plan.evidence), [])
        reference = next(value for value in plan.evidence.propositions if value.meaning.subject == "i_CYCLES")
        self.assertEqual(len(reference.spans), 2)
        self.assertIn("o_START stays high", reference.meaning.objects[0])
        self.assertIn("from 0 up to 15", reference.spans[0].text)
        self.assertTrue(_ipos_evidence_semantic_findings(replace(plan.evidence, dependencies=())))
        altered = replace(reference, meaning=replace(reference.meaning, subject="i_OTHER"))
        self.assertTrue(_ipos_evidence_semantic_findings(replace(plan.evidence,
            propositions=tuple(altered if value == reference else value for value in plan.evidence.propositions))))
        ambiguous = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "ambiguous", "statement": text.replace("the user with", "i_MODE with")}], domain="digital")
        self.assertTrue(build_ipos_claim_plan(ambiguous).evidence.unresolved_contexts)

    def test_ipos_structured_rows_do_not_hide_functional_behavior(self):
        from workflow_routing import descriptive_nonprose_output_path
        for text, expected in (
            ("0xA000 CONTROL_REG 32 0xA040 reset_n, software_reset_n", "register"),
            ("0xA044 RESERVED 0xAFFF reset_n", "register"),
            ("bus_mode 1'b0 To activate the serial interface", "configuration"),
            ("device_address 7'h5f Device address", "configuration"),
            ('interface_mode "serial" To choose if the interface is serial only or both.', "configuration"),
            ("Suggested val. 1200", "configuration"),
            ("Controller 32 processes samples", ""),
            ("When enabled, the result shall be computed using Table 3", ""),
            ("timeout 6000 If the timer expires the transaction is aborted", ""),
        ):
            with self.subTest(text=text):
                record = NormalizedSourceRecord("field", text, "approved_normative", "block_local",
                                                ("approved snapshot",), "retained", "approved same-block evidence")
                span = DescriptiveSourceSpan("field", 0, len(text), text)
                self.assertEqual(descriptive_nonprose_output_path(span, record, profile=IPOS_DESCRIPTIVE_PROFILE), expected)
                self.assertEqual(descriptive_nonprose_output_path(span, replace(record, decision="rejected"),
                                                                 profile=IPOS_DESCRIPTIVE_PROFILE), "")
                self.assertEqual(descriptive_nonprose_output_path(span, record, profile=SRS_DESCRIPTIVE_PROFILE), "")

    def test_register_catalog_headers_do_not_create_prose_obligations(self):
        from workflow_routing import descriptive_nonprose_output_path
        rows = (
            "Device Register Map 2.3.4.",
            "Device Register Map Here is described the list of registers in the register map.",
            "20 **Update-IR** Run-Test/Idle Select-DR-Scan 2.3.",
            "BASE START ADDRESS ADDRESS BLOCK REGISTER WIDTH BASE END ADDRESS RESET REQ_ID 0xA000 CONTROL_REG 32 0xA040 reset_n",
            "The control registers are accessible by the serial interface, while the others are only accessible by the host bus.",
        )
        text = " ".join(rows)
        record = NormalizedSourceRecord("catalog", text, "approved_normative", "block_local",
                                        ("approved snapshot",), "retained", "approved register catalog")
        for index, row in enumerate(rows):
            start = text.index(row)
            span = DescriptiveSourceSpan("catalog", start, start + len(row), row)
            self.assertEqual(descriptive_nonprose_output_path(span, record, profile=IPOS_DESCRIPTIVE_PROFILE),
                             "register" if index < 4 else "")

    def test_data_format_tables_route_without_exempting_computation_or_behavior(self):
        from workflow_routing import descriptive_nonprose_output_path
        for text, expected in (
            ("The stored data shall have the following tag: DATA TYPE BIT: 7 - 0 SAMPLE DATA Table 4: Data Tags", "table"),
            ("The stored data shall have the following tag: DATA TYPE BIT: 7 - 0 SAMPLE DATA", "table"),
            ("When in segmented mode, the offsets shall coincide with the following table: OFFSET DATA TYPE 0 SAMPLE Table 5: Offset Data Table", "table"),
            ("Example: Memory Base Address = 0xA000, then: Buffer 0 memory location 0xA000 Buffer 1 memory location 0xA100 Table 6: Memory Address", "table"),
            ("The result shall be computed as follows: INPUT OUTPUT 1 2 Table 7: Result Computation", ""),
            ("The stored data shall have the following tag: DATA TYPE BIT: 7 - 0 SAMPLE DATA Table 4: Data Tags and Module shall start sampling", ""),
            ("The stored data shall have the following tag: DATA TYPE BIT: 7 - 0 SAMPLE DATA Table 4: Data Tags and start sampling", ""),
        ):
            with self.subTest(text=text):
                record = NormalizedSourceRecord("structured", text, "approved_normative", "block_local",
                                                ("approved snapshot",), "retained", "approved structured evidence")
                span = DescriptiveSourceSpan("structured", 0, len(text), text)
                self.assertEqual(descriptive_nonprose_output_path(span, record, profile=IPOS_DESCRIPTIVE_PROFILE), expected)

    def test_guarded_format_table_routes_as_one_contiguous_structured_obligation(self):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings
        functional = build_ipos_functional_input("Module",
            {"Function": "Store samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "format", "statement":
              "When in segmented mode, the offsets shall coincide with the following table: OFFSET DATA TYPE 0 SAMPLE Table 5: Offset Data Table"}],
            domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        disposition = next(value for value in plan.dispositions if value.reason == "nonprose_output_path:table")
        self.assertEqual(len(disposition.spans), 2)
        self.assertIn("When in segmented mode,", disposition.spans[0].text)

    def test_identifier_retrieval_is_bounded_approved_scoped_and_exact_first(self):
        from workflow_routing import retrieve_descriptive_identifier_support, _ipos_admitted_records
        functional = build_ipos_functional_input("Module Core",
            {"Function": "Acquire samples", "Inputs": "sample ready: Ready indication", "Outputs": "results"},
            [{"mapped_block": "Module Core", "source_req_id": "local", "statement": "When sample_ready is high, Module Core shall acquire samples."}],
            domain="digital", candidate_requirement_rows=[
                {"mapped_block": "Other", "source_req_id": "foreign", "statement": "sample_ready controls unrelated behavior."},
                {"mapped_block": "Module Core", "source_req_id": "local", "statement": "When sample_ready is high, Module Core shall acquire samples."}])
        records = _ipos_admitted_records(functional)
        matches = retrieve_descriptive_identifier_support("sample_ready", records, profile=IPOS_DESCRIPTIVE_PROFILE)
        self.assertEqual([value["match"] for value in matches], ["exact", "normalized_variant"])
        self.assertEqual(matches[1]["record_id"], "inventory:inputs")
        self.assertFalse(any("Other" in value["record_id"] for value in matches))
        self.assertEqual(retrieve_descriptive_identifier_support("sample_rdy", records, profile=IPOS_DESCRIPTIVE_PROFILE), ())
        self.assertEqual(retrieve_descriptive_identifier_support("o_sample_ready", records, profile=IPOS_DESCRIPTIVE_PROFILE), ())
        self.assertEqual(retrieve_descriptive_identifier_support("sample_ready",
            tuple(replace(record, decision="rejected") for record in records), profile=IPOS_DESCRIPTIVE_PROFILE), ())
        with self.assertRaises(ValueError):
            retrieve_descriptive_identifier_support("sample_ready", records, profile=IPOS_DESCRIPTIVE_PROFILE, max_candidates=1)

    def test_normalized_guard_actor_preserves_authored_spelling_and_authority(self):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings
        functional = build_ipos_functional_input("Module Core",
            {"Function": "Acquire samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "guard", "statement": "During startup MODULE_CORE shall acquire samples."}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        proposition = next(value for value in plan.evidence.propositions if value.meaning.predicate == "acquire")
        self.assertEqual(proposition.meaning.subject, "MODULE_CORE")
        self.assertEqual(proposition.owner, "Module Core")

    def test_ipos_normalized_io_actor_requires_a_unique_approved_description(self):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings, _ipos_evidence_semantic_findings
        inventory = {"Function": "Acquire samples", "Inputs": "samples", "Outputs": "o_sample_ready: Completion indication"}
        statement = "When acquisition finishes O-SAMPLE-READY shall be set high."
        functional = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "guard", "statement": statement}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        proposition = next(value for value in plan.evidence.propositions if value.meaning.subject == "O-SAMPLE-READY")
        self.assertEqual(proposition.meaning.contexts[0].text, "When acquisition finishes")
        changed = replace(proposition, meaning=replace(proposition.meaning, subject="OTHER-SIGNAL"))
        self.assertTrue(_ipos_evidence_semantic_findings(replace(plan.evidence,
            propositions=tuple(changed if value == proposition else value for value in plan.evidence.propositions))))
        for outputs in ("o_sample_ready", "o_sample_ready: Completion indication; o-sample-ready: Error indication"):
            with self.subTest(outputs=outputs):
                unsupported = build_ipos_functional_input("Module", dict(inventory, Outputs=outputs),
                    [{"source_req_id": "guard", "statement": statement}], domain="digital")
                self.assert_nonessential_ipos_residual(unsupported)

    def test_ipos_formula_tail_is_structured_not_a_descriptive_obligation(self):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings, ipos_descriptive_structural_audit
        functional = build_ipos_functional_input("Module",
            {"Function": "Acquire samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "interval", "statement":
              "The interval shall define how many slots pass between acquisitions following this formula: interval = slots * scale"}],
            domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        self.assertTrue(any(value.reason == "nonprose_output_path:parameter" and value.decision == "excluded"
                            for value in plan.dispositions))
        self.assertTrue(any(value.meaning.predicate == "define" for value in plan.evidence.propositions))
        self.assertIn("interval = slots * scale", plan.evidence.records[-1].text)
        self.assertEqual(ipos_descriptive_structural_audit(functional)["source_findings"], [])

    def test_ipos_formula_tail_with_behavior_retains_nonessential_residual(self):
        from workflow_routing import ipos_claim_preflight_findings
        functional = build_ipos_functional_input("Module",
            {"Function": "Acquire samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "interval", "statement":
              "The interval shall define how many slots pass between acquisitions following this formula: interval = slots * scale and then sampling starts"}],
            domain="digital")
        self.assert_nonessential_ipos_residual(functional)

    def test_ipos_priority_list_preserves_ranks_order_and_direction(self):
        from workflow_routing import (
            build_ipos_claim_plan, ipos_claim_preflight_findings, _ipos_evidence_semantic_findings,
            realize_descriptive_functional_claim, _ipos_readback_section_findings,
        )
        inventory = {"Function": "Process commands", "Inputs": "commands", "Outputs": "results"}
        source = ("If multiple command bits are written, Module shall execute only the routine at higher priority and discard the others. "
                  "Priority is listed below, with first routine at higher priority: 1. RESET routine 2. WRITE routine 3. TEST routine 4. BOOT routine")
        functional = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "priority", "statement": source}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        priority = next(value for value in plan.evidence.propositions if value.proposition_id.endswith(":priority"))
        self.assertEqual(priority.meaning.objects, ("1: RESET routine", "2: WRITE routine", "3: TEST routine", "4: BOOT routine"))
        self.assertTrue(any(edge.target_id == priority.proposition_id for edge in plan.evidence.dependencies))
        for changed in (
            replace(priority, meaning=replace(priority.meaning, objects=tuple(reversed(priority.meaning.objects)))),
            replace(priority, meaning=replace(priority.meaning, qualifiers=(", from lower to higher priority",))),
        ):
            mutated = replace(plan.evidence, propositions=tuple(changed if value == priority else value for value in plan.evidence.propositions))
            self.assertTrue(_ipos_evidence_semantic_findings(mutated))
        for claim in plan.claims:
            rendered = realize_descriptive_functional_claim(claim)
            self.assertEqual(_ipos_readback_section_findings(rendered, (claim,)), [])
        broken = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "broken", "statement": source.replace("3. TEST", "2. TEST")}], domain="digital")
        self.assert_nonessential_ipos_residual(broken)

    def test_ipos_case_reference_binds_only_explicit_check_in_same_requirement(self):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings, _ipos_evidence_semantic_findings
        inventory = {"Function": "Process commands", "Inputs": "commands", "Outputs": "results"}
        source = ("During verification, Module shall compare each stored value with its register value to check if programming has been successful. "
                  "In that case Module shall write 0xFF into the last memory location.")
        functional = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "case", "statement": source}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        consequence = next(value for value in plan.evidence.propositions if value.meaning.predicate == "write")
        self.assertEqual(consequence.meaning.contexts[0].text, "If programming has been successful")
        self.assertEqual(len(consequence.spans), 2)
        self.assertTrue(any(edge.target_id == consequence.proposition_id for edge in plan.evidence.dependencies))
        self.assertTrue(any(len(claim.proposition_ids) == 2 for claim in plan.claims))
        changed = replace(consequence, meaning=replace(consequence.meaning,
            contexts=(DescriptiveContext("atom", "If programming has failed"),)))
        mutated = replace(plan.evidence, propositions=tuple(changed if value == consequence else value for value in plan.evidence.propositions))
        self.assertTrue(_ipos_evidence_semantic_findings(mutated))
        self.assertTrue(_ipos_evidence_semantic_findings(replace(plan.evidence, dependencies=())))
        unresolved = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "no-antecedent", "statement": "In that case Module shall write 0xFF into memory."}], domain="digital")
        self.assert_nonessential_ipos_residual(unresolved)
        ambiguous = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "ambiguous", "statement": source.replace("to check if programming has been successful", "during programming")}], domain="digital")
        self.assert_nonessential_ipos_residual(ambiguous)

    def test_ipos_comma_less_guard_requires_exact_approved_owner_and_preserves_condition(self):
        from workflow_routing import build_ipos_claim_plan, ipos_claim_preflight_findings, _ipos_evidence_semantic_findings
        inventory = {"Function": "Process commands", "Inputs": "commands", "Outputs": "results"}
        statement = "During programming routine Module shall copy the register contents into memory following the memory procedure."
        functional = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "guard", "statement": statement}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertEqual(ipos_claim_preflight_findings(functional), [])
        proposition = next(value for value in plan.evidence.propositions if value.spans[0].record_id.startswith("requirement:"))
        self.assertEqual(proposition.meaning.subject, "Module")
        self.assertEqual(proposition.meaning.contexts[0].text, "During programming routine")
        changed = replace(proposition, meaning=replace(proposition.meaning, contexts=()))
        mutated = replace(plan.evidence, propositions=tuple(changed if value == proposition else value for value in plan.evidence.propositions))
        self.assertTrue(_ipos_evidence_semantic_findings(mutated))
        ambiguous = build_ipos_functional_input("Module", inventory,
            [{"source_req_id": "ambiguous", "statement": statement.replace("routine Module", "routine Another unit")}], domain="digital")
        self.assert_nonessential_ipos_residual(ambiguous)

    def test_ipos_modal_predicates_preserve_inflection_order_and_processing_actions(self):
        from workflow_routing import (
            build_ipos_claim_plan, ipos_claim_preflight_findings,
            realize_descriptive_functional_claim, _ipos_readback_section_findings,
        )
        statements = (
            "The initialization routine shall starts when setup_ready rises.",
            "The initialization routine shall completes when setup_done rises.",
            "During initialization, Processor shall first read the last memory location at address 0x7F.",
            "During verification, Processor shall compare each stored value with the corresponding register value.",
            "After startup, Processor shall ignore further start transactions.",
            "When multiple commands arrive, Processor shall execute only the command at higher priority and discard the others.",
        )
        for index, statement in enumerate(statements):
            with self.subTest(statement=statement):
                functional = build_ipos_functional_input("Module",
                    {"Function": "Process commands", "Inputs": "commands", "Outputs": "results"},
                    [{"source_req_id": str(index), "statement": statement}], domain="digital")
                plan = build_ipos_claim_plan(functional)
                self.assertFalse(plan.evidence.unresolved_contexts)
                self.assertEqual(ipos_claim_preflight_findings(functional), [])
                self.assertEqual(plan.evidence.records[-1].text, statement)
                for claim in plan.claims:
                    rendered = realize_descriptive_functional_claim(claim)
                    self.assertEqual(_ipos_readback_section_findings(rendered, (claim,)), [])
                    if any("first" in meaning.qualifiers for meaning in claim.meanings):
                        self.assertTrue(_ipos_readback_section_findings(rendered.replace("first ", ""), (claim,)))

    def test_guarded_passive_relationship_is_not_excluded_as_an_atomic_signal(self):
        from workflow_routing import ipos_root_cause_report
        functional = build_ipos_functional_input("Module", {"Function": "Process samples", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "bypass", "statement": "When override is enabled, the calibration result shall be bypassed in favor of the configured value."},
             {"source_req_id": "write", "statement": "When processing completes, the data shall be written inside the buffer using the local interface."},
             {"source_req_id": "signal", "statement": "When enabled, Output shall be set high."}], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        output = " ".join(sections["internal_functions"])
        self.assertIn("calibration result is bypassed in favor of the configured value", output)
        self.assertIn("data is written inside the buffer using the local interface", output)
        self.assertIn("output is set high", output)
        self.assertNotIn("output is set high", " ".join(sections["scope"]))
        self.assertFalse(any(row["decision"] == "projection_gap" for row in audit))
        report = ipos_root_cause_report(functional)
        self.assertEqual(len(report["classes"]), 5)
        self.assertEqual(report["release_status"], "NOT_CERTIFIED")
        self.assertEqual(sum(value["counts"]["unresolved_context"] for value in report["classes"]), report["unresolved_contexts"])

    def test_ipos_interpretation_preserves_exact_independent_clause_before_prose(self):
        from workflow_routing import build_ipos_claim_plan
        source = "Processor shall filter samples. Receiver shall wait. When configured ("
        functional = build_ipos_functional_input("Module", {"Function": "Coordinate sampling", "Inputs": "samples", "Outputs": "results"},
                     [{"source_req_id": "one", "statement": source}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        claim = next(value for value in plan.claims if value.responsibilities == ("Processor",))
        self.assertEqual(claim.meanings[0].predicate, "filter")
        self.assertEqual(claim.spans[0].text, "Processor shall filter samples.")
        self.assertEqual(source[claim.spans[0].start:claim.spans[0].end], claim.spans[0].text)
        self.assert_nonessential_ipos_residual(functional)
        self.assertFalse(hasattr(claim, "text"))
        self.assertEqual(plan.acceptance, DescriptiveAcceptance())

    def test_ipos_unresolved_reference_withholds_dependent_not_independent_responsibility(self):
        from workflow_routing import build_ipos_claim_plan
        functional = build_ipos_functional_input("Module", {"Function": "Coordinate sampling", "Inputs": "samples", "Outputs": "results"},
                     [{"source_req_id": "one", "statement": "Processor shall filter samples. Receiver shall wait. It shall wait until"}], domain="digital")
        plan = build_ipos_claim_plan(functional)
        self.assertTrue(any(value.responsibilities == ("Processor",) for value in plan.claims))
        self.assertTrue(any(value.responsibilities == ("Receiver",) for value in plan.claims))
        self.assertEqual(plan.evidence.dependencies, ())
        self.assertTrue(plan.evidence.unresolved_contexts)

    def test_claim_scope_cannot_promote_initiation_into_result_or_interaction(self):
        evidence = self.synthesis_evidence()
        first = replace(evidence.propositions[0], meaning=DescriptiveMeaning("behavior", "Processor", "start", ("acquisition",)))
        evidence = replace(evidence, propositions=(first,), dependencies=(), unresolved_contexts=())
        intent = DescriptiveClaimIntent("start", "functions", (first.proposition_id,), scope="initiation")
        plan = build_descriptive_synthesis_plan(evidence, (intent,), profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(plan.claims[0].scope, "initiation")
        self.assertEqual(plan.claims[0].responsibilities, ("Processor",))
        for scope in ("result", "local_interaction", "exception", "lifecycle"):
            with self.subTest(scope=scope):
                changed = build_descriptive_synthesis_plan(evidence, (replace(intent, scope=scope),),
                          profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
                self.assertEqual(changed.claims, ())
                self.assertEqual(changed.decisions[0].decision, "unresolved")

    def test_ipos_publication_checks_actual_text_independently_of_audit_and_composer(self):
        self.assertEqual(ipos_publication_quality_findings({
            "1.1": "The module processes samples.",
            "1.2": "The acquisition interval is configurable.",
            "1.4": "When READY_STATE is active, SAMPLE_LIMIT defines the acquisition interval.",
        }), [])
        for text, finding in (
            ("The DEVICE FSM FSM enters the error state.", "duplicated_entity_label"),
            ("When the counter reach its limit, the FSM raises an interrupt.", "subject_verb_agreement"),
            ("The user starts measurement writing register START at 1.", "missing_action_link"),
            ("When configured (", "unbalanced_delimiters"),
            ("The module transfers samples into.", "unfinished_clause"),
            ("- **Acquire samples.** The module is designed to acquire samples.", "heading_body_tautology"),
        ):
            with self.subTest(text=text):
                self.assertTrue(any(finding in value for value in ipos_publication_quality_findings({"1.4": text})))
        self.assertTrue(any("duplicated_explanation" in value for value in ipos_publication_quality_findings({
            "1.1": "The module processes samples.", "1.2": "The module processes samples.",
        })))

    def test_ipos_projection_retains_exact_identifiers_and_does_not_gate_supported_verbs(self):
        functional = build_ipos_functional_input("Module",
            {"Function": "Coordinate sampling", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "When READY_STATE is active, SAMPLE_LIMIT shall define the acquisition interval."}],
            domain="digital")
        flow = _ipos_shared_flow(functional)
        fact = next(value for value in flow.facts if value.evidence_kind == "approved_normative")
        self.assertEqual(fact.subject, "SAMPLE_LIMIT")
        self.assertIn("READY_STATE", fact.condition)
        self.assertIn("SAMPLE_LIMIT", " ".join(value.text for value in flow.paragraphs))
        changed = replace(fact, subject="sAMPLE_LIMIT")
        self.assertTrue(any("ipos_source_identifier_changed" in finding for finding in
                   ipos_source_preservation_findings(flow.records,
                   tuple(changed if value.fact_id == fact.fact_id else value for value in flow.facts))))
        application = ipos_descriptive_structural_audit(functional)["application"]
        self.assertEqual(next(value["decision"] for value in application["decisions"]
                              if "requirement:Module:one:" in value["proposition_id"]), "retained")

    def test_section_contributions_reject_relabelled_meaning_and_retain_supported_detail(self):
        meaning = DescriptiveMeaning("behavior", "Module", "process", ("samples",))
        purpose = DescriptiveFunctionalClaim("purpose", "purpose", ("source",), (meaning,), (), ())
        repeated = replace(purpose, claim_id="new-title", role="main_functions", proposition_ids=("other-id",))
        detail = replace(purpose, claim_id="detail", role="functions",
                         meanings=(replace(meaning, contexts=(DescriptiveContext("atom", "while enabled"),),
                                           results=("filtered samples",)),))
        decisions = descriptive_section_contributions((purpose, repeated, detail))
        self.assertEqual(tuple(value.decision for value in decisions), ("retained", "excluded", "retained"))
        self.assertEqual(decisions[1].reason, "equivalent_content_already_expressed")
        self.assertEqual(decisions[2].added_meanings, detail.meanings)

    def test_functional_claim_is_explicit_and_independent_of_prose_units(self):
        meaning = DescriptiveMeaning("transformation", "Processor", "filter", ("samples",),
                                     results=("filtered samples",))
        source = "Processor filters samples."
        span = DescriptiveSourceSpan("source", 0, len(source), source)
        intent = DescriptiveClaimIntent("claim", "functions", ("filtering",))
        claim = DescriptiveFunctionalClaim(intent.claim_id, intent.role, intent.proposition_ids,
                                           (meaning,), (), (span,))
        self.assertFalse(hasattr(claim, "unit_id"))
        self.assertEqual(claim.meanings[0].results, ("filtered samples",))
        self.assertEqual(claim.synthesis_rule, "explicit_relationship_bundle-v1")

    def test_descriptive_acceptance_requires_three_independent_passes(self):
        pending = DescriptiveAcceptance()
        self.assertEqual(pending.verification_status, "UNRESOLVED")
        accepted = replace(pending, semantic_fidelity="PASS", publication_quality="PASS", materialization="PASS")
        self.assertEqual(accepted.verification_status, "PASS")
        for boundary in ("semantic_fidelity", "publication_quality", "materialization"):
            self.assertEqual(replace(accepted, **{boundary: "UNRESOLVED"}).verification_status, "UNRESOLVED")
            self.assertEqual(replace(accepted, **{boundary: "FAIL"}).verification_status, "FAIL")
            self.assertEqual(replace(accepted, **{boundary: "model_valid"}).verification_status, "FAIL")
        for boundary in ("semantic_findings", "publication_findings", "materialization_findings"):
            self.assertEqual(replace(accepted, **{boundary: ("unresolved_obligation",)}).verification_status, "FAIL")

    def synthesis_evidence(self, profile=IPOS_DESCRIPTIVE_PROFILE):
        text = "Processor filters samples. Receiver waits. When configured ("
        record = NormalizedSourceRecord("record", text, profile.authority_tiers[0], profile.allowed_scopes[0],
                                        ("source",), "retained", "admitted")
        first_end = text.index(" Receiver")
        second_start = first_end + 1
        second_end = text.index(" When")
        spans = (DescriptiveSourceSpan("record", 0, first_end, text[:first_end]),
                 DescriptiveSourceSpan("record", second_start, second_end, text[second_start:second_end]))
        meanings = (DescriptiveMeaning("transformation", "Processor", "filter", ("samples",)),
                    DescriptiveMeaning("behavior", "Receiver", "wait"))
        propositions = tuple(DescriptiveProposition(identity, meaning, (span,), "Module",
                             profile.allowed_domains[0] if profile.allowed_domains else "digital",
                             profile.allowed_layers[0] if profile.allowed_layers else "system")
                             for identity, meaning, span in zip(("filtering", "waiting"), meanings, spans))
        context_start = second_end + 1
        unresolved = DescriptiveUnresolvedContext("configuration",
                     (DescriptiveSourceSpan("record", context_start, len(text), text[context_start:]),),
                     "incomplete_condition")
        edge = DescriptiveDependency("guard", "referenced_context", "waiting", "configuration", unresolved.spans)
        return DescriptiveEvidenceModel((record,), propositions, (edge,), (unresolved,))

    def test_synthesis_preserves_independent_clause_in_same_record_with_unresolved_context(self):
        for policy in DESCRIPTIVE_APPLICATION_PROFILES.values():
            profile = next(value for value in DESCRIPTIVE_PROFILES.values() if value.name == policy.writing_profile)
            with self.subTest(policy=policy.name):
                evidence = self.synthesis_evidence(profile)
                evidence = replace(evidence, propositions=tuple(replace(value, domain=policy.domains[0],
                                   layer=policy.layers[0]) for value in evidence.propositions))
                intents = (DescriptiveClaimIntent("filter", "functions", ("filtering",)),
                           DescriptiveClaimIntent("wait", "functions", ("waiting",)))
                plan = build_descriptive_synthesis_plan(evidence, intents, profile=profile, owner="Module")
                self.assertEqual(plan.findings, ())
                self.assertEqual(tuple(claim.claim_id for claim in plan.claims), ("filter",))
                self.assertEqual(plan.decisions[0].decision, "preserved")
                self.assertIn("unresolved_dependency:configuration", plan.decisions[1].findings)
                self.assertEqual(plan.acceptance, DescriptiveAcceptance())
                self.assertEqual(plan.acceptance.verification_status, "UNRESOLVED")
                findings = validate_descriptive_synthesis_plan(plan, expected_evidence=evidence,
                           expected_intents=intents, profile=profile, owner="Module")
                if profile.document_type == "IPOS":
                    self.assertEqual(findings, [])
                    self.assertTrue(any(value.reason.startswith("omitted_nonessential:")
                                        and value.decision == "excluded" for value in plan.dispositions))
                else:
                    self.assertTrue(findings)

    def test_synthesis_transitive_unresolved_context_blocks_only_related_claims(self):
        evidence = self.synthesis_evidence()
        edge = DescriptiveDependency("sequence", "before", "filtering", "waiting", evidence.propositions[0].spans)
        evidence = replace(evidence, dependencies=(edge, *evidence.dependencies))
        intent = DescriptiveClaimIntent("filter", "functions", ("filtering",))
        plan = build_descriptive_synthesis_plan(evidence, (intent,), profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(plan.claims, ())
        self.assertIn("unresolved_dependency:configuration", plan.decisions[0].findings)
        reverse = replace(edge, dependency_id="reverse", source_id="waiting", target_id="filtering")
        invalid = replace(evidence, dependencies=(edge, reverse))
        plan = build_descriptive_synthesis_plan(invalid, (intent,), profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertIn("descriptive_dependency_cycle", plan.findings)

    def test_synthesis_audit_cannot_promote_model_integrity_to_publication_acceptance(self):
        evidence = self.synthesis_evidence()
        evidence = replace(evidence, records=(replace(evidence.records[0], text="Processor filters samples."),),
                           propositions=evidence.propositions[:1], dependencies=(), unresolved_contexts=())
        intents = (DescriptiveClaimIntent("filter", "functions", ("filtering",)),)
        plan = build_descriptive_synthesis_plan(evidence, intents, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        audit = descriptive_synthesis_plan_audit(plan, expected_evidence=evidence, expected_intents=intents,
                                                profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(audit["typed_model_status"], "model_valid")
        self.assertEqual(audit["source_semantic_validation"], "not_implemented")
        self.assertEqual(audit["final_prose_validation"], "not_implemented")
        self.assertEqual(audit["acceptance"]["publication_quality"], "UNRESOLVED")
        self.assertEqual(json.loads(json.dumps(audit)), audit)
        forged = replace(plan, acceptance=DescriptiveAcceptance("PASS", "PASS", "PASS"))
        audit = descriptive_synthesis_plan_audit(forged, expected_evidence=evidence, expected_intents=intents,
                                                profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(audit["typed_model_status"], "blocked")
        self.assertIn("descriptive_synthesis_unverified_acceptance", audit["findings"])
        self.assertEqual(audit["acceptance"]["publication_quality"], "UNRESOLVED")

    def test_synthesis_closure_preserves_direction_context_and_many_source_provenance(self):
        for profile in DESCRIPTIVE_PROFILES.values():
            evidence = self.synthesis_evidence(profile)
            record = replace(evidence.records[0], record_id="guard-source", text="After filtering, Receiver waits.")
            edge_span = DescriptiveSourceSpan(record.record_id, 0, len(record.text), record.text)
            edge = DescriptiveDependency("sequence", "after", "waiting", "filtering", (edge_span,))
            evidence = replace(evidence, records=(replace(evidence.records[0], text="Processor filters samples. Receiver waits."),
                                                  record), dependencies=(edge,), unresolved_contexts=())
            intents = (DescriptiveClaimIntent("pipeline", "functions", ("filtering",), ("sequence",)),)
            plan = build_descriptive_synthesis_plan(evidence, intents, profile=profile, owner="Module")
            self.assertEqual(plan.claims[0].proposition_ids, ("filtering", "waiting"))
            self.assertEqual(plan.claims[0].dependencies, (edge,))
            self.assertEqual({span.record_id for span in plan.claims[0].spans}, {"record", "guard-source"})
            self.assertEqual(validate_descriptive_synthesis_plan(plan, expected_evidence=evidence,
                             expected_intents=intents, profile=profile, owner="Module"), [])
            claim = plan.claims[0]
            changed = replace(claim.meanings[0], negative=True, coordination="or")
            for mutation in (replace(plan, claims=(replace(claim, meanings=(changed, claim.meanings[1])),)),
                             replace(plan, claims=(replace(claim, dependencies=()),)),
                             replace(plan, claims=(replace(claim, role="organization"),)),
                             replace(plan, version="unknown"),
                             replace(plan, dispositions=()),
                             replace(plan, evidence=replace(evidence, dependencies=())),
                             replace(plan, acceptance=DescriptiveAcceptance("PASS", "PASS", "PASS"))):
                self.assertTrue(validate_descriptive_synthesis_plan(mutation, expected_evidence=evidence,
                                expected_intents=intents, profile=profile, owner="Module"))

    def test_synthesis_rejects_name_based_grouping_and_records_all_unmodeled_text(self):
        evidence = self.synthesis_evidence()
        disconnected = replace(evidence, dependencies=())
        intents = (DescriptiveClaimIntent("combined", "functions", ("filtering", "waiting")),)
        plan = build_descriptive_synthesis_plan(disconnected, intents, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(plan.claims, ())
        self.assertIn("disconnected_claim_bundle", plan.decisions[0].findings)
        missing = replace(evidence, unresolved_contexts=(), dependencies=())
        plan = build_descriptive_synthesis_plan(missing, (), profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        omitted = [disposition for disposition in plan.dispositions
               if disposition.reason == "omitted_nonessential:unmodeled_source_span"]
        self.assertEqual(len(omitted), 1)
        self.assertEqual(omitted[0].decision, "excluded")
        self.assertEqual(omitted[0].spans[0].text.strip(), "When configured (")
        covered = {offset for disposition in plan.dispositions for span in disposition.spans
                   for offset in range(span.start, span.end)}
        self.assertTrue(all(offset in covered for offset, character in enumerate(evidence.records[0].text)
                            if not character.isspace()))
        self.assertTrue(validate_descriptive_synthesis_plan(replace(plan, dispositions=()),
                expected_evidence=missing, expected_intents=(), profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_synthesis_rejects_scope_span_identity_role_and_dependency_mutations(self):
        evidence = self.synthesis_evidence()
        intent = DescriptiveClaimIntent("filter", "functions", ("filtering",))
        for mutated in (replace(evidence, propositions=(replace(evidence.propositions[0], owner="Other"),)),
                        replace(evidence, dependencies=(replace(evidence.dependencies[0], target_id="missing"),)),
                        replace(evidence, unresolved_contexts=(replace(evidence.unresolved_contexts[0], context_id="waiting"),)),
                        replace(evidence, propositions=(replace(evidence.propositions[0], spans=(
                            replace(evidence.propositions[0].spans[0], text="Invented"),)),)),
                        replace(evidence, propositions=(replace(evidence.propositions[0], spans=(
                            replace(evidence.propositions[0].spans[0], end="invalid"),)),))):
            plan = build_descriptive_synthesis_plan(mutated, (intent,), profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
            self.assertTrue(plan.findings)
            self.assertEqual(plan.claims, ())
        for mutated in (replace(intent, role="missing"), replace(intent, role="organization"),
                        replace(intent, proposition_ids=()), replace(intent, required_dependency_ids=("missing",)),
                        replace(intent, required_dependency_ids=("guard",))):
            plan = build_descriptive_synthesis_plan(evidence, (mutated,), profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
            self.assertEqual(plan.claims, ())
            self.assertEqual(plan.decisions[0].decision, "unresolved")

    def test_application_profiles_keep_runtime_activation_and_writing_fingerprints_unchanged(self):
        self.assertEqual(len(DESCRIPTIVE_APPLICATION_PROFILES), 5)
        self.assertFalse(DRS_DESCRIPTIVE_PROFILE.runtime_enabled)
        self.assertFalse(ARS_DESCRIPTIVE_PROFILE.runtime_enabled)
        for policy in DESCRIPTIVE_APPLICATION_PROFILES.values():
            profile = next(value for value in DESCRIPTIVE_PROFILES.values() if value.name == policy.writing_profile)
            self.assertTrue(set(dict(policy.role_sections).values()).issubset(profile.topic_roles))
            self.assertEqual(policy.presentation_order, ("behavior", "support", "clarification"))

    def test_application_keeps_composition_explicit_and_rejects_a_capability_role(self):
        structure = self.structure()
        proposition = structure.propositions[0]
        policy = DESCRIPTIVE_APPLICATION_PROFILES["digital-ipos-structure-v1"]
        assessment = DescriptiveAssessment("composition", "organization", "functional", ("organization",))
        decision = apply_descriptive_profile((proposition,), (assessment,), policy=policy,
                   profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")[0]
        self.assertEqual(decision.decision, "retained")
        self.assertEqual(decision.rule, "essential_meaning")
        changed = apply_descriptive_profile((proposition,), (replace(assessment, role="functions"),), policy=policy,
                  profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")[0]
        self.assertEqual((changed.decision, changed.rule), ("unresolved", "relation_role_mismatch"))

    def test_application_excludes_local_detail_from_srs_but_retains_it_for_ipos(self):
        proposition = self.structure().propositions[1]
        assessment = DescriptiveAssessment("waiting", "operating_concept", "encoding", ())
        srs = apply_descriptive_profile((proposition,), (assessment,),
              policy=DESCRIPTIVE_APPLICATION_PROFILES["srs-structure-v1"], profile=SRS_DESCRIPTIVE_PROFILE)[0]
        ipos = apply_descriptive_profile((proposition,), (assessment,),
               policy=DESCRIPTIVE_APPLICATION_PROFILES["digital-ipos-structure-v1"],
               profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")[0]
        self.assertEqual((srs.decision, srs.rule), ("excluded", "outside_profile_layer"))
        self.assertEqual(ipos.decision, "retained")
        system = replace(proposition, layer="system")
        self.assertEqual(apply_descriptive_profile((system,), (assessment,),
            policy=DESCRIPTIVE_APPLICATION_PROFILES["srs-structure-v1"], profile=SRS_DESCRIPTIVE_PROFILE)[0].rule,
            "nonessential_encoding")

    def test_application_retains_essential_qualifiers_and_dependencies_over_detail_exclusions(self):
        original = self.structure().propositions[1]
        policy = DESCRIPTIVE_APPLICATION_PROFILES["digital-ipos-structure-v1"]
        assessment = DescriptiveAssessment("waiting", "operating_concept", "procedure", ())
        for meaning in (replace(original.meaning, qualifiers=("at least 2 ms",)),
                        replace(original.meaning, contexts=(DescriptiveContext("atom", "enabled"),)),
                        replace(original.meaning, negative=True), replace(original.meaning, modality="may")):
            proposition = replace(original, meaning=meaning)
            self.assertEqual(apply_descriptive_profile((proposition,), (assessment,), policy=policy,
                profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")[0].decision, "retained")
        second = self.structure().propositions[0]
        edge = DescriptiveDependency("edge", "before", "composition", "waiting", second.spans)
        decisions = apply_descriptive_profile((second, original),
            (DescriptiveAssessment("composition", "organization", "functional", ("organization",)), assessment),
            policy=policy, profile=IPOS_DESCRIPTIVE_PROFILE, dependencies=(edge,), owner="Module")
        self.assertEqual(decisions[1].decision, "retained")

    def test_application_distinguishes_digital_analog_and_explicit_missing_metadata(self):
        original = replace(self.structure().propositions[1], layer="integration")
        assessment = DescriptiveAssessment("waiting", "operating_concept", "functional", ("behavior",))
        for name, profile, domain, decision in (
            ("drs-structure-v1", DRS_DESCRIPTIVE_PROFILE, "digital", "retained"),
            ("ars-structure-v1", ARS_DESCRIPTIVE_PROFILE, "digital", "excluded"),
            ("ars-structure-v1", ARS_DESCRIPTIVE_PROFILE, "analog", "retained"),
            ("drs-structure-v1", DRS_DESCRIPTIVE_PROFILE, "analog", "excluded")):
            self.assertEqual(apply_descriptive_profile((replace(original, domain=domain),), (assessment,),
                policy=DESCRIPTIVE_APPLICATION_PROFILES[name], profile=profile)[0].decision, decision)
        for domain in ("", "digital"):
            proposition = replace(original, domain=domain, layer="block")
            decision = apply_descriptive_profile((proposition,), (assessment,),
                policy=DESCRIPTIVE_APPLICATION_PROFILES["analog-ipos-structure-v1"],
                profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")[0]
            self.assertEqual(decision.decision, "unresolved" if not domain else "excluded")

    def test_application_never_excludes_unsupported_syntax_by_profile(self):
        proposition = self.structure().propositions[1]
        for policy in DESCRIPTIVE_APPLICATION_PROFILES.values():
            profile = next(value for value in DESCRIPTIVE_PROFILES.values() if value.name == policy.writing_profile)
            assessment = DescriptiveAssessment("waiting", "operating_concept", "procedure", (), "unsupported_construction")
            decision = apply_descriptive_profile((proposition,), (assessment,), policy=policy, profile=profile, owner="Module")[0]
            self.assertEqual((decision.decision, decision.rule), ("unresolved", "unsupported_construction"))

    def test_application_inventory_and_audit_fail_closed_without_runtime_rollout(self):
        structure = self.structure()
        policy = DESCRIPTIVE_APPLICATION_PROFILES["digital-ipos-structure-v1"]
        assessments = (DescriptiveAssessment("composition", "organization", "functional", ("organization",)),
                       DescriptiveAssessment("waiting", "operating_concept", "functional", ("behavior",)))
        for values in ((), assessments[:1], assessments * 2):
            with self.assertRaises(ValueError):
                apply_descriptive_profile(structure.propositions, values, policy=policy,
                                          profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        with self.assertRaises(ValueError):
            apply_descriptive_profile(structure.propositions, assessments, policy=replace(policy, domains=()),
                                      profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        audit = descriptive_profile_application_audit(structure.propositions, assessments,
                policy=policy, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(audit["profile_status"], "blocked")
        self.assertIn("descriptive_required_role_missing:purpose", audit["findings"])
        self.assertEqual(audit["runtime_migration"], "not_implemented")
        changed = {**audit, "final_prose_validation": "PASS"}
        self.assertIn("descriptive_profile_application_audit_changed", validate_descriptive_profile_application_audit(
            changed, propositions=structure.propositions, assessments=assessments,
            policy=policy, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_application_audit_requires_common_integrity_even_for_excluded_detail(self):
        structure = self.structure(SRS_DESCRIPTIVE_PROFILE)
        policy = DESCRIPTIVE_APPLICATION_PROFILES["srs-structure-v1"]
        assessments = (DescriptiveAssessment("composition", "organization", "functional", ("organization",)),
                       DescriptiveAssessment("waiting", "operating_concept", "encoding", ()))
        audit = descriptive_profile_application_audit(structure.propositions, assessments,
                policy=policy, profile=SRS_DESCRIPTIVE_PROFILE, structure=structure)
        self.assertEqual(audit["profile_status"], "policy_valid")
        self.assertEqual(audit["shared_integrity_status"], "model_valid")
        self.assertEqual(audit["decisions"][1]["decision"], "excluded")
        self.assertEqual(validate_descriptive_profile_application_audit(json.loads(json.dumps(audit)),
            propositions=structure.propositions, assessments=assessments, policy=policy,
            profile=SRS_DESCRIPTIVE_PROFILE, structure=structure), [])
        missing = descriptive_profile_application_audit(structure.propositions, assessments,
                  policy=policy, profile=SRS_DESCRIPTIVE_PROFILE)
        self.assertIn("descriptive_profile_structure_missing", missing["findings"])
        broken = replace(structure, dispositions=())
        blocked = descriptive_profile_application_audit(structure.propositions, assessments,
                  policy=policy, profile=SRS_DESCRIPTIVE_PROFILE, structure=broken)
        self.assertEqual(blocked["shared_integrity_status"], "blocked")
        self.assertEqual(blocked["profile_status"], "blocked")
        inventory = descriptive_profile_application_audit(structure.propositions[:1], assessments[:1],
                    policy=policy, profile=SRS_DESCRIPTIVE_PROFILE, structure=structure)
        self.assertIn("descriptive_profile_structure_inventory_changed", inventory["findings"])

    def structure(self, profile=IPOS_DESCRIPTIVE_PROFILE):
        source = "Coordinator contains Receiver and Sender. Receiver waits."
        record = NormalizedSourceRecord("record-one", source, profile.authority_tiers[0],
                                        profile.allowed_scopes[0], ("source-one",), "retained", "admitted")
        span = DescriptiveSourceSpan(record.record_id, 0, len(source), source)
        meanings = (DescriptiveMeaning("containment", "Coordinator", "contain", ("Receiver", "Sender")),
                    DescriptiveMeaning("behavior", "Receiver", "wait"))
        propositions = tuple(DescriptiveProposition(identity, meaning, (span,), "Module",
                              profile.allowed_domains[0] if profile.allowed_domains else "digital",
                              profile.allowed_layers[0] if profile.allowed_layers else "system")
                             for identity, meaning in zip(("composition", "waiting"), meanings))
        unit = SemanticUnit("unit-one", profile.topic_roles[0], source, (record.record_id,), 1)
        claim = DescriptiveClaim("claim-one", unit.unit_id, ("composition", "waiting"), meanings)
        disposition = DescriptiveDisposition((span,), claim.proposition_ids, "preserved", "supported", (claim.claim_id,))
        return DescriptiveStructure(DESCRIPTIVE_STRUCTURE_VERSION, stable_payload_hash(asdict(profile)),
                                    (record,), propositions, (), (claim,), (disposition,), (unit,))

    def test_structure_supports_composition_and_intransitive_behavior_across_profiles(self):
        for profile in DESCRIPTIVE_PROFILES.values():
            with self.subTest(profile=profile.name):
                structure = self.structure(profile)
                self.assertEqual(validate_descriptive_structure(structure, profile=profile, owner="Module"), [])
                self.assertEqual(structure.propositions[1].meaning.objects, ())
        self.assertFalse(DRS_DESCRIPTIVE_PROFILE.runtime_enabled)
        self.assertFalse(ARS_DESCRIPTIVE_PROFILE.runtime_enabled)

    def test_structure_rejects_claim_mutations_without_losing_names_or_ids(self):
        for profile in DESCRIPTIVE_PROFILES.values():
            structure = self.structure(profile)
            original = structure.claims[0].meanings[0]
            for mutation in (replace(original, subject="Receiver", objects=("Coordinator", "Sender")),
                             replace(original, relation="interaction"), replace(original, subject="Module"),
                             replace(original, coordination="or"), replace(original, negative=True)):
                with self.subTest(profile=profile.name, mutation=mutation):
                    claim = replace(structure.claims[0], meanings=(mutation, structure.claims[0].meanings[1]))
                    self.assertIn("descriptive_claim_meaning_changed:claim-one", validate_descriptive_structure(
                        replace(structure, claims=(claim,)), profile=profile, owner="Module"))

    def test_structure_rejects_span_loss_and_unresolved_evidence(self):
        structure = self.structure()
        span = structure.dispositions[0].spans[0]
        shortened = replace(span, end=span.end - 1, text=span.text[:-1])
        for disposition in (replace(structure.dispositions[0], spans=(shortened,)),
                            replace(structure.dispositions[0], decision="unresolved", reason="unsupported_construction", claim_ids=())):
            self.assertTrue(validate_descriptive_structure(replace(structure, dispositions=(disposition,)),
                            profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_structure_dependencies_are_explicit_and_cannot_be_invented_or_dropped(self):
        structure = self.structure()
        edge = DescriptiveDependency("edge-one", "before", "composition", "waiting", structure.propositions[0].spans)
        claim = replace(structure.claims[0], dependencies=(edge,))
        supported = replace(structure, dependencies=(edge,), claims=(claim,))
        self.assertEqual(validate_descriptive_structure(supported, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"), [])
        for mutation in (replace(supported, claims=structure.claims),
                         replace(structure, claims=(claim,)),
                         replace(supported, claims=(replace(claim, dependencies=(replace(edge, relation="after"),)),))):
            self.assertTrue(validate_descriptive_structure(mutation, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_structure_preserves_boolean_context_optional_behavior_and_quantities(self):
        structure = self.structure()
        context = DescriptiveContext("and", operands=(DescriptiveContext("atom", "enabled"),
                   DescriptiveContext("not", operands=(DescriptiveContext("atom", "busy"),))))
        meaning = replace(structure.propositions[1].meaning, contexts=(context,), modality="may", qualifiers=("at least 2 ms",))
        proposition = replace(structure.propositions[1], meaning=meaning)
        claim = replace(structure.claims[0], meanings=(structure.propositions[0].meaning, meaning))
        supported = replace(structure, propositions=(structure.propositions[0], proposition), claims=(claim,))
        self.assertEqual(validate_descriptive_structure(supported, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"), [])
        for mutation in (replace(meaning, contexts=()), replace(meaning, modality=""),
                         replace(meaning, contexts=(replace(context, operator="or"),)),
                         replace(meaning, qualifiers=("at most 2 ms",))):
            mutated_claim = replace(claim, meanings=(structure.propositions[0].meaning, mutation))
            self.assertTrue(validate_descriptive_structure(replace(supported, claims=(mutated_claim,)),
                            profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_structure_requires_explicit_version_profile_and_complete_dispositions(self):
        structure = self.structure()
        for mutation in (replace(structure, version="unknown"), replace(structure, profile_fingerprint="changed"),
                         replace(structure, dispositions=()), replace(structure, claims=()),
                         replace(structure, records=structure.records * 2)):
            self.assertTrue(validate_descriptive_structure(mutation, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_structure_duplicate_exclusion_retains_all_provenance_and_requires_equivalence(self):
        structure = self.structure()
        record = replace(structure.records[0], record_id="record-two", provenance=("source-two",))
        span = replace(structure.propositions[0].spans[0], record_id=record.record_id)
        duplicate = replace(structure.propositions[0], proposition_id="duplicate", spans=(span,))
        disposition = DescriptiveDisposition((span,), (duplicate.proposition_id,), "excluded",
                                             "equivalent_duplicate", representative_id="composition")
        unit = replace(structure.units[0], source_record_ids=("record-one", "record-two"))
        supported = replace(structure, records=(*structure.records, record),
                            propositions=(*structure.propositions, duplicate),
                            dispositions=(*structure.dispositions, disposition), units=(unit,))
        self.assertEqual(validate_descriptive_structure(supported, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"), [])
        for mutation in (replace(supported, units=structure.units),
                         replace(supported, dispositions=(*structure.dispositions, replace(disposition, reason="parser_gap"))),
                         replace(supported, propositions=(*structure.propositions, replace(duplicate,
                                 meaning=replace(duplicate.meaning, negative=True))))):
            self.assertTrue(validate_descriptive_structure(mutation, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_structure_requires_dependency_provenance_and_rejects_causal_cycles(self):
        structure = self.structure()
        record = replace(structure.records[0], record_id="dependency-source", text="Receiver waits after composition.")
        span = DescriptiveSourceSpan(record.record_id, 0, len(record.text), record.text)
        edge = DescriptiveDependency("edge-one", "after", "waiting", "composition", (span,))
        claim = replace(structure.claims[0], dependencies=(edge,))
        unit = replace(structure.units[0], source_record_ids=("record-one", "dependency-source"))
        disposition = replace(structure.dispositions[0], spans=(*structure.dispositions[0].spans, span))
        supported = replace(structure, records=(*structure.records, record), dependencies=(edge,),
                            claims=(claim,), units=(unit,), dispositions=(disposition,))
        self.assertEqual(validate_descriptive_structure(supported, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"), [])
        self.assertTrue(validate_descriptive_structure(replace(supported, dispositions=structure.dispositions),
                        profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))
        reverse = replace(edge, dependency_id="edge-two", relation="before")
        cyclic = replace(supported, dependencies=(edge, reverse), claims=(replace(claim, dependencies=(edge, reverse)),))
        self.assertIn("descriptive_dependency_cycle", validate_descriptive_structure(
            cyclic, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_structure_versioned_audit_roundtrip_is_not_semantic_or_publication_pass(self):
        structure = self.structure()
        payload = json.loads(json.dumps(descriptive_structure_audit(structure, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")))
        self.assertEqual(validate_descriptive_structure_audit(payload, expected=structure,
                         profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"), [])
        self.assertEqual(payload["source_semantic_validation"], "not_implemented")
        self.assertEqual(payload["final_prose_validation"], "not_implemented")
        for field, value in (("contract_version", "legacy"), ("phase", "runtime"),
                             ("source_semantic_validation", "PASS"), ("profile", {}), ("structure", {})):
            changed = {**payload, field: value}
            self.assertIn("descriptive_structure_audit_changed", validate_descriptive_structure_audit(
                changed, expected=structure, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))
        blocked = replace(structure, dispositions=(replace(structure.dispositions[0], decision="unresolved",
                          reason="incomplete_source", claim_ids=()),))
        audit = descriptive_structure_audit(blocked, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(audit["typed_model_status"], "blocked")
        self.assertTrue(validate_descriptive_structure_audit(audit, expected=blocked,
                        profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))

    def test_structure_cannot_account_for_unprojected_text_with_a_preserved_disposition(self):
        structure = self.structure()
        extra = replace(structure.records[0], record_id="extra", text="Unknown construction.")
        span = DescriptiveSourceSpan(extra.record_id, 0, len(extra.text), extra.text)
        disposition = replace(structure.dispositions[0], spans=(*structure.dispositions[0].spans, span))
        mutated = replace(structure, records=(*structure.records, extra), dispositions=(disposition,))
        self.assertIn("descriptive_disposition_unmodeled_span", validate_descriptive_structure(
            mutated, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))
        unresolved = DescriptiveDisposition((span,), (), "unresolved", "unsupported_construction")
        diagnosed = replace(mutated, dispositions=(*structure.dispositions, unresolved))
        findings = validate_descriptive_structure(diagnosed, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module")
        self.assertEqual(findings, ["descriptive_structure_unresolved:unsupported_construction"])

    def test_structure_rejects_invalid_spans_contexts_order_scope_and_unknown_profiles(self):
        structure = self.structure()
        span = structure.propositions[0].spans[0]
        for changed_span in (replace(span, start=-1), replace(span, end=True), replace(span, text="changed"),
                             replace(span, record_id="missing")):
            proposition = replace(structure.propositions[0], spans=(changed_span,))
            self.assertTrue(validate_descriptive_structure(replace(structure,
                            propositions=(proposition, structure.propositions[1])),
                            profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))
        for context in (DescriptiveContext("atom"), DescriptiveContext("not"),
                        DescriptiveContext("and", operands=(DescriptiveContext("atom", "enabled"),)),
                        DescriptiveContext("unknown", "enabled")):
            meaning = replace(structure.propositions[0].meaning, contexts=(context,))
            proposition = replace(structure.propositions[0], meaning=meaning)
            self.assertIn("descriptive_context_invalid", validate_descriptive_structure(replace(structure,
                          propositions=(proposition, structure.propositions[1])),
                          profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))
        for mutation in (replace(structure, units=(replace(structure.units[0], output_order=0),)),
                         replace(structure, propositions=(replace(structure.propositions[0], owner="Other"),
                                                         structure.propositions[1]))):
            self.assertTrue(validate_descriptive_structure(mutation, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))
        unknown = replace(IPOS_DESCRIPTIVE_PROFILE, name="unknown")
        self.assertIn("descriptive_structure_contract_invalid", validate_descriptive_structure(
                      replace(structure, profile_fingerprint=stable_payload_hash(asdict(unknown))),
                      profile=unknown, owner="Module"))

    def profile(self, name, scopes):
        return DescriptiveWritingProfile(name, scopes, "TEST", ("descriptive",), ("admitted_context",),
                                         "unrestricted", (), (), ("synthetic",), (), (), True)

    def test_ipos_identity_text_pairing_survives_duplicate_text_and_fallback(self):
        statements = ("The controller shall store samples.", "The controller shall store samples.",
                      "The controller shall process results.")
        rows = [{"mapped_block": "Module", "source_req_id": identity, "statement": statement}
                for identity, statement in zip(("req-one", "req-two", "req-three"), statements)]
        record = build_ipos_functional_input(
            "Module", {"Function": "Store and process data", "Inputs": "samples", "Outputs": "results"}, rows,
        )
        self.assertEqual(record.local_requirement_ids, ("req-one", "req-two", "req-three"))
        self.assertEqual(record.local_requirement_evidence, statements)
        fallback = replace(record, candidate_requirement_records=())
        _sections, audit = compose_ipos_overview(fallback)
        candidates = {row["candidate_evidence_requirement_id"]: row["candidate_evidence_statement"]
                      for row in audit if row.get("candidate_evidence_statement")}
        self.assertEqual(candidates, dict(zip(record.local_requirement_ids, statements)))

    def test_structural_relation_mutations_fail_with_unchanged_names_and_contributors(self):
        profile = self.profile("synthetic", ("system",))
        fact = DescriptiveFact("Coordinator", "control", ("Receiver",), "source-one",
                               "Coordinator controls Receiver.", "system")
        for mutation in ("Receiver controls Coordinator.", "Receiver and Coordinator are present."):
            paragraph = DescriptiveParagraph(mutation, (fact,))
            findings = validate_descriptive_paragraphs([paragraph], facts=[fact], profile=profile)
            self.assertTrue(any("descriptive_relation_changed" in finding for finding in findings))
        for valid in ("Coordinator controls Receiver.", "Coordinator controls Receiver during operation."):
            self.assertEqual(validate_descriptive_paragraphs(
                [DescriptiveParagraph(valid, (fact,))], facts=[fact], profile=profile,
            ), [])

    def test_ipos_projection_has_no_character_cap_and_preserves_inventory_input_meaning(self):
        payload = "the acquisition results " + "from the independently controlled sampling domains " * 4
        statement = "Coordinator_FSM shall store " + payload + "with a dedicated tag."
        facets = _ipos_structural_behavior_facets(statement, "Coordinator_FSM")
        self.assertIn("with dedicated tag", facets["actions"][0])
        record = build_ipos_functional_input(
            "Module", {"Function": "Store results", "Inputs": "Configuration register settings",
                       "Outputs": "Stored results"},
        )
        sections, _audit = compose_ipos_overview(record)
        self.assertIn("configuration register settings", sections["scope"][0].casefold())

    def test_sequence_requires_explicit_source_support_not_contributor_order(self):
        profile = self.profile("synthetic", ("system",))
        fact = DescriptiveFact("Coordinator", "control", ("Receiver",), "source-one",
                               "Coordinator controls Receiver.", "system")
        for statement in (fact.statement, "After startup, Coordinator controls Receiver."):
            source_fact = replace(fact, statement=statement)
            findings = validate_descriptive_paragraphs(
                [DescriptiveParagraph("Coordinator controls Receiver, then starts processing.", (source_fact,))],
                facts=[source_fact], profile=profile,
            )
            self.assertTrue(any("descriptive_sequence_not_supported" in finding for finding in findings))

    def test_ipos_shared_migration_preserves_guards_polarity_and_composition(self):
        rows = [{"mapped_block": "Module", "source_req_id": identity, "statement": statement}
                for identity, statement in (
                    ("one", "When enabled, Coordinator shall transfer samples to Receiver."),
                    ("two", "Unless ready, Coordinator shall not perform acquisition."),
                    ("three", "Coordinator shall contain Receiver and Sender."))]
        sections, audit = compose_shared_ipos_overview(build_ipos_functional_input(
            "Module", {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"}, rows,
            domain="analog",
        ))
        output = " ".join((*sections["scope"], *sections["internal_structure"], *sections["internal_functions"]))
        for expected in ("When enabled", "samples to Receiver", "Unless ready", "does not perform acquisition",
                         "contains Receiver and Sender"):
            self.assertIn(expected, output)
        self.assertEqual({row["candidate_evidence_requirement_id"] for row in audit
                          if row["decision"] == "selected"} - {""}, {"one", "two", "three"})
        self.assertTrue(all(row["profile"] == "block_local_v1" for row in audit))

    def test_ipos_shared_migration_retains_unsupported_source_as_nonessential_residual(self):
        statement = "Coordinator shall store samples and Receiver shall process them."
        sections, audit = compose_shared_ipos_overview(build_ipos_functional_input(
            "Module", {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": statement}], domain="digital",
        ))
        gaps = [row for row in audit if row["decision"] == "omitted"]
        self.assertEqual(len(gaps), 1)
        self.assertEqual(gaps[0]["candidate_evidence_statement"], statement)
        self.assertTrue(gaps[0]["reason"].startswith("omitted_nonessential:"))
        self.assertNotIn("Receiver", " ".join(sections["scope"]))

    def test_ipos_runtime_adapter_materializes_common_structure_and_profile_roles_for_both_domains(self):
        rows = [{"source_req_id": "organization", "statement": "Coordinator shall contain Receiver and Sender."},
                {"source_req_id": "waiting", "statement": "Unless enabled, Receiver shall not toggle."}]
        for domain in ("digital", "analog"):
            functional_input = build_ipos_functional_input("Module",
                {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"}, rows, domain=domain)
            structure = build_ipos_descriptive_structure(functional_input)
            self.assertEqual(validate_descriptive_structure(structure, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"), [])
            self.assertEqual(structure.propositions[-3].meaning.relation, "containment")
            self.assertEqual(structure.propositions[-2].meaning.objects, ())
            audit = ipos_descriptive_structural_audit(functional_input)
            self.assertEqual(audit["application"]["profile_status"], "policy_valid")
            self.assertFalse(any(value["decision"] == "unresolved" for value in audit["section_plan"]))
            self.assertEqual(audit["source_findings"], [])
            self.assertEqual(audit["publication_status"], "requires_independent_final_artifact_checks")
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "descriptive_structure_audit.json"
                write_ipos_structural_audit(path, functional_input)
                self.assertEqual(validate_ipos_structural_audit(path, functional_input), audit["application"]["findings"])
                changed = json.loads(path.read_text(encoding="utf-8"))
                changed["model"]["plan"]["evidence"]["propositions"][0]["meaning"]["subject"] = "Receiver"
                path.write_text(json.dumps(changed), encoding="utf-8")
                self.assertIn("ipos_structural_audit_changed", validate_ipos_structural_audit(path, functional_input))

    def test_ipos_runtime_adapter_keeps_incomplete_source_as_audited_nonessential_residual(self):
        functional_input = build_ipos_functional_input("Module",
            {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "Coordinator shall transfer samples into the."}], domain="digital")
        audit = ipos_descriptive_structural_audit(functional_input)
        self.assertEqual(audit["model"]["typed_model_status"], "model_valid")
        self.assertEqual(audit["application"]["profile_status"], "policy_valid")
        self.assertEqual(audit["source_semantic_status"], "controlled_constructions_valid")
        self.assertEqual(audit["model"]["plan"]["dispositions"][-1]["decision"], "excluded")
        self.assertEqual(audit["publication_status"], "requires_independent_final_artifact_checks")
        self.assert_nonessential_ipos_residual(functional_input)
        partial = build_ipos_functional_input("Module",
            {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "Coordinator shall write samples to Receiver. The value shall refer to the cycle start (for"}],
            domain="digital")
        flow = _ipos_shared_flow(partial)
        self.assertEqual(len(flow.facts), 4)
        self.assertIn("unbalanced source delimiters", flow.audit[3]["reason"])
        self.assertTrue(ipos_source_preservation_findings(flow.records, flow.facts))
        self.assertEqual(ipos_descriptive_structural_audit(partial)["source_semantic_status"], "controlled_constructions_valid")
        self.assert_nonessential_ipos_residual(partial)

    def test_ipos_runtime_adapter_preserves_each_complete_sentence_and_audits_a_trailing_fragment(self):
        statement = "Coordinator shall store samples. Receiver shall process samples."
        functional_input = build_ipos_functional_input(
            "Module", {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": statement}], domain="digital",
        )
        flow = _ipos_shared_flow(functional_input)
        self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])
        self.assertEqual(len(flow.facts), 6)
        self.assertEqual(ipos_descriptive_structural_audit(functional_input)["application"]["profile_status"], "policy_valid")
        changed = replace(flow.facts[-2], subject="Coordinator")
        self.assertTrue(ipos_source_preservation_findings(flow.records, (*flow.facts[:-2], changed, flow.facts[-1])))
        statement += " This"
        sections, audit = compose_shared_ipos_overview(build_ipos_functional_input(
            "Module", {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": statement}], domain="digital",
        ))
        gaps = [row for row in audit if row["decision"] == "omitted"]
        self.assertTrue(gaps)
        self.assertEqual(gaps[0]["candidate_evidence_statement"], statement)
        self.assertTrue(gaps[0]["reason"].startswith("omitted_nonessential:"))
        self.assertIn("Receiver processes samples.", " ".join(sections["internal_functions"]))
        self.assertIn("Coordinator stores samples.", sections["internal_functions"])

    def test_ipos_overview_never_publishes_adverbs_as_verbs_flattened_tables_or_partial_objects(self):
        for statement in (
            "The user shall unexpectedly select the channel.",
            "The data shall have the following tag: CHANNEL BIT 0 VALUE 1.",
            "The phase shall wait for sampling the new.",
            "The only registers that the user shall change are the listed ones.",
            "Coordinator shall store samples. Receiver shall go into the.",
            "Outputs shall be high when the sequencer is.",
            "Output shall go high,",
        ):
            with self.subTest(statement=statement):
                functional = build_ipos_functional_input("Module",
                    {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
                    [{"source_req_id": "one", "statement": statement}], domain="digital")
                sections, audit = compose_shared_ipos_overview(functional)
                self.assertEqual(sections["internal_functions"], ["Coordinator stores samples."]
                                 if statement.startswith("Coordinator shall store samples.")
                                 else ["Output goes high."] if statement == "Output shall go high," else [])
                self.assertEqual(sections["internal_structure"], [])
                omissions = [row for row in audit if row["decision"] == "omitted"]
                self.assertEqual(len(omissions), 1)
                self.assertTrue(omissions[0]["reason"].startswith("omitted_nonessential:"))
                self.assertEqual(omissions[0]["candidate_evidence_statement"], statement)
                self.assert_nonessential_ipos_residual(functional)

    def test_ipos_scope_expands_inventory_capabilities_without_changing_purpose(self):
        functional = build_ipos_functional_input("Module",
            {"Function": "Acquire samples, process data, and write results", "Inputs": "samples", "Outputs": "results"},
            [], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        self.assertEqual(len(sections["scope"]), 4)
        for responsibility in ("acquire samples", "process data", "write results"):
            self.assertEqual(sum(responsibility in text for text in sections["scope"]), 1)
        self.assertIn("acquire samples, process data, and write results", sections["functionality"][0])
        boundaries = [text for text in sections["scope"] if "accepts samples" in text and "produces results" in text]
        self.assertEqual(len(boundaries), 1)
        self.assertNotIn("designed to", boundaries[0])
        plan = ipos_descriptive_structural_audit(functional)["section_plan"]
        self.assertFalse(any(value["decision"] == "excluded" for value in plan))
        self.assertFalse(any(row["source_record_id"] == "inventory:main_functions" for row in audit))
        flow = _ipos_shared_flow(functional)
        self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])
        self.assertTrue(ipos_source_preservation_findings(flow.records, flow.facts[:-1]))
        self.assertTrue(ipos_source_preservation_findings(flow.records, (*flow.facts, flow.facts[-1])))
        changed = replace(flow.facts[-1], objects=("to acquire samples and store results",))
        self.assertTrue(ipos_source_preservation_findings(flow.records, (*flow.facts[:-1], changed)))

    def test_ipos_main_function_list_keeps_shared_objects_and_parenthetical_clauses_intact(self):
        functional = build_ipos_functional_input("Module",
            {"Function": "Generate and receive commands, process values (read inputs, write outputs), and store results",
             "Inputs": "samples", "Outputs": "results"}, [], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        self.assertEqual(len(sections["scope"]), 4)
        for responsibility in ("generate and receive commands", "process values (read inputs, write outputs)", "store results"):
            self.assertEqual(sum(responsibility in text for text in sections["scope"]), 1)
        self.assertIn("generate and receive commands", sections["functionality"][0])
        self.assertIn("process values (read inputs, write outputs)", sections["functionality"][0])
        self.assertIn("store results", sections["functionality"][0])
        flow = _ipos_shared_flow(functional)
        self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])
        for changed in (replace(flow.facts[-3], objects=("designed to generate or receive commands",)),
                        replace(flow.facts[-3], mode="service mode")):
            self.assertTrue(ipos_source_preservation_findings(flow.records, (*flow.facts[:-3], changed, *flow.facts[-2:])))

    def test_ipos_passive_control_with_explicit_functional_effect_is_retained(self):
        statement = ("When in INITIAL state, when averaging completes, the ramp trigger shall be set high "
                     "to start the rising of the digital ramp and go into the RISING state.")
        functional = build_ipos_functional_input("Module",
            {"Function": "Control acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "ramp", "statement": statement}], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        self.assertIn("to start the rising of the digital ramp and go into the RISING state",
                      " ".join(sections["internal_functions"]))
        self.assertIn("When in INITIAL state, when averaging completes",
                      " ".join(sections["internal_functions"]))
        self.assertFalse(any(row["decision"] == "projection_gap" for row in audit))
        flow = _ipos_shared_flow(functional)
        self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])
        claim = json.loads(next(row["claim_json"] for row in audit if row["source_record_id"].startswith("requirement:")))
        self.assertEqual(claim["scope"], "initiation")
        self.assertEqual(claim["meanings"][0]["effects"][0]["action"], "start")
        self.assertNotIn("**", " ".join(sections["internal_functions"]))

    def test_ipos_named_function_groups_use_literal_entities_and_effects_not_project_catalogues(self):
        rows = [{"source_req_id": identity, "statement": statement} for identity, statement in (
            ("start", "Operator shall start measurement writing CONTROL at 1."),
            ("sequence", "When ready, Acquisition FSM shall go into ACTIVE state."),
            ("wait", "Unless ready, acquisition FSM shall wait for conversion completion."),
            ("control", "When ready, trigger shall be set high to start the falling of the oscillator ramp."))]
        functional = build_ipos_functional_input("Unrelated module",
            {"Function": "Control acquisition", "Inputs": "samples", "Outputs": "results"}, rows, domain="analog")
        sections, audit = compose_shared_ipos_overview(functional)
        output = " ".join((*sections["scope"], *sections["internal_functions"]))
        self.assertNotIn("**", output)
        claim = json.loads(next(row["claim_json"] for row in audit if row["candidate_evidence_requirement_id"] == "start"))
        self.assertEqual(claim["scope"], "initiation")
        self.assertEqual(claim["meanings"][0]["subject"], "Operator")
        self.assertIn("Unless ready", output)
        self.assertIn("waits for conversion completion", output)
        self.assertFalse(any(row["decision"] == "projection_gap" for row in audit))

    def test_ipos_flag_lifecycle_preserves_nested_guard_threshold_and_clearing(self):
        statement = ("In case the optical sample data are equal at 0x8000 or 0x7FFF for at least half "
                     "of the selected average (example, if 128 acquisitions are selected, the threshold "
                     "shall be greater or equal to 64), o_saturation_flag shall be set to 1 "
                     "and shall go to 0 at the start of the next frame or frame repetition.")
        functional = build_ipos_functional_input("Module",
            {"Function": "Process data", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "flag", "statement": statement}], domain="digital")
        sections, audit = compose_shared_ipos_overview(functional)
        from workflow_routing import build_ipos_claim_plan, realize_descriptive_functional_claim
        claims = [claim for claim in build_ipos_claim_plan(functional).claims
                  if any(span.record_id.endswith(":flag") for span in claim.spans)]
        output = " ".join(realize_descriptive_functional_claim(claim) for claim in claims)
        demoted = [row for row in audit if row["candidate_evidence_requirement_id"] == "flag"]
        self.assertTrue(demoted)
        self.assertTrue(all(row["decision"] == "demoted" and not row["rendered_summary"] for row in demoted))
        self.assertNotIn("o_saturation_flag", " ".join(sections["internal_functions"]))
        self.assertIn("o_saturation_flag", output)
        self.assertNotIn("o saturation flag", output)
        self.assertIn("threshold is greater or equal to 64", output)
        self.assertIn("set to 1 and goes to 0 at the start of the next frame or frame repetition", output)
        self.assertNotIn("shall", output)
        self.assertFalse(any(row["decision"] == "projection_gap" for row in audit))
        flow = _ipos_shared_flow(functional)
        self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])
        fact = flow.facts[-2]
        for changed in (replace(fact, condition=fact.condition.replace("64", "63")),
                        replace(fact, objects=("set to 1",))):
            self.assertTrue(ipos_source_preservation_findings(flow.records, (*flow.facts[:-2], changed, flow.facts[-1])))

    def test_ipos_overview_does_not_turn_actors_or_signals_into_titles(self):
        for domain in ("digital", "analog"):
            with self.subTest(domain=domain):
                functional = build_ipos_functional_input("Module",
                    {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
                    [{"source_req_id": "operator", "statement": "Coordinator shall process samples."},
                     {"source_req_id": "signal", "statement": "When enabled, Output shall be set high."}],
                    domain=domain)
                sections, audit = compose_shared_ipos_overview(functional)
                rendered = " ".join(text for paragraphs in sections.values() for text in paragraphs)
                self.assertNotIn("**Coordinator.", rendered)
                self.assertNotIn("**Output.", rendered)
                self.assertIn("Coordinator processes samples.", rendered)
                self.assertIn("output is set high", rendered)
                self.assertNotIn("output is set high", " ".join(sections["scope"]))
                self.assertFalse(any(row["decision"] == "projection_gap" for row in audit))
                selected = [json.loads(row["semantic_unit_json"])["text"] for row in audit
                            if row["decision"] == "selected"]
                self.assertIn("When enabled, output is set high.", selected)
                flow = _ipos_shared_flow(functional)
                self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])

    def test_ipos_independent_checker_rejects_malformed_predicate_even_when_fact_copies_source(self):
        statement = "Operator shall always select the channel."
        record = NormalizedSourceRecord("one", statement, "approved_normative", "block_local",
                                        ("approved:one",), "retained", "approved")
        fact = DescriptiveFact("Operator", "always", ("select the channel",), "approved:one", statement,
                               "block_local", owner="Module", domain="digital", layer="block",
                               evidence_kind="approved_normative")
        self.assertTrue(ipos_source_preservation_findings((record,), (fact,)))
        self.assertIn("descriptive_predicate_grammar_unsupported",
                      descriptive_fact_findings(fact, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Module"))
        functional = build_ipos_functional_input("Module",
            {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "Coordinator shall have an input buffer."}], domain="digital")
        flow = _ipos_shared_flow(functional)
        self.assertIn("Coordinator has an input buffer.", [unit.text for unit in flow.units])
        self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])

    def test_ipos_shared_independent_source_check_rejects_joint_fact_mutations(self):
        statement = "When ready, Coordinator shall not transfer samples to Receiver or Sender."
        record = NormalizedSourceRecord("requirement:Module:one", statement, "approved_normative",
                                        "block_local", ("approved:one",), "retained", "approved")
        fact = DescriptiveFact("Coordinator", "transfer", ("samples to Receiver or Sender",),
                               "approved:one", statement, "block_local", condition="When ready", negative=True)
        self.assertEqual(ipos_source_preservation_findings([record], [fact]), [])
        for mutation in (replace(fact, condition=""), replace(fact, negative=False),
                         replace(fact, objects=("samples to Receiver and Sender",)),
                         replace(fact, subject="Receiver", objects=("samples to Coordinator or Sender",)),
                         replace(fact, objects=("samples to Receiver",))):
            self.assertTrue(ipos_source_preservation_findings([record], [mutation]))
        self.assertTrue(ipos_source_preservation_findings([record], []))

    def test_ipos_shared_independent_source_check_accepts_explicit_equivalence(self):
        record = NormalizedSourceRecord("requirement:Module:one", "Coordinator shall comprise Receiver and Sender.",
                                        "approved_normative", "block_local", ("approved:one",), "retained", "approved")
        fact = DescriptiveFact("Coordinator", "contain", ("Receiver", "Sender"), "approved:one",
                               record.text, "block_local")
        self.assertEqual(ipos_source_preservation_findings([record], [fact]), [])

    def test_ipos_shared_claim_rejects_unsupported_child_and_accepts_equivalence(self):
        fact = DescriptiveFact("Coordinator", "comprise", ("Receiver", "Sender"), "approved:one",
                               "Coordinator shall comprise Receiver and Sender.", "block_local")
        paragraph = DescriptiveParagraph("Coordinator contains Receiver and Sender.", (fact,))
        self.assertEqual(_ipos_claim_preservation_findings(paragraph), [])
        self.assertTrue(_ipos_claim_preservation_findings(replace(
            paragraph, text="Coordinator contains Receiver, Sender and Processor.",
        )))

    def test_ipos_shared_guard_with_actor_and_passive_prohibition(self):
        sections, audit = compose_shared_ipos_overview(build_ipos_functional_input(
            "Module", {"Function": "Control acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "When Coordinator is ready, Coordinator shall transfer samples."},
             {"source_req_id": "two", "statement": "Unless ready, Output shall not be set high."}], domain="digital",
        ))
        output = " ".join(sections["internal_functions"])
        self.assertIn("When Coordinator is ready, coordinator transfers samples", output)
        self.assertIn("output is not set high", output)
        self.assertIn("Unless ready, output is not set high.",
                  [json.loads(row["semantic_unit_json"])["text"] for row in audit if row["decision"] == "selected"])

    def test_ipos_shared_guard_negation_does_not_invert_action_and_decimal_is_not_truncated(self):
        sections, _audit = compose_shared_ipos_overview(build_ipos_functional_input(
            "Module", {"Function": "Control acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "When Coordinator is not ready, Coordinator shall transfer 2.5 samples."}], domain="digital",
        ))
        self.assertIn("When Coordinator is not ready, coordinator transfers 2.5 samples", " ".join(sections["internal_functions"]))

    def test_ipos_shared_context_and_subject_list_preserve_roles(self):
        rows = [{"source_req_id": "one", "statement": "During initialization, Coordinator shall copy samples to Receiver. [Vpriority High]"},
                {"source_req_id": "two", "statement": "Receiver, Sender and Processor shall be reset by Coordinator."}]
        record = build_ipos_functional_input(
            "Module", {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"}, rows,
        )
        flow = _ipos_shared_flow(record)
        self.assertEqual(ipos_source_preservation_findings(flow.records, flow.facts), [])
        self.assertEqual(flow.facts[3].subject, "Coordinator")
        self.assertEqual(flow.facts[4].subject, "Receiver, Sender and Processor")
        self.assertNotIn("Vpriority", " ".join(unit.text for unit in flow.units))

    def test_ipos_shared_intransitive_behavior_preserves_negation_and_audits_incomplete_source(self):
        rows = [{"source_req_id": "one", "statement": "Unless enabled, Receiver shall not toggle."},
                {"source_req_id": "two", "statement": "Coordinator shall transfer samples into the."}]
        functional_input = build_ipos_functional_input(
            "Module", {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"}, rows, domain="digital",
        )
        sections, audit = compose_shared_ipos_overview(functional_input)
        self.assertEqual({row["candidate_evidence_requirement_id"] for row in audit
                          if row["decision"] == "omitted"}, {"two"})
        self.assert_nonessential_ipos_residual(functional_input)
        self.assertIn("receiver does not toggle", " ".join(sections["internal_functions"]))
        flow = _ipos_shared_flow(functional_input)
        self.assertIn("Unless enabled, receiver does not toggle.", [unit.text for unit in flow.units])
        findings = ipos_source_preservation_findings(flow.records, flow.facts)
        self.assertEqual(findings, ["ipos_source_projection_unresolved:requirement:Module:two"])

    def test_ipos_shared_threshold_operator_and_symbol_mutations_are_detected(self):
        unresolved = build_ipos_functional_input("Module",
            {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "When ready, it shall transfer samples."}], domain="digital")
        sections, audit = compose_shared_ipos_overview(unresolved)
        self.assertEqual(sections["internal_functions"], [])
        self.assertTrue(any("unresolved subject reference" in row["reason"] for row in audit))
        missing_domain = build_ipos_functional_input("Module",
            {"Function": "Coordinate acquisition", "Inputs": "samples", "Outputs": "results"},
            [{"source_req_id": "one", "statement": "Coordinator shall wait."}])
        with self.assertRaisesRegex(ValueError, "explicit approved digital/analog domain"):
            compose_shared_ipos_overview(missing_domain)
        source = "If level <= 16, Coordinator shall calculate \u03b4."
        record = NormalizedSourceRecord("one", source, "approved_normative", "block_local",
                                        ("approved:one",), "retained", "approved")
        fact = DescriptiveFact("Coordinator", "calculate", ("\u03b4",), "approved:one", source,
                               "block_local", condition="If level <= 16")
        self.assertEqual(ipos_source_preservation_findings([record], [fact]), [])
        self.assertTrue(ipos_source_preservation_findings([record], [replace(fact, condition="If level > 16")]))
        self.assertTrue(ipos_source_preservation_findings([record], [replace(fact, objects=("\u03b8",))]))

    def test_fact_prose_aggregates_compatible_objects_with_all_contributors(self):
        profile = self.profile("synthetic", ("system",))
        facts = [DescriptiveFact("The system", "monitor", (value,), f"source {index}",
                                 f"The system monitors {value}.", "system")
                 for index, value in enumerate(("temperature", "humidity", "pressure"))]
        paragraphs = compose_descriptive_paragraphs(facts, profile=profile)
        self.assertEqual(paragraphs[0].text, "The system monitors temperature, humidity and pressure.")
        self.assertEqual(paragraphs[0].facts, tuple(facts))
        self.assertEqual(validate_descriptive_paragraphs(paragraphs, facts=facts, profile=profile), [])
        self.assertEqual(paragraphs, compose_descriptive_paragraphs(facts, profile=profile))

    def test_fact_prose_flow_admission_projection_audit_and_determinism(self):
        profile = self.profile("synthetic", ("system",))
        records = [NormalizedSourceRecord("record-1", "The system monitors temperature.", "admitted_context",
                                          "system", ("source-1",), "selected", "eligible"),
                   NormalizedSourceRecord("record-2", "Unparsed construction.", "admitted_context",
                                          "system", ("source-2",), "retained", "eligible"),
                   NormalizedSourceRecord("record-3", "Excluded text.", "unapproved",
                                          "system", ("source-3",), "selected", "not admissible")]
        called = []
        def adapter(record):
            called.append(record.record_id)
            return ([DescriptiveFact("The system", "monitor", ("temperature",), record.provenance[0],
                                     record.text, record.ownership_scope)], "") if record.record_id == "record-1" else ([], "unsupported syntax")
        result = run_descriptive_flow(records, profile=profile, adapter_name="synthetic", adapter=adapter)
        self.assertEqual(called, ["record-1", "record-2"])
        self.assertEqual(result.records, tuple(records))
        self.assertEqual(result.units[0].source_record_ids, ("record-1",))
        self.assertEqual(result.paragraphs[0].text, "The system monitors temperature.")
        self.assertEqual({row["decision"] for row in result.audit}, {"admitted", "selected", "projection_gap", "rejected_evidence"})
        self.assertEqual(result, run_descriptive_flow(records, profile=profile, adapter_name="synthetic", adapter=adapter))
        changed = tuple({**row, "contributor_ids": "[]"} if row["decision"] == "selected" else row for row in result.audit)
        self.assertTrue(validate_descriptive_flow(replace(result, audit=changed), profile=profile))
        self.assertTrue(validate_descriptive_flow(replace(result, audit=()), profile=profile))
        self.assertTrue(validate_descriptive_flow(replace(result, units=()), profile=profile))
        with self.assertRaises(ValueError):
            run_descriptive_flow(records, profile=profile, adapter_name="undeclared", adapter=adapter)
        with self.assertRaises(ValueError):
            run_descriptive_flow([records[0], records[0]], profile=profile, adapter_name="synthetic", adapter=adapter)

    def test_fact_prose_flow_rejects_forged_adapter_provenance(self):
        record = NormalizedSourceRecord("record", "Original description.", "admitted_context", "system",
                                        ("original source",), "selected", "eligible")
        forged = DescriptiveFact("The system", "monitor", ("temperature",), "different source", record.text, "system")
        with self.assertRaises(ValueError):
            run_descriptive_flow([record], profile=self.profile("synthetic", ("system",)), adapter_name="synthetic",
                                 adapter=lambda source: ([forged], ""))

    def test_fact_prose_profiles_keep_future_runtime_inactive(self):
        self.assertEqual(descriptive_profile("SRS"), SRS_DESCRIPTIVE_PROFILE)
        self.assertEqual(set(DESCRIPTIVE_PROFILES), {"SRS", "DRS", "ARS", "IPOS"})
        self.assertEqual(descriptive_profile("IPOS"), IPOS_DESCRIPTIVE_PROFILE)
        for profile in (DRS_DESCRIPTIVE_PROFILE, ARS_DESCRIPTIVE_PROFILE):
            with self.subTest(profile=profile.document_type):
                with self.assertRaises(ValueError):
                    descriptive_profile(profile.document_type)
                for candidate in (profile, replace(profile, runtime_enabled=True)):
                    with self.assertRaises(ValueError):
                        run_descriptive_flow([], profile=candidate, adapter_name=profile.projection_adapters[0],
                                             adapter=lambda record: ([], ""))
        with self.assertRaises(ValueError):
            descriptive_profile("UNKNOWN")
        with self.assertRaises(TypeError):
            DESCRIPTIVE_PROFILES["DRS"] = SRS_DESCRIPTIVE_PROFILE

    def test_fact_prose_profile_detail_and_owner_rules_do_not_leak(self):
        fact = DescriptiveFact("The controller", "provide", ("register access",), "source", "Original requirement.",
                               "block_local", section="function", evidence_kind="approved_normative",
                               owner="Controller A", layer="block")
        self.assertEqual(descriptive_fact_findings(fact, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Controller A"), [])
        self.assertIn("descriptive_owner_invalid", descriptive_fact_findings(fact, profile=IPOS_DESCRIPTIVE_PROFILE, owner="Controller B"))
        self.assertTrue(descriptive_fact_findings(fact, profile=SRS_DESCRIPTIVE_PROFILE))
        srs_fact = replace(fact, scope="system", section=SRS_DESCRIPTIVE_PROFILE.topic_roles[0],
                   evidence_kind="descriptive", objects=("measurements",), condition="When register access is enabled")
        self.assertIn("descriptive_profile_detail_excluded", descriptive_fact_findings(srs_fact, profile=SRS_DESCRIPTIVE_PROFILE))
        synthetic = replace(fact, evidence_kind="descriptive", section="")
        self.assertIn("register access", compose_descriptive_paragraphs([synthetic], profile=self.profile("local", ("block_local",)))[0].text)
        for profile, domain in ((DRS_DESCRIPTIVE_PROFILE, "digital"), (ARS_DESCRIPTIVE_PROFILE, "analog")):
            integration = replace(fact, scope="architecture", section="integration", evidence_kind="descriptive",
                                  domain=domain, layer="integration")
            self.assertEqual(descriptive_fact_findings(integration, profile=profile), [])
            self.assertIn("descriptive_domain_invalid", descriptive_fact_findings(replace(integration, domain="wrong"), profile=profile))
            self.assertIn("descriptive_layer_invalid", descriptive_fact_findings(replace(integration, layer="block"), profile=profile))
            self.assertIn("descriptive_scope_or_kind_invalid", descriptive_fact_findings(replace(integration, evidence_kind="approved_normative"), profile=profile))
        with self.assertRaises(ValueError):
            replace(IPOS_DESCRIPTIVE_PROFILE, owner_policy="guess")
        policy = replace(self.profile("local", ("block_local",)), evidence_kinds=("approved_normative",),
                         authority_tiers=("admitted_context", "approved_normative"))
        record = NormalizedSourceRecord("record", fact.statement, "approved_normative", "block_local",
                                        (fact.source,), "selected", "already admitted")
        result = run_descriptive_flow([record], profile=policy, adapter_name="synthetic", adapter=lambda source: ([fact], ""))
        self.assertEqual(validate_descriptive_flow(result, profile=policy), [])
        self.assertTrue(validate_descriptive_flow(replace(result, records=(replace(record, authority_tier="admitted_context"),)), profile=policy))

    def test_fact_prose_preserves_conditions_modalities_and_alternatives(self):
        profile = self.profile("synthetic", ("block_local",))
        facts = [DescriptiveFact("The controller", "transfer", ("messages", "commands"),
                                 "source 1", "Original relationship.", "block_local",
                                 condition="When enabled", modality="may", coordination="or"),
                 DescriptiveFact("The controller", "transfer", ("messages",),
                                 "source 2", "Original prohibition.", "block_local", negative=True)]
        paragraphs = compose_descriptive_paragraphs(facts, profile=profile)
        self.assertEqual(len(paragraphs), 2)
        self.assertEqual(paragraphs[0].text, "When enabled, the controller may transfer messages or commands.")
        self.assertEqual(paragraphs[1].text, "The controller does not transfer messages.")
        self.assertEqual(validate_descriptive_paragraphs(paragraphs, facts=facts, profile=profile), [])

    def test_fact_prose_does_not_merge_different_modes_or_conditions(self):
        profile = self.profile("synthetic", ("system",))
        facts = [DescriptiveFact("The system", "store", ("measurements",),
                                 f"source {index}", "Original relationship.", "system",
                                 mode=mode, condition=condition)
                 for index, (mode, condition) in enumerate(
                     (("Mode A", "When enabled"), ("Mode B", "When enabled"), ("Mode B", "When requested")))]
        self.assertEqual(len(compose_descriptive_paragraphs(facts, profile=profile)), 3)

    def test_fact_prose_validator_detects_loss_without_recomposing(self):
        profile = self.profile("synthetic", ("system",))
        fact = DescriptiveFact("The system", "monitor", ("temperature",),
                               "source", "Original relationship.", "system", modality="may",
                               qualifiers=("only during maintenance",))
        paragraph = DescriptiveParagraph("The system monitors humidity.", (fact,))
        findings = validate_descriptive_paragraphs([paragraph], facts=[fact], profile=profile)
        self.assertTrue(any("content_lost" in finding for finding in findings))
        self.assertTrue(any("modality_lost" in finding for finding in findings))
        self.assertTrue(validate_descriptive_paragraphs([], facts=[fact], profile=profile))
        normative_link = DescriptiveParagraph("The system may monitor temperature only during maintenance. Covers: SOURCE-ID", (fact,))
        self.assertIn("descriptive_normative_wording", validate_descriptive_paragraphs([normative_link], facts=[fact], profile=profile))
        with self.assertRaises(ValueError):
            compose_descriptive_paragraphs([fact], profile=self.profile("local", ("block_local",)))

    def test_drs_top_level_overview_breaks_long_description_into_short_paragraphs(self):
        overview = _drs_top_level_overview(
            [("Sensor", "FIFO", "sample data"), ("Host", "Regmap", "SPI host transactions")],
            ["The device integrates ECG and PPG.", "The source describes an embedded engine."],
        )
        paragraphs = overview.split("\n\n")
        self.assertEqual(len(paragraphs), 3)
        self.assertTrue(all(len(re.findall(r"(?<=[.!?])(?:\s|$)", paragraph)) <= 2 for paragraph in paragraphs))
        self.assertIn("The device integrates ECG and PPG.", overview)
        self.assertIn("The source material does not define a separate application role.", overview)

    def test_drs_section_quality_rejects_process_and_counts_in_both_formats(self):
        for prose in ("Approved integration evidence groups the exchanges.",
                      "Approved block functions identify Storage.",
                      "Approved clock routes reach destinations.",
                      "The fabric connects 4 source endpoints across 8 paths.",
                      "The architecture also records a status exchange."):
            with self.subTest(prose=prose):
                self.assertTrue(drs_section_prose_findings(prose))
                markdown = "## 3. Top Level Overview\n" + prose
                document = ET.fromstring(
                    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:body><w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr>'
                    '<w:r><w:t>3. Top Level Overview</w:t></w:r></w:p>'
                    '<w:p><w:r><w:t>' + prose + '</w:t></w:r></w:p></w:body></w:document>'
                )
                findings = _drs_rendered_prose_findings(markdown, document)
                self.assertTrue(any(":markdown:3:" in finding for finding in findings))
                self.assertTrue(any(":docx:3:" in finding for finding in findings))
        self.assertEqual(drs_section_prose_findings("The fabric distributes 16 MHz clocks without clock gating."), [])
        self.assertEqual(_drs_rendered_prose_findings("## 6. Requirement catalog\nThe unit shall operate.\n"), [])

    def test_drs_section_grouping_preserves_edges_and_runtime_names(self):
        rows = [("Producer", "Consumer", "samples"), ("Producer", "Monitor", "samples"),
                ("Peer", "Consumer", "samples"), ("Peer", "Monitor", "samples"),
                ("Host", "Consumer", "register writes")]
        original = list(rows)
        summary = compose_drs_interaction_summary(rows + rows)
        self.assertEqual(summary, compose_drs_interaction_summary(rows))
        self.assertIn("Producer and Peer provide samples to Consumer and Monitor.", summary)
        self.assertIn("Host provides register writes to Consumer.", summary)
        self.assertNotIn("writes to Consumer and Monitor", summary)
        self.assertEqual(rows, original)
        renamed = [(source.replace("Producer", "Capture"), target, exchange) for source, target, exchange in rows]
        self.assertEqual(compose_drs_interaction_summary(renamed), summary.replace("Producer", "Capture"))
        role = compose_drs_role_summary([("Capture", "Perform sample filtering, and write results to memory.")], ("filtering", "results"))
        self.assertIn("Capture performs sample filtering.", role)
        self.assertIn("Capture writes results to memory.", role)
        self.assertEqual(drs_section_prose_findings(role), [])
        coordinated = compose_drs_role_summary([
            ("Control", "Generate and receive control signals, perform filtering, and write results."),
            ("Bridge", "Accept requests and translate them into bus transactions."),
        ], ())
        self.assertIn("Control generates and receives control signals.", coordinated)
        self.assertIn("Control performs filtering.", coordinated)
        self.assertIn("Bridge accepts requests and translates them into bus transactions.", coordinated)
        self.assertNotIn("signals, perform", coordinated)
        missing_role = _drs_processing_summary([], [("Compute", "Storage", "bus connection")])
        self.assertTrue(missing_role.startswith("Compute connects to Storage through bus."))
        self.assertTrue(missing_role.endswith("Processing behavior: need clarification."))

    def test_drs_section_synthesis_groups_approved_exchanges_without_promoting_endpoints(self):
        rows = [("Source A", "Sink A", "digital samples"),
                ("Source B", "Sink B", "FIFO watermark flags"),
                ("Host", "Registers", "SPI register transactions"),
                ("Unit", "Endpoint", "XBAR connection")]
        summary = _drs_section_interaction_summary(rows)
        self.assertIn("Source A provides digital samples to Sink A", summary)
        self.assertIn("Host provides SPI register transactions to Registers", summary)
        self.assertIn("Unit connects to Endpoint through XBAR", summary)
        self.assertNotIn("Source A -> Sink A", summary)
        self.assertNotIn("approved", summary.lower())
        self.assertNotRegex(summary, r"\d+ (?:paths|source endpoints|destinations|interconnect links)")
        self.assertNotIn("Endpoint block", summary)
        self.assertIn("need clarification", _drs_section_interaction_summary([]))
        clock_summary = _drs_clock_section_summary([
            ("REQ-1", "u_power.clock", "u_core.clk", "16 MHz", "Not gated"),
            ("REQ-2", "need clarification", "need clarification", "32 kHz", "Gated"),
        ])
        self.assertIn("16 MHz", clock_summary)
        self.assertIn("32 kHz", clock_summary)
        self.assertIn("without clock gating", clock_summary)
        self.assertIn("need clarification", clock_summary)
        self.assertNotIn("approved", clock_summary.lower())
        self.assertNotIn("route(s)", clock_summary)

    def test_drs_shared_facets_and_power_routes_preserve_evidence_boundary(self):
        storage = BlockInfo("Store", "Provide FIFO storage and AHB/memory access in mutually exclusive modes.", "", "", "", "concrete_block")
        power = BlockInfo("Power", "Sequence boot clocks and reset.", "", "", "", "concrete_block")
        facets = dict(_drs_shared_facet_rows([storage, power], [("Host", "Store", "XBAR connection")]))
        self.assertIn("Store", facets["Shared-resource ownership"])
        self.assertNotIn("Provide FIFO storage", facets["Shared-resource ownership"])
        self.assertIn("Concurrent-access ordering policy: need clarification", facets["Register and memory access coordination"])
        for unsupported in ("Flow control and backpressure", "Synchronization", "Latency and bandwidth"):
            self.assertEqual(facets[unsupported], "need clarification")
        paths = [("REQ-1", "top.u_power.clock", "top.core.clock", "16 MHz", "Not gated"),
                 ("REQ-2", "need clarification", "need clarification", "16 MHz", "Gated")]
        self.assertEqual(_drs_power_clock_routes(paths, [power]), paths[:1])
        ports = [{"Owner": "Power", "Port name": "resetn_core", "Direction": "output",
                  "Type / details": "reset core", "Source page": "41", "Ownership status": "approved"},
                 {"Owner": "Other", "Port name": "resetn_io", "Direction": "output",
                  "Type / details": "reset io", "Source page": "42", "Ownership status": "approved"}]
        self.assertEqual(_drs_power_reset_outputs(ports, [power]), [("Power", "resetn_core", "reset core", "41")])

    def test_drs_clock_paths_require_approved_explicit_routes(self):
        route = {
            "source_req_id": "REQ-1", "approved_classification": "Digital", "lifecycle_state": "approved",
            "requirement_statement": "The top.pmu.clock signal shall be connected to the top.core.clock signal, with frequency 16 MHz and with clock gating top.pmu.enable.",
        }
        self.assertEqual(_drs_clock_path_rows([route]), [
            ("REQ-1", "top.pmu.clock", "top.core.clock", "16 MHz", "Gated; control expression: top.pmu.enable")
        ])
        split_control = {**route, "requirement_statement": route["requirement_statement"].replace(
            "top.pmu.enable", "top.u_pm u.enable | ~top.u_p ad_mux.scan_enable")}
        self.assertEqual(_drs_clock_path_rows([split_control])[0][4],
                         "Gated; control expression: top.u_pm u.enable | ~top.u_p ad_mux.scan_enable; signal spelling: need clarification")
        self.assertIn(r"u.enable \| ~top", _drs_clock_path_markdown_row(_drs_clock_path_rows([split_control])[0]))
        without_control = {**route, "requirement_statement": route["requirement_statement"].replace(
            " top.pmu.enable", "")}
        self.assertEqual(_drs_clock_path_rows([without_control])[0][4], "Gated; control expression: need clarification")
        ungated = {**route, "requirement_statement": route["requirement_statement"].replace(
            "with clock gating top.pmu.enable", "without clock gating")}
        self.assertEqual(_drs_clock_path_rows([ungated])[0][4], "Not gated")
        malformed = {**route, "requirement_statement": route["requirement_statement"].replace("top.core.clock", "top.u.core.clock")}
        self.assertEqual(_drs_clock_path_rows([malformed])[0][1:3], ("need clarification", "need clarification"))
        self.assertEqual(_drs_clock_path_rows([{**route, "lifecycle_state": "staged"},
                                               {**route, "requirement_statement": "A clock might reach core."}]), [])
        sync = {**route, "requirement_statement": "The u_p\nmu.resetn_32k_main_ctrl signal shall be connected to the top.i_rstn_sync\n_32 signal."}
        ports = [
            {"Owner": "Power", "Port name": "resetn_32k_main_ctrl", "Direction": "output", "Ownership status": "approved"},
            {"Owner": "Controller", "Port name": "i_rstn_sync_32", "Direction": "input", "Ownership status": "approved"},
        ]
        self.assertEqual(_drs_sync_connection_rows([sync], ports),
                         [("REQ-1", "Power.resetn_32k_main_ctrl", "Controller.i_rstn_sync_32")])
        truncated = {**sync, "requirement_statement": "The u_p signal shall be connected to the mu.resetn_32k_main_ctrlu_top.i_rstn_sync_32 signal."}
        self.assertEqual(_drs_sync_connection_rows([truncated], ports),
                 [("REQ-1", "Power.resetn_32k_main_ctrl", "Controller.i_rstn_sync_32")])
        self.assertEqual(_drs_sync_connection_rows([sync], ports[:1]),
                         [("REQ-1", "Power.resetn_32k_main_ctrl", "need clarification")])
        self.assertEqual(_drs_sync_connection_rows([{**sync, "lifecycle_state": "staged"}], ports), [])

    def test_drs_function_and_topic_summaries_do_not_promote_trace_only_roles(self):
        trace_only = BlockInfo("Core", "Represent the digital architecture block described by Stage 2 interaction evidence.", "", "", "", "concrete_block")
        power = BlockInfo("Power unit", "Sequence POR, clock-ready signaling, and digital reset for power modes.", "", "", "", "concrete_block")
        storage = BlockInfo("Store", "Provide FIFO storage in mutually exclusive modes.", "", "", "", "concrete_block")
        self.assertIn("need clarification", _drs_function_summary(trace_only, []))
        self.assertIn("An interrupt triggers Core wake-up", _drs_function_summary(trace_only, ["To wake up the Core, trigger an interrupt."]))
        self.assertNotIn("Core", _drs_topic_narrative("Power and clock islands", [trace_only, power, storage]))
        self.assertIn("Power unit", _drs_topic_narrative("Power and clock islands", [trace_only, power, storage]))
        self.assertIn("mutually exclusive", _drs_topic_narrative("Arbitration and control", [trace_only, power, storage]))

    @unittest.skipUnless(os.name == "nt", "Windows display-name API")
    def test_document_author_follows_current_windows_user(self):
        for full_name in ("First User", "Second User"):
            def display_name(_format, buffer, _length):
                buffer.value = full_name
                return 1

            with patch("ctypes.windll.secur32", SimpleNamespace(GetUserNameExW=display_name)):
                self.assertEqual(document_author_name(), full_name)

        with patch("ctypes.windll.secur32", SimpleNamespace(GetUserNameExW=lambda *_args: 0)):
            with self.assertRaisesRegex(RuntimeError, "full display name is unavailable"):
                document_author_name()

    def test_document_author_validator_rejects_mismatched_title(self):
        markdown_path = Path(__file__).resolve().parents[1] / "artifacts/stage5_drs/digital_requirements_specification.md"
        self.assertEqual(validate_document_author_fields(markdown_path), [])
        original_read_text = Path.read_text

        def changed_author(path, *args, **kwargs):
            content = original_read_text(path, *args, **kwargs)
            return content.replace("Author: " + document_author_name(), "Author: Different User") if path == markdown_path else content

        with patch.object(Path, "read_text", changed_author):
            self.assertIn(f"document_docx_author_mismatch:{markdown_path}", validate_document_author_fields(markdown_path))

    def test_drs_domain_projection_and_block_relationships_use_structured_evidence(self):
        records = [{
            "evidence_kind": "power_domain",
            "statement": "PD_TEST: Included block(s): core. Domain type: Retention. Control mode: firmware. • Function: Preserves state. • Characteristics: Idle.",
        }, {
            "evidence_kind": "architecture",
            "statement": "PD_FALSE: Included block(s): core. Domain type: Switchable. Control mode: SW. • Function: Invalid.",
        }]
        self.assertEqual(_drs_power_domain_rows(records), [("PD_TEST", "Retention", "firmware", "Preserves state")])
        self.assertEqual(_drs_domain_function_sentence("PD_TEST", "Preserves state"), "PD_TEST has the documented function: Preserves state.")
        self.assertEqual(_drs_domain_function_sentence("PD_SUPPLY", "power supply to memory"), "PD_SUPPLY supplies power to memory.")
        interactions = [("Block A", "Regmap", "configuration writes"), ("ISPU", "Block A", "sample ready")]
        self.assertEqual(_drs_block_relationships("Block A", interactions),
                         "to Regmap (configuration writes); from ISPU (sample ready)")
        self.assertIn("need clarification", _drs_block_relationships("Other block", interactions))

    def test_drs_central_validator_rejects_missing_power_table(self):
        from docx import Document
        import run_drs_gen_spec_agent as generator
        power_header = ["Domain", "Type", "Control", "Function"]
        power_row = ["PD_TEST", "Retention", "firmware", "Preserves state"]
        clock_header = ["Source requirement", "Source clock", "Destination clock", "Frequency", "Clock gating"]
        clock_rows = [
            ("ROUTE-1", "Power.clk", "Core.clk", "10 MHz", "enable | mode"),
            ("ROUTE-2", "need clarification", "Core.other", "need clarification", "need clarification"),
        ]
        clock_summary = "The clock network distributes the fixture clocks."
        clock_table = "| " + " | ".join(clock_header) + " |\n" + "\n".join(
            _drs_clock_path_markdown_row(row) for row in clock_rows)
        from spec_document_contract import apply_clean_document_presentation, clean_descriptive_review_text
        from workflow_routing import IPOS_CONNECTION_TABLES, render_connection_table_sections
        connection_rows = [tuple(clean_descriptive_review_text(value) for value in row) for row in clock_rows]
        connection_sections = render_connection_table_sections("4.7", {
            "clock": [_drs_clock_path_markdown_row(row) for row in connection_rows]}, first_table_number=8)
        baseline = "\n".join([
            "Snapshot ID: fixture-snapshot", "## 3. Top Level Overview", "Fixture overview.",
            "### 3.1 Digital Main Functions", "#### **Storage** {#storage}", "Storage buffers samples.",
            "Priority between simultaneous resource requesters: need clarification.",
            "## 4. Functional requirements", "### 4.1 Digital block list",
            "| Block | Role | Architectural relationship |", "| Storage | Buffer samples. | none |",
            "### 4.2 Arbitration, control, and shared resources",
            "Flow control and backpressure: need clarification", "### 4.3 Other functions",
            "| Domain | Type | Control | Function |", "| " + " | ".join(power_row) + " |",
            "### 4.6 Digital performance requirements", clock_table,
            "Clock-path endpoints requiring clarification: ROUTE-2.",
            "### 4.7 Clock and synchronization across domains", clock_summary, clock_table,
            "Clock-domain crossing, reset-domain crossing, and measured interrupt response time: need clarification.",
            *connection_sections,
            "## 5. Other requirements",
        ])
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            root = Path(directory)
            output = root / "artifacts/stage5_drs"
            output.mkdir(parents=True)
            markdown_path = output / "digital_requirements_specification.md"
            (output / "descriptive_summary_audit.csv").write_text("section\n", encoding="utf-8")
            (output / "descriptive_low_power_audit.csv").write_text("decision,query_seed,statement\n", encoding="utf-8")
            (output / "drs_traceability_matrix.csv").write_text("requirement_statement\n", encoding="utf-8")
            baseline = "\n".join(apply_clean_document_presentation(baseline.splitlines(), markdown_path, "DRS"))
            publication_audit_path = output / "descriptive_publication_audit.json"
            publication_audit = json.loads(publication_audit_path.read_text(encoding="utf-8"))
            document = Document()
            document.add_paragraph(baseline)
            for values in (power_header, power_row, clock_header, *clock_rows,
                           ["Block", "Role", "Architectural relationship"], ["Storage", "Buffer samples.", "none"]):
                table = document.add_table(rows=1, cols=len(values))
                for cell, value in zip(table.rows[0].cells, values):
                    cell.text = clean_descriptive_review_text(value)
            document.add_heading("4.7.1 Clock connections", level=4)
            document.add_heading("Table 8. Clock connections", level=5)
            table = document.add_table(rows=1, cols=len(IPOS_CONNECTION_TABLES["clock"][1]))
            for cell, value in zip(table.rows[0].cells, IPOS_CONNECTION_TABLES["clock"][1]):
                cell.text = value
            for values in connection_rows:
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = value
            document.save(markdown_path.with_suffix(".docx"))
            block = BlockInfo("Storage", "Buffer samples.", "", "", "", "concrete_block")
            mocks = {
                "_read_block_inventory": [], "_read_interface_rows": [], "_source_interface_tables_by_block": {},
                "_drs_descriptive_blocks": [block], "validate_drs_block_descriptions": [],
                "_top_digital_coverage_records": [],
                "_drs_intro_rows": {"3. Top Level Overview": [], generator.DRS_INTRO_TOPICS[0][0]: []},
                "_drs_rendered_prose_findings": [], "_drs_power_domain_rows": [tuple(power_row)],
                "power_domain_records_from_ocr": [], "_drs_block_relationships": "none",
                "_stage1_overview_descriptive_records": ([], {"Storage": [{"statement": "Buffer samples."}]}),
                "load_sqlite_snapshot_input": SimpleNamespace(rows=()),
                "_drs_clock_section_summary": clock_summary, "_drs_power_clock_routes": clock_rows,
                "_drs_section_interaction_summary": "", "_drs_shared_facet_rows": [("Flow control and backpressure", "need clarification")],
                "_drs_power_reset_outputs": [], "_drs_sync_connection_rows": [],
            }
            for name, value in mocks.items():
                stack.enter_context(patch.object(generator, name, side_effect=lambda *args, value=value, **kwargs:
                                                list(value) if isinstance(value, list) else value))
            stack.enter_context(patch.object(generator, "_drs_clock_path_rows", side_effect=lambda *args, **kwargs:
                                            connection_rows if kwargs.get("preserve_source_fields") else clock_rows))
            stack.enter_context(patch.object(generator, "DRS_NARRATIVE_ROLE_TERMS", ()))
            markdown_path.write_text(baseline, encoding="utf-8")
            self.assertEqual(validate_drs_descriptive_artifacts(root), [])
            mutations = (
                ("| Domain | Type | Control | Function |", "", ["drs_power_domain_table_missing"]),
                ("#### **Storage** {#storage}", "", ["drs_digital_function_missing:Storage"]),
                (clock_summary, "Clock summary removed.", ["drs_clock_section_summary_missing"]),
                ("Flow control and backpressure: need clarification", "", ["drs_shared_facet_missing:Flow control and backpressure"]),
                ("Clock-path endpoints requiring clarification: ROUTE-2.", "", ["drs_power_clock_endpoint_boundary_missing"]),
                (r"enable \| mode", "changed gate", ["drs_power_clock_route_missing:ROUTE-1", "drs_clock_path_missing:ROUTE-1"]),
            )
            for original, replacement, expected in mutations:
                with self.subTest(original=original):
                    markdown_path.write_text(baseline.replace(original, replacement), encoding="utf-8")
                    changed_audit = dict(publication_audit, omissions=[item for item in publication_audit["omissions"]
                                        if original in baseline or original not in item["original"]])
                    publication_audit_path.write_text(json.dumps(changed_audit), encoding="utf-8")
                    findings = validate_drs_descriptive_artifacts(root)
                    for finding in expected:
                        self.assertIn(finding, findings)

    def test_drs_intro_uses_only_structured_approved_interactions(self):
        entries = [
            "Interaction: PMU -> Main Controller (power/clock status; mode transition)",
            "Interaction: PMU -> Main Controller (power/clock status; duplicate note)",
            "Interaction: SPI interface -> Regmap (SPI register transactions; host access)",
            "Interaction: Regmap -> FIFO (XBAR connection; source-derived cell)",
            "An unsupported requirement shall not be copied.",
        ]
        rows = project_architecture_interactions(entries)
        self.assertEqual(rows, [
            ("PMU", "Main Controller", "power/clock status"),
            ("SPI interface", "Regmap", "SPI register transactions"),
            ("Regmap", "FIFO", "XBAR connection"),
        ])
        self.assertIn("PMU provides power/clock status to Main Controller", summarize_architecture_interactions(rows))
        self.assertEqual(summarize_architecture_interactions([("ADSP", "Regmap", "XBAR connection")]),
                 "ADSP connects to Regmap through XBAR connection.")
        self.assertNotIn("shall", summarize_architecture_interactions(rows))
        sections = _drs_intro_rows(entries)
        self.assertIn(rows[0], sections["4.6 Digital performance requirements"])
        self.assertIn(rows[1], sections["4.3 Register and configuration requirements"])
        self.assertIn(rows[1], sections["4.4 Interface and communication requirements"])
        self.assertNotIn(rows[2], sections["4.3 Register and configuration requirements"])
        self.assertNotIn(rows[2], sections["4.5 Data-path and buffering requirements"])
        self.assertIn(rows[2], _drs_intro_rows(entries, {"Regmap"})["7. Interfaces and mixed-signal interactions"])
        self.assertNotIn("IRQ logic", str(sections))
        self.assertEqual(project_architecture_interfaces([
            "Approved interface VDD owned by PMU: Core supply",
            "Approved interface VDD owned by PMU: Core supply",
        ]), [("VDD", "PMU", "Core supply")])

    def test_drs_block_descriptions_use_approved_concrete_inventory_only(self):
        blocks = [
            BlockInfo("Block A", "Route digital samples.", "", "", "", "concrete_block"),
            BlockInfo("Control logic", "Route digital samples.", "", "", "", "logic_context"),
        ]
        selected = _drs_descriptive_blocks(blocks, [], {})
        self.assertEqual([block.name for block in selected], ["Block A"])
        self.assertEqual(
            compose_technical_block_purpose(selected[0].name, selected[0].function),
            "The Block A block is designed to route digital samples.",
        )

    def test_drs_block_validator_requires_provenance_and_tabular_ports(self):
        sentence = "The Block A block is designed to route digital samples."
        markdown = (
            "### **4.1 Digital block list** {#41-digital-block-list}\n"
            "The digital subsystem includes Block A.\n"
            "## 9. Top-level integration requirements\n"
            "### **9.1 Block A** {#91-block-a}\n" + sentence + "\n\n"
            "| Port name | Direction | Type / details | Source page |\n"
            "|---|---|---|---|\n| SDA | input | pad signal | 42 |\n"
        )
        blocks = [{"name": "Block A", "function": "Route digital samples."}]
        audit = [{
            "section": "DRS digital block descriptions", "mapped_block": "Block A",
            "source_evidence": "Route digital samples.", "statement": sentence,
            "source": "artifacts/stage2_mirco_arc/block_inventory.csv",
            "scope_decision": "approved_concrete_block", "output_order": "1",
        }]
        ports = [{"Owner": "Block A", "Ownership status": "approved", "Port name": "SDA",
                  "Direction": "input", "Type / details": "pad signal", "Source page": "42"}]
        self.assertEqual(validate_drs_block_descriptions(markdown, blocks, audit, ports), [])
        self.assertEqual(_render_approved_port_table(ports)[-1], "| SDA | input | pad signal | 42 |")
        variants = (
            (markdown.replace("| SDA | input | pad signal | 42 |", "- Port name: SDA"), "drs_port_name_bullet_list"),
            (markdown + "### **9.2 Control logic** {#92-control-logic}\nUnapproved text.\n", "unapproved_drs_block_paragraph"),
            (markdown.replace(sentence, "General functional description: approved digital function"),
             "description_uses_internal_process_language"),
        )
        for changed, expected in variants:
            with self.subTest(expected=expected):
                self.assertIn(expected, " ".join(validate_drs_block_descriptions(changed, blocks, audit, ports)))
        self.assertIn("drs_block_description_provenance_mismatch", " ".join(
            validate_drs_block_descriptions(markdown, blocks, [audit[0] | {"source_evidence": "other"}], ports)
        ))

    def test_ipos_evidence_summary_restores_names_and_guards_from_accepted_statements(self):
        cases = (
            (
                ("A new operation shall start only when quokka_run is low, and one between OTP_TEST bit, OTP_WRITE bit or OTP_BOOT bit is written to 1",
                 "To perform a BOOT routine, OTP_BOOT bit shall be written when quokka_run is low."),
                "Operation start conditions", ("quokka_run is low", "OTP_BOOT initiates the BOOT routine"),
            ),
            (
                ("During boot, the ADSP shall copy the content of the OTP memory into the OTP registers, following the access procedure.",
                 "During WRITE, the ADSP shall copy the content of the OTP regmap into the OTP memory, following the burn procedure."),
                "OTP memory transfer", ("OTP memory to OTP registers", "OTP regmap to OTP memory"),
            ),
            (
                ("A double byte write access to the address 0x68-0x6f shall start the data preload from AHB in order to have the internal fifo not empty when the read request will be done by the master.",),
                "Read data preloading", ("0x68-0x6f", "AHB", "master read request"),
            ),
            (
                ("At the release of the POR signal from the analog domain, the controller shall manage the turn-on of the LDO1V8 and starts clocks Start-Up Sequence (reported in spec).",),
                "Power-on start-up", ("POR", "LDO1V8", "clock start-up sequence"),
            ),
            (
                ("All the SENSOR_HUB_x registers shall be reset every time a new I2C operation starts.",),
                "Register reset on operation start", ("SENSOR_HUB_x registers", "I2C operation"),
            ),
            (
                ("If input signal HMASTER is set to 0 or 1, the FIFO Controller shall enter Memory mode.",),
                "Memory mode entry", ("HMASTER", "0 or 1"),
            ),
        )
        for statements, expected_title, fragments in cases:
            with self.subTest(title=expected_title):
                title, description = _ipos_evidence_summary([
                    {"candidate_evidence_statement": statement} for statement in statements
                ])
                self.assertEqual(title, expected_title)
                for fragment in fragments:
                    self.assertIn(fragment, description)

    def test_snapshot_candidates_follow_approved_ipos_allocation_not_block_label(self):
        snapshot = SimpleNamespace(source_path=Path("approved.csv"), rows=[
            {"source_req_id": "LOCAL", "requirement_statement": "Local approved behavior."},
            {"source_req_id": "ARCH", "requirement_statement": "Approved architecture behavior."},
            {"source_req_id": "OTHER", "requirement_statement": "Another block's behavior."},
        ])

        class Contract:
            def allocation(self, source_id):
                return {"approved_block": "Block A" if source_id != "OTHER" else "Block B",
                        "owning_target": "DRS" if source_id == "ARCH" else "Digital IPOS"}

        candidates = approved_snapshot_candidate_rows(snapshot, Contract())
        self.assertEqual([row["source_req_id"] for row in candidates if row["mapped_block"] == "Block A"],
                         ["LOCAL"])
        self.assertEqual([row["source_req_id"] for row in candidates], ["LOCAL", "OTHER"])
        self.assertEqual(approved_snapshot_candidate_rows(snapshot, Contract(), domain="analog"), [])

    def test_ipos_shared_source_conflicts_are_unresolved_not_silently_suppressed(self):
        source = "Coordinator shall transfer samples."
        first = DescriptiveFact("Coordinator", "transfer", ("samples",), "approved:one", source, "block_local")
        second = replace(first, source="approved:two", negative=True, statement=source.replace("shall", "shall not"))
        records = [NormalizedSourceRecord(str(index), fact.statement, "approved_normative", "block_local",
                                         (fact.source,), "retained", "approved") for index, fact in enumerate((first, second))]
        self.assertTrue(any("ipos_approved_source_conflict" in finding
                            for finding in ipos_source_preservation_findings(records, [first, second])))

    def test_detail_projection_keeps_multiple_actions_and_conditions_without_name_rules(self):
        statement = "When the channel is active, the Sample_FSM shall sample incoming data and store processed results."
        facets, _rejected = detail_evidence_facets(statement, "")
        self.assertEqual(len(facets), 2)
        self.assertIn("Samples incoming data", facets[0]["text"])
        self.assertIn("Stores processed results", facets[1]["text"])
        self.assertTrue(all("when the channel is active" in facet["text"] for facet in facets))
        renamed, _rejected = detail_evidence_facets(statement.replace("Sample_FSM", "Other_FSM"), "")
        self.assertEqual(facets, renamed)

    def test_detail_groups_do_not_invent_peer_interactions(self):
        rows = [{"decision": "accepted_candidate", "candidate_evidence_requirement_id": "REQ-1",
                 "candidate_evidence_statement": "The Sample_FSM shall sample incoming data.",
                 "candidate_refinement": "Samples incoming data."}]
        summary = compose_detail_groups(rows)
        self.assertEqual(len(summary), 1)
        self.assertNotIn("Interacts with", summary[0])
        self.assertEqual(rows[0]["summary_group_contributor_ids"], "REQ-1")

    def test_ipos_composer_is_inventory_bound_and_always_populates_sections(self):
        record = build_ipos_functional_input(
            "Block A",
            {"Function": "Store samples.", "Inputs": "samples", "Outputs": "status"},
            [{"mapped_block": "Block A", "source_req_id": "REQ-A", "requirement_statement": "The block shall store samples."}],
            domain="digital",
            provenance=("approved snapshot", "block_inventory.csv"),
        )
        sections, audit = compose_ipos_overview(record)
        self.assertTrue(sections["functionality"])
        self.assertTrue(sections["scope"])
        self.assertEqual(audit[0]["derivation_mode"], "direct_function")
        self.assertEqual(audit[1]["derivation_mode"], "function_plus_io_scope")
        self.assertEqual(audit[2]["decision"], "suppressed")
        self.assertNotIn("shall", " ".join(sections["scope"]).casefold())
        self.assertNotIn("register", " ".join(sections["scope"]).casefold())

    def test_ipos_scope_preserves_register_map_compound(self):
        record = build_ipos_functional_input(
            "Block A", {"Function": "Route device signals", "Inputs": "Mode selection, register-map controls",
                        "Outputs": "Routed signals"}, [],
        )
        sections, _audit = compose_ipos_overview(record)
        self.assertIn("register-map controls", sections["scope"][0])
        self.assertNotIn("-map controls", sections["scope"][0].replace("register-map", ""))

    def test_ipos_composer_does_not_use_other_block_or_drs_prose(self):
        record = build_ipos_functional_input(
            "Block A",
            {"Function": "Manage local buffering.", "Inputs": "samples", "Outputs": "buffered data"},
            [{"mapped_block": "Block B", "source_req_id": "REQ-B", "requirement_statement": "The DRS says to process ECG."}],
        )
        sections, audit = compose_ipos_overview(record)
        rendered = " ".join(sections["functionality"] + sections["scope"])
        self.assertNotIn("ECG", rendered)
        candidate = next(item for item in audit if item.get("candidate_evidence_statement"))
        self.assertEqual(candidate["decision"], "suppressed_due_to_cross_block_risk")

    def test_same_block_supported_refinement_is_rendered_and_audited(self):
        row = {
            "mapped_block": "Block A", "source_req_id": "REQ-A",
            "requirement_statement": "The block shall sample incoming data and store processed results.",
        }
        record = build_ipos_functional_input(
            "Block A",
            {"Function": "Acquire and process incoming data", "Inputs": "incoming data", "Outputs": "processed results"},
            [row], domain="digital", provenance=("approved snapshot", "block_inventory.csv"),
        )
        sections, audit = compose_ipos_overview(record)
        candidate = next(item for item in audit if item.get("candidate_refinement"))
        self.assertEqual(candidate["decision"], "accepted_candidate")
        self.assertEqual(candidate["target_section"], "supported_functions_and_scope")
        self.assertEqual(candidate["supporting_local_requirement_ids"], "REQ-A")
        self.assertIn(candidate["rendered_summary"], sections["scope"])
        self.assertTrue(candidate["rendered_summary"].startswith("**Local data processing."))
        self.assertNotIn("Supports approved", candidate["candidate_refinement"])
        self.assertIn("incoming data", candidate["candidate_refinement"])
        self.assertTrue(candidate["output_order"])

    def test_ipos_refinement_preserves_action_object_and_safe_context(self):
        rows = [
            {
                "mapped_block": "Block A", "source_req_id": "SAMPLE",
                "requirement_statement": "When the channel is selected, the controller shall sample incoming ECG data.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "WRITE",
                "requirement_statement": "The controller shall write processed results to the output buffer.",
            },
        ]
        record = build_ipos_functional_input(
            "Block A",
            {
                "Function": "Acquire and process sampled data",
                "Inputs": "channel configuration and sampled data",
                "Outputs": "processed results and output buffer",
            },
            rows,
        )
        sections, audit = compose_ipos_overview(record)
        candidates = {
            row["candidate_evidence_requirement_id"]: row["candidate_refinement"]
            for row in audit if row.get("candidate_evidence_statement")
        }
        self.assertIn("Samples incoming ECG data when the channel is selected.", candidates["SAMPLE"])
        self.assertIn("Writes processed results to output buffer.", candidates["WRITE"])
        self.assertNotEqual(candidates["SAMPLE"], candidates["WRITE"])
        rendered = {
            row["candidate_evidence_requirement_id"]: row["rendered_summary"]
            for row in audit if row.get("decision") == "accepted_candidate"
        }
        self.assertIn("Samples incoming ECG data", rendered["SAMPLE"])
        self.assertIn("Stores processed results to output buffer", rendered["WRITE"])

    def test_conflicts_require_the_same_action_and_target(self):
        inventory = {"Function": "Acquire incoming data", "Inputs": "incoming data", "Outputs": "output memory"}
        rows = [
            {"mapped_block": "Block A", "source_req_id": "STORE", "requirement_statement": "The controller shall store incoming data in output memory."},
            {"mapped_block": "Block A", "source_req_id": "SAMPLE", "requirement_statement": "The controller shall not sample incoming data in output memory."},
        ]
        _sections, audit = compose_ipos_overview(build_ipos_functional_input("Block A", inventory, rows))
        decisions = {row["candidate_evidence_requirement_id"]: row["decision"] for row in audit if row.get("candidate_evidence_statement")}
        self.assertEqual(decisions, {"STORE": "accepted_candidate", "SAMPLE": "accepted_candidate"})

    def test_ipos_projection_omits_parameter_only_and_unparsed_predicates(self):
        rows = [
            {
                "mapped_block": "Block A", "source_req_id": "TIMING",
                "requirement_statement": "The user shall configure time slot duration in the operation register.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "VERB",
                "requirement_statement": "The controller shall change state after setting the mode register.",
            },
        ]
        record = build_ipos_functional_input(
            "Block A",
            {"Function": "Manage sampled data", "Inputs": "configuration and samples", "Outputs": "control state"},
            rows,
        )
        sections, audit = compose_ipos_overview(record)
        candidates = [row for row in audit if row.get("candidate_evidence_statement")]
        self.assertFalse(any(row["candidate_refinement"] for row in candidates))
        self.assertFalse(any("time slot" in item.casefold() or "mode register" in item.casefold() for item in sections["scope"]))

    def test_ipos_channel_refinements_use_configuration_boundary_and_keep_distinct_targets(self):
        rows = [
            {
                "mapped_block": "Block A", "source_req_id": "CHANNEL-A",
                "requirement_statement": "The user shall select the active channel, in this case CH_A_SEL.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "CHANNEL-B",
                "requirement_statement": "The user shall select the active channel, in this case CH_B_SEL.",
            },
        ]
        record = build_ipos_functional_input(
            "Block A",
            {"Function": "Process sampled data", "Inputs": "configuration settings", "Outputs": "processed results"},
            rows,
        )
        sections, audit = compose_ipos_overview(record)
        candidates = [row for row in audit if row.get("decision") == "accepted_candidate"]
        summaries = {row["candidate_evidence_requirement_id"]: row["candidate_refinement"] for row in candidates}
        self.assertEqual(len(summaries), 2)
        self.assertIn("CH A", summaries["CHANNEL-A"])
        self.assertIn("CH B", summaries["CHANNEL-B"])
        self.assertNotEqual(summaries["CHANNEL-A"], summaries["CHANNEL-B"])
        rendered_groups = {row["rendered_summary"] for row in candidates if row.get("rendered_summary")}
        self.assertEqual(len(rendered_groups), 1)
        self.assertEqual(len(sections["scope"]), 2)
        self.assertIn("CH A", next(iter(rendered_groups)))
        self.assertIn("CH B", next(iter(rendered_groups)))

    def test_ipos_refinement_grouping_is_bounded_and_preserves_provenance(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "ONE", "requirement_statement": "The user shall select the active channel, in this case CH_A_SEL."},
            {"mapped_block": "Block A", "source_req_id": "TWO", "requirement_statement": "The user shall select the active channel, in this case CH_B_SEL."},
            {"mapped_block": "Block A", "source_req_id": "THREE", "requirement_statement": "The controller shall write processed results to the output buffer."},
        ]
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Process sampled data", "Inputs": "configuration settings", "Outputs": "processed results and output buffer"}, rows,
        ))
        accepted = [row for row in audit if row.get("decision") == "accepted_candidate"]
        self.assertLess(len(sections["scope"]) - 1, len(accepted))
        selection_rows = [row for row in accepted if row["candidate_evidence_requirement_id"] in {"ONE", "TWO"}]
        self.assertEqual({row["summary_group_id"] for row in selection_rows}, {"local-scope-01"})
        self.assertTrue(all(row["rendered_summary"].startswith("**Local configuration.") for row in selection_rows))
        self.assertTrue(all(row["candidate_evidence_requirement_id"] in row["summary_group_contributor_ids"] for row in accepted))

    def test_ipos_refinement_grouping_keeps_distinct_actions_separate(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "SAMPLE", "requirement_statement": "The controller shall sample incoming data."},
            {"mapped_block": "Block A", "source_req_id": "WRITE", "requirement_statement": "The controller shall write processed results to the output buffer."},
        ]
        sections, _audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Acquire and process incoming data", "Inputs": "incoming data", "Outputs": "processed results and output buffer"}, rows,
        ))
        self.assertTrue(any(entry.startswith("**Local data processing.") for entry in sections["scope"]))
        self.assertTrue(any(entry.startswith("**Operation control and data routing.") for entry in sections["scope"]))

    def test_ipos_refinement_groups_source_exact_fsm_macro_subfunctions(self):
        rows = [
            {
                "mapped_block": "Block A", "source_req_id": "DEVICE",
                "requirement_statement": "When Device_FSM FSM is active, the controller shall control sampled data.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "ELAB",
                "requirement_statement": "When ELAB_FSM is active, the controller shall process sampled data.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "ADC",
                "requirement_statement": "When ADC_Phases is active, the ADC start signal shall stay high.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "PPG",
                "requirement_statement": "When PPG_FSM is active, the PPG output shall be high.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "PPG-VARIANT",
                "requirement_statement": "When PPG FSM is active, the PPG output shall be high.",
            },
        ]
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A",
            {"Function": "Control and process sampled data", "Inputs": "sampled data", "Outputs": "control state"},
            rows,
        ))
        summaries = sections["scope"][1:]
        self.assertEqual(len(summaries), 4)
        self.assertTrue(any(summary.startswith("**Device FSM.") for summary in summaries))
        self.assertTrue(any(summary.startswith("**ELAB FSM.") for summary in summaries))
        self.assertTrue(any(summary.startswith("**ADC phases.") for summary in summaries))
        self.assertEqual(sum(summary.startswith("**PPG FSM.") for summary in summaries), 1)
        self.assertTrue(all(summary.count("\n") <= 1 for summary in summaries))
        self.assertTrue(any("ADC start" in summary for summary in summaries))
        self.assertTrue(any("PPG output" in summary for summary in summaries))
        accepted = [row for row in audit if row.get("decision") == "accepted_candidate"]
        self.assertEqual(
            {row["summary_group_id"] for row in accepted},
            {"local-scope-01", "local-scope-02", "local-scope-03", "local-scope-04"},
        )
        structural_rows = [row for row in accepted if row.get("structural_facets_json")]
        self.assertTrue(all(row["structural_selected_evidence_ids"] for row in structural_rows))

    def test_structural_aggregation_uses_all_accepted_same_block_facets(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "ONE", "requirement_statement": "When ADC_Phases is active, the controller shall start ADC sampling."},
            {"mapped_block": "Block A", "source_req_id": "TWO", "requirement_statement": "When ADC_Phases is active, the controller shall wait for i_ADC_EOC."},
        ]
        sections, audit = compose_ipos_overview(
            build_ipos_functional_input(
                "Block A", {"Function": "Control sampled data", "Inputs": "sampled data", "Outputs": "control state"}, rows,
            ),
        )
        summary = next(item for item in sections["scope"] if item.startswith("**ADC phases."))
        self.assertIn("Starts ADC sampling", summary)
        self.assertIn("ADC conversion completion", summary)
        structural = [row for row in audit if row.get("structural_entity_display") == "ADC_Phases"]
        self.assertEqual(structural[0]["structural_selected_evidence_ids"], "ONE; TWO")
        self.assertEqual([row["structural_aggregation_order"] for row in structural], ["1", "2"])

    def test_deterministic_aggregation_uses_all_accepted_entity_evidence(self):
        rows = [
            {
                "mapped_block": "Block A", "source_req_id": "ELAB-ONE",
                "requirement_statement": "When ELAB_FSM is active, the controller shall process ECG samples in PPG phase and wait for Device_FSM.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "ELAB-TWO",
                "requirement_statement": "When ELAB FSM is active, the controller shall sequence BIA samples in ECG phase.",
            },
        ]
        sections, audit = compose_ipos_overview(
            build_ipos_functional_input(
                "Block A",
                {"Function": "Process sampled data", "Inputs": "sampled data", "Outputs": "control state"},
                rows,
            ),
        )
        summary = next(item for item in sections["scope"] if item.startswith("**ELAB FSM."))
        self.assertIn("processes ECG samples", summary)
        self.assertIn("sequences BIA samples", summary)
        self.assertIn("Interacts with Device FSM", summary)
        self.assertNotIn("PPG phase", summary)
        self.assertNotIn("local acquisition and processing sequence", summary)
        self.assertNotIn("Provides", summary)
        structural = [row for row in audit if row.get("structural_entity_display") == "ELAB_FSM"]
        self.assertTrue(all(row["structural_selection_mode"] == "deterministic_aggregation" for row in structural))
        self.assertTrue(all(row["structural_selected_evidence_ids"] == "ELAB-ONE; ELAB-TWO" for row in structural))

    def test_deterministic_aggregation_reports_explicit_block_and_signal_dependencies(self):
        rows = [{
            "mapped_block": "Block A", "source_req_id": "ADC-ONE",
            "requirement_statement": "When ADC_Phases is in acquisition phase, the controller shall wait for i_ADC_EOC from the ADC block.",
        }]
        sections, _audit = compose_ipos_overview(
            build_ipos_functional_input(
                "Block A",
                {"Function": "Process sampled data", "Inputs": "sampled data", "Outputs": "control state"},
                rows,
            ),
        )
        summary = next(item for item in sections["scope"] if item.startswith("**ADC phases."))
        self.assertIn("Interacts with ADC block", summary)
        self.assertNotIn("i_ADC_EOC", summary)

    def test_aggregation_merges_channel_setup_with_measurement_start_and_omits_interaction_filler(self):
        rows = [
            {
                "mapped_block": "Block A", "source_req_id": "CONFIG",
                "requirement_statement": "The user shall select ECG0, BIA, GSR, PPG, and ECG1 ECG2 channels and Data Storage Mode.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "START",
                "requirement_statement": "The user shall start measurement.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "GSR-AVG",
                "requirement_statement": "The average block shall perform the average on 16 sample acquisition for GSR.",
            },
        ]
        sections, _audit = compose_ipos_overview(
            build_ipos_functional_input(
                "Block A",
                {"Function": "Control measurements", "Inputs": "configuration and samples", "Outputs": "measurement results"},
                rows,
            ),
        )
        summaries = [item for item in sections["scope"] if item.startswith("**")]
        setup = next(item for item in summaries if item.startswith("**Measurement setup and initiation."))
        averaging = next(item for item in summaries if "16 samples for GSR" in item)
        self.assertIn("ECG0", setup)
        self.assertIn("Data Storage Mode", setup)
        self.assertIn("and starts measurement", setup)
        self.assertNotIn("then", setup)
        self.assertNotIn("No direct dependency", "\n".join(summaries))
        self.assertNotIn("Local configuration.", "\n".join(summaries))
        self.assertNotIn("Operation control and data routing.", "\n".join(summaries))
        self.assertIn("Averages 16 samples for GSR", averaging)

    def test_approved_gsr_formula_and_cds_averages_remain_in_summary(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "INTERVAL", "requirement_statement":
             "The i_m_sensor shall define how many time slots pass between two sensor acquisitions following this formula: sensor_data = time_slot 2 * m_sensor"},
            {"mapped_block": "Block A", "source_req_id": "CDS", "requirement_statement":
             "If en_sensor_cds is set high the SENSOR_SNS and SENSOR_CDS shall be sampled 16 times, then mediated to get one value for each of them, and then the final result shall be the subtraction between SENSOR_SNS and SENSOR_CDS."},
            {"mapped_block": "Block A", "source_req_id": "AVERAGE", "requirement_statement":
             "The average block shall perform the average on 16 sample acquisition for sensor."},
        ]
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Control sensor acquisition", "Inputs": "configuration and samples", "Outputs": "measurement results"}, rows,
        ))
        summary = " ".join(sections["scope"]).casefold()
        for meaning in ("time slots between two sensor acquisitions", "without resolving its arithmetic",
                        "16 acquisitions of sensor sns and sensor cds", "subtracts their averaged values",
                        "16 samples for sensor"):
            self.assertIn(meaning, summary)
        self.assertEqual(
            {row["candidate_evidence_requirement_id"] for row in audit if row.get("decision") == "accepted_candidate"},
            {"INTERVAL", "CDS", "AVERAGE"},
        )

    def test_approved_procedure_name_is_not_reduced_to_following_steps(self):
        statement = (
            "To perform a SOFT Reset procedure, the user shall perform the following steps: "
            "1. If Mode Operation has been set to 1, reset it to 0 "
            "2. Set the Soft Reset bit to 0 for at least 1 ms "
            "3. Set the Soft Reset bit to 1 "
            "4. Wait for Device Status to reach IDLE"
        )
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Control reset", "Inputs": "reset settings", "Outputs": "device status"},
            [{"mapped_block": "Block A", "source_req_id": "RESET", "requirement_statement": statement}],
        ))
        candidate = next(row for row in audit if row.get("candidate_evidence_requirement_id") == "RESET")
        self.assertEqual(candidate["decision"], "accepted_candidate")
        self.assertEqual(candidate["materialization_status"], "suppressed_atomic")
        self.assertNotIn("following steps", " ".join(sections["scope"]).casefold())
        self.assertEqual(candidate["candidate_evidence_statement"], statement)

    def test_prose_before_table_and_explicit_clamp_remain_eligible(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "PPG", "requirement_statement":
             "Once the PPG Raw Data and PPG Noise reach the average value, depending on the ALC Mask bit value, the two accumulated shall be subtracted and final results shall be computed as follow: Table 18: PPG Result Table"},
            {"mapped_block": "Block A", "source_req_id": "FRAMES", "requirement_statement":
             "After the last repetition has been accumulated, the data output shall be obtained dividing by the number of repeated frames the accumulated data to obtain a 20 bit data. Table 20: Frame results"},
            {"mapped_block": "Block A", "source_req_id": "CLAMP", "requirement_statement":
             "If the user set a value greater than 1024 into ones of the m_channel, the value shall clamp at 1024 (check the input i_m_ecg"},
        ]
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Process PPG and frame data", "Inputs": "configuration and sampled data", "Outputs": "results"}, rows,
        ))
        summary = " ".join(sections["scope"]).casefold()
        for meaning in ("ppg raw data and ppg noise", "alc mask", "number of repetitions", "20-bit", "1024"):
            self.assertIn(meaning, summary)
        self.assertEqual(
            {row["candidate_evidence_requirement_id"] for row in audit if row.get("decision") == "accepted_candidate"},
            {"PPG", "FRAMES", "CLAMP"},
        )

    def test_explicit_ordered_state_progression_is_not_lost(self):
        statement = (
            "When wave velocity mode is selected and the device FSM is in WAIT_SU, "
            "the Elaboration FSM shall go into the TIMER BUFFER State and then into the PPG ALC STORAGE"
        )
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Control PPG elaboration", "Inputs": "wave velocity mode", "Outputs": "control state"},
            [{"mapped_block": "Block A", "source_req_id": "TRANSITION", "requirement_statement": statement}],
        ))
        candidate = next(row for row in audit if row.get("candidate_evidence_requirement_id") == "TRANSITION")
        self.assertEqual(candidate["decision"], "accepted_candidate")
        self.assertEqual(candidate["materialization_status"], "suppressed_atomic")
        self.assertNotIn("TIMER BUFFER then PPG ALC STORAGE".casefold(), " ".join(sections["scope"]).casefold())

    def test_orphaned_ppg_handoff_does_not_become_a_standalone_function(self):
        statement = (
            "When in ALC Compensation, the PPG shall wait that the rise of "
            "i_end_compensation signal from the ALC Compensation block to go into "
            "the RX_Start_UP state. This shall happen in OPERATIVE."
        )
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Control PPG", "Inputs": "PPG samples", "Outputs": "PPG state"},
            [{"mapped_block": "Block A", "source_req_id": "HANDOFF", "requirement_statement": statement}],
        ))
        candidate = next(row for row in audit if row.get("candidate_evidence_requirement_id") == "HANDOFF")
        self.assertEqual(candidate["materialization_status"], "suppressed_atomic")
        self.assertNotIn("handoff", " ".join(sections["scope"]).casefold())

    def test_state_entry_does_not_fold_prior_state_into_destination(self):
        from workflow_routing import _ipos_refinement_parts

        statement = (
            "When only GSR is selected, if i_m_gsr is greater than 0, the Device FSM "
            "shall go in SLEEP after the IDLE State and shall go back in Operative State "
            "only when the measurement runs."
        )
        action, destination, _qualifier, _negative = _ipos_refinement_parts(statement)
        self.assertNotIn("SLEEP after the IDLE", destination)
        self.assertNotIn("Operative state", destination)

    def test_return_from_sleep_keeps_following_guard(self):
        from workflow_routing import _ipos_refinement_parts

        statement = (
            "When in Sleep Mode, the device shall go back to the Operative State "
            "if no error (o_error_irq stuck at 0) occurs and all the operations are ended."
        )
        action, destination, _qualifier, _negative = _ipos_refinement_parts(statement)
        self.assertEqual(action, "enter")
        self.assertIn("Operative state when in Sleep Mode if no error", destination)
        self.assertIn("all the operations are ended", destination)
        self.assertNotIn("the device if", destination)

    def test_control_effects_preserve_condition_reset_and_timed_step(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "FLAG", "requirement_statement":
             "In case the samples are equal to 0x8000 or 0x7FFF for at least half of the selected average, "
             "o_saturation_flag shall be set to 1 and shall go to 0 at the start of the next frame or frame repetition."},
            {"mapped_block": "Block A", "source_req_id": "CLOCK", "requirement_statement":
             "When ECG, BIA or GSR is sampled, i_ADC_en is equal to 1, the o_ADC_clk_en shall be set to 1."},
            {"mapped_block": "Block A", "source_req_id": "GSR", "requirement_statement":
             "If en_gsr_cds is set low [bit 3 of GSR_CTRL], the gsr_curr_on shall set high in every time slot"},
            {"mapped_block": "Block A", "source_req_id": "STEP", "requirement_statement":
             "When the input signal o_Tx_preset is set to 1, the o_Tx_Vref shall increment the first step "
             "at value i_N_start at a time i_T_preset."},
        ]
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Control sampling and outputs", "Inputs": "samples and timing control", "Outputs": "control state and data"}, rows,
        ))
        summary = " ".join(sections["scope"]).casefold()
        for meaning in ("saturation flag", "at least half", "clears at the start of the next frame",
                        "adc clk en", "ecg, bia or gsr is sampled", "gsr curr on", "every time slot",
                        "tx vref", "n start", "t preset"):
            self.assertIn(meaning, summary)
        self.assertEqual(
            {row["candidate_evidence_requirement_id"] for row in audit if row.get("decision") == "accepted_candidate"},
            {"FLAG", "CLOCK", "GSR", "STEP"},
        )
        summaries = {row["candidate_evidence_requirement_id"]: row["rendered_summary"] for row in audit if row.get("decision") == "accepted_candidate"}
        self.assertNotEqual(summaries["FLAG"], summaries["CLOCK"])

    def test_structural_summary_does_not_split_a_state_path_list(self):
        from workflow_routing import _ipos_structural_aggregate_summary

        members = [{"candidate_evidence_requirement_id": "PPG", "candidate_evidence_statement":
                    "When PPG_FSM is in Measurement state, it shall stay in this state until the rise of i_end_average. "
                    "When PPG_FSM is in Rising Ramp state, it shall stay until Rising Time ends. "
                    "When PPG_FSM is in Falling Ramp state, it shall stay until Falling Time ends. "
                    "When PPG_FSM is in SECOND ALC SAMPLES state, it shall wait for ADC completion."}]
        summary, _facets, _reason = _ipos_structural_aggregate_summary("PPG_FSM", members)
        self.assertNotIn("\n  SECOND ALC SAMPLES paths.", summary)

    def test_aggregation_device_lifecycle_and_elaboration_summary_keep_all_paths_and_relations(self):
        rows = [
            {
                "mapped_block": "Block A", "source_req_id": "DEVICE-BOOT",
                "requirement_statement": "When Device_FSM is in BOOT state, it shall wait for i_quokka_boot_end and then go to IDLE state, with clock generated from Clock generator block.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "DEVICE-OPERATIVE",
                "requirement_statement": "When Device_FSM is in IDLE state, it shall enter WAIT_SU state and OPERATIVE state.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "DEVICE-SLEEP-ERROR",
                "requirement_statement": "When Device_FSM is in OPERATIVE state, it shall enter SLEEP state and ERROR state, and raise o_error_time_slot to Digital_Top when the time slot expires.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "ELAB-ECG",
                "requirement_statement": "When ELAB_FSM is in ECG state, it shall process ECG samples and set ADC_Mux output for ADC PHASE acquisition.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "ELAB-PPG",
                "requirement_statement": "When ELAB_FSM is in PPG state, it shall wait for PPG_START_UP_TIME and start PPG operation for the selected frame.",
            },
            {
                "mapped_block": "Block A", "source_req_id": "ELAB-BIA",
                "requirement_statement": "When ELAB_FSM is in BIA state, it shall sequence BIA samples and enable ADC acquisition.",
            },
        ]
        sections, _audit = compose_ipos_overview(
            build_ipos_functional_input(
                "Block A",
                {"Function": "Control local measurements", "Inputs": "configuration and samples", "Outputs": "measurement data"},
                rows,
            ),
        )
        device = next(item for item in sections["scope"] if item.startswith("**Device FSM."))
        elaboration = next(item for item in sections["scope"] if item.startswith("**ELAB FSM."))
        device_casefold = device.casefold()
        for state in ("Boot", "Idle", "start-up", "Operative", "Sleep", "Error"):
            self.assertIn(state.casefold(), device_casefold)
        self.assertIn("ADSP boot completion", device)
        self.assertIn("Digital Top error interrupt", device)
        self.assertIn("ECG", elaboration)
        self.assertIn("BIA", elaboration)
        self.assertIn("PPG", elaboration)
        self.assertIn("ADC phase acquisition", elaboration)
        self.assertIn("PPG operation control", elaboration)
        self.assertNotIn("i_quokka_boot_end", device)
        self.assertNotIn("ADC_Mux", elaboration)

    def test_aggregation_responsibility_keeps_all_accepted_action_phrases(self):
        summary = _ipos_aggregate_responsibility(
            [
                "Processes ECG samples",
                "Sequences BIA samples",
                "Starts PPG operation",
                "Enables ADC acquisition",
            ],
            [],
        )
        for phrase in ("Processes ECG samples", "sequences BIA samples", "starts PPG operation", "Enables ADC acquisition"):
            self.assertIn(phrase.casefold(), summary.casefold())

    def test_structural_aggregation_renders_supported_responsibility_without_interaction(self):
        rows = [{
            "mapped_block": "Block A", "source_req_id": "ELAB-ONE",
            "requirement_statement": "When ELAB_FSM is active, the controller shall process sampled data.",
        }]
        sections, audit = compose_ipos_overview(
            build_ipos_functional_input(
                "Block A",
                {"Function": "Process sampled data", "Inputs": "sampled data", "Outputs": "control state"},
                rows,
            ),
        )
        self.assertTrue(any(item.startswith("**ELAB FSM.") for item in sections["scope"]))
        structural = next(row for row in audit if row.get("structural_entity_display") == "ELAB_FSM")
        self.assertEqual(structural["structural_summary_status"], "selected")
        self.assertEqual(structural["structural_summary_reason"], "deterministic aggregation")

    def test_ipos_validator_rejects_requirement_expansion_and_boilerplate(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "ONE", "requirement_statement": "The user shall select the active channel, in this case CH_A_SEL."},
            {"mapped_block": "Block A", "source_req_id": "TWO", "requirement_statement": "The user shall select the active channel, in this case CH_B_SEL."},
            {"mapped_block": "Block A", "source_req_id": "THREE", "requirement_statement": "The user shall select the active channel, in this case CH_C_SEL."},
        ]
        record = build_ipos_functional_input(
            "Block A", {"Function": "Process sampled data", "Inputs": "configuration settings", "Outputs": "processed results"}, rows,
        )
        sections, audit = compose_ipos_overview(record)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            markdown_path = root / "ipos.md"
            audit_path = root / "descriptive_summary_audit.csv"
            markdown_path.write_text(
                "## 1. Block overview\n\n### 1.1 Functionality\n\n" + sections["functionality"][0]
                + "\n\n### 1.2 Supported functions and scope\n\n- " + sections["scope"][0]
                + "\n- Supports approved local operations.\n\n## 2. Source I/O\n", encoding="utf-8",
            )
            for index, row in enumerate(audit):
                if row.get("decision") == "accepted_candidate":
                    row["summary_group_id"] = f"expanded-{index}"
                    row["rendered_summary"] = row["candidate_refinement"]
                    row["summary_group_contributor_ids"] = row["candidate_evidence_requirement_id"]
            write_ipos_descriptive_audit(audit_path, audit)
            findings = validate_ipos_descriptive_output(
                markdown_path, audit_path, "Block A", [row["requirement_statement"] for row in rows],
                [row["source_req_id"] for row in rows], DESCRIPTIVE_DIGITAL_TOPIC_SPECS, record,
            )
        joined = " ".join(findings)
        self.assertIn("requirement_by_requirement_scope_expansion", joined)
        self.assertIn("generic_action_family_placeholder", joined)

    def test_dry_run_rejects_table_signal_cross_block_conflict_and_normative_only_evidence(self):
        inventory = {"Function": "Store incoming data", "Inputs": "incoming data", "Outputs": "output memory"}
        cases = (
            ([{"mapped_block": "Block A", "source_req_id": "TABLE", "requirement_statement": "The block shall route values as shown in the following table."}], "suppressed_due_to_table_or_signal_only_evidence"),
            ([{"mapped_block": "Block A", "source_req_id": "SIGNALS", "requirement_statement": "ECG, PPG, GSR"}], "suppressed_due_to_table_or_signal_only_evidence"),
            ([{"mapped_block": "Block B", "source_req_id": "OTHER", "requirement_statement": "The block shall store incoming data in output memory."}], "suppressed_due_to_cross_block_risk"),
            ([{"mapped_block": "Block A", "source_req_id": "NORM", "requirement_statement": "The block shall comply with its assigned parameters."}], "suppressed_due_to_normative_reuse"),
        )
        for rows, expected in cases:
            with self.subTest(expected=expected):
                record = build_ipos_functional_input("Block A", inventory, rows)
                _sections, audit = compose_ipos_overview(record)
                candidate = next(item for item in audit if item.get("candidate_evidence_statement"))
                self.assertEqual(candidate["decision"], expected)

        bounded = build_ipos_functional_input(
            "Block A",
            {"Function": "Acquire incoming data", "Inputs": "incoming data", "Outputs": "status"},
            [{"mapped_block": "Block A", "source_req_id": "BOUNDED", "requirement_statement": "The block shall sample incoming data."}],
        )
        bounded_sections, bounded_audit = compose_ipos_overview(bounded)
        bounded_candidate = next(item for item in bounded_audit if item.get("candidate_evidence_statement"))
        self.assertEqual(bounded_candidate["decision"], "accepted_candidate")
        self.assertIn(bounded_candidate["rendered_summary"], bounded_sections["scope"])

        conflict_rows = [
            {"mapped_block": "Block A", "source_req_id": "POS", "requirement_statement": "The block shall store incoming data in output memory."},
            {"mapped_block": "Block A", "source_req_id": "NEG", "requirement_statement": "The block shall not store incoming data in output memory."},
        ]
        _sections, conflict_audit = compose_ipos_overview(build_ipos_functional_input("Block A", inventory, conflict_rows))
        self.assertEqual(
            {item["decision"] for item in conflict_audit if item.get("candidate_evidence_statement")},
            {"suppressed_due_to_conflict"},
        )

    def test_conditional_passive_set_retains_functional_effect_and_transition(self):
        inventory = {"Function": "Control acquisition", "Inputs": "sampled data", "Outputs": "control state"}
        statement = (
            "When in FIRST_SAMPLES state, when i_end_average is set to 1, "
            "the o_start_ramp signal shall be set high to start the rising of the ramp "
            "and go into the RISING_RAMP state."
        )
        record = build_ipos_functional_input(
            "Block A", inventory,
            [{"mapped_block": "Block A", "source_req_id": "TRANSITION", "requirement_statement": statement},
             {"mapped_block": "Block A", "source_req_id": "BARE", "requirement_statement":
              "When in READY state, when i_go is set to 1, the o_flag signal shall be set high."}],
        )
        sections, audit = compose_ipos_overview(record)
        transition = next(item for item in audit if item.get("candidate_evidence_requirement_id") == "TRANSITION")
        bare = next(item for item in audit if item.get("candidate_evidence_requirement_id") == "BARE")
        self.assertEqual(transition["decision"], "accepted_candidate")
        summary = transition["rendered_summary"].casefold()
        for meaning in ("rising of ramp", "rising ramp state", "end average is set to 1", "first samples state"):
            self.assertIn(meaning, summary)
        self.assertIn(transition["rendered_summary"], sections["scope"])
        self.assertEqual(bare["decision"], "suppressed_due_to_normative_reuse")

    def test_approved_acquisition_interval_survives_define_and_time_filters(self):
        inventory = {"Function": "Control acquisition", "Inputs": "configuration settings", "Outputs": "control state"}
        rows = [
            {"mapped_block": "Block A", "source_req_id": "INTERVAL", "requirement_statement":
             "The i_m_sensor [SENSOR_FREQ] shall define how many time slots pass between two sensor acquisitions following this formula: unreadable"},
            {"mapped_block": "Block A", "source_req_id": "UNSUPPORTED", "requirement_statement":
             "The block shall define the configuration."},
        ]
        sections, audit = compose_ipos_overview(build_ipos_functional_input("Block A", inventory, rows))
        interval = next(item for item in audit if item.get("candidate_evidence_requirement_id") == "INTERVAL")
        unsupported = next(item for item in audit if item.get("candidate_evidence_requirement_id") == "UNSUPPORTED")
        self.assertEqual(interval["decision"], "accepted_candidate")
        self.assertIn("time slots between two sensor acquisitions", interval["rendered_summary"])
        self.assertIn(interval["rendered_summary"], sections["scope"])
        self.assertNotIn("unreadable", " ".join(sections["scope"]))
        self.assertEqual(unsupported["decision"], "suppressed_due_to_normative_reuse")

    def test_digital_and_analog_use_identical_dry_run_classification(self):
        row = {"mapped_block": "Block A", "source_req_id": "REQ-A", "requirement_statement": "The block shall sample incoming data and store processed results."}
        inventory = {"Function": "Acquire and process incoming data", "Inputs": "incoming data", "Outputs": "processed results"}
        decisions = []
        for domain in ("digital", "analog"):
            record = build_ipos_functional_input("Block A", inventory, [row], domain=domain)
            _sections, audit = compose_ipos_overview(record)
            decisions.append((audit[-1]["decision"], audit[-1]["candidate_refinement"]))
        self.assertEqual(decisions[0], decisions[1])

    def test_global_candidate_audit_keeps_cross_block_evidence_suppressed(self):
        local_row = {
            "mapped_block": "Block A", "source_req_id": "REQ-A",
            "requirement_statement": "The controller shall sample incoming data and store processed results in the local output buffer.",
        }
        other_block_row = {
            "mapped_block": "Block B", "source_req_id": "REQ-B",
            "requirement_statement": "The controller shall sample incoming data and store processed results in the local output buffer.",
        }
        candidate_rows = [local_row, other_block_row]
        record = build_ipos_functional_input(
            "Block A",
            {"Function": "Acquire and process incoming data", "Inputs": "incoming data", "Outputs": "processed results"},
            candidate_rows,
            domain="digital",
            candidate_requirement_rows=candidate_rows,
        )

        _sections, audit = compose_ipos_overview(record)
        candidates = {
            item["candidate_evidence_requirement_id"]: item
            for item in audit
            if item.get("candidate_evidence_statement")
        }

        self.assertEqual(record.local_requirement_ids, ("REQ-A",))
        self.assertEqual(candidates["REQ-A"]["decision"], "accepted_candidate")
        self.assertEqual(candidates["REQ-B"]["decision"], "suppressed_due_to_cross_block_risk")
        self.assertEqual(candidates["REQ-B"]["supporting_local_requirement_ids"], "")

    def test_shared_validator_rejects_loss_of_accepted_candidate(self):
        row = {"mapped_block": "Block A", "source_req_id": "REQ-A", "requirement_statement": "The block shall sample incoming data and store processed results."}
        record = build_ipos_functional_input(
            "Block A",
            {"Function": "Acquire and process incoming data", "Inputs": "incoming data", "Outputs": "processed results"},
            [row], provenance=("approved snapshot", "block_inventory.csv"),
        )
        sections, audit = compose_ipos_overview(record)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            markdown_path = root / "ipos.md"
            audit_path = root / "descriptive_summary_audit.csv"
            markdown_path.write_text(
                "## 1. Block overview\n\n### 1.1 Functionality\n\n" + sections["functionality"][0]
                + "\n\n### 1.2 Supported functions and scope\n\n"
                + "\n".join("- " + item for item in sections["scope"])
                + "\n\n## 2. Source I/O\n", encoding="utf-8",
            )
            write_ipos_descriptive_audit(audit_path, audit)
            findings = validate_ipos_descriptive_output(
                markdown_path, audit_path, "Block A", [row["requirement_statement"]], ["REQ-A"],
                DESCRIPTIVE_DIGITAL_TOPIC_SPECS, record,
            )
            self.assertEqual(findings, [])
            with audit_path.open("r", encoding="utf-8", newline="") as handle:
                audit_rows = list(csv.DictReader(handle))
                fieldnames = list(audit_rows[0])
            candidate = next(item for item in audit_rows if item.get("candidate_evidence_requirement_id") == "REQ-A")
            candidate["rendered_summary"] = "**Operation control and data routing.** Starts only."
            with audit_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(audit_rows)
            findings = validate_ipos_descriptive_output(
                markdown_path, audit_path, "Block A", [row["requirement_statement"]], ["REQ-A"],
                DESCRIPTIVE_DIGITAL_TOPIC_SPECS, record,
            )
            self.assertIn("low_quality_subfunction_summary:Operation control and data routing", " ".join(findings))
            write_ipos_descriptive_audit(audit_path, audit)
            markdown_path.write_text(markdown_path.read_text(encoding="utf-8").replace(
                "- " + audit[-1]["rendered_summary"] + "\n", ""
            ), encoding="utf-8")
            findings = validate_ipos_descriptive_output(
                markdown_path, audit_path, "Block A", [row["requirement_statement"]], ["REQ-A"],
                DESCRIPTIVE_DIGITAL_TOPIC_SPECS, record,
            )
        self.assertIn("accepted_refinement_group_missing_from_rendered_scope", " ".join(findings))

    def test_ipos_validator_rejects_governance_prose_in_final_overview(self):
        record = build_ipos_functional_input(
            "Block A", {"Function": "Route samples", "Inputs": "incoming samples", "Outputs": "routed samples"},
            [], provenance=("approved snapshot",),
        )
        sections, audit = compose_ipos_overview(record)
        with tempfile.TemporaryDirectory() as directory:
            markdown_path = Path(directory) / "ipos.md"
            audit_path = Path(directory) / "descriptive_summary_audit.csv"
            write_ipos_descriptive_audit(audit_path, audit)
            for legacy_text in (
                "The block provides this approved macro-purpose: route samples.",
                "The local scope covers the approved Block A function, accepting incoming samples.",
                "The block uses approved functional evidence to route samples.",
                "The block provides approved signal handling for routed samples.",
            ):
                with self.subTest(legacy_text=legacy_text):
                    markdown_path.write_text(
                        "## 1. Block overview\n\n### 1.1 Functionality\n\n" + legacy_text
                        + "\n\n### 1.2 Supported functions and scope\n\n- " + sections["scope"][0]
                        + "\n\n## 2. Source I/O\n", encoding="utf-8",
                    )
                    findings = validate_ipos_descriptive_output(
                        markdown_path, audit_path, "Block A", [], [], DESCRIPTIVE_DIGITAL_TOPIC_SPECS, record,
                    )
                    self.assertTrue(any("description_uses_internal_process_language" in finding for finding in findings))

    def test_validator_rejects_raw_standalone_heading(self):
        statement = "The controller shall write processed results to the FIFO with a dedicated tag."
        record = build_ipos_functional_input(
            "Block A", {"Function": "Store results", "Inputs": "processed results", "Outputs": "FIFO data"},
            [{"mapped_block": "Block A", "source_req_id": "STORE", "requirement_statement": statement}],
            provenance=("approved snapshot",),
        )
        sections, audit = compose_ipos_overview(record)
        with tempfile.TemporaryDirectory() as directory:
            markdown_path = Path(directory) / "ipos.md"
            audit_path = Path(directory) / "descriptive_summary_audit.csv"
            markdown_path.write_text(
                "## 1. Block overview\n\n### 1.1 Functionality\n\n" + sections["functionality"][0]
                + "\n\n### 1.2 Supported functions and scope\n\n"
                + "\n".join("- " + item for item in sections["scope"])
                + "\n\n## 2. Source I/O\n", encoding="utf-8",
            )
            write_ipos_descriptive_audit(audit_path, audit)
            with audit_path.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
                fieldnames = list(rows[0])
            candidate = next(row for row in rows if row.get("candidate_evidence_requirement_id") == "STORE")
            self.assertEqual(candidate["materialization_status"], "standalone")
            candidate["rendered_summary"] = "**raw_signal_name.** Enables output when set."
            with audit_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows)
            findings = validate_ipos_descriptive_output(
                markdown_path, audit_path, "Block A", [statement], ["STORE"],
                DESCRIPTIVE_DIGITAL_TOPIC_SPECS, record,
            )
        self.assertIn("low_quality_standalone_summary:STORE", " ".join(findings))

    def test_validator_fails_when_parser_silently_suppresses_supported_reset_behavior(self):
        statement = (
            "In case the sampled data reach the threshold, o_saturation_flag shall be set to 1 "
            "and shall go to 0 at the start of the next frame."
        )
        record = build_ipos_functional_input(
            "Block A", {"Function": "Control sampling", "Inputs": "sampled data", "Outputs": "control outputs"},
            [{"mapped_block": "Block A", "source_req_id": "FLAG", "requirement_statement": statement}],
            provenance=("approved snapshot",),
        )
        with patch("workflow_routing._ipos_refinement_candidate", return_value=""):
            sections, audit = compose_ipos_overview(record)
            with tempfile.TemporaryDirectory() as directory:
                markdown_path = Path(directory) / "ipos.md"
                audit_path = Path(directory) / "descriptive_summary_audit.csv"
                markdown_path.write_text(
                    "## 1. Block overview\n\n### 1.1 Functionality\n\n" + sections["functionality"][0]
                    + "\n\n### 1.2 Supported functions and scope\n\n"
                    + "\n".join("- " + item for item in sections["scope"])
                    + "\n\n## 2. Source I/O\n", encoding="utf-8",
                )
                write_ipos_descriptive_audit(audit_path, audit)
                findings = validate_ipos_descriptive_output(
                    markdown_path, audit_path, "Block A", [statement], ["FLAG"],
                    DESCRIPTIVE_DIGITAL_TOPIC_SPECS, record,
                )
        self.assertIn("approved_behavior_missing_from_scope:FLAG:saturation flag", " ".join(findings))

    def test_ipos_subfunction_summary_has_title_purpose_context_and_provenance(self):
        rows = [
            {"mapped_block": "Block A", "source_req_id": "SELECT", "requirement_statement": "The user shall select the active channel, in this case CH_A_SEL."},
            {"mapped_block": "Block A", "source_req_id": "WRITE", "requirement_statement": "The controller shall write processed results to the output buffer."},
        ]
        sections, audit = compose_ipos_overview(build_ipos_functional_input(
            "Block A", {"Function": "Process sampled data", "Inputs": "configuration settings and sampled data", "Outputs": "processed results and output buffer"}, rows,
        ))
        summaries = sections["scope"][1:]
        self.assertTrue(all(summary.startswith("**") and summary.count("\n") <= 1 for summary in summaries))
        self.assertTrue(all("No direct dependency" not in summary for summary in summaries))
        accepted = [row for row in audit if row.get("decision") == "accepted_candidate"]
        self.assertTrue(all(row["rendered_summary"] in summaries for row in accepted))
        self.assertTrue(all(row["candidate_evidence_requirement_id"] in row["summary_group_contributor_ids"] for row in accepted))

    def test_workflow_rule_states_end_to_end_contract_without_project_examples(self):
        rule_path = Path(__file__).resolve().parents[1] / ".github/skills/workflow-stage-gate/SKILL.md"
        text = rule_path.read_text(encoding="utf-8")
        rule_lines = text.splitlines()
        contract_start = next(index for index, line in enumerate(rule_lines) if "materialized IPOS Block Overviews" in line)
        contract = " ".join(" ".join(rule_lines[contract_start:contract_start + 16]).casefold().split())
        for required in (
            "approved snapshot", "function`, `inputs`, and `outputs", "same-block approved ipos requirement evidence",
            "acronyms alone", "signal lists alone", "i/o names alone", "tables alone", "ordered semantic units",
            "non-normative", "ownership", "responsibility", "system-level behavior", "suppressed candidate",
            "preserves its supported action and object", "generic action-family placeholder",
        ):
            self.assertIn(required, contract)
        project_vocabulary = r"main controller|smart fifo|\becg\b|\bbia\b|\bppg\b|\bgsr\b"
        self.assertNotRegex(text.casefold(), project_vocabulary)
        rule_start = next(index for index, line in enumerate(rule_lines) if "Shared IPOS refinement-audit input contract" in line)
        rule_end = next((index for index in range(rule_start + 1, len(rule_lines)) if rule_lines[index].startswith("- ")), len(rule_lines))
        shared_audit_rule = " ".join(rule_lines[rule_start:rule_end]).casefold()
        self.assertIn("identical complete candidate-requirement set", shared_audit_rule)
        self.assertIn("parity tests for both ipos domains", shared_audit_rule)

        orchestrator_path = Path(__file__).resolve().parents[1] / ".github/agents/workflow-orchestrator.agent.md"
        orchestrator = orchestrator_path.read_text(encoding="utf-8").casefold()
        self.assertNotRegex(orchestrator, project_vocabulary)
        self.assertIn("identical complete snapshot-scoped candidate-requirement set", orchestrator)
        self.assertIn("cross-block candidates are audit-only suppressed findings", orchestrator)

    def test_valid_empty_ipos_input_does_not_force_sections(self):
        record = build_ipos_functional_input(
            "Empty", {"Function": "", "Inputs": "", "Outputs": ""}, materialized=False
        )
        sections, audit = compose_ipos_overview(record)
        self.assertEqual(sections["functionality"], [])
        self.assertEqual(sections["scope"], [])
        self.assertEqual(audit, [])

    def test_ipos_scope_validator_rejects_content_outside_evidence_projection(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            markdown_path = root / "ipos.md"
            audit_path = root / "descriptive_summary_audit.csv"
            markdown_path.write_text(
                "## 1. Block overview\n\n"
                "### 1.2 Supported functions and scope\n\n"
                "- **SPI handling:** Supports serial-interface transactions.\n\n"
                "## 2. Source I/O\n",
                encoding="utf-8",
            )
            empty_assembly = assemble_descriptive_summary(
                [], DESCRIPTIVE_DIGITAL_TOPIC_SPECS, include_general=False,
                summary_only=True, ipos_topic_mode=True,
            )
            write_descriptive_audit(audit_path, empty_assembly, section="IPOS overview")
            findings = validate_ipos_descriptive_output(
                markdown_path, audit_path, "Control Block",
                ["The block controls I2C transactions."], [],
                DESCRIPTIVE_DIGITAL_TOPIC_SPECS,
            )
        self.assertTrue(any("does not match the deterministic projection" in finding for finding in findings))

    def test_normative_evidence_is_retained_as_descriptive_provenance(self):
        assembly = assemble_descriptive_summary(
            [{"statement": "The controller shall enter standby state.", "source": "approved:req-1"}],
            DESCRIPTIVE_MODE_TOPIC_SPECS,
        )
        self.assertIn("Modes and transitions", assembly.topics)
        self.assertNotIn("approved:req-1", assembly.topics["Modes and transitions"][0])
        self.assertNotIn("(source=", assembly.topics["Modes and transitions"][0])
        self.assertEqual(assembly.audit[0]["normative_evidence"], "yes")
        self.assertEqual(assembly.audit[0]["decision"], "selected")
        self.assertNotIn("Approved evidence:", assembly.topics["Modes and transitions"][0])
        self.assertNotIn("(source=", assembly.topics["Modes and transitions"][0])
        self.assertEqual(assembly.audit[0]["source"], "approved:req-1")

    def test_duplicate_and_limit_decisions_are_audited(self):
        records = [
            {"statement": "Standby mode uses a clock.", "source": "a"},
            {"statement": "Standby mode uses a clock.", "source": "a"},
            {"statement": "Sleep mode uses a clock.", "source": "b"},
        ]
        assembly = assemble_descriptive_summary(
            records, DESCRIPTIVE_MODE_TOPIC_SPECS, max_items_per_topic=1
        )
        reasons = [row["reason"] for row in assembly.audit]
        self.assertIn("duplicate descriptive evidence", reasons)
        self.assertIn("topic output limit", reasons)
        self.assertEqual(len(assembly.topics["Modes and transitions"]), 1)

    def test_empty_topics_are_omitted_by_default(self):
        assembly = assemble_descriptive_summary(
            [{"statement": "A calibration reference exists.", "source": "approved:context"}],
            DESCRIPTIVE_POWER_TOPIC_SPECS,
            include_general=False,
        )
        self.assertNotIn("Power states and sequencing", assembly.topics)
        self.assertEqual(assembly.topics, {})

    def test_power_domain_categories_use_approved_descriptive_signals(self):
        assembly = assemble_descriptive_summary(
            [
                {"statement": "The retention domain preserves register state when the main logic is powered down.", "source": "approved:power-1"},
                {"statement": "The power domain is enabled after the supply rail is stable.", "source": "approved:power-2"},
            ],
            DESCRIPTIVE_POWER_TOPIC_SPECS,
            include_general=False,
        )
        self.assertIn("Supply-Domain Relationships", assembly.topics)
        self.assertIn("Retention, Standby, Sleep, and Active Behavior", assembly.topics)
        self.assertTrue(all("Approved evidence:" not in item for items in assembly.topics.values() for item in items))
        self.assertEqual({row["source"] for row in assembly.audit if row["decision"] == "selected"}, {"approved:power-1", "approved:power-2"})

    def test_audit_is_reproducible_and_reviewable(self):
        records = [
            {"statement": "The power domain supports retention.", "source": "approved:req-2"},
            {"statement": "The block contains shared memory.", "source": "approved:block-1"},
        ]
        first = assemble_descriptive_summary(records, DESCRIPTIVE_POWER_TOPIC_SPECS)
        second = assemble_descriptive_summary(records, DESCRIPTIVE_POWER_TOPIC_SPECS)
        self.assertEqual(first, second)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "audit.csv"
            write_descriptive_audit(path, first, section="Power States")
            with path.open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["section"], "Power States")
            self.assertIn("assignment_reason", rows[0])
            self.assertIn("output_order", rows[0])

    def test_ipos_topic_assignment_uses_functional_evidence_not_provenance(self):
        assignment = assign_ipos_descriptive_topic(
            "The controller sequences clock reset release and power-island enable.",
            DESCRIPTIVE_DIGITAL_TOPIC_SPECS,
        )
        self.assertEqual(assignment["topic"], "Power and clock islands")
        metadata_only = assemble_descriptive_summary(
            [{
                "summary": "Metadata-only entry",
                "statement": "Metadata-only entry",
                "source": "power clock island provenance",
                "functional_evidence": "",
            }],
            DESCRIPTIVE_DIGITAL_TOPIC_SPECS,
            include_general=False,
            summary_only=True,
            ipos_topic_mode=True,
        )
        self.assertEqual(metadata_only.topics, {})
        self.assertEqual(metadata_only.audit[0]["reason"], "no functional topic evidence")

    def test_ipos_topic_assignment_omits_ambiguous_functional_evidence(self):
        assembly = assemble_descriptive_summary(
            [{
                "summary": "Ambiguous entry",
                "statement": "Ambiguous entry",
                "source": "approved provenance",
                "functional_evidence": "clock reset",
            }],
            (("Clock", ("clock",)), ("Reset", ("reset",))),
            include_general=False,
            summary_only=True,
            ipos_topic_mode=True,
        )
        self.assertEqual(assembly.topics, {})
        self.assertEqual(assembly.audit[0]["reason"], "ambiguous functional topic evidence")

    def test_low_power_assembly_is_scope_filtered_and_config_driven(self):
        repo_root = Path(__file__).resolve().parents[1]
        grouped, audit = assemble_low_power_descriptive(
            [
                {
                    "statement": "The retention domain preserves register state when the main logic is powered down.",
                    "source": "approved:architecture",
                    "scope": "architecture",
                    "domain": "system",
                },
                {
                    "statement": "Ultra-low-power hardware is a state-of-the-art feature.",
                    "source": "approved:marketing",
                    "scope": "architecture",
                    "domain": "system",
                },
                {
                    "statement": "The analog bias domain enters sleep and restores on wake-up.",
                    "source": "approved:analog",
                    "scope": "architecture",
                    "domain": "analog",
                },
            ],
            repo_root=repo_root,
            profile="srs",
        )
        self.assertIn("Power-domain architecture", grouped)
        self.assertNotIn("Ultra-low-power hardware is a state-of-the-art feature.", grouped["Power-domain architecture"])
        self.assertEqual(
            next(row["reason"] for row in audit if row["source"] == "approved:marketing"),
            "excluded phrase: state-of-the-art feature",
        )
        self.assertEqual(
            next(row["reason"] for row in audit if row["source"] == "approved:analog"),
            "eligible descriptive evidence",
        )

    def test_low_power_domain_scope_excludes_block_local_drs_evidence(self):
        repo_root = Path(__file__).resolve().parents[1]
        grouped, audit = assemble_low_power_descriptive(
            [{"statement": "The retention domain preserves state.", "source": "block-local", "scope": "block_local", "domain": "digital"}],
            repo_root=repo_root,
            profile="drs",
        )
        self.assertEqual(grouped, {})
        self.assertEqual(audit[0]["reason"], "scope not allowed")

    def test_drs_top_level_predicate_rejects_block_local_evidence(self):
        self.assertFalse(is_drs_top_level_record({"scope": "block_local", "domain": "digital"}))
        self.assertTrue(is_drs_top_level_record({"scope": "integration", "domain": "digital"}))

    def test_function_rows_preserve_system_scope_and_reject_block_scope(self):
        records = architecture_records_from_function_rows([
            (
                "The retention domain preserves register state when powered down.",
                "architecture_function; source=Section 10.2, paragraph 011 (page 32)",
                "System",
            ),
            (
                "The local block enters retention mode.",
                "architecture_function; source=Section 17.2, paragraph 052 (page 157)",
                "BIST Controller",
            ),
        ])
        self.assertEqual(records[0]["scope"], "architecture")
        self.assertEqual(records[0]["source"], "Section 10.2, paragraph 011 (page 32)")
        self.assertEqual(records[1]["scope"], "block_local")

    def test_selected_low_power_audit_is_the_only_downstream_handoff(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "descriptive_low_power_audit.csv"
            write_low_power_audit(path, [
                {
                    "decision": "selected",
                    "statement": "The power domain can be powered independently.",
                    "source": "Section 10.2",
                    "scope": "architecture",
                    "domain": "system",
                },
                {
                    "decision": "rejected",
                    "statement": "The local block enters retention.",
                    "source": "Section 17.2",
                    "scope": "block_local",
                    "domain": "BIST Controller",
                },
            ])
            records = read_selected_low_power_audit(path)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["scope"], "architecture")
        self.assertEqual(records[0]["source"], "Section 10.2")

    def test_power_domain_evidence_preserves_names_attributes_and_provenance(self):
        records = power_domain_records_from_ocr(Path(__file__).resolve().parents[1])
        combined = "\n".join(row["statement"] for row in records)
        self.assertIn("PD_TOP1V2", combined)
        self.assertIn("PD_STREDL", combined)
        self.assertIn("PD_TOP3V3", combined)
        self.assertIn("OTP memory block", combined)
        self.assertNotIn("PUMP_85AL05_2: Included block(s)", combined)
        self.assertNotIn("ca_fifo_bist_c", combined)
        self.assertNotIn("ca_qk_bist_c", combined)
        self.assertNotIn("ca_qk_rom_bist", combined)
        self.assertNotIn("ta_afe_bist_ctrl", combined)
        self.assertNotIn("u_ca_fifo_bist_c", combined)
        self.assertIn("Switchable", combined)
        self.assertIn("Never powered down", combined)
        self.assertTrue(all("Stage 1 OCR power-domain" in row["source"] for row in records))

    def test_version_history_and_signal_level_rows_are_rejected(self):
        grouped, audit = assemble_low_power_descriptive(
            [
                {"statement": "Version History: added Power Domain.", "source": "Section 0.1 Version History", "scope": "architecture", "domain": "system"},
                {"statement": "ca_stredl_bist_c_sif_reg_e enables a local test path.", "source": "Section 18", "scope": "architecture", "domain": "system"},
            ],
            repo_root=Path(__file__).resolve().parents[1],
            profile="srs",
        )
        self.assertEqual(grouped, {})
        self.assertIn("excluded phrase: version history", [row["reason"] for row in audit])
        self.assertIn("signal-level detail not eligible for top-level descriptive text", [row["reason"] for row in audit])

    def test_srs_low_power_renderer_has_one_domain_section_and_clarifies_missing_sequence(self):
        lines = []
        _render_srs_low_power_section(
            lines,
            [{"statement": "PD_STREDL: Domain type: Switchable. Control mode: SW.", "source": "approved", "scope": "architecture", "domain": "power", "evidence_kind": "power_domain"}],
            repo_root=Path(__file__).resolve().parents[1],
        )
        rendered = "\n".join(lines)
        self.assertEqual(rendered.count("#### Power-domain architecture"), 1)
        self.assertIn("Entry / power-up:", rendered)
        self.assertIn("Wake-up / restore:", rendered)
        self.assertIn("no retention, state-restore, or wake-up timing is specified", rendered)


if __name__ == "__main__":
    unittest.main()
