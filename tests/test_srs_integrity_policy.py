import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from canonical_store import connect
from run_srs_crosscheck_agent import _validate_srs_allocation_scope
from run_srs_gen_spec_agent import (
    Requirement,
    _block_navigation_target,
    _build_heading_registry,
    _deduplicate_authored_requirement_blocks,
    _is_srs_allocated,
    _read_snapshot_srs_allocations,
    _rewrite_internal_links,
    _terminate_authored_requirement_blocks,
    _write_srs_markdown,
)
from allocation_ledger import build_source_ledger
from validate_downstream_coherence import DownstreamContract, srs_catalog_from_context
from workflow_routing import (
    SRS_SYSTEM_OVERVIEW_HEADINGS,
    SRS_INTRODUCTORY_SECTIONS,
    SRS_INTRODUCTORY_WRITING_RULE,
    SRS_INTRODUCTORY_WRITING_RULE_ID,
    SRS_DOCUMENT_CONTENT_PROFILE,
    compose_srs_system_overview,
    compose_srs_support_content,
    validate_srs_support_content,
    validate_srs_system_overview,
)


class SrsIntegrityPolicyTests(unittest.TestCase):
    def test_srs_catalog_uses_explicit_context_without_requirement_count(self):
        contract = object.__new__(DownstreamContract)
        contract.snapshot_id = "synthetic"
        contract.expected = {"source": {"approved_block": "Conflict", "owning_domain": "analog"}}
        context = {"fingerprint": "profile", "payload": {"block_defs": {
            "Unused": {"classification": "Digital", "function": "Buffer measurements."},
            "Unknown": {"function": "Route measurements."},
            "Conflict": {"classification": "Digital", "function": "Convert measurements."},
        }}}
        entries, findings = srs_catalog_from_context(
            [{"Block": name} for name in ("Unused", "Unknown", "Conflict")], contract, context)
        self.assertEqual([entry["block"] for entry in entries], ["Unused"])
        self.assertIn("SRS_CATALOG_CLASSIFICATION_MISSING: Unknown", findings)
        self.assertIn("SRS_CATALOG_CLASSIFICATION_CONFLICT: Conflict", findings)

    def test_srs_overview_validator_rejects_empty_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            markdown_path = Path(directory) / "srs.md"
            sections = "\n\n".join(f"### {heading}\n" for heading in SRS_SYSTEM_OVERVIEW_HEADINGS)
            markdown_path.write_text(
                f"## 3. System Overview\n\n{sections}\n## 4. Detail\n",
                encoding="utf-8",
            )

            findings = validate_srs_system_overview(markdown_path)

        self.assertTrue(any("section_empty" in finding for finding in findings))

    def test_srs_overview_projects_only_scoped_system_level_prose(self):
        sections = compose_srs_system_overview(
            [
                {
                    "statement": "DeviceZ is an analog front end with embedded processing capabilities.",
                    "source": "Stage 1 OCR overview",
                    "scope": "system",
                    "evidence_kind": "descriptive",
                },
                {
                    "statement": "The ECG signal chain is designed for ECG measurement, such as lead detection.",
                    "source": "Stage 1 OCR overview",
                    "scope": "system",
                    "evidence_kind": "descriptive",
                },
                {
                    "statement": "The BIA signal chain is designed for body impedance and human breathing measurements with a configurable excitation path. BIA channel delivers real and imaginary impedance values.",
                    "source": "Stage 1 OCR overview",
                    "scope": "system",
                    "evidence_kind": "descriptive",
                },
                {
                    "statement": "The ECG signal chain has complementary features supporting ECG measurement, such as driven reference and lead-off detection.",
                    "source": "Stage 1 OCR overview",
                    "scope": "system",
                    "evidence_kind": "descriptive",
                },
                {
                    "statement": "A clock gating cell is forced transparent in ATPG mode.",
                    "source": "Stage 1 OCR architecture paragraph",
                    "scope": "architecture",
                    "evidence_kind": "descriptive",
                },
                {
                    "statement": "Sensor-Hub supports external sensor acquisition.",
                    "source": "Stage 1 OCR overview",
                    "scope": "architecture",
                    "evidence_kind": "descriptive",
                },
            ],
            mode_rows=[{"mode": "Normal Mode", "evidence": "approved source mode"}],
            interface_rows=[
                {"Interface": "SENSOR_ANALOG_INPUT", "Direction": "input", "Type": "analog", "Owner": "Sensor-Hub", "Purpose": "Raw sensor-domain analog input"},
                {"Interface": "SERIAL_CLK", "Direction": "input", "Type": "digital", "Owner": "SPI interface", "Purpose": "Serial clock input"},
                {"Interface": "SERIAL_DATA_OUT", "Direction": "output", "Type": "digital", "Owner": "SPI interface", "Purpose": "Serial data output"},
                {"Interface": "INT1", "Direction": "output", "Type": "digital", "Owner": "IRQ logic", "Purpose": "Primary interrupt output"},
                {"Interface": "VDD", "Direction": "input", "Type": "power", "Owner": "PMU", "Purpose": "Core supply"},
            ],
            power_records=[
                {
                    "statement": "Power domains can be powered independently; an always-on domain remains active.",
                    "source": "Stage 1 OCR power-domain description",
                    "scope": "architecture",
                    "evidence_kind": "power_domain",
                },
                {
                    "statement": "A retention domain preserves state when the main logic is powered down.",
                    "source": "Stage 1 OCR power-domain description",
                    "scope": "architecture",
                    "evidence_kind": "power_domain",
                },
            ],
            block_names=("Sensor-Hub",),
        )

        rendered = "\n".join(item for values in sections.values() for item in values)
        self.assertIn("The system is an analog front end with embedded processing capabilities.", rendered)
        self.assertIn("The system supports ECG measurement.", rendered)
        self.assertIn("The system supports body impedance and human breathing measurements with a configurable excitation path.", rendered)
        self.assertIn("with a configurable excitation path", rendered)
        self.assertNotIn("lead-off detection", rendered)
        self.assertNotIn("clock gating cell", rendered)
        self.assertNotIn("Sensor-Hub", rendered)
        self.assertNotIn("SERIAL_CLK", rendered)
        self.assertIn("Normal Mode", rendered)
        self.assertIn("system-level role and transitions are not specified", rendered)
        self.assertIn("sensing input, serial communication, interrupt reporting and power supply roles", rendered)
        self.assertIn("| Boundary role | Direction | Medium |", rendered)
        self.assertIn("Retention behavior", rendered)
        self.assertNotIn("Switchable behavior", rendered)
        self.assertIn("Top-level clock and reset coordination are not specified.", rendered)
        self.assertNotIn("SENSOR_ANALOG_INPUT", rendered)
        self.assertNotIn("Sensor-Hub", rendered)

    def test_srs_overview_interface_roles_require_described_purpose(self):
        sections = compose_srs_system_overview([], interface_rows=[
            {"Interface": "UNKNOWN_PIN", "Owner": "SPI interface", "Purpose": "Unspecified role"},
            {"Interface": "VDD", "Type": "power", "Purpose": "Core supply"},
        ])
        rendered = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[3]])
        self.assertIn("power supply roles", rendered)
        self.assertNotIn("Serial communication", rendered)
        self.assertNotIn("interrupt reporting", rendered)
        self.assertNotIn("UNKNOWN_PIN", rendered)

    def test_srs_overview_recovers_mode_function_without_local_detail(self):
        records = [{
            "statement": "Acquisition unit samples ECG signals from ADC, average and stores sampled data in FIFO.",
            "source": "Section 8.1 Data Storage Mode (under Section 8 Functional Modes), paragraph 008",
            "scope": "architecture",
        }]
        audit = []
        sections = compose_srs_system_overview(records, block_names=("Acquisition unit",), audit_rows=audit)
        mode_content = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[4]])
        self.assertIn("**Data Storage Mode**", mode_content)
        self.assertIn("acquires ECG measurements, averages the samples and stores the sampled data", mode_content)
        self.assertNotIn("ADC", mode_content)
        self.assertNotIn("FIFO", mode_content)
        self.assertNotIn("Acquisition unit", mode_content)
        contributors = [row for row in audit if row["section"] == SRS_SYSTEM_OVERVIEW_HEADINGS[4]]
        self.assertEqual(len(contributors), 1)
        self.assertEqual(contributors[0]["source"], records[0]["source"])
        self.assertEqual(contributors[0]["statement"], records[0]["statement"])
        records[0]["scope"] = "block_local"
        self.assertNotIn("acquires ECG", str(compose_srs_system_overview(records)))

    def test_srs_overview_power_glossary_is_not_design_authority(self):
        sections = compose_srs_system_overview([], power_records=[{
            "statement": "Retention Domain: Domain that retains register state even when the main logic is powered down.",
            "source": "Section 10.2 Definitions",
            "scope": "architecture",
        }])
        self.assertNotIn("Retention behavior", str(sections))

    def test_srs_overview_recovers_descriptive_clause_beside_normative_detail(self):
        records = [{
            "statement": "Set a control field if data shall be elaborated. • Acquisition unit samples of ECG, BIA or GSR from ADC, averages and stores sampled data in registers.",
            "source": "Section 8.2 Normal Mode (under Section 8 Functional Modes)",
            "scope": "architecture",
        }, {
            "statement": "ECG, BIA or GSR data raw in SYSTEM REGISTERS are processed by ProcessorZ.",
            "source": "Section 8.2 Normal Mode (under Section 8 Functional Modes)",
            "scope": "architecture",
        }]
        sections = compose_srs_system_overview(records, block_names=("Acquisition unit", "ProcessorZ"))
        rendered = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[4]])
        self.assertIn("**Normal Mode**", rendered)
        self.assertIn("acquires ECG, BIA or GSR measurements", rendered)
        self.assertIn("processes ECG, BIA or GSR measurements", rendered)
        self.assertNotIn("shall", rendered)
        self.assertNotIn("register", rendered.casefold())
        self.assertNotIn("ProcessorZ", rendered)

    def test_srs_overview_raw_storage_is_not_merged_with_processed_storage(self):
        records = [{"statement": "Collector samples pressure signals from ADC and stores raw data in FIFO.",
                    "source": "Section 8.1 Storage Mode, paragraph 1", "scope": "architecture"},
                   {"statement": "Collector samples temperature signals from ADC, averages and stores data in FIFO.",
                    "source": "Section 8.1 Storage Mode, paragraph 2", "scope": "architecture"}]
        sections = compose_srs_system_overview(records)
        content = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[4]])
        self.assertIn("acquires pressure measurements and stores the raw sampled data", content)
        self.assertIn("acquires temperature measurements, averages the samples and stores the sampled data", content)
        self.assertNotIn("pressure and temperature measurements", content)

    def test_srs_overview_fact_aggregation_preserves_all_source_contributors(self):
        records = [{"statement": f"Collector samples {value} signals from ADC and stores data in FIFO.",
                    "source": f"Section 8.1 Storage Mode, paragraph {index}", "scope": "architecture"}
                   for index, value in enumerate(("temperature", "humidity", "pressure"))]
        audit = []
        sections = compose_srs_system_overview(records, audit_rows=audit)
        rendered = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[4]])
        self.assertIn("temperature, humidity and pressure measurements", rendered)
        contributors = [row for row in audit if row["fact_id"] and row["decision"] == "selected"]
        self.assertEqual(len(contributors), 3)
        self.assertEqual({row["source"] for row in contributors}, {row["source"] for row in records})
        self.assertTrue(all(len(json.loads(row["contributor_ids"])) == 3 for row in contributors))
        self.assertEqual(sections, compose_srs_system_overview(records))

    def test_srs_overview_sensor_collection_processing_and_impedance_results_are_retained(self):
        records = [{
            "statement": "Sensor collector can read external sensors by means of an SPI master protocol, enabling data collection in FIFO and embedded elaboration also on external data domain. The data collected by sensor collector is accessible from ProcessorA, for sensor fusion algorithm elaboration.",
            "source": "System description, paragraph 1", "scope": "system",
        }, {
            "statement": "The impedance signal chain is designed for impedance measurement with a configurable path. Impedance channel delivers both the real and the imaginary parts of the measured impedance.",
            "source": "Architecture description, paragraph 2", "scope": "architecture",
        }]
        audit = []
        sections = compose_srs_system_overview(records, block_names=("ProcessorA",), audit_rows=audit)
        content = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[1]])
        self.assertIn("can acquire external sensor measurements through SPI master protocol with data collection and embedded processing", content)
        self.assertIn("provides collected external sensor measurements for sensor-fusion processing", content)
        self.assertIn("provides the real and the imaginary parts of the measured impedance", content)
        self.assertIn("supports impedance measurement with a configurable path", content)
        self.assertNotIn("ProcessorA", content)
        self.assertTrue(all(row["source"] in {record["source"] for record in records} for row in audit if row["fact_id"]))

    def test_srs_overview_fact_qualifiers_do_not_become_unconditional_capabilities(self):
        records = [{"statement": "When enabled, The system may acquire temperature measurements.",
                    "source": "Functional description", "scope": "system"},
                   {"statement": "Collector never samples pressure signals from ADC and stores data in FIFO.",
                    "source": "Section 8.1 Storage Mode", "scope": "architecture"}]
        audit = []
        sections = compose_srs_system_overview(records, audit_rows=audit)
        rendered = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[1]])
        self.assertIn("When enabled, the system may acquire temperature measurements.", rendered)
        self.assertNotIn("acquires pressure", str(sections))
        self.assertTrue(any("projection_gap" in row["reason"] for row in audit))

    def test_srs_overview_persisted_audit_detects_changed_documents_and_contributors(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "srs.md"
            audit = []
            sections = compose_srs_system_overview([{
                "statement": "When enabled, The system may acquire temperature measurements according to the user's configuration.",
                "source": "Functional description", "scope": "system",
            }], audit_rows=audit)
            text = "## 3. System Overview\n\n" + "\n\n".join(
                f"### {heading}\n\n" + "\n".join(sections[heading]) for heading in SRS_SYSTEM_OVERVIEW_HEADINGS
            ) + "\n\n## 4. Detail\n"
            output.write_text(text, encoding="utf-8")
            audit_path = root / "descriptive_system_overview_audit.csv"
            def write_audit():
                with audit_path.open("w", encoding="utf-8", newline="") as handle:
                    writer = csv.DictWriter(handle, fieldnames=list(audit[0]))
                    writer.writeheader()
                    writer.writerows(audit)
            write_audit()
            self.assertEqual(validate_srs_system_overview(output), [])
            output.write_text(text.replace("may acquire", "acquires"), encoding="utf-8")
            self.assertTrue(any("missing_from_markdown" in finding for finding in validate_srs_system_overview(output)))
            output.write_text(text, encoding="utf-8")
            selected = next(row for row in audit if row["fact_id"])
            selected["contributor_ids"] = "[]"
            write_audit()
            self.assertTrue(any("contributors_invalid" in finding for finding in validate_srs_system_overview(output)))
            selected["contributor_ids"] = json.dumps([selected["fact_id"]])
            write_audit()
            fingerprint = selected["profile_fingerprint"]
            selected["profile_fingerprint"] = "changed-profile"
            write_audit()
            self.assertTrue(any("fact_invalid" in finding for finding in validate_srs_system_overview(output)))
            selected["profile_fingerprint"] = fingerprint
            write_audit()
            document = Document()
            document.add_paragraph("The system acquires humidity measurements.")
            document.save(output.with_suffix(".docx"))
            self.assertTrue(any("missing_from_docx" in finding for finding in validate_srs_system_overview(output)))
            document = Document()
            for row in audit:
                if row["fact_id"]:
                    document.add_paragraph(row["rendered_summary"].replace("'", "\u2019"))
            document.save(output.with_suffix(".docx"))
            self.assertEqual(validate_srs_system_overview(output), [])

    def test_srs_overview_full_renderer_has_unique_headings_and_audit(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "srs.md"
            template = root / "template.md"
            template.write_text(
                "## SRS Document Format\n## 1. Introduction\n### 1.1 Purpose\n### 1.2 Scope\nCover template directives.\n"
                "### 1.3 Intended audience\n### 1.4 References\n## 2. Definitions and terminology\n"
                "### 2.1 System terminology\n### 2.2 Analog terminology\n### 2.3 Digital terminology\n"
                "### 2.4 Measurement and acceptance terms\n### 2.5 Category Convention\n## 3. System Overview\n" +
                "\n".join(f"### {heading}" for heading in SRS_SYSTEM_OVERVIEW_HEADINGS) +
                "\n## 4. Analog sub-system\n### 4.1 Analog block list\n## 5. Digital sub-system\n### 5.1 Digital block list\n## 6. Cross-domain requirements\n### 6.4 Power sequencing across domains\n## Required CSV Output Format\n",
                encoding="utf-8",
            )
            entries = [
                {"block": "Converter", "category": "analog", "function": "Convert measurements without averaging.",
                 "source": "snapshot architecture profile:synthetic", "classification_sources": "[]", "snapshot_id": "synthetic-snapshot"},
                {"block": "Unused buffer", "category": "digital", "function": "Store measurements only during acquisition.",
                 "source": "snapshot architecture profile:synthetic", "classification_sources": "[]", "snapshot_id": "synthetic-snapshot"},
            ]
            power = [{"statement": "PD_A: Included block(s): converter_core. Domain type: Always-On. Voltage: 3A. Control mode: HW. Notes: Active after 20 ms. Function: Supplies the converter. Characteristics: Never powered down.",
                      "source": "Architecture power description", "scope": "architecture", "evidence_kind": "power_domain"}]
            review = ["SRS_CATALOG_CLASSIFICATION_MISSING: Unresolved"]
            with patch("run_srs_gen_spec_agent._structured_srs_low_power_records", return_value=power):
                _write_srs_markdown(
                    output_path=output, template_path=template, source_spec="source.pdf",
                    project_name="Synthetic project", document_author="Test author",
                    requirements=[], srs_id_map={}, req_to_block={}, block_inventory_rows=[],
                    interface_rows=[], source_port_rows=[], interaction_rows=[],
                    connection_matrix_srs_ids={}, function_rows=[], descriptive_records=[],
                    block_requirements={}, unmapped_requirements=[], user_specific_requirements=[],
                    non_block_context_rows=[], crosscheck_decision="PASS", missing_inputs=[],
                    mode_rows=[], approved_srs_source_ids=set(), snapshot_id="synthetic-snapshot",
                    contract_fingerprint="synthetic-fingerprint", document_version="1.0",
                    catalog_entries=entries, catalog_findings=review,
                )
            self.assertEqual(validate_srs_system_overview(output), [])
            self.assertTrue(output.with_name("descriptive_system_overview_audit.csv").is_file())
            self.assertIn("This document presents the system-level requirements and supporting context for Synthetic project", output.read_text(encoding="utf-8"))
            original = output.read_text(encoding="utf-8")
            content_sections, _ = compose_srs_support_content("Synthetic project", "synthetic-snapshot", entries, review, power)
            audience_text = content_sections["1.3"][0]
            self.assertIn(audience_text, original)
            self.assertNotIn("Cover template directives", original)
            self.assertIn("This category denotes that the content of the object text is a general comment.", original)
            for section in SRS_INTRODUCTORY_SECTIONS:
                self.assertTrue(content_sections[section])
                for paragraph in content_sections[section]:
                    self.assertIn(paragraph, original)
            self.assertIn("- **Converter**", original)
            self.assertIn("- **Unused buffer**", original)
            self.assertIn("Active after 20 ms", original)
            self.assertNotIn("Supply voltage: 3A", original)
            self.assertNotIn("- Voltage: 3A", original)
            self.assertIn("[section 3.6](#36-power-clock-and-reset-overview)", original)
            self.assertIn("SRS_POWER_VOLTAGE_UNRELIABLE", original)
            for old, new, expected in (
                ("This document presents", "This document omits", "srs_purpose_missing"),
                (audience_text, "", "srs_intended_audience_missing"),
                (content_sections["1.2"][0], "", "srs_introductory_section_missing:1.2"),
                ("- **Unused buffer**", "- **Changed block**", "srs_catalog_coverage_invalid:5.1"),
                ("| PD_A | Always-On", "| PD_A | Switchable", "srs_domain_row_missing_or_changed"),
            ):
                output.write_text(original.replace(old, new), encoding="utf-8")
                self.assertIn(expected, validate_srs_support_content(output))
            output.write_text(original, encoding="utf-8")
            audit_path = output.with_name("descriptive_srs_content_audit.json")
            audit_original = audit_path.read_text(encoding="utf-8")
            audit = json.loads(audit_original)
            self.assertEqual(audit["writing_rule_id"], SRS_INTRODUCTORY_WRITING_RULE_ID)
            self.assertEqual(SRS_DOCUMENT_CONTENT_PROFILE.document_type, "SRS")
            self.assertIsInstance(SRS_INTRODUCTORY_WRITING_RULE, tuple)
            self.assertIn("scope is limited to system-level requirements allocated to this document", content_sections["1.2"][0])
            self.assertNotIn("This specification covers system-level functionality", original)
            audit["writing_rule_fingerprint"] = "changed-rule"
            audit_path.write_text(json.dumps(audit), encoding="utf-8")
            self.assertTrue(any("writing rule" in finding for finding in validate_srs_support_content(output)))
            audit_path.write_text(audit_original, encoding="utf-8")
            audit = json.loads(audit_original)
            audit["records"][1]["text"] = "Invented function."
            audit_path.write_text(json.dumps(audit), encoding="utf-8")
            self.assertTrue(any("provenance_invalid" in finding for finding in validate_srs_support_content(output)))
            audit_path.write_text(audit_original, encoding="utf-8")
            audit = json.loads(audit_original)
            audit["profile_fingerprint"] = "changed-profile"
            audit_path.write_text(json.dumps(audit), encoding="utf-8")
            self.assertTrue(any("audit_invalid" in finding for finding in validate_srs_support_content(output)))
            audit_path.write_text(audit_original, encoding="utf-8")
            from docx import Document
            sections, _ = compose_srs_support_content("Synthetic project", "synthetic-snapshot", entries, review, power)
            document = Document()
            for section in (*SRS_INTRODUCTORY_SECTIONS, "4.1", "5.1"):
                document.add_heading(section + " Heading", level=3)
                for unit in sections[section]:
                    document.add_paragraph(unit.split("\n\n", 1)[-1].strip())
            document.add_heading("3.6 Power", level=3)
            table = document.add_table(rows=2, cols=5)
            for cell, value in zip(table.rows[0].cells, ("Domain", "Type", "Control", "Functional Role", "Power Conditions")):
                cell.text = value
            for cell, value in zip(table.rows[1].cells, sections["3.6"][0].split("|")[1:-1]):
                cell.text = value.strip()
            document.save(output.with_suffix(".docx"))
            self.assertEqual(validate_srs_support_content(output), [])
            table.cell(1, 1).text = "Switchable"
            document.save(output.with_suffix(".docx"))
            self.assertIn("srs_domain_row_missing_from_docx", validate_srs_support_content(output))
            table.cell(1, 1).text = "Always-On"
            document.add_heading("6.4 Sequencing", level=3)
            document.element.body.append(table._tbl)
            document.save(output.with_suffix(".docx"))
            self.assertIn("srs_domain_row_missing_from_docx", validate_srs_support_content(output))

    def test_srs_overview_recovers_power_candidate_without_topic_output_cap(self):
        power_records = [
            {
                "statement": f"Descriptive power-domain relationship {index} is documented.",
                "source": f"Stage 1 OCR description {index}",
                "scope": "architecture",
                "evidence_kind": "power_domain",
            }
            for index in range(7)
        ]
        power_records.append({
            "statement": "A switchable domain can be powered down when its associated logic is not in use.",
            "source": "Stage 1 OCR power-domain description",
            "scope": "architecture",
            "evidence_kind": "power_domain",
        })
        sections = compose_srs_system_overview([], power_records=power_records)

        self.assertIn("Switchable behavior", "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[5]]))
        self.assertIn("can enter a powered-down state when its associated logic is not in use", str(sections))

    def test_srs_overview_power_retains_modality_and_does_not_invent_low_power_timing(self):
        sections = compose_srs_system_overview([], power_records=[{
            "statement": "A switchable domain may be powered down in idle when its associated logic is not in use.",
            "source": "Architecture power description", "scope": "architecture",
        }, {
            "statement": "An always-on domain is never powered down.",
            "source": "Architecture power description", "scope": "architecture",
        }])
        content = "\n".join(sections[SRS_SYSTEM_OVERVIEW_HEADINGS[5]])
        self.assertIn("may enter a powered-down state in idle when its associated logic is not in use", content)
        self.assertNotIn("during low-power operation", content)

    def test_srs_overview_validator_rejects_audit_and_raw_interface_content(self):
        with tempfile.TemporaryDirectory() as directory:
            markdown_path = Path(directory) / "srs.md"
            sections = []
            for heading in SRS_SYSTEM_OVERVIEW_HEADINGS:
                body = "System-level technical description."
                if heading.startswith("3.2"):
                    body = "Approved evidence identifies a capability."
                elif heading.startswith("3.4"):
                    body = "| Interface | Direction | Type | Owner |\n|---|---|---|---|\n| PIN | input | digital | block |"
                sections.append(f"### {heading}\n\n{body}")
            markdown_path.write_text(
                "## 3. System Overview\n\n" + "\n\n".join(sections) + "\n\n## 4. Detail\n",
                encoding="utf-8",
            )

            findings = validate_srs_system_overview(markdown_path)

        self.assertIn("srs_system_overview_contains_audit_or_filler_language", findings)
        self.assertIn("srs_system_overview_contains_raw_interface_dump", findings)

    def _allocation_db(self, root, rows):
        connection = connect(root)
        connection.execute(
            "INSERT INTO projects(project_id, project_name, created_at) VALUES ('P', 'P', '2026-01-01T00:00:00+00:00')"
        )
        for req_id, target in rows:
            connection.execute(
                """INSERT INTO requirement_allocations(
                    allocation_id, project_id, req_id, source_type, spec_level,
                    owning_target, requirement_class, lineage_mode, coverage_status,
                    snapshot_id, created_at)
                    VALUES (?, 'P', ?, 'primary', ?, ?, 'system_level',
                            'normal_hierarchical', 'covered', 'snap-1', ?)""",
                ("allocation-" + req_id, req_id, target, target, "2026-01-01T00:00:00+00:00"),
            )
        connection.commit()
        connection.close()

    def test_srs_placement_uses_allocation_target_not_source_id_shape(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self._allocation_db(root, [("DDS_STBIO1_0013", "SRS"), ("IPOS_MAIN_CTRL_0001", "DRS")])
            allocations = _read_snapshot_srs_allocations(
                root,
                project_id="P",
                snapshot_id="snap-1",
                requirement_ids={"DDS_STBIO1_0013", "IPOS_MAIN_CTRL_0001"},
            )
            primary = Requirement("DDS_STBIO1_0013", "SYS", "shall", "", "", "")
            supplementary = Requirement("IPOS_MAIN_CTRL_0001", "DIG", "shall", "", "", "")
            self.assertTrue(_is_srs_allocated(primary, allocations))
            self.assertFalse(_is_srs_allocated(supplementary, allocations))

    def test_missing_allocation_rows_fail_explicitly(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self._allocation_db(root, [("REQ-1", "SRS")])
            with self.assertRaisesRegex(RuntimeError, "do not match selected snapshot"):
                _read_snapshot_srs_allocations(
                    root,
                    project_id="P",
                    snapshot_id="snap-1",
                    requirement_ids={"REQ-1", "REQ-2"},
                )

    def test_zero_allocation_rows_fail_explicitly(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            connect(root).close()
            with self.assertRaisesRegex(RuntimeError, "non-empty allocation rows"):
                _read_snapshot_srs_allocations(
                    root,
                    project_id="P",
                    snapshot_id="snap-1",
                    requirement_ids={"REQ-1"},
                )

    def test_duplicate_authored_blocks_are_removed_but_distinct_ids_remain(self):
        lines = [
            "**[SRS-REQ-001] Requirement:**",
            "First.",
            "**[SRS-REQ-002] Requirement:**",
            "Second.",
            "**[SRS-REQ-001] Requirement:**",
            "Duplicate.",
        ]
        result = _deduplicate_authored_requirement_blocks(lines)
        self.assertEqual(result.count("**[SRS-REQ-001] Requirement:**"), 1)
        self.assertEqual(result.count("**[SRS-REQ-002] Requirement:**"), 1)
        self.assertNotIn("Duplicate.", result)

    def test_duplicate_blocks_are_removed_after_authored_normalization(self):
        normalized = _terminate_authored_requirement_blocks([
            "[SRS-REQ-001] Requirement:",
            "First.",
            "[SRS-REQ-001] Requirement:",
            "Duplicate.",
        ])
        result = _deduplicate_authored_requirement_blocks(normalized)
        self.assertEqual(result.count("**[SRS-REQ-001] Requirement:**"), 1)
        self.assertNotIn("Duplicate.", result)

    def test_duplicate_heading_titles_resolve_to_final_registry_anchor(self):
        registry = _build_heading_registry([
            "## First {#same}",
            "## Second {#same}",
        ])
        rewritten, resolved, non_links = _rewrite_internal_links(
            ["[first](#same)", "[missing](#does-not-exist)"], registry
        )
        self.assertEqual(registry[0][2], "same")
        self.assertEqual(registry[1][2], "same-1")
        self.assertEqual(rewritten[0], "[first](#same)")
        self.assertEqual(rewritten[1], "N/A")
        self.assertEqual((resolved, non_links), (1, 1))

    def test_final_registry_rewrite_removes_links_to_missing_rendered_headings(self):
        registry = _build_heading_registry(["## Rendered heading {#rendered-heading}"])
        rewritten, resolved, non_links = _rewrite_internal_links(
            ["[Rendered](#rendered-heading)", "[Not emitted](#not-emitted)"], registry
        )
        self.assertEqual(rewritten, ["[Rendered](#rendered-heading)", "N/A"])
        self.assertEqual((resolved, non_links), (1, 1))

    def test_disabled_block_navigation_is_plain_text(self):
        self.assertEqual(_block_navigation_target("91-sensor-hub", False), "N/A")
        self.assertEqual(_block_navigation_target("91-sensor-hub", True), "[Jump](#91-sensor-hub)")

    def test_stage3_partition_requires_srs_rows_once_but_allows_downstream_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self._allocation_db(root, [("REQ-S", "SRS"), ("REQ-D", "DRS"), ("REQ-A", "ARS")])
            manifest_path = root / "artifacts/traceability_reports/requirement_allocation_materialization.json"
            manifest_path.parent.mkdir(parents=True, exist_ok=True)
            manifest_path.write_text(json.dumps({"snapshot_id": "snap-1", "row_count": 3}), encoding="utf-8")
            findings = _validate_srs_allocation_scope(
                root,
                project_id="P",
                snapshot_id="snap-1",
                snapshot_ids={"REQ-S", "REQ-D", "REQ-A"},
                trace_source_ids=["REQ-S"],
            )
            self.assertEqual(findings, [])

    def test_stage3_partition_reports_missing_srs_coverage(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self._allocation_db(root, [("REQ-S", "SRS")])
            manifest_path = root / "artifacts/traceability_reports/requirement_allocation_materialization.json"
            manifest_path.parent.mkdir(parents=True, exist_ok=True)
            manifest_path.write_text(json.dumps({"snapshot_id": "snap-1", "row_count": 1}), encoding="utf-8")
            findings = _validate_srs_allocation_scope(
                root,
                project_id="P",
                snapshot_id="snap-1",
                snapshot_ids={"REQ-S"},
                trace_source_ids=[],
            )
            self.assertTrue(any("exactly once" in finding for finding in findings))

    def test_ledger_maps_canonical_snapshot_ids_to_source_ids(self):
        rows = build_source_ledger(
            [{
                "canonical_id": "can-1",
                "source_req_id": "DDS_STBIO1_0013",
                "requirement_statement": "The ADSP shall write the register.",
                "approved_classification": "Digital",
            }],
            [{"Requirement ID": "DDS_STBIO1_0013", "approved_block": "ADSP"}],
            snapshot_id="snap-1",
            allocation_rows=[{
                "source_req_id": "DDS_STBIO1_0013",
                "allocation_class": "block_local_digital",
                "owning_target": "Digital IPOS",
                "lineage_mode": "normal_hierarchical",
                "source_origin_req_ids": "DDS_STBIO1_0013",
            }],
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["owning_block"], "ADSP")
        self.assertEqual(rows[0]["owning_target"], "Digital IPOS")

    def test_block_owner_does_not_override_authoritative_srs_or_drs_target(self):
        source = [{
            "source_req_id": "REQ-S",
            "requirement_statement": "The system shall coordinate blocks.",
            "approved_classification": "System",
        }]
        mapping = [{"Requirement ID": "REQ-S", "approved_block": "Main Controller"}]
        for allocation_class, target in (("system_level", "SRS"), ("top_digital_architecture", "DRS")):
            rows = build_source_ledger(
                source,
                mapping,
                snapshot_id="snap-1",
                allocation_rows=[{
                    "source_req_id": "REQ-S",
                    "allocation_class": allocation_class,
                    "owning_target": target,
                    "lineage_mode": "normal_hierarchical",
                    "source_origin_req_ids": "REQ-S",
                }],
            )
            self.assertEqual(rows[0]["owning_target"], target)
            self.assertEqual(rows[0]["owning_block"], "Main Controller")

    def test_missing_authoritative_allocation_metadata_fails(self):
        with self.assertRaisesRegex(ValueError, "authoritative allocation metadata"):
            build_source_ledger(
                [{"source_req_id": "REQ-1", "requirement_statement": "shall", "approved_classification": "Digital"}],
                [{"Requirement ID": "REQ-1", "approved_block": "ADSP"}],
                snapshot_id="snap-1",
                allocation_rows=[],
            )


if __name__ == "__main__":
    unittest.main()
