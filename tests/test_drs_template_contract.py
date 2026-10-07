import copy
import sys
import tempfile
import unittest
from importlib import import_module
from pathlib import Path

import yaml
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
contract_module = import_module("spec_document_contract")
DRS_TEMPLATE = contract_module.DRS_TEMPLATE
drs_conventions_markdown = contract_module.drs_conventions_markdown
load_drs_contract = contract_module.load_drs_contract
drs_document_events = contract_module.drs_document_events
drs_table_inventory = contract_module.drs_table_inventory
validate_drs_events = contract_module.validate_drs_events
arrange_drs_contract_lines = contract_module.arrange_drs_contract_lines


class DrsTemplateContractTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(__file__).resolve().parents[1]
        self.contract = load_drs_contract(self.root / DRS_TEMPLATE)

    def test_load_and_render(self):
        self.assertEqual(self.contract["version"], 1)
        self.assertEqual(len(self.contract["sha256"]), 64)
        self.assertEqual(drs_conventions_markdown(self.contract)[0], "### 2.1 Conventions")
        self.assertIn("#### Requirement", drs_conventions_markdown(self.contract))

    def test_reject_missing_template(self):
        with self.assertRaises(FileNotFoundError):
            load_drs_contract(self.root / "nonexistent-template.md")

    def fixture_events(self):
        events = []
        for title in self.contract["navigation"]:
            events.append({"kind": "heading", "level": 4 if title.startswith("Table") else 3, "text": title})
            if title == "0.1 Table of contents":
                events.extend({"kind": "paragraph", "text": item} for item in self.contract["sections"] + ["2.1 Conventions"])
            if title.startswith("Table"):
                events.append({"kind": "table", "rows": [["Name", "Value"], ["Source", "retained"]]})
        events.extend([
            {"kind": "heading", "level": 4, "text": "Table 3. Section navigation index"},
            {"kind": "table", "rows": [["Section", "Paragraph"], ["2.1 Conventions", "2.1"]]},
        ])
        for title in self.contract["sections"]:
            events.append({"kind": "heading", "level": 2, "text": title})
            if title == "2. Definitions and terminology":
                events.append({"kind": "heading", "level": 3, "text": "2.1 Conventions"})
                for category in self.contract["conventions"]["categories"]:
                    events.extend([{"kind": "heading", "level": 4, "text": category["name"]}, {"kind": "paragraph", "text": category["definition"]}])
        return events

    def test_negative_final_content(self):
        original = self.fixture_events()
        baseline = drs_table_inventory(original)
        self.assertEqual(validate_drs_events(original, self.contract, baseline), [])
        for mutation in ("missing", "moved", "document_control", "duplicate", "definition", "category_order", "table", "section", "navigation", "toc", "counting", "snippet", "clarification_first"):
            with self.subTest(mutation=mutation):
                events = copy.deepcopy(original)
                start = next(index for index, event in enumerate(events) if event.get("text") == "2.1 Conventions" and event["kind"] == "heading")
                if mutation == "missing":
                    del events[start:start + 9]
                elif mutation == "moved":
                    events.extend(events[start:start + 9])
                    del events[start:start + 9]
                elif mutation == "document_control":
                    block = events[start:start + 9]
                    del events[start:start + 9]
                    insertion = next(index for index, event in enumerate(events) if event.get("text") == "Table 6. Category convention")
                    events[insertion:insertion] = block
                elif mutation == "duplicate":
                    events.extend(events[start:start + 9])
                elif mutation == "definition":
                    events[start + 2]["text"] = "Changed meaning."
                elif mutation == "category_order":
                    events[start + 1], events[start + 3] = events[start + 3], events[start + 1]
                elif mutation == "table":
                    next(event for event in events if event["kind"] == "table" and event["rows"][0] == ["Name", "Value"])["rows"].pop()
                elif mutation == "section":
                    events[-1], events[-2] = events[-2], events[-1]
                elif mutation == "navigation":
                    events.pop(0)
                elif mutation == "toc":
                    events.pop(next(index for index, event in enumerate(events) if event.get("text") == "2.1 Conventions" and event["kind"] == "paragraph"))
                else:
                    position = next(index for index, event in enumerate(events) if event.get("text") == "3. Top Level Overview" and event["kind"] == "heading") + 1
                    events[position:position] = ([{"kind": "paragraph", "text": "Approved integration evidence groups 12 paths."}] if mutation == "counting" else [{"kind": "paragraph", "text": "Module A -> Module B -> Module C"}] if mutation == "snippet" else [
                        {"kind": "paragraph", "text": "Transition ordering: need clarification."},
                        {"kind": "table", "rows": [["Source", "Destination"], ["Module A", "Module B"]]},
                    ])
                self.assertTrue(validate_drs_events(events, self.contract, baseline))

    def test_version_history_allows_append_but_rejects_loss_or_rewrite(self):
        original = self.fixture_events()
        baseline = drs_table_inventory(original)
        table_index = next(index + 1 for index, event in enumerate(original) if event.get("text") == "Table 1. Version history")
        appended = copy.deepcopy(original)
        appended[table_index]["rows"].extend([["Run", "new"], ["Run", "new"]])
        self.assertEqual(validate_drs_events(appended, self.contract, baseline), [])
        for mutation in ("remove", "rewrite", "reorder", "duplicate_table"):
            with self.subTest(mutation=mutation):
                events = copy.deepcopy(appended)
                rows = events[table_index]["rows"]
                if mutation == "remove":
                    rows.pop(1)
                elif mutation == "rewrite":
                    rows[1][1] = "changed"
                elif mutation == "reorder":
                    rows[1], rows[2] = rows[2], rows[1]
                else:
                    events.extend(copy.deepcopy(events[table_index - 1:table_index + 1]))
                self.assertTrue(any("version history" in finding for finding in validate_drs_events(events, self.contract, baseline)))

    def test_docx_only_definition_divergence(self):
        events = self.fixture_events()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "spec.docx"
            document = Document()
            for event in events:
                if event["kind"] == "heading":
                    document.add_heading(event["text"], level=event["level"])
                elif event["kind"] == "paragraph":
                    document.add_paragraph(event["text"])
                else:
                    table = document.add_table(rows=0, cols=len(event["rows"][0]))
                    for row in event["rows"]:
                        for cell, value in zip(table.add_row().cells, row):
                            cell.text = value
            document.save(path)
            baseline = drs_table_inventory(events)
            self.assertEqual(validate_drs_events(drs_document_events(path), self.contract, baseline), [])
            for paragraph in document.paragraphs:
                if paragraph.text == self.contract["conventions"]["categories"][3]["definition"]:
                    paragraph.text = "A requirement does not require verification."
            document.save(path)
            self.assertEqual(validate_drs_events(events, self.contract, baseline), [])
            findings = validate_drs_events(drs_document_events(path), self.contract, baseline)
            self.assertTrue(any("conventions" in finding for finding in findings))

    def test_arrangement_preserves_separate_tables(self):
        lines = []
        for title in self.contract["sections"]:
            lines.extend(["## " + title, ""])
            if title.startswith("4."):
                lines.extend(["### 4.7 Clock and synchronization across domains", "Clock drives the interface.", "Timing: need clarification.", "", "| Clock | Frequency |", "|---|---|", "| Shared | 20 MHz |", "", "| Focus | Reference |", "|---|---|", "| Timing | Source |", ""])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "spec.md"
            path.write_text("\n".join(arrange_drs_contract_lines(lines, self.contract)), encoding="utf-8")
            events = drs_document_events(path)
            self.assertEqual(len(drs_table_inventory(events)), 2)
            timing = [event for event in events if event.get("text") == "Timing: need clarification."][0]
            self.assertGreater(events.index(timing), max(index for index, event in enumerate(events) if event["kind"] == "table"))

    def test_internal_index_rejects_late_zero_number(self):
        events = self.fixture_events()
        table = next(event for event in events if event["kind"] == "table" and event["rows"][0] == ["Section", "Paragraph"])
        table["rows"] = [["Section", "Paragraph"], ["Introduction", "1"], ["2.1 Conventions", "0"]]
        self.assertTrue(any("numerical order" in finding for finding in validate_drs_events(events, self.contract, drs_table_inventory(events))))

    def test_reject_invalid_contract(self):
        for mutation in ("version", "old_placement", "unknown_rule", "missing_rule", "duplicate_key", "missing", "malformed"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as folder:
                payload = copy.deepcopy(self.contract)
                payload.pop("sha256")
                if mutation == "version":
                    payload["version"] = 999
                elif mutation == "old_placement":
                    payload["conventions"].update(heading="Conventions", after="Table 2. Reference documents", before="Table 6. Category convention")
                elif mutation == "unknown_rule":
                    payload["mandatory_rules"].append("unimplemented")
                elif mutation == "missing_rule":
                    payload["mandatory_rules"].pop()
                text = yaml.safe_dump(payload)
                if mutation == "duplicate_key":
                    text += "version: 1\n"
                elif mutation == "malformed":
                    text = "version: ["
                path = Path(folder) / "template.md"
                path.write_text("" if mutation == "missing" else "```yaml drs-document-contract\n" + text + "```\n", encoding="utf-8")
                with self.assertRaises((ValueError, yaml.YAMLError)):
                    load_drs_contract(path)


if __name__ == "__main__":
    unittest.main()