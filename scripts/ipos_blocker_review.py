"""Non-authoritative, block-local Stage 6 review workbook boundary."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation


REVIEW_COLUMNS = (
    "Reference req_id", "Block", "Section or claim", "Failure type",
    "Expected supported meaning", "Observed problem", "Likely failing layer",
    "Flow recommendation", "Why", "User decision", "Note", "Resulting action",
)
REVIEW_CHOICES = (
    "omit (as non-blocking)", "user clarified", "need source/spec update - clarification",
)
RECOMMENDATIONS = frozenset(REVIEW_CHOICES)
ACTIONS = frozenset((
    "re-evaluate after user note", "remain blocked pending future source/spec update",
    "generate supported content; retain source/spec update gap for readiness",
))


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _require_closed(path: Path) -> None:
    if (path.parent / ("~$" + path.name)).exists() or (path.parent / (".~lock." + path.name + "#")).exists():
        raise ValueError("Save and close the review workbook in Excel before re-evaluation")
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes

        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = (wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p,
                                      wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE)
        kernel.CreateFileW.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = (wintypes.HANDLE,)
        handle = kernel.CreateFileW(str(path), 0xC0000000, 0, None, 3, 0, None)
        if handle == ctypes.c_void_p(-1).value:
            raise ValueError("Save and close the review workbook; exclusive access is unavailable")
        kernel.CloseHandle(handle)
    else:
        with path.open("r+b"):
            pass


def review_paths(repo_root: Path, block: str) -> tuple[Path, Path, Path]:
    from generate_ipos_specs import _block_slug

    folder = repo_root / "artifacts/orchestrator/ipos_blocker_reviews" / _block_slug(block)
    return folder / "blocker_review.xlsx", folder / "blocker_review.json", folder / "blocker_review.md"


def _add_decision_dropdown(sheet) -> None:
    decision = DataValidation(type="list", formula1='"' + ",".join(REVIEW_CHOICES) + '"',
                              allow_blank=True, showDropDown=False, showErrorMessage=True,
                              errorStyle="stop", errorTitle="Choose a flow option",
                              error="Select one of the defined flow choices.")
    sheet.add_data_validation(decision)
    decision.add("J2:J" + str(sheet.max_row))
    for row in sheet.iter_rows(min_row=2):
        row[9].protection = Protection(locked=False)
        row[9].fill = PatternFill("solid", fgColor="E7F3D8")


def write_blocker_review(repo_root: Path, snapshot_id: str, block: str, blockers: list[dict]) -> Path:
    if not blockers:
        raise ValueError("A review workbook requires active blockers")
    if any(value.get("Block") != block for value in blockers):
        raise ValueError("A review workbook must contain exactly one IPOS block")
    rows = [[str(value.get(column, "")) for column in REVIEW_COLUMNS] for value in blockers]
    invalid_choices = any(row[7] not in RECOMMENDATIONS or row[11] not in ACTIONS or (row[9] and row[9] not in RECOMMENDATIONS) for row in rows)
    if invalid_choices:
        raise ValueError("Invalid review recommendation, decision or resulting action")
    workbook_path, manifest_path, companion_path = review_paths(repo_root, block)
    if manifest_path.exists():
        previous = json.loads(manifest_path.read_text(encoding="utf-8"))
        if previous["snapshot_id"] == snapshot_id and previous["rows"] == rows and workbook_path.exists():
            return upgrade_blocker_review_dropdown(repo_root, snapshot_id=snapshot_id, block=block)
        if workbook_path.exists():
            read_blocker_review(repo_root, snapshot_id=previous["snapshot_id"], block=block, require_saved=False)
            archive = workbook_path.with_name("blocker_review_" + _digest(workbook_path)[:16] + ".xlsx")
            if not archive.exists():
                archive.write_bytes(workbook_path.read_bytes())
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Active IPOS blockers"
    sheet.append(REVIEW_COLUMNS)
    for row in rows:
        sheet.append(row)
    widths = (24, 24, 36, 30, 48, 48, 30, 38, 48, 20, 48, 48)
    for position, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, position).column_letter].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A63")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.protection = Protection(locked=cell.column != 11)
        row[10].fill = PatternFill("solid", fgColor="E7F3D8")
        lines = max((len(str(cell.value or "")) + int(width) - 1) // int(width) for cell, width in zip(row, widths))
        sheet.row_dimensions[row[0].row].height = max(60, min(240, lines * 18 + 6))
    _add_decision_dropdown(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.protection.sheet = True
    workbook.save(workbook_path)
    workbook.close()
    manifest = {"snapshot_id": snapshot_id, "block": block, "rows": rows,
                "initial_sha256": _digest(workbook_path), "authority": False, "gate_override": False}
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    companion_path.write_text(
        "# " + block + " - BLOCKED\n\nActive blockers: " + str(len(rows))
        + "\n\nExcel: [Open review](" + workbook_path.name + ")\n\nPath: " + str(workbook_path.resolve())
        + "\n\nSelect User decision and edit Note; save and close Excel, then use Re-evaluate in the GUI. Review-only; no approval or gate override.\n",
        encoding="utf-8")
    return workbook_path


def read_blocker_review(repo_root: Path, *, snapshot_id: str, block: str, require_saved: bool = True, history_sha256: str = "") -> dict:
    workbook_path, manifest_path, _companion = review_paths(repo_root, block)
    if history_sha256:
        if not re.fullmatch(r"[0-9a-f]{64}", history_sha256):
            raise ValueError("Invalid archived review fingerprint")
        history = workbook_path.parent / "history" / history_sha256[:16]
        workbook_path = history / workbook_path.name
        manifest_path = history / manifest_path.name
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest["snapshot_id"] != snapshot_id or manifest["block"] != block:
        raise ValueError("Review snapshot or block does not match")
    _require_closed(workbook_path)
    before = _digest(workbook_path)
    if history_sha256 and before != history_sha256:
        raise ValueError("Archived omit decision workbook changed")
    if require_saved and before == manifest["initial_sha256"]:
        raise ValueError("The review workbook has not been saved with a user decision or note")
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        if workbook.sheetnames != ["Active IPOS blockers"]:
            raise ValueError("Review workbook sheets changed")
        sheet = workbook.active
        if sheet.max_column != len(REVIEW_COLUMNS) or tuple(cell.value for cell in sheet[1]) != REVIEW_COLUMNS:
            raise ValueError("Review workbook columns changed")
        rows = [[str(cell.value or "") for cell in row] for row in sheet.iter_rows(min_row=2)]
        if len(rows) != len(manifest["rows"]):
            raise ValueError("Review workbook blocker rows changed")
        notes = []
        future_updates = []
        decisions = []
        for old, new in zip(manifest["rows"], rows):
            if any(old[index] != new[index] for index in range(len(REVIEW_COLUMNS)) if index not in (9, 10)):
                raise ValueError("Only User decision and Note may be edited; review facts are not authority")
            if new[9] and new[9] not in RECOMMENDATIONS:
                raise ValueError("User decision must be one of the defined flow choices")
            future_update = "need source/spec update - clarification" in (new[7], new[9])
            needs_note = new[9] == "user clarified" or (not new[9] and new[7] == "user clarified")
            if not future_update and needs_note and require_saved and not new[10].strip():
                raise ValueError("A user clarified blocker requires a Note")
            reviewed = dict(zip(REVIEW_COLUMNS, new))
            if future_update:
                reviewed["Resulting action"] = "generate supported content; retain source/spec update gap for readiness"
                future_updates.append(reviewed)
            if new[9]:
                decisions.append(reviewed)
            if new[10].strip():
                notes.append(reviewed)
        if any(cell.data_type == "f" for row in sheet.iter_rows(min_row=2) for cell in row):
            raise ValueError("Review notes must be literal text, not Excel formulas")
    finally:
        workbook.close()
    if _digest(workbook_path) != before:
        raise ValueError("Review workbook changed during checking; save and close it again")
    _require_closed(workbook_path)
    return {"snapshot_id": snapshot_id, "block": block, "workbook_sha256": before,
            "notes": notes, "decisions": decisions, "future_source_updates": future_updates,
            "authority": False, "gate_override": False}


def upgrade_blocker_review_dropdown(repo_root: Path, *, snapshot_id: str, block: str) -> Path:
    workbook_path, manifest_path, companion_path = review_paths(repo_root, block)
    context = read_blocker_review(repo_root, snapshot_id=snapshot_id, block=block, require_saved=False)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook.active
        formula = '"' + ",".join(REVIEW_CHOICES) + '"'
        expected_range = "J2" if sheet.max_row == 2 else "J2:J" + str(sheet.max_row)
        has_dropdown = any(value.type == "list" and value.formula1 == formula and str(value.sqref) == expected_range and not value.showDropDown and value.showErrorMessage and value.errorStyle == "stop" for value in sheet.data_validations.dataValidation)
        editable = all(not row[9].protection.locked for row in sheet.iter_rows(min_row=2))
        if has_dropdown and editable:
            return workbook_path
        _add_decision_dropdown(sheet)
        _require_closed(workbook_path)
        if _digest(workbook_path) != context["workbook_sha256"]:
            raise ValueError("Review workbook changed during dropdown upgrade; save and close it again")
        workbook.save(workbook_path)
    finally:
        workbook.close()
    if context["workbook_sha256"] == manifest["initial_sha256"]:
        manifest["initial_sha256"] = _digest(workbook_path)
        manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    if companion_path.exists():
        companion_path.write_text(companion_path.read_text(encoding="utf-8").replace(
            "Edit Note only;", "Select User decision and edit Note;"), encoding="utf-8")
    return workbook_path


def apply_descriptive_omissions(repo_root: Path, *, snapshot_id: str, functional_input, review_context=None):
    from dataclasses import replace
    from workflow_routing import build_ipos_claim_plan, _ipos_document_contract, _ipos_claim_presentation_role, IPOS_MAIN_FUNCTION_WRITING_RULE

    block = functional_input.block
    workbook_path, manifest_path, _companion = review_paths(repo_root, block)
    contexts = []
    if review_context is None:
        if manifest_path.exists() and json.loads(manifest_path.read_text(encoding="utf-8"))["snapshot_id"] == snapshot_id:
            review_context = read_blocker_review(repo_root, snapshot_id=snapshot_id, block=block, require_saved=False)
    if review_context is not None:
        contexts.append((review_context, None))
    receipt_path = workbook_path.with_name("descriptive_omissions.json")
    if receipt_path.exists():
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        if receipt["snapshot_id"] == snapshot_id:
            if receipt["block"] != block or receipt.get("scope") != "descriptive_only" or receipt.get("authority") is not False or receipt.get("gate_override") is not False:
                raise ValueError("Invalid descriptive omission receipt")
            if receipt.get("writing_rule", "legacy_claim_sections_v1") not in {"legacy_claim_sections_v1", IPOS_MAIN_FUNCTION_WRITING_RULE}:
                raise ValueError("Unknown descriptive omission writing rule")
            archived = read_blocker_review(repo_root, snapshot_id=snapshot_id, block=block,
                                           history_sha256=receipt["workbook_sha256"])
            contexts.append((archived, receipt))
    if not any(row["User decision"] == "omit (as non-blocking)" for context, _receipt in contexts for row in context.get("decisions", ())):
        if any(receipt is not None and receipt["omitted_claim_ids"] for _context, receipt in contexts):
            raise ValueError("Archived omit decision no longer matches the current claim plan")
        return functional_input
    plan = build_ipos_claim_plan(functional_input)
    sections = {section["title"]: section["role"] for section in _ipos_document_contract()["sections"]}
    omitted = set(functional_input.descriptive_omitted_claim_ids)
    for context, receipt in contexts:
        local = set()
        for row in context.get("decisions", ()):
            if row["User decision"] != "omit (as non-blocking)":
                continue
            title = row["Section or claim"]
            if title in sections:
                matches = [claim for claim in plan.claims if _ipos_claim_presentation_role(claim,
                           main_function_summary=receipt is None or receipt.get("writing_rule") == IPOS_MAIN_FUNCTION_WRITING_RULE) == sections[title]]
            else:
                matches = [claim for claim in plan.claims if ", ".join(claim.responsibilities) == title]
                if len(matches) != 1:
                    raise ValueError("Omit decision cannot be uniquely bound to a current claim: " + title)
            if not matches:
                raise ValueError("Omit decision has no current descriptive claims: " + title)
            local.update(claim.claim_id for claim in matches)
        if receipt is not None and sorted(local) != receipt["omitted_claim_ids"]:
            raise ValueError("Archived omit decision no longer matches the current claim plan")
        omitted.update(local)
    return replace(functional_input, descriptive_omitted_claim_ids=tuple(sorted(omitted)))


def blocker_rows(block: str, findings: list[str], functional_input, *, user_notes=(), user_decisions=()) -> list[dict]:
    from workflow_routing import build_ipos_claim_plan, _ipos_claim_output_role, _ipos_document_contract

    plan = build_ipos_claim_plan(functional_input)
    claims = {claim.claim_id: claim for claim in plan.claims}
    sections = {section["title"].split()[0]: section for section in _ipos_document_contract()["sections"]}
    source_ids = set(functional_input.local_requirement_ids)
    points = {}
    point_findings = {}
    for finding in findings:
        clean = re.sub(r":[^:;]*\.(?:md|docx)(?=:|$)", "", str(finding))
        code, separator, location = clean.partition(":")
        point = location if separator else code
        points.setdefault(point, [])
        point_findings.setdefault(point, []).append(str(finding))
        if code not in points[point]:
            points[point].append(code)
    labels = {"ipos_final_unmodeled_content": "Unmodeled content",
              "ipos_final_claim_missing_or_changed": "Missing/changed supported claim",
              "ipos_final_required_section_missing": "Required supported section missing"}
    rows = []
    for location, codes in points.items():
        operational = all(code in {"ipos_generation_failure", "ipos_publication_failure"} for code in codes)
        claim = claims.get(location.removeprefix("audit:"))
        section = sections.get(location)
        selected = [claim] if claim else [value for value in plan.claims
                    if section and _ipos_claim_output_role(value.role) == section["role"]]
        meanings = list(dict.fromkeys(meaning for value in selected for meaning in value.meanings))
        expected = "; ".join(meaning.subject + ": " + ("not " if meaning.negative else "") + meaning.predicate
                              + " " + ", ".join(meaning.objects) for meaning in meanings[:2])
        referenced = sorted({span.record_id.rsplit(":", 1)[-1] for value in selected for span in value.spans
                             if span.record_id.rsplit(":", 1)[-1] in source_ids})
        source_update = any("relationship_changed" in value or "missing_source" in value for value in codes)
        recommendation = "need source/spec update - clarification" if source_update else "user clarified"
        title = section["title"] if section else (", ".join(claim.responsibilities) if claim else location)
        notes = [value["Note"] for value in user_notes if value.get("Section or claim") == title and value.get("Note")]
        decisions = [value["User decision"] for value in user_decisions if value.get("Section or claim") == title]
        why = "Current strict checker failure; a user note is review context only and cannot waive the gate."
        if operational:
            why = "Resolve the reported tooling error and retry; a source/spec change or user note cannot repair an operational failure."
        if notes:
            why += " The saved user note was inspected; the checker still reports this point."
        explicit_reference = location.rsplit(":", 1)[-1]
        reference = "; ".join(referenced) if claim else explicit_reference if explicit_reference in source_ids else ""
        rows.append({"Reference req_id": reference, "Block": block,
                     "Section or claim": title,
                     "Failure type": "; ".join(labels.get(code, code) for code in codes),
                     "Expected supported meaning": "No missing source meaning is inferred from this operational failure." if operational else expected or "See the current checker evidence; no missing meaning is inferred.",
                     "Observed problem": "\n".join(point_findings[location]) if operational else "Supported claim/section fails the existing checker; exact lost fact is not localized.",
                     "Likely failing layer": "generation/publication tooling" if operational else "source/model relationship review" if source_update else "realization/readback/section mapping; not localized",
                     "Flow recommendation": recommendation, "Why": why,
                     "User decision": decisions[0] if len(set(decisions)) == 1 else "",
                     "Note": "\n".join(notes),
                     "Resulting action": "generate supported content; retain source/spec update gap for readiness" if source_update else "re-evaluate after user note"})
    return rows


def assess_ipos_block(repo_root: Path, *, snapshot_id: str, kind: str, block: str,
                      rows: list[dict], candidate_rows: list[dict], writer=None,
                      descriptive_source_records=()) -> dict:
    import generate_ipos_specs as ipos
    from workflow_routing import build_ipos_functional_input, ipos_claim_preflight_findings

    config = dict(ipos.CONFIG[kind])
    config["docx"] = ipos._block_docx_name(kind, block)
    functional_input = build_ipos_functional_input(
        block, ipos._block_inventory_record(repo_root, block), rows, domain=kind, materialized=True,
        provenance=("approved snapshot", "artifacts/stage2_mirco_arc/block_inventory.csv"),
        candidate_requirement_rows=candidate_rows, descriptive_source_records=descriptive_source_records)
    workbook_path, manifest_path, _companion = review_paths(repo_root, block)
    review_context = {"notes": [], "future_source_updates": []}
    if manifest_path.exists():
        previous = json.loads(manifest_path.read_text(encoding="utf-8"))
        if previous["snapshot_id"] == snapshot_id:
            try:
                upgrade_blocker_review_dropdown(repo_root, snapshot_id=snapshot_id, block=block)
                review_context = read_blocker_review(repo_root, snapshot_id=snapshot_id, block=block, require_saved=False)
            except (ValueError, OSError) as exc:
                return {"block": block, "status": "BLOCKED", "review_path": str(workbook_path.resolve()),
                        "blocker_count": len(previous["rows"]), "waiting_reason": str(exc)}
    try:
        functional_input = apply_descriptive_omissions(repo_root, snapshot_id=snapshot_id,
                                                      functional_input=functional_input, review_context=review_context if manifest_path.exists() else None)
    except (ValueError, OSError) as exc:
        return {"block": block, "status": "BLOCKED", "review_path": str(workbook_path.resolve()),
                "blocker_count": len(previous["rows"]) if manifest_path.exists() else 1, "waiting_reason": str(exc)}
    from validate_downstream_coherence import classify_descriptive_findings
    classified = classify_descriptive_findings(ipos_claim_preflight_findings(functional_input))
    residuals = classified["residuals"]
    findings = classified["hard_findings"]
    if not findings:
        probe_root = repo_root / "artifacts/orchestrator"
        probe_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".ipos-check-", dir=probe_root) as directory:
            output = Path(directory) / ipos._block_slug(block)
            existing = repo_root / ipos.CONFIG[kind]["output_dir"] / "blocks" / ipos._block_slug(block)
            if existing.exists():
                shutil.copytree(existing, output, ignore=shutil.ignore_patterns("~$*"))
            try:
                (writer or ipos._write_ipos_document)(
                    repo_root, config, kind, config["title"].format(block_name=block), output, block, rows,
                    snapshot_id, require_docx=True, candidate_requirement_rows=candidate_rows)
            except (RuntimeError, ValueError, OSError) as exc:
                report_path = output / "document_contract_report.json"
                if report_path.exists():
                    report = json.loads(report_path.read_text(encoding="utf-8"))
                    findings = [str(value) for check in report["checks"].values() for value in check.get("findings", [])]
                findings = findings or ["ipos_generation_failure:" + str(exc)]
    if residuals and not findings:
        residual_rows = blocker_rows(block, [value["finding"] for value in residuals], functional_input,
                                     user_notes=review_context["notes"], user_decisions=review_context.get("decisions", []))
        write_blocker_review(repo_root, snapshot_id, block, residual_rows)
    if findings:
        blockers = blocker_rows(block, findings, functional_input, user_notes=review_context["notes"],
                                user_decisions=review_context.get("decisions", []))
        path = write_blocker_review(repo_root, snapshot_id, block, blockers)
        result = {"block": block, "status": "BLOCKED", "review_path": str(path.resolve()),
                  "blocker_count": len(blockers), "review_context": review_context, "findings": findings}
    else:
        result = {"block": block, "status": "READY", "blocker_count": 0, "review_context": review_context}
    result.update({"residuals": residuals, "generation_status": "GENERATABLE" if not findings else "BLOCKED",
                   "publication_status": "NOT_ATTEMPTED", "readiness_status": "NOT_ASSESSED"})
    if manifest_path.exists():
        evaluation = manifest_path.with_name("last_evaluation.json")
        evaluation.write_text(json.dumps({"snapshot_id": snapshot_id, **result, "authority": False,
                                         "gate_override": False}, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
        if (result["status"] == "READY" and not residuals and not review_context.get("future_source_updates")
            and not workbook_path.with_name("descriptive_omissions.json").exists()):
            try:
                _require_closed(workbook_path)
            except (ValueError, OSError) as exc:
                return {"block": block, "status": "BLOCKED", "review_path": str(workbook_path.resolve()),
                        "blocker_count": len(previous["rows"]), "waiting_reason": str(exc)}
            history = workbook_path.parent / "history" / _digest(workbook_path)[:16]
            history.mkdir(parents=True, exist_ok=True)
            omitted = getattr(functional_input, "descriptive_omitted_claim_ids", ())
            if omitted:
                from workflow_routing import IPOS_MAIN_FUNCTION_WRITING_RULE
                receipt = {"snapshot_id": snapshot_id, "block": block, "scope": "descriptive_only",
                           "workbook_sha256": _digest(workbook_path), "omitted_claim_ids": list(omitted),
                           "writing_rule": IPOS_MAIN_FUNCTION_WRITING_RULE,
                           "authority": False, "gate_override": False}
                workbook_path.with_name("descriptive_omissions.json").write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
            for path in review_paths(repo_root, block):
                if path.exists():
                    path.replace(history / path.name)
    return result


def emit_block_review(result: dict, snapshot_id: str) -> None:
    if result["status"] == "BLOCKED":
        print("[IPOS_REVIEW] " + json.dumps({"block": result["block"], "snapshot_id": snapshot_id,
              "path": result["review_path"], "blocker_count": result["blocker_count"],
              "reason": result.get("waiting_reason", "Current active generation blockers")}, ensure_ascii=True), flush=True)


def assess_ipos_blocks(repo_root: Path, *, snapshot_id: str, kind: str, blocks: list[str],
                       rows: list[dict], candidate_rows: list[dict], descriptive_source_records=()) -> list[dict]:
    results = []
    for block in blocks:
        result = assess_ipos_block(repo_root, snapshot_id=snapshot_id, kind=kind, block=block,
                                   rows=[row for row in rows if row["mapped_block"] == block], candidate_rows=candidate_rows,
                                   descriptive_source_records=descriptive_source_records)
        results.append(result)
        emit_block_review(result, snapshot_id)
    return results


def publish_ipos_document(writer, *args, **kwargs) -> None:
    repo_root, config, kind, _title, output, block = args[:6]
    import generate_ipos_specs as ipos

    filenames = (ipos._block_markdown_name(kind, block), config["docx"], "descriptive_summary_audit.csv",
                 "descriptive_structure_audit.json", "materialization_audit.json", "document_contract_report.json")
    with tempfile.TemporaryDirectory(prefix=".ipos-publish-", dir=repo_root / "artifacts/orchestrator") as directory:
        backup = Path(directory)
        for name in filenames:
            if (output / name).exists():
                shutil.copy2(output / name, backup / name)
        try:
            writer(*args, **kwargs)
        except Exception:
            for name in filenames:
                target = output / name
                original = backup / name
                if original.exists() and (not target.exists() or _digest(target) != _digest(original)):
                    shutil.copy2(original, target)
                elif not original.exists() and target.exists():
                    target.unlink()
            raise


def publication_failure_review(repo_root: Path, *, snapshot_id: str, kind: str, block: str,
                               rows: list[dict], candidate_rows: list[dict], failure: Exception) -> dict:
    import generate_ipos_specs as ipos
    from workflow_routing import build_ipos_functional_input

    functional = build_ipos_functional_input(block, ipos._block_inventory_record(repo_root, block), rows,
                  domain=kind, materialized=True, candidate_requirement_rows=candidate_rows)
    findings = ["ipos_publication_failure:" + str(failure)]
    blockers = blocker_rows(block, findings, functional)
    path = write_blocker_review(repo_root, snapshot_id, block, blockers)
    result = {"block": block, "status": "BLOCKED", "review_path": str(path.resolve()),
              "blocker_count": len(blockers), "findings": findings}
    emit_block_review(result, snapshot_id)
    return result