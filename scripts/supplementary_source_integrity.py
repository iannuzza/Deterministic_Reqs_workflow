"""Source-bounded supplementary text preservation checks, without authority writes."""

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


TAGGED_REQUIREMENT_HEADER_RE = re.compile(
    r"\[\s*(?P<id>[A-Z][A-Z0-9]*(?:[_-][A-Z0-9]+)*_?\s*\d+)\s*\]\s*REQUIREMENT\s*: ?",
    re.IGNORECASE,
)


def _normalize_tagged_lines(lines):
    normalized = []
    index = 0
    while index < len(lines):
        current = " ".join(lines[index].split())
        if index + 1 < len(lines) and "[" in current and "]" not in current:
            current = f"{current} {lines[index + 1].strip()}"
            normalized.extend([current, ""])
            index += 2
            continue
        normalized.append(current)
        index += 1
    return normalized


def bounded_source_bodies(page_texts):
    lines = []
    locations = []
    for page_no, text in enumerate(page_texts, start=1):
        for line_no, line in enumerate(_normalize_tagged_lines(text.splitlines()), start=1):
            lines.append(line)
            locations.append((page_no, line_no))
    text = "\n".join(lines)
    headers = list(TAGGED_REQUIREMENT_HEADER_RE.finditer(text))
    bodies = []
    for position, header in enumerate(headers):
        boundary = headers[position + 1].start() if position + 1 < len(headers) else len(text)
        end = re.search(r"\[END\]", text[header.end():boundary], re.IGNORECASE)
        source_id = re.sub(r"\s+", "", header.group("id"))
        if end is None:
            raise ValueError(f"Supplementary body lacks a bounded [END]: {source_id}")
        page_no, line_no = locations[text.count("\n", 0, header.start())]
        body = " ".join(text[header.end():header.end() + end.start()].split())
        if not body:
            raise ValueError(f"Empty supplementary body: {source_id}")
        bodies.append({"source_req_id": source_id, "requirement_statement": body,
                       "source_page": str(page_no), "source_chunk_id": f"p{page_no:03d}:l{line_no}"})
    return bodies


def validate_staged_bodies(staged, page_texts):
    bodies = bounded_source_bodies(page_texts)
    for row in staged:
        matches = [body for body in bodies if body["source_req_id"] == row.get("source_req_id")]
        if (len(matches) != 1 or matches[0]["source_page"] != row.get("source_page")
                or " ".join(row.get("requirement_statement", "").split()) != matches[0]["requirement_statement"]):
            raise ValueError(f"Supplementary body preservation validation failed: {row.get('source_req_id')}@p{row.get('source_page')}")


def validate_ingestion_batch(ingestion_dir, staged_rows=None, review_rows=None):
    def read_rows(name):
        with (ingestion_dir / name).open(encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))

    staged = read_rows("normalized_requirements.csv") if staged_rows is None else staged_rows
    reviews = read_rows("comparison_results.csv") if review_rows is None else review_rows
    review_sets = [reviews]
    workbook_path = ingestion_dir / "comparison_results.xlsx"
    if review_rows is None and workbook_path.exists():
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
        try:
            values = list(workbook["Source Ingestion Review"].iter_rows(values_only=True))
            headers = [str(value or "").strip() for value in values[0]] if values else []
            if len(set(headers)) != len(headers):
                raise ValueError("Duplicate supplementary workbook columns")
            review_sets.append([dict(zip(headers, ["" if value is None else str(value).strip() for value in row]))
                                for row in values[1:] if any(value is not None and str(value).strip() for value in row)])
        finally:
            workbook.close()
    pages = read_rows("raw_ocr_pages.csv")
    page_numbers = [int(row["page"]) for row in pages]
    if sorted(page_numbers) != list(range(1, len(pages) + 1)):
        raise ValueError("Supplementary extracted pages are missing or duplicated")
    index_rows = read_rows("ocr_index.csv")
    indexed_pages = {int(row["page"]): row for row in index_rows}
    if len(indexed_pages) != len(index_rows) or set(indexed_pages) != set(page_numbers):
        raise ValueError("Supplementary extraction index/page mismatch")
    for row in pages:
        text_path = ingestion_dir / "ocr_extracts" / Path(indexed_pages[int(row["page"])]["text_file"]).name
        if not text_path.exists() or text_path.read_text(encoding="utf-8").strip() != row["text"].strip():
            raise ValueError("Supplementary extracted text/raw-page mismatch")
    validate_staged_bodies(staged, [row["text"] for row in sorted(pages, key=lambda row: int(row["page"]))])
    staged_by_id = {row["staged_id"]: row for row in staged}
    if len(staged_by_id) != len(staged):
        raise ValueError("Duplicate supplementary staged/review IDs")
    for reviews in review_sets:
        if len({row.get("staged_id", "") for row in reviews}) != len(reviews):
            raise ValueError("Duplicate supplementary staged/review IDs")
        if set(staged_by_id) != {row.get("staged_id", "") for row in reviews}:
            raise ValueError("Supplementary review/staging identity mismatch")
        for review in reviews:
            staged_row = staged_by_id[review["staged_id"]]
            for field in ("source_req_id", "requirement_statement", "source_spec", "source_revision", "source_page", "source_chunk_id"):
                if review.get(field, "").strip() != staged_row.get(field, "").strip():
                    raise ValueError(f"Supplementary review changed staged {field}: {review['staged_id']}")
    return staged