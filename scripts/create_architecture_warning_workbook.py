#!/usr/bin/env python3
"""Create a diagnostic workbook for current Architecture Map placement warnings."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


WARNING_FIELDS = [
    "warning",
    "requirement_id",
    "review_decision",
    "approved_classification",
    "approved_block",
    "allocation_class",
    "owning_target",
    "allocation_rationale",
    "lineage_mode",
    "source_origin_req_ids",
    "hierarchy_parent_req_ids",
    "owning_domain",
    "reviewer_notes",
]


def create(repo_root: Path) -> Path:
    mapping_path = repo_root / "artifacts/stage1_specs/architecture_mapping_preview.csv"
    output_path = repo_root / "artifacts/stage1_specs/architecture_mapping_warnings.xlsx"
    with mapping_path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    warnings = [
        row for row in rows
        if (row.get("review_decision") or "").strip().casefold() in {"approved", "reassigned"}
        and (row.get("approved_block") or "").strip().casefold() in {"system", "digital", "analog"}
        and not (row.get("hierarchy_parent_req_ids") or "").strip()
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Architecture Placement Warnings"
    sheet.append(WARNING_FIELDS)
    for row in sorted(warnings, key=lambda item: (item.get("requirement_id") or "")):
        sheet.append([
            "BLOCKER: hierarchy_parent_req_ids missing; do not infer from approved_block or source type.",
            *(row.get(field, "") for field in WARNING_FIELDS[1:]),
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="F4B183")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, 1).fill = warning_fill
        for cell in sheet[row_number]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {1: 78, 2: 22, 3: 18, 4: 22, 5: 18, 6: 28, 7: 20, 8: 58, 9: 22, 10: 28, 11: 28, 12: 18, 13: 45}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    workbook.save(output_path)
    return output_path


if __name__ == "__main__":
    path = create(Path(__file__).resolve().parents[1])
    print(path)
