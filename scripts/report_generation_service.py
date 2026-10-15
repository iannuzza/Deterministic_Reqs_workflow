"""Snapshot-only report generation foundations."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from audit_log import record_event
from approved_snapshot_resolver import resolve_complete_authoritative_input
from canonical_store import connect, fingerprint, utc_now


REPORT_FIELD_PRIORITY = (
    "requirement_id", "source_req_id", "canonical_id", "parent_id", "block",
    "classification", "requirement_statement", "statement", "status", "notes",
)


def _ordered_fields(rows: list[dict[str, object]]) -> list[str]:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    priority = {name: index for index, name in enumerate(REPORT_FIELD_PRIORITY)}
    return sorted(fields, key=lambda name: (priority.get(name, 1000), fields.index(name)))


def _style_data_sheet(worksheet, fields: list[str], row_count: int, *, table_name: str = "ReportData") -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 30
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index, field in enumerate(fields, start=1):
        values = [str(worksheet.cell(row, column_index).value or "") for row in range(2, row_count + 2)]
        width = min(60, max(12, len(field) + 2, max((len(value) for value in values), default=0) + 2))
        worksheet.column_dimensions[chr(64 + column_index) if column_index <= 26 else worksheet.cell(1, column_index).column_letter].width = width
    if row_count:
        table = Table(displayName=table_name, ref=f"A1:{worksheet.cell(row_count + 1, len(fields)).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        worksheet.add_table(table)
    else:
        worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "1:1"


def _write_metadata_sheet(
    workbook: Workbook, report_kind: str, resolved, output_path: Path, row_count: int,
    source_catalog_path: Path | None = None,
) -> None:
    worksheet = workbook.create_sheet("Report Metadata")
    worksheet.sheet_view.showGridLines = False
    worksheet.append(["Catalog export metadata" if source_catalog_path else "Snapshot report metadata", ""])
    worksheet.append(["Report kind", report_kind])
    worksheet.append(["Generated at (UTC)", utc_now()])
    worksheet.append(["Snapshot ID", resolved.snapshot_id])
    worksheet.append(["Snapshot hash", resolved.snapshot_hash])
    worksheet.append(["Resolved source", str(resolved.source_path)])
    worksheet.append(["Output path", str(output_path)])
    worksheet.append(["Row count", row_count])
    worksheet.append(["Authority", "Stage 1 extracted catalog; not canonical snapshot authority" if source_catalog_path else "Approved immutable snapshot; derived artifact"])
    if source_catalog_path:
        worksheet.append(["Catalog source", str(source_catalog_path)])
    worksheet.append(["Partial export", "No"])
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 96
    worksheet.freeze_panes = "A2"
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in worksheet.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True, color="1F4E78")
        row[1].alignment = Alignment(vertical="top", wrap_text=True)


def write_readiness_rule_workbook(output_path: Path | None = None, *, expected_workbook_sha256: str | None = None) -> Path:
    import hashlib
    from readiness_checklist import CHECKLIST_PATH, checklist_definition, checklist_fingerprint, sync_checklist_from_excel

    output_path = output_path or CHECKLIST_PATH.with_suffix(".xlsx")
    if expected_workbook_sha256 is not None:
        if output_path.with_name("~$" + output_path.name).exists() or hashlib.sha256(output_path.read_bytes()).hexdigest() != expected_workbook_sha256:
            raise ValueError("Readiness workbook changed or is locked; refusing to overwrite rule edits")
    elif output_path.resolve() == CHECKLIST_PATH.with_suffix(".xlsx").resolve():
        sync_checklist_from_excel()
    definition = checklist_definition()
    fingerprint = (hashlib.sha256(CHECKLIST_PATH.read_bytes()).hexdigest() if expected_workbook_sha256 is not None
                   else checklist_fingerprint())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Checklist Rules"
    fields = ["Category", "Checklist ID", "Checklist Rule"]
    sheet.append(fields)
    for category, items in definition["categories"].items():
        for identity, description in items.items():
            sheet.append([category, identity, description])
    _style_data_sheet(sheet, fields, sheet.max_row - 1, table_name="ReadinessRules")
    for column, width in (("A", 30), ("B", 30), ("C", 100)):
        sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[cells[0].row].height = max(42, 15 * ((len(str(cells[2].value)) + 85) // 86))
        for cell in cells:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.sheet_properties.outlinePr.summaryRight = False
    metadata = workbook.create_sheet("Policy Metadata")
    metadata.append(["Field", "Value"])
    for key, value in (("Policy version", definition["version"]), ("Source file", CHECKLIST_PATH.name),
                       ("Source SHA256", fingerprint), ("Rule count", sheet.max_row - 1),
                       ("Authority", "false; editable checklist reference, not source requirement approval"),
                       ("Policy source", "Saved Excel rule edits are validated and synchronized to readiness_checklist.json before final SysML execution"),
                       ("Synchronization", "Save and close, then run python -B scripts/readiness_checklist.py --sync-from-excel, or start the final SysML phase"),
                       ("Applicability", "Same catalog for each governed model; non-applicability requires evidence. DFT/DFX and diagnostics are informational unless required by governed policy"),
                       ("Standard interfaces", "Admitted applicable standard references cover their defined semantics, handshakes and timing without redundant restatement; roles, modes, deviations, sidebands and integration constraints remain checked"),
                       ("Required optional rules", json.dumps(definition["required_optional_rule_ids"])),
                       ("Assessment results", "Not included; this workbook defines rules, not per-model outcomes")):
        metadata.append([key, value])
    _style_data_sheet(metadata, ["Field", "Value"], metadata.max_row - 1, table_name="ReadinessRuleMetadata")
    metadata.column_dimensions["A"].width = 26
    metadata.column_dimensions["B"].width = 100
    for cells in metadata.iter_rows(min_row=2):
        metadata.row_dimensions[cells[0].row].height = max(32, 15 * ((len(str(cells[1].value)) + 85) // 86))
        for cell in cells:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(output_path.stem + ".export.tmp.xlsx")
    try:
        workbook.save(temporary)
        if expected_workbook_sha256 is not None:
            if output_path.with_name("~$" + output_path.name).exists() or hashlib.sha256(output_path.read_bytes()).hexdigest() != expected_workbook_sha256:
                raise ValueError("Readiness workbook changed or is locked; refusing to overwrite rule edits")
        temporary.replace(output_path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return output_path


def write_readiness_workbooks(directory: Path, summary: dict, reports: dict[Path, dict]) -> None:
    def add_sheet(workbook, title, fields, rows, table_name):
        worksheet = workbook.create_sheet(title)
        worksheet.append(fields)
        for row in rows:
            worksheet.append([json.dumps(row.get(field), sort_keys=True) if isinstance(row.get(field), (list, dict))
                              else row.get(field, "") for field in fields])
        for cells in worksheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        _style_data_sheet(worksheet, fields, len(rows), table_name=table_name)
        return worksheet

    def new_workbook():
        workbook = Workbook()
        workbook.remove(workbook.active)
        return workbook

    def metadata(workbook, model, status, binding):
        from readiness_checklist import checklist_reference

        reference = checklist_reference()
        values = {"Model": model, "Decision": status, "Snapshot ID": binding["snapshot_id"],
                  "Policy version": summary["policy_version"], "Checklist fingerprint": binding["checklist_fingerprint"],
                  "Publication status": summary["publication_status"], "Authority": "false; derived report only",
                  "Evidence": "Governed snapshot only; candidate references do not establish support",
                  "Report edits": "Never consumed as approval or readiness evidence",
                  "Checklist reference": reference["excel"], "Checklist obligation": reference["purpose"],
                  "Rule editing": reference["editing"], "Rule synchronization": reference["sync_command"],
                  "Limitations": "Uncertain semantics block READY; a missing match is not proof of source absence"}
        add_sheet(workbook, "Metadata", ["Field", "Value"], [{"Field": key, "Value": value} for key, value in values.items()], "MetadataTable")
        for cells in workbook["Metadata"].iter_rows(min_row=2):
            if cells[0].value == "Checklist reference":
                cells[1].hyperlink = "../../../" + reference["excel"]
                cells[1].style = "Hyperlink"

    for path, report in reports.items():
        workbook = new_workbook()
        fields = ["Category", "Checklist item", "Assessment result", "Evidence status", "Reason", "Readiness impact",
                  "Checklist ID", "Evidence references", "Candidate evidence references"]
        add_sheet(workbook, "Checklist", fields, report["checklist"], "ChecklistTable")
        add_sheet(workbook, "Blocking Items", ["Checklist ID", "Reason"], report["blocking_items_summary"], "BlockingItemsTable")
        add_sheet(workbook, "Residual Gaps", ["Gap"], [{"Gap": gap} for gap in report["residual_gaps_preventing_READY"]], "ResidualGapsTable")
        metadata(workbook, report["model"], report["overall_decision"], report["binding"])
        workbook.save(path.with_suffix(".xlsx"))
        workbook.close()
    workbook = new_workbook()
    fields = ["Block/model", "Readiness", "Pass", "Fail", "Not applicable", "Uncertain", "Blocking items", "Main blockers", "Official model", "Excel report"]
    for title, key, table_name in (("Block Readiness", "blocks", "BlockReadinessTable"),
                                    ("System and Shared", "structural_models", "StructuralReadinessTable")):
        rows = [{"Block/model": row["block"], "Readiness": row["readiness"], "Pass": row["checklist_results"]["pass"],
                 "Fail": row["checklist_results"]["fail"], "Not applicable": row["checklist_results"]["not_applicable"],
                 "Uncertain": row["checklist_results"]["uncertain"], "Blocking items": row["blocking_items"],
                 "Main blockers": row["main_blockers"], "Official model": row["model"],
                 "Excel report": Path(row["report"]).with_suffix(".xlsx").name} for row in summary[key]]
        worksheet = add_sheet(workbook, title, fields, rows, table_name)
        for index, row in enumerate(rows, start=2):
            worksheet.cell(index, len(fields)).hyperlink = row["Excel report"]
            worksheet.cell(index, len(fields)).style = "Hyperlink"
    metadata(workbook, "Official centrally scoped model set", summary["status"], summary["binding"])
    workbook.save(directory / "readiness_summary.xlsx")
    workbook.close()


def write_relation_reports(directory: Path, report: dict, inventory: dict) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    payload = {**report, "inventory": inventory}
    (directory / "relation_summary.json").write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [("View Coverage", ["view", "expected", "serialized", "gaps", "status"],
               [{"view": key, **value} for key, value in report["views"].items()]),
              ("Block Matrix", ["block", "containment", "block_relation", "signal_connection", "source_signals", "block_peer_references", "descriptions_not_established",
                                "peer_references_resolved", "peer_references_not_established", "gaps", "relation_ids"], report["blocks"]),
              ("Gaps", ["view", "status", "resolution_status", "relation_id", "kind", "source", "target", "evidence", "reason", "candidate", "statement"], report["gaps"]),
              ("Integrity Findings", ["finding"], [{"finding": value} for value in report["findings"]]),
              ("Relations", ["identity", "relation_class", "kind", "source", "target", "status", "evidence", "qualifiers", "snapshot_id", "reason"], inventory["relations"]),
              ("Signal Checks", ["block", "port", "kind", "name_status", "direction", "direction_status", "width", "description", "description_status", "peer_endpoint",
                                 "connection_status", "resolution_status", "peer_block", "peer_role", "reference_level", "peer_port_status",
                                 "source_table", "source_page", "source_file", "evidence", "snapshot_id", "reason"], inventory.get("signal_checks", [])),
              ("Signal Status Crosscheck", ["block", "port", "kind", "direction", "source", "target", "resolution_status", "model_status",
                                           "relation_ids", "snapshot_id", "reason"], report.get("signal_status_crosscheck", [])),
              ("Metadata", ["field", "value"], [{"field": key, "value": value} for key, value in
                {"snapshot_id": report["snapshot_id"], "inventory_fingerprint": report["inventory_fingerprint"],
                 "authority": False, "policy_version": inventory["policy_version"],
                 "limitation": "Candidate accounting and block coverage do not prove complete signal connectivity; report edits never confer approval."}.items()])]
    for index, (title, fields, rows) in enumerate(sheets):
        sheet = workbook.create_sheet(title)
        sheet.append(fields)
        for row in rows:
            sheet.append([json.dumps(row.get(field), sort_keys=True) if isinstance(row.get(field), (dict, list, tuple)) else row.get(field, "") for field in fields])
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        _style_data_sheet(sheet, fields, len(rows), table_name="RelationReport" + str(index))
    workbook.save(directory / "relation_summary.xlsx")
    workbook.close()


def generate_snapshot_report(
    repo_root: Path,
    *,
    project_id: str,
    snapshot_id: str | None,
    use_latest_approved: bool,
    report_kind: str,
    output_path: Path,
    rows: Iterable[Mapping[str, object]],
    source_catalog_path: Path | None = None,
    additional_sheets: Mapping[str, Iterable[Mapping[str, object]]] | None = None,
) -> Path:
    resolved, selection = resolve_complete_authoritative_input(
        repo_root, report_kind, project_id=project_id, snapshot_id=snapshot_id
    )
    materialized = [dict(row) for row in rows]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fields = (list(materialized[0]) if source_catalog_path and materialized else _ordered_fields(materialized)) or ["snapshot_id"]
    if output_path.suffix.lower() == ".xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = report_kind[:31]
        worksheet.append(fields)
        for row in materialized:
            worksheet.append([row.get(field, "") for field in fields])
        _style_data_sheet(worksheet, fields, len(materialized))
        _write_metadata_sheet(workbook, report_kind, resolved, output_path, len(materialized), source_catalog_path)
        for sheet_index, (title, sheet_rows) in enumerate((additional_sheets or {}).items(), start=2):
            summary_rows = [dict(row) for row in sheet_rows]
            summary_fields = list(dict.fromkeys(key for row in summary_rows for key in row)) or ["snapshot_id"]
            summary = workbook.create_sheet(title)
            summary.append(summary_fields)
            for row in summary_rows:
                summary.append([row.get(field, "") for field in summary_fields])
            _style_data_sheet(summary, summary_fields, len(summary_rows), table_name=f"ReportData{sheet_index}")
            for column_index, field in enumerate(summary_fields, start=1):
                if field.endswith("_percentage"):
                    for column in summary.iter_cols(min_col=column_index, max_col=column_index, min_row=2):
                        for cell in column:
                            cell.number_format = '0.0"%"'
        workbook.properties.title = f"{report_kind} catalog export" if source_catalog_path else f"{report_kind} snapshot report"
        workbook.properties.subject = f"Stage 1 catalog: {source_catalog_path}" if source_catalog_path else f"Derived from approved snapshot {resolved.snapshot_id}"
        workbook.save(output_path)
    else:
        with output_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(materialized)
    connection = connect(repo_root)
    try:
        report_material = {"snapshot_id": resolved.snapshot_id, "kind": report_kind, "output": str(output_path), "row_count": len(materialized)}
        report_id = f"report-{fingerprint(report_material)[:24]}"
        metadata = {"row_count": len(materialized), "fields": fields, "authority": "stage1_extracted_catalog" if source_catalog_path else "approved_snapshot", "source_path": str(source_catalog_path or resolved.source_path), "selection": selection}
        if additional_sheets:
            metadata["additional_sheets"] = list(additional_sheets)
        connection.execute("INSERT OR REPLACE INTO report_metadata(report_id, snapshot_id, report_kind, output_format, output_path, source_fingerprint, generated_at, status, metadata_json) VALUES (?, ?, ?, ?, ?, ?, ?, 'generated', ?)", (report_id, resolved.snapshot_id, report_kind, output_path.suffix.lower().lstrip(".") or "csv", str(output_path), fingerprint(materialized) if source_catalog_path else resolved.snapshot_hash, utc_now(), json.dumps(metadata, sort_keys=True)))
        connection.commit()
    finally:
        connection.close()
    record_event(repo_root, event_type="report_generated", entity_type="report", entity_id=report_id, actor="system", message=f"{report_kind} exported from Stage 1 catalog" if source_catalog_path else f"{report_kind} report generated from approved snapshot", project_id=project_id, payload={"snapshot_id": resolved.snapshot_id, "output_path": str(output_path)})
    return output_path
