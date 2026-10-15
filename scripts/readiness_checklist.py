"""Non-authoritative checklist mechanics used only by the central readiness assessor."""

import hashlib
import json
import re
from pathlib import Path


CHECKLIST_PATH = Path(__file__).resolve().parents[1] / "config/readiness_checklist.json"
RESULTS = {"pass", "fail", "not_applicable", "uncertain"}
EVIDENCE = {"supported", "partially_supported", "missing", "contradictory"}
IMPACTS = {"blocking", "non_blocking", "informational"}
REQUIRED_IDS = frozenset("""
scope.boundary scope.purpose scope.behavior io.names io.direction io.width io.meaning io.interpretation io.handshake io.deviations
clock.identity clock.synchrony reset.identity reset.polarity reset.semantics reset.initialization clock.domains
control.identity control.states control.transitions control.initialization control.interactions
function.main function.triggers function.sequence function.priority function.errors
data.transformations data.arithmetic data.signedness data.precision data.latency
config.effects config.illegal config.dependencies timing.sequence timing.local timing.standard
structure.parts structure.interactions structure.support consistency.contradictions consistency.ambiguity consistency.traceability consistency.snapshot residual.material
io.grouping io.standard io.role io.mode io.sideband io.integration clock.cdc clock.power control.recovery
data.format data.overflow config.csr config.memory config.buffering timing.performance evidence.authority residual.classification diagnostics.behavior readiness.decision
""".split())
OPTIONAL_IDS = frozenset({"diagnostics.behavior"})
STANDARD_PREREQUISITES = frozenset({"io.standard", "io.role", "io.mode", "io.deviations", "io.sideband", "io.integration"})
STANDARD_COVERED_IDS = frozenset({"io.meaning", "io.interpretation", "io.handshake", "timing.local", "timing.standard"})


def _validate_definition(definition):
    if definition.get("authority") is not False or definition.get("version") != "governed-ai-readiness-checklist-v2":
        raise ValueError("Invalid readiness checklist policy")
    categories = definition.get("categories")
    if not isinstance(categories, dict) or not all(isinstance(category, str) and category.strip() and isinstance(items, dict)
                                                  for category, items in categories.items()):
        raise ValueError("Invalid readiness checklist categories")
    identities = [identity for items in categories.values() for identity in items]
    if not REQUIRED_IDS.issubset(identities) or len(identities) != len(set(identities)):
        raise ValueError("Readiness checklist obligations are missing or duplicate")
    if any(not isinstance(identity, str) or not re.fullmatch(r"[a-z][a-z0-9_]*(?:\.[a-z][a-z0-9_]*)+", identity) for identity in identities):
        raise ValueError("Invalid readiness checklist ID")
    if not all(isinstance(text, str) and text.strip() for items in categories.values() for text in items.values()):
        raise ValueError("Readiness checklist descriptions are missing")
    optional = definition.get("optional_rule_ids")
    required_optional = definition.get("required_optional_rule_ids")
    if (not isinstance(optional, list) or any(not isinstance(identity, str) for identity in optional)
            or set(optional) != OPTIONAL_IDS or len(optional) != len(OPTIONAL_IDS)
            or not isinstance(required_optional, list) or any(not isinstance(identity, str) or identity not in OPTIONAL_IDS for identity in required_optional)
            or len(required_optional) != len(set(required_optional))):
        raise ValueError("Invalid optional readiness rule policy")
    return definition


def checklist_definition():
    return _validate_definition(json.loads(CHECKLIST_PATH.read_text(encoding="utf-8")))


def checklist_reference() -> dict:
    return {"excel": "config/readiness_checklist.xlsx", "json": "config/readiness_checklist.json",
            "purpose": "Checklist reference to be satisfied for every hierarchy-level model and block",
            "editing": "Users may edit the Excel rules. Save and close the workbook; validated edits are synchronized to JSON before final SysML execution.",
            "sync_command": "python -B scripts/readiness_checklist.py --sync-from-excel",
            "safety": "Mandatory built-in IDs cannot be removed. Modified and added rules require current evidence; old assessments are invalidated. Diagnostics are informational unless required by governed policy; standard coverage never waives local integration checks."}


def _workbook_categories(workbook):
    if "Checklist Rules" not in workbook or "Policy Metadata" not in workbook:
        raise ValueError("Readiness workbook requires Checklist Rules and Policy Metadata sheets")
    sheet = workbook["Checklist Rules"]
    if tuple(cell.value for cell in sheet[1][:3]) != ("Category", "Checklist ID", "Checklist Rule"):
        raise ValueError("Readiness workbook rule headers have changed")
    categories, identities = {}, set()
    for cells in sheet.iter_rows(min_row=2):
        if all(cell.value is None for cell in cells):
            continue
        if any(cell.data_type == "f" for cell in cells) or any(cell.value is not None for cell in cells[3:]):
            raise ValueError("Readiness rules require literal values in exactly three columns")
        values = [cell.value for cell in cells[:3]]
        if len(values) != 3 or any(not isinstance(value, str) or not value.strip() for value in values):
            raise ValueError("Readiness rules cannot contain blank categories, IDs or descriptions")
        category, identity, description = values
        if identity in identities:
            raise ValueError("Duplicate readiness checklist ID: " + identity)
        identities.add(identity)
        categories.setdefault(category, {})[identity] = description
    return categories


def sync_checklist_from_excel() -> bool:
    from openpyxl import load_workbook

    source = CHECKLIST_PATH
    path = source.with_suffix(".xlsx")
    if not path.exists():
        return False
    if path.with_name("~$" + path.name).exists():
        raise ValueError("Save and close the readiness checklist workbook before synchronization")
    original = source.read_bytes()
    original_workbook = path.read_bytes()
    definition = _validate_definition(json.loads(original))
    workbook = load_workbook(path, data_only=False)
    try:
        categories = _workbook_categories(workbook)
        updated = _validate_definition({**definition, "categories": categories})
        if categories == definition["categories"]:
            return False
        metadata = {row[0].value: row[1] for row in workbook["Policy Metadata"].iter_rows(min_row=2, max_col=2)}
        if "Source SHA256" not in metadata or metadata["Source SHA256"].value != hashlib.sha256(original).hexdigest():
            raise ValueError("Readiness workbook and JSON changed independently; resolve the policy conflict before synchronization")
        previous = {identity: text for items in definition["categories"].values() for identity, text in items.items()}
        customized = set(definition.get("customized_rule_ids", ()))
        customized.update(identity for items in categories.values() for identity, text in items.items() if previous.get(identity) != text)
        updated["customized_rule_ids"] = sorted(customized)
        encoded = (json.dumps(updated, indent=2, ensure_ascii=True) + "\n").encode("utf-8")
        metadata["Source SHA256"].value = hashlib.sha256(encoded).hexdigest()
        if "Rule count" in metadata:
            metadata["Rule count"].value = sum(len(items) for items in categories.values())
        temporary_workbook = path.with_name(path.stem + ".sync.tmp.xlsx")
        temporary_json = source.with_name(source.name + ".sync.tmp")
        try:
            workbook.save(temporary_workbook)
            temporary_json.write_bytes(encoded)
            if source.read_bytes() != original or path.read_bytes() != original_workbook:
                raise ValueError("Readiness policy changed during synchronization; no update applied")
            temporary_json.replace(source)
            try:
                temporary_workbook.replace(path)
            except OSError:
                if source.read_bytes() == encoded:
                    temporary_json.write_bytes(original)
                    temporary_json.replace(source)
                raise
        finally:
            temporary_workbook.unlink(missing_ok=True)
            temporary_json.unlink(missing_ok=True)
        return True
    finally:
        workbook.close()


def checklist_fingerprint():
    from openpyxl import load_workbook

    original = CHECKLIST_PATH.read_bytes()
    definition = _validate_definition(json.loads(original))
    path = CHECKLIST_PATH.with_suffix(".xlsx")
    if path.with_name("~$" + path.name).exists():
        raise ValueError("Save and close the readiness checklist workbook before assessment")
    if path.exists():
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if _workbook_categories(workbook) != definition["categories"]:
                raise ValueError("Readiness Excel rules differ from JSON; run scripts/readiness_checklist.py --sync-from-excel")
        finally:
            workbook.close()
    return hashlib.sha256(original).hexdigest()


def build_checklist(checks=None, *, candidate_evidence_refs=()):
    checks = checks or {}
    definition = checklist_definition()
    known_ids = {identity for items in definition["categories"].values() for identity in items}
    if set(checks) - known_ids:
        raise ValueError("Unknown readiness checklist result")
    rows = []
    for category, items in definition["categories"].items():
        for identity, text in items.items():
            row = {
                "Checklist ID": identity, "Category": category, "Checklist item": text,
                "Assessment result": "uncertain", "Evidence status": "partially_supported" if candidate_evidence_refs else "missing",
                "Reason": "No completeness-capable central evidence adapter establishes this item or its non-applicability; candidate evidence is not proof.",
                "Readiness impact": "blocking", "Evidence references": [],
                "Candidate evidence references": list(candidate_evidence_refs),
            }
            if identity in definition.get("customized_rule_ids", ()):
                row["Reason"] = "User-modified checklist rule requires a matching central evidence adapter; previous results do not establish support."
            else:
                row.update(checks.get(identity, {}))
            if identity in OPTIONAL_IDS and identity not in definition["required_optional_rule_ids"]:
                if checks.get(identity, {}).get("Readiness impact") != "blocking":
                    row["Readiness impact"] = "informational"
                    row["Reason"] = "Optional under the active checklist policy; not a universal blocker. " + row["Reason"]
            elif identity in OPTIONAL_IDS:
                row["Readiness impact"] = "blocking"
            rows.append(row)
    by_id = {row["Checklist ID"]: row for row in rows}
    for row in rows:
        if row.get("Coverage basis") == "standard" and not _standard_coverage_valid(row, by_id):
            row.update({"Assessment result": "uncertain", "Evidence status": "missing", "Evidence references": [],
                        "Reason": "Standard coverage requires traceable applicability, role, subset or mode, deviations, sidebands and local integration checks; otherwise full custom checks apply."})
    decision = by_id["readiness.decision"]
    if "readiness.decision" in definition.get("customized_rule_ids", ()):
        return rows
    supported = all(_row_satisfied(row, definition) for row in rows if row is not decision)
    decision.update({"Assessment result": "pass" if supported else "uncertain",
                     "Evidence status": "supported" if supported else "missing", "Readiness impact": "blocking",
                     "Evidence references": sorted({reference for row in rows if row is not decision
                                                     for reference in row["Evidence references"]}) if supported else [],
                     "Reason": "All required checklist obligations are supported; central coherence, scope and dependency gates still apply." if supported
                               else "Required checklist obligations remain unresolved; generation or publication does not establish READY."})
    return rows


def _supported_result(row):
    references = row.get("Evidence references")
    return (row.get("Assessment result") in {"pass", "not_applicable"} and row.get("Evidence status") == "supported"
            and isinstance(references, list) and bool(references)
            and all(isinstance(reference, str) and reference.strip() for reference in references))


def _standard_coverage_valid(row, by_id):
    scope = row.get("Standard-defined scope")
    if row["Checklist ID"] not in STANDARD_COVERED_IDS or not isinstance(scope, str) or not scope.strip():
        return False
    return all(_supported_result(by_id[identity]) and by_id[identity].get("Reason")
               and (identity not in {"io.standard", "io.role"} or by_id[identity].get("Assessment result") == "pass")
               for identity in STANDARD_PREREQUISITES)


def _row_satisfied(row, definition):
    result, evidence, impact = (row.get(field) for field in ("Assessment result", "Evidence status", "Readiness impact"))
    if result not in RESULTS or evidence not in EVIDENCE or impact not in IMPACTS or not row.get("Reason"):
        return False
    if (row.get("Checklist ID") in OPTIONAL_IDS and row["Checklist ID"] not in definition["required_optional_rule_ids"]
            and impact == "informational" and result in {"fail", "uncertain"}):
        return True
    return _supported_result(row) and (row.get("Checklist ID") != "readiness.decision" or result == "pass")


def checklist_decision(rows):
    definition = checklist_definition()
    expected = {identity: (category, text) for category, items in definition["categories"].items()
                for identity, text in items.items()}
    if not isinstance(rows, list) or len(rows) != len(expected):
        return "NOT_READY"
    seen = set()
    for row in rows:
        if not isinstance(row, dict):
            return "NOT_READY"
        identity = row.get("Checklist ID")
        if identity not in expected or identity in seen:
            return "NOT_READY"
        seen.add(identity)
        if (row.get("Category"), row.get("Checklist item")) != expected[identity]:
            return "NOT_READY"
        if not _row_satisfied(row, definition):
            return "NOT_READY"
    by_id = {row["Checklist ID"]: row for row in rows}
    if any(row.get("Coverage basis") == "standard" and not _standard_coverage_valid(row, by_id) for row in rows):
        return "NOT_READY"
    return "READY"


def blocking_items(rows):
    return [row for row in rows if row.get("Readiness impact") == "blocking"
            and row.get("Assessment result") in {"fail", "uncertain"}]


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Synchronize saved readiness checklist rules from Excel to JSON")
    parser.add_argument("--sync-from-excel", action="store_true", required=True)
    parser.parse_args()
    print("Readiness checklist JSON updated from Excel" if sync_checklist_from_excel() else "Readiness checklist already synchronized")