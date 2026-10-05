#!/usr/bin/env python3
"""Read-only, project-agnostic coherence check from source review to Stage 2B."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from pre_freeze_coherence_gate import authority_fingerprint
from requirement_allocation_policy import requires_hierarchy_parent


SEVERITIES = ("ERROR", "WARNING", "INFO")
ALLOCATION_TARGETS = {
    "system_level": "SRS",
    "top_digital_architecture": "DRS",
    "top_analog_architecture": "ARS",
    "block_local_digital": "Digital IPOS",
    "block_local_analog": "Analog IPOS",
    "descriptive_only": "none",
}
REQUIRED_ALLOCATION_FIELDS = (
    "allocation_class", "owning_target", "allocation_rationale", "lineage_mode",
    "source_origin_req_ids", "owning_domain",
)
ID_FIELDS = ("requirement_id", "source_req_id", "id", "Requirement ID")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [{key: _text(value) for key, value in row.items()} for row in csv.DictReader(handle)]


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _id(row: dict[str, str]) -> str:
    for field in ID_FIELDS:
        if _text(row.get(field)):
            return _text(row.get(field))
    return ""


def _ids(rows: list[dict[str, str]]) -> tuple[set[str], list[str]]:
    values = [_id(row) for row in rows]
    present = [value for value in values if value]
    counts = Counter(present)
    return set(present), sorted(value for value, count in counts.items() if count > 1)


def _add(findings: list[dict[str, str]], severity: str, category: str, message: str, path: str = "") -> None:
    findings.append({"severity": severity, "category": category, "message": message, "path": path})


def _table_columns(connection: sqlite3.Connection, table: str) -> set[str]:
    return {row[1] for row in connection.execute(f"PRAGMA table_info({table})").fetchall()}


def _rows(connection: sqlite3.Connection, query: str, parameters: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
    cursor = connection.execute(query, parameters)
    names = [description[0] for description in cursor.description or []]
    return [dict(zip(names, row)) for row in cursor.fetchall()]


def _workbook_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if "Architecture Mapping Review" not in workbook.sheetnames:
                return []
            worksheet = workbook["Architecture Mapping Review"]
            iterator = worksheet.iter_rows(values_only=True)
            headers = [_text(value) for value in next(iterator, ())]
            return [
                {header: _text(value) for header, value in zip(headers, row)}
                for row in iterator
                if any(_text(value) for value in row)
            ]
        finally:
            workbook.close()
    except (ImportError, OSError, ValueError, KeyError):
        pass
    try:
        with zipfile.ZipFile(path) as archive:
            shared = []
            if "xl/sharedStrings.xml" in archive.namelist():
                import xml.etree.ElementTree as element_tree
                root = element_tree.fromstring(archive.read("xl/sharedStrings.xml"))
                ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                for item in root.findall("x:si", ns):
                    shared.append("".join(node.text or "" for node in item.iter() if node.tag.endswith("}t")))
            import xml.etree.ElementTree as element_tree
            workbook = element_tree.fromstring(archive.read("xl/workbook.xml"))
            rels = element_tree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            rel_map = {
                relationship.attrib["Id"]: relationship.attrib["Target"]
                for relationship in rels
            }
            ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
            sheet = next((item for item in workbook.findall("x:sheets/x:sheet", ns) if item.attrib.get("name") == "Architecture Mapping Review"), None)
            if sheet is None:
                return []
            target = rel_map[sheet.attrib["{" + ns["r"] + "}id"]]
            sheet_path = "xl/" + target.lstrip("/") if not target.startswith("/") else target.lstrip("/")
            root = element_tree.fromstring(archive.read(sheet_path))
            rows: list[list[str]] = []
            for row in root.findall(".//x:sheetData/x:row", ns):
                values: list[str] = []
                for cell in row.findall("x:c", ns):
                    value = cell.find("x:v", ns)
                    if cell.attrib.get("t") == "inlineStr":
                        text = "".join(
                            node.text or ""
                            for node in cell.findall(".//x:t", ns)
                        )
                    else:
                        text = "" if value is None else value.text or ""
                    if cell.attrib.get("t") == "s" and text.isdigit() and int(text) < len(shared):
                        text = shared[int(text)]
                    values.append(_text(text))
                rows.append(values)
            if not rows:
                return []
            headers = rows[0]
            return [dict(zip(headers, row + [""] * (len(headers) - len(row)))) for row in rows[1:] if any(row)]
    except (OSError, KeyError, ValueError, zipfile.BadZipFile):
        return []


def _historical_warning_metadata(repo_root: Path) -> dict[str, dict[str, str]]:
    path = repo_root / "artifacts/stage1_specs/historical_ingestion_warnings.csv"
    if not path.exists():
        return {}
    rows = _read_csv(path)
    return {
        row.get("path", ""): row
        for row in rows
        if row.get("path")
    }


def _is_proven_non_authoritative_historical(
    metadata: dict[str, str], historical_path: Path, active_path: Path
) -> bool:
    triage = _text(metadata.get("workflow_triage")).casefold()
    current_effect = _text(metadata.get("current_authority_effect")).casefold()
    return (
        _text(metadata.get("historical_batch_id")) == historical_path.name
        and Path(_text(metadata.get("active_batch_path"))) == active_path
        and (current_effect == "none" or current_effect.startswith("none;"))
        and (
            triage.startswith("info_only")
            or (
                _text(metadata.get("status")) == "historical_non_authoritative"
                and _text(metadata.get("cleanup_action")) == "mark_historical_non_authoritative"
            )
        )
    )


def _discover_batches(repo_root: Path, findings: list[dict[str, str]]) -> tuple[Path | None, dict[str, Any]]:
    root = repo_root / "artifacts" / "source_ingestion"
    context = _read_json(repo_root / "config/project_context.json")
    configured_primary = Path(_text(context.get("source_spec_path"))).name.casefold()
    candidates: list[tuple[str, Path, dict[str, Any]]] = []
    for path in sorted(root.iterdir()) if root.exists() else []:
        if not path.is_dir():
            continue
        metadata_path = path / "source_metadata.json"
        metadata = _read_json(metadata_path) if metadata_path.exists() else {}
        if metadata.get("ingestion_id") or (path / "normalized_requirements.csv").exists():
            candidates.append((_text(metadata.get("staged_at")) or path.name, path, metadata))
    if not candidates:
        _add(findings, "ERROR", "active_batch", "No source-review ingestion batch was discoverable", str(root))
        return None, {}
    supplementary = [
        item for item in candidates
        if _text(item[2].get("source_kind")).casefold() == "supplementary"
        or (
            not _text(item[2].get("source_kind"))
            and configured_primary
            and Path(_text(item[2].get("source_spec"))).name.casefold() != configured_primary
        )
    ]
    if supplementary:
        candidates = supplementary
    candidates.sort(key=lambda item: (item[0], item[1].name))
    newest_key = candidates[-1][0]
    newest = [item for item in candidates if item[0] == newest_key]
    if len(newest) != 1:
        _add(findings, "ERROR", "active_batch", "Active source-review batch is not uniquely identifiable")
    active = newest[-1]
    historical_metadata = _historical_warning_metadata(repo_root)
    for _, path, _ in candidates[:-1]:
        finding_message = "Historical/superseded ingestion batch coexists with active batch"
        if _is_proven_non_authoritative_historical(historical_metadata.get(str(path), {}), path, active[1]):
            _add(
                findings,
                "INFO",
                "stale_state",
                f"Historical/superseded ingestion batch retained for audit; non-authoritative by workflow metadata: {path.name}",
                str(path),
            )
        else:
            _add(findings, "WARNING", "stale_state", finding_message, str(path))
    metadata = dict(active[2])
    metadata["batch_path"] = str(active[1])
    metadata["candidate_count"] = len(candidates)
    if not _text(metadata.get("source_revision")) or not _text(metadata.get("source_revision_date")):
        _add(findings, "ERROR", "revision", "Active batch lacks source revision identity", str(active[1] / "source_metadata.json"))
    return active[1], metadata


def _validate_csv_ids(findings: list[dict[str, str]], path: Path, category: str) -> set[str]:
    if not path.exists():
        _add(findings, "WARNING", category, "Artifact is absent", str(path))
        return set()
    rows = _read_csv(path)
    values, duplicates = _ids(rows)
    if duplicates:
        _add(findings, "ERROR", "id_chain", f"Duplicate IDs: {', '.join(duplicates[:20])}", str(path))
    if not values:
        _add(findings, "ERROR", "id_chain", "Artifact has no usable requirement IDs", str(path))
    return values


def _check_allocation(findings: list[dict[str, str]], rows: list[dict[str, str]], path: Path, label: str) -> dict[str, dict[str, str]]:
    by_id: dict[str, dict[str, str]] = {}
    for row in rows:
        requirement_id = _id(row)
        if not requirement_id:
            continue
        by_id[requirement_id] = row
        decision = _text(row.get("review_decision")).casefold()
        if decision in {"approved", "reassigned"}:
            missing = [field for field in REQUIRED_ALLOCATION_FIELDS if not _text(row.get(field))]
            if missing:
                _add(findings, "ERROR", "hierarchical_placement", f"Approved row {requirement_id} is missing: {', '.join(missing)}", str(path))
            allocation_class = _text(row.get("allocation_class")).casefold()
            target = _text(row.get("owning_target"))
            if allocation_class not in ALLOCATION_TARGETS:
                _add(findings, "ERROR", "hierarchical_placement", f"Approved row {requirement_id} has invalid allocation_class", str(path))
            elif target.casefold() != ALLOCATION_TARGETS[allocation_class].casefold():
                _add(findings, "ERROR", "hierarchical_placement", f"Allocation target contradicts class for {requirement_id}: {allocation_class} -> {target}", str(path))
            if requires_hierarchy_parent(allocation_class, _text(row.get("lineage_mode"))) and not _text(row.get("hierarchy_parent_req_ids")):
                _add(findings, "ERROR", "hierarchical_placement", f"Lower-level row {requirement_id} requires hierarchy_parent_req_ids after synchronized CSV authority crosscheck and authoritative candidate remediation", str(path))
    return by_id


def validate(repo_root: Path) -> dict[str, Any]:
    findings: list[dict[str, str]] = []
    checked: list[str] = []
    active_path, active = _discover_batches(repo_root, findings)
    if active_path:
        checked.append(str(active_path))
        manifest = active_path / "merge_manifest.json"
        if not manifest.exists():
            _add(findings, "ERROR", "merge_readiness", "Active batch is missing merge_manifest.json", str(manifest))
        else:
            payload = _read_json(manifest)
            required = ("staging_batch_id", "source_revision_id")
            missing = [field for field in required if not _text(payload.get(field))]
            if not _text(payload.get("merge_readiness") or payload.get("status")):
                missing.append("merge_readiness/status")
            if missing:
                _add(findings, "ERROR", "merge_readiness", f"Merge manifest is missing: {', '.join(missing)}", str(manifest))
            if _text(payload.get("merge_readiness") or payload.get("status")).casefold() not in {"ready", "ready_for_approval", "approved", "complete", "completed"}:
                _add(findings, "ERROR", "merge_readiness", "Merge manifest does not declare a complete/ready state", str(manifest))

    mapping_path = repo_root / "artifacts/stage1_specs/architecture_mapping_preview.csv"
    workbook_path = mapping_path.with_suffix(".xlsx")
    mapping_rows = _read_csv(mapping_path) if mapping_path.exists() else []
    mapping_ids = _validate_csv_ids(findings, mapping_path, "architecture_map")
    mapping_by_id = _check_allocation(findings, mapping_rows, mapping_path, "architecture_map")
    checked.extend([str(mapping_path), str(workbook_path)])
    if workbook_path.exists():
        workbook_rows = _workbook_rows(workbook_path)
        workbook_ids, workbook_duplicates = _ids(workbook_rows)
        if not workbook_rows:
            _add(findings, "ERROR", "stage2a", "Mapping workbook cannot be read or has no review rows", str(workbook_path))
        if workbook_duplicates:
            _add(findings, "ERROR", "stage2a", f"Workbook duplicate IDs: {', '.join(workbook_duplicates[:20])}", str(workbook_path))
        if workbook_ids != mapping_ids:
            _add(findings, "ERROR", "stage2a", "Workbook and mapping CSV requirement IDs differ", str(workbook_path))
        for requirement_id, row in mapping_by_id.items():
            workbook_row = next((candidate for candidate in workbook_rows if _id(candidate) == requirement_id), None)
            if workbook_row and any(_text(workbook_row.get(field)) != _text(row.get(field)) for field in ("approved_block", "allocation_class", "owning_target", "allocation_rationale", "lineage_mode", "source_origin_req_ids", "hierarchy_parent_req_ids", "owning_domain", "review_decision")):
                _add(findings, "ERROR", "stage2a", f"Workbook/CSV approval fields differ for {requirement_id}", str(workbook_path))
    else:
        _add(findings, "ERROR", "stage2a", "Editable mapping workbook is missing", str(workbook_path))

    stage2a_path = repo_root / "artifacts/stage2_mirco_arc/requirement_to_block_traceability.csv"
    stage2a_ids = _validate_csv_ids(findings, stage2a_path, "stage2a")
    if mapping_ids and stage2a_ids and mapping_ids != stage2a_ids:
        _add(findings, "ERROR", "id_chain", "Architecture Map and Stage 2A traceability IDs differ", str(stage2a_path))

    profile_path = repo_root / "config/stage2_mirco_arc_profile.json"
    profile = _read_json(profile_path)
    checked.append(str(profile_path))
    approval = profile.get("approval", {}) if isinstance(profile.get("approval"), dict) else {}
    if _text(approval.get("status")).casefold() != "approved":
        _add(findings, "ERROR", "stage2a", "Architecture profile approval is not currently approved", str(profile_path))
    if _text(approval.get("status")).casefold() == "approved" and mapping_path.exists() and _text(approval.get("reviewed_mapping_preview_csv_sha256")) != _sha256(mapping_path):
        _add(findings, "ERROR", "stage2a", "Approved mapping CSV hash is stale", str(mapping_path))
    draft_path = repo_root / "artifacts/stage1_specs/architecture_profile_draft.json"
    draft = _read_json(draft_path)
    review_package = draft.get("review_package", {}) if isinstance(draft.get("review_package"), dict) else {}
    reviewed_draft_hash = _text(approval.get("reviewed_draft_profile_sha256"))
    current_draft_hash = _text(review_package.get("draft_profile_sha256"))
    if _text(approval.get("status")).casefold() == "approved" and reviewed_draft_hash != current_draft_hash:
        _add(findings, "ERROR", "stage2a", "Approved profile draft hash does not match the current draft review package", str(draft_path))

    database_path = repo_root / "data/canonical/canonical_store.sqlite"
    checked.append(str(database_path))
    canonical_ids: set[str] = set()
    snapshot_ids: set[str] = set()
    if not database_path.exists():
        _add(findings, "ERROR", "canonical", "Canonical SQLite database is missing", str(database_path))
    else:
        connection = sqlite3.connect(f"file:{database_path.as_posix()}?mode=ro", uri=True)
        connection.row_factory = sqlite3.Row
        try:
            projects = _rows(connection, "SELECT project_id FROM projects ORDER BY project_id")
            project_id = _text(projects[0].get("project_id")) if len(projects) == 1 else _text(_read_json(repo_root / "config/project_context.json").get("project_name"))
            if not project_id:
                _add(findings, "ERROR", "canonical", "Canonical project identity is ambiguous")
            canonical_rows = _rows(connection, "SELECT source_req_id FROM canonical_requirements WHERE project_id = ?", (project_id,)) if project_id else []
            canonical_ids = {_text(row.get("source_req_id")) for row in canonical_rows if _text(row.get("source_req_id"))}
            if mapping_ids and canonical_ids and not mapping_ids.issubset(canonical_ids):
                _add(findings, "ERROR", "id_chain", "Architecture Map contains IDs absent from canonical requirements")
            if active:
                staged = _rows(connection, "SELECT staged_id, source_req_id, source_revision_id, staging_batch_id FROM staged_requirements WHERE staging_batch_id LIKE ?", ("%" + _text(active.get("ingestion_id")) + "%",))
                if not staged:
                    _add(findings, "WARNING", "active_batch", "Canonical SQLite has no staged rows discoverable by active ingestion identity")
                revision_ids = {_text(row.get("source_revision_id")) for row in staged if _text(row.get("source_revision_id"))}
                if len(revision_ids) > 1:
                    _add(findings, "ERROR", "revision", "Active staged requirements reference multiple source_revision_id values")
                revision_rows = _rows(
                    connection,
                    "SELECT source_revision_id, revision_number, revision_date FROM source_revisions WHERE source_revision_id IN ({})".format(",".join("?" for _ in revision_ids)),
                    tuple(sorted(revision_ids)),
                ) if revision_ids else []
                if len(revision_rows) != len(revision_ids):
                    _add(findings, "ERROR", "revision", "Active staged requirements reference an unregistered source revision")
                if revision_rows and any(_text(row.get("revision_number")) != _text(active.get("source_revision")) for row in revision_rows):
                    _add(findings, "ERROR", "revision", "Active staged source revision number does not match source-review metadata")
            snapshots = _rows(connection, "SELECT snapshot_id, status, canonical_revision_set_json FROM snapshots WHERE project_id = ? ORDER BY created_at", (project_id,)) if project_id else []
            snapshot_ids = {_text(row.get("snapshot_id")) for row in snapshots}
            approved_snapshots = [row for row in snapshots if _text(row.get("status")).casefold() == "approved"]
            if not approved_snapshots:
                _add(findings, "ERROR", "stage2b", "No approved canonical snapshot is available for Stage 2B prerequisites", str(database_path))
            else:
                current_revision_rows = _rows(connection, "SELECT current_revision_id FROM canonical_requirements WHERE project_id = ? AND current_revision_id IS NOT NULL", (project_id,))
                current_revisions = {str(row["current_revision_id"]) for row in current_revision_rows}
                for snapshot in approved_snapshots:
                    try:
                        snapshot_revisions = set(json.loads(_text(snapshot.get("canonical_revision_set_json")) or "[]"))
                    except json.JSONDecodeError:
                        snapshot_revisions = set()
                    if snapshot_revisions != current_revisions:
                        _add(findings, "WARNING", "stale_state", f"Approved snapshot {_text(snapshot.get('snapshot_id'))} does not match current canonical revision set", str(database_path))
        finally:
            connection.close()

    supplementary_staged = {row.get("supplementary_staged_id") for row in mapping_rows if row.get("supplementary_staged_id")}
    if active and supplementary_staged and not any(_text(value).find(_text(active.get("ingestion_id"))) >= 0 for value in supplementary_staged):
        _add(findings, "ERROR", "stale_state", "Architecture Map supplementary staged IDs do not identify the active ingestion batch", str(mapping_path))
    if mapping_ids and stage2a_ids:
        stale = sorted(mapping_ids - stage2a_ids)
        if stale:
            _add(findings, "ERROR", "id_chain", f"Current mapping IDs missing from Stage 2A: {', '.join(stale[:20])}", str(stage2a_path))

    errors = sorted((item for item in findings if item["severity"] == "ERROR"), key=lambda item: (item["category"], item["message"], item["path"]))
    warnings = sorted((item for item in findings if item["severity"] == "WARNING"), key=lambda item: (item["category"], item["message"], item["path"]))
    infos = sorted((item for item in findings if item["severity"] == "INFO"), key=lambda item: (item["category"], item["message"], item["path"]))
    result = {
        "checker": "validate_source_to_stage2b_coherence",
        "mode": "read_only_local_deterministic",
        "phase": "Phase 1",
        "authority_fingerprint": authority_fingerprint(repo_root),
        "discovery_assumptions": [
            "The newest deterministically ordered supplementary source-review batch is active.",
            "When source_kind is absent, supplementary status is determined by comparison with config/project_context.json source_spec_path.",
            "Canonical SQLite is opened through a read-only URI and is authoritative when present.",
            "The workbook is inspected read-only and the synchronized mapping CSV is compared as the machine-readable handoff.",
            "Unresolved hierarchy-parent rows are categorized from synchronized mapping CSV authority before user escalation.",
        ],
        "active_batch": active,
        "checked_paths": sorted(set(checked)),
        "findings": errors + warnings + infos,
        "blockers": errors,
        "stale_state_findings": [item for item in errors + warnings + infos if item["category"] == "stale_state"],
        "id_chain_findings": [item for item in errors + warnings + infos if item["category"] == "id_chain"],
        "hierarchical_placement_findings": [item for item in errors + warnings + infos if item["category"] == "hierarchical_placement"],
        "stage2a_findings": [item for item in errors + warnings + infos if item["category"] == "stage2a"],
        "stage2b_findings": [item for item in errors + warnings + infos if item["category"] == "stage2b"],
        "counts": {"ERROR": len(errors), "WARNING": len(warnings), "INFO": len(infos)},
        "result": "FAIL" if errors else "PASS",
        "next_recommended_step": "Review Phase 1 blockers and obtain approval for a state-changing Phase 2 fix." if errors else "Obtain approval before any state-changing Phase 2 work.",
    }
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Read-only Source-to-Stage-2B coherence checker")
    parser.add_argument("--report", default="artifacts/validation/source_to_stage2b_coherence.json")
    args = parser.parse_args()
    repo_root = Path(__file__).resolve().parents[1]
    result = validate(repo_root)
    report_path = repo_root / args.report
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    counts = result["counts"]
    print(f"Phase 1 Source-to-Stage-2B Coherence: {result['result']}")
    print(f"ERROR={counts['ERROR']} WARNING={counts['WARNING']} INFO={counts['INFO']}")
    if result["active_batch"]:
        active = result["active_batch"]
        print(f"active_batch={active.get('ingestion_id', '')} revision={active.get('source_revision', '')} date={active.get('source_revision_date', '')}")
    print(f"report={report_path}")
    for finding in result["findings"][:12]:
        print(f"{finding['severity']}: {finding['message']}")
    return 1 if result["result"] == "FAIL" else 0


if __name__ == "__main__":
    raise SystemExit(main())
