"""Regenerate the workflow main-steps v3 presentation from current workflow policy."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.util import Inches, Pt


def _project_name(repo_root: Path) -> str:
    try:
        context = json.loads((repo_root / "config/project_context.json").read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        context = {}
    return str(context.get("project_name") or "Engineering")


def _set_title(slide, text: str) -> None:
    if slide.shapes.title is not None:
        slide.shapes.title.text = text
        return
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.2), Inches(0.65))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(30)
    paragraph.font.bold = True


def _add_bullets(slide, title: str, items: list[str]) -> None:
    _set_title(slide, title)
    box = slide.shapes.add_textbox(Inches(0.75), Inches(1.25), Inches(11.8), Inches(5.8))
    frame = box.text_frame
    frame.clear()
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(19)
        paragraph.space_after = Pt(10)


def _add_box(slide, text: str, x: float, y: float, width: float, height: float, color: tuple[int, int, int]):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(width), Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*color)
    shape.line.color.rgb = RGBColor(74, 96, 117)
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(13)
    paragraph.font.bold = True
    return shape


def _add_gate(slide, text: str, x: float, y: float):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.DIAMOND,
        Inches(x), Inches(y), Inches(0.75), Inches(0.75),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 243, 205)
    shape.line.color.rgb = RGBColor(138, 90, 18)
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(8)
    paragraph.font.bold = True
    return shape


def _connector(slide, source, target, label: str = "") -> None:
    x1 = source.left + source.width
    y1 = source.top + source.height // 2
    x2 = target.left
    y2 = target.top + target.height // 2
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = RGBColor(74, 96, 117)
    line.line.width = Pt(1.5)
    if label:
        box = slide.shapes.add_textbox((x1 + x2) // 2 - Inches(0.35), (y1 + y2) // 2 - Inches(0.25), Inches(0.7), Inches(0.35))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = label
        paragraph.font.size = Pt(10)
        paragraph.font.bold = True


def _add_stage_table(slide) -> None:
    _set_title(slide, "Main Steps: Stage 0 to Stage 5, Optional Stage 6")
    rows = [
        ("0", "Ontology baseline", "Gate 0"),
        ("1", "Source requirements extraction", "Gate 1"),
        ("2", "Formalization and mapping preview", "Gate 2"),
        ("Review", "Architecture Map Review", "Reviewer approval"),
        ("2A", "Micro-architecture mapping", "Stage 2A gate"),
        ("3", "System Requirements Specification", "Stage 3 gate"),
        ("4", "Architecture Requirements Specification", "Stage 4 gate"),
        ("5", "Detailed Requirements Specification", "Stage 5 gate"),
        ("6 optional", "Cross-project architecture comparison", "Standalone post-Stage 5 command"),
    ]
    table = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.7), Inches(1.2), Inches(11.9), Inches(5.6)).table
    for index, heading in enumerate(("Step", "Primary activity", "Decision point")):
        table.cell(0, index).text = heading
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    for row_index in range(len(rows) + 1):
        for column_index in range(3):
            cell = table.cell(row_index, column_index)
            if row_index == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(217, 237, 247)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(12 if row_index else 13)
                    run.font.bold = row_index == 0


def _add_pipeline_diagram(slide) -> None:
    _set_title(slide, "Stage-Gated Workflow")
    rows = [
        [
            ("stage", "S0\nOntology"),
            ("gate", "Gate\n0"),
            ("stage", "S1\nRequirements"),
            ("gate", "Gate\n1"),
            ("stage", "S2\nSpecs"),
            ("gate", "Gate\n2"),
        ],
        [
            ("review", "Architecture\nMap Review"),
            ("stage", "S2A\nMicro-Arch"),
            ("gate", "Gate\n2A"),
            ("stage", "S3\nSRS"),
            ("gate", "Gate\n3"),
        ],
        [
            ("stage", "S4\nARS"),
            ("gate", "Gate\n4"),
            ("stage", "S5\nDRS"),
            ("gate", "Gate\n5"),
            ("optional", "S6\nOptional Compare"),
        ],
    ]
    colors = {
        "stage": (234, 240, 246),
        "review": (255, 243, 205),
        "optional": (232, 222, 248),
    }
    all_shapes = []
    y_positions = [1.35, 2.75, 4.15]
    for row_index, row in enumerate(rows):
        row_shapes = []
        x = 0.55
        for kind, text in row:
            if kind == "gate":
                shape = _add_gate(slide, text, x, y_positions[row_index] + 0.08)
                x += 1.15
            else:
                width = 1.85 if kind == "review" else 1.45
                shape = _add_box(slide, text, x, y_positions[row_index], width, 0.9, colors[kind])
                x += width + 0.35
            row_shapes.append(shape)
            all_shapes.append(shape)
            if len(row_shapes) > 1:
                _connector(slide, row_shapes[-2], row_shapes[-1])
        if row_index > 0:
            _connector(slide, rows_last_shape, row_shapes[0])
        rows_last_shape = row_shapes[-1]

    note = slide.shapes.add_textbox(Inches(0.7), Inches(5.55), Inches(11.8), Inches(0.8))
    paragraph = note.text_frame.paragraphs[0]
    paragraph.text = "Each Stage 0-5 generation step is followed by its gate before the next stage consumes its artifacts. Architecture Map Review is the approval step before Stage 2A; Stage 6 remains optional after Gate 5."
    paragraph.font.size = Pt(16)


def _add_traceability_slide(slide) -> None:
    _set_title(slide, "System Traceability And Coverage")
    source = _add_box(slide, "Source Specification\nStage 1 source catalog", 0.6, 2.5, 2.3, 1.0, (234, 240, 246))
    srs = _add_box(slide, "SRS\nSource-to-SRS coverage", 3.8, 2.5, 2.3, 1.0, (217, 237, 247))
    ars = _add_box(slide, "ARS\nScoped coverage or N/A", 7.3, 1.35, 2.3, 1.0, (241, 243, 245))
    drs = _add_box(slide, "DRS\nScoped coverage", 7.3, 3.75, 2.3, 1.0, (217, 234, 211))
    _connector(slide, source, srs, "direct")
    _connector(slide, srs, ars, "direct")
    _connector(slide, srs, drs, "direct")
    note = slide.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(11.5), Inches(0.9))
    paragraph = note.text_frame.paragraphs[0]
    paragraph.text = "Stage 5 writes the hierarchy coverage report. The System Traceability GUI tab reads it, shows direct-link percentages on arrows, and excludes context-only rows from coverage metrics."
    paragraph.font.size = Pt(16)


def main() -> None:
    parser = argparse.ArgumentParser(description="Regenerate the workflow main-steps v3 presentation.")
    parser.add_argument(
        "--output-path",
        type=Path,
        default=None,
        help="Optional output path. Defaults to docs/workflow_main_steps_template_v3.pptx.",
    )
    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parents[1]
    template_path = repo_root / "templates" / "Presentation_template.pptx"
    output_path = args.output_path or repo_root / "docs" / "workflow_main_steps_template_v3.pptx"
    if not output_path.is_absolute():
        output_path = repo_root / output_path
    presentation = Presentation(str(template_path))

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    _set_title(title_slide, f"{_project_name(repo_root)} Workflow Main Steps")
    subtitle = title_slide.placeholders[1] if len(title_slide.placeholders) > 1 else None
    if subtitle is not None:
        subtitle.text = "Stage-gated requirements engineering workflow"

    _add_stage_table(presentation.slides.add_slide(presentation.slide_layouts[5]))
    _add_pipeline_diagram(presentation.slides.add_slide(presentation.slide_layouts[5]))
    _add_bullets(presentation.slides.add_slide(presentation.slide_layouts[1]), "Architecture Map Review", [
        "Review the editable requirement-to-block mapping preview after Stage 2.",
        "Approve or reassign every mapping with current evidence hashes and reviewer identity.",
        "Stage 2A remains blocked until all decisions are approved and approval evidence matches.",
    ])
    _add_traceability_slide(presentation.slides.add_slide(presentation.slide_layouts[5]))
    _add_bullets(presentation.slides.add_slide(presentation.slide_layouts[1]), "Core Workflow Rules", [
        "Stage 1 is the only direct reader of the source specification.",
        "Stage 2 and later consume artifacts only and enforce the independence guard.",
        "Traceability coverage joins documents through source_req_id against the Stage 1 source catalog.",
        "A scoped document with no in-scope requirements is reported as not applicable, not as zero coverage.",
        "Optional Stage 6 compares current generated architecture artifacts and does not block Stage 0-5 gates.",
    ])
    _add_bullets(presentation.slides.add_slide(presentation.slide_layouts[1]), "Operator Actions", [
        "Use workflow_cli.py to run one stage or an ascending contiguous stage range.",
        "Use Architecture Map Review to review and approve mapping evidence before Stage 2A.",
        "Use the System Traceability GUI tab to inspect the current report and direct-link coverage values.",
        "Use run --stage 6 --project-to-compare after Stage 5 when sibling-project comparison is needed.",
        "Use validate --all to confirm the available stage gates after an approved regeneration.",
    ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()