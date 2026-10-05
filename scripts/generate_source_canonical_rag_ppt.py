#!/usr/bin/env python3
"""Generate a PowerPoint presentation for the source-to-canonical RAG pipeline."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(25, 48, 76)
BLUE = RGBColor(47, 111, 159)
TEAL = RGBColor(35, 132, 135)
GREEN = RGBColor(64, 139, 91)
ORANGE = RGBColor(211, 128, 43)
RED = RGBColor(176, 76, 68)
INK = RGBColor(34, 43, 52)
MUTED = RGBColor(91, 105, 116)
PALE_BLUE = RGBColor(231, 241, 248)
PALE_TEAL = RGBColor(229, 244, 242)
PALE_ORANGE = RGBColor(250, 241, 224)
PALE_GRAY = RGBColor(244, 247, 249)
WHITE = RGBColor(255, 255, 255)

WIDE = (13.333, 7.5)


def add_text(slide, text, x, y, w, h, *, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.55, 0.28, 12.2, 0.52, size=26, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 0.86, 12.0, 0.34, size=11, color=MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.22), Inches(12.15), Inches(0.025))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()


def add_footer(slide, number):
    add_text(slide, "Source-to-Canonical RAG Pipeline", 0.58, 7.15, 5.5, 0.18, size=8, color=MUTED)
    add_text(slide, str(number), 12.35, 7.15, 0.35, 0.18, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, title, body, x, y, w, h, *, fill=PALE_GRAY, accent=BLUE, body_size=13):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = accent
    card.line.width = Pt(1.2)
    add_text(slide, title, x + 0.12, y + 0.10, w - 0.24, 0.34, size=15, color=accent, bold=True)
    add_text(slide, body, x + 0.12, y + 0.52, w - 0.24, h - 0.62, size=body_size, color=INK)
    return card


def add_flow_box(slide, label, x, y, w, h, *, fill=PALE_BLUE, accent=BLUE, size=12):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = accent
    box.line.width = Pt(1.2)
    add_text(slide, label, x + 0.06, y + 0.12, w - 0.12, h - 0.18, size=size, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    return box


def add_arrow(slide, x1, y1, x2, y2, color=MUTED):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.6)
    line.line.end_arrowhead = True
    return line


def add_bullets(slide, items, x, y, w, h, *, size=16, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
        p.bullet = True
    return box


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    add_text(slide, "From PDF Specification\nto Canonical Database and RAG", 0.75, 1.25, 11.6, 1.35, size=32, color=WHITE, bold=True)
    add_text(slide, "Deterministic, local-only requirements and retrieval workflow", 0.8, 2.9, 10.8, 0.45, size=18, color=RGBColor(205, 226, 236))
    add_flow_box(slide, "SOURCE", 0.9, 4.45, 1.55, 0.72, fill=PALE_BLUE, accent=BLUE)
    add_flow_box(slide, "EVIDENCE", 3.0, 4.45, 1.75, 0.72, fill=PALE_TEAL, accent=TEAL)
    add_flow_box(slide, "AUTHORITY", 5.35, 4.45, 1.85, 0.72, fill=PALE_ORANGE, accent=ORANGE)
    add_flow_box(slide, "GENERATED SPECS", 7.8, 4.45, 2.35, 0.72, fill=PALE_BLUE, accent=BLUE)
    add_arrow(slide, 2.48, 4.81, 2.92, 4.81, WHITE)
    add_arrow(slide, 4.8, 4.81, 5.27, 4.81, WHITE)
    add_arrow(slide, 7.25, 4.81, 7.72, 4.81, WHITE)
    add_text(slide, "Stage 0 → Stage 1 → Stage 2 → S2H → SRS / ARS / DRS / IPOS", 0.85, 6.15, 11.5, 0.35, size=15, color=RGBColor(205, 226, 236), align=PP_ALIGN.CENTER)
    add_footer(slide, 1)


def slide_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "End-to-end workflow", "The source is transformed through evidence, approval, and generation boundaries.")
    labels = [
        ("Source PDF", BLUE),
        ("S0\nBootstrap + ontology", TEAL),
        ("S1\nOCR + requirements", ORANGE),
        ("S2 / S2A\nArchitecture mapping", GREEN),
        ("S2H\nApproved snapshot", RED),
        ("SRS / ARS / DRS", BLUE),
        ("IPOS per block", TEAL),
    ]
    x = 0.55
    for index, (label, accent) in enumerate(labels):
        add_flow_box(slide, label, x, 2.85, 1.55, 1.0, fill=PALE_GRAY, accent=accent, size=12)
        if index < len(labels) - 1:
            add_arrow(slide, x + 1.58, 3.35, x + 1.82, 3.35)
        x += 1.82
    add_card(slide, "Invariant", "Stage 2+ consumes approved artifacts and snapshots. It does not reread the PDF directly.", 1.25, 5.05, 4.7, 0.95, fill=PALE_ORANGE, accent=ORANGE, body_size=14)
    add_card(slide, "Reproducibility", "Same approved snapshot + same configuration produces the same downstream result.", 7.25, 5.05, 4.7, 0.95, fill=PALE_TEAL, accent=TEAL, body_size=14)
    add_footer(slide, 2)


def slide_stage0(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Stage 0: source bootstrap and ontology", "Stage 0 establishes source identity and architectural evidence before requirements extraction.")
    add_card(slide, "Bootstrap", "• Resolve configured PDF\n• Verify readability\n• Calculate source fingerprint\n• Register project, source, revision, and ingestion batch", 0.65, 1.65, 3.65, 3.7, fill=PALE_BLUE, accent=BLUE)
    add_card(slide, "Ontology evidence", "• Entities and blocks\n• Roles and functions\n• Properties and relations\n• Interfaces and protocols\n• Analog/digital indicators\n• Tables, figures, modes, and power states", 4.85, 1.65, 3.65, 3.7, fill=PALE_TEAL, accent=TEAL)
    add_card(slide, "Gate 0", "Validates source presence, fingerprint, provenance, ontology artifacts, and blocking defects.\n\nFAIL stops the workflow. PASS allows Stage 1.", 9.05, 1.65, 3.65, 3.7, fill=PALE_ORANGE, accent=ORANGE)
    add_text(slide, "Files: s0_source_bootstrap.py · generate_stage0_ontology_outputs.py · validate_stage0_gate.py", 0.75, 6.05, 11.8, 0.35, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, 3)


def slide_ocr(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Stage 1: per-page extraction and requirement reconstruction", "The PDF is read once at the source-ingestion boundary and converted into page-addressable evidence.")
    add_flow_box(slide, "PDF\nsource", 0.65, 2.0, 1.55, 0.95, fill=PALE_BLUE, accent=BLUE)
    add_flow_box(slide, "pypdf\nPdfReader", 2.75, 2.0, 1.75, 0.95, fill=PALE_TEAL, accent=TEAL)
    add_flow_box(slide, "Page TXT\n_p001.txt...", 5.05, 2.0, 1.85, 0.95, fill=PALE_GRAY, accent=BLUE)
    add_flow_box(slide, "index.csv\npage provenance", 7.45, 2.0, 1.95, 0.95, fill=PALE_ORANGE, accent=ORANGE)
    add_flow_box(slide, "Requirement\nsummary", 9.95, 2.0, 2.0, 0.95, fill=PALE_TEAL, accent=TEAL)
    for x in (2.25, 4.6, 7.0, 9.5):
        add_arrow(slide, x, 2.48, x + 0.42, 2.48)
    add_bullets(slide, [
        "Configured ID and Requirement-marker recognition",
        "Normative language, limits, ranges, conditions, and table-row rules",
        "Continuation reconstruction across lines and pages",
        "Page, line, table, section, and source provenance",
        "Gate 1 blocks duplicate IDs, truncation, and missing coverage",
    ], 1.05, 4.15, 11.1, 1.8, size=15)
    add_footer(slide, 4)


def slide_rag_build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "RAG construction: deterministic indexing", "The RAG index is local, rebuildable, and separate from workflow authority.")
    add_card(slide, "Chunking", "Sliding-window chunks\n\nCurrent documented configuration:\n• 1200-character chunk size\n• 150-character overlap\n• Structural table and figure chunks", 0.55, 1.55, 2.85, 4.4, fill=PALE_BLUE, accent=BLUE, body_size=13)
    add_card(slide, "Normalization", "Algorithms:\n• _tokenize_terms\n• _simple_lemma\n• _controlled_terms\n\nAdds lemmas, acronyms, phrase mappings, units, and approved synonyms.", 3.75, 1.55, 2.85, 4.4, fill=PALE_TEAL, accent=TEAL, body_size=13)
    add_card(slide, "SQLite schema", "chunks\n\nchunk_fts\nOriginal text in SQLite FTS5\n\nchunk_terms_fts\nNormalized terms in SQLite FTS5", 6.95, 1.55, 2.85, 4.4, fill=PALE_ORANGE, accent=ORANGE, body_size=13)
    add_card(slide, "Manifest", "Records source, chunking, vocabulary, normalized-term index, and retrieval configuration.\n\nRAG output is derived and can be rebuilt.", 10.15, 1.55, 2.65, 4.4, fill=PALE_GRAY, accent=GREEN, body_size=13)
    add_footer(slide, 5)


def slide_rag_query(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "RAG querying: BM25, normalized retrieval, and RRF", "Three retrieval paths provide exact matching, recall expansion, and stable rank fusion.")
    add_flow_box(slide, "Query", 0.6, 2.65, 1.25, 0.72, fill=PALE_GRAY, accent=BLUE)
    add_flow_box(slide, "Lexical\nFTS5 + BM25", 2.45, 1.65, 2.05, 0.9, fill=PALE_BLUE, accent=BLUE)
    add_flow_box(slide, "Normalized\nFTS5 + BM25", 2.45, 3.55, 2.05, 0.9, fill=PALE_TEAL, accent=TEAL)
    add_flow_box(slide, "RRF fusion\nk = 60", 5.65, 2.55, 1.9, 0.9, fill=PALE_ORANGE, accent=ORANGE)
    add_flow_box(slide, "Top-k chunks\n+ provenance", 8.7, 2.55, 2.2, 0.9, fill=PALE_TEAL, accent=TEAL)
    add_arrow(slide, 1.9, 3.0, 2.35, 2.1)
    add_arrow(slide, 1.9, 3.0, 2.35, 4.0)
    add_arrow(slide, 4.55, 2.1, 5.55, 2.85)
    add_arrow(slide, 4.55, 4.0, 5.55, 3.15)
    add_arrow(slide, 7.6, 3.0, 8.6, 3.0)
    add_text(slide, "RRF(d) = Σ 1 / (k + rankᵢ(d))", 3.75, 5.45, 5.9, 0.45, size=21, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Optional local semantic retrieval may widen or recover candidates; semantic similarity alone is never authoritative.", 1.2, 6.15, 10.9, 0.35, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, 6)


def slide_canonical(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Canonical database: workflow authority", "The canonical SQLite store owns review state, revisions, provenance, mappings, and approvals.")
    steps = [
        ("staged_requirements", PALE_BLUE, BLUE),
        ("canonical_requirements", PALE_TEAL, TEAL),
        ("canonical_requirement_revisions", PALE_ORANGE, ORANGE),
        ("requirement_provenance", PALE_GRAY, GREEN),
    ]
    x = 0.65
    for index, (label, fill, accent) in enumerate(steps):
        add_flow_box(slide, label, x, 2.25, 2.55, 0.9, fill=fill, accent=accent, size=11)
        if index < len(steps) - 1:
            add_arrow(slide, x + 2.58, 2.7, x + 2.92, 2.7)
        x += 3.15
    add_card(slide, "Stored authority", "Source IDs · revisions · classifications · mapping approvals · workflow states · audit events", 1.0, 4.3, 5.1, 1.15, fill=PALE_BLUE, accent=BLUE, body_size=14)
    add_card(slide, "Not authority", "RAG SQLite, CSV, JSON, Markdown, XLSX, and DOCX are derived, review, or compatibility artifacts.", 7.2, 4.3, 5.1, 1.15, fill=PALE_ORANGE, accent=ORANGE, body_size=14)
    add_text(slide, "Database: data/canonical/canonical_store.sqlite", 2.5, 6.35, 8.3, 0.35, size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)


def slide_snapshot(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "S2H: approved immutable snapshot", "The snapshot freezes the exact approved input used by downstream generators.")
    add_card(slide, "Snapshot material", "• Canonical revision set\n• Approved mapping set\n• Vocabulary and taxonomy hashes\n• Architecture profile hashes\n• Retrieval configuration hash\n• Semantic flags", 0.8, 1.7, 3.65, 3.9, fill=PALE_BLUE, accent=BLUE)
    add_card(slide, "Resolver", "approved_snapshot_resolver.py\n\n--snapshot-id <id>\nor\n--use-latest-approved\n\nRejects incomplete, impacted, or unapproved snapshots.", 4.85, 1.7, 3.65, 3.9, fill=PALE_ORANGE, accent=ORANGE)
    add_card(slide, "Downstream contract", "SRS / ARS / DRS / Digital IPOS / Analog IPOS consume the approved snapshot and approved architectural evidence.", 8.9, 1.7, 3.65, 3.9, fill=PALE_TEAL, accent=TEAL)
    add_text(slide, "Same snapshot + same configuration = reproducible generation", 1.5, 6.25, 10.3, 0.4, size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 8)


def slide_generation(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Downstream generation and block-level IPOS", "All generated specifications preserve upstream provenance and remain downstream of snapshot authority.")
    add_flow_box(slide, "Approved\nsnapshot", 0.7, 2.35, 1.7, 0.9, fill=PALE_ORANGE, accent=ORANGE)
    add_flow_box(slide, "SRS\nSystem", 3.0, 1.35, 1.55, 0.82, fill=PALE_BLUE, accent=BLUE)
    add_flow_box(slide, "ARS\nAnalog", 3.0, 2.55, 1.55, 0.82, fill=PALE_TEAL, accent=TEAL)
    add_flow_box(slide, "DRS\nDigital", 3.0, 3.75, 1.55, 0.82, fill=PALE_BLUE, accent=BLUE)
    add_flow_box(slide, "Digital IPOS\nper block", 6.1, 2.0, 2.05, 0.9, fill=PALE_TEAL, accent=TEAL)
    add_flow_box(slide, "Analog IPOS\nper block", 9.1, 2.0, 2.05, 0.9, fill=PALE_ORANGE, accent=ORANGE)
    add_arrow(slide, 2.48, 2.8, 2.9, 1.75)
    add_arrow(slide, 2.48, 2.8, 2.9, 2.95)
    add_arrow(slide, 2.48, 2.8, 2.9, 4.15)
    add_arrow(slide, 4.65, 2.95, 6.0, 2.45)
    add_arrow(slide, 4.65, 2.95, 9.0, 2.45)
    add_card(slide, "Output locations", "artifacts/stage6_digital_ipos/blocks/<block-slug>/\n\nartifacts/stage7_analog_ipos/blocks/<block-slug>/", 2.0, 5.0, 9.4, 1.0, fill=PALE_GRAY, accent=GREEN, body_size=14)
    add_footer(slide, 9)


def slide_authority(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Authority boundaries", "Retrieval can find evidence; only approved workflow state can authorize generation.")
    add_card(slide, "RAG can", "• Search chunks\n• Return page and source provenance\n• Support crosschecks\n• Improve recall\n• Provide review evidence", 0.8, 1.65, 3.7, 3.9, fill=PALE_TEAL, accent=TEAL)
    add_card(slide, "RAG cannot", "• Approve a requirement\n• Change canonical text\n• Change mapping authority\n• Create an approved snapshot\n• Change coverage denominators", 4.8, 1.65, 3.7, 3.9, fill=PALE_ORANGE, accent=ORANGE)
    add_card(slide, "Generation requires", "• Approved snapshot\n• Approved mapping\n• Current profiles\n• Valid provenance\n• Passing gates and crosschecks", 8.8, 1.65, 3.7, 3.9, fill=PALE_BLUE, accent=BLUE)
    add_text(slide, "Deterministic · local-only · auditable · reproducible", 1.5, 6.25, 10.3, 0.4, size=21, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 10)


def slide_files(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Key files and commands", "The workflow is implemented by local Python scripts and reviewable artifacts.")
    add_card(slide, "Source and Stage 0", "s0_source_bootstrap.py\ngenerate_stage0_ontology_outputs.py\nvalidate_stage0_gate.py", 0.65, 1.45, 3.0, 2.05, fill=PALE_BLUE, accent=BLUE, body_size=13)
    add_card(slide, "Stage 1 and RAG", "extract_requirements_ocr.py\ngenerate_stage1_requirements.py\nbuild_rag_index.py\nquery_rag_index.py", 3.9, 1.45, 3.0, 2.05, fill=PALE_TEAL, accent=TEAL, body_size=13)
    add_card(slide, "Authority", "canonical_store.py\nmerge_engine.py\nsnapshot_manager.py\napproved_snapshot_resolver.py", 7.15, 1.45, 2.6, 2.05, fill=PALE_ORANGE, accent=ORANGE, body_size=13)
    add_card(slide, "Specifications", "run_srs_gen_spec_agent.py\nrun_ars_gen_spec_agent.py\nrun_drs_gen_spec_agent.py\ngenerate_ipos_specs.py", 10.0, 1.45, 2.7, 2.05, fill=PALE_GRAY, accent=GREEN, body_size=13)
    add_text(slide, "Main commands", 0.75, 4.15, 2.0, 0.3, size=17, color=NAVY, bold=True)
    add_bullets(slide, [
        "python scripts/run_stage0_gate0.py",
        "python scripts/run_stage1_requirements_gate1.py",
        "python scripts/build_rag_index.py",
        "python scripts/query_rag_index.py --query \"shall OR must\" --mode hybrid --top-k 10",
        "python scripts/generate_ipos_specs.py --kind digital --use-latest-approved",
        "python scripts/generate_ipos_specs.py --kind analog --use-latest-approved",
    ], 0.85, 4.55, 11.2, 1.65, size=14)
    add_footer(slide, 11)


def generate(output: Path, template: Path | None = None) -> Path:
    prs = Presentation(str(template)) if template and template.exists() else Presentation()
    prs.slide_width = Inches(WIDE[0])
    prs.slide_height = Inches(WIDE[1])
    while prs.slides:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]
    for builder in (
        slide_title,
        slide_pipeline,
        slide_stage0,
        slide_ocr,
        slide_rag_build,
        slide_rag_query,
        slide_canonical,
        slide_snapshot,
        slide_generation,
        slide_authority,
        slide_files,
    ):
        builder(prs)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate the source-to-canonical RAG pipeline presentation.")
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "docs" / "source-to-canonical-rag-pipeline.pptx")
    parser.add_argument("--template", type=Path, default=Path(__file__).resolve().parents[1] / "templates" / "Presentation_template.pptx")
    args = parser.parse_args()
    print(generate(args.output, args.template))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
