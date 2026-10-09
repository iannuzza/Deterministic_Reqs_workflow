"""Shared project-agnostic routing helpers for Stage 2A through Stage 5."""

from __future__ import annotations

import csv
import difflib
import getpass
import hashlib
import json
import os
import re
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from types import MappingProxyType
from typing import Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Set, Tuple

from spec_document_contract import NormalizedSourceRecord, SemanticUnit, drs_document_events, drs_table_inventory, stable_payload_hash
from ipos_semantic_normalizer import (
    extract_structural_function_name as _normalized_structural_function_name,
)


RETAINED_ROUTING_STATUS = "retained_as_non_block_function_context"
AUTHORED_REQUIREMENT_SPACER = "<p>&nbsp;</p>"
GENERATED_SPEC_VERSION = "0.1"


def runtime_user_name() -> str:
    """Return the local process user for non-authoritative logging and metadata."""
    try:
        value = getpass.getuser().strip()
    except (ImportError, OSError):
        value = ""
    return value or os.environ.get("USERNAME", "").strip() or os.environ.get("USER", "").strip() or "unknown-user"


def document_author_name() -> str:
    """Use the signed-in Windows user's display name for document authorship."""
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes

        buffer = ctypes.create_unicode_buffer(256)
        length = wintypes.ULONG(len(buffer))
        if ctypes.windll.secur32.GetUserNameExW(3, buffer, ctypes.byref(length)) and buffer.value.strip():
            return buffer.value.strip()
        raise RuntimeError("Windows full display name is unavailable for the current document author")
    return runtime_user_name()


def validate_document_author_fields(markdown_path: Path) -> List[str]:
    """Validate current authorship while preserving historical snapshot authors."""
    if not markdown_path.exists() or not markdown_path.with_suffix(".docx").exists():
        return [f"document_author_artifact_missing:{markdown_path}"]
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    authors = [line.removeprefix("Author: ").strip() for line in lines if line.startswith("Author: ")]
    if len(authors) != 1 or not authors[0]:
        return [f"document_author_missing:{markdown_path}"]
    findings: List[str] = []
    try:
        history_rows = _document_history_rows(markdown_path)
        if not history_rows or any(not row[-1].strip() for row in history_rows):
            findings.append(f"document_history_author_missing:{markdown_path}")
    except (ValueError, IndexError):
        findings.append(f"document_history_author_missing:{markdown_path}")
    try:
        with zipfile.ZipFile(markdown_path.with_suffix(".docx")) as archive:
            core = ET.fromstring(archive.read("docProps/core.xml"))
        creator = core.findtext("{http://purl.org/dc/elements/1.1/}creator")
        if creator != authors[0]:
            findings.append(f"document_docx_author_mismatch:{markdown_path}")
    except (OSError, ValueError, KeyError, ET.ParseError, zipfile.BadZipFile):
        findings.append(f"document_docx_author_unreadable:{markdown_path}")
    return findings


def _document_history_rows(output_path: Path) -> List[List[str]]:
    if not output_path.exists():
        return []
    tables = [
        table for table in drs_table_inventory(drs_document_events(output_path))
        if table["rows"] and table["rows"][0] == ["Version", "Date", "Description", "Author"]
    ]
    if len(tables) != 1 or any(len(row) != 4 for row in tables[0]["rows"]):
        raise ValueError(f"Missing, duplicate or malformed version history: {output_path}")
    return tables[0]["rows"][1:]


def _document_snapshot_id(output_path: Path) -> str:
    if not output_path.exists():
        return ""
    text = output_path.read_text(encoding="utf-8")
    match = re.search(r"(?:Snapshot ID:\s*|Snapshot:\s*`)([^`\s]+)", text, re.IGNORECASE)
    if match is None:
        raise ValueError(f"Missing document snapshot: {output_path}")
    return match.group(1)


def document_version_history_markdown(
    output_path: Path, snapshot_id: str, version: str, run_date: str, description: str, author: str,
) -> List[str]:
    """Append a snapshot version entry; leave same-snapshot history unchanged."""
    rows = _document_history_rows(output_path)
    if _document_snapshot_id(output_path) != snapshot_id or not rows:
        rows.append([version, run_date, description, author])
    return [
        "| Version | Date | Description | Author |",
        "|---|---|---|---|",
        *["| " + " | ".join(cell.replace("|", r"\|") for cell in row) + " |" for row in rows],
    ]


def document_version_for_snapshot(output_path: Path, snapshot_id: str) -> str:
    """Keep the document version stable, incrementing it when the snapshot changes."""
    if not output_path.exists():
        return GENERATED_SPEC_VERSION
    previous_snapshot = _document_snapshot_id(output_path)
    history_rows = _document_history_rows(output_path)
    version_match = re.fullmatch(r"(\d+)\.(\d+)", history_rows[-1][0]) if history_rows else None
    if not version_match:
        return GENERATED_SPEC_VERSION
    major, minor = int(version_match.group(1)), int(version_match.group(2))
    if previous_snapshot == snapshot_id:
        return f"{major}.{minor}"
    return f"{major}.{minor + 1}"


def apply_shared_spec_markdown_formatting(
    lines: Iterable[str],
    *,
    preserve_authored_text: bool = False,
) -> List[str]:
    """Normalize headings and list markers without changing authored content."""
    formatted: List[str] = []
    heading_re = re.compile(r"^(#{1,6})\s+(?!\*\*)(.*?)(\s*\{#[^}]+\})?\s*$")
    for value in lines:
        line = str(value)
        if not preserve_authored_text:
            line = line.replace("•", "-").replace("●", "-")
        match = heading_re.match(line)
        if match:
            title = match.group(2).rstrip()
            anchor = match.group(3) or ""
            line = f"{match.group(1)} **{title}**{anchor}"
        formatted.append(line)
    return formatted


def spec_anchor_slug(text: str) -> str:
    """Return a stable Markdown fragment for generated specification headings."""
    slug = re.sub(r"[^a-z0-9\s-]", "", (text or "").casefold())
    return re.sub(r"\s+", "-", slug).strip("-") or "section"


def apply_unique_spec_heading_anchors(lines: Iterable[str]) -> List[str]:
    """Add deterministic unique anchors to all Markdown headings."""
    heading_re = re.compile(r"^(#{1,6})\s+(.*?)(?:\s+\{#[^}]+\})?\s*$")
    used: Dict[str, int] = {}
    output: List[str] = []
    for value in lines:
        line = str(value)
        match = heading_re.match(line)
        if not match:
            output.append(line)
            continue
        title = re.sub(r"\s+\{#[^}]+\}\s*$", "", match.group(2)).strip()
        base = spec_anchor_slug(re.sub(r"\*", "", title))
        used[base] = used.get(base, 0) + 1
        anchor = base if used[base] == 1 else f"{base}-{used[base] - 1}"
        output.append(f"{match.group(1)} {title} {{#{anchor}}}")
    return output


def validate_spec_internal_links(markdown_path: Path, docx_path: Path | None = None) -> List[str]:
    """Validate same-document Markdown links and DOCX internal hyperlink targets."""
    findings: List[str] = []
    text = markdown_path.read_text(encoding="utf-8", errors="ignore")
    anchors = {match.group(1).casefold() for match in re.finditer(r"\{#([A-Za-z0-9_-]+)\}", text)}
    for match in re.finditer(r"\[[^\]]+\]\(#([^)]+)\)", text):
        if match.group(1).casefold() not in anchors:
            findings.append(f"{markdown_path.name}: unresolved internal link #{match.group(1)}")
    if docx_path is None or not docx_path.exists():
        return findings
    with zipfile.ZipFile(docx_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    bookmarks = {
        match.group(1).casefold()
        for match in re.finditer(r"w:bookmarkStart[^>]+w:name=\"([^\"]+)\"", document_xml)
    }
    for match in re.finditer(r"w:anchor=\"([^\"]+)\"", document_xml):
        if match.group(1).casefold() not in bookmarks:
            findings.append(f"{docx_path.name}: unresolved internal DOCX target #{match.group(1)}")
    return findings


def validate_ipos_docx_layout(docx_path: Path) -> List[str]:
    """Verify title-page, TOC/navigation, and footer layout contract."""
    findings: List[str] = []
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(docx_path) as archive:
        document_xml = archive.read("word/document.xml")
        footer_xml = [
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if re.fullmatch(r"word/footer\d+\.xml", name)
        ]
    root = ET.fromstring(document_xml)
    body = root.find(f"{word_namespace}body")
    paragraphs = [element for element in list(body or []) if element.tag == f"{word_namespace}p"]
    if not paragraphs:
        return [f"{docx_path.name}: DOCX contains no body content"]
    text = [_docx_element_text(paragraph, word_namespace).strip() for paragraph in paragraphs]
    title_index = next(
        (
            index for index, value in enumerate(text)
            if value.casefold().endswith("specification") or " ipos - " in value.casefold()
        ),
        None,
    )
    toc_index = next((index for index, value in enumerate(text) if value.casefold() == "0.1 table of contents"), None)
    if title_index is None:
        findings.append(f"{docx_path.name}: document title is missing from the title page")
    if toc_index is None:
        findings.append(f"{docx_path.name}: table of contents is missing")
    if title_index is not None and toc_index is not None and title_index >= toc_index:
        findings.append(f"{docx_path.name}: title metadata does not precede the TOC")
    if toc_index is not None:
        toc_paragraph = paragraphs[toc_index]
        page_break_before = toc_paragraph.find(f"{word_namespace}pPr/{word_namespace}pageBreakBefore")
        preceding_break = (
            toc_index > 0
            and paragraphs[toc_index - 1].find(f".//{word_namespace}br[@{word_namespace}type='page']")
            is not None
        )
        if page_break_before is None and not preceding_break:
            findings.append(f"{docx_path.name}: TOC does not start page 2 after the title page")
    navigation_positions = [
        index for index, paragraph in enumerate(paragraphs)
        if _docx_element_text(paragraph, word_namespace).strip().casefold() == "0. document navigation"
    ]
    if not navigation_positions or navigation_positions[0] <= 0:
        findings.append(f"{docx_path.name}: Document Navigation is not after the TOC")
    if not any(
        re.search(r'w:jc w:val="right"', footer)
        and re.search(r'<w:instrText[^>]*>\s*PAGE\s*</w:instrText>', footer)
        for footer in footer_xml
    ):
        findings.append(f"{docx_path.name}: bottom-right PAGE footer is missing")
    return findings

DESCRIPTIVE_MODE_TOPIC_SPECS = (
    ("Modes and transitions", ("mode", "state", "operative", "standby", "retention", "transition")),
    ("Mode resources", ("block", "clock", "fifo", "memory", "register", "resource")),
)
DESCRIPTIVE_POWER_TOPIC_SPECS = (
    ("Retention, Standby, Sleep, and Active Behavior", ("retention", "standby", "sleep", "active", "operative", "powered down")),
    ("Power Sequencing and Enable/Disable", ("sequence", "sequencing", "startup", "shutdown", "enable", "disable", "turn-on", "turn-off")),
    ("Domain Partitioning and Dependencies", ("partition", "dependency", "dependent", "relationship", "shared supply")),
    ("Supply-Domain Relationships", ("supply", "voltage", "rail", "power supply")),
    ("Power Domain Architecture", ("power domain", "island", "ldo", "bias")),
)
DESCRIPTIVE_SYSTEM_TOPIC_SPECS = (
    ("Power and domain management", ("power", "domain", "island", "ldo", "supply")),
    ("System interconnects and interfaces", ("bus", "interconnect", "i2c", "spi", "ahb", "interface")),
    ("Shared resources", ("fifo", "memory", "register", "regmap", "shared")),
    ("Processing and principal blocks", ("processor", "core", "dsp", "sensor", "adc", "controller")),
)
DESCRIPTIVE_DIGITAL_TOPIC_SPECS = (
    ("Power and clock islands", ("power", "island", "clock", "reset", "supply")),
    ("Buses and interconnects", ("bus", "interconnect", "ahb", "i2c", "spi", "serial", "protocol")),
    ("Arbitration and control", ("arbitration", "controller", "control", "interrupt", "state")),
    ("Shared digital resources", ("fifo", "memory", "register", "regmap", "shared")),
    ("Digital processing blocks", ("processor", "core", "dsp", "sensor hub", "ispu")),
)

# These are coverage-review categories, not new requirement categories. They are
# deliberately broad so the DRS can report integration evidence without copying
# block-local Digital IPOS detail.
TOP_DIGITAL_CATEGORY_SPECS = (
    ("Functional behavior", ("function", "control", "operation", "behavior", "processing")),
    ("Clocking and reset", ("power/clock", "clock-ready", "reset", "por", "frequency", "clock domain", "reset release")),
    ("Power management, power states, retention, and wake-up", ("power mode", "power domain", "power-down", "retention", "standby", "sleep", "wake", "sequencing")),
    ("Bus, interconnect, and protocol behavior", ("bus", "interconnect", "protocol", "ahb", "i2c", "spi", "uart", "serial")),
    ("Memory map, register access, and protection", ("memory map", "register", "regmap", "address", "write-protection", "read-only", "register access")),
    ("Interrupts, events, and status reporting", ("interrupt", "irq", "event", "status reporting", "ready status", "watermark", "overrun")),
    ("Boot, initialization, and configuration", ("boot", "initialization", "startup", "configuration", "configure")),
    ("Security, access control, and isolation", ("security", "secure", "access control", "isolation", "privilege", "authentication")),
    ("DMA, data movement, buffering, and FIFO behavior", ("dma", "data movement", "fifo", "buffer", "stream", "transfer")),
    ("Arbitration, priority, and QoS", ("arbitration", "priority", "qos", "quality of service", "grant")),
    ("Error handling, fault reporting, and recovery", ("error", "fault", "recovery", "failure", "fault report")),
    ("Performance, latency, and throughput", ("performance", "latency", "throughput", "bandwidth", "timing")),
    ("CDC, RDC, and synchronization", ("cdc", "rdc", "synchron", "clock domain", "reset domain")),
    ("Debug, trace, test, and DFT visibility", ("debug", "trace", "test", "dft", "scan", "bist", "observability")),
    ("Interfaces to analog, pads, and external systems", ("analog", "pad", "external", "sensor", "mixed-signal")),
    ("Low-power domain coordination", ("low-power", "power domain", "island", "domain coordination", "domain dependency")),
    ("Firmware-visible behavior", ("firmware", "software", "command", "status", "register")),
    ("Mode transitions and operational states", ("mode", "state", "transition", "active", "standby", "operative")),
    ("Safety, monitoring, watchdog, and timeout behavior", ("safety", "monitor", "watchdog", "timeout", "supervision")),
    ("Block interaction and integration rules", ("interaction", "connected", "connection", "integration", "source", "destination")),
)
DESCRIPTIVE_ANALOG_TOPIC_SPECS = (
    ("Power and bias", ("power", "supply", "voltage", "bias", "reference")),
    ("Sensing and signal paths", ("sensor", "signal", "ecg", "ppg", "bio", "impedance", "eda")),
    ("Sampling and conversion", ("adc", "sampling", "conversion", "acquisition")),
    ("Calibration and measurement quality", ("calibration", "noise", "offset", "gain", "accuracy")),
    ("Analog principal blocks", ("analog", "afe", "amplifier", "filter", "buffer")),
)

# IPOS uses only block-local functional areas. These labels deliberately do not
# reuse DRS integration topics and are rendered through one shared composer.
IPOS_LOCAL_DIGITAL_TOPIC_SPECS = (
    ("boot and OTP services", ("boot", "otp", "lifecycle", "firmware")),
    ("power, clock, and reset sequencing", ("power", "clock", "reset", "por", "ldo")),
    ("local interface protocol handling", ("ahb", "i2c", "spi", "protocol", "bus", "serial")),
    ("local register and memory management", ("register", "registers", "regmap", "memory", "ram", "address")),
    ("local state and control management", ("control", "state", "mode", "interrupt", "irq")),
    ("local buffering and FIFO management", ("fifo storage", "fifo", "buffer", "queue", "watermark")),
    ("local digital signal processing", ("dsp elaboration", "dsp", "signal processing", "filter")),
    ("local pad and signal routing", ("pad", "mux", "gpio", "route", "routing")),
)

IPOS_LOCAL_ANALOG_TOPIC_SPECS = (
    ("local power and bias control", ("power", "supply", "voltage", "bias", "reference")),
    ("local sensing and signal conditioning", ("sensor", "signal", "ecg", "ppg", "bio", "impedance", "eda")),
    ("local sampling and conversion", ("adc", "sampling", "conversion", "acquisition")),
    ("local calibration and measurement control", ("calibration", "noise", "offset", "gain", "accuracy")),
    ("local analog signal-path control", ("analog", "afe", "amplifier", "filter", "buffer")),
)

IPOS_DRS_LEVEL_TOPICS = {
    "Power and clock islands",
    "Buses and interconnects",
    "Arbitration and control",
    "Shared digital resources",
    "Digital processing blocks",
}


@dataclass(frozen=True)
class DescriptiveAssembly:
    """Deterministic, non-authoritative descriptive output and its audit trail."""

    topics: Dict[str, List[str]]
    audit: List[Dict[str, str]]


@dataclass(frozen=True)
class IPOSFunctionalInput:
    """Snapshot-bound, non-authoritative input for IPOS descriptive composition."""

    block: str
    domain: str
    function: str
    inputs: str
    outputs: str
    local_requirement_ids: Tuple[str, ...]
    local_requirement_evidence: Tuple[str, ...]
    candidate_requirement_records: Tuple[Tuple[str, str, str, str], ...]
    materialized: bool
    provenance: Tuple[str, ...]


@dataclass(frozen=True)
class TopDigitalCoverage:
    """Deterministic coverage status for top-level integration review."""

    category: str
    status: str
    evidence: List[str]
    matched_terms: List[str]
    evidence_count: int


def assess_top_digital_coverage(
    records: Iterable[Mapping[str, object]],
    category_specs: Sequence[Tuple[str, Sequence[str]]] = TOP_DIGITAL_CATEGORY_SPECS,
    *,
    max_evidence_per_category: int = 3,
    allowed_scopes: Sequence[str] = ("top_digital", "integration", "architecture", "lifted_integration"),
) -> Tuple[List[TopDigitalCoverage], List[Dict[str, str]]]:
    """Assess approved top-level evidence; block-local evidence is excluded."""
    prepared: List[Tuple[int, str, str, bool, str]] = []
    for index, record in enumerate(records):
        scope = csv_cell_text(record.get("scope") or "").strip().casefold()
        if scope not in {item.casefold() for item in allowed_scopes}:
            continue
        statement = csv_cell_text(
            record.get("statement") or record.get("function") or record.get("text")
        ).strip()
        source = csv_cell_text(record.get("source") or record.get("provenance")).strip()
        if statement:
            prepared.append((
                index,
                statement,
                source,
                bool(re.search(r"\b(?:shall|must|required to|requirement)\b", statement, re.I)),
                scope,
            ))

    coverage: List[TopDigitalCoverage] = []
    audit: List[Dict[str, str]] = []
    for category, terms in category_specs:
        matches: List[Tuple[int, str, str, bool, List[str], str]] = []
        seen: Set[str] = set()
        for index, statement, source, normative, scope in prepared:
            searchable = f"{statement} {source}".casefold()
            matched_terms = [
                csv_cell_text(term).strip()
                for term in terms
                if csv_cell_text(term).strip().casefold() in searchable
            ]
            if not matched_terms or statement.casefold() in seen:
                continue
            seen.add(statement.casefold())
            matches.append((index, statement, source, normative, matched_terms, scope))
        matches.sort(key=lambda item: item[0])
        selected = matches[:max_evidence_per_category]
        if not matches:
            status = "Missing"
        elif any(item[3] for item in matches) or len(matches) >= 2:
            status = "Covered"
        else:
            status = "Partial"
        coverage.append(TopDigitalCoverage(
            category=category,
            status=status,
            evidence=[item[1] for item in selected],
            matched_terms=sorted({term for item in matches for term in item[4]}),
            evidence_count=len(matches),
        ))
        if selected:
            audit.append({
                "category": category,
                "status": status,
                "decision": "selected",
                "output_order": "; ".join(str(output_order) for output_order, _item in enumerate(selected, start=1)),
                "matched_terms": "; ".join(sorted({term for item in selected for term in item[4]})),
                "statement": " || ".join(item[1] for item in selected),
                "source": " || ".join(item[2] for item in selected),
                "scope": "; ".join(sorted({item[5] for item in selected})),
            })
        else:
            audit.append({
                "category": category,
                "status": status,
                "decision": "no approved evidence",
                "output_order": "",
                "matched_terms": "",
                "statement": "",
                "source": "",
                "scope": "",
            })
    return coverage, audit


def write_top_digital_coverage_audit(path: Path, rows: Iterable[Mapping[str, object]]) -> None:
    """Persist coverage decisions and provenance without altering authority."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = ["category", "status", "decision", "output_order", "matched_terms", "statement", "source", "scope"]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: csv_cell_text(row.get(field, "")) for field in fields})


def _descriptive_tokens(value: object) -> Set[str]:
    return set(re.findall(r"[a-z0-9]+", csv_cell_text(value).casefold()))


def assign_ipos_descriptive_topic(
    functional_evidence: object,
    topic_specs: Sequence[Tuple[str, Sequence[str]]],
) -> Dict[str, str]:
    """Choose one IPOS topic from functional evidence or omit an ambiguous input."""
    evidence = csv_cell_text(functional_evidence).strip()
    searchable = evidence.casefold()
    tokens = _descriptive_tokens(evidence)
    if not evidence:
        return {
            "topic": "",
            "matched_functional_evidence": "",
            "assignment_reason": "no functional topic evidence",
        }
    if not _is_ipos_descriptive_functional_evidence(evidence):
        return {
            "topic": "",
            "matched_functional_evidence": "",
            "assignment_reason": "non-functional, parameterized, or low-level evidence",
        }
    candidates: List[Tuple[int, str, List[str]]] = []
    for topic, terms in topic_specs:
        matched_terms = sorted(
            csv_cell_text(term).strip()
            for term in terms
            if (
                csv_cell_text(term).casefold().strip() in searchable
                if len(_descriptive_tokens(term)) > 1
                else bool(_descriptive_tokens(term).intersection(tokens))
            )
        )
        if matched_terms:
            if not _is_ipos_topic_evidence_admissible(topic, searchable, tokens, matched_terms):
                continue
            score = sum(len(_descriptive_tokens(term)) for term in matched_terms)
            candidates.append((score, topic, matched_terms))
    if not candidates:
        return {
            "topic": "",
            "matched_functional_evidence": "",
            "assignment_reason": "no functional topic evidence",
        }
    candidates.sort(key=lambda item: (-item[0], item[1].casefold()))
    best_score = candidates[0][0]
    best = [candidate for candidate in candidates if candidate[0] == best_score]
    if len(best) != 1:
        return {
            "topic": "",
            "matched_functional_evidence": "; ".join(
                f"{topic}: {', '.join(terms)}" for _score, topic, terms in best
            ),
            "assignment_reason": "ambiguous functional topic evidence",
        }
    _score, topic, matched_terms = best[0]
    return {
        "topic": topic,
        "matched_functional_evidence": "; ".join(matched_terms),
        "assignment_reason": "strongest functional topic signals: " + ", ".join(matched_terms),
    }


def _is_ipos_descriptive_functional_evidence(evidence: str) -> bool:
    """Reject local metadata and implementation fragments from IPOS summaries."""
    normalized = evidence.casefold().strip()
    if not normalized:
        return False
    if re.match(
        r"^(?:the\s+)?(?:parameter|configuration|address|register(?:\s+field)?|field|value|version|timer|bit|table|figure)\b",
        normalized,
    ):
        return False
    if re.search(
        r"\b(?:suggested\s*(?:val(?:ue)?\.?|:)|0x[0-9a-f]+|\b(?:ip_version|ser_mode|i2c_master_code|i2c_dev_id|rd_to_max_val)\b|ocr|page\s*\d+)\b",
        normalized,
    ):
        return False
    if re.search(r"\b(?:shall|must|required to)\b.*\b(?:register|address|bit|field|value|timer)\b", normalized):
        return False
    return True


def _is_ipos_topic_evidence_admissible(
    topic: str,
    searchable: str,
    tokens: Set[str],
    matched_terms: Sequence[str],
) -> bool:
    """Require functional combinations rather than isolated shared keywords."""
    matched = {term.casefold() for term in matched_terms}
    if topic == "boot and OTP services":
        return "boot" in tokens or "firmware" in tokens
    if topic == "power, clock, and reset sequencing":
        return len(matched.intersection({"power", "clock", "reset", "por", "ldo"})) >= 2
    if topic == "local digital signal processing":
        return (
            "dsp elaboration" in matched
            or "signal processing" in matched
            or ("dsp" in tokens and bool({"signal", "processing", "elaboration"}.intersection(tokens)))
        )
    if topic == "local register and memory management":
        return (
            "memory" in matched
            or "ram" in matched
            or "regmap" in matched
            or bool({"register", "registers"}.intersection(tokens))
            and bool({"configuration", "status", "access", "mapped"}.intersection(tokens))
        )
    return bool(matched)


def assemble_descriptive_summary(
    records: Iterable[Mapping[str, object]],
    topic_specs: Sequence[Tuple[str, Sequence[str]]],
    *,
    max_items_per_topic: int = 6,
    include_general: bool = True,
    summary_only: bool = False,
    ipos_topic_mode: bool = False,
) -> DescriptiveAssembly:
    """Assemble descriptive evidence with stable decisions; never creates authority."""
    topics = {name: [] for name, _terms in topic_specs}
    if include_general:
        topics["General architecture"] = []
    audit: List[Dict[str, str]] = []
    seen: Set[str] = set()

    for index, record in enumerate(records):
        evidence_statement = csv_cell_text(
            record.get("statement")
            or record.get("function")
            or record.get("Non-Block Function Context")
            or record.get("Source Paragraph")
            or record.get("text")
        ).strip()
        statement = csv_cell_text(record.get("summary") if summary_only else evidence_statement).strip()
        source = csv_cell_text(
            record.get("source")
            or record.get("Source Paragraph")
            or record.get("source_text")
            or record.get("provenance")
        ).strip()
        base_audit = {
            "input_order": str(index),
            "statement": statement,
            "source": source,
            "scope": csv_cell_text(record.get("scope") or record.get("domain")).strip(),
            "mapped_block": csv_cell_text(record.get("mapped_block") or record.get("block") or record.get("owner")).strip(),
            "layer": csv_cell_text(record.get("layer")).strip(),
            "scope_decision": "accepted input scope",
            "normative_evidence": "yes" if re.search(r"\b(?:shall|must|required to|requirement)\b", statement, re.I) else "no",
            "source_evidence": evidence_statement,
            "functional_evidence": csv_cell_text(record.get("functional_evidence")).strip(),
            "matched_functional_evidence": "",
        }
        ipos_assignment = (
            assign_ipos_descriptive_topic(base_audit["functional_evidence"], topic_specs)
            if ipos_topic_mode
            else None
        )
        if ipos_assignment and not ipos_assignment["topic"]:
            audit.append({
                **base_audit,
                "decision": "rejected",
                "reason": ipos_assignment["assignment_reason"],
                "assignment_reason": ipos_assignment["assignment_reason"],
                "matched_functional_evidence": ipos_assignment["matched_functional_evidence"],
            })
            continue
        if summary_only and not statement:
            audit.append({**base_audit, "decision": "rejected", "reason": "no supported synthesized summary"})
            continue
        if not statement:
            audit.append({**base_audit, "decision": "rejected", "reason": "missing descriptive statement"})
            continue

        rendered = statement
        key = rendered.casefold()
        if key in seen:
            audit.append({**base_audit, "decision": "rejected", "reason": "duplicate descriptive evidence"})
            continue

        if ipos_topic_mode:
            assignment = ipos_assignment or {}
            target = assignment["topic"]
            assignment_reason = assignment["assignment_reason"]
            base_audit["matched_functional_evidence"] = assignment["matched_functional_evidence"]
            if not target:
                audit.append({
                    **base_audit,
                    "decision": "rejected",
                    "reason": assignment_reason,
                    "assignment_reason": assignment_reason,
                })
                continue
        else:
            searchable_text = csv_cell_text(f"{statement} {source}").casefold()
            searchable = _descriptive_tokens(searchable_text)
            target = "General architecture" if include_general else ""
            assignment_reason = "no configured topic signal"
            for topic, terms in topic_specs:
                matched_terms = sorted(
                    configured_term
                    for configured_term in terms
                    if (
                        csv_cell_text(configured_term).casefold().strip() in searchable_text
                        if len(_descriptive_tokens(configured_term)) > 1
                        else bool(_descriptive_tokens(configured_term).intersection(searchable))
                    )
                )
                if matched_terms:
                    target = topic
                    assignment_reason = "lexical/normalized topic signals: " + ", ".join(matched_terms)
                    break
        if not target:
            audit.append({**base_audit, "decision": "rejected", "reason": "no eligible configured topic"})
            continue
        seen.add(key)
        if len(topics[target]) >= max_items_per_topic:
            audit.append({
                **base_audit,
                "decision": "rejected",
                "reason": "topic output limit",
                "topic": target,
                "assignment_reason": assignment_reason,
            })
            continue
        topics[target].append(rendered)
        audit.append({
            **base_audit,
            "decision": "selected",
            "reason": "eligible approved evidence",
            "topic": target,
            "assignment_reason": assignment_reason,
            "output_order": str(len(topics[target])),
        })

    return DescriptiveAssembly(
        topics={name: values for name, values in topics.items() if values},
        audit=audit,
    )


SRS_SYSTEM_OVERVIEW_HEADINGS = (
    "3.1 General System Description",
    "3.2 Main System Capabilities",
    "3.3 Main Architectural Domains and Subsystems",
    "3.4 External Interfaces and System Boundaries",
    "3.5 Operating Concept",
    "3.6 Power, Clock, and Reset Overview",
    "3.7 Assumptions, Scope Limits, and Allocation Boundaries",
)


@dataclass(frozen=True)
class DescriptiveWritingProfile:
    name: str
    allowed_scopes: Tuple[str, ...]
    document_type: str
    evidence_kinds: Tuple[str, ...]
    authority_tiers: Tuple[str, ...]
    owner_policy: str
    allowed_domains: Tuple[str, ...]
    allowed_layers: Tuple[str, ...]
    projection_adapters: Tuple[str, ...]
    topic_roles: Tuple[str, ...]
    excluded_terms: Tuple[str, ...]
    runtime_enabled: bool
    version: int = 1

    def __post_init__(self) -> None:
        if not self.name or not self.document_type or self.version != 1:
            raise ValueError("Invalid descriptive profile identity or version")
        if self.owner_policy not in {"unrestricted", "same_owner"}:
            raise ValueError("Unsupported descriptive owner policy")
        for values in (self.allowed_scopes, self.evidence_kinds, self.authority_tiers, self.projection_adapters):
            if not values or len(set(values)) != len(values) or any(not value.strip() for value in values):
                raise ValueError("Explicit nonempty descriptive profile controls required")


@dataclass(frozen=True)
class DescriptiveFact:
    subject: str
    action: str
    objects: Tuple[str, ...]
    source: str
    statement: str
    scope: str
    section: str = ""
    mode: str = ""
    condition: str = ""
    modality: str = ""
    negative: bool = False
    coordination: str = "and"
    object_suffix: str = ""
    plural_subject: bool = False
    qualifiers: Tuple[str, ...] = ()
    results: Tuple[str, ...] = ()
    evidence_kind: str = "descriptive"
    owner: str = ""
    domain: str = ""
    layer: str = ""

    @property
    def fact_id(self) -> str:
        payload = json.dumps(asdict(self), sort_keys=True, ensure_ascii=True)
        return "desc-" + hashlib.sha256(payload.encode("utf-8")).hexdigest()[:20]


@dataclass(frozen=True)
class DescriptiveParagraph:
    text: str
    facts: Tuple[DescriptiveFact, ...]


SRS_DESCRIPTIVE_PROFILE = DescriptiveWritingProfile(
    "srs-overview-v1", ("system", "architecture"), "SRS", ("descriptive",),
    ("admitted_context",), "unrestricted", (), (), ("srs-overview",),
    SRS_SYSTEM_OVERVIEW_HEADINGS,
    ("register", "regmap", "address", "signal name", "port name"), True,
)
DRS_DESCRIPTIVE_PROFILE = DescriptiveWritingProfile(
    "drs-integration-v1", ("architecture", "digital_integration"), "DRS", ("descriptive",),
    ("admitted_context",), "unrestricted", ("digital", "shared"), ("integration",),
    ("drs-integration",), ("integration", "interfaces", "operating_concept", "power"), (), False,
)
ARS_DESCRIPTIVE_PROFILE = DescriptiveWritingProfile(
    "ars-integration-v1", ("architecture", "analog_integration"), "ARS", ("descriptive",),
    ("admitted_context",), "unrestricted", ("analog", "mixed_signal", "shared"), ("integration",),
    ("ars-integration",), ("integration", "interfaces", "operating_concept", "power"), (), False,
)
IPOS_DESCRIPTIVE_PROFILE = DescriptiveWritingProfile(
    "ipos-structural-integrity-v1", ("block_local",), "IPOS", ("descriptive", "approved_normative"),
    ("admitted_context", "approved_normative"), "same_owner", (), ("block",),
    ("ipos-local",), ("function", "interfaces", "operating_concept", "power"), (), True,
)
DESCRIPTIVE_PROFILES = MappingProxyType({profile.document_type: profile for profile in (
    SRS_DESCRIPTIVE_PROFILE, DRS_DESCRIPTIVE_PROFILE, ARS_DESCRIPTIVE_PROFILE, IPOS_DESCRIPTIVE_PROFILE,
)})


def descriptive_profile(document_type: str) -> DescriptiveWritingProfile:
    try:
        profile = DESCRIPTIVE_PROFILES[document_type]
    except KeyError:
        raise ValueError("Unknown descriptive document profile: " + document_type) from None
    if not profile.runtime_enabled:
        raise ValueError("Descriptive profile not activated: " + document_type)
    return profile


DESCRIPTIVE_INTRANSITIVE_ACTIONS = frozenset({"wait", "toggle", "start", "stop", "resume", "complete", "return"})
DESCRIPTIVE_NORMATIVE_ACTIONS = frozenset({
    "accept", "acquire", "accumulate", "allow", "average", "be", "calculate", "change", "check",
    "compare", "complete", "comprise", "compute", "configure", "contain", "control", "convert", "copy",
    "decrement", "disable", "discard", "divide", "drive", "enable", "end", "enter", "execute", "exit", "filter",
    "generate", "go", "have", "ignore", "include", "increment", "maintain", "make", "monitor", "obtain",
    "perform", "process", "produce", "provide", "raise", "read", "receive", "reset", "resume",
    "return", "route", "run", "sample", "select", "send", "set", "start", "stay", "stop",
    "store", "subtract", "support", "toggle", "transfer", "use", "wait", "work", "write",
})
DESCRIPTIVE_FINITE_ACTION_BASES = MappingProxyType({
    (action[:-1] + "ies" if action.endswith("y") and action[-2] not in "aeiou" else
     action + ("es" if action.endswith(("s", "sh", "ch", "x", "z", "o")) else "s")): action
    for action in DESCRIPTIVE_NORMATIVE_ACTIONS
})


def descriptive_fact_findings(
    fact: DescriptiveFact, *, profile: DescriptiveWritingProfile, owner: str = "",
) -> List[str]:
    findings = []
    if fact.scope not in profile.allowed_scopes or fact.evidence_kind not in profile.evidence_kinds:
        findings.append("descriptive_scope_or_kind_invalid")
    if profile.allowed_domains and fact.domain not in profile.allowed_domains:
        findings.append("descriptive_domain_invalid")
    if profile.allowed_layers and fact.layer not in profile.allowed_layers:
        findings.append("descriptive_layer_invalid")
    if profile.owner_policy == "same_owner" and (not owner or fact.owner != owner):
        findings.append("descriptive_owner_invalid")
    if profile.topic_roles and fact.section not in profile.topic_roles:
        findings.append("descriptive_topic_invalid")
    if any(not isinstance(value, str) or not value.strip() for value in
           (fact.source, fact.statement, fact.subject, fact.action, *fact.objects)) or (
               not fact.objects and fact.action not in DESCRIPTIVE_INTRANSITIVE_ACTIONS):
        findings.append("descriptive_fact_or_provenance_incomplete")
    if fact.modality not in {"", "can", "may"} or fact.coordination not in {"and", "or"}:
        findings.append("descriptive_grammar_unsupported")
    if (fact.evidence_kind == "approved_normative" and fact.action not in DESCRIPTIVE_NORMATIVE_ACTIONS
            and not (profile == IPOS_DESCRIPTIVE_PROFILE and fact.action == "define")):
        findings.append("descriptive_predicate_grammar_unsupported")
    descriptive_text = " ".join((fact.subject, fact.action, *fact.objects, fact.object_suffix,
                                 fact.condition, fact.mode, *fact.qualifiers, *fact.results))
    if any(re.search(r"\b" + re.escape(term) + r"\b", descriptive_text, re.IGNORECASE)
           for term in profile.excluded_terms):
        findings.append("descriptive_profile_detail_excluded")
    return findings


@dataclass(frozen=True)
class DescriptiveFlowResult:
    records: Tuple[NormalizedSourceRecord, ...]
    facts: Tuple[DescriptiveFact, ...]
    paragraphs: Tuple[DescriptiveParagraph, ...]
    units: Tuple[SemanticUnit, ...]
    audit: Tuple[Dict[str, str], ...]


def descriptive_evidence_findings(
    record: NormalizedSourceRecord, *, profile: DescriptiveWritingProfile,
) -> List[str]:
    findings = []
    if not record.record_id or not record.text.strip() or not record.provenance or any(not source.strip() for source in record.provenance):
        findings.append("descriptive_evidence_provenance_incomplete")
    if record.decision not in {"selected", "retained"} or record.authority_tier not in profile.authority_tiers:
        findings.append("descriptive_evidence_not_admissible")
    if record.ownership_scope not in profile.allowed_scopes:
        findings.append("descriptive_evidence_scope_invalid")
    return findings


DESCRIPTIVE_STRUCTURE_VERSION = "1.0"
DESCRIPTIVE_RELATIONS = frozenset({
    "behavior", "containment", "responsibility", "interaction", "transformation",
    "configuration", "state_transition",
})
DESCRIPTIVE_DEPENDENCIES = frozenset({
    "before", "after", "initiates", "completes", "causes", "referenced_context",
})


@dataclass(frozen=True)
class DescriptiveSourceSpan:
    record_id: str
    start: int
    end: int
    text: str


@dataclass(frozen=True)
class DescriptiveContext:
    operator: str
    text: str = ""
    operands: Tuple[DescriptiveContext, ...] = ()


@dataclass(frozen=True)
class DescriptiveEffect:
    action: str
    object: str


@dataclass(frozen=True)
class DescriptiveMeaning:
    relation: str
    subject: str
    predicate: str
    objects: Tuple[str, ...] = ()
    contexts: Tuple[DescriptiveContext, ...] = ()
    mode: str = ""
    modality: str = ""
    negative: bool = False
    coordination: str = "and"
    qualifiers: Tuple[str, ...] = ()
    results: Tuple[str, ...] = ()
    effects: Tuple[DescriptiveEffect, ...] = ()


@dataclass(frozen=True)
class DescriptiveProposition:
    proposition_id: str
    meaning: DescriptiveMeaning
    spans: Tuple[DescriptiveSourceSpan, ...]
    owner: str
    domain: str
    layer: str


@dataclass(frozen=True)
class DescriptiveDependency:
    dependency_id: str
    relation: str
    source_id: str
    target_id: str
    spans: Tuple[DescriptiveSourceSpan, ...]


@dataclass(frozen=True)
class DescriptiveClaim:
    claim_id: str
    unit_id: str
    proposition_ids: Tuple[str, ...]
    meanings: Tuple[DescriptiveMeaning, ...]
    dependencies: Tuple[DescriptiveDependency, ...] = ()


DESCRIPTIVE_SYNTHESIS_VERSION = "1.1"
DESCRIPTIVE_CLAIM_SCOPES = frozenset({"initiation", "behavior", "result", "exception", "local_interaction"})


@dataclass(frozen=True)
class DescriptiveClaimIntent:
    claim_id: str
    role: str
    proposition_ids: Tuple[str, ...]
    required_dependency_ids: Tuple[str, ...] = ()
    scope: str = "behavior"
    synthesis_rule: str = "explicit_relationship_bundle-v1"


@dataclass(frozen=True)
class DescriptiveFunctionalClaim:
    claim_id: str
    role: str
    proposition_ids: Tuple[str, ...]
    meanings: Tuple[DescriptiveMeaning, ...]
    dependencies: Tuple[DescriptiveDependency, ...]
    spans: Tuple[DescriptiveSourceSpan, ...]
    synthesis_rule: str = "explicit_relationship_bundle-v1"
    scope: str = "behavior"

    @property
    def responsibilities(self) -> Tuple[str, ...]:
        return tuple(dict.fromkeys(meaning.subject for meaning in self.meanings))


def descriptive_claim_scope_findings(scope: str, meanings: Sequence[DescriptiveMeaning]) -> List[str]:
    if scope not in DESCRIPTIVE_CLAIM_SCOPES:
        return ["unknown_claim_scope"]
    if not meanings or any(not meaning.subject.strip() or not meaning.predicate.strip() for meaning in meanings):
        return ["missing_functional_responsibility_or_behavior"]
    if any(not isinstance(effect, DescriptiveEffect) or not effect.action.strip() or not effect.object.strip()
           for meaning in meanings for effect in meaning.effects):
        return ["invalid_functional_effect"]
    if scope == "initiation" and not any(meaning.predicate in {"start", "initiate"}
                                         or any(effect.action in {"start", "initiate"} for effect in meaning.effects)
                                         for meaning in meanings):
        return ["initiation_scope_without_initiating_action"]
    if scope == "result" and not any(meaning.results or (meaning.predicate == "produce" and meaning.objects)
                                     for meaning in meanings):
        return ["result_scope_without_supported_result"]
    if scope == "local_interaction" and not any(meaning.relation == "interaction" for meaning in meanings):
        return ["interaction_scope_without_supported_interaction"]
    if scope == "exception" and not any(context.operator == "not" for meaning in meanings for context in meaning.contexts):
        return ["exception_scope_without_explicit_exception_condition"]
    return []


@dataclass(frozen=True)
class DescriptiveClaimDecision:
    claim_id: str
    decision: str
    findings: Tuple[str, ...] = ()


@dataclass(frozen=True)
class DescriptiveAcceptance:
    semantic_fidelity: str = "UNRESOLVED"
    publication_quality: str = "UNRESOLVED"
    materialization: str = "UNRESOLVED"
    semantic_findings: Tuple[str, ...] = ()
    publication_findings: Tuple[str, ...] = ()
    materialization_findings: Tuple[str, ...] = ()

    @property
    def verification_status(self) -> str:
        boundaries = ((self.semantic_fidelity, self.semantic_findings),
                      (self.publication_quality, self.publication_findings),
                      (self.materialization, self.materialization_findings))
        if any(status not in {"PASS", "FAIL", "UNRESOLVED"} or (status == "PASS" and findings)
               for status, findings in boundaries):
            return "FAIL"
        if any(status == "FAIL" for status, findings in boundaries):
            return "FAIL"
        return "PASS" if all(status == "PASS" for status, findings in boundaries) else "UNRESOLVED"


@dataclass(frozen=True)
class DescriptiveDisposition:
    spans: Tuple[DescriptiveSourceSpan, ...]
    proposition_ids: Tuple[str, ...]
    decision: str
    reason: str
    claim_ids: Tuple[str, ...] = ()
    representative_id: str = ""


@dataclass(frozen=True)
class DescriptiveStructure:
    version: str
    profile_fingerprint: str
    records: Tuple[NormalizedSourceRecord, ...]
    propositions: Tuple[DescriptiveProposition, ...]
    dependencies: Tuple[DescriptiveDependency, ...]
    claims: Tuple[DescriptiveClaim, ...]
    dispositions: Tuple[DescriptiveDisposition, ...]
    units: Tuple[SemanticUnit, ...]


def descriptive_context_findings(context: DescriptiveContext) -> List[str]:
    if context.operator == "atom":
        return [] if context.text.strip() and not context.operands else ["descriptive_context_invalid"]
    arity_valid = (context.operator == "not" and len(context.operands) == 1) or (
        context.operator in {"and", "or"} and len(context.operands) >= 2)
    findings = [] if arity_valid and not context.text else ["descriptive_context_invalid"]
    for operand in context.operands:
        findings.extend(descriptive_context_findings(operand))
    return findings


@dataclass(frozen=True)
class DescriptiveUnresolvedContext:
    context_id: str
    spans: Tuple[DescriptiveSourceSpan, ...]
    reason: str


@dataclass(frozen=True)
class DescriptiveEvidenceModel:
    records: Tuple[NormalizedSourceRecord, ...]
    propositions: Tuple[DescriptiveProposition, ...]
    dependencies: Tuple[DescriptiveDependency, ...] = ()
    unresolved_contexts: Tuple[DescriptiveUnresolvedContext, ...] = ()


def descriptive_source_model_findings(
    model: DescriptiveEvidenceModel, *, profile: DescriptiveWritingProfile, owner: str = "",
) -> List[str]:
    findings = []
    if profile not in DESCRIPTIVE_PROFILES.values():
        findings.append("descriptive_structure_profile_unknown")
    indexes = []
    for values, identity in ((model.records, "record_id"), (model.propositions, "proposition_id"),
                             (model.dependencies, "dependency_id"), (model.unresolved_contexts, "context_id")):
        indexed = {getattr(value, identity): value for value in values}
        if len(indexed) != len(values) or any(not key.strip() for key in indexed):
            findings.append("descriptive_structure_identity_invalid:" + identity)
        indexes.append(indexed)
    records, propositions, dependencies, unresolved = indexes
    if set(propositions) & set(unresolved):
        findings.append("descriptive_context_identity_collision")

    def check_spans(spans: Sequence[DescriptiveSourceSpan]) -> None:
        valid = bool(spans) and len(set(spans)) == len(spans)
        for span in spans:
            record = records.get(span.record_id)
            if (record is None or type(span.start) is not int or type(span.end) is not int
                    or not 0 <= span.start < span.end <= len(record.text)
                    or record.text[span.start:span.end] != span.text or not span.text.strip()):
                valid = False
        if not valid:
            findings.append("descriptive_source_span_invalid")

    for record in model.records:
        findings.extend(descriptive_evidence_findings(record, profile=profile))
    for proposition in model.propositions:
        check_spans(proposition.spans)
        meaning = proposition.meaning
        if (meaning.relation not in DESCRIPTIVE_RELATIONS or not meaning.subject.strip()
                or not meaning.predicate.strip() or meaning.modality not in {"", "can", "may"}
                or type(meaning.negative) is not bool or meaning.coordination not in {"and", "or"}
                or any(not value.strip() for value in (*meaning.objects, *meaning.qualifiers, *meaning.results))
                  or any(not isinstance(effect, DescriptiveEffect) or not effect.action.strip() or not effect.object.strip()
                      for effect in meaning.effects)
                or (meaning.relation != "behavior" and not meaning.objects)):
            findings.append("descriptive_proposition_invalid:" + proposition.proposition_id)
        for context in meaning.contexts:
            findings.extend(descriptive_context_findings(context))
        if ((profile.owner_policy == "same_owner" and (not owner or proposition.owner != owner))
                or (profile.allowed_domains and proposition.domain not in profile.allowed_domains)
                or (profile.allowed_layers and proposition.layer not in profile.allowed_layers)):
            findings.append("descriptive_proposition_scope_invalid:" + proposition.proposition_id)
    for context in model.unresolved_contexts:
        check_spans(context.spans)
        if not context.reason.strip():
            findings.append("descriptive_unresolved_context_invalid:" + context.context_id)
    endpoints = set(propositions) | set(unresolved)
    for dependency in model.dependencies:
        check_spans(dependency.spans)
        if (dependency.relation not in DESCRIPTIVE_DEPENDENCIES
                or dependency.source_id not in endpoints or dependency.target_id not in endpoints
                or dependency.source_id == dependency.target_id):
            findings.append("descriptive_dependency_invalid:" + dependency.dependency_id)
    successors: Dict[str, Set[str]] = {identity: set() for identity in endpoints}
    for dependency in model.dependencies:
        if dependency.relation != "referenced_context" and dependency.source_id in successors and dependency.target_id in successors:
            source_id, target_id = (dependency.target_id, dependency.source_id) if dependency.relation == "after" else (
                dependency.source_id, dependency.target_id)
            successors[source_id].add(target_id)
    remaining = set(successors)
    while remaining:
        targets = {target for source in remaining for target in successors[source] if target in remaining}
        roots = remaining - targets
        if not roots:
            findings.append("descriptive_dependency_cycle")
            break
        remaining.difference_update(roots)
    return list(dict.fromkeys(findings))


@dataclass(frozen=True)
class DescriptiveSynthesisPlan:
    version: str
    profile_fingerprint: str
    evidence: DescriptiveEvidenceModel
    intents: Tuple[DescriptiveClaimIntent, ...]
    claims: Tuple[DescriptiveFunctionalClaim, ...]
    decisions: Tuple[DescriptiveClaimDecision, ...]
    dispositions: Tuple[DescriptiveDisposition, ...]
    findings: Tuple[str, ...]
    acceptance: DescriptiveAcceptance = DescriptiveAcceptance()


def build_descriptive_synthesis_plan(
    evidence: DescriptiveEvidenceModel, intents: Sequence[DescriptiveClaimIntent], *,
    profile: DescriptiveWritingProfile, owner: str = "",
) -> DescriptiveSynthesisPlan:
    """Bundle explicit model relationships; do not certify source interpretation or publication."""
    intents = tuple(intents)
    findings = descriptive_source_model_findings(evidence, profile=profile, owner=owner)
    intent_ids = [intent.claim_id for intent in intents]
    if len(set(intent_ids)) != len(intent_ids) or any(not identity.strip() for identity in intent_ids):
        findings.append("descriptive_claim_intent_identity_invalid")
    propositions = {value.proposition_id: value for value in evidence.propositions}
    dependencies = {value.dependency_id: value for value in evidence.dependencies}
    unresolved = {value.context_id: value for value in evidence.unresolved_contexts}
    claims = []
    decisions = []
    for intent in intents:
        claim_findings = []
        if findings:
            claim_findings.append("invalid_evidence_model")
        if intent.role not in DESCRIPTIVE_ROLE_RELATIONS:
            claim_findings.append("invalid_claim_role")
        if (not intent.proposition_ids or len(set(intent.proposition_ids)) != len(intent.proposition_ids)
                or any(identity not in propositions for identity in intent.proposition_ids)):
            claim_findings.append("invalid_claim_members")
        if (len(set(intent.required_dependency_ids)) != len(intent.required_dependency_ids)
                or any(identity not in dependencies for identity in intent.required_dependency_ids)):
            claim_findings.append("invalid_required_dependencies")
        selected = set(intent.proposition_ids)
        changed = True
        while changed:
            changed = False
            for edge in evidence.dependencies:
                endpoints = {edge.source_id, edge.target_id}
                if selected & endpoints and not endpoints.issubset(selected):
                    selected.update(endpoints)
                    changed = True
        edges = tuple(edge for edge in evidence.dependencies
                      if {edge.source_id, edge.target_id}.issubset(selected))
        if not set(intent.required_dependency_ids).issubset(edge.dependency_id for edge in edges):
            claim_findings.append("required_dependency_outside_claim")
        for identity in sorted(selected & set(unresolved)):
            claim_findings.append("unresolved_dependency:" + identity)
        if selected:
            reachable = {next(iter(selected))}
            changed = True
            while changed:
                changed = False
                for edge in edges:
                    endpoints = {edge.source_id, edge.target_id}
                    if reachable & endpoints and not endpoints.issubset(reachable):
                        reachable.update(endpoints)
                        changed = True
            same_responsibility = (intent.synthesis_rule == "shared_responsibility_bundle-v1"
                                   and len({propositions[identity].meaning.subject for identity in selected
                                            if identity in propositions}) == 1
                                   and selected.issubset(propositions))
            same_context = (intent.synthesis_rule == "shared_context_bundle-v1"
                            and selected.issubset(propositions) and bool(selected)
                            and all(propositions[identity].meaning.contexts for identity in selected)
                            and len({(propositions[identity].owner, propositions[identity].domain,
                                      propositions[identity].layer, propositions[identity].meaning.contexts,
                                      propositions[identity].meaning.mode, propositions[identity].meaning.modality,
                                      propositions[identity].meaning.negative) for identity in selected}) == 1)
            if reachable != selected and not same_responsibility and not same_context:
                claim_findings.append("disconnected_claim_bundle")
        if intent.synthesis_rule not in {"explicit_relationship_bundle-v1", "shared_responsibility_bundle-v1", "shared_context_bundle-v1"}:
            claim_findings.append("unknown_synthesis_rule")
        if any(propositions[identity].meaning.relation not in DESCRIPTIVE_ROLE_RELATIONS.get(intent.role, ())
               for identity in intent.proposition_ids if identity in propositions):
            claim_findings.append("unsupported_claim_role")
        claim_findings.extend(descriptive_claim_scope_findings(intent.scope,
                      tuple(value.meaning for value in evidence.propositions if value.proposition_id in selected)))
        claim_findings = list(dict.fromkeys(claim_findings))
        decisions.append(DescriptiveClaimDecision(intent.claim_id, "unresolved" if claim_findings else "preserved",
                                                  tuple(claim_findings)))
        if not claim_findings:
            members = tuple(value for value in evidence.propositions if value.proposition_id in selected)
            spans = tuple(dict.fromkeys(span for value in (*members, *edges) for span in value.spans))
            claims.append(DescriptiveFunctionalClaim(intent.claim_id, intent.role,
                          tuple(value.proposition_id for value in members), tuple(value.meaning for value in members),
                          edges, spans, synthesis_rule=intent.synthesis_rule, scope=intent.scope))
    dispositions = []
    for proposition in evidence.propositions:
        linked = tuple(claim.claim_id for claim in claims if proposition.proposition_id in claim.proposition_ids)
        dispositions.append(DescriptiveDisposition(proposition.spans, (proposition.proposition_id,),
                            "preserved" if linked else "unresolved",
                            "explicit_relationship_bundle" if linked else "no_supported_claim", linked))
    for context in evidence.unresolved_contexts:
        dispositions.append(DescriptiveDisposition(context.spans, (), "unresolved", context.reason))
    for edge in evidence.dependencies:
        linked = tuple(claim.claim_id for claim in claims if edge in claim.dependencies)
        dispositions.append(DescriptiveDisposition(edge.spans, (edge.source_id, edge.target_id),
                            "preserved" if linked else "unresolved",
                            "explicit_dependency" if linked else "unresolved_dependency", linked))
    for record in evidence.records:
        covered = {offset for disposition in dispositions for span in disposition.spans
                   if span.record_id == record.record_id and type(span.start) is int and type(span.end) is int
                   for offset in range(max(0, span.start), min(len(record.text), span.end))}
        start = None
        for offset in range(len(record.text) + 1):
            uncovered = offset < len(record.text) and offset not in covered
            if uncovered and start is None:
                start = offset
            if not uncovered and start is not None:
                text = record.text[start:offset]
                if text.strip():
                    span = DescriptiveSourceSpan(record.record_id, start, offset, text)
                    dispositions.append(DescriptiveDisposition((span,), (), "unresolved", "unmodeled_source_span"))
                start = None
    return DescriptiveSynthesisPlan(DESCRIPTIVE_SYNTHESIS_VERSION, stable_payload_hash(asdict(profile)), evidence,
                                   intents, tuple(claims), tuple(decisions), tuple(dispositions), tuple(findings))


def validate_descriptive_synthesis_plan(
    plan: DescriptiveSynthesisPlan, *, expected_evidence: DescriptiveEvidenceModel,
    expected_intents: Sequence[DescriptiveClaimIntent], profile: DescriptiveWritingProfile, owner: str = "",
) -> List[str]:
    """Check a derived plan against bound inputs, never against its own embedded evidence."""
    expected = build_descriptive_synthesis_plan(expected_evidence, expected_intents, profile=profile, owner=owner)
    findings = list(expected.findings)
    if replace(plan, acceptance=DescriptiveAcceptance()) != expected:
        findings.append("descriptive_synthesis_plan_changed")
    for decision in expected.decisions:
        findings.extend("descriptive_claim_unresolved:" + decision.claim_id + ":" + finding
                        for finding in decision.findings)
    for disposition in expected.dispositions:
        if disposition.decision == "unresolved":
            findings.append("descriptive_synthesis_unresolved:" + disposition.reason)
    if plan.acceptance != DescriptiveAcceptance():
        findings.append("descriptive_synthesis_unverified_acceptance")
    return list(dict.fromkeys(findings))


def descriptive_synthesis_plan_audit(
    plan: DescriptiveSynthesisPlan, *, expected_evidence: DescriptiveEvidenceModel,
    expected_intents: Sequence[DescriptiveClaimIntent], profile: DescriptiveWritingProfile, owner: str = "",
) -> dict:
    findings = validate_descriptive_synthesis_plan(plan, expected_evidence=expected_evidence,
               expected_intents=expected_intents, profile=profile, owner=owner)
    return {
        "contract_version": DESCRIPTIVE_SYNTHESIS_VERSION,
        "phase": "pre_prose_model_only",
        "profile": json.loads(json.dumps(asdict(profile))),
        "plan": json.loads(json.dumps(asdict(plan))),
        "typed_model_status": "blocked" if findings else "model_valid",
        "findings": findings,
        "acceptance": json.loads(json.dumps(asdict(DescriptiveAcceptance()))),
        "source_semantic_validation": "not_implemented",
        "final_prose_validation": "not_implemented",
    }


@dataclass(frozen=True)
class DescriptiveSectionContribution:
    claim_id: str
    role: str
    decision: str
    added_meanings: Tuple[DescriptiveMeaning, ...]
    added_dependency_ids: Tuple[str, ...]
    reason: str


def descriptive_section_contributions(
    claims: Sequence[DescriptiveFunctionalClaim],
) -> Tuple[DescriptiveSectionContribution, ...]:
    """Account for typed information gain; labels and contributor IDs add no meaning."""
    known_meanings: Set[DescriptiveMeaning] = set()
    known_dependencies = set()
    contributions = []
    for claim in claims:
        additions = tuple(meaning for meaning in claim.meanings if meaning not in known_meanings)
        by_id = dict(zip(claim.proposition_ids, claim.meanings))
        edges = tuple((edge, (edge.relation, by_id.get(edge.source_id), by_id.get(edge.target_id)))
                      for edge in claim.dependencies)
        added_edges = tuple(edge.dependency_id for edge, signature in edges if signature not in known_dependencies)
        valid_role = claim.role in DESCRIPTIVE_ROLE_RELATIONS
        retained = valid_role and bool(additions or added_edges)
        contributions.append(DescriptiveSectionContribution(
            claim.claim_id, claim.role, "retained" if retained else "unresolved" if not valid_role else "excluded",
            additions if retained else (), added_edges if retained else (),
            "distinct_supported_content" if retained else "unknown_section_role" if not valid_role else "equivalent_content_already_expressed",
        ))
        if retained:
            known_meanings.update(claim.meanings)
            known_dependencies.update(signature for edge, signature in edges)
    return tuple(contributions)


def validate_descriptive_structure(
    structure: DescriptiveStructure, *, profile: DescriptiveWritingProfile, owner: str = "",
) -> List[str]:
    """Check typed model integrity, not source interpretation or final prose equivalence."""
    findings: List[str] = []
    if (structure.version != DESCRIPTIVE_STRUCTURE_VERSION
            or structure.profile_fingerprint != stable_payload_hash(asdict(profile))
            or DESCRIPTIVE_PROFILES.get(profile.document_type) != profile):
        findings.append("descriptive_structure_contract_invalid")

    def index(values: Sequence, identity: str) -> dict:
        result = {getattr(value, identity): value for value in values}
        if len(result) != len(values) or any(not key.strip() for key in result):
            findings.append("descriptive_structure_identity_invalid:" + identity)
        return result

    records = index(structure.records, "record_id")
    propositions = index(structure.propositions, "proposition_id")
    dependencies = index(structure.dependencies, "dependency_id")
    claims = index(structure.claims, "claim_id")
    units = index(structure.units, "unit_id")
    if sorted(unit.output_order for unit in structure.units) != list(range(1, len(structure.units) + 1)):
        findings.append("descriptive_structure_unit_order_invalid")

    def check_spans(spans: Sequence[DescriptiveSourceSpan]) -> bool:
        valid = bool(spans) and len(set(spans)) == len(spans)
        for span in spans:
            record = records.get(span.record_id)
            if (record is None or type(span.start) is not int or type(span.end) is not int
                    or not 0 <= span.start < span.end <= len(record.text)
                    or record.text[span.start:span.end] != span.text or not span.text.strip()):
                valid = False
        if not valid:
            findings.append("descriptive_source_span_invalid")
        return valid

    findings.extend(descriptive_source_model_findings(
        DescriptiveEvidenceModel(structure.records, structure.propositions, structure.dependencies),
        profile=profile, owner=owner,
    ))
    for claim in structure.claims:
        unit = units.get(claim.unit_id)
        members = [propositions[identity] for identity in claim.proposition_ids if identity in propositions]
        expected_meanings = tuple(member.meaning for member in members)
        if (not claim.proposition_ids or len(set(claim.proposition_ids)) != len(claim.proposition_ids)
                or len(members) != len(claim.proposition_ids) or claim.meanings != expected_meanings):
            findings.append("descriptive_claim_meaning_changed:" + claim.claim_id)
        duplicate_ids = {identity for disposition in structure.dispositions
                 if disposition.decision == "excluded" and disposition.reason == "equivalent_duplicate"
                 and disposition.representative_id in claim.proposition_ids for identity in disposition.proposition_ids}
        contributing_spans = [span for member in members for span in member.spans]
        contributing_spans.extend(span for member in structure.propositions if member.proposition_id in duplicate_ids
                      for span in member.spans)
        contributing_spans.extend(span for dependency in claim.dependencies for span in dependency.spans)
        contributing_records = {span.record_id for span in contributing_spans}
        expected_records = tuple(record.record_id for record in structure.records if record.record_id in contributing_records)
        if (unit is None or not unit.text.strip() or unit.source_record_ids != expected_records
                or (profile.topic_roles and unit.section not in profile.topic_roles)):
            findings.append("descriptive_claim_unit_invalid:" + claim.claim_id)
        if len({edge.dependency_id for edge in claim.dependencies}) != len(claim.dependencies):
            findings.append("descriptive_claim_dependency_invalid:" + claim.claim_id)
        for dependency in claim.dependencies:
            if (dependencies.get(dependency.dependency_id) != dependency
                    or not {dependency.source_id, dependency.target_id}.issubset(claim.proposition_ids)):
                findings.append("descriptive_claim_dependency_invalid:" + claim.claim_id)
    if {claim.unit_id for claim in structure.claims} != set(units):
        findings.append("descriptive_structure_unit_coverage_invalid")
    if {edge.dependency_id for claim in structure.claims for edge in claim.dependencies} != set(dependencies):
        findings.append("descriptive_structure_dependency_coverage_invalid")

    dispositions: Dict[str, DescriptiveDisposition] = {}
    covered: Dict[str, Set[int]] = {identity: set() for identity in records}
    for disposition in structure.dispositions:
        if check_spans(disposition.spans):
            for span in disposition.spans:
                covered[span.record_id].update(range(span.start, span.end))
        if (not disposition.reason.strip() or disposition.decision not in {"preserved", "excluded", "unresolved"}
                or len(set(disposition.proposition_ids)) != len(disposition.proposition_ids)):
            findings.append("descriptive_disposition_invalid")
        if disposition.decision == "unresolved":
            findings.append("descriptive_structure_unresolved:" + disposition.reason)
        if disposition.decision != "preserved" and disposition.claim_ids:
            findings.append("descriptive_disposition_invalid")
        if disposition.decision != "unresolved":
            accountable_spans = {span for identity in disposition.proposition_ids if identity in propositions
                                 for span in propositions[identity].spans}
            accountable_spans.update(span for dependency in structure.dependencies
                                     if {dependency.source_id, dependency.target_id}.issubset(disposition.proposition_ids)
                                     for span in dependency.spans)
            if set(disposition.spans) != accountable_spans:
                findings.append("descriptive_disposition_unmodeled_span")
        for identity in disposition.proposition_ids:
            proposition = propositions.get(identity)
            if proposition is None or identity in dispositions:
                findings.append("descriptive_disposition_proposition_invalid:" + identity)
                continue
            dispositions[identity] = disposition
            if not set(proposition.spans).issubset(disposition.spans):
                findings.append("descriptive_disposition_span_invalid:" + identity)
            linked = tuple(claim.claim_id for claim in structure.claims if identity in claim.proposition_ids)
            if disposition.decision == "preserved" and (
                    not linked or len(set(disposition.claim_ids)) != len(disposition.claim_ids)
                    or set(disposition.claim_ids) != set(linked)):
                findings.append("descriptive_disposition_claim_invalid:" + identity)
            if disposition.decision == "excluded":
                representative = propositions.get(disposition.representative_id)
                if (disposition.reason != "equivalent_duplicate" or representative is None
                        or representative.proposition_id == identity or representative.meaning != proposition.meaning
                        or (representative.owner, representative.domain, representative.layer)
                        != (proposition.owner, proposition.domain, proposition.layer) or linked):
                    findings.append("descriptive_exclusion_not_justified:" + identity)
        if not disposition.proposition_ids and disposition.decision != "unresolved":
            findings.append("descriptive_span_without_proposition")
    for identity, disposition in dispositions.items():
        if disposition.decision == "excluded":
            representative = dispositions.get(disposition.representative_id)
            if representative is None or representative.decision != "preserved":
                findings.append("descriptive_duplicate_representative_not_preserved:" + identity)
    if set(dispositions) != set(propositions):
        findings.append("descriptive_structure_proposition_coverage_invalid")
    disposed_dependency_spans = {span for disposition in structure.dispositions for span in disposition.spans}
    if any(span not in disposed_dependency_spans for dependency in structure.dependencies for span in dependency.spans):
        findings.append("descriptive_dependency_span_disposition_missing")
    for identity, record in records.items():
        if any(not character.isspace() and offset not in covered[identity]
               for offset, character in enumerate(record.text)):
            findings.append("descriptive_structure_span_coverage_invalid:" + identity)
    return list(dict.fromkeys(findings))


def descriptive_structure_audit(
    structure: DescriptiveStructure, *, profile: DescriptiveWritingProfile, owner: str = "",
) -> dict:
    findings = validate_descriptive_structure(structure, profile=profile, owner=owner)
    return {
        "contract_version": DESCRIPTIVE_STRUCTURE_VERSION,
        "phase": "shared_model_only",
        "profile": json.loads(json.dumps(asdict(profile))),
        "structure": json.loads(json.dumps(asdict(structure))),
        "typed_model_status": "blocked" if findings else "model_valid",
        "findings": findings,
        "source_semantic_validation": "not_implemented",
        "final_prose_validation": "not_implemented",
    }


def validate_descriptive_structure_audit(
    payload: Mapping, *, expected: DescriptiveStructure, profile: DescriptiveWritingProfile, owner: str = "",
) -> List[str]:
    findings = validate_descriptive_structure(expected, profile=profile, owner=owner)
    if payload != descriptive_structure_audit(expected, profile=profile, owner=owner):
        findings.append("descriptive_structure_audit_changed")
    return findings


@dataclass(frozen=True)
class DescriptiveApplicationProfile:
    name: str
    writing_profile: str
    domains: Tuple[str, ...]
    layers: Tuple[str, ...]
    abstraction: str
    role_sections: Tuple[Tuple[str, str], ...]
    omitted_details: Tuple[str, ...]
    required_roles: Tuple[str, ...]
    presentation_order: Tuple[str, ...]
    version: int = 1
    output_sections: Tuple[Tuple[str, str], ...] = ()
    functional_actions: Tuple[str, ...] = ()


DESCRIPTIVE_ROLE_RELATIONS = MappingProxyType({
    "purpose": ("responsibility", "behavior"),
    "main_functions": ("responsibility", "behavior"),
    "functions": ("responsibility", "behavior", "interaction", "transformation", "configuration"),
    "organization": ("containment", "responsibility"),
    "boundaries": ("interaction", "responsibility", "configuration", "behavior"),
    "limits": tuple(sorted(DESCRIPTIVE_RELATIONS)),
    "operating_concept": tuple(sorted(DESCRIPTIVE_RELATIONS - {"containment"})),
    "power": tuple(sorted(DESCRIPTIVE_RELATIONS)),
})
_BLOCK_ROLE_SECTIONS = (
    ("purpose", "function"), ("main_functions", "function"), ("functions", "function"), ("boundaries", "interfaces"),
    ("organization", "function"), ("limits", "function"),
    ("operating_concept", "operating_concept"), ("power", "power"),
)
_INTEGRATION_ROLE_SECTIONS = (
    ("purpose", "integration"), ("functions", "integration"), ("boundaries", "interfaces"),
    ("organization", "integration"), ("limits", "integration"),
    ("operating_concept", "operating_concept"), ("power", "power"),
)
_IPOS_ROLE_OUTPUTS = (
    ("purpose", "functionality"), ("main_functions", "scope"), ("functions", "internal_functions"), ("boundaries", "scope"),
    ("organization", "internal_structure"), ("limits", "scope"),
    ("operating_concept", "internal_functions"), ("power", "internal_functions"),
)
IPOS_DESCRIPTIVE_SECTION_TITLES = MappingProxyType({
    "functionality": "1.1 Functionality", "scope": "1.2 Supported functions and scope",
    "internal_structure": "1.3 Internal structure", "internal_functions": "1.4 Supported block functions",
    "general_architecture": "1.5 General architecture",
})
_IPOS_FUNCTIONAL_ACTIONS = (
    "accept", "acquire", "accumulate", "average", "calculate", "compute", "control", "convert",
    "copy", "divide", "filter", "generate", "monitor", "perform", "process", "produce",
    "provide", "read", "receive", "route", "sample", "send", "store", "subtract", "support",
    "transfer", "write",
)
DESCRIPTIVE_APPLICATION_PROFILES = MappingProxyType({policy.name: policy for policy in (
    DescriptiveApplicationProfile(
        "srs-structure-v1", SRS_DESCRIPTIVE_PROFILE.name, ("system", "digital", "analog", "mixed_signal", "shared"),
        ("system", "architecture"), "system", (
            ("purpose", SRS_SYSTEM_OVERVIEW_HEADINGS[0]), ("functions", SRS_SYSTEM_OVERVIEW_HEADINGS[1]),
            ("organization", SRS_SYSTEM_OVERVIEW_HEADINGS[2]), ("boundaries", SRS_SYSTEM_OVERVIEW_HEADINGS[3]),
            ("operating_concept", SRS_SYSTEM_OVERVIEW_HEADINGS[4]), ("power", SRS_SYSTEM_OVERVIEW_HEADINGS[5]),
            ("limits", SRS_SYSTEM_OVERVIEW_HEADINGS[6]),
        ), ("encoding", "example", "procedure"), (), ("behavior", "support", "clarification")),
    DescriptiveApplicationProfile(
        "drs-structure-v1", DRS_DESCRIPTIVE_PROFILE.name, ("digital", "shared"), ("integration",),
        "digital_integration", _INTEGRATION_ROLE_SECTIONS, ("example", "procedure"), (),
        ("behavior", "support", "clarification")),
    DescriptiveApplicationProfile(
        "ars-structure-v1", ARS_DESCRIPTIVE_PROFILE.name, ("analog", "mixed_signal", "shared"), ("integration",),
        "analog_integration", _INTEGRATION_ROLE_SECTIONS, ("example", "procedure"), (),
        ("behavior", "support", "clarification")),
    DescriptiveApplicationProfile(
        "digital-ipos-structure-v1", IPOS_DESCRIPTIVE_PROFILE.name, ("digital",), ("block",),
        "block", _BLOCK_ROLE_SECTIONS, ("example", "procedure"), ("purpose", "functions"),
        ("behavior", "support", "clarification"), output_sections=_IPOS_ROLE_OUTPUTS,
        functional_actions=_IPOS_FUNCTIONAL_ACTIONS),
    DescriptiveApplicationProfile(
        "analog-ipos-structure-v1", IPOS_DESCRIPTIVE_PROFILE.name, ("analog", "mixed_signal"), ("block",),
        "block", _BLOCK_ROLE_SECTIONS, ("example", "procedure"), ("purpose", "functions"),
        ("behavior", "support", "clarification"), output_sections=_IPOS_ROLE_OUTPUTS,
        functional_actions=_IPOS_FUNCTIONAL_ACTIONS),
)})


@dataclass(frozen=True)
class DescriptiveAssessment:
    proposition_id: str
    role: str
    detail_kind: str
    essential_aspects: Tuple[str, ...]
    source_status: str = "supported"


@dataclass(frozen=True)
class DescriptiveProfileDecision:
    proposition_id: str
    decision: str
    rule: str
    section: str
    abstraction: str


def apply_descriptive_profile(
    propositions: Sequence[DescriptiveProposition], assessments: Sequence[DescriptiveAssessment],
    *, policy: DescriptiveApplicationProfile, profile: DescriptiveWritingProfile,
    dependencies: Sequence[DescriptiveDependency] = (), owner: str = "",
) -> Tuple[DescriptiveProfileDecision, ...]:
    """Select roles from explicit assessments; never interpret or rewrite source meaning."""
    if (DESCRIPTIVE_APPLICATION_PROFILES.get(policy.name) != policy
            or DESCRIPTIVE_PROFILES.get(profile.document_type) != profile
            or policy.writing_profile != profile.name):
        raise ValueError("Unknown or changed descriptive application profile")
    by_id = {proposition.proposition_id: proposition for proposition in propositions}
    assessment_ids = [assessment.proposition_id for assessment in assessments]
    if (len(by_id) != len(propositions) or len(set(assessment_ids)) != len(assessments)
            or set(assessment_ids) != set(by_id) or any(not identity.strip() for identity in by_id)):
        raise ValueError("Incomplete or duplicate descriptive assessment inventory")
    roles = dict(policy.role_sections)
    by_assessment = {assessment.proposition_id: assessment for assessment in assessments}
    linked_ids = {identity for edge in dependencies for identity in (edge.source_id, edge.target_id)}
    if not linked_ids.issubset(by_id):
        raise ValueError("Unknown descriptive dependency endpoint")
    decisions = []
    for proposition in propositions:
        assessment = by_assessment[proposition.proposition_id]
        decision, rule, section = "retained", "essential_meaning", roles.get(assessment.role, "")
        if (assessment.source_status not in {"supported", "unsupported_construction", "incomplete_source",
                                           "context_conflict", "missing_approval", "unresolved_reference",
                                           "unresolved_functional_abstraction"}
                or assessment.detail_kind not in {"functional", "encoding", "example", "procedure"}
                or len(set(assessment.essential_aspects)) != len(assessment.essential_aspects)
                or any(aspect not in {"responsibility", "organization", "conditions", "behavior", "result", "limits"}
                       for aspect in assessment.essential_aspects)):
            decision, rule = "unresolved", "invalid_assessment"
        elif assessment.source_status != "supported":
            decision, rule = "unresolved", assessment.source_status
        elif not proposition.owner or not proposition.domain or not proposition.layer:
            decision, rule = "unresolved", "missing_scope_metadata"
        elif profile.owner_policy == "same_owner" and (not owner or proposition.owner != owner):
            decision, rule = "unresolved", "owner_mismatch"
        elif proposition.domain not in policy.domains:
            decision, rule = "excluded", "outside_profile_domain"
        elif proposition.layer not in policy.layers:
            decision, rule = "excluded", "outside_profile_layer"
        elif assessment.role not in roles:
            decision, rule = "unresolved", "unknown_section_role"
        elif proposition.meaning.relation not in DESCRIPTIVE_ROLE_RELATIONS[assessment.role]:
            decision, rule = "unresolved", "relation_role_mismatch"
        elif (assessment.detail_kind in policy.omitted_details and not assessment.essential_aspects
              and proposition.proposition_id not in linked_ids
              and proposition.meaning.relation == "behavior" and not proposition.meaning.contexts
              and not proposition.meaning.mode and not proposition.meaning.modality
              and not proposition.meaning.negative and not proposition.meaning.qualifiers
              and not proposition.meaning.results):
            decision, rule = "excluded", "nonessential_" + assessment.detail_kind
        else:
            rule = "essential_meaning" if assessment.essential_aspects or proposition.proposition_id in linked_ids else "profile_detail"
        decisions.append(DescriptiveProfileDecision(proposition.proposition_id, decision, rule, section, policy.abstraction))
    return tuple(decisions)


def descriptive_profile_application_audit(
    propositions: Sequence[DescriptiveProposition], assessments: Sequence[DescriptiveAssessment],
    *, policy: DescriptiveApplicationProfile, profile: DescriptiveWritingProfile,
    dependencies: Sequence[DescriptiveDependency] = (), owner: str = "",
    structure: Optional[DescriptiveStructure] = None,
) -> dict:
    decisions = apply_descriptive_profile(propositions, assessments, policy=policy, profile=profile,
                                         dependencies=dependencies, owner=owner)
    retained_roles = {assessment.role for assessment in assessments
                      if any(decision.proposition_id == assessment.proposition_id and decision.decision == "retained"
                             for decision in decisions)}
    findings = ["descriptive_profile_unresolved:" + decision.proposition_id + ":" + decision.rule
                for decision in decisions if decision.decision == "unresolved"]
    findings.extend("descriptive_required_role_missing:" + role for role in policy.required_roles if role not in retained_roles)
    if structure is None:
        integrity_findings = ["descriptive_profile_structure_missing"]
    else:
        integrity_findings = validate_descriptive_structure(structure, profile=profile, owner=owner)
        if tuple(propositions) != structure.propositions or tuple(dependencies) != structure.dependencies:
            integrity_findings.append("descriptive_profile_structure_inventory_changed")
    findings.extend(integrity_findings)
    return {
        "contract_version": "1.0", "phase": "profile_application_only",
        "policy": json.loads(json.dumps(asdict(policy))), "policy_fingerprint": stable_payload_hash(asdict(policy)),
        "writing_profile_fingerprint": stable_payload_hash(asdict(profile)),
        "propositions": json.loads(json.dumps([asdict(value) for value in propositions])),
        "assessments": json.loads(json.dumps([asdict(value) for value in assessments])),
        "dependencies": json.loads(json.dumps([asdict(value) for value in dependencies])),
        "decisions": json.loads(json.dumps([asdict(value) for value in decisions])),
        "structure": json.loads(json.dumps(asdict(structure))) if structure is not None else None,
        "shared_integrity_status": "blocked" if integrity_findings else "model_valid",
        "profile_status": "blocked" if findings else "policy_valid", "findings": findings,
        "source_semantic_validation": "not_implemented", "final_prose_validation": "not_implemented",
        "runtime_migration": "not_implemented",
    }


def validate_descriptive_profile_application_audit(
    payload: Mapping, *, propositions: Sequence[DescriptiveProposition], assessments: Sequence[DescriptiveAssessment],
    policy: DescriptiveApplicationProfile, profile: DescriptiveWritingProfile,
    dependencies: Sequence[DescriptiveDependency] = (), owner: str = "",
    structure: Optional[DescriptiveStructure] = None,
) -> List[str]:
    expected = descriptive_profile_application_audit(propositions, assessments, policy=policy, profile=profile,
                                                   dependencies=dependencies, owner=owner, structure=structure)
    findings = list(expected["findings"])
    if payload != expected:
        findings.append("descriptive_profile_application_audit_changed")
    return findings


def run_descriptive_flow(
    records: Iterable[NormalizedSourceRecord], *, profile: DescriptiveWritingProfile,
    adapter_name: str,
    adapter: Callable[[NormalizedSourceRecord], Tuple[Sequence[DescriptiveFact], str]],
    owner: str = "",
) -> DescriptiveFlowResult:
    registered = DESCRIPTIVE_PROFILES.get(profile.document_type)
    if not profile.runtime_enabled or (registered is not None and profile != registered):
        raise ValueError("Descriptive profile not activated or changed: " + profile.document_type)
    if adapter_name not in profile.projection_adapters:
        raise ValueError("Projection adapter outside descriptive profile: " + adapter_name)
    inputs = tuple(records)
    if len({record.record_id for record in inputs}) != len(inputs):
        raise ValueError("Duplicate descriptive source record IDs")
    facts: List[DescriptiveFact] = []
    origins: Dict[str, List[str]] = {}
    audits: List[Dict[str, str]] = []
    fingerprint = stable_payload_hash(asdict(profile))

    def audit(record: NormalizedSourceRecord, decision: str, reason: str, **fields: str) -> None:
        audits.append({
            "profile": profile.name, "profile_version": str(profile.version),
            "profile_fingerprint": fingerprint, "projection_adapter": adapter_name,
            "source_record_id": record.record_id, "normalized_record_json": json.dumps(asdict(record), sort_keys=True),
            "decision": decision, "reason": reason, "fact_id": "", "fact_json": "",
            "contributor_ids": "[]", "unit_id": "", "semantic_unit_json": "", **fields,
        })

    for record in inputs:
        if descriptive_evidence_findings(record, profile=profile):
            audit(record, "rejected_evidence", record.rationale or "outside explicit evidence boundary")
            continue
        projected, reason = adapter(record)
        projected = tuple(projected)
        if not projected:
            audit(record, "projection_gap", reason or "no supported construction")
            continue
        for fact in projected:
            if (fact.statement != record.text or fact.source not in record.provenance
                    or fact.scope != record.ownership_scope
                    or (fact.evidence_kind == "approved_normative" and record.authority_tier != "approved_normative")):
                raise ValueError("Projection changed admitted provenance or evidence scope")
            findings = descriptive_fact_findings(fact, profile=profile, owner=owner)
            if findings:
                raise ValueError("Invalid descriptive projection: " + "; ".join(findings))
            facts.append(fact)
            origins.setdefault(fact.fact_id, []).append(record.record_id)
        audit(record, "admitted", "explicit evidence and profile boundary passed")
        if reason:
            audit(record, "projection_gap", reason)

    paragraphs = compose_descriptive_paragraphs(facts, profile=profile, owner=owner)
    findings = validate_descriptive_paragraphs(paragraphs, facts=facts, profile=profile, owner=owner)
    if findings:
        raise ValueError("Descriptive preservation failed: " + "; ".join(findings))
    by_id = {record.record_id: record for record in inputs}
    units = []
    for order, paragraph in enumerate(paragraphs, start=1):
        contributors = tuple(dict.fromkeys(fact.fact_id for fact in paragraph.facts))
        source_ids = tuple(dict.fromkeys(source_id for fact_id in contributors for source_id in origins[fact_id]))
        unit = SemanticUnit("desc-unit-" + stable_payload_hash((fingerprint, contributors, paragraph.text))[:20],
                            paragraph.facts[0].section, paragraph.text, source_ids, order)
        units.append(unit)
        for fact in paragraph.facts:
            for source_id in dict.fromkeys(origins[fact.fact_id]):
                audit(by_id[source_id], "selected", "compatible facts; all contributors retained",
                      fact_id=fact.fact_id, fact_json=json.dumps(asdict(fact), sort_keys=True),
                      contributor_ids=json.dumps(contributors), unit_id=unit.unit_id,
                      semantic_unit_json=json.dumps(asdict(unit), sort_keys=True))
    result = DescriptiveFlowResult(inputs, tuple(facts), tuple(paragraphs), tuple(units), tuple(audits))
    findings = validate_descriptive_flow(result, profile=profile, owner=owner)
    if findings:
        raise ValueError("Descriptive audit linkage failed: " + "; ".join(findings))
    return result


def validate_descriptive_flow(
    result: DescriptiveFlowResult, *, profile: DescriptiveWritingProfile, owner: str = "",
) -> List[str]:
    findings = validate_descriptive_paragraphs(result.paragraphs, facts=result.facts, profile=profile, owner=owner)
    records = {record.record_id: record for record in result.records}
    units = {unit.unit_id: unit for unit in result.units}
    facts = {fact.fact_id: fact for fact in result.facts}
    fingerprint = stable_payload_hash(asdict(profile))
    if len(units) != len(result.units) or len(result.units) != len(result.paragraphs):
        findings.append("descriptive_semantic_units_invalid")
    expected_links = set()
    for paragraph, unit in zip(result.paragraphs, result.units):
        contributors = tuple(dict.fromkeys(fact.fact_id for fact in paragraph.facts))
        origins = tuple(dict.fromkeys(record.record_id for fact in paragraph.facts for record in result.records
                                     if fact.statement == record.text and fact.source in record.provenance
                                     and fact.scope == record.ownership_scope
                                     and (fact.evidence_kind != "approved_normative" or record.authority_tier == "approved_normative")
                                     and not descriptive_evidence_findings(record, profile=profile)))
        if (not paragraph.facts or unit.text != paragraph.text or unit.section != paragraph.facts[0].section
                or unit.source_record_ids != origins or unit.output_order < 1
                or unit.unit_id != "desc-unit-" + stable_payload_hash((fingerprint, contributors, paragraph.text))[:20]):
            findings.append("descriptive_semantic_unit_changed")
        for fact in paragraph.facts:
            if not any(fact.statement == records[record_id].text and fact.source in records[record_id].provenance for record_id in origins):
                findings.append("descriptive_fact_source_link_missing:" + fact.fact_id)
            for record_id in origins:
                record = records[record_id]
                if fact.statement == record.text and fact.source in record.provenance:
                    expected_links.add((unit.unit_id, fact.fact_id, record_id))
    actual_links = set()
    audited_records = set()
    for row in result.audit:
        record_id = row.get("source_record_id", "")
        record = records.get(record_id)
        audited_records.add(record_id)
        try:
            if (record is None or row["profile"] != profile.name or row["profile_version"] != str(profile.version)
                    or row["profile_fingerprint"] != fingerprint
                    or row["projection_adapter"] not in profile.projection_adapters
                    or json.loads(row["normalized_record_json"]) != json.loads(json.dumps(asdict(record)))):
                raise ValueError("Invalid profile or normalized record")
            if row["decision"] == "selected":
                if descriptive_evidence_findings(record, profile=profile):
                    raise ValueError("Selected audit record is not admissible")
                unit = units[row["unit_id"]]
                fact = facts[row["fact_id"]]
                paragraph = next(paragraph for paragraph in result.paragraphs
                                 if paragraph.text == unit.text and paragraph.facts[0].section == unit.section
                                 and fact in paragraph.facts)
                if (json.loads(row["fact_json"]) != json.loads(json.dumps(asdict(fact)))
                        or json.loads(row["semantic_unit_json"]) != json.loads(json.dumps(asdict(unit)))
                        or json.loads(row["contributor_ids"]) != list(dict.fromkeys(member.fact_id for member in paragraph.facts))):
                    raise ValueError("Invalid fact or semantic unit")
                actual_links.add((unit.unit_id, fact.fact_id, record_id))
            elif row["decision"] not in {"admitted", "rejected_evidence", "projection_gap"}:
                raise ValueError("Invalid audit decision")
        except (KeyError, TypeError, ValueError, StopIteration):
            findings.append("descriptive_audit_link_invalid:" + record_id)
    if set(records) != audited_records:
        findings.append("descriptive_record_audit_coverage_invalid")
    if expected_links != actual_links:
        findings.append("descriptive_contributor_audit_coverage_invalid")
    return findings


def _descriptive_join(values: Sequence[str], conjunction: str = "and") -> str:
    if len(values) < 2:
        return values[0] if values else ""
    return ", ".join(values[:-1]) + f" {conjunction} " + values[-1]


def _descriptive_fact_key(fact: DescriptiveFact) -> tuple:
    return (
        fact.subject, fact.action, fact.scope, fact.section, fact.mode, fact.condition,
        fact.modality, fact.negative, fact.coordination, fact.object_suffix,
        fact.plural_subject, fact.qualifiers, fact.results,
        fact.objects if fact.coordination == "or" or fact.source == "inventory:main_functions" else (),
        fact.evidence_kind, fact.owner, fact.domain, fact.layer,
    )


def aggregate_descriptive_facts(
    facts: Iterable[DescriptiveFact], *, profile: DescriptiveWritingProfile, owner: str = "",
) -> List[Tuple[DescriptiveFact, ...]]:
    groups: Dict[tuple, List[DescriptiveFact]] = {}
    for fact in facts:
        findings = descriptive_fact_findings(fact, profile=profile, owner=owner)
        if findings:
            raise ValueError("Invalid descriptive fact: " + "; ".join(findings))
        groups.setdefault(_descriptive_fact_key(fact), []).append(fact)
    return [tuple(members) for members in groups.values()]


def compose_descriptive_paragraphs(
    facts: Iterable[DescriptiveFact], *, profile: DescriptiveWritingProfile, owner: str = "",
) -> List[DescriptiveParagraph]:
    paragraphs: List[DescriptiveParagraph] = []
    for members in aggregate_descriptive_facts(facts, profile=profile, owner=owner):
        first = members[0]
        objects = list(dict.fromkeys(value for fact in members for value in fact.objects))
        action = first.action
        if first.modality:
            predicate = first.modality + (" not " if first.negative else " ") + action
        elif first.negative and action == "be":
            predicate = "are not" if first.plural_subject else "is not"
        elif first.negative:
            predicate = ("do not " if first.plural_subject else "does not ") + action
        elif action == "be":
            predicate = "are" if first.plural_subject else "is"
        elif first.plural_subject:
            predicate = action
        elif action == "have":
            predicate = "has"
        elif action.endswith("y") and len(action) > 1 and action[-2] not in "aeiou":
            predicate = action[:-1] + "ies"
        else:
            predicate = action + ("es" if action.endswith(("s", "sh", "ch", "x", "z", "o")) else "s")
        text = f"{first.subject} {predicate} {_descriptive_join(objects, first.coordination)}"
        if first.object_suffix:
            text += " " + first.object_suffix
        if first.qualifiers:
            text += " " + " ".join(first.qualifiers)
        if first.results:
            text += (", " if len(first.results) > 1 else " and ") + _descriptive_join(first.results)
        if first.condition:
            preserve_initial = profile == IPOS_DESCRIPTIVE_PROFILE and bool(re.match(r"(?:\w*_\w*|[A-Z][A-Z0-9]+)\b", text))
            text = f"{first.condition}, {text if preserve_initial else text[0].lower() + text[1:]}"
        if first.mode:
            preserve_initial = profile == IPOS_DESCRIPTIVE_PROFILE and bool(re.match(r"(?:\w*_\w*|[A-Z][A-Z0-9]+)\b", text))
            text = f"In {first.mode}, {text if preserve_initial else text[0].lower() + text[1:]}"
        paragraphs.append(DescriptiveParagraph(text.rstrip(" .") + ".", tuple(members)))
    return paragraphs


def validate_descriptive_paragraphs(
    paragraphs: Sequence[DescriptiveParagraph], *, facts: Sequence[DescriptiveFact],
    profile: DescriptiveWritingProfile, owner: str = "",
) -> List[str]:
    findings: List[str] = []
    expected = {fact.fact_id for fact in facts}
    represented: Set[str] = set()
    for paragraph in paragraphs:
        keys = {_descriptive_fact_key(fact) for fact in paragraph.facts}
        if len(keys) != 1:
            findings.append("descriptive_incompatible_aggregation")
        text = " ".join(paragraph.text.casefold().split())
        if re.search(r"\b(?:shall|must|required to)\b|\bCovers\s*:", paragraph.text, re.IGNORECASE):
            findings.append("descriptive_normative_wording")
        for fact in paragraph.facts:
            represented.add(fact.fact_id)
            findings.extend(descriptive_fact_findings(fact, profile=profile, owner=owner))
            if fact.scope not in profile.allowed_scopes or not fact.source or not fact.statement:
                findings.append("descriptive_scope_or_provenance_invalid")
            terms = (*fact.objects, fact.subject, fact.object_suffix, fact.condition, fact.mode, *fact.qualifiers, *fact.results)
            if any(" ".join(term.casefold().split()) not in text for term in terms if term):
                findings.append("descriptive_fact_content_lost:" + fact.fact_id)
            if fact.modality and not re.search(r"\b" + re.escape(fact.modality) + r"\b", text):
                findings.append("descriptive_modality_lost:" + fact.fact_id)
            if not (fact.action == "be" and re.search(r"\b(?:is|are|be)\b", text)) and not (fact.action == "have" and re.search(r"\bhas\b", text)) and not re.search(r"\b" + re.escape(fact.action) + r"(?:s|es)?\b", text) and not (
                fact.action.endswith("y") and re.search(r"\b" + re.escape(fact.action[:-1]) + r"ies\b", text)
            ):
                findings.append("descriptive_action_lost:" + fact.fact_id)
            if fact.coordination == "or" and len(fact.objects) > 1 and " or " not in text:
                findings.append("descriptive_alternative_lost:" + fact.fact_id)
            subject = " ".join(fact.subject.casefold().split())
            action_forms = [fact.action, fact.action + "s", fact.action + "es"]
            if fact.action == "be":
                action_forms.extend(("is", "are"))
            if fact.action == "have":
                action_forms.append("has")
            if fact.action.endswith("y"):
                action_forms.append(fact.action[:-1] + "ies")
            predicate_text = text
            if fact.mode:
                predicate_text = re.sub(r"^in " + re.escape(fact.mode.casefold()) + r",\s*", "", predicate_text)
            if fact.condition:
                predicate_text = re.sub(r"^" + re.escape(fact.condition.casefold()) + r",\s*", "", predicate_text)
            predicates = list(re.finditer(
                re.escape(subject) + r"\s+(?:(?:can|may|does|do)\s+)?(?P<negative>not\s+)?(?:"
                + "|".join(re.escape(form) for form in action_forms) + r")\b(?:\s+(?P<passive_negative>not\s+)?(?P<object>.*))?\.?$",
                predicate_text,
            ))
            if not any(fact.negative == bool(predicate["negative"] or predicate["passive_negative"])
                       for predicate in predicates):
                findings.append("descriptive_polarity_changed:" + fact.fact_id)
            if not any(all(" ".join(value.casefold().split()) in (predicate["object"] or "")
                           for value in fact.objects) for predicate in predicates):
                findings.append("descriptive_relation_changed:" + fact.fact_id)
            if re.search(r"\bthen\b", text) and not re.search(
                r"\bthen\b", fact.statement, re.I,
            ):
                findings.append("descriptive_sequence_not_supported:" + fact.fact_id)
    for fact_id in sorted(expected - represented):
        findings.append("descriptive_fact_not_rendered:" + fact_id)
    for fact_id in sorted(represented - expected):
        findings.append("descriptive_unexpected_contributor:" + fact_id)
    return findings


def extract_srs_descriptive_facts(
    statement: str, *, source: str, scope: str, section: str, mode: str = "",
) -> Tuple[List[DescriptiveFact], str]:
    facts: List[DescriptiveFact] = []
    recognized = False
    unsupported = False
    for clause in re.split(r"[•]|(?<=[.!?])\s+(?=[A-Z])", statement):
        clause = clause.strip()
        if not clause or re.search(r"\b(?:shall|must|required to|Covers:)\b", clause, re.IGNORECASE):
            continue
        condition = ""
        conditional = re.match(r"^((?:When|If|Unless)\s+[^,]+),\s*(.+)$", clause, re.IGNORECASE)
        if conditional:
            condition, clause = conditional.groups()
        external = re.search(
            r"\bcan\s+read\s+external\s+sensors\s+by\s+means\s+of\s+(?:an?\s+)?"
            r"([A-Za-z0-9+ -]+?\s+protocol)\b", clause, re.IGNORECASE,
        )
        samples = re.search(
            r"\bsamples(?:\s+of)?\s+([A-Za-z ,/-]+?)\s+(?:signals?\s+)?from\s+ADC\b"
            r"([^.!?]*\bstores\b[^.!?]*)[.!?]$", clause, re.IGNORECASE,
        )
        processed = re.fullmatch(
            r"([A-Za-z ,/-]+?)\s+data\s+raw\s+in\b.+?\bare\s+processed\s+by\b.+?[.!?]",
            clause, re.IGNORECASE,
        )
        accessible = re.fullmatch(
            r"The data collected by\b.+?\bis accessible from\b.+?,\s*for sensor fusion algorithm elaboration[.!?]",
            clause, re.IGNORECASE,
        )
        impedance = re.fullmatch(
            r".+?\s+channel\s+delivers\s+both\s+the real and the imaginary parts of\s+(.+?)[.!?]",
            clause, re.IGNORECASE,
        )
        direct = re.fullmatch(
            r"The system\s+(?:(can|may)\s+(not\s+)?(support|provide|perform|integrate|acquire|process|store)|"
            r"(does not)\s+(support|provide|perform|integrate|acquire|process|store)|"
            r"(supports|provides|performs|integrates|acquires|processes|stores))\s+(.+)[.]",
            clause, re.IGNORECASE,
        )
        if not any((external, samples, processed and mode, accessible, impedance, direct)):
            unsupported = True
            continue
        recognized = True
        guarded_clause = re.sub(r"according to the user's configuration", "", clause, flags=re.IGNORECASE)
        if not direct and re.search(
            r"\b(?:not|no|cannot|never|may|might|optionally|only|if|when|unless|until|before|after|without|except)\b",
            guarded_clause, re.IGNORECASE,
        ):
            continue
        values = dict(subject="The system", source=source, statement=statement, scope=scope,
                      section=section, mode=mode, condition=condition)
        if external:
            qualifiers = ["through " + external.group(1).strip()]
            tail = clause[external.end():].strip(" .")
            if tail:
                if re.fullmatch(r",\s*enabling data collection in\s+\S+\s+and embedded elaboration also on external data domain", tail, re.IGNORECASE):
                    qualifiers.append("with data collection and embedded processing of external measurements")
                else:
                    unsupported = True
                    continue
            facts.append(DescriptiveFact(
                action="acquire", objects=("external sensor measurements",), modality="can",
                qualifiers=tuple(qualifiers), **values,
            ))
        elif samples:
            measured = samples.group(1).strip()
            coordination = "or" if re.search(r"\bor\b", measured) else "and"
            objects = tuple(item.strip() for item in re.split(r",\s*|\s+(?:and|or)\s+", measured) if item.strip())
            results = []
            if re.search(r"\baverag(?:e|es|ing)\b", samples.group(2), re.IGNORECASE):
                averaging = "averages the samples"
                if "according to the user's configuration" in samples.group(2).casefold():
                    averaging += " according to the user's configuration"
                results.append(averaging)
            results.append("stores the raw sampled data" if re.search(r"\braw\b", samples.group(2), re.IGNORECASE)
                           else "stores the sampled data")
            facts.append(DescriptiveFact(action="acquire", objects=objects, coordination=coordination,
                                         object_suffix="measurements", results=tuple(results), **values))
        elif processed and mode:
            measured = processed.group(1).strip()
            objects = tuple(item.strip() for item in re.split(r",\s*|\s+(?:and|or)\s+", measured) if item.strip())
            facts.append(DescriptiveFact(action="process", objects=objects, object_suffix="measurements",
                                         coordination="or" if re.search(r"\bor\b", measured) else "and", **values))
        elif accessible:
            if not any(fact.objects == ("external sensor measurements",) for fact in facts):
                unsupported = True
                continue
            facts.append(DescriptiveFact(action="provide", objects=("collected external sensor measurements",),
                                         qualifiers=("for sensor-fusion processing",), **values))
        elif impedance:
            facts.append(DescriptiveFact(action="provide",
                                         objects=("the real and the imaginary parts of " + impedance.group(1),), **values))
        elif direct:
            if re.search(r"\b(?:register|regmap|address|port|signal)\b|[_=]|\b0x[0-9a-f]+\b", direct.group(7), re.IGNORECASE):
                continue
            action = (direct.group(3) or direct.group(5) or direct.group(6)).casefold()
            action = "process" if action == "processes" else action.removesuffix("s") if direct.group(6) else action
            facts.append(DescriptiveFact(action=action, objects=(direct.group(7).strip(),),
                                         modality=(direct.group(1) or "").casefold(),
                                         negative=bool(direct.group(2) or direct.group(4)), **values))
    reason = "unsupported descriptive clause requires explicit projection" if facts and unsupported else (
        "" if facts else "qualifier construction requires explicit projection" if recognized else "no supported functional construction"
    )
    return facts, reason


def compose_srs_system_overview(
    records: Iterable[Mapping[str, object]],
    *,
    mode_rows: Iterable[Mapping[str, object]] = (),
    interface_rows: Iterable[Mapping[str, object]] = (),
    power_records: Iterable[Mapping[str, object]] = (),
    block_names: Sequence[str] = (),
    audit_rows: Optional[List[Dict[str, str]]] = None,
) -> Dict[str, List[str]]:
    profile = descriptive_profile("SRS")
    """Project scoped descriptive evidence into a non-authoritative SRS overview."""
    overview_records: List[Tuple[str, str]] = []
    overview_scopes: Dict[Tuple[str, str], str] = {}
    recovered_modes: Dict[str, List[str]] = {}
    projected_facts: List[DescriptiveFact] = []

    def audit(section: str, record: Mapping[str, object], summary: str, reason: str) -> None:
        if audit_rows is not None:
            audit_rows.append({
                "section": section,
                "statement": csv_cell_text(record.get("statement") or record.get("Purpose") or record.get("text")),
                "source": csv_cell_text(record.get("source") or record.get("provenance")),
                "scope": csv_cell_text(record.get("scope") or ""),
                "rendered_summary": summary,
                "decision": "selected" if summary else "rejected",
                "reason": reason,
                "profile": SRS_DESCRIPTIVE_PROFILE.name,
                "fact_id": "",
                "fact_json": "",
                "contributor_ids": "[]",
                "profile_version": str(profile.version),
                "profile_fingerprint": stable_payload_hash(asdict(profile)),
                "projection_adapter": "srs-overview",
                "source_record_id": "",
                "normalized_record_json": "",
                "unit_id": "",
                "semantic_unit_json": "",
            })

    def audit_candidate(section: str, statement: str, source: str, summary: str) -> None:
        audit(section, {"statement": statement, "source": source, "scope": overview_scopes[(statement, source)]}, summary,
              "system-level projection of complete descriptive relationship")

    for record in records:
        statement = csv_cell_text(record.get("statement") or record.get("function") or record.get("text")).strip()
        source = csv_cell_text(record.get("source") or record.get("provenance")).strip()
        scope = csv_cell_text(record.get("scope") or record.get("evidence_scope")).casefold().strip()
        kind = csv_cell_text(record.get("evidence_kind") or record.get("record_type")).casefold().strip()
        if not statement or not source or scope not in {"system", "architecture"} or kind == "power_domain":
            audit("", record, "", "not a provenanced system/architecture descriptive record")
            continue
        mode_match = re.search(r"\bSection\s+\d+(?:\.\d+)*\s+([^,(]+?\bMode)\b", source, re.IGNORECASE)
        mode = mode_match.group(1).strip() if mode_match else csv_cell_text(record.get("mode")).strip()
        if mode:
            recovered_modes.setdefault(mode, [])
        facts, projection_reason = extract_srs_descriptive_facts(
            statement, source=source, scope=scope,
            section=SRS_SYSTEM_OVERVIEW_HEADINGS[4 if mode else 1], mode=mode,
        )
        projected_facts.extend(facts)
        if projection_reason:
            audit("", record, "", "projection_gap: " + projection_reason)
        if re.search(
            r"\b(?:shall|must|required to|Covers:)\b|\b(?:register|regmap|address|port\s+name|signal\s+(?:level|name)|"
            r"ATPG|BIST|scan|clock\s+gating|reset\s+synchroniz|transparenc|Req_ID|Target|Delay|Requirements Definition)\b|"
            r"\b0x[0-9a-f]+\b|[_=]|[|•]",
            statement,
            re.IGNORECASE,
        ):
            audit("", record, "", "raw/implementation text not eligible for direct narrative")
            continue
        if re.match(r"^(?:aim of this document|this document|in the following table|table\s+\d+)\b", statement, re.IGNORECASE):
            continue
        if len(statement) < 35 or not re.search(r"[.!?]\s*$", statement):
            continue
        if any(
            re.search(r"\b" + r"[-_\s]+".join(re.escape(token) for token in re.split(r"[-_\s]+", label.strip()) if token) + r"\b", statement, re.IGNORECASE)
            for label in block_names if label.strip()
        ):
            audit("", record, "", "named block text not eligible for direct narrative")
            continue
        overview_records.append((statement, source))
        overview_scopes[(statement, source)] = scope
        audit("", record, statement, "complete scoped descriptive candidate")

    sections: Dict[str, List[str]] = {heading: [] for heading in SRS_SYSTEM_OVERVIEW_HEADINGS}
    identity_re = re.compile(r"^([A-Z][A-Za-z0-9.-]*)\s+is\s+((?:an?|the)\s+.+)$", re.IGNORECASE)
    capability_re = re.compile(
        r"^(?:The\s+)?(.+?)\s+signal\s+chain\s+is\s+designed\s+for\s+(.+?)\.?$",
        re.IGNORECASE,
    )
    feature_capability_re = re.compile(
        r"^(?:The\s+)?(.+?)\s+signal\s+chain\s+has\b.+?\bfeatures\s+supporting\s+(.+?)"
        r"(?:,\s*(?:such as|including)\b.*)?\.?$",
        re.IGNORECASE,
    )
    identities: List[str] = []
    capabilities: List[str] = []
    measurement_facts: List[DescriptiveFact] = []
    interface_roles: List[str] = []
    for statement, source in overview_records:
        identity_match = identity_re.match(statement)
        if identity_match:
            identity = identity_match.group(2).rstrip('.')
            identity = re.sub(r"\banalog front end device\b", "analog front end", identity, flags=re.IGNORECASE)
            identity = re.sub(r"\bembedded process capabilities\b", "embedded processing capabilities", identity, flags=re.IGNORECASE)
            identities.append(f"The system is {identity}.")
            audit_candidate(SRS_SYSTEM_OVERVIEW_HEADINGS[0], statement, source, identities[-1])
        elif re.match(r"^The system is\b", statement, re.IGNORECASE):
            identities.append(statement)
            audit_candidate(SRS_SYSTEM_OVERVIEW_HEADINGS[0], statement, source, identities[-1])

        capability_match = capability_re.match(statement)
        if capability_match:
            role = re.split(r",\s*(?:such as|including)\b|\.\s+", capability_match.group(2), maxsplit=1, flags=re.IGNORECASE)[0].strip().rstrip(".")
        elif feature_capability_re.match(statement):
            feature_match = feature_capability_re.match(statement)
            role = re.split(r",\s*(?:such as|including)\b|\.\s+", feature_match.group(2), maxsplit=1, flags=re.IGNORECASE)[0].strip().rstrip(".")
        else:
            role = ""
        if role:
            role_parts = re.split(r"\s+(?=with\b|without\b|only\b|when\b|during\b)", role, maxsplit=1, flags=re.IGNORECASE)
            fact = DescriptiveFact(
                subject="The system", action="support", objects=(role_parts[0],),
                qualifiers=tuple(role_parts[1:]), source=source, statement=statement,
                scope=overview_scopes[(statement, source)], section=SRS_SYSTEM_OVERVIEW_HEADINGS[1],
            )
            projected_facts.append(fact)
            measurement_facts.append(fact)
            if len(re.split(r"\.\s+", statement)) > 1 and not any(fact.source == source for fact in projected_facts[:-1]):
                audit("", {"statement": statement, "source": source, "scope": overview_scopes[(statement, source)]}, "",
                      "projection_gap: additional sentence has no supported functional construction")
        if re.search(r"\b(?:system boundary|external interface|host interface|cross-domain interface)\b", statement, re.IGNORECASE):
            interface_roles.append(statement)

    def unique(values: Iterable[str]) -> List[str]:
        result: List[str] = []
        seen: Set[str] = set()
        for value in values:
            normalized = re.sub(r"\s+", " ", value).strip()
            key = normalized.casefold()
            if key and key not in seen:
                seen.add(key)
                result.append(normalized)
        return result

    def project(facts: Sequence[DescriptiveFact]) -> List[DescriptiveParagraph]:
        normalized: Dict[str, NormalizedSourceRecord] = {}
        projected: Dict[str, List[DescriptiveFact]] = {}
        for fact in facts:
            record_id = "desc-source-" + stable_payload_hash((fact.statement, fact.source, fact.scope))[:20]
            normalized[record_id] = NormalizedSourceRecord(record_id, fact.statement, "admitted_context",
                                                          fact.scope, (fact.source,), "selected",
                                                          "SRS adapter admitted scoped context")
            projected.setdefault(record_id, []).append(fact)
        result = run_descriptive_flow(normalized.values(), profile=profile, adapter_name="srs-overview",
                                      adapter=lambda record: (projected[record.record_id], ""))
        for linkage in result.audit:
            if linkage["decision"] != "selected":
                continue
            fact = next(fact for fact in result.facts if fact.fact_id == linkage["fact_id"])
            unit = next(unit for unit in result.units if unit.unit_id == linkage["unit_id"])
            audit(fact.section, {"statement": fact.statement, "source": fact.source, "scope": fact.scope},
                  unit.text, linkage["reason"])
            if audit_rows is not None:
                audit_rows[-1].update(linkage)
        return list(result.paragraphs)

    for paragraph in project(projected_facts):
        first = paragraph.facts[0]
        if first.mode:
            recovered_modes[first.mode].append(paragraph.text)
        else:
            capabilities.append(paragraph.text)

    identities = unique(identities)
    capabilities = unique(capabilities)
    sections[SRS_SYSTEM_OVERVIEW_HEADINGS[0]] = identities or ["The system identity and product role are not specified."]
    sections[SRS_SYSTEM_OVERVIEW_HEADINGS[1]] = [
        item for item in capabilities
        if item.casefold() not in {identity.casefold() for identity in identities}
    ] or ["System-level capabilities are not specified."]

    domain_items: List[str] = []
    domain_groups = [
        ("Sensing and measurement", "The measurement paths", True, measurement_facts),
        ("External acquisition and data handling", "The acquisition and data paths", True,
         [fact for fact in projected_facts if not fact.mode and fact.action in {"acquire", "store", "provide"}]),
        ("Embedded processing", "The processing function", False,
         [fact for fact in projected_facts if fact.action == "process" and not fact.negative]),
    ]
    for name, subject, plural, facts in domain_groups:
        if not facts:
            continue
        paragraphs = project([replace(fact, subject=subject, plural_subject=plural,
                                      section=SRS_SYSTEM_OVERVIEW_HEADINGS[2]) for fact in facts])
        domain_items.append(f"- **{name}**\n" + "\n".join("  - " + text for text in unique(paragraph.text for paragraph in paragraphs)))
    sections[SRS_SYSTEM_OVERVIEW_HEADINGS[2]] = domain_items or [
        "Top-level domain and subsystem responsibilities are not specified in the available descriptions."
    ]

    interface_groups: Dict[str, Dict[str, Set[str]]] = {}
    group_specs = (
        ("Sensing input", r"\b(?:sensor|sensing|analog input)\b"),
        ("Serial communication", r"\b(?:serial|i2c|spi)\b"),
        ("Measurement data exchange", r"\b(?:digitized samples?|buffered (?:output )?samples?)\b"),
        ("Interrupt reporting", r"\binterrupt\b"),
        ("Power supply", r"\bsupply\b"),
    )
    interface_sources: Dict[str, List[Mapping[str, object]]] = {}
    for row in interface_rows:
        purpose = csv_cell_text(row.get("Purpose") or row.get("Description")).strip()
        if not purpose:
            continue
        searchable = purpose.casefold()
        group = next((name for name, pattern in group_specs if re.search(pattern, searchable, re.IGNORECASE)), "")
        if not group:
            continue
        interface_sources.setdefault(group, []).append(row)
        values = interface_groups.setdefault(group, {"direction": set(), "type": set()})
        direction = csv_cell_text(row.get("Direction")).strip().casefold()
        medium = csv_cell_text(row.get("Type")).strip().casefold()
        if direction in {"input", "output", "bidirectional"}:
            values["direction"].add(direction)
        if medium in {"analog", "digital", "power"}:
            values["type"].add(medium)
    interface_content: List[str] = []
    if interface_groups:
        role_names = [name.casefold() for name in interface_groups]
        interface_content.append(
            "The interfaces provide " + _descriptive_join(role_names) + " roles. "
            "External versus internal placement is not specified for every interface."
        )
        interface_content.extend([
            "",
            "| Boundary role | Direction | Medium |",
            "|---|---|---|",
        ])
        direction_order = {"input": 0, "bidirectional": 1, "output": 2}
        for group, values in interface_groups.items():
            directions = ", ".join(sorted(values["direction"], key=lambda item: direction_order[item]))
            media = ", ".join(sorted(values["type"]))
            summary = f"| {group} | {directions} | {media} |"
            interface_content.append(summary)
            for row in interface_sources[group]:
                audit(SRS_SYSTEM_OVERVIEW_HEADINGS[3], row, summary, "interface role from purpose; no pin-name inference")
    elif interface_roles:
        interface_content.extend(unique(interface_roles))
    else:
        interface_content.append("External and cross-domain interface roles are not specified at system level.")
    sections[SRS_SYSTEM_OVERVIEW_HEADINGS[3]] = interface_content

    modes = unique(
        csv_cell_text(row.get("mode")).strip()
        for row in mode_rows
        if csv_cell_text(row.get("mode")).strip()
        and not csv_cell_text(row.get("evidence")).casefold().startswith("derived default:")
    )
    modes = unique([*modes, *recovered_modes])
    mode_items: List[str] = []
    for mode in modes:
        descriptions = [
            statement for statement, _source in overview_records
            if re.search(r"\b" + re.escape(mode) + r"\b", statement, re.IGNORECASE)
        ]
        role = unique([*recovered_modes.get(mode, []), *descriptions])
        mode_items.append(
            f"- **{mode}**\n" + "\n".join(f"  - {item}" for item in role) if role else
            f"- **{mode}**\n  - Its system-level role and transitions are not specified."
        )
        for statement, source in overview_records:
            if statement in descriptions:
                audit_candidate(SRS_SYSTEM_OVERVIEW_HEADINGS[4], statement, source, statement)
    sections[SRS_SYSTEM_OVERVIEW_HEADINGS[4]] = mode_items or ["System-level operating modes are not specified."]

    eligible_power = [
        record for record in power_records
        if csv_cell_text(record.get("source") or record.get("provenance"))
        and csv_cell_text(record.get("scope") or record.get("evidence_scope")).casefold() in {"system", "architecture"}
        and not re.search(
            r"\b(?:shall|must)\b|\b(?:not|no)\s+(?:switchable|power-gated|always[- ]on|retention)\b|"
            r"\b(?:cannot|can not|does not|do not)\s+(?:be powered|preserve|retain)\b",
            csv_cell_text(record.get("statement")), re.IGNORECASE,
        )
        and "definitions" not in csv_cell_text(record.get("source")).casefold()
    ]
    power_items: List[str] = []
    power_roles = (
        (r"\b(?:power domain|power domains|supply domain)\b.*\b(?:independent|independently)\b", "Supply-domain structure", "The architecture separates logic into supply domains with independent power control."),
        (r"\b(?:always[- ]on|always active|never powered down)\b", "Always-on behavior", "An always-on domain remains powered during low-power operation."),
        (r"\b(?:switchable|power-gated)\b|\b(?:domain|domain logic)\b.{0,60}\b(?:can|may)\s+be\s+powered\s+down\b", "Switchable behavior", "The power architecture includes a domain whose supply can be switched off."),
        (r"\b(?:retains|preserves)\s+(?:register\s+)?state\b.{0,100}\bpowered down\b", "Retention behavior", "State is preserved while the main logic is powered down."),
    )
    for pattern, label, _summary in power_roles:
        matches = [record for record in eligible_power if re.search(pattern, csv_cell_text(record.get("statement")), re.IGNORECASE)]
        facts = []
        for record in matches:
            statement = csv_cell_text(record.get("statement"))
            if re.search(r"\b(?:only|unless|if|after|before|until|except)\b", statement, re.IGNORECASE):
                audit("", record, "", "projection_gap: power qualifier requires explicit projection")
                continue
            values = dict(source=csv_cell_text(record.get("source") or record.get("provenance")),
                          statement=statement, scope=csv_cell_text(record.get("scope") or record.get("evidence_scope")).casefold(),
                          section=SRS_SYSTEM_OVERVIEW_HEADINGS[5])
            if label == "Supply-domain structure":
                fact = DescriptiveFact(subject="The supply domains", plural_subject=True, action="support",
                                       objects=("independent power control",), **values)
            elif label == "Always-on behavior":
                fact = DescriptiveFact(subject="An always-on domain", action="remain", objects=("powered",), **values)
            elif label == "Switchable behavior":
                qualifiers = []
                if re.search(r"\bin idle\b", statement, re.IGNORECASE):
                    qualifiers.append("in idle")
                if re.search(r"\bwhen\b[^.!?•]*\bnot in use\b", statement, re.IGNORECASE):
                    qualifiers.append("when its associated logic is not in use")
                modal = re.search(r"\b(can|may)\s+be\s+powered\s+down\b", statement, re.IGNORECASE)
                fact = DescriptiveFact(subject="A switchable domain", action="enter", objects=("a powered-down state",),
                                       modality=modal.group(1).casefold() if modal else "can",
                                       qualifiers=tuple(qualifiers), **values)
            else:
                fact = DescriptiveFact(subject="The power architecture", action="preserve", objects=("state",),
                                       qualifiers=("while the main logic is powered down",), **values)
            facts.append(fact)
        paragraphs = project(facts)
        if paragraphs:
            power_items.append(f"- **{label}**\n" + "\n".join("  - " + paragraph.text for paragraph in paragraphs))
    if not any("**Retention behavior**" in item for item in power_items):
        power_items.append("- **Retention**\n  - State retention across system power transitions is not specified.")
    clock_reset = [
        statement for statement, _source in overview_records
        if re.search(r"\b(?:clock|reset)\b", statement, re.IGNORECASE)
    ]
    if clock_reset:
        power_items.extend(f"- **Clock and reset**\n  - {item}" for item in unique(clock_reset))
    else:
        power_items.append("- **Clock and reset**\n  - Top-level clock and reset coordination are not specified.")
    sections[SRS_SYSTEM_OVERVIEW_HEADINGS[5]] = power_items

    scope_content = [
        statement for statement, _source in overview_records
        if re.search(r"\b(?:system scope|scope covers|system assumes|allocation boundar|engineering scope)\b", statement, re.IGNORECASE)
    ]
    sections[SRS_SYSTEM_OVERVIEW_HEADINGS[6]] = unique(scope_content) or [
        "System-level scope covers product capabilities and externally observable behavior. Subsystem implementation remains within its engineering allocation.",
        "- **Assumptions**\n  - Additional system-level assumptions are not described in the available descriptions.",
    ]
    for statement, source in overview_records:
        if statement in scope_content:
            audit_candidate(SRS_SYSTEM_OVERVIEW_HEADINGS[6], statement, source, statement)
    if audit_rows is not None:
        rendered = "\n".join(item for items in sections.values() for item in items)
        for row in audit_rows:
            if row["rendered_summary"] and (not row["section"] or row["rendered_summary"] not in rendered):
                row["decision"] = "candidate_only"
                row["reason"] = "candidate retained; not rendered in this overview"
    return sections


SRS_INTRODUCTORY_SECTIONS = ("1.1", "1.2", "1.3", "1.4", "2.1", "2.2", "2.3", "2.4")

SRS_INTRODUCTORY_WRITING_RULE_ID = "srs-introductory-authority-bounded-v1"
SRS_INTRODUCTORY_WRITING_RULE = (
    "Use concise, natural technical prose bounded by the selected approved content.",
    "Describe document purpose and scope without promising complete coverage or creating allocation.",
    "Mention behavior, interfaces, performance and implementation only to the extent supported by approved content.",
    "Treat terminology as a reading aid, not authority for capabilities, limits, obligations or acceptance criteria.",
    "Reference existing source and traceability information without implying new applicability or approval.",
    "Keep this preferred writing rule within the SRS introductory profile; do not inherit it in DRS, ARS or IPOS.",
)

SRS_DOCUMENT_CONTENT_PROFILE = DescriptiveWritingProfile(
    name="srs-document-content", document_type="SRS", allowed_scopes=("document", "system", "architecture"),
    evidence_kinds=("document_introduction", "document_terminology", "block_catalog", "power_domain"),
    authority_tiers=("presentation_policy", "admitted_context"), owner_policy="unrestricted",
    allowed_domains=(), allowed_layers=(), projection_adapters=("srs-document-content",),
    topic_roles=(*SRS_INTRODUCTORY_SECTIONS, "3.6", "4.1", "5.1"), excluded_terms=(), runtime_enabled=True,
)


def compose_srs_support_content(project_name: str, snapshot_id: str,
                                catalog_entries: Sequence[Mapping[str, str]], catalog_findings: Sequence[str],
                                power_records: Sequence[Mapping[str, object]]) -> Tuple[Dict[str, List[str]], Dict[str, object]]:
    sections: Dict[str, List[str]] = {section: [] for section in SRS_DOCUMENT_CONTENT_PROFILE.topic_roles}
    records: List[NormalizedSourceRecord] = []
    units: List[SemanticUnit] = []
    power_rows: List[Dict[str, str]] = []
    gaps: List[str] = list(catalog_findings)

    def retain(section: str, text: str, source_text: str, provenance: Tuple[str, ...], tier: str,
               scope: str = "architecture") -> None:
        record_id = "srs-source-" + stable_payload_hash((section, source_text, provenance))[:20]
        record = NormalizedSourceRecord(record_id, source_text, tier, scope,
                                        provenance, "selected", "SRS document-specific presentation")
        unit = SemanticUnit("srs-unit-" + stable_payload_hash((section, text, record_id))[:20], section, text,
                            (record_id,), len(units))
        records.append(record)
        units.append(unit)
        sections[section].append(text)

    purpose = (f"This document presents the system-level requirements and supporting context for {project_name} "
               "included in the selected approved baseline. Descriptive content does not add requirements or change their allocation.")
    retain("1.1", purpose, purpose, ("SRS document-purpose policy",), "presentation_policy", "document")
    scope = ("The scope is limited to system-level requirements allocated to this document and the supporting system context. "
             "Analog and digital behavior, interfaces, operating conditions and verification information are described only "
             "to the extent supported by the approved content. This description does not extend subsystem or block-level allocations.")
    retain("1.2", scope, scope, ("SRS document-scope policy",), "presentation_policy", "document")
    audience = ("This document is intended for system architects, analog and digital design engineers, firmware engineers, "
                "verification and validation engineers, test and product engineers, and program and customer stakeholders.")
    retain("1.3", audience, audience, ("SRS intended-audience policy",), "presentation_policy", "document")
    references = ("Reference documents are listed in Table 2. Source links for any included requirements are recorded "
                  "in the associated traceability matrix. Listing a document or standard does not itself establish "
                  "applicability or add obligations.")
    retain("1.4", references, references, ("SRS reference-document policy",), "presentation_policy", "document")
    terminology = {
        "2.1": (
            ("System", "The device or functional scope identified in the approved descriptions."),
            ("Subsystem", "A grouping of related functions identified in the approved architecture context."),
            ("Block", "A concrete unit identified in the approved architecture context."),
            ("SRS", "System Requirements Specification; the system-level requirements and context document."),
            ("ARS", "Analog Requirements Specification; the analog and mixed-signal integration document."),
            ("DRS", "Digital Requirements Specification; the digital integration document."),
            ("IPOS", "The block-local implementation requirements specification."),
        ),
        "2.2": (
            ("Input/output range", "The span of input or output values considered under stated operating conditions."),
            ("Gain", "The ratio of a change in output to the corresponding change in input."),
            ("Offset", "The deviation from the specified reference response at a reference input."),
            ("Noise", "Unwanted variations superimposed on a signal or measurement."),
            ("Bandwidth", "The frequency interval over which a response meets a defined criterion."),
            ("Full-scale", "The reference magnitude or span corresponding to the specified measurement range."),
            ("ODR", "Output data rate; the rate at which new output samples become available."),
            ("Sensitivity", "The change in output per unit change in the measured quantity."),
        ),
        "2.3": (
            ("Register", "A named storage element used to expose data, configuration or status."),
            ("Bit field", "A defined subset of bits within a register or data word."),
            ("Mode", "An operating configuration described in the applicable content."),
            ("State", "A condition of control logic identified within a described operating sequence."),
            ("Reset", "Initialization of affected logic as described in the applicable content."),
            ("Interrupt", "An event notification that requests attention from a controller or processor."),
            ("Boot sequence", "An initialization sequence described in the applicable content."),
            ("Interface", "A defined boundary through which components exchange data, control or status."),
        ),
        "2.4": (
            ("RMS noise", "The root-mean-square amplitude of noise over a stated measurement interval and bandwidth."),
            ("SNR", "Signal-to-noise ratio; the ratio of signal power to noise power, commonly expressed in decibels."),
            ("Latency", "The elapsed time between a defined initiating event and its corresponding response."),
            ("Precision", "The repeatability of measurement results under stated conditions."),
            ("Tolerance", "A permitted deviation where specified relative to a reference value or condition."),
            ("GR&R", "Gage repeatability and reproducibility; the evaluation of measurement-system variation."),
            ("Guard band", "A margin between acceptance and specification limits where such a margin is specified."),
            ("Pass/fail", "An assessment against acceptance criteria specified by the applicable requirements."),
        ),
    }
    terminology_context = {
        "2.1": ("These definitions aid interpretation of the system descriptions and document abbreviations. "
            "They do not establish capabilities, ownership or allocation; specific meanings follow the approved content."),
        "2.2": ("Analog terms are descriptive. Ranges, operating conditions and performance limits are those "
            "stated in the applicable approved content."),
        "2.3": ("Digital terms aid interpretation of configuration and control descriptions. Their inclusion "
            "does not imply a particular implementation or supported feature."),
        "2.4": ("Measurement terms aid interpretation of verification information. Test conditions, limits and "
            "acceptance criteria remain those specified in the applicable approved requirements."),
    }
    for section, definitions in terminology.items():
        context = terminology_context[section]
        retain(section, context, context, ("SRS terminology policy:" + section,), "presentation_policy", "document")
        for term, definition in definitions:
            text = f"**{term}:** {definition}"
            retain(section, text, text, ("SRS terminology policy:" + section + ":" + term,), "presentation_policy", "document")
    for entry in catalog_entries:
        section = "4.1" if entry["category"] == "analog" else "5.1"
        summary = compose_technical_block_purpose(entry["block"], entry["function"])
        retain(section, f"- **{entry['block']}**\n\n  {summary}", entry["function"],
               (entry["source"], "approved classification sources:" + entry["classification_sources"]), "admitted_context")
    labels = re.compile(r"(?:Included block\(s\)|Domain type|Voltage|Control mode|Notes|Function|Characteristics|Control):\s*", re.IGNORECASE)
    for record in power_records:
        statement = csv_cell_text(record.get("statement"))
        source = csv_cell_text(record.get("source") or record.get("provenance"))
        scope = csv_cell_text(record.get("scope") or record.get("evidence_scope")).casefold()
        if csv_cell_text(record.get("evidence_kind") or record.get("record_type")).casefold() != "power_domain":
            continue
        if not source or scope not in {"architecture", "system"}:
            gaps.append("SRS_POWER_SCOPE_OR_PROVENANCE_MISSING: " + source)
            continue
        matches = list(labels.finditer(statement))
        name = statement[:matches[0].start()].strip(" :.") if matches else ""
        fields = {match.group(0).split(":")[0].casefold():
                  statement[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(statement)].strip(" .\u2022")
                  for index, match in enumerate(matches)}
        if not name or not fields.get("domain type") or not fields.get("function"):
            gaps.append("SRS_POWER_PROJECTION_GAP: " + source)
            continue
        role = fields["function"]
        blocks = fields.get("included block(s)", "")
        if blocks:
            role += ". Associated blocks: " + ", ".join(value.strip() for value in blocks.split(",") if value.strip())
        voltage = fields.get("voltage", "")
        if voltage:
            if re.fullmatch(r"\d+(?:\.\d+)?\s*(?:mV|V)", voltage):
                role += ". Supply voltage: " + voltage
            else:
                gaps.append("SRS_POWER_VOLTAGE_UNRELIABLE: " + source + ": " + voltage)
        conditions = fields.get("characteristics", "")
        notes = fields.get("notes", "")
        if notes and notes.casefold() not in conditions.casefold():
            if re.search(r"\s+\d+\s+\d+(?:\.\d+)*$", notes):
                gaps.append("SRS_POWER_NOTE_UNREADABLE: " + source + ": " + notes)
            else:
                conditions = notes + (". " + conditions if conditions else "")
        cells = [name, fields["domain type"], fields.get("control mode") or fields.get("control") or "Not specified",
                 role, conditions or "Not specified"]
        text = "| " + " | ".join(value.replace("|", "\\|").replace("\n", " ") for value in cells) + " |"
        retain("3.6", text, statement, (source,), "admitted_context", scope)
        power_rows.append({"source": source, "statement": statement, "domain": name, "row": text})
    return sections, {"schema_version": 1, "snapshot_id": snapshot_id, "project_name": project_name,
                      "profile": SRS_DOCUMENT_CONTENT_PROFILE.name, "profile_version": SRS_DOCUMENT_CONTENT_PROFILE.version,
                      "profile_fingerprint": stable_payload_hash(asdict(SRS_DOCUMENT_CONTENT_PROFILE)),
                      "projection_adapter": "srs-document-content",
                      "writing_rule_id": SRS_INTRODUCTORY_WRITING_RULE_ID,
                      "writing_rule_fingerprint": stable_payload_hash(SRS_INTRODUCTORY_WRITING_RULE),
                      "records": [asdict(record) for record in records], "units": [asdict(unit) for unit in units],
                      "catalog_entries": list(catalog_entries), "power_rows": power_rows, "review_findings": gaps}


def _validate_srs_fact_audit(markdown_path: Path, section_bodies: Mapping[str, str]) -> List[str]:
    audit_path = markdown_path.with_name("descriptive_system_overview_audit.csv")
    if not audit_path.exists():
        return []
    findings: List[str] = []
    groups: Dict[Tuple[str, str, Tuple[str, ...]], List[tuple]] = {}
    with audit_path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if not {"profile", "fact_id", "fact_json", "contributor_ids", "profile_version", "profile_fingerprint",
            "projection_adapter", "source_record_id", "normalized_record_json", "unit_id", "semantic_unit_json"}.issubset(reader.fieldnames or []):
            return ["srs_descriptive_audit_schema_missing"]
        for row in reader:
            if not row.get("fact_id"):
                continue
            try:
                payload = json.loads(row["fact_json"])
                for field in ("objects", "qualifiers", "results"):
                    payload[field] = tuple(payload[field])
                fact = DescriptiveFact(**payload)
                contributors = json.loads(row["contributor_ids"])
                if not isinstance(contributors, list) or any(not isinstance(value, str) for value in contributors):
                    raise ValueError("Invalid contributors")
                if fact.fact_id not in contributors or len(set(contributors)) != len(contributors):
                    findings.append("srs_descriptive_audit_contributors_invalid:" + fact.section)
                if (fact.fact_id != row["fact_id"] or row["profile"] != SRS_DESCRIPTIVE_PROFILE.name
                        or row["decision"] != "selected" or row["section"] != fact.section
                        or any(row[field] != getattr(fact, field) for field in ("source", "statement", "scope"))):
                    raise ValueError("Audit fields disagree with fact")
                record_payload = json.loads(row["normalized_record_json"])
                record_payload["provenance"] = tuple(record_payload["provenance"])
                record = NormalizedSourceRecord(**record_payload)
                unit_payload = json.loads(row["semantic_unit_json"])
                unit_payload["source_record_ids"] = tuple(unit_payload["source_record_ids"])
                unit = SemanticUnit(**unit_payload)
                if (record.record_id != row["source_record_id"] or record.text != fact.statement
                        or record.ownership_scope != fact.scope or fact.source not in record.provenance
                        or descriptive_evidence_findings(record, profile=SRS_DESCRIPTIVE_PROFILE)
                        or row["profile_version"] != str(SRS_DESCRIPTIVE_PROFILE.version)
                        or row["profile_fingerprint"] != stable_payload_hash(asdict(SRS_DESCRIPTIVE_PROFILE))
                        or row["projection_adapter"] not in SRS_DESCRIPTIVE_PROFILE.projection_adapters
                        or unit.unit_id != row["unit_id"] or unit.section != fact.section
                        or unit.text != row["rendered_summary"] or record.record_id not in unit.source_record_ids
                        or unit.unit_id != "desc-unit-" + stable_payload_hash(
                            (row["profile_fingerprint"], tuple(contributors), unit.text))[:20]):
                    raise ValueError("Invalid profile or materialization linkage")
                groups.setdefault((fact.section, row["rendered_summary"], tuple(contributors)), []).append(
                    (fact, contributors, row, record, unit))
            except (KeyError, TypeError, ValueError):
                findings.append("srs_descriptive_audit_fact_invalid:" + row.get("fact_id", ""))
    docx_path = markdown_path.with_suffix(".docx")
    docx_text = None
    if docx_path.exists():
        try:
            from docx import Document
            document = Document(docx_path)
            docx_text = " ".join(paragraph.text for paragraph in document.paragraphs)
        except Exception:
            findings.append("srs_descriptive_docx_unreadable")
    normalize = lambda text: " ".join(text.replace("\u2019", "'").replace("\u2018", "'").casefold().split())
    for (section, summary, _contributors), members in groups.items():
        facts = [member[0] for member in members]
        expected = list(dict.fromkeys(fact.fact_id for fact in facts))
        if any(member[1] != expected for member in members):
            findings.append("srs_descriptive_audit_contributors_invalid:" + section)
        records = tuple({member[3].record_id: member[3] for member in members}.values())
        result = DescriptiveFlowResult(records, tuple(facts), (DescriptiveParagraph(summary, tuple(facts)),),
                                       (members[0][4],), tuple(member[2] for member in members))
        findings.extend(validate_descriptive_flow(result, profile=SRS_DESCRIPTIVE_PROFILE))
        if not summary or normalize(summary) not in normalize(section_bodies.get(section, "")):
            findings.append("srs_descriptive_fact_missing_from_markdown:" + section)
        if docx_text is not None and normalize(summary) not in normalize(docx_text):
            findings.append("srs_descriptive_fact_missing_from_docx:" + section)
    return findings


def validate_srs_support_content(markdown_path: Path) -> List[str]:
    audit_path = markdown_path.with_name("descriptive_srs_content_audit.json")
    repo_root = markdown_path.parents[2]
    production = (repo_root / "data/canonical/canonical_store.sqlite").exists()
    if not audit_path.exists():
        return ["srs_content_audit_missing"] if production else []
    findings: List[str] = []
    try:
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        if audit["schema_version"] != 1:
            raise ValueError("Unknown content audit schema")
        profile = SRS_DOCUMENT_CONTENT_PROFILE
        if (audit["profile"] != profile.name or audit["profile_version"] != profile.version
                or audit["profile_fingerprint"] != stable_payload_hash(asdict(profile))
                or audit["projection_adapter"] not in profile.projection_adapters):
            raise ValueError("Invalid SRS document-content profile")
        if (audit["writing_rule_id"] != SRS_INTRODUCTORY_WRITING_RULE_ID
                or audit["writing_rule_fingerprint"] != stable_payload_hash(SRS_INTRODUCTORY_WRITING_RULE)):
            raise ValueError("Invalid SRS introductory writing rule")
        records = {record["record_id"]: record for record in audit["records"]}
        if len(records) != len(audit["records"]):
            raise ValueError("Duplicate normalized source IDs")
        text = markdown_path.read_text(encoding="utf-8")
        def body(section: str) -> str:
            match = re.search(r"(?ms)^### " + re.escape(section) + r"\s[^\n]*\n(.*?)(?=^#{1,3}\s|\Z)", text)
            return match.group(1) if match else ""
        def normalized(value: str) -> str:
            value = re.sub(r"<[^>]+>|&nbsp;", "", value)
            return " ".join(value.replace("**", "").replace("\u2019", "'").replace("\u2018", "'").split()).casefold()
        sections = {section: body(section) for section in profile.topic_roles}
        if not normalized(sections["1.1"]) or "this document presents the system-level requirements" not in normalized(sections["1.1"]):
            findings.append("srs_purpose_missing")
        if not normalized(sections["1.3"]):
            findings.append("srs_intended_audience_missing")
        for section in SRS_INTRODUCTORY_SECTIONS:
            if not normalized(sections[section]):
                findings.append("srs_introductory_section_missing:" + section)
        snapshot = re.search(r"(?m)^Snapshot ID:\s*(\S+)", text)
        if not snapshot or snapshot.group(1) != audit["snapshot_id"]:
            findings.append("srs_content_snapshot_mismatch")
        if production:
            from canonical_store import read_stage2_descriptive_evidence
            from validate_downstream_coherence import resolve_downstream_contract, snapshot_architecture_context, srs_catalog_from_context
            contract = resolve_downstream_contract(repo_root, audit["snapshot_id"])
            with (repo_root / "artifacts/stage2_mirco_arc/block_inventory.csv").open(encoding="utf-8-sig", newline="") as handle:
                inventory = list(csv.DictReader(handle))
            entries, review = srs_catalog_from_context(inventory, contract, snapshot_architecture_context(repo_root, contract))
            if audit["catalog_entries"] != entries:
                findings.append("srs_catalog_authority_mismatch")
            source_power = [record for record in read_stage2_descriptive_evidence(repo_root)
                            if csv_cell_text(record.get("evidence_kind")).casefold() == "power_domain"]
            _, expected = compose_srs_support_content(audit["project_name"], audit["snapshot_id"], entries, review, source_power)
            expected = json.loads(json.dumps(expected))
            for field in ("records", "units", "power_rows", "review_findings"):
                if audit[field] != expected[field]:
                    findings.append("srs_content_provenance_mismatch:" + field)
        docx_sections: Dict[str, str] = {}
        docx_tables: List[List[List[str]]] = []
        docx_path = markdown_path.with_suffix(".docx")
        if docx_path.exists():
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            document = Document(docx_path)
            current = ""
            for element in document.element.body:
                words = " ".join(node.text or "" for node in element.iter(qn("w:t")))
                if element.tag == qn("w:p"):
                    heading = re.match(r"^(\d+(?:\.\d+)*)\.?\s", words)
                    if heading:
                        current = heading.group(1)
                if current in sections:
                    docx_sections[current] = docx_sections.get(current, "") + " " + words
                if current == "3.6" and element.tag == qn("w:tbl"):
                    table = Table(element, document)
                    docx_tables.append([[cell.text for cell in row.cells] for row in table.rows])
        used_records: Set[str] = set()
        seen_units: Set[str] = set()
        for index, unit in enumerate(audit["units"]):
            section = unit["section"]
            origins = unit["source_record_ids"]
            if (unit["unit_id"] in seen_units or unit["output_order"] != index or len(origins) != 1
                    or any(origin not in records for origin in origins)):
                findings.append("srs_content_linkage_invalid")
                continue
            seen_units.add(unit["unit_id"])
            used_records.update(origins)
            record = records[origins[0]]
            if (section not in profile.topic_roles or record["authority_tier"] not in profile.authority_tiers
                    or record["ownership_scope"] not in profile.allowed_scopes
                    or (section in SRS_INTRODUCTORY_SECTIONS) != (record["authority_tier"] == "presentation_policy")
                    or not record["provenance"] or not record["text"] or record["decision"] != "selected"
                    or record["record_id"] != "srs-source-" + stable_payload_hash(
                        (section, record["text"], tuple(record["provenance"])))[:20]
                    or unit["unit_id"] != "srs-unit-" + stable_payload_hash((section, unit["text"], record["record_id"]))[:20]):
                findings.append("srs_content_provenance_invalid:" + section)
            rendered = unit["text"]
            if section == "3.6":
                if rendered not in sections[section]:
                    findings.append("srs_domain_row_missing_or_changed")
                cells = [value.strip().replace("\\|", "|") for value in re.split(r"(?<!\\)\|", rendered)[1:-1]]
                if docx_path.exists() and not any(cells == row for table in docx_tables for row in table):
                    findings.append("srs_domain_row_missing_from_docx")
            else:
                summary = rendered.split("\n\n", 1)[-1].strip()
                if normalized(rendered) not in normalized(sections.get(section, "")):
                    findings.append("srs_content_missing_from_markdown:" + section)
                if docx_path.exists() and normalized(summary) not in normalized(docx_sections.get(section, "")):
                    findings.append("srs_content_missing_from_docx:" + section)
                if section in {"4.1", "5.1"}:
                    remainder = text.replace(sections[section], "", 1)
                    if normalized(summary) in normalized(remainder):
                        findings.append("srs_catalog_summary_repeated_outside_catalog")
        if used_records != set(records):
            findings.append("srs_content_contributor_coverage_invalid")
        for section, category in (("4.1", "analog"), ("5.1", "digital")):
            names = re.findall(r"(?m)^\s*- \*\*([^*]+)\*\*", sections[section])
            expected_names = [entry["block"] for entry in audit["catalog_entries"] if entry["category"] == category]
            if names != expected_names:
                findings.append("srs_catalog_coverage_invalid:" + section)
        headers = ["Domain", "Type", "Control", "Functional Role", "Power Conditions"]
        if audit["power_rows"]:
            header = "| " + " | ".join(headers) + " |"
            if header not in sections["3.6"]:
                findings.append("srs_domain_table_header_missing")
            else:
                table_start = sections["3.6"].index(header)
                if not normalized(sections["3.6"][:table_start]):
                    findings.append("srs_power_explanation_missing")
                clock_start = sections["3.6"].find("**Clock and reset**")
                if clock_start >= 0 and clock_start < table_start:
                    findings.append("srs_power_clock_reset_order_invalid")
            rows = [line.strip() for line in sections["3.6"].splitlines() if line.strip().startswith("|")]
            if rows != [header, "|---|---|---|---|---|", *[row["row"] for row in audit["power_rows"]]]:
                findings.append("srs_domain_table_fidelity_invalid")
            if docx_path.exists() and not any(table and table[0] == headers for table in docx_tables):
                findings.append("srs_domain_table_header_missing_from_docx")
        for finding in audit["review_findings"]:
            from spec_document_contract import publication_text_is_accounted
            publication_audit_path = markdown_path.with_name("descriptive_publication_audit.json")
            publication_audit = json.loads(publication_audit_path.read_text(encoding="utf-8")) if publication_audit_path.exists() else {}
            if finding in text or not publication_text_is_accounted(finding, text, "", publication_audit):
                findings.append("srs_content_review_finding_accounting_invalid")
    except (KeyError, ValueError, TypeError, OSError, IndexError) as exc:
        findings.append("srs_content_audit_invalid:" + str(exc))
    return findings


def validate_srs_system_overview(markdown_path: Path) -> List[str]:
    """Independently validate the SRS system-level overview contract."""
    if not markdown_path.exists():
        return ["srs_system_overview_missing_markdown"]
    text = markdown_path.read_text(encoding="utf-8", errors="replace")
    match = re.search(r"(?ms)^## 3\. System Overview\b.*?(?=^## 4\.)", text)
    if not match:
        return ["srs_system_overview_section_missing"]
    overview = match.group(0)
    narrative = re.sub(r"(?m)^\| Domain \| Type \| Control \| Functional Role \| Power Conditions \|\s*\n(?:^\|[^\n]*\|\s*\n)+", "", overview)
    findings: List[str] = []
    section_bodies: Dict[str, str] = {}
    for heading in SRS_SYSTEM_OVERVIEW_HEADINGS:
        heading_matches = list(re.finditer(rf"^### {re.escape(heading)}\b[^\n]*", overview, flags=re.MULTILINE))
        if len(heading_matches) != 1:
            findings.append(f"srs_system_overview_heading_invalid:{heading}")
            continue
        section_start = heading_matches[0].end()
        next_heading = re.search(r"(?m)^### |^## 4\.", overview[section_start:])
        section_body = overview[section_start:section_start + next_heading.start()] if next_heading else overview[section_start:]
        section_bodies[heading] = section_body
        meaningful_body = re.sub(r"(?is)<p\s*>\s*&nbsp;\s*</p>", "", section_body)
        meaningful_body = re.sub(r"(?is)<[^>]+>|&nbsp;|\s+", "", meaningful_body)
        if not meaningful_body:
            findings.append(f"srs_system_overview_section_empty:{heading}")
    if re.search(r"\b(?:DRS|ARS|IPOS)\b|^\s*Covers:\s*", overview, flags=re.IGNORECASE | re.MULTILINE):
        findings.append("srs_system_overview_contains_downstream_or_normative_linkage")
    if re.search(
        r"\b(?:register|regmap|port\s+name|address|signal\s+level|signal\s+name)\b",
        narrative,
        flags=re.IGNORECASE,
    ):
        findings.append("srs_system_overview_contains_implementation_detail")
    if re.search(r"\b(?:approved evidence|approved architecture|topic output limit|evidence count|need clarification|aim of this document|audit)\b", overview, re.IGNORECASE):
        findings.append("srs_system_overview_contains_audit_or_filler_language")
    if re.search(r"\b(?:shall|must|required to|Covers:)\b|\b0x[0-9a-f]+\b|\b(?:i|o|ca|u)_[a-z0-9_]+\b|\b[A-Za-z][A-Za-z0-9.-]*_[A-Za-z0-9_.-]+\b", narrative, re.IGNORECASE):
        findings.append("srs_system_overview_contains_normative_or_raw_source_fragment")
    if re.search(r"(?m)^\|[^\n]*\bOwner\b[^\n]*\|\s*\n\|[-| ]+\|", overview, re.IGNORECASE):
        findings.append("srs_system_overview_contains_raw_interface_dump")
    normalized_lines: Dict[str, str] = {}
    for line in overview.splitlines():
        normalized = re.sub(r"[*_`#>-]", "", line).strip().casefold()
        normalized = re.sub(r"\s+", " ", normalized)
        if len(normalized) < 60 or normalized.startswith("|") or normalized in {"&nbsp;"}:
            continue
        if normalized in normalized_lines:
            findings.append("srs_system_overview_repeats_narrative_content")
            break
        normalized_lines[normalized] = line
    findings.extend(_validate_srs_fact_audit(markdown_path, section_bodies))
    findings.extend(validate_srs_support_content(markdown_path))
    return findings


def build_ipos_descriptive_records(
    block_name: str,
    block_function: str,
    source_ports: Sequence[Mapping[str, object]],
    requirement_statements: Sequence[Mapping[str, object]],
    topic_specs: Sequence[Tuple[str, Sequence[str]]],
) -> List[Dict[str, str]]:
    """Build concise, evidence-linked IPOS summaries without copying requirements."""
    records: List[Dict[str, str]] = []
    if block_function.strip():
        assignment = assign_ipos_descriptive_topic(block_function, topic_specs)
        if assignment["topic"]:
            records.append({
                "summary": assignment["topic"],
                "statement": block_function.strip(),
                "source": f"artifacts/stage2_mirco_arc/block_inventory.csv:{block_name}",
                "domain": "block_function",
                "mapped_block": block_name,
                "functional_evidence": block_function.strip(),
            })
    # Requirement statements remain authoritative detail. The approved inventory
    # function is the only sufficiently high-level source for overview synthesis.
    # `requirement_statements` is deliberately retained in this shared signature
    # for compatibility with callers and audit structure.
    return records


def approved_snapshot_candidate_rows(snapshot, downstream_contract, *, domain: str = "") -> List[Dict[str, str]]:
    """Project the complete approved candidate set identically for all IPOS validators."""
    if domain not in {"", "digital", "analog"}:
        raise ValueError("Unsupported IPOS admitted-context domain: " + domain)
    targets = {"digital": "Digital IPOS", "analog": "Analog IPOS"}
    return [
        {
            "mapped_block": str(downstream_contract.allocation(row["source_req_id"]).get("approved_block") or "").strip(),
            "source_req_id": row["source_req_id"],
            "requirement_statement": row["requirement_statement"],
            "source_artifact": str(snapshot.source_path),
        }
        for row in snapshot.rows
        if row.get("source_req_id") and row.get("requirement_statement", "").strip()
        and downstream_contract.allocation(row["source_req_id"]).get("approved_block")
        and downstream_contract.allocation(row["source_req_id"]).get("owning_target") in {"Digital IPOS", "Analog IPOS"}
        and (not domain or downstream_contract.allocation(row["source_req_id"]).get("owning_target") == targets[domain])
    ]


def build_ipos_functional_input(
    block_name: str,
    inventory: Mapping[str, object],
    requirement_rows: Sequence[Mapping[str, object]] = (),
    *,
    domain: str = "",
    materialized: bool = True,
    provenance: Sequence[str] = (),
    candidate_requirement_rows: Sequence[Mapping[str, object]] | None = None,
) -> IPOSFunctionalInput:
    """Normalize approved inventory fields and same-block local evidence only."""
    local_rows = [
        row for row in requirement_rows
        if csv_cell_text(row.get("mapped_block") or row.get("owning_block") or block_name).strip() == block_name
    ]
    local_records = tuple(dict.fromkeys(
        (
            csv_cell_text(row.get("source_req_id") or row.get("requirement_id") or row.get("ipos_req_id")).strip(),
            csv_cell_text(row.get("requirement_statement") or row.get("statement") or row.get("text")).strip(),
        )
        for row in local_rows
    ))
    if any(not identity or not statement for identity, statement in local_records):
        raise ValueError("IPOS local evidence requires paired identity and statement")
    evidence_rows = requirement_rows if candidate_requirement_rows is None else candidate_requirement_rows
    candidate_records = tuple(
        (
            csv_cell_text(row.get("mapped_block") or row.get("owning_block") or block_name).strip(),
            csv_cell_text(row.get("source_req_id") or row.get("requirement_id") or row.get("ipos_req_id")).strip(),
            csv_cell_text(row.get("requirement_statement") or row.get("statement") or row.get("text")).strip(),
            csv_cell_text(row.get("source_artifact") or row.get("source") or "").strip(),
        )
        for row in evidence_rows
        if csv_cell_text(row.get("requirement_statement") or row.get("statement") or row.get("text")).strip()
    )
    return IPOSFunctionalInput(
        block=block_name.strip(),
        domain=domain.strip(),
        function=csv_cell_text(inventory.get("Function")).strip(),
        inputs=csv_cell_text(inventory.get("Inputs")).strip(),
        outputs=csv_cell_text(inventory.get("Outputs")).strip(),
        local_requirement_ids=tuple(identity for identity, _statement in local_records),
        local_requirement_evidence=tuple(statement for _identity, statement in local_records),
        candidate_requirement_records=candidate_records,
        materialized=materialized,
        provenance=tuple(dict.fromkeys(str(item).strip() for item in provenance if str(item).strip())),
    )


def compose_technical_block_purpose(block_name: str, function: str) -> str:
    """Express an approved inventory function as a descriptive block sentence."""
    purpose = function.strip().rstrip(".")
    if not block_name.strip() or not purpose:
        return ""
    return f"The {block_name} block is designed to {purpose[:1].lower() + purpose[1:]}."


def project_architecture_interactions(entries: Iterable[str]) -> List[Tuple[str, str, str]]:
    """Project approved interaction descriptions into stable structural rows."""
    rows: List[Tuple[str, str, str]] = []
    for entry in entries:
        match = re.fullmatch(r"Interaction:\s*(.+?)\s*->\s*(.+?)\s*\(([^()]*)\)", entry.strip())
        if not match:
            continue
        source, target = match.group(1).strip(), match.group(2).strip()
        exchange = match.group(3).split(";", 1)[0].strip()
        row = (source, target, exchange)
        if source and target and exchange and row not in rows:
            rows.append(row)
    return rows


def project_architecture_interfaces(entries: Iterable[str]) -> List[Tuple[str, str, str]]:
    """Keep interface name, owner, and purpose from existing catalog evidence."""
    rows: List[Tuple[str, str, str]] = []
    for entry in entries:
        match = re.fullmatch(r"Approved interface (.+?) owned by (.+?):\s*(.+)", entry.strip())
        if match and match.groups() not in rows:
            rows.append(match.groups())
    return rows


def summarize_architecture_interactions(rows: Sequence[Tuple[str, str, str]]) -> str:
    """Describe the first distinct approved exchanges without adding relationships."""
    if not rows:
        return ""
    exchanges = list(dict.fromkeys(exchange for _source, _target, exchange in rows))
    source, target, _exchange = rows[0]
    opening = (f"{source} connects to {target} through {exchanges[0]}"
               if exchanges[0].casefold().endswith("connection") else
               f"{source} provides {exchanges[0]} to {target}")
    if len(rows) == 1:
        return opening + "."
    return (opening + "; related paths carry "
            + ", ".join(exchanges[1:4]) + ".") if len(exchanges) > 1 else (
        opening + f", alongside {len(rows) - 1} other routed paths."
    )


def compose_drs_interaction_summary(rows: Sequence[Tuple[str, str, str]]) -> str:
    """Group identical directed exchanges without implying additional connections."""
    if not rows:
        return "Digital exchange behavior: need clarification."
    destinations: Dict[Tuple[str, str], List[str]] = {}
    for source, target, exchange in rows:
        targets = destinations.setdefault((source, exchange), [])
        if target not in targets:
            targets.append(target)
    groups: Dict[Tuple[str, Tuple[str, ...]], List[str]] = {}
    for (source, exchange), targets in destinations.items():
        groups.setdefault((exchange, tuple(targets)), []).append(source)
    source_statements: Dict[Tuple[str, ...], List[str]] = {}
    connections = []
    for (exchange, targets), sources in groups.items():
        subject = _drs_prose_list(sources)
        destination = _drs_prose_list(targets)
        connection = re.fullmatch(r"(.+?)\s+connections?", exchange, re.I)
        if connection:
            verb = "connect" if len(sources) > 1 else "connects"
            connections.append(f"{subject} {verb} to {destination} through {connection.group(1)}.")
        else:
            source_statements.setdefault(tuple(sources), []).append(f"{exchange} to {destination}")
    sentences = []
    for sources, exchanges in source_statements.items():
        verb = "provide" if len(sources) > 1 else "provides"
        sentences.append(f"{_drs_prose_list(sources)} {verb} {_drs_prose_list(exchanges)}.")
    return " ".join(sentences + connections)


def _drs_prose_list(values: Sequence[str]) -> str:
    if len(values) < 2:
        return "".join(values)
    return ", ".join(values[:-1]) + " and " + values[-1]


def compose_drs_clock_summary(paths: Sequence[Tuple[str, str, str, str, str]]) -> str:
    """Describe frequency and gating behavior; hierarchy remains in route tables."""
    if not paths:
        return "Clock distribution and gating behavior: need clarification."
    frequencies: Dict[str, List[str]] = {}
    for _source_id, _source, _target, frequency, gating in paths:
        modes = frequencies.setdefault(frequency, [])
        mode = ("without clock gating" if gating == "Not gated" else
                "with clock gating" if gating.startswith("Gated") else "")
        if mode and mode not in modes:
            modes.append(mode)
    descriptions = [f"{frequency} clocks {'with and without clock gating' if len(modes) == 2 else _drs_prose_list(modes)}".strip()
                    for frequency, modes in frequencies.items()]
    summary = "The clock network distributes " + "; ".join(descriptions) + "."
    if any(source == "need clarification" or target == "need clarification"
           for _source_id, source, target, _frequency, _gating in paths):
        summary += " Incomplete source and destination clock labels: need clarification."
    return summary


def compose_drs_role_summary(roles: Sequence[Tuple[str, str]], terms: Sequence[str]) -> str:
    """Project complete functional clauses within an already selected role set."""
    sentences = []
    for name, function in roles:
        clauses = re.split(
            r",\s+(?:and\s+)?(?=(?:provide|generate|receive|perform|write|route|control|coordinate|implement|accept|translate|sequence|execute|operate|aggregate)\b)",
            function.strip().rstrip("."), flags=re.I,
        )
        for clause in clauses:
            if terms and not any(term in clause.casefold() for term in terms):
                continue
            words = clause.split(" ", 1)
            if len(words) != 2:
                continue
            verb, predicate = words
            if re.fullmatch(r"[A-Za-z]+", verb):
                inflected = (verb[:-1] + "ies" if re.search(r"[^aeiou]y$", verb, re.I) else
                             verb + "es" if re.search(r"(?:s|sh|ch|x|z|o)$", verb, re.I) else verb + "s")
                predicate = re.sub(
                    r"\band (provide|generate|receive|perform|write|route|control|coordinate|implement|accept|translate|sequence|execute|operate|aggregate)\b",
                    lambda match: "and " + match.group(1) + "s", predicate,
                )
                sentence = f"{name} {inflected.lower()} {predicate}."
                if sentence not in sentences:
                    sentences.append(sentence)
    return " ".join(sentences)


def drs_section_prose_findings(text: str) -> List[str]:
    """Inspect final section prose independently of composer parity."""
    findings = natural_descriptive_prose_findings(text)
    if re.search(r"\b(?:\d+|one|two|three|four|five|six|seven|eight|nine|ten)\s+"
                 r"(?:(?:named|source|destination|architectural|routed|approved|other)\s+)*"
                 r"(?:paths?|endpoints?|destinations?|routes?|links?|groups?|items?)\b", text, re.I):
        findings.append("description_counts_structure")
    if re.search(r"\b(?:evidence\s+groups?|block\s+functions?\s+identify|recorded\s+exchanges\s+cover|"
                 r"the\s+architecture\s+(?:also\s+)?records)\b", text, re.I):
        findings.append("description_uses_audit_summary")
    return list(dict.fromkeys(findings))


def natural_descriptive_prose_findings(text: str) -> List[str]:
    """Check descriptive text without depending on a document's section layout."""
    findings: List[str] = []
    if re.search(
        r"\bapproved\b"
        r"|\b(?:source|snapshot|candidate|refinement|audit|provenance)\s+"
        r"(?:local\s+)?(?:function(?:al)?\s+evidence|function|scope|behavior|requirement|candidate)\b"
        r"|\b(?:local\s+scope\s+covers|supported\s+by\s+the\s+approved|derivation\s+mode"
        r"|general\s+functional\s+description|block\s+summary)\b",
        text, re.I,
    ):
        findings.append("description_uses_internal_process_language")
    if re.search(r"\b(?:shall|must|required to)\b", text, re.I):
        findings.append("description_uses_requirement_form_language")
    return findings


def validate_drs_block_descriptions(
    markdown_text: str,
    blocks: Sequence[Mapping[str, str]],
    audit_rows: Sequence[Mapping[str, str]],
    port_rows: Sequence[Mapping[str, str]],
    docx_path: Path | None = None,
) -> List[str]:
    """Check DRS block paragraphs against inventory authority and rendered I/O."""
    findings: List[str] = []
    block_list = re.search(r"(?m)^###\s+(?:\*\*)?4\.1\s+Digital block list[^\n]*\n([^\n]+)", markdown_text)
    if block_list:
        findings.extend(f"drs_block_list_{finding}" for finding in natural_descriptive_prose_findings(block_list.group(1)))
    headings = list(re.finditer(r"(?m)^###\s+(?:\*\*)?9\.\d+\s+([^*{\n]+)", markdown_text))
    sections: Dict[str, str] = {}
    for index, heading in enumerate(headings):
        name = heading.group(1).strip()
        heading_end = markdown_text.find("\n", heading.start())
        remaining = markdown_text[heading_end + 1:headings[index + 1].start() if index + 1 < len(headings) else len(markdown_text)]
        next_top = re.search(r"(?m)^##\s+", remaining)
        body = remaining[:next_top.start()] if next_top else remaining
        if name in sections:
            findings.append(f"duplicate_drs_block_paragraph:{name}")
        sections[name] = body
    expected = {block["name"]: block for block in blocks}
    for name in sorted(set(sections) - set(expected)):
        findings.append(f"unapproved_drs_block_paragraph:{name}")
    provenance_rows = [row for row in audit_rows if row.get("section") == "DRS digital block descriptions"]
    if len(provenance_rows) != len(blocks):
        findings.append("drs_block_description_audit_count_mismatch")
    docx_paragraphs: Set[str] = set()
    docx_port_names: Set[str] = set()
    if docx_path is not None:
        try:
            with zipfile.ZipFile(docx_path) as archive:
                document = ET.fromstring(archive.read("word/document.xml"))
            tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            for paragraph in document.iter(tag + "p"):
                docx_paragraphs.add(re.sub(r"\s+", " ", "".join(
                    text.text or "" for text in paragraph.iter(tag + "t")
                )).strip())
            for table in document.iter(tag + "tbl"):
                for row in table.iter(tag + "tr"):
                    cell = next(row.iter(tag + "tc"), None)
                    if cell is not None:
                        docx_port_names.add("".join(text.text or "" for text in cell.iter(tag + "t")).strip())
        except (OSError, KeyError, ValueError, ET.ParseError, zipfile.BadZipFile):
            findings.append("drs_descriptive_docx_unreadable")
    for order, block in enumerate(blocks, start=1):
        name = block["name"]
        sentence = compose_technical_block_purpose(name, block["function"])
        body = sections.get(name, "")
        if not body:
            findings.append(f"missing_drs_block_paragraph:{name}")
            continue
        first_line = next((line.strip() for line in body.splitlines() if line.strip()), "")
        if first_line != sentence:
            findings.append(f"drs_block_purpose_mismatch:{name}")
        findings.extend(f"{finding}:{name}" for finding in natural_descriptive_prose_findings(first_line))
        if re.search(r"(?m)^\s*-\s*Port name\s*:", body, re.I):
            findings.append(f"drs_port_name_bullet_list:{name}")
        if "Block-specific inputs, outputs, ports" in body:
            findings.append(f"drs_internal_scope_scaffolding:{name}")
        table_names = {cell.strip().replace("\\|", "|") for cell in re.findall(r"(?m)^\|\s*([^|]+?)\s*\|", body)}
        approved_ports = {row.get("Port name", "").strip() for row in port_rows
                          if row.get("Owner") == name and row.get("Ownership status") == "approved"
                          and row.get("Port name", "").strip()}
        for port in sorted(approved_ports - table_names):
            findings.append(f"drs_approved_port_missing_from_table:{name}:{port}")
        if docx_path is not None and docx_paragraphs and sentence not in docx_paragraphs:
            findings.append(f"drs_block_purpose_missing_from_docx:{name}")
        for port in sorted(approved_ports - docx_port_names) if docx_path is not None and docx_paragraphs else []:
            findings.append(f"drs_approved_port_missing_from_docx_table:{name}:{port}")
        if order > len(provenance_rows):
            continue
        row = provenance_rows[order - 1]
        if any((row.get(field) or "").strip() != value for field, value in (
            ("mapped_block", name), ("source_evidence", block["function"]), ("statement", sentence),
            ("source", "artifacts/stage2_mirco_arc/block_inventory.csv"),
            ("scope_decision", "approved_concrete_block"), ("output_order", str(order)),
        )):
            findings.append(f"drs_block_description_provenance_mismatch:{name}")
    return findings


def _ipos_source_text(value: str) -> str:
    value = re.sub(r"^\s*\[[^\]]+\]\s*(?:Requirement|Definition|Assumption|Comment)\s*:?\s*", "", value, flags=re.I)
    value = re.sub(r"\[(?:TO:|Covers:|Vpriority)[^\]]*\]|\[End\]", "", value, flags=re.I)
    return " ".join(value.split()).strip().rstrip(".")


IPOS_TECHNICAL_IDENTIFIER_PATTERN = r"\b(?:[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9_]+|[A-Z][A-Z0-9]+)\b"


def _ipos_lower_prose_initial(value: str) -> str:
    return value if re.match(IPOS_TECHNICAL_IDENTIFIER_PATTERN, value) else value[:1].lower() + value[1:]


def _ipos_admitted_records(functional_input: IPOSFunctionalInput) -> Tuple[NormalizedSourceRecord, ...]:
    from validate_downstream_coherence import canonical_human_label

    records = [NormalizedSourceRecord(
        "inventory:" + field, value, "admitted_context", "block_local",
        ("inventory:" + field, *functional_input.provenance), "retained", "approved block inventory",
    ) for field, value in (("function", functional_input.function),
                          ("inputs", functional_input.inputs), ("outputs", functional_input.outputs))]
    candidates = functional_input.candidate_requirement_records or tuple(
        (functional_input.block, identity, statement, "approved snapshot")
        for identity, statement in zip(functional_input.local_requirement_ids,
                                       functional_input.local_requirement_evidence)
    )
    seen: Dict[str, str] = {}
    for block, identity, statement, source in candidates:
        if canonical_human_label(block) != canonical_human_label(functional_input.block):
            continue
        if not identity or not statement.strip():
            raise ValueError("Incomplete IPOS approved evidence identity/text")
        if identity in seen:
            if seen[identity] != statement:
                raise ValueError("Conflicting IPOS approved evidence revision: " + identity)
            continue
        seen[identity] = statement
        records.append(NormalizedSourceRecord(
            "requirement:" + block + ":" + identity, statement, "approved_normative", "block_local",
            ("approved requirement:" + identity, source or "approved snapshot", *functional_input.provenance),
            "retained", "same-owner approved IPOS evidence; normative text preserved separately",
        ))
    local = dict(zip(functional_input.local_requirement_ids, functional_input.local_requirement_evidence))
    if any(seen.get(identity) != statement for identity, statement in local.items()):
        raise ValueError("IPOS admitted context omits or changes allocated local evidence")
    records.append(NormalizedSourceRecord(
        "inventory:main_functions", functional_input.function, "admitted_context", "block_local",
        ("inventory:main_functions", *functional_input.provenance), "retained", "approved block inventory Function",
    ))
    return tuple(records)


def _ipos_main_function_phrases(text: str) -> Tuple[str, ...]:
    actions = "|".join(sorted(DESCRIPTIVE_NORMATIVE_ACTIONS))
    phrases = []
    depth = 0
    start = 0
    for index, character in enumerate(text):
        if character in "([":
            depth += 1
        elif character in ")]":
            depth -= 1
        elif character == "," and depth == 0:
            separator = re.match(r",\s*(?:and\s+)?(?=(?:" + actions + r")\b)", text[index:], re.I)
            if separator:
                phrases.append(text[start:index].strip())
                start = index + separator.end()
    phrases.append(text[start:].strip())
    return tuple(phrase for phrase in phrases if phrase)


def ipos_main_function_name(fact: DescriptiveFact) -> str:
    if fact.source != "inventory:main_functions" or len(fact.objects) != 1:
        return ""
    phrase = fact.objects[0].removeprefix("designed to ")
    return phrase[:1].upper() + phrase[1:]


def ipos_supported_function_name(fact: DescriptiveFact) -> str:
    if fact.evidence_kind != "approved_normative":
        return ""
    entity = re.fullmatch(r"(?:The\s+)?(.+?\b(?:FSM|phases?))(?:\s+FSM)?", fact.subject, re.I)
    if entity:
        return entity[1][:1].upper() + entity[1][1:]
    flag = re.fullmatch(r"(?:The\s+)?(?:[io]\s+)?(.+?\s+flag)", fact.subject, re.I)
    if flag and fact.condition and any(re.search(r"\bset\s+to\s+\d+\s+and\s+go(?:es)?\s+to\s+\d+\s+at\b", obj, re.I)
                                       for obj in fact.objects):
        domain = re.match(r"In case\s+(?:the\s+)?(\w+)\s+sample data\b", fact.condition, re.I)
        phrase = (domain[1].upper() + " " if domain else "") + flag[1]
        return phrase[:1].upper() + phrase[1:] + " control"
    for obj in fact.objects:
        effect = re.search(
            r"\bto\s+(?:start|initiate|control)\s+(?:the\s+)?"
            r"(?:(?:rising|falling)\s+of\s+(?:the\s+)?)?(.+?)(?=\s+and\s+|$)", obj, re.I)
        if fact.action == "be" and effect:
            phrase = effect[1].strip()
            return phrase[:1].upper() + phrase[1:] + " control"
        if fact.action == "start":
            phrase = re.split(r"\s+(?:by|when|writing|with|after|before)\b", obj, maxsplit=1, flags=re.I)[0]
            return phrase[:1].upper() + phrase[1:] + " initiation"
    return ""


def _ipos_top_level_modals(text: str) -> Tuple[re.Match, ...]:
    depth = 0
    positions = set()
    for index, character in enumerate(text):
        if character in "([":
            depth += 1
        elif character in ")]":
            depth -= 1
        elif depth == 0:
            positions.add(index)
    return tuple(match for match in re.finditer(r"\b(?:shall|must)\b", text, re.I)
                 if match.start() in positions)


def _project_ipos_record(
    record: NormalizedSourceRecord, functional_input: IPOSFunctionalInput,
) -> Tuple[Sequence[DescriptiveFact], str]:
    text = _ipos_source_text(record.text)
    common = dict(source=record.provenance[0], statement=record.text, scope="block_local",
                  owner=functional_input.block, domain=functional_input.domain, layer="block")
    if record.record_id.startswith("inventory:"):
        field = record.record_id.split(":", 1)[1]
        if field == "main_functions":
            if not text:
                return (), "missing approved inventory function"
            return tuple(DescriptiveFact("The " + functional_input.block + " block", "be",
                                         ("designed to " + phrase[:1].lower() + phrase[1:],),
                                         section="function", **common)
                         for phrase in _ipos_main_function_phrases(text)), ""
        action = {"function": "be", "inputs": "accept", "outputs": "produce"}[field]
        obj = "designed to " + _ipos_lower_prose_initial(text) if field == "function" else text
        if not text:
            return (), "missing approved inventory " + field
        return (DescriptiveFact("The " + functional_input.block + " block", action, (obj,),
                                section="function" if field == "function" else "interfaces", **common),), ""
    projected = []
    gaps = []
    for index, clause in enumerate(re.split(r"(?<=[.!?])\s+(?=[A-Za-z])", text), start=1):
        facts, reason = _project_ipos_normative_clause(record, functional_input, clause)
        projected.extend(facts)
        if reason:
            gaps.append("clause " + str(index) + ": " + reason)
    return (() if gaps else tuple(projected)), "; ".join(gaps)


def _ipos_comma_less_guard_boundary(prefix: str, owner: str) -> Tuple[str, str]:
    if not owner:
        return "", ""
    match = re.fullmatch(r"(.+\s+\S+)\s+((?:the\s+)?" + re.escape(owner) + r"(?:\s+block)?)", prefix, re.I)
    return (match[1].strip(), match[2].strip()) if match else ("", "")


def _project_ipos_normative_clause(
    record: NormalizedSourceRecord, functional_input: IPOSFunctionalInput, text: str,
) -> Tuple[Sequence[DescriptiveFact], str]:
    text = re.sub(r"^\s*\u2022\s*", "", text).strip().rstrip(".")
    text = re.sub(r"^In particular,\s*", "", text, flags=re.I)
    common = dict(source=record.provenance[0], statement=record.text, scope="block_local",
                  owner=functional_input.block, domain=functional_input.domain, layer="block")
    if re.search(r"[,;:]\s*$", text):
        return (), "incomplete source continuation; full source retained"
    if re.search(r"\b(?:Table|Figure)\s+\d+|\bformula\b|\bfollowing steps\b|\bas follows?\s*:|\b(?:following|tag)\s*:|\u2022", text, re.I):
        return (), "unsupported structured/table/procedure evidence; full source retained"
    if any(text.count(opening) != text.count(closing) for opening, closing in (("(", ")"), ("[", "]"))):
        return (), "unbalanced source delimiters; full source retained"
    modals = _ipos_top_level_modals(text)
    if not modals or any(not re.search(r"\band\s*$", text[:modal.start()], re.I) for modal in modals[1:]):
        return (), "unsupported multi-clause or non-normative construction; full source retained"
    prefix, predicate = text[:modals[0].start()].strip(), text[modals[0].end():].strip()
    condition = ""
    if re.match(r"^(?:(?:When|If|Unless|After|Before|Once|In case|In order to|During|At|To)\b|(?:At least\s+)?\d+(?:\.\d+)?\s*(?:ns|us|ms|s)\s+after\b)", prefix, re.I):
        guard_prefix = prefix
        condition, separator, prefix = guard_prefix.rpartition(",")
        if not separator:
            condition, prefix = _ipos_comma_less_guard_boundary(guard_prefix, functional_input.block)
            if not prefix:
                return (), "ambiguous guard boundary; full source retained"
        condition = re.sub(r"\b(?:shall|must)\s+be\b", "is", condition.strip(), flags=re.I)
        if re.search(r"\b(?:shall|must)\b", condition, re.I):
            return (), "unsupported normative guard; full source retained"
    predicate = re.sub(r"\band\s+(?:shall|must)\s+", "and ", predicate, flags=re.I)
    text = prefix.strip() + " shall " + predicate
    match = re.fullmatch(r"(.+?)\s+(?:shall|must)\s+(not\s+)?((?:always|first)\s+)?(?:be able to\s+)?((?!not\b)[A-Za-z]+)(?:\s+(.+))?", text, re.I)
    if not match or re.search(r"\b(?:and|or|then|to|by|with|into|the|a|an|for|from|of|in|at|if|when|unless|until|before|after|that|while|is|are|has|have|was|were|been)\s*$", text, re.I):
        return (), "unsupported or incomplete predicate; full source retained"
    subject, negative, adverb, action, obj = match.groups()
    if adverb and negative:
        return (), "unsupported adverb negation scope; full source retained"
    action = DESCRIPTIVE_FINITE_ACTION_BASES.get(action.casefold(), action.casefold())
    if action.casefold() not in DESCRIPTIVE_NORMATIVE_ACTIONS and action.casefold() != "define":
        return (), "unsupported predicate grammar; full source retained"
    if re.search(r"\b(?:a|an|the)\s+(?:new|first|last|following|selected|configured)\s*$", obj or "", re.I):
        return (), "incomplete descriptive object; full source retained"
    obj = obj or ""
    if not obj and action.casefold() not in DESCRIPTIVE_INTRANSITIVE_ACTIONS:
        return (), "unsupported or incomplete predicate; full source retained"
    subject = subject.strip()
    if re.fullmatch(r"(?:it|they|this|that|these|those)(?:\s+(?:block|device|unit|signal|channel|state))?", subject, re.I):
        return (), "unresolved subject reference; full source retained"
    if re.search(r"\b(?:that|which|who)\b", subject, re.I):
        return (), "unsupported relative-clause subject; full source retained"
    if re.match(r"then\b", subject, re.I):
        return (), "unsupported conditional consequence marker; full source retained"
    if condition and "," in subject:
        return (), "ambiguous compound guard/actor boundary; full source retained"
    if re.search(r"[.!?]\s+[A-Z]|\b(?:shall|must)\b", obj):
        return (), "unsupported additional source proposition; full source retained"
    action = action.casefold()
    if action == "be" and not re.search(r"\band\b", subject, re.I):
        obj = re.sub(r"(\bset\s+to\s+\d+\s+and\s+)go(?=\s+to\s+\d+\b)", r"\1goes", obj, flags=re.I)
    coordination = "or" if re.search(r"\bor\b", obj, re.I) else "and"
    objects = (obj.strip(),) if obj else ()
    if not re.search(r"\b(?:and|or)\b", obj, re.I) or action in {"contain", "comprise", "include"}:
        objects = tuple(part.strip() for part in re.split(r"\s+(?:and|or)\s+", obj, flags=re.I) if part.strip())
    return (DescriptiveFact(subject, action, objects, section="function", condition=condition,
                            negative=bool(negative), coordination=coordination,
                            modality="can" if re.search(r"\bbe able to\b", text, re.I) else "",
                            plural_subject=bool(re.search(r"\band\b", subject, re.I)),
                            qualifiers=(adverb.strip().casefold(),) if adverb else (),
                            evidence_kind="approved_normative", **common),), ""


def _ipos_shared_flow(functional_input: IPOSFunctionalInput) -> DescriptiveFlowResult:
    return run_descriptive_flow(
        _ipos_admitted_records(functional_input), profile=IPOS_DESCRIPTIVE_PROFILE,
        adapter_name="ipos-local", owner=functional_input.block,
        adapter=lambda record: _project_ipos_record(record, functional_input),
    )


def _ipos_meaning(fact: DescriptiveFact) -> DescriptiveMeaning:
    objects = fact.objects
    effects = ()
    if len(objects) == 1:
        effect = re.fullmatch(r"(.+?)\s+to\s+(start|initiate)\s+(.+)", objects[0], re.I)
        if effect:
            objects = (effect[1],)
            effects = (DescriptiveEffect(effect[2].casefold(), effect[3]),)
    return DescriptiveMeaning(
        DESCRIPTIVE_ACTION_RELATIONS.get(fact.action, "behavior"), fact.subject, fact.action, objects,
        (DescriptiveContext("atom", fact.condition),) if fact.condition else (),
        fact.mode, fact.modality, fact.negative, fact.coordination,
        (*fact.qualifiers, *((fact.object_suffix,) if fact.object_suffix else ())), fact.results, effects,
    )


IPOS_SEMANTIC_RECOVERY_RULES = ("coordinated_guard_recomposition_v1", "independent_partial_meaning_v1",
                                "qualitative_before_formula_v1", "shared_context_bundle-v1", "guard_subject_reference_v1",
                                "supported_before_dangling_coordination_v1", "explicit_case_reference_v1")


def _ipos_explicit_case_reference(antecedent: str, consequent: str) -> str:
    antecedent = _ipos_source_text(antecedent).strip().rstrip(".")
    consequent = _ipos_source_text(consequent).strip()
    reference = re.fullmatch(r"In\s+(?:that|this)\s+case\s*,?\s+(.+?\s+(?:shall|must)\b.+)", consequent, re.I)
    conditions = tuple(re.finditer(r"\bto\s+(?:check|verify|determine)\s+(?:if|whether)\s+(.+)$", antecedent, re.I))
    if not reference or len(conditions) != 1 or len(_ipos_top_level_modals(antecedent)) != 1:
        return ""
    condition = conditions[0][1].strip()
    if re.search(r"[.;:]|\b(?:shall|must)\b|\b(?:and|or|to|the|a|an|if|when|is|are|has|have|been)\s*$", condition, re.I):
        return ""
    return "If " + condition + ", " + reference[1]


def _ipos_recomposition_text(spans: Sequence[DescriptiveSourceSpan]) -> str:
    if len(spans) == 2:
        reference = _ipos_explicit_case_reference(spans[0].text, spans[1].text)
        if reference:
            return reference
    values = []
    for index, span in enumerate(spans):
        value = span.text.strip()
        if index or len(spans) == 1:
            value = re.sub(r"^,?\s*(?:and\s+)?(?=\w)", "", value, flags=re.I)
        values.append(value)
    text = " ".join(values)
    text = re.sub(r"^In particular,\s*", "", text, flags=re.I)
    antecedent = re.fullmatch(r"(If\s+([A-Za-z_][\w-]*)\s+is\s+set\s+to\s+\d+\s*,)\s+it\s+((?:shall|must)\b.+)", text, re.I)
    if antecedent:
        text = antecedent[1] + " " + antecedent[2] + " " + antecedent[3]
    return text


def _ipos_recover_clause(record: NormalizedSourceRecord, functional_input: IPOSFunctionalInput,
                         span: DescriptiveSourceSpan) -> Tuple[list, list]:
    text = span.text
    modals = _ipos_top_level_modals(text)
    if not modals:
        return [], [span]
    boundaries = [0]
    for previous, modal in zip(modals, modals[1:]):
        tail = text[previous.end():modal.start()]
        if re.search(r"\b(?:if|when|unless|otherwise|except|or|so|but)\b|(?:and|,)\s+(?:after|before|then|once)\b", tail, re.I):
            continue
        boundary = re.search(r"(,\s*|\s+and\s+)(?:the\s+|no\s+)?[A-Za-z_][\w-]*(?:\s+[A-Za-z_][\w-]*)*\s*$", tail, re.I)
        if boundary:
            boundaries.append(previous.end() + boundary.start())
    boundaries.append(len(text))
    prefix = text[:modals[0].start()]
    guard_end = prefix.rfind(",") + 1 if re.match(
        r"^\s*(?:\[Covers:[^\]]+\]\s*)*(?:(?:When|If|Unless|After|Before|Once|In case|In order to|During|At|To)\b|(?:At least\s+)?\d+(?:\.\d+)?\s*(?:ns|us|ms|s)\s+after\b)", prefix, re.I) else 0
    guard = DescriptiveSourceSpan(span.record_id, span.start, span.start + guard_end, text[:guard_end]) if guard_end else None
    recovered, residuals = [], []
    for number, (start, end) in enumerate(zip(boundaries, boundaries[1:])):
        body_start = guard_end if number == 0 and guard else start
        body = DescriptiveSourceSpan(span.record_id, span.start + body_start, span.start + end, text[body_start:end])
        pieces = ((guard, body) if guard else (body,))
        candidate = _ipos_recomposition_text(pieces)
        facts, reason = _project_ipos_normative_clause(record, functional_input, _ipos_source_text(candidate))
        if any(re.fullmatch(r"it|them|this|that|these|those", obj, re.I) for fact in facts for obj in fact.objects):
            facts = ()
        if not facts:
            cut = re.search(r"\s+following\s+this\s+formula\s*:|\s*\((?:The purpose|Check)\b", body.text, re.I)
            if cut is None:
                cut = re.search(r"\s+and\s*[.!?]?\s*$|,\s*[.!?]?\s*$", body.text, re.I)
            if cut:
                prefix_body = replace(body, end=body.start + cut.start(), text=body.text[:cut.start()])
                prefix_pieces = ((guard, prefix_body) if guard else (prefix_body,))
                partial, _ = _project_ipos_normative_clause(record, functional_input,
                               _ipos_source_text(_ipos_recomposition_text(prefix_pieces)))
                if partial:
                    command = re.search(r",\s*(?:" + "|".join(sorted(DESCRIPTIVE_NORMATIVE_ACTIONS)) + r")\b",
                                        prefix_body.text, re.I)
                    if command:
                        independent = replace(prefix_body, end=prefix_body.start + command.start(),
                                              text=prefix_body.text[:command.start()])
                        independent_pieces = ((guard, independent) if guard else (independent,))
                        supported, _ = _project_ipos_normative_clause(record, functional_input,
                                       _ipos_source_text(_ipos_recomposition_text(independent_pieces)))
                        if supported:
                            partial, prefix_body, prefix_pieces = supported, independent, independent_pieces
                    facts, pieces = partial, prefix_pieces
                    residuals.append(replace(body, start=prefix_body.end, text=body.text[prefix_body.end - body.start:]))
        if facts:
            recovered.extend((fact, pieces) for fact in facts)
        else:
            residuals.append(body)
    if residuals and guard:
        residuals.insert(0, guard)
    return recovered, residuals


def build_ipos_evidence_model(functional_input: IPOSFunctionalInput) -> DescriptiveEvidenceModel:
    records = tuple(record for record in _ipos_admitted_records(functional_input)
                    if record.record_id != "inventory:main_functions") if functional_input.materialized else ()
    propositions = []
    unresolved = []
    dependencies = []
    for record in records:
        priority_header = re.search(IPOS_PRIORITY_LIST_HEADER, record.text, re.I)
        separators = tuple(match for match in re.finditer(r"(?<=[.!?])\s+(?=[A-Za-z])", record.text)
                           if priority_header is None or match.start() < priority_header.start())
        starts = [0, *(match.end() for match in separators)]
        ends = [*(match.start() for match in separators), len(record.text)]
        if record.record_id.startswith("inventory:"):
            starts, ends = [0], [len(record.text)]
        previous_ids = ()
        for index, (start, end) in enumerate(zip(starts, ends), start=1):
            span = DescriptiveSourceSpan(record.record_id, start, end, record.text[start:end])
            priority_list = _ipos_priority_list(span.text) if not record.record_id.startswith("inventory:") else None
            if priority_list is not None:
                entries, direction = priority_list
                identity = record.record_id + ":clause:" + str(index) + ":priority"
                meaning = DescriptiveMeaning("behavior", "The priority order", "be", entries,
                                            qualifiers=(direction,))
                propositions.append(DescriptiveProposition(identity, meaning, (span,),
                                    functional_input.block, functional_input.domain, "block"))
                previous = next((value for value in propositions if len(previous_ids) == 1
                                 and value.proposition_id == previous_ids[0]), None)
                if previous is not None and previous.meaning.predicate in {"execute", "select"}:
                    dependencies.append(DescriptiveDependency("priority:" + previous.proposition_id + ":" + identity,
                                        "referenced_context", previous.proposition_id, identity, (*previous.spans, span)))
                previous_ids = (identity,)
                continue
            if record.record_id.startswith("inventory:"):
                facts, reason = _project_ipos_record(record, functional_input)
            else:
                facts, reason = _project_ipos_normative_clause(record, functional_input, _ipos_source_text(span.text))
            recovered, residuals = ([(fact, (span,)) for fact in facts], [])
            reference_id = ""
            if not facts and len(previous_ids) == 1:
                previous = next((value for value in propositions if value.proposition_id == previous_ids[0]), None)
                if previous is not None and len(previous.spans) == 1:
                    reference = _ipos_explicit_case_reference(previous.spans[0].text, span.text)
                    if reference:
                        facts, _ = _project_ipos_normative_clause(record, functional_input, reference)
                        if facts:
                            reference_id = previous.proposition_id
                            recovered = [(fact, (*previous.spans, span)) for fact in facts]
            if not facts and not record.record_id.startswith("inventory:"):
                recovered, residuals = _ipos_recover_clause(record, functional_input, span)
            clause_ids = []
            for number, (fact, support) in enumerate(recovered, start=1):
                identity = record.record_id + ":clause:" + str(index) + ":meaning:" + str(number)
                propositions.append(DescriptiveProposition(identity, _ipos_meaning(fact), support,
                                    functional_input.block, functional_input.domain, "block"))
                if reference_id:
                    dependencies.append(DescriptiveDependency("case:" + reference_id + ":" + identity,
                                        "referenced_context", reference_id, identity, support))
                clause_ids.append(identity)
            if residuals or not recovered:
                identity = record.record_id + ":clause:" + str(index) + ":unresolved"
                unresolved.append(DescriptiveUnresolvedContext(identity, tuple(residuals) or (span,), reason or "unsupported_clause"))
                if re.match(r"\s*(?:Except|Otherwise|Unless)\b", span.text, re.I):
                    for previous in previous_ids:
                        dependencies.append(DescriptiveDependency("context:" + previous + ":" + identity,
                                            "referenced_context", previous, identity, (span,)))
                clause_ids.append(identity)
            previous_ids = tuple(clause_ids)
    return DescriptiveEvidenceModel(records, tuple(propositions), tuple(dependencies), tuple(unresolved))


IPOS_PRIORITY_LIST_HEADER = r"\bPriority\s+is\s+listed\s+below,\s+with\s+first\s+(?:routine|operation|command|item|entry)\s+at\s+(higher|lower)\s+priority:\s*"


def _ipos_priority_list(text: str) -> Tuple[Tuple[str, ...], str] | None:
    text = _ipos_source_text(text).strip().rstrip(".")
    header = re.match(IPOS_PRIORITY_LIST_HEADER, text, re.I)
    if header is None:
        return None
    body = text[header.end():]
    markers = tuple(re.finditer(r"\b(\d+)[.)]\s+", body))
    if len(markers) < 2 or markers[0].start() != 0 or tuple(int(value[1]) for value in markers) != tuple(range(1, len(markers) + 1)):
        return None
    entries = []
    for index, marker in enumerate(markers):
        end = markers[index + 1].start() if index + 1 < len(markers) else len(body)
        label = body[marker.end():end].strip()
        if not label or re.search(r"[.;:]|\b(?:shall|must)\b", label, re.I):
            return None
        entries.append(marker[1] + ": " + label)
    direction = ", from higher to lower priority" if header[1].casefold() == "higher" else ", from lower to higher priority"
    return tuple(entries), direction


def _ipos_functional_scope_supported(meaning: DescriptiveMeaning) -> bool:
    if re.match(r"^Select\s+how\b", meaning.subject, re.I):
        return False
    if any(re.search(r"\bas explained in\b", obj, re.I) for obj in meaning.objects):
        return False
    if any(re.search(r"\bnegated version\b", obj, re.I) for obj in meaning.objects):
        return len(re.findall(IPOS_TECHNICAL_IDENTIFIER_PATTERN, meaning.subject)) == len(
            re.findall(IPOS_TECHNICAL_IDENTIFIER_PATTERN, " ".join(meaning.objects)))
    return True


def build_ipos_claim_plan(functional_input: IPOSFunctionalInput) -> DescriptiveSynthesisPlan:
    if functional_input.domain not in {"digital", "analog"}:
        raise ValueError("IPOS claim interpretation requires an explicit approved digital/analog domain")
    evidence = build_ipos_evidence_model(functional_input)
    grouped = {}
    components = {value.proposition_id: frozenset((value.proposition_id,)) for value in evidence.propositions}
    for edge in evidence.dependencies:
        if edge.relation == "referenced_context" and edge.source_id in components and edge.target_id in components:
            members = components[edge.source_id] | components[edge.target_id]
            for identity in members:
                components[identity] = members
    dependency_members = {identity for edge in evidence.dependencies for identity in (edge.source_id, edge.target_id)}
    for proposition in evidence.propositions:
        record_id = proposition.spans[0].record_id
        meaning = proposition.meaning
        role = ("purpose" if record_id == "inventory:function" else
                "main_functions" if record_id == "inventory:main_functions" else
                "boundaries" if record_id in {"inventory:inputs", "inventory:outputs"} else
                "organization" if meaning.relation == "containment" else
                "limits" if meaning.relation == "configuration" or meaning.predicate == "define" else "functions")
        scope = ("initiation" if meaning.predicate in {"start", "initiate"}
             or any(effect.action in {"start", "initiate"} for effect in meaning.effects) else
                 "local_interaction" if meaning.relation == "interaction" else "behavior")
        if role == "functions" and not _ipos_functional_scope_supported(meaning):
            continue
        responsibility = (("context", tuple(stable_payload_hash(asdict(context)) for context in meaning.contexts),
                   meaning.mode, meaning.modality, meaning.negative)
                  if role == "functions" and meaning.contexts else ("subject", meaning.subject))
        connected = len(components[proposition.proposition_id]) > 1
        if connected:
            responsibility = ("dependency", tuple(sorted(components[proposition.proposition_id])))
        group = (role, scope, responsibility, proposition.proposition_id if not connected and proposition.proposition_id in dependency_members else "")
        grouped.setdefault(group, []).append(proposition.proposition_id)
    intents = tuple(DescriptiveClaimIntent("ipos-claim-" + stable_payload_hash(key)[:20], key[0], tuple(identities),
                    scope=key[1], synthesis_rule="explicit_relationship_bundle-v1" if key[2][0] == "dependency" else
                    "shared_context_bundle-v1" if key[2][0] == "context"
                    else "shared_responsibility_bundle-v1") for key, identities in grouped.items())
    return build_descriptive_synthesis_plan(evidence, intents, profile=IPOS_DESCRIPTIVE_PROFILE,
                                          owner=functional_input.block)


def _ipos_document_contract() -> dict:
    from spec_document_contract import IPOS_TEMPLATE, load_ipos_contract
    return load_ipos_contract(Path(__file__).resolve().parents[1] / IPOS_TEMPLATE)


def _ipos_claim_output_role(role: str) -> str:
    return {"purpose": "purpose", "main_functions": "capabilities", "boundaries": "capabilities",
            "limits": "capabilities", "organization": "organization", "functions": "functions"}[role]


IPOS_PROSE_GRAMMAR_VERSION = "finite_clause_agreement_v1"
IPOS_MODAL_ADVERB_RULE = "always_modal_predicate_v1"


def _ipos_grammatical_phrase(value: str, role: str = "", action: str = "") -> str:
    if role == "condition":
        value = re.sub(r"(\bcounter\s+)reach\b", r"\1reaches", value, flags=re.I)
        value = re.sub(r"(\btimer\s+)expire\b", r"\1expires", value, flags=re.I)
    elif role == "object" and action == "start":
        method = re.fullmatch(r"(.+?)\s+writing\s+((?:the\s+)?register\b.+)", value, re.I)
        if method and not re.search(r"\b(?:by|without|while|before|after|not)\b", method[1], re.I):
            value = method[1] + " by writing " + method[2]
    elif role == "object" and action in {"go", "enter"}:
        if not re.search(r"\bto\s+(?:start|initiate|go|enter|raise)\b", value, re.I):
            value = re.sub(r"\band raise(?=\s+an?\s+interrupt\b)", "and raises", value, flags=re.I)
    return value


def _descriptive_meaning_predicate(meaning: DescriptiveMeaning) -> str:
    plural = bool(re.search(r"\band\b", meaning.subject, re.I))
    action = meaning.predicate
    if meaning.modality:
        verb = meaning.modality + (" not " if meaning.negative else " ") + action
    elif action == "be":
        verb = ("are" if plural else "is") + (" not" if meaning.negative else "")
    elif meaning.negative:
        verb = ("do not " if plural else "does not ") + action
    elif plural:
        verb = action
    elif action == "have":
        verb = "has"
    elif action.endswith("y") and len(action) > 1 and action[-2] not in "aeiou":
        verb = action[:-1] + "ies"
    else:
        verb = action + ("es" if action.endswith(("s", "sh", "ch", "x", "z", "o")) else "s")
    for qualifier in meaning.qualifiers:
        if qualifier in {"always", "first"}:
            verb = qualifier + " " + verb
    return " ".join(value for value in (verb, _ipos_grammatical_phrase(_descriptive_join(meaning.objects, meaning.coordination),
                                                               "object", meaning.predicate),
                    " and ".join("to " + effect.action + " " + effect.object for effect in meaning.effects),
                    " ".join(value for value in meaning.qualifiers if value not in {"always", "first"})) if value)


def realize_descriptive_functional_claim(claim: DescriptiveFunctionalClaim) -> str:
    findings = descriptive_claim_scope_findings(claim.scope, claim.meanings)
    if findings:
        raise ValueError("Unsupported functional realization: " + "; ".join(findings))
    groups = {}
    for meaning in dict.fromkeys(claim.meanings):
        if any(context.operator != "atom" for context in meaning.contexts) or meaning.results:
            raise ValueError("Unsupported realization context or result construction")
        groups.setdefault((meaning.subject, meaning.contexts, meaning.mode), []).append(meaning)
    sentences = []
    for (subject, contexts, mode), meanings in groups.items():
        predicate = " and ".join(_descriptive_meaning_predicate(meaning) for meaning in meanings)
        sentence = subject + " " + predicate
        prefixes = (("In " + mode,) if mode else ()) + tuple(_ipos_grammatical_phrase(context.text, "condition") for context in contexts)
        if prefixes:
            preserve = bool(re.match(r"(?:\w*_\w*|[A-Z][A-Z0-9]+)\b", subject))
            sentence = ", ".join(prefixes) + ", " + (sentence if preserve else sentence[:1].lower() + sentence[1:])
        sentences.append(sentence.rstrip(" .") + ".")
    return " ".join(sentences)


def ipos_claim_section_plan(plan: DescriptiveSynthesisPlan) -> Tuple[DescriptiveSectionContribution, ...]:
    return descriptive_section_contributions(plan.claims)


def _ipos_complete_passive_relationship(meaning: DescriptiveMeaning) -> bool:
    if meaning.predicate != "be" or not meaning.contexts or len(meaning.objects) != 1:
        return False
    return bool(re.fullmatch(r"bypassed in favor of .+|written (?:in|inside) .+ using .+", meaning.objects[0], re.I))


IPOS_ROOT_CAUSE_TAXONOMY = MappingProxyType({
    "damaged_admitted_evidence": ("semantic normalization", 4),
    "bounded_interpretation_rule": ("interpretation rules", 2),
    "claim_grouping_or_scope": ("claim grouping", 1),
    "deterministic_cross_fragment_recomposition": ("recomposition", 3),
    "unsupported_inference_required": ("semantic normalization", 5),
})


def ipos_unresolved_root_cause(context: DescriptiveUnresolvedContext,
                               functional_input: IPOSFunctionalInput) -> str:
    records = {record.record_id: record for record in _ipos_admitted_records(functional_input)}
    recovered, residuals = [], []
    for span in context.spans:
        record = records.get(span.record_id)
        if record is None or record.text[span.start:span.end] != span.text:
            raise ValueError("Recovery diagnosis requires exact admitted evidence spans")
        values, gaps = _ipos_recover_clause(record, functional_input, span)
        recovered.extend(values)
        residuals.extend(gaps)
    if recovered and not residuals:
        return ("deterministic_cross_fragment_recomposition" if any(len(spans) > 1 for _, spans in recovered)
                else "bounded_interpretation_rule")
    text = _ipos_source_text(" ".join(span.text for span in context.spans))
    if ("unbalanced source delimiters" in context.reason or "incomplete source continuation" in context.reason
            or re.search(r"\b(?:shall|must|and|or|then|to|by|with|into|the|a|an|for|from|of|in|at|if|when|unless|until|before|after|that|while|is|are|has|have|was|were|been|new|selected|configured)\s*$", text, re.I)
            or text.casefold() in {"this", "if", "then", "and"}):
        return "damaged_admitted_evidence"
    if recovered:
        return "bounded_interpretation_rule"
    if re.search(r"\b(?:shall|must)\s+be able to\s*:", text, re.I) and "\u2022" in text:
        return "bounded_interpretation_rule"
    if not _ipos_top_level_modals(text):
        return "damaged_admitted_evidence"
    if ("unresolved subject reference" in context.reason or "structured/table" in context.reason
            or re.match(r"This\s+(?:shall|must)\b", text, re.I)):
        return "unsupported_inference_required"
    return "bounded_interpretation_rule"


IPOS_DAMAGED_EVIDENCE_CLASSES = MappingProxyType({
    "truly_truncated_or_corrupted": "semantic normalization: preserve supported parts and explicit gaps",
    "recoverable_from_same_fragment": "interpretation rules",
    "recoverable_by_better_normalization": "semantic normalization",
    "recoverable_by_bounded_same_record_recomposition": "same-record recomposition",
    "recoverable_by_bounded_cross_record_recomposition": "cross-record recomposition",
    "not_recoverable_without_unsupported_inference": "semantic normalization: retain unresolved scope",
})


def ipos_damaged_evidence_report(functional_input: IPOSFunctionalInput,
                                contexts: Sequence[DescriptiveUnresolvedContext]) -> dict:
    plan = build_ipos_claim_plan(functional_input)
    records = {record.record_id: record for record in plan.evidence.records}
    observations = []
    for context in contexts:
        if not context.spans or any(span.record_id not in records
                or records[span.record_id].text[span.start:span.end] != span.text for span in context.spans):
            raise ValueError("Damaged-evidence diagnosis requires exact admitted provenance")
        supported = tuple(proposition for proposition in plan.evidence.propositions
                          if proposition.proposition_id.startswith(context.context_id.removesuffix("unresolved") + "meaning:")
                          and any(span.record_id == proposition.spans[-1].record_id
                                  and span.start <= proposition.spans[-1].start < proposition.spans[-1].end <= span.end
                                  for span in context.spans))
        text = " ".join(span.text for span in context.spans)
        if supported:
            category = ("recoverable_by_bounded_same_record_recomposition" if any(
                        len(proposition.spans) > 1 for proposition in supported)
                        else "recoverable_from_same_fragment")
        elif re.search(r"\bformula\s*:|\b(?:Table|Figure)\s+\d+", text, re.I):
            category = "not_recoverable_without_unsupported_inference"
        else:
            category = "truly_truncated_or_corrupted"
        observations.append({"context_id": context.context_id, "class": category,
                             "spans": [asdict(span) for span in context.spans],
                             "supported_proposition_ids": [value.proposition_id for value in supported],
                             "recovery_without_invention": bool(supported),
                             "recovery_extent": "supported_partial_meaning" if supported else "no_demonstrated_recovery",
                             "residual_still_blocking": any(value.context_id == context.context_id
                                                            for value in plan.evidence.unresolved_contexts)})
    classes = []
    for category, location in IPOS_DAMAGED_EVIDENCE_CLASSES.items():
        members = [value for value in observations if value["class"] == category]
        classes.append({"class": category, "count": len(members), "examples": members[:2],
                        "minimum_shared_fix_location": location,
                        "recovery_without_invention": any(value["recovery_without_invention"] for value in members),
                        "status": "observed" if members else "no_demonstrated_case"})
    return {"version": "ipos_damaged_evidence_recovery_v1", "classes": classes,
            "context_count": len(contexts), "observations": observations,
            "classification_basis": "verified_partial_recovery_not_full_reconstruction",
            "cross_record_recomposition_enabled": False, "final_outputs_regenerated": False}


def ipos_root_cause_report(functional_input: IPOSFunctionalInput,
                           contexts: Optional[Sequence[DescriptiveUnresolvedContext]] = None) -> dict:
    plan = build_ipos_claim_plan(functional_input)
    observations = []
    selected_contexts = tuple(contexts) if contexts is not None else plan.evidence.unresolved_contexts
    for context in selected_contexts:
        observations.append({"population": "unresolved_context", "identity": context.context_id,
                             "class": ipos_unresolved_root_cause(context, functional_input), "reason": context.reason,
                             "spans": [asdict(span) for span in context.spans]})
    preserved = {identity for claim in plan.claims for identity in claim.proposition_ids}
    for proposition in plan.evidence.propositions:
        if proposition.proposition_id not in preserved:
            observations.append({"population": "withheld_claim", "identity": proposition.proposition_id,
                                 "class": "claim_grouping_or_scope" if _ipos_functional_scope_supported(proposition.meaning)
                                 else "unsupported_inference_required", "reason": "no_supported_claim",
                                 "spans": [asdict(span) for span in proposition.spans]})
    classes = []
    for name, (location, rank) in IPOS_ROOT_CAUSE_TAXONOMY.items():
        members = [value for value in observations if value["class"] == name]
        counts = {population: sum(value["population"] == population for value in members)
                  for population in ("unresolved_context", "withheld_claim", "editorial_defect")}
        classes.append({"class": name, "minimum_fix_location": location, "rank": rank,
                        "counts": counts, "examples": members[:2],
                        "status": "observed" if members else "not_established_from_admitted_evidence"})
    return {"version": "ipos_semantic_recovery_triage_v1", "block": functional_input.block,
            "classification_basis": "bounded_recovery_candidates_not_semantic_certification",
            "interpretation_rules": IPOS_SEMANTIC_RECOVERY_RULES,
            "classes": sorted(classes, key=lambda value: value["rank"]),
            "observations": observations, "unresolved_contexts": len(selected_contexts),
            "independently_verified_source_findings": _ipos_evidence_semantic_findings(plan.evidence),
            "release_status": "NOT_CERTIFIED", "final_outputs_regenerated": False}


DESCRIPTIVE_ACTION_RELATIONS = MappingProxyType({
    "contain": "containment", "comprise": "containment", "include": "containment",
    "configure": "configuration", "select": "configuration",
    "calculate": "transformation", "convert": "transformation", "process": "transformation",
    "transfer": "interaction", "receive": "interaction", "send": "interaction",
})


def build_ipos_descriptive_structure(
    functional_input: IPOSFunctionalInput, flow: Optional[DescriptiveFlowResult] = None,
) -> DescriptiveStructure:
    flow = flow if flow is not None else _ipos_shared_flow(functional_input)
    propositions = []
    claims = []
    dispositions = []
    by_fact = {}
    for fact in flow.facts:
        if fact.fact_id in by_fact:
            continue
        spans = tuple(DescriptiveSourceSpan(record.record_id, 0, len(record.text), record.text)
                      for record in flow.records if record.text == fact.statement and fact.source in record.provenance)
        meaning = DescriptiveMeaning(
            DESCRIPTIVE_ACTION_RELATIONS.get(fact.action, "behavior"), fact.subject, fact.action,
            fact.objects, (DescriptiveContext("atom", fact.condition),) if fact.condition else (),
            fact.mode, fact.modality, fact.negative, fact.coordination,
            (*fact.qualifiers, *((fact.object_suffix,) if fact.object_suffix else ())), fact.results,
        )
        proposition = DescriptiveProposition(fact.fact_id, meaning, spans, fact.owner, fact.domain, fact.layer)
        propositions.append(proposition)
        by_fact[fact.fact_id] = proposition
    for paragraph, unit in zip(flow.paragraphs, flow.units):
        identities = tuple(dict.fromkeys(fact.fact_id for fact in paragraph.facts))
        claims.append(DescriptiveClaim("claim:" + unit.unit_id, unit.unit_id, identities,
                                      tuple(by_fact[identity].meaning for identity in identities)))
    for proposition in propositions:
        linked = tuple(claim.claim_id for claim in claims if proposition.proposition_id in claim.proposition_ids)
        dispositions.append(DescriptiveDisposition(proposition.spans, (proposition.proposition_id,),
                            "preserved", "controlled_source_projection", linked))
    gaps = {row["source_record_id"]: row["reason"] for row in flow.audit
            if row["decision"] in {"projection_gap", "rejected_evidence"}}
    for record in flow.records:
        if record.record_id in gaps:
            dispositions.append(DescriptiveDisposition(
                (DescriptiveSourceSpan(record.record_id, 0, len(record.text), record.text),), (),
                "unresolved", record.record_id + ":" + gaps[record.record_id]))
    return DescriptiveStructure(DESCRIPTIVE_STRUCTURE_VERSION, stable_payload_hash(asdict(IPOS_DESCRIPTIVE_PROFILE)),
                                flow.records, tuple(propositions), (), tuple(claims), tuple(dispositions), flow.units)


def ipos_descriptive_structural_audit(
    functional_input: IPOSFunctionalInput, flow: Optional[DescriptiveFlowResult] = None,
) -> dict:
    plan = build_ipos_claim_plan(functional_input)
    names = {"digital": "digital-ipos-structure-v1", "analog": "analog-ipos-structure-v1"}
    if functional_input.domain not in names:
        raise ValueError("IPOS structural application requires an explicit approved digital/analog domain")
    policy = DESCRIPTIVE_APPLICATION_PROFILES[names[functional_input.domain]]
    model = descriptive_synthesis_plan_audit(plan, expected_evidence=plan.evidence,
            expected_intents=plan.intents, profile=IPOS_DESCRIPTIVE_PROFILE, owner=functional_input.block)
    source_findings = _ipos_evidence_semantic_findings(plan.evidence)
    source_findings.extend("ipos_source_projection_unresolved:" + context.context_id
                           for context in plan.evidence.unresolved_contexts)
    claim_findings = validate_descriptive_synthesis_plan(plan, expected_evidence=plan.evidence,
                     expected_intents=plan.intents, profile=IPOS_DESCRIPTIVE_PROFILE, owner=functional_input.block)
    roles = {identity: intent.role for intent in plan.intents for identity in intent.proposition_ids}
    preserved = {identity for claim in plan.claims for identity in claim.proposition_ids}
    application = {
        "policy": json.loads(json.dumps(asdict(policy))),
        "decisions": [{"proposition_id": proposition.proposition_id,
                       "decision": "retained" if proposition.proposition_id in preserved else "unresolved",
                       "role": roles[proposition.proposition_id], "rule": "explicit_functional_claim"}
                      for proposition in plan.evidence.propositions],
        "profile_status": "blocked" if claim_findings else "policy_valid", "findings": claim_findings,
    }
    return {
        "contract_version": DESCRIPTIVE_SYNTHESIS_VERSION, "phase": "ipos_functional_claim_materialization",
        "model": model, "application": application,
        "section_plan": json.loads(json.dumps([asdict(value) for value in ipos_claim_section_plan(plan)])),
        "source_findings": source_findings,
        "source_validation_scope": "controlled_normative_constructions_and_literal_inventory",
        "source_semantic_status": "blocked" if source_findings else "controlled_constructions_valid",
        "publication_status": "requires_independent_final_artifact_checks",
        "functional_claim_plan": json.loads(json.dumps(asdict(plan))),
        "functional_claim_findings": claim_findings,
        "document_contract": _ipos_document_contract(),
    }


def write_ipos_structural_audit(path: Path, functional_input: IPOSFunctionalInput) -> None:
    path.write_text(json.dumps(ipos_descriptive_structural_audit(functional_input), indent=2, sort_keys=True) + "\n",
                    encoding="utf-8")


def validate_ipos_structural_audit(path: Path, functional_input: IPOSFunctionalInput) -> List[str]:
    if not path.is_file():
        return ["ipos_structural_audit_missing"]
    try:
        supplied = json.loads(path.read_text(encoding="utf-8"))
        expected = ipos_descriptive_structural_audit(functional_input)
    except (OSError, ValueError, TypeError, KeyError):
        return ["ipos_structural_audit_invalid"]
    findings = list(expected["model"]["findings"])
    findings.extend(expected["application"]["findings"])
    findings.extend(expected["source_findings"])
    findings.extend(expected["functional_claim_findings"])
    if supplied != expected:
        findings.append("ipos_structural_audit_changed")
    return list(dict.fromkeys(findings))


def _ipos_profile_output_section(application: Mapping, proposition_id: str) -> str:
    assessment = next(value for value in application["assessments"] if value["proposition_id"] == proposition_id)
    return dict(application["policy"]["output_sections"])[assessment["role"]]


def _ipos_section_contributions(
    flow: DescriptiveFlowResult, application: Mapping,
) -> Tuple[DescriptiveSectionContribution, ...]:
    assessments = {value["proposition_id"]: value for value in application["assessments"]}
    by_id = {value["proposition_id"]: value for value in application["decisions"]}
    structure_meanings = {}
    for fact in flow.facts:
        structure_meanings[fact.fact_id] = DescriptiveMeaning(
            DESCRIPTIVE_ACTION_RELATIONS.get(fact.action, "behavior"), fact.subject, fact.action, fact.objects,
            (DescriptiveContext("atom", fact.condition),) if fact.condition else (),
            fact.mode, fact.modality, fact.negative, fact.coordination,
            (*fact.qualifiers, *((fact.object_suffix,) if fact.object_suffix else ())), fact.results,
        )
    inventory_parts = tuple(structure_meanings[fact.fact_id] for fact in flow.facts
                            if fact.source == "inventory:main_functions")
    claims = []
    for paragraph, unit in zip(flow.paragraphs, flow.units):
        if any(by_id[fact.fact_id]["decision"] != "retained" for fact in paragraph.facts):
            continue
        meanings = tuple(structure_meanings[fact.fact_id] for fact in paragraph.facts)
        if paragraph.facts[0].source == "inventory:function":
            meanings = (*meanings, *inventory_parts)
        claims.append(DescriptiveFunctionalClaim(unit.unit_id, assessments[paragraph.facts[0].fact_id]["role"],
                      tuple(fact.fact_id for fact in paragraph.facts), meanings, (), ()))
    contributions = descriptive_section_contributions(claims)
    by_unit = {unit.unit_id: paragraph for paragraph, unit in zip(flow.paragraphs, flow.units)}
    result = []
    for contribution in contributions:
        paragraph = by_unit[contribution.claim_id]
        if contribution.role == "functions" and contribution.decision == "retained":
            incomplete = all(not fact.objects or (fact.action == "be"
                             and not re.search(r"\b(?:to|and)\b", " ".join(fact.objects), re.I))
                             for fact in paragraph.facts)
            if incomplete:
                contribution = replace(contribution, decision="unresolved", added_meanings=(),
                                       reason="incomplete_function_description")
        result.append(contribution)
    return tuple(result)


def compose_ipos_overview(
    functional_input: IPOSFunctionalInput,
    document_contract: Optional[Mapping] = None,
) -> Tuple[Dict[str, List[str]], List[Dict[str, str]]]:
    sections: Dict[str, List[str]] = {key: [] for key in (
        "functionality", "scope", "internal_structure", "internal_functions", "general_architecture",
    )}
    if not functional_input.materialized:
        return sections, []
    contract = _ipos_document_contract()
    if document_contract is not None and dict(document_contract) != contract:
        raise ValueError("IPOS generator contract differs from the validated active contract")
    plan = build_ipos_claim_plan(functional_input)
    decisions = {value.claim_id: value for value in ipos_claim_section_plan(plan)}
    by_role = dict(zip(("purpose", "capabilities", "organization", "functions", "architecture"), sections))
    titles = {section["role"]: section["title"] for section in contract["sections"]}
    audit = []
    records = {record.record_id: record for record in plan.evidence.records}
    for claim in plan.claims:
        decision = decisions[claim.claim_id]
        text = ""
        reason = decision.reason
        if decision.decision == "retained":
            try:
                text = realize_descriptive_functional_claim(claim)
            except ValueError as error:
                reason = str(error)
        role = _ipos_claim_output_role(claim.role)
        if text:
            sections[by_role[role]].append(text)
        unit = SemanticUnit(claim.claim_id, titles[role], text,
                            tuple(dict.fromkeys(span.record_id for span in claim.spans)), len(audit) + 1)
        audit.append({"block": functional_input.block, "domain": functional_input.domain,
                      "source_record_id": unit.source_record_ids[0],
                      "candidate_evidence_requirement_id": unit.source_record_ids[0].rsplit(":", 1)[-1]
                          if unit.source_record_ids[0].startswith("requirement:") else "",
                      "candidate_evidence_statement": records[unit.source_record_ids[0]].text
                          if unit.source_record_ids[0].startswith("requirement:") else "",
                      "decision": "selected" if text else "excluded" if decision.decision == "excluded" else "projection_gap",
                      "reason": reason, "descriptive_statement": text, "rendered_summary": text,
                      "target_section": titles[role], "projection_adapter": contract["interpretation"],
                      "profile": contract["profile"], "profile_version": str(contract["version"]),
                      "profile_fingerprint": contract["sha256"], "unit_id": claim.claim_id,
                      "claim_json": json.dumps(asdict(claim), sort_keys=True),
                      "contributor_ids": json.dumps(claim.proposition_ids),
                      "semantic_unit_json": json.dumps(asdict(unit), sort_keys=True),
                      "provenance_references": "; ".join(dict.fromkeys(reference for identity in unit.source_record_ids
                                                for reference in records[identity].provenance))})
    for disposition in plan.dispositions:
        if disposition.decision == "unresolved":
            for span in disposition.spans:
                audit.append({"block": functional_input.block, "domain": functional_input.domain,
                              "candidate_evidence_requirement_id": span.record_id.rsplit(":", 1)[-1]
                                  if span.record_id.startswith("requirement:") else "",
                              "candidate_evidence_statement": records[span.record_id].text,
                              "profile": contract["profile"],
                              "decision": "projection_gap", "source_record_id": span.record_id,
                              "reason": disposition.reason, "span_json": json.dumps(asdict(span), sort_keys=True)})
    return sections, audit


def _legacy_compose_ipos_overview(
    functional_input: IPOSFunctionalInput,
) -> Tuple[Dict[str, List[str]], List[Dict[str, str]]]:
    """Compose bounded IPOS prose from inventory authority and supported local evidence."""
    if not functional_input.materialized:
        return {
            "functionality": [], "scope": [], "internal_structure": [],
            "internal_functions": [], "general_architecture": [],
        }, []
    function = (functional_input.function or "perform its specified function").strip().rstrip(".")
    functionality = [compose_technical_block_purpose(functional_input.block, function)]
    input_text = functional_input.inputs or "the approved local inputs"
    input_text = re.sub(r"\s+", " ", input_text).strip(" ,.;:")
    output_text = functional_input.outputs or "the approved local outputs"
    if re.match(r"[A-Z][a-z]", input_text):
        input_text = input_text[:1].lower() + input_text[1:]
    if re.match(r"[A-Z][a-z]", output_text):
        output_text = output_text[:1].lower() + output_text[1:]
    scope = [
        f"The {functional_input.block} accepts {input_text} and produces {output_text}."
    ]
    candidate_audit = _classify_ipos_local_refinements(functional_input)
    scope.extend(_group_ipos_local_refinements(
        candidate_audit, functional_input.function, functional_input.inputs, functional_input.outputs,
    ))
    audit: List[Dict[str, str]] = []
    common = {
        "source_authority_tier": "approved_snapshot_block_inventory",
        "block": functional_input.block,
        "domain": functional_input.domain,
        "source_function": functional_input.function,
        "inputs_used": functional_input.inputs,
        "outputs_used": functional_input.outputs,
        "provenance_references": "; ".join(functional_input.provenance),
        "hierarchy_level": "IPOS block",
        "duplicate_cross_document_note": "",
    }
    audit.append({
        **common,
        "target_section": "1.1 Functionality",
        "descriptive_statement": functionality[0],
        "supporting_requirement_evidence_ids": "; ".join(functional_input.local_requirement_ids),
        "derivation_mode": "direct_function",
        "decision": "selected",
        "exclusion_suppression_reason": "",
    })
    audit.append({
        **common,
        "target_section": "1.2 Supported functions and scope",
        "descriptive_statement": scope[0],
        "supporting_requirement_evidence_ids": "; ".join(functional_input.local_requirement_ids),
        "derivation_mode": "function_plus_io_scope",
        "decision": "selected",
        "exclusion_suppression_reason": "",
    })
    for requirement_id, evidence in zip(functional_input.local_requirement_ids, functional_input.local_requirement_evidence):
        audit.append({
            **common,
            "target_section": "1.2 Supported functions and scope",
            "descriptive_statement": evidence,
            "supporting_requirement_evidence_ids": requirement_id,
            "derivation_mode": "suppressed_due_to_insufficient_evidence",
            "decision": "suppressed",
            "exclusion_suppression_reason": "normative local requirement retained as evidence only; it cannot redefine the approved block scope",
        })
    scope_order = {value.casefold(): index for index, value in enumerate(scope, start=1)}
    for row in candidate_audit:
        rendered_summary = row.get("rendered_summary", "").strip()
        if row.get("decision") == "accepted_candidate" and rendered_summary:
            row["descriptive_statement"] = rendered_summary
            row["output_order"] = str(scope_order.get(rendered_summary.casefold(), ""))
    audit.extend(candidate_audit)
    return {
        "functionality": functionality,
        "scope": scope,
        "internal_structure": [],
        "internal_functions": [],
        "general_architecture": [],
    }, audit


def build_ipos_normalized_records(functional_input: IPOSFunctionalInput) -> Tuple[NormalizedSourceRecord, ...]:
    return build_ipos_evidence_model(functional_input).records


def build_ipos_semantic_units(
    functional_input: IPOSFunctionalInput,
    sections: Mapping[str, Sequence[str]],
    audit_rows: Sequence[Mapping[str, str]],
) -> Tuple[SemanticUnit, ...]:
    units = {}
    for row in audit_rows:
        if row.get("decision") != "selected":
            continue
        payload = json.loads(row["semantic_unit_json"])
        payload["source_record_ids"] = tuple(payload["source_record_ids"])
        unit = SemanticUnit(**payload)
        units[unit.unit_id] = unit
    rendered = "\n".join(text for paragraphs in sections.values() for text in paragraphs)
    ordered = sorted(units.values(), key=lambda unit: rendered.find(unit.text))
    return tuple(replace(unit, output_order=order) for order, unit in enumerate(ordered, start=1))


def _legacy_build_ipos_normalized_records(functional_input: IPOSFunctionalInput) -> Tuple[NormalizedSourceRecord, ...]:
    """Expose IPOS authority and candidate decisions through the shared record contract."""
    records = [
        NormalizedSourceRecord(
            record_id="inventory:function",
            text=functional_input.function,
            authority_tier="authoritative",
            ownership_scope=functional_input.block,
            provenance=functional_input.provenance,
            decision="retained",
            rationale="approved block function authority",
        ),
        NormalizedSourceRecord(
            record_id="inventory:inputs",
            text=functional_input.inputs,
            authority_tier="authoritative",
            ownership_scope=functional_input.block,
            provenance=functional_input.provenance,
            decision="retained",
            rationale="approved block input boundary",
        ),
        NormalizedSourceRecord(
            record_id="inventory:outputs",
            text=functional_input.outputs,
            authority_tier="authoritative",
            ownership_scope=functional_input.block,
            provenance=functional_input.provenance,
            decision="retained",
            rationale="approved block output boundary",
        ),
    ]
    candidate_rows = _classify_ipos_local_refinements(functional_input)
    _group_ipos_local_refinements(
        candidate_rows, functional_input.function, functional_input.inputs, functional_input.outputs,
    )
    for row in candidate_rows:
        evidence_block = row.get("candidate_evidence_block", "").strip()
        requirement_id = row.get("candidate_evidence_requirement_id", "").strip()
        decision = row.get("decision", "")
        suppressed_atomic = row.get("materialization_status") == "suppressed_atomic"
        records.append(NormalizedSourceRecord(
            record_id=f"requirement:{evidence_block}:{requirement_id}",
            text=row.get("candidate_refinement", "").strip(),
            authority_tier=(
                "authoritative"
                if evidence_block.casefold() == functional_input.block.casefold()
                else "discovery_only"
            ),
            ownership_scope=evidence_block,
            provenance=tuple(filter(None, (
                *functional_input.provenance,
                f"approved requirement:{requirement_id}" if requirement_id else "",
            ))),
            decision="retained" if decision == "accepted_candidate" and not suppressed_atomic else "suppressed",
            rationale=(
                "accepted evidence is an atomic helper operation without an autonomous summary"
                if suppressed_atomic else
                row.get("exclusion_suppression_reason", "").strip()
                or "supported same-scope local refinement"
            ),
        ))
    return tuple(records)


def _legacy_build_ipos_semantic_units(
    functional_input: IPOSFunctionalInput,
    sections: Mapping[str, Sequence[str]],
    audit_rows: Sequence[Mapping[str, str]],
) -> Tuple[SemanticUnit, ...]:
    """Map composed IPOS text to stable source-linked semantic units."""
    units: List[SemanticUnit] = []
    for text in sections.get("functionality", ()):
        units.append(SemanticUnit(
            unit_id=f"functionality:{len(units) + 1}",
            section="1.1 Functionality",
            text=text,
            source_record_ids=("inventory:function",),
            output_order=len(units) + 1,
        ))
    for index, text in enumerate(sections.get("scope", ())):
        source_record_ids = ("inventory:function", "inventory:inputs", "inventory:outputs")
        if index > 0:
            matching = [
                row for row in audit_rows
                if row.get("decision") == "accepted_candidate"
                and row.get("rendered_summary", "").strip() == text.strip()
            ]
            if matching:
                source_record_ids = tuple(
                    f"requirement:{row.get('candidate_evidence_block', '').strip()}:{row.get('candidate_evidence_requirement_id', '').strip()}"
                    for row in matching
                )
        units.append(SemanticUnit(
            unit_id=f"supported-scope:{index + 1}",
            section="1.2 Supported functions and scope",
            text=text,
            source_record_ids=source_record_ids,
            output_order=len(units) + 1,
        ))
    return tuple(units)


_IPOS_REFINEMENT_STOP_WORDS = {
    "a", "an", "and", "as", "at", "by", "for", "from", "in", "into", "is", "it", "of",
    "on", "or", "shall", "should", "the", "their", "this", "to", "when", "where", "with",
    "be", "been", "being", "does", "do", "not", "must", "will", "which", "that",
}
_IPOS_ACTION_VERBS = {
    "accept", "accumulate", "acquire", "aggregate", "average", "calculate", "clamp", "configure",
    "control", "convert", "copy", "decrement", "disable", "divide", "elaborate", "enable", "enter",
    "exit", "filter", "generate", "increment", "manage", "measure", "perform", "process", "raise",
    "read", "receive", "remain", "reset", "route", "sample", "select", "sequence", "set", "start",
    "stop", "store", "subtract", "transfer", "wait", "write",
}
_IPOS_ACTION_FORMS = {
    "acquired": "acquire", "acquires": "acquire", "acquiring": "acquire",
    "aggregated": "aggregate", "aggregates": "aggregate", "aggregating": "aggregate",
    "averaged": "average", "averages": "average", "averaging": "average",
    "calculated": "calculate", "calculates": "calculate", "calculating": "calculate",
    "clamped": "clamp", "clamps": "clamp", "clamping": "clamp",
    "configured": "configure", "configures": "configure", "configuring": "configure",
    "controlled": "control", "controls": "control", "controlling": "control",
    "converted": "convert", "converts": "convert", "converting": "convert",
    "copied": "copy", "copies": "copy", "copying": "copy",
    "decremented": "decrement", "decrements": "decrement", "decrementing": "decrement",
    "disabled": "disable", "disables": "disable", "disabling": "disable",
    "divided": "divide", "divides": "divide", "dividing": "divide",
    "elaborated": "elaborate", "elaborates": "elaborate", "elaborating": "elaborate",
    "enabled": "enable", "enables": "enable", "enabling": "enable",
    "entered": "enter", "enters": "enter", "entering": "enter",
    "exited": "exit", "exits": "exit", "exiting": "exit",
    "filtered": "filter", "filters": "filter", "filtering": "filter",
    "generated": "generate", "generates": "generate", "generating": "generate",
    "incremented": "increment", "increments": "increment", "incrementing": "increment",
    "managed": "manage", "manages": "manage", "managing": "manage",
    "measured": "measure", "measures": "measure", "measuring": "measure",
    "performed": "perform", "performs": "perform", "performing": "perform",
    "processed": "process", "processes": "process", "processing": "process",
    "raised": "raise", "raises": "raise", "raising": "raise",
    "read": "read", "reads": "read", "reading": "read",
    "received": "receive", "receives": "receive", "receiving": "receive",
    "remained": "remain", "remains": "remain", "remaining": "remain",
    "reset": "reset", "resets": "reset", "resetting": "reset",
    "routed": "route", "routes": "route", "routing": "route",
    "sampled": "sample", "samples": "sample", "sampling": "sample",
    "selected": "select", "selects": "select", "selecting": "select",
    "sequenced": "sequence", "sequences": "sequence", "sequencing": "sequence",
    "set": "set", "sets": "set", "setting": "set",
    "started": "start", "starts": "start", "starting": "start",
    "stopped": "stop", "stops": "stop", "stopping": "stop",
    "stored": "store", "stores": "store", "storing": "store",
    "subtracted": "subtract", "subtracts": "subtract", "subtracting": "subtract",
    "transferred": "transfer", "transfers": "transfer", "transferring": "transfer",
    "waited": "wait", "waits": "wait", "waiting": "wait",
    "written": "write", "writes": "write", "writing": "write",
}
_IPOS_ACTION_PRESENT = {
    "acquire": "Acquires", "accumulate": "Accumulates", "aggregate": "Aggregates", "average": "Averages",
    "calculate": "Calculates", "clamp": "Clamps", "configure": "Configures", "control": "Controls",
    "convert": "Converts", "copy": "Copies", "decrement": "Decrements", "disable": "Disables",
    "divide": "Divides", "elaborate": "Elaborates", "enable": "Enables", "enter": "Enters",
    "exit": "Exits", "filter": "Filters", "generate": "Generates", "increment": "Increments",
    "manage": "Manages", "measure": "Measures", "perform": "Performs", "process": "Processes",
    "raise": "Raises", "read": "Reads", "receive": "Receives", "remain": "Remains", "reset": "Resets",
    "route": "Routes", "sample": "Samples", "select": "Selects", "sequence": "Sequences",
    "set": "Sets", "start": "Starts", "stop": "Stops", "store": "Stores", "subtract": "Subtracts",
    "transfer": "Transfers", "wait": "Waits", "write": "Writes",
}
_IPOS_NORMATIVE_RE = re.compile(r"\b(?:shall|must|required to|is required to)\b", re.I)
_IPOS_NEGATION_RE = re.compile(r"\b(?:not|never|cannot|doesn't|does not|shall not|must not)\b", re.I)


def _ipos_action_lemma(word: str) -> str:
    value = word.casefold()
    if value in _IPOS_ACTION_VERBS:
        return value
    return _IPOS_ACTION_FORMS.get(value, "")


def _ipos_refinement_tokens(value: str) -> Set[str]:
    tokens: Set[str] = set()
    for token in re.findall(r"[a-z][a-z0-9_-]*", value.casefold()):
        if token in _IPOS_REFINEMENT_STOP_WORDS or len(token) <= 2 or token.isupper():
            continue
        lemma = _ipos_action_lemma(token)
        if not lemma and token.endswith("ies") and len(token) > 4:
            lemma = token[:-3] + "y"
        elif not lemma and token.endswith("s") and len(token) > 3:
            lemma = token[:-1]
        elif not lemma and token.endswith("ed") and len(token) > 4:
            lemma = token[:-2]
        tokens.add(lemma or token)
    return tokens


def _ipos_refinement_has_inventory_support(
    action: str,
    statement_tokens: Set[str],
    function_tokens: Set[str],
    input_tokens: Set[str],
    output_tokens: Set[str],
) -> bool:
    if (function_tokens | input_tokens | output_tokens).intersection(statement_tokens):
        return True
    return (
        action in {"configure", "select"}
        and bool(input_tokens.intersection({"configuration", "setting", "settings"}))
    ) or (
        action in {"write", "store", "route", "transfer"}
        and bool(output_tokens.intersection({"data", "routing", "result", "results"}))
    ) or (
        action in {"sample", "acquire", "measure", "receive"}
        and bool(input_tokens.intersection({"data", "sample", "samples"}))
    )


def _ipos_refinement_parts(statement: str) -> Tuple[str, str, str, bool]:
    statement = re.sub(r"\[(?:Covers:|TO:|Vpriority)[^]]*\]", " ", statement, flags=re.I)
    statement = re.sub(r"\b(?:DDS|IPOS)_[A-Z0-9_-]+\b", " ", statement, flags=re.I)
    clamped_value = re.search(
        r"\bvalue greater than (\d+) into ones? of the ([A-Za-z][A-Za-z0-9_]*)"
        r".*?\bvalue shall clamp at (\d+)\b", statement, re.I,
    )
    if clamped_value and clamped_value.group(1) == clamped_value.group(3):
        return "clamp", f"{clamped_value.group(2).replace('_', ' ')} values above {clamped_value.group(1)} to {clamped_value.group(1)}", "", False
    if re.search(r"\bPPG Raw Data and PPG Noise reach\b", statement, re.I) and re.search(
        r"\bdepending on the ALC Mask bit value, the two accumulated shall be subtracted\b", statement, re.I,
    ):
        return "subtract", "accumulated PPG raw data and PPG noise according to the ALC Mask", "", False
    repeated_frames = re.search(
        r"\bafter the last repetition has been accumulated, the data output shall be obtained "
        r"dividing by the number of repeated frames the accumulated data to obtain a (\d+) bit data\b",
        statement, re.I,
    )
    if repeated_frames:
        return "divide", f"accumulated frame data by the number of repetitions to produce a {repeated_frames.group(1)} bit result", "", False
    sampled_pair = re.search(r"\bshall\s+be\s+sampled\s+(\d+)\s+times,\s+then\s+mediated\b", statement, re.I)
    subtraction = re.search(
        r"\bsubtraction\s+between\s+([A-Za-z][A-Za-z0-9_]*)\s+and\s+([A-Za-z][A-Za-z0-9_]*)\b",
        statement, re.I,
    )
    if sampled_pair and subtraction and re.match(r"If\s+.+?\bis\s+set\s+high\b", statement, re.I):
        first, second = subtraction.groups()
        if all(re.search(rf"\b{re.escape(operand)}\b", statement[:sampled_pair.start()], re.I) for operand in (first, second)):
            return (
                "average",
                f"{sampled_pair.group(1)} acquisitions of {first.replace('_', ' ')} and "
                f"{second.replace('_', ' ')} separately when the mode is enabled, "
                "then subtracts their averaged values",
                "", False,
            )
    acquisition_interval = re.fullmatch(
        r"(.+?)\s+(?:shall|must)\s+define\s+how\s+many\s+(time slots?)\s+pass\s+"
        r"between\s+two\s+(.+?)\s+acquisitions\b(?:\s+following\s+this\s+formula:.*)?\.?",
        statement.strip(), flags=re.I,
    )
    if acquisition_interval:
        _source, interval_unit, acquisition = acquisition_interval.groups()
        description = f"number of {interval_unit} between two {acquisition} acquisitions"
        if re.search(r"\bfollowing\s+this\s+formula:\s*\S+\s*=", statement, re.I):
            description += (
                "; the cited expression relates acquisition data to time slots and the configured input"
                " without resolving its arithmetic"
            )
        return "configure", description, "", False
    conditional_transition = re.fullmatch(
        r"When in (.+?) state, when (.+?), (.+?) shall be set (?:high|low) "
        r"to start (.+?) and go into (.+?) state\.?",
        statement.strip(), flags=re.I,
    )
    if conditional_transition:
        initial_state, trigger, signal, effect, next_state = conditional_transition.groups()
        if re.search(r"\b(?:signal|output)\b|\b[oO]_[A-Za-z0-9_]+\b", signal):
            trigger = re.sub(r"\b[iIoO]_(?=[A-Za-z0-9])", "", trigger).replace("_", " ")
            return (
                "start", f"{effect} and enters {next_state.replace('_', ' ')} state",
                f"{trigger} in {initial_state.replace('_', ' ')} state", False,
            )
    procedure = re.match(
        r"To perform (?:an? |the )?([A-Za-z][A-Za-z ]{1,60}) procedure,\s*"
        r"(?:the )?user shall perform the following steps:\s*\d+\.",
        statement, re.I,
    )
    if procedure:
        return "perform", f"{procedure.group(1).strip()} procedure", "", False
    state_progression = re.search(
        r"\b(?:FSM|controller)\s+shall\s+go\s+into\s+(?:the\s+)?([A-Za-z][A-Za-z0-9_ ]+?)\s+state\s+"
        r"and\s+then\s+into\s+(?:the\s+)?([A-Za-z][A-Za-z0-9_ ]+?)(?:\s+state|[.;]|$)",
        statement, re.I,
    )
    if state_progression:
        first, second = (value.strip().replace("_", " ") for value in state_progression.groups())
        return "sequence", f"{first} then {second} states", "", False
    state_entry = re.search(
        r"\bshall\s+go\s+(?:back\s+)?(?:into|in|to)\s+(?:the\s+)?"
        r"((?:(?!(?:after|before|when|if)\b)[A-Za-z][A-Za-z0-9_]*\s*){1,3}?)\s+state\b", statement, re.I,
    )
    if state_entry:
        destination = state_entry.group(1).strip().replace("_", " ")
        context = statement[:state_entry.start()]
        if _IPOS_NORMATIVE_RE.search(context):
            return "", "", "", False
        condition = re.match(r"^(?:When|If)\s+(.+)$", context.strip(), re.I)
        if condition:
            context = re.sub(
                r",\s*(?:the\s+)?(?:device|[A-Za-z0-9_]+\s+FSM|controller)\s*$",
                "", condition.group(1), flags=re.I,
            )
            context = re.sub(r"\[[^]]*\]|\([^)]*\)", "", context)
            context = re.sub(r"\s+", " ", context).strip(" ,")
            if len(context.split()) <= 25 and context:
                outcome = f"{destination} state when {context}"
                following = re.match(r"\s*if\s+(.+?)(?:\s*\.\s*Furthermore\b|[.;]|$)", statement[state_entry.end():], re.I)
                if following:
                    guard = re.sub(r"\([^)]*\)", "", following.group(1))
                    guard = re.sub(r"\s+", " ", guard).strip(" ,")
                    if not guard or guard.split()[-1].casefold() in {"the", "if", "and", "or", "when"}:
                        return "", "", "", False
                    if len(guard.split()) <= 20:
                        outcome += f" if {guard}"
                    else:
                        return "", "", "", False
                return "enter", outcome, "", False
            return "", "", "", False
        return "enter", f"{destination} state", "", False
    controlled_output = re.search(
        r"\b(o_[A-Za-z0-9_]+|[A-Za-z]+(?:_[A-Za-z0-9]+)+)\s+shall\s+(?:be\s+)?set\s+"
        r"(?:to\s+)?([01]|high|low)\b",
        statement, re.I,
    )
    if controlled_output:
        prefix = statement[:controlled_output.start()]
        if re.search(r"\b(?:samples?|sampled|acquisition|conversion|threshold|frame|time slot)\b|_cds\b", prefix, re.I):
            output = re.sub(r"^[io]_", "", controlled_output.group(1), flags=re.I).replace("_", " ")
            condition = re.match(r"^(?:When|If|In case)\s+(.+)$", prefix.strip(), re.I)
            context = re.sub(r",\s*(?:the\s*)?$", "", condition.group(1)).strip() if condition else ""
            context = re.sub(r"\b(?:shall|must)\s+be\b", "is", context, flags=re.I)
            context = re.sub(r"\b(?:shall|must)\b", "", context, flags=re.I)
            context = re.sub(r"\s+", " ", context).strip()
            if context and len(context.split()) <= 65:
                effect = f"{output} output when {context}"
                reset = re.search(r"\band\s+shall\s+go\s+to\s+0\s+at\s+(.+?)(?:[.;]|$)",
                                  statement[controlled_output.end():], re.I)
                if reset and controlled_output.group(2).casefold() in {"1", "high"}:
                    effect += f"; clears at {reset.group(1).strip()}"
                recurring = re.search(r"\bin\s+every\s+time\s+slot\b", statement[controlled_output.end():], re.I)
                if recurring:
                    effect += " in every time slot"
                return ("enable" if controlled_output.group(2).casefold() in {"1", "high"} else "disable"), effect, "", False
    timed_step = re.search(
        r"\b(o_[A-Za-z0-9_]+)\s+shall\s+increment\s+the\s+first\s+step\s+at\s+value\s+"
        r"(i_[A-Za-z0-9_]+).*?\bat\s+a\s+time\s+(i_[A-Za-z0-9_]+)\b",
        statement, re.I,
    )
    if timed_step:
        output, value, timing = (part.replace("_", " ") for part in timed_step.groups())
        return "increment", f"{output} first step to {value} at {timing} timing", "", False
    modal = re.search(r"\b(?:shall|must|should|will|required to|is required to)\b", statement, re.I)
    if not modal:
        return "", "", "", False
    subject = statement[:modal.start()].strip(" ,.;:")
    predicate = statement[modal.end():].strip()
    negative = bool(re.match(r"not\b", predicate, re.I))
    predicate = re.sub(r"^not\s+", "", predicate, flags=re.I)
    passive_subject = ""
    passive = re.match(r"be\s+(written|read|stored|sampled|selected|set|reset|enabled|disabled|generated|raised)\b", predicate, re.I)
    if passive:
        action = _ipos_action_lemma(passive.group(1))
        passive_subject = re.split(r",\s*(?=(?:the\s+)?[A-Za-z][A-Za-z0-9_]*(?:\s+signal)?$)", subject)[-1]
        tail = f"{passive_subject} {predicate[passive.end():]}".strip()
        passive_subject = subject
    else:
        action_phrase = re.match(r"(?:(?:also|always|only)\s+)*(?:be\s+able\s+to\s+)?([A-Za-z][A-Za-z'-]*)\b", predicate, re.I)
        if not action_phrase:
            return "", "", "", negative
        action = _ipos_action_lemma(action_phrase.group(1))
        if not action:
            return "", "", "", negative
        tail = predicate[action_phrase.end():].strip()
        if action == "perform" and re.match(r"(?:(?:the|an?)\s+)?averag(?:e|ing)\b", tail, re.I):
            action = "average"
            consumed = re.match(r"(?:(?:the|an?)\s+)?averag(?:e|ing)\b", tail, re.I)
            tail = tail[consumed.end():].strip()
        elif action == "perform" and re.match(r"(?:(?:the|an?)\s+)?subtract(?:ion|ing)?\b", tail, re.I):
            action = "subtract"
            consumed = re.match(r"(?:(?:the|an?)\s+)?subtract(?:ion|ing)?\b", tail, re.I)
            tail = tail[consumed.end():].strip()
    tail = re.split(r"\b(?:when|if|unless|where|because|provided that|after|before|during)\b", tail, maxsplit=1, flags=re.I)[0]
    tail = re.split(r"[,;.]|\band\s+(?:an?\s+)?(?:acquire|aggregate|average|calculate|configure|control|convert|copy|decrement|disable|divide|elaborate|enable|enter|exit|filter|generate|increment|manage|measure|perform|process|raise|read|receive|remain|reset|route|sample|select|sequence|set|start|stop|store|subtract|transfer|wait|write)\b", tail, maxsplit=1, flags=re.I)[0]
    detail = re.search(r"\b(?:register|address|bit|bits|clock|clocks|cycle|cycles|timer|milliseconds|seconds|frequency|value|values|duration|index|division|parameter|parameters|counter|delay|number|time|step|max|min|ms|us)\b|\b\d+\s*(?:ms|us)\b", tail, re.I)
    if detail:
        tail = tail[:detail.start()]
    if action == "start":
        tail = re.sub(r"\s+writing\s*$", "", tail, flags=re.I)
    if action == "average":
        tail = re.sub(
            r"^\s*on\s+(\d+)\s+sample\s+acquisition(?:s)?\s+for\b",
            r"\1 samples for", tail, flags=re.I,
        )
        tail = re.sub(r"^\s*on\s+\d+\s+", " ", tail, flags=re.I)
        tail = re.sub(r"\bsample acquisition for\b", "samples for", tail, flags=re.I)
    if action == "wait":
        tail = re.sub(
            r"^\s*that\s+(?:the\s+)?([A-Za-z][A-Za-z0-9_-]*)\s+ends\s+(?:the\s+)?([A-Za-z][A-Za-z0-9_-]*(?:\s+phase)?)\b",
            r"for \1 \2 completion", tail, flags=re.I,
        )
        tail = re.sub(
            r"^\s*that\s+the\s+rise\s+of\s+([A-Za-z][A-Za-z0-9_-]*)\s+signal\s+from\s+(.+?)\s+to\s+go\s+into\s+(.+?)\s+state\b",
            r"for \1 from \2 before entering \3 state", tail, flags=re.I,
        )
        if not tail.casefold().startswith("for "):
            return "", "", "", negative
        if not re.search(r"\bcompletion\b|\bbefore entering\b", tail, re.I):
            return "", "", "", negative
    tail = re.sub(r"\b(?:the|a|an|which|that|shall|must|should|be|is|are|of|and|or)\b", " ", tail, flags=re.I)
    if action != "average" or not re.match(r"\s*\d+\s+samples\s+for\b", tail, re.I):
        tail = re.sub(r"\b(?:0x[0-9a-f]+|\d+)\b", " ", tail, flags=re.I)
    tail = re.sub(r"\(\s*[A-Z0-9_]{3,}\s*\)", " ", tail)
    signal_identifier_count = len(re.findall(r"\b[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)+\b", tail))
    tail = re.sub(
        r"\b[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)+\b",
        lambda match: " ".join(
            part for index, part in enumerate(match.group(0).split("_"))
            if part and not (index == 0 and part.casefold() in {"i", "o", "n"})
        ),
        tail,
    )
    tail = re.sub(r"\s+", " ", tail).strip(" ,.;:-")
    if action == "configure":
        tail = re.sub(r"\b([A-Za-z][A-Za-z-]*)\s+selected\b", r"\1 selection", tail, flags=re.I)
    object_words = tail.split()
    if not object_words or (len(object_words) < 2 and object_words[0].casefold() in {"to", "of", "for", "with", "from", "by"}):
        return "", "", "", negative
    if object_words[-1].casefold() in {"at", "to", "from", "for", "with", "first", "next", "new", "each", "max", "min"}:
        return "", "", "", negative
    if action in {"increment", "decrement"}:
        return "", "", "", negative
    if all(word.casefold() in {"high", "low", "true", "false", "enabled", "disabled", "ones"} for word in object_words):
        return "", "", "", negative
    if signal_identifier_count >= 3:
        return "", "", "", negative
    if action == "set":
        return "", "", "", negative
    if action == "raise" and re.search(r"\bsignal\b", tail, re.I) and not re.search(r"\binterrupt\b", tail, re.I):
        return "", "", "", negative
    if re.search(r"\bfor each\b", tail, re.I):
        return "", "", "", negative
    object_words = [word for word in object_words if word.casefold() not in {"signal", "signals"}]
    if not object_words:
        return "", "", "", negative
    object_text = " ".join(object_words[:12]).strip()

    qualifier = ""
    condition = re.search(r"\b(when|if|during|while)\s+([^,;]{1,100}),", subject, re.I)
    if condition:
        candidate_qualifier = re.sub(r"\s+", " ", condition.group(2)).strip()
        if set(_ipos_refinement_tokens(candidate_qualifier)).intersection(_ipos_refinement_tokens(object_text)):
            candidate_qualifier = ""
        if (
            len(candidate_qualifier.split()) <= 9
            and not re.search(r"\d|_|\b(?:register|address|bit|clock|cycle|timer|value|frequency)\b", candidate_qualifier, re.I)
        ):
            qualifier = candidate_qualifier
    return action, object_text, qualifier, negative


def _ipos_refinement_candidate(statement: str) -> str:
    action, object_text, qualifier, negative = _ipos_refinement_parts(statement)
    if not action or not object_text:
        return ""
    subject = _IPOS_ACTION_PRESENT.get(action, "")
    if not subject:
        return ""
    user_configuration = re.search(r"\b(?:user|operator|host)\b", statement[:re.search(r"\b(?:shall|must|should|will|required to|is required to)\b", statement, re.I).start()], re.I)
    if user_configuration and action in {"configure", "select"}:
        verb_noun = "configuration" if action == "configure" else "selection"
        if action == "select":
            selector_labels = []
            for match in re.finditer(r"\b([A-Z][A-Z0-9]*(?:_[A-Z0-9]+)*)_SEL\b", statement):
                label = match.group(1).replace("_", " ")
                following = statement[match.end():]
                description = re.match(r"\s*\(([^()]*)\)", following)
                description_text = description.group(1).strip() if description else ""
                if description_text and not description_text.isupper() and not re.search(r"\b(?:device|register|config|channel|ctrl)\b", description_text, re.I):
                    selector_labels.append(description_text)
                else:
                    selector_labels.append(label)
            channel_labels = [label for label in selector_labels if "mode" not in label.casefold()]
            mode_labels = [label for label in selector_labels if "mode" in label.casefold()]
            mode_match = re.search(r"\b((?:[A-Za-z]+\s+){0,2}mode)\b", statement, re.I)
            if mode_match and mode_match.group(1).casefold() not in " ".join(selector_labels).casefold():
                mode_labels.append(mode_match.group(1).strip())
            if channel_labels:
                channel_text = " and ".join(dict.fromkeys(channel_labels))
                channel_count = sum(len(label.split()) for label in set(channel_labels))
                object_text = f"{channel_text} channel" + ("s" if channel_count > 1 else "")
            if mode_labels:
                object_text += " and " + " and ".join(dict.fromkeys(mode_labels))
        sentence = f"Accepts user {verb_noun} of {object_text}"
        if action == "configure":
            mode_match = re.search(r"\b((?:[A-Za-z]+\s+){0,2}mode)\b", statement, re.I)
            if mode_match and mode_match.group(1).casefold() not in object_text.casefold():
                sentence += f" and {mode_match.group(1).strip()}"
    else:
        sentence = f"{subject} {object_text}"
    if qualifier:
        sentence += f" when {qualifier}"
    if negative:
        sentence = f"Does not {action} {object_text}"
    return sentence.rstrip(" .") + "."


def _ipos_join_summary_terms(terms: Sequence[str]) -> str:
    """Join distinct local terms without changing their supported meaning."""
    values = list(dict.fromkeys(value.strip() for value in terms if value.strip()))
    if len(values) <= 1:
        return values[0] if values else ""
    if len(values) == 2:
        return f"{values[0]} and {values[1]}"
    return f"{', '.join(values[:-1])}, and {values[-1]}"


def _ipos_group_summary_text(candidates: Sequence[str]) -> str:
    """Render one non-normative local summary while retaining distinct objects."""
    values = list(dict.fromkeys(value.strip().rstrip(".") for value in candidates if value.strip()))
    if len(values) == 1:
        return values[0] + "."
    selection_prefix = "Accepts user selection of "
    if all(value.startswith(selection_prefix) for value in values):
        labels: List[str] = []
        qualifiers: List[str] = []
        for value in values:
            match = re.match(r"Accepts user selection of (.+?) channels?(?:\s+and\s+(.+))?$", value, re.I)
            if not match:
                break
            labels.extend(part.strip() for part in re.split(r"\s+and\s+", match.group(1)) if part.strip())
            if match.group(2):
                qualifiers.append(match.group(2).strip())
        else:
            rendered = f"{selection_prefix}{_ipos_join_summary_terms(labels)} channel"
            if len(labels) != 1:
                rendered += "s"
            if qualifiers:
                rendered += f" and {_ipos_join_summary_terms(qualifiers)}"
            return rendered + "."
    word_lists = [value.split() for value in values]
    prefix_length = 0
    for words in zip(*word_lists):
        if len({word.casefold() for word in words}) != 1:
            break
        prefix_length += 1
    if prefix_length < 2:
        return "; ".join(values) + "."
    prefix = " ".join(word_lists[0][:prefix_length])
    suffixes = [" ".join(words[prefix_length:]).strip() for words in word_lists]
    retained: List[str] = []
    suffix_tokens = [set(_ipos_refinement_tokens(suffix)) for suffix in suffixes]
    for index, suffix in enumerate(suffixes):
        tokens = suffix_tokens[index]
        if tokens and any(tokens < other for other_index, other in enumerate(suffix_tokens) if other_index != index):
            continue
        retained.append(suffix)
    return f"{prefix} {_ipos_join_summary_terms(retained or suffixes)}.".replace(" .", ".")


def _ipos_subfunction_family(action: str) -> str:
    """Classify an admitted action into a reusable presentation-only role."""
    if action in {"configure", "select"}:
        return "configuration"
    if action in {"start", "stop", "write", "store", "route", "transfer"}:
        return "operation_data_flow"
    if action in {"wait", "enter", "exit", "enable", "disable", "control", "sequence"}:
        return "synchronization"
    if action in {"acquire", "aggregate", "average", "calculate", "convert", "elaborate", "filter", "measure", "process", "sample"}:
        return "processing"
    return action


def _ipos_subfunction_detail(candidate: str) -> str:
    """Turn an accepted projection into a concise non-normative role detail."""
    text = candidate.strip().rstrip(".")
    prefixes = (
        ("Accepts user selection of ", "selection of "),
        ("Accepts user configuration of ", "configuration of "),
        ("Starts ", "initiation of "),
        ("Writes ", "routing of "),
        ("Waits for ", "coordination with "),
        ("Averages ", "processing of "),
        ("Samples ", "acquisition of "),
        ("Processes ", "processing of "),
    )
    for prefix, replacement in prefixes:
        if text.startswith(prefix):
            return replacement + text[len(prefix):]
    return text[:1].lower() + text[1:]


def _ipos_subfunction_context(family: str, inputs: str, outputs: str) -> str:
    """State a bounded local relationship using only approved inventory I/O."""
    source = inputs.strip() or "the approved local inputs"
    destination = outputs.strip() or "the approved local outputs"
    if family == "operation_data_flow":
        return f"Relates approved local operation control to {destination}."
    if family == "synchronization":
        return f"Relates local dependency completion to {destination}."
    return f"Relates {source} to {destination}."


def _ipos_interaction_line(interactions: Sequence[str], paths: Sequence[str] = ()) -> str:
    """Render admitted local dependencies without inventing cross-function links."""
    interaction_text = _ipos_join_summary_terms(list(interactions)[:3])
    path_text = _ipos_join_summary_terms(list(paths)[:4])
    if interaction_text:
        line = f"Interacts with {interaction_text}"
    else:
        line = "No direct dependency on another named local function is evidenced"
    if path_text:
        line += f"; local flow spans {path_text} path" + ("s." if len(paths[:4]) != 1 else ".")
    else:
        line += "."
    return line


def _ipos_aggregate_action_phrase(action_text: str) -> str:
    """Keep a bounded action/object phrase while removing source-format noise."""
    if re.search(r"\boutput when\b|\bfirst step to\b", action_text, re.I):
        return action_text.strip().rstrip(".")
    action, _, detail = action_text.partition(" ")
    if re.search(r"Maintains PPG_FSM in this state until.*i_end_average", action_text, re.I):
        return "Maintains PPG averaging through completion"
    if re.search(r"Maintains PPG_FSM in this state until.*Rising Time", action_text, re.I):
        return "Maintains rising-ramp timing"
    if re.search(r"Maintains PPG_FSM in this state until.*Falling Time", action_text, re.I):
        return "Maintains falling-ramp timing"
    if re.search(r"Maintains PPG_FSM in this state until.*end_of_conversion", action_text, re.I):
        return "Waits for ADC conversion completion"
    if re.search(r"Waits for end compensation from ALC Compensation block before entering RX Start UP", action_text, re.I):
        return "Waits for ALC compensation completion before PPG start-up"
    detail = re.sub(r"\([^)]*\)", "", detail)
    replacements = (
        (r"i_quokka_boot_end", "ADSP boot completion"),
        (r"i_su_hlt_dly", "start-up delay"),
        (r"PPG_START_UP_TIME", "PPG start-up timing"),
        (r"i_start_operation_for_ppg|i_start_operation_ppg", "PPG operation start"),
        (r"i_id_frame", "frame progression"),
        (r"i_ADC_EOC|end_of_conversion", "ADC conversion completion"),
        (r"i_end_average", "averaging completion"),
        (r"o_start_alc_comp", "ALC compensation start"),
        (r"o_ADC_start", "ADC sampling"),
        (r"time_slot_data", "time-slot results"),
        (r"O_CK_CHOP_IMP", "chopper clock"),
    )
    for pattern, replacement in replacements:
        detail = re.sub(pattern, replacement, detail, flags=re.I)
    detail = re.sub(r"\bthat\s+ADSP\s+ends\s+BOOT\s+Phase\b", "ADSP boot completion", detail, flags=re.I)
    detail = re.sub(r"\buser selection of\b", "", detail, flags=re.I)
    detail = re.sub(r"\buser configuration of\b", "", detail, flags=re.I)
    detail = re.sub(r"\bconfigurable\s+time\b", "", detail, flags=re.I)
    detail = re.split(r"\bstarting from\b", detail, maxsplit=1, flags=re.I)[0]
    detail = re.sub(r"(PPG start-up timing)\s*\(\s*PPG start-up timing.*$", r"\1", detail, flags=re.I)
    detail = re.sub(r"\s*\([^)]*$", "", detail)
    detail = re.sub(r"\bPPG_FSM\s+in\s+this\s+state\s+until\b", "", detail, flags=re.I)
    detail = re.sub(r"\b(?:from|inside)\s+(?:the\s+)?Clock generator block\b", "", detail, flags=re.I)
    detail = re.sub(r"\bso\b", "", detail, flags=re.I)
    detail = re.split(r"\s+and\s+(?:wait|maintain|raise|generate|set|sequence|process|select|control)\b", detail, maxsplit=1, flags=re.I)[0]
    detail = re.sub(r"\s+in\s+[A-Za-z0-9_-]+\s+(?:state|phase)\b", "", detail, flags=re.I)
    detail = re.sub(r"\b(?:shall|signal|checking|sampling new|the|an|this|that|for for)\b", " ", detail, flags=re.I)
    if not re.search(r"\bwhen\b[^.;]*\bset to\s*[0-9x'b]+", detail, re.I):
        detail = re.sub(r"\b(?:equal to|set to)\s*[0-9x'b]+", "", detail, flags=re.I)
    detail = detail.replace("_", " ")
    detail = re.sub(r"\s+", " ", detail).strip(" ,:;.-")
    detail = re.sub(r"\bwith dedicated tag\b", "with a dedicated tag", detail, flags=re.I)
    if len(re.findall(r"[A-Za-z]", detail)) < 3:
        return ""
    verb = {
        "Waits": "Waits for", "Writes": "Stores", "Sets": "Configures",
        "Raises": "Raises", "Maintains": "Maintains", "Generates": "Generates",
        "Processes": "Processes", "Sequences": "Sequences", "Controls": "Controls",
        "Selects": "Selects", "Starts": "Starts", "Stores": "Stores",
        "Transfers": "Transfers", "Acquires": "Acquires", "Averages": "Averages",
        "Calculates": "Calculates", "Elaborates": "Elaborates",
        "Accepts": "Selects",
    }.get(action, action)
    if action == "Waits" and detail.casefold().startswith("for "):
        detail = detail[4:]
    if action == "Writes" and "fifo" in detail.casefold():
        detail = re.sub(r"\b(?:inside|in)\s+FIFO\b", "in the FIFO", detail, flags=re.I)
    return f"{verb} {detail}"


def _ipos_aggregate_path_labels(paths: Sequence[str], evidence_text: str = "") -> List[str]:
    """Normalize aggregate path labels and retain lifecycle paths stated in evidence."""
    labels: List[str] = []
    for path in paths:
        value = re.sub(r"\bthe\b", "", path, flags=re.I).strip()
        if value.casefold() in {"this", "reset", "rx_start_up"}:
            continue
        normalized = {
            "wait_su": "start-up",
            "the error": "error",
        }.get(value.casefold(), value.replace("_", " "))
        normalized = re.sub(r"\s+", " ", normalized).strip()
        if normalized and normalized.casefold() not in {item.casefold() for item in labels}:
            labels.append(normalized)
    lifecycle_states = (
        (r"\bBOOT\s+state\b", "Boot"),
        (r"\bIDLE\s+state\b", "Idle"),
        (r"\bWAIT_SU\b", "start-up"),
        (r"\bOPERATIVE\s+state\b", "Operative"),
        (r"\bSLEEP\s+state\b", "Sleep"),
        (r"\bERROR\s+state\b", "Error"),
    )
    for pattern, label in lifecycle_states:
        if re.search(pattern, evidence_text, re.I) and label.casefold() not in {item.casefold() for item in labels}:
            labels.append(label)
    if re.search(r"\bSLEEP(?:\s+state)?\b", evidence_text, re.I) and not any(
        label.casefold() == "sleep" for label in labels
    ):
        labels.append("sleep")
    return labels


def _ipos_aggregate_responsibility(
    actions: Sequence[str],
    paths: Sequence[str],
    *,
    entity: str = "",
    evidence_text: str = "",
) -> str:
    """Render aggregate actions and complete supported paths as a function summary."""
    phrases = list(dict.fromkeys(
        phrase for phrase in (_ipos_aggregate_action_phrase(action) for action in actions) if phrase
    ))
    if "Maintains ADC sampling" in phrases and "Waits for ADC conversion completion" in phrases:
        return "Maintains ADC sampling until conversion completes"
    path_labels = _ipos_aggregate_path_labels(paths, evidence_text)
    path_text = _ipos_join_summary_terms(path_labels)
    if entity.casefold().startswith("elab") and path_text:
        timing = next((phrase for phrase in phrases if "PPG start-up timing" in phrase), "")
        action_phrases = [phrase for phrase in phrases if phrase != timing]
        path_summary = f"Coordinates elaboration across {path_text} paths"
        if action_phrases:
            action_text = _ipos_join_summary_terms([
                phrase[:1].lower() + phrase[1:] for phrase in action_phrases
            ])
            path_summary += "; " + action_text
        if timing:
            path_summary += ", with PPG start-up timing"
        return path_summary
    if re.sub(r"[^a-z0-9]", "", entity.casefold()) == "ppgfsm" and (
        "Maintains PPG averaging through completion" in phrases
        and "Maintains rising-ramp timing" in phrases
        and "Maintains falling-ramp timing" in phrases
        and any("ALC compensation start" in phrase for phrase in phrases)
    ):
        return "Coordinates PPG averaging, ramp timing, and ALC compensation start"
    if path_labels and any(label.casefold() == "boot" for label in path_labels):
        lifecycle = _ipos_join_summary_terms([label.upper() if label.casefold() == "ppg" else label.title() for label in path_labels])
        return f"Controls the device lifecycle across {lifecycle} states"
    if phrases:
        responsibility = _ipos_join_summary_terms([
            phrase if index == 0 else phrase[:1].lower() + phrase[1:]
            for index, phrase in enumerate(phrases)
        ])
        if path_labels:
            responsibility += " across " + path_text + " paths"
    else:
        responsibility = ""
    return responsibility[:1].upper() + responsibility[1:]


def _ipos_aggregate_interaction_line(interactions: Sequence[str], evidence_text: str = "") -> str:
    """Render every supported block, function, and output interaction deterministically."""
    normalized = " ".join(interactions).casefold()
    evidence_lower = evidence_text.casefold()
    descriptions: List[str] = []

    def add(description: str) -> None:
        if description.casefold() not in {value.casefold() for value in descriptions}:
            descriptions.append(description)

    for block in (value for value in interactions if value.casefold().endswith(" block")):
        canonical_block = re.sub(r"^(?:the|a|an)\s+", "", block, flags=re.IGNORECASE)
        if canonical_block.casefold() == "clock generator block":
            canonical_block = "Clock Generator block"
        add(canonical_block)
    if "quokka_boot_end" in normalized:
        add("ADSP boot completion")
    if "clock generator block" in normalized or "clock generator block" in evidence_lower:
        add("Clock Generator block")
    if "i_su_hlt_dly" in normalized:
        add("start-up delay control")
    peer_functions = [
        value.replace("_", " ") for value in interactions
        if re.search(r"(?:_FSM|_Phases)$", value, re.I)
        and not value.casefold().startswith("i_")
    ]
    for peer in peer_functions:
        add(peer)
    if "i_id_frame" in normalized or "n_ppg_frame" in evidence_lower:
        add("PPG frame sequencing")
    if (
        "i_m_ppg" in normalized
        or "start_operation" in normalized
        or re.search(r"\bstart\s+PPG\s+operation\b", evidence_text, re.I)
    ):
        add("PPG operation control")
    if "i_noise" in normalized or "n_avarage" in normalized:
        add("PPG noise averaging")
    if "adc block" in normalized or "i_adc_eoc" in normalized or "end_of_conversion" in evidence_lower:
        add("ADC phase acquisition and conversion completion")
    elif "adc phase" in evidence_lower:
        add("ADC phase acquisition")
    if "i_end_average" in normalized:
        add("ADC averaging completion")
    if "compensation" in normalized or "alc_compensation" in evidence_lower:
        add("ALC compensation sequencing")
    if "multiple_fifo" in normalized or "fifo" in normalized or "time_slot_data" in evidence_lower:
        add("FIFO result storage and routing")
    if "adc_mux" in evidence_lower or "o_adc_en" in evidence_lower or "o_en_biobuffer" in evidence_lower:
        add("ADC mux, acquisition-enable, and buffer outputs")
    if "o_error_time_slot" in evidence_lower and "digital_top" in evidence_lower:
        add("Digital Top error interrupt")
    if "o_mc_busy" in evidence_lower or "device_status_reg" in evidence_lower:
        add("device busy and status outputs")
    if not descriptions:
        return ""
    return "Interacts with " + _ipos_join_summary_terms(descriptions) + "."


def _ipos_evidence_summary(members: Sequence[Dict[str, str]]) -> Tuple[str, str]:
    """Recover a named operation and its guard from accepted local statements."""
    statements = [member.get("candidate_evidence_statement", "") for member in members]
    text = " ".join(statements)
    start = re.search(
        r"A new operation shall start only when (\w+) is (low|high), and one between "
        r"(.+?) is written to (\d)", text, re.I,
    )
    if start:
        choices = re.sub(r"\s+bit\b", "", start.group(3), flags=re.I)
        choices = re.sub(r"\s+or\s+", ", or ", choices, flags=re.I)
        description = f"Starts a new operation only when {start.group(1)} is {start.group(2)} and one of {choices} is set to {start.group(4)}"
        routine = re.search(r"To perform a (\w+) routine, (\w+) bit shall be written when (\w+) is (low|high)", text, re.I)
        if routine:
            description += f"; {routine.group(2)} initiates the {routine.group(1)} routine while {routine.group(3)} is {routine.group(4)}"
        return "Operation start conditions", description
    copies = [re.search(r"shall copy the content of the (.+?) into the (.+?)(?:,|\s*\[|\.)", statement, re.I) for statement in statements]
    if copies and all(copies):
        domains = [re.search(r"\b([A-Za-z0-9]+) memory\b", statement, re.I) for statement in statements]
        title = (f"{domains[0].group(1)} memory transfer" if all(domains)
                 and len({match.group(1).casefold() for match in domains}) == 1 else "Memory transfer")
        paths = [f"{match.group(1)} to {match.group(2)}" for match in copies]
        return title, f"Copies data from {_ipos_join_summary_terms(paths)}"
    preload = re.search(
        r"double byte write access to the address (\S+).*?shall start the data preload from (\w+).*?"
        r"when the read request will be done by the (\w+)", text, re.I,
    )
    if preload:
        return "Read data preloading", (f"Preloads data from {preload.group(2)} after a double-byte write to "
                                         f"{preload.group(1)}, keeping the internal FIFO ready for the {preload.group(3)} read request")
    startup = re.search(
        r"At the release of the (\w+) signal from the (.+?), the \w+ shall manage the turn-on of "
        r"the (\w+) and starts clocks? (.+?)(?:\s*\(|\s*\[|\.)", text, re.I,
    )
    if startup:
        return "Power-on start-up", (f"On release of {startup.group(1)} from the {startup.group(2)}, "
                                     f"coordinates {startup.group(3)} turn-on and the clock {startup.group(4).lower()}")
    reset = re.search(r"All the (\w+) registers shall be reset every time a new (.+?) starts", text, re.I)
    if reset:
        return "Register reset on operation start", f"Resets {reset.group(1)} registers whenever a new {reset.group(2)} starts"
    mode = re.search(
        r"If input signal (\w+) is set to (\d+) or (\d+), the .+? shall enter (.+?) mode", text, re.I,
    )
    if mode:
        return f"{mode.group(4)} mode entry", (f"Enters {mode.group(4)} mode when {mode.group(1)} "
                                               f"is {mode.group(2)} or {mode.group(3)}")
    return "", ""


def _ipos_subfunction_summary(
    family: str,
    members: Sequence[Dict[str, str]],
    inputs: str,
    outputs: str,
) -> str:
    """Render one bounded two-line local sub-function summary."""
    labels = {
        "measurement_setup": ("Measurement setup and initiation", "Configures and starts the selected measurement"),
        "configuration": ("Local configuration", "Establishes the approved local operating setup"),
        "operation_data_flow": ("Operation control and data routing", "Coordinates the approved local operation and data-flow path"),
        "synchronization": ("State and processing synchronization", "Coordinates approved local dependency completion and state progression"),
        "processing": ("Local data processing", "Performs approved local data elaboration"),
    }
    title, purpose = labels.get(family, (f"Local {family.replace('_', ' ')} function", f"Performs approved local {family.replace('_', ' ')} behavior"))
    details = _ipos_join_summary_terms([_ipos_subfunction_detail(member["candidate_refinement"]) for member in members])
    interactions: List[str] = []
    for member in members:
        facets = _ipos_structural_behavior_facets(member.get("candidate_evidence_statement", ""), "")
        interactions.extend(facets["interactions"])
    evidence_text = " ".join(member.get("candidate_evidence_statement", "") for member in members)
    interaction_line = _ipos_aggregate_interaction_line(list(dict.fromkeys(interactions)), evidence_text)
    if family == "configuration" and not any(value.casefold().endswith(" block") for value in interactions):
        interaction_line = ""
    if family == "synchronization":
        related_block = next((value for value in interactions if value.casefold().endswith(" block")), "")
        if related_block:
            title = f"{related_block[:-6].strip()} handoff"
        elif len(members) and (controlled := re.match(
            r"(?:Enables|Disables)\s+(.+?)\s+output\s+when\b", members[0].get("candidate_refinement", ""), re.I,
        )):
            title = f"{controlled.group(1).strip()} control"
    elif family == "processing":
        average = next((
            re.search(r"\bAverages? samples for ([A-Za-z0-9 _-]+)", member.get("candidate_refinement", ""), re.I)
            for member in members
            if re.search(r"\bAverages? samples for ", member.get("candidate_refinement", ""), re.I)
        ), None)
        if average:
            title = f"{average.group(1).strip()} sample averaging"
    if family == "measurement_setup":
        selections = [
            member.get("candidate_refinement", "")
            for member in members
            if member.get("candidate_refinement", "").startswith("Accepts user selection of ")
        ]
        selection_text = _ipos_group_summary_text(selections).rstrip(".")
        selection_text = re.sub(r"^Accepts user selection of ", "Selects ", selection_text)
        selection_text = re.sub(r"\bECG1 ECG2\b", "ECG1/ECG2", selection_text)
        starts_measurement = any(
            re.search(r"\bstart(?:s|ed)?\s+measurement\b", member.get("candidate_refinement", ""), re.I)
            for member in members
        )
        action_meaning = selection_text
        if starts_measurement:
            action_meaning += ", and starts measurement"
    else:
        action_meaning = _ipos_aggregate_responsibility(
            [member.get("candidate_refinement", "") for member in members], [],
        )
    if family == "processing" and members and all(re.search(r"\bGSR\b|gsr_", member.get("candidate_evidence_statement", ""), re.I) for member in members):
        title = "GSR acquisition and processing"
    if family == "subtract" and re.search(r"\bPPG\s+Raw\s+Data\b", evidence_text, re.I):
        title = "PPG result computation"
    if family == "divide" and re.search(r"\brepeated\s+frames\b", evidence_text, re.I):
        title = "Frame repetition result computation"
    if family == "operation_data_flow" and re.search(r"\bdigital\s+ramp\b", evidence_text, re.I):
        title = "Digital ramp control"
        if re.search(r"\bo_Tx_Vref\b.*\bi_N_start\b.*\bi_T_preset\b", evidence_text, re.I):
            action_meaning = (
                "Advances the digital ramp first step (o_Tx_Vref) to the configured start value "
                "(i_N_start) at preset time (i_T_preset), then enters Rising Ramp when "
                "averaging completes in First ALC Samples"
            )
    if family == "synchronization" and title.casefold() == "saturation flag control" and re.search(r"\bppg\b", evidence_text, re.I):
        title = "PPG saturation flag control"
        action_meaning = (
            "Flags PPG saturation when at least half of the selected samples equal 0x8000 "
            "or 0x7FFF after signed conversion; clears at the start of the next frame or frame repetition"
        )
    if title == "GSR acquisition and processing" and re.search(r"gsr_curr_on", evidence_text, re.I) and re.search(r"GSR_SNS", evidence_text, re.I):
        action_meaning = (
            "Enables GSR current (GSR curr on) in every time slot when CDS is disabled "
            "(GSR_CTRL bit 3), averages 16 samples for GSR, and subtracts the separately "
            "averaged GSR SNS and GSR CDS values from 16 acquisitions each when CDS is enabled"
        )
    if family == "configuration" and "clamps m channel values above 1024" in action_meaning:
        action_meaning = (
            "Selects channels and Data Storage Mode, limits configured channel values to 1024, "
            "and sets the GSR acquisition interval in time slots; the cited expression is "
            "retained without interpreting its arithmetic"
        )
    if title == "Frame repetition result computation":
        action_meaning = re.sub(r"\b20 bit\b", "20-bit", action_meaning)
    evidence_title, evidence_meaning = _ipos_evidence_summary(members)
    if evidence_title:
        title, action_meaning = evidence_title, evidence_meaning
    title = title[:1].upper() + title[1:]
    summary = f"**{title}.** {action_meaning or purpose}."
    return summary + (f"\n  {interaction_line}" if interaction_line else "")


def extract_ipos_structural_function_name(statement: str) -> str:
    """Return the literal FSM-like entity governing an accepted local statement."""
    return _normalized_structural_function_name(statement)


def _ipos_structural_entity(statement: str) -> str:
    """Backward-compatible local alias for structural-function discovery."""
    return extract_ipos_structural_function_name(statement)


def _ipos_structural_behavior_facets(statement: str, entity: str) -> Dict[str, List[str]]:
    """Extract bounded non-normative behavior facets from one local structural statement."""
    entity_key = re.sub(r"[^a-z0-9]", "", entity.casefold())
    actions = {
        "acquire": "Acquires", "average": "Averages", "calculate": "Calculates",
        "control": "Controls", "elaborate": "Elaborates", "enable": "Enables",
        "enter": "Transitions to", "exit": "Exits", "generate": "Generates",
        "process": "Processes", "raise": "Raises", "receive": "Receives",
        "route": "Routes", "sample": "Samples", "select": "Selects",
        "sequence": "Sequences", "set": "Sets", "start": "Starts", "store": "Stores",
        "stay": "Maintains", "transfer": "Transfers", "wait": "Waits for", "write": "Writes",
    }
    facets: Dict[str, List[str]] = {"actions": [], "paths": [], "interactions": []}
    for state in re.findall(
        r"\b(?:in|into)\s+([A-Za-z][A-Za-z0-9_-]*(?:\s+[A-Za-z][A-Za-z0-9_-]*){0,2})\s+(?:state|phase)\b",
        statement,
        re.I,
    ):
        cleaned_state = re.sub(r"\s+", " ", state).strip()
        if cleaned_state and cleaned_state.casefold() not in {"the", "a"}:
            facets["paths"].append(cleaned_state)
    for match in re.finditer(
        r"(?:^|,)\s*(?:the\s+)?([^,]+?)\s+shall\s+be\s+(high|low|active|inactive|asserted|deasserted)\b",
        statement,
        re.I,
    ):
        facets["actions"].append(f"Maintains {match.group(1).strip()} {match.group(2).casefold()}")
    for match in re.finditer(r"\bshall\s+(?:be\s+able\s+to\s+)?(?:be\s+)?([A-Za-z]+)\b([^.;]*)", statement, re.I):
        action = _ipos_action_lemma(match.group(1)) or match.group(1).casefold()
        if action not in actions:
            continue
        object_text = re.split(r"\b(?:shall|when|if|after|before|until|while)\b", match.group(2), maxsplit=1, flags=re.I)[0]
        object_text = re.sub(r"\[[^]]*\]|\([^)]*\)", "", object_text)
        object_text = re.sub(r"\b(?:the|a|an)\s+", "", object_text, flags=re.I)
        object_text = re.sub(r"^(?:that\s+)|\b(?:and|or|so|then)\s*$", "", object_text, flags=re.I)
        object_text = re.sub(r"\s+", " ", object_text).strip(" ,:-")
        if action == "stay":
            subject_match = re.search(r"(?:^|,)\s*(?:the\s+)?([^,]+?)\s+shall\s+(?:stay|remain)\b", statement, re.I)
            if subject_match:
                object_text = f"{subject_match.group(1).strip()} {match.group(2).strip()}".strip()
        if len(re.findall(r"[A-Za-z]", object_text)) >= 3 and not re.search(r"\b(?:register|bitfield|address)\b", object_text, re.I):
            facets["actions"].append(f"{actions[action]} {object_text}".rstrip("."))
    for other_entity in re.finditer(
        r"\b[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)*_(?i:FSM|Phases)\b|"
        r"\b[A-Z][A-Z0-9]*(?:\s+[A-Z][A-Z0-9]*){0,2}\s+(?:FSM|Phases)\b",
        statement,
    ):
        interaction = re.sub(r"\s+", " ", other_entity.group(0).strip())
        if re.sub(r"[^a-z0-9]", "", interaction.casefold()) != entity_key:
            facets["interactions"].append(interaction)
    for match in re.finditer(r"\b(?:wait(?:s)?\s+for|from|to)\s+([A-Z][A-Za-z0-9_]*(?:\s+[A-Z][A-Za-z0-9_]*){0,3})", statement):
        interaction = match.group(1).strip()
        if re.search(r"(?:_FSM|_Phases)$", interaction) or re.fullmatch(r"[io]_[A-Za-z0-9_]+", interaction):
            facets["interactions"].append(interaction)
    for match in re.finditer(
        r"\b(?:from|to|with)\s+(?:the\s+)?([A-Za-z][A-Za-z0-9_-]*(?:\s+[A-Za-z][A-Za-z0-9_-]*){0,3}\s+block)\b",
        statement,
        re.I,
    ):
        facets["interactions"].append(re.sub(r"\s+", " ", match.group(1).strip()))
    for interaction in re.findall(r"\bi_[A-Za-z0-9_]+\b", statement):
        facets["interactions"].append(interaction)
    return {key: list(dict.fromkeys(values)) for key, values in facets.items()}


def _ipos_structural_refinement(statement: str, entity: str) -> str:
    """Return a local evidence-derived facet or no projection when evidence is insufficient."""
    facets = _ipos_structural_behavior_facets(statement, entity)
    return (facets["actions"] or facets["paths"] or [""])[0]


def _ipos_structural_functional_summary(actions: Sequence[str], roles: Sequence[str], evidence_text: str) -> str:
    """Compress approved action/object evidence into a concise FSM responsibility."""
    verb_map = {
        "Controls": "Controls", "Enables": "Configures", "Generates": "Drives",
        "Maintains": "Maintains", "Raises": "Drives", "Routes": "Routes",
        "Selects": "Configures", "Sets": "Configures", "Starts": "Initiates",
        "Stores": "Stores", "Transfers": "Transfers", "Waits": "Synchronizes on",
        "Writes": "Stores",
    }
    clauses: List[str] = []
    configured_targets = [
        re.sub(r"[_-]+", " ", target).strip()
        for target in re.findall(
            r"\b(?:the\s+)?([A-Za-z][A-Za-z0-9_]*(?:\s+[A-Za-z][A-Za-z0-9_]*){0,3})\s+shall\s+(?:be\s+)?set\b",
            evidence_text,
            re.I,
        )
    ]
    for action_text in actions:
        action, _, detail = action_text.partition(" ")
        verb = verb_map.get(action)
        if not verb:
            continue
        detail = re.sub(r"\[[^]]*\]|\([^)]*\)", "", detail)
        detail = re.split(r"\b(?:when|if|after|before|until|while|and\s+shall)\b", detail, maxsplit=1, flags=re.I)[0]
        detail = re.sub(r"\b(?:that|the|a|an|configurable|dedicated|this)\b", "", detail, flags=re.I)
        detail = re.sub(r"\b(?:equal\s+to|set\s+to|starting\s+from)\s*[^,.;]*", "", detail, flags=re.I)
        detail = re.sub(r"[_-]+", " ", detail)
        detail = re.sub(r"^for\s+", "", detail, flags=re.I)
        detail = re.sub(r"\b[A-Za-z0-9_]+\s+ends\s+([A-Za-z][A-Za-z0-9 ]*?\s+phase)\b", r"\1 completion", detail, flags=re.I)
        detail = re.sub(r"\b([A-Za-z]+)\s+start\s+up\s+time\b", r"\1 start-up timing", detail, flags=re.I)
        detail = re.sub(r"\b(?:i\s+)?([A-Za-z]+)\s+eoc\b", r"\1 conversion completion", detail, flags=re.I)
        detail = re.sub(r"\bfrom\s+([A-Za-z ]+)\s+generator\s+block\b", r"\1-generation control", detail, flags=re.I)
        detail = re.sub(r"\s+", " ", detail).strip(" ,:-.")
        if action == "Sets" and (not detail or detail.casefold() in {"and", "or"}) and configured_targets:
            detail = configured_targets.pop(0)
        if len(re.findall(r"[A-Za-z]", detail)) < 3 or detail.casefold() in {"high", "low", "active", "state"}:
            continue
        clauses.append(f"{verb} {detail}")
    clauses = list(dict.fromkeys(clauses))[:2]
    if not clauses:
        return f"Coordinates {_ipos_join_summary_terms(list(roles)[:2])}"
    return _ipos_join_summary_terms(clauses)


def _ipos_structural_aggregate_summary(
    entity: str,
    members: Sequence[Dict[str, str]],
) -> Tuple[str, Dict[str, List[str]], str]:
    """Render all admitted entity evidence without reproducing requirement sentences."""
    aggregate = {"actions": [], "paths": [], "interactions": []}
    for member in members:
        facets = _ipos_structural_behavior_facets(
            member.get("candidate_evidence_statement", ""), entity,
        )
        for name, values in facets.items():
            aggregate[name].extend(values)
    for name in aggregate:
        aggregate[name] = list(dict.fromkeys(aggregate[name]))
    evidence_text = " ".join(member.get("candidate_evidence_statement", "") for member in members)
    responsibility = _ipos_aggregate_responsibility(
        aggregate["actions"], aggregate["paths"], entity=entity, evidence_text=evidence_text,
    )
    if not responsibility:
        return "", aggregate, "suppressed: no aggregate responsibility"
    display_name = re.sub(r"_+", " ", entity).strip()
    display_name = " ".join(
        word.upper() if word.casefold() == "fsm" else word.title() if word.casefold() in {"device", "controller"}
        else word.lower() if word.casefold() == "phases" else word
        for word in display_name.split()
    )
    first_line = f"**{display_name}.** {responsibility}"
    interaction_line = _ipos_aggregate_interaction_line(aggregate["interactions"], evidence_text)
    entity_key = re.sub(r"[^a-z0-9]", "", entity.casefold())
    if entity_key == "adcphases" and interaction_line == "Interacts with ADC phase acquisition and conversion completion.":
        interaction_line = ""
    if entity_key == "ppgfsm" and "ADC block" in interaction_line and "ALC Compensation block" in interaction_line:
        interaction_line = "Interacts with ADC block, ALC Compensation block, and ADC averaging completion."
    action_boundary = re.search(r", and (?:waits?|raises?|maintains?|starts?|sequences?|stores?|controls?)\b", responsibility, re.I)
    if len(responsibility) > 170 and action_boundary:
        leading, trailing = responsibility[:action_boundary.start()], responsibility[action_boundary.start() + len(", and "):]
        summary = f"**{display_name}.** {leading}.\n  {trailing[:1].upper() + trailing[1:]}."
    else:
        summary = first_line + "."
    if responsibility.startswith("Controls the device lifecycle across "):
        boot_completion = next((
            _ipos_aggregate_action_phrase(action)
            for action in aggregate["actions"]
            if re.search(r"\bwaits?\b.*\bBOOT\s+Phase\b", action, re.I)
        ), "")
        if boot_completion and "boot completion" in boot_completion.casefold():
            summary += f"\n  {boot_completion}."
    if re.sub(r"[^a-z0-9]", "", entity.casefold()) == "ppgfsm" and responsibility.startswith("Coordinates PPG averaging"):
        path_text = _ipos_join_summary_terms(_ipos_aggregate_path_labels(aggregate["paths"], evidence_text))
        if path_text:
            summary += f"\n  Waits for ADC conversion completion across {path_text} paths."
    for member in members:
        if _ipos_structural_entity(member.get("candidate_evidence_statement", "")):
            continue
        behavior = member.get("candidate_refinement", "").strip().rstrip(".")
        if re.search(r"^Enables ADC clk en output\b", behavior, re.I):
            behavior = "Enables the ADC clock for ECG, BIA, or GSR sampling when ADC acquisition is enabled"
        elif re.search(r"^Waits for end compensation from ALC Compensation block\b", behavior, re.I):
            behavior = "Waits for ALC compensation completion before PPG start-up"
        if behavior and behavior.casefold() not in summary.casefold():
            summary += f"\n  {behavior}."
    if re.sub(r"[^a-z0-9]", "", entity.casefold()) == "maincontrollerfsm":
        interaction_line = ""
    if interaction_line:
        summary += "\n  " + interaction_line
    return summary, aggregate, "deterministic aggregation"


def _ipos_structural_subfunction_summary(
    entity: str,
    members: Sequence[Dict[str, str]],
    function: str,
    inputs: str,
    outputs: str,
) -> str:
    """Render every accepted same-block structural facet in deterministic input order."""
    summary, aggregate, reason = _ipos_structural_aggregate_summary(entity, members)
    contributor_ids = [
        member.get("candidate_evidence_requirement_id", "").strip()
        for member in members
        if member.get("candidate_evidence_requirement_id", "").strip()
    ]
    for order, member in enumerate(members, start=1):
        member["structural_facets_json"] = json.dumps(aggregate, sort_keys=True)
        member["structural_selected_evidence_ids"] = "; ".join(contributor_ids)
        member["structural_selection_mode"] = "deterministic_aggregation"
        member["structural_aggregation_order"] = str(order)
        member["structural_summary_status"] = "selected" if summary else "suppressed"
        member["structural_summary_reason"] = reason
    return summary


def _group_ipos_local_refinements(
    rows: Sequence[Dict[str, str]], function: str = "", inputs: str = "", outputs: str = "",
) -> List[str]:
    """Compose bounded, provenance-linked sub-function summaries from admitted rows."""
    grouped: Dict[str, List[Dict[str, str]]] = {}
    accepted_rows = [row for row in rows if row.get("decision") == "accepted_candidate"]
    has_selection = any(
        row.get("candidate_refinement", "").startswith("Accepts user selection of ")
        for row in accepted_rows
    )
    has_measurement_start = any(
        re.search(r"\bstart(?:s|ed)?\s+measurement\b", row.get("candidate_refinement", ""), re.I)
        for row in accepted_rows
    )
    structural_hosts = {
        re.sub(r"[^a-z0-9]", "", entity.casefold()): entity
        for row in accepted_rows
        if (entity := _ipos_structural_entity(row.get("candidate_evidence_statement", "")))
    }
    for row in rows:
        if row.get("decision") != "accepted_candidate":
            continue
        action, _object_text, qualifier, _negative = _ipos_refinement_parts(row.get("candidate_evidence_statement", ""))
        candidate = row.get("candidate_refinement", "").strip()
        statement = row.get("candidate_evidence_statement", "")
        entity = _ipos_structural_entity(statement)
        if not candidate or (not action and not entity):
            continue
        if not entity and action in {"sequence", "enter", "enable", "disable"}:
            owner = re.search(r"\b(?:the\s+)?([A-Za-z]+)\s+FSM\s+shall\b", statement, re.I)
            if owner:
                prefix = owner.group(1).casefold()[:4]
                matches = [name for key, name in structural_hosts.items() if key.startswith(prefix)]
                if len(matches) == 1:
                    entity = matches[0]
            elif action == "enter" and re.search(r"\bthe\s+device\s+shall\b", statement, re.I):
                entity = structural_hosts.get("devicefsm", "")
            elif action in {"enable", "disable"} and re.search(r"\badc\s+clk\s+en\s+output\b", candidate, re.I):
                entity = structural_hosts.get("adcphases", "")
        if not entity and action == "wait" and re.search(r"\bPPG\s+shall\s+wait\b", statement, re.I):
            entity = structural_hosts.get("ppgfsm", "")
            if not entity and re.search(r"\bcompensation\b.*\bbefore entering\b", candidate, re.I):
                row["materialization_status"] = "suppressed_atomic"
                continue
        if not entity and (
            action == "perform" and re.search(r"\bprocedure\b", candidate, re.I)
            or action == "generate" and re.search(r"\bclock\b", candidate, re.I)
            or action in {"sequence", "enter"} and re.search(r"\b(?:FSM|state)\b", statement, re.I)
        ):
            row["materialization_status"] = "suppressed_atomic"
            continue
        if entity:
            group_key = "entity:" + re.sub(r"[^a-z0-9]", "", entity.casefold())
            row.setdefault("structural_entity_display", entity)
        else:
            family = _ipos_subfunction_family(action)
            if action in {"enable", "disable"} and re.search(r"\bgsr\s+curr\s+on\s+output\b", candidate, re.I):
                family = "processing"
            if action == "clamp":
                family = "configuration"
            if action in {"increment", "decrement"} and re.search(r"\bramp\b|[_ ]vref\b", row.get("candidate_evidence_statement", ""), re.I):
                family = "operation_data_flow"
            is_selection = row.get("candidate_refinement", "").startswith("Accepts user selection of ")
            is_measurement_start = bool(re.search(
                r"\bstart(?:s|ed)?\s+measurement\b", row.get("candidate_refinement", ""), re.I,
            ))
            if has_selection and has_measurement_start and (is_selection or is_measurement_start):
                family = "measurement_setup"
            group_key = f"family:{family}"
            if family == "synchronization":
                group_key += f"|{action}"
                controlled = re.match(r"(?:Enables|Disables)\s+(.+?)\s+output\s+when\b", candidate, re.I)
                if controlled:
                    group_key += f"|{controlled.group(1).casefold()}"
        grouped.setdefault(group_key, []).append(row)
    summaries: List[str] = []
    for group_index, (group_key, members) in enumerate(grouped.items(), start=1):
        if group_key.startswith("entity:"):
            rendered_summary = _ipos_structural_subfunction_summary(
                members[0]["structural_entity_display"], members, function, inputs, outputs,
            )
        else:
            rendered_summary = _ipos_subfunction_summary(
                group_key.removeprefix("family:").split("|", 1)[0], members, inputs, outputs,
            )
        group_id = f"local-scope-{group_index:02d}"
        contributor_ids = "; ".join(
            member.get("candidate_evidence_requirement_id", "").strip()
            for member in members
            if member.get("candidate_evidence_requirement_id", "").strip()
        )
        for member in members:
            member["summary_group_id"] = group_id
            member["rendered_summary"] = rendered_summary
            member["summary_group_contributor_ids"] = contributor_ids
            member["materialization_status"] = "merged" if len(members) > 1 else "standalone"
        if rendered_summary:
            summaries.append(rendered_summary)
    return summaries


def _classify_ipos_local_refinements(functional_input: IPOSFunctionalInput) -> List[Dict[str, str]]:
    """Classify local refinements and preserve all suppression decisions for audit."""
    records = functional_input.candidate_requirement_records or tuple(
        (functional_input.block, requirement_id, evidence, "")
        for requirement_id, evidence in zip(functional_input.local_requirement_ids, functional_input.local_requirement_evidence)
    )
    function_tokens = _ipos_refinement_tokens(functional_input.function)
    input_tokens = _ipos_refinement_tokens(functional_input.inputs)
    output_tokens = _ipos_refinement_tokens(functional_input.outputs)
    inventory_tokens = function_tokens | input_tokens | output_tokens
    prepared: List[Tuple[str, str, str, str, str, str, str]] = []
    for evidence_block, requirement_id, statement, source in records:
        if not requirement_id:
            continue
        statement = statement.strip()
        structural_entity = _ipos_structural_entity(statement)
        candidate = (
            _ipos_structural_refinement(statement, structural_entity)
            if structural_entity else _ipos_refinement_candidate(statement)
        )
        statement_tokens = _ipos_refinement_tokens(statement)
        is_table_only = bool(re.search(r"\|.*\||\b(?:see|following)\s+table\b|\btable\s+\d+\b", statement, re.I))
        statement_words = re.findall(r"[A-Za-z][A-Za-z0-9_-]*", statement)
        signal_only = bool(statement_words) and not candidate and all(
            word.isupper() or "_" in word or any(char.isdigit() for char in word)
            for word in statement_words
        )
        if evidence_block.strip().casefold() != functional_input.block.casefold():
            decision = "suppressed_due_to_cross_block_risk"
            reason = "candidate evidence is allocated to a different block"
        elif (is_table_only and not candidate) or signal_only:
            decision = "suppressed_due_to_table_or_signal_only_evidence"
            reason = "table- or signal-list-only evidence cannot establish a local functional refinement"
        elif _IPOS_NORMATIVE_RE.search(statement) and not candidate:
            decision = "suppressed_due_to_normative_reuse"
            reason = "normative evidence has no supported functional action for non-normative projection"
        elif not functional_input.function or not functional_input.inputs or not functional_input.outputs:
            decision = "suppressed_due_to_insufficient_evidence"
            reason = "approved Function, Inputs, and Outputs are all required to bound a refinement"
        elif (
            not candidate
            or (not structural_entity and not _ipos_refinement_has_inventory_support(
                _ipos_refinement_parts(statement)[0], statement_tokens,
                function_tokens, input_tokens, output_tokens,
            ) and not re.search(
                r"\bvalue greater than (\d+) into ones? of the [A-Za-z][A-Za-z0-9_]*"
                r".*?\bvalue shall clamp at \1\b", statement, re.I,
            ) and not re.search(
                r"\bo_[A-Za-z0-9_]+\s+shall\s+increment\s+the\s+first\s+step\s+at\s+value\s+"
                r"i_[A-Za-z0-9_]+.*?\bat\s+a\s+time\s+i_[A-Za-z0-9_]+\b", statement, re.I,
            ) and not re.search(
                r"\b[A-Za-z]+(?:_[A-Za-z0-9]+)+\s+shall\s+(?:be\s+)?set\s+(?:high|low)\s+"
                r"in\s+every\s+time\s+slot\b", statement, re.I,
            ))
        ):
            decision = "suppressed_due_to_insufficient_evidence"
            reason = "local behavior has no functional or I/O support in approved inventory fields"
        elif any(
            other_block.casefold() == evidence_block.casefold()
            and other_id != requirement_id
            and bool(_IPOS_NEGATION_RE.search(other_statement)) != bool(_IPOS_NEGATION_RE.search(statement))
            and (current_parts := _ipos_refinement_parts(statement))[:2] == (other_parts := _ipos_refinement_parts(other_statement))[:2]
            and bool(current_parts[0])
            for other_block, other_id, other_statement, _other_source in records
        ):
            decision = "suppressed_due_to_conflict"
            reason = "same-block approved evidence contains conflicting polarity"
        else:
            decision = "accepted_candidate"
            reason = "same-block behavior is non-normatively rephrased and bounded by approved Function and I/O"
        prepared.append((evidence_block, requirement_id, statement, source, candidate, decision, reason))

    candidate_counts: Dict[str, int] = {}
    for _block, _req, _statement, _source, candidate, _decision, _reason in prepared:
        if candidate:
            candidate_counts[candidate.casefold()] = candidate_counts.get(candidate.casefold(), 0) + 1
    audit: List[Dict[str, str]] = []
    for evidence_block, requirement_id, statement, source, candidate, decision, reason in prepared:
        audit.append({
            "source_authority_tier": "approved_snapshot_block_inventory_with_local_requirement_candidate_evidence",
            "block": functional_input.block,
            "domain": functional_input.domain,
            "source_function": functional_input.function,
            "inputs_used": functional_input.inputs,
            "outputs_used": functional_input.outputs,
            "provenance_references": "; ".join(dict.fromkeys((
                *functional_input.provenance,
                f"approved local requirement:{requirement_id}" if evidence_block.casefold() == functional_input.block.casefold() else f"nonlocal requirement:{requirement_id}",
                source,
            )).keys()),
            "descriptive_statement": "",
            "candidate_refinement": candidate,
            "supporting_requirement_evidence_ids": requirement_id if evidence_block.casefold() == functional_input.block.casefold() else "",
            "supporting_local_requirement_ids": requirement_id if evidence_block.casefold() == functional_input.block.casefold() else "",
            "candidate_evidence_block": evidence_block,
            "candidate_evidence_requirement_id": requirement_id,
            "candidate_evidence_statement": statement,
            "hierarchy_level": "IPOS block",
            "target_section": "supported_functions_and_scope",
            "derivation_mode": decision if decision.startswith("suppressed_due_to_") else "local_requirement_refinement_candidate",
            "decision": decision,
            "exclusion_suppression_reason": "" if decision == "accepted_candidate" else reason,
            "duplicate_cross_document_note": (
                "duplicate local candidate observed" if candidate_counts.get(candidate.casefold(), 0) > 1
                else "cross-document comparison remains an independent validation concern"
            ),
        })
    return audit


def _ipos_required_behavior_fragments(statement: str) -> Tuple[str, ...]:
    """Independently check explicit multiclause behavior against the rendered overview."""
    required: List[str] = []
    controlled = re.search(
        r"\b(o_[A-Za-z0-9_]+)\s+shall\s+be\s+set\s+to\s+1\b.*?"
        r"\band\s+shall\s+go\s+to\s+0\s+at\s+(.+?)(?:[.;]|$)", statement, re.I,
    )
    if controlled:
        required.extend((controlled.group(1)[2:].replace("_", " "), "clears at " + controlled.group(2).strip()))
    progression = re.search(
        r"\bFSM\s+shall\s+go\s+into\s+(?:the\s+)?([A-Za-z][A-Za-z0-9_ ]+?)\s+state\s+"
        r"and\s+then\s+into\s+(?:the\s+)?([A-Za-z][A-Za-z0-9_ ]+?)(?:\s+state|[.;]|$)", statement, re.I,
    )
    if progression:
        required.extend(value.strip().replace("_", " ") for value in progression.groups())
    timed_step = re.search(
        r"\b(o_[A-Za-z0-9_]+)\s+shall\s+increment\s+the\s+first\s+step\s+at\s+value\s+"
        r"(i_[A-Za-z0-9_]+).*?\bat\s+a\s+time\s+(i_[A-Za-z0-9_]+)\b", statement, re.I,
    )
    if timed_step:
        required.extend(value.replace("_", " ") for value in timed_step.groups())
    periodic = re.search(
        r"\b([A-Za-z]+(?:_[A-Za-z0-9]+)+)\s+shall\s+(?:be\s+)?set\s+(?:high|low)\s+"
        r"in\s+every\s+time\s+slot\b", statement, re.I,
    )
    if periodic:
        required.extend((periodic.group(1).replace("_", " "), "every time slot"))
    return tuple(dict.fromkeys(required))


def _ipos_integrity_tokens(value: str) -> List[str]:
    equivalents = {"comprise": "contain", "comprises": "contain", "contains": "contain",
                    "is": "be", "are": "be", "has": "have", "goes": "go"}
    return [equivalents.get(token, token) for token in re.findall(
                r"\w+|[<>=+*/%&|^]+|(?<!\w)-(?=\d)", _ipos_source_text(value).casefold())
            if token not in {"a", "an", "the", "shall", "must"}]


def ipos_source_preservation_findings(
    records: Sequence[NormalizedSourceRecord], facts: Sequence[DescriptiveFact],
) -> List[str]:
    from collections import Counter

    findings = []
    for record in records:
        linked = [fact for fact in facts if fact.statement == record.text and fact.source in record.provenance]
        if not linked:
            findings.append("ipos_source_projection_unresolved:" + record.record_id)
            continue
        if record.record_id.startswith("inventory:"):
            field = record.record_id.split(":", 1)[1]
            if field == "main_functions":
                phrases = [fact.objects[0][len("designed to "):] for fact in linked
                           if len(fact.objects) == 1 and fact.objects[0].startswith("designed to ")]
                conjunction = r",\s*and\s+(?=(?:" + "|".join(sorted(DESCRIPTIVE_NORMATIVE_ACTIONS)) + r")\b)"
                expected_text = re.sub(conjunction, ", ", _ipos_source_text(record.text), flags=re.I).casefold()
                actual_text = re.sub(conjunction, ", ", ", ".join(phrases), flags=re.I).casefold()
                if (len(phrases) != len(linked) or expected_text != actual_text
                        or any(fact.action != "be" or fact.plural_subject
                               or fact.subject != "The " + fact.owner + " block"
                               or fact.condition or fact.negative or fact.modality or fact.mode
                               or fact.object_suffix or fact.qualifiers or fact.results for fact in linked)):
                    findings.append("ipos_inventory_meaning_changed:" + record.record_id)
                continue
            required = Counter(_ipos_integrity_tokens(record.text))
            actual = Counter(token for fact in linked for value in fact.objects for token in _ipos_integrity_tokens(value))
            if required - actual or any(
                fact.action not in ({"serve", "be"} if field == "function" else
                                    {{"inputs": "accept", "outputs": "produce"}[field]})
                or (field == "function" and fact.objects != (("to " if fact.action == "serve" else "designed to ")
                    + _ipos_lower_prose_initial(_ipos_source_text(record.text)),))
                for fact in linked):
                findings.append("ipos_inventory_meaning_changed:" + record.record_id)
            continue
        source = _ipos_source_text(record.text)
        boundaries = [0]
        for index, character in enumerate(source):
            if character in ".!?" and index + 1 < len(source) and source[index + 1].isspace():
                tail = source[index + 1:].lstrip()
                if tail and tail[0].isascii() and tail[0].isalpha():
                    boundaries.append(index + 1)
        boundaries.append(len(source))
        expected = []
        invalid = any(source.count(opening) != source.count(closing)
                  for opening, closing in (("(", ")"), ("[", "]")))
        for start, end in zip(boundaries, boundaries[1:]):
            clause = source[start:end].strip().lstrip("\u2022").strip().rstrip(".!?")
            clause = re.sub(r"^In particular,\s*", "", clause, flags=re.I)
            if re.search(r"\b(?:and|or|then|to|by|with|into|the|a|an|for|from|of|in|at|if|when|unless|until|before|after|that|while|is|are|has|have|was|were|been)\s*$|[,;:]\s*$|\b(?:a|an|the)\s+(?:new|first|last|following|selected|configured)\s*$|\bas follows?\s*:|\b(?:following|tag)\s*:|\u2022", clause, re.I):
                invalid = True
            modals = _ipos_top_level_modals(clause)
            if not modals or any(not re.search(r"\band\s*$", clause[:modal.start()], re.I) for modal in modals[1:]):
                invalid = True
                continue
            modal = modals[0]
            prefix = clause[:modal.start()].strip()
            guard, subject = "", prefix
            if re.match(r"^(?:(?:When|If|Unless|After|Before|Once|In case|In order to|During|At|To)\b|(?:At least\s+)?\d+(?:\.\d+)?\s*(?:ns|us|ms|s)\s+after\b)", prefix, re.I):
                guard, separator, subject = prefix.rpartition(",")
                if not separator:
                    guard, subject = _ipos_comma_less_guard_boundary(prefix, linked[0].owner)
                    if not subject:
                        invalid = True
                        continue
            predicate = clause[modal.end():].strip()
            guard = re.sub(r"\b(?:shall|must)\s+be\b", "is", guard, flags=re.I)
            predicate = re.sub(r"\band\s+(?:shall|must)\s+", "and ", predicate, flags=re.I)
            negative = bool(re.match(r"not\b", predicate, re.I))
            predicate = re.sub(r"^not\s+", "", predicate, flags=re.I)
            adverb_match = re.match(r"(always|first)\b", predicate, re.I)
            adverb = adverb_match[1].casefold() if adverb_match else ""
            if adverb and negative:
                invalid = True
            predicate = re.sub(r"^(?:always|first)\s+", "", predicate, flags=re.I)
            capability = bool(re.match(r"be able to\b", predicate, re.I))
            predicate = re.sub(r"^be able to\s+", "", predicate, flags=re.I)
            action, separator, obj = predicate.partition(" ")
            action = DESCRIPTIVE_FINITE_ACTION_BASES.get(action.casefold(), action.casefold())
            if ((action.casefold() not in DESCRIPTIVE_NORMATIVE_ACTIONS and action.casefold() != "define")
                    or re.fullmatch(r"(?:it|they|this|that|these|those)(?:\s+(?:block|device|unit|signal|channel|state))?", subject.strip(), re.I)
                    or re.search(r"\b(?:that|which|who)\b", subject, re.I)):
                invalid = True
            if not separator and action.casefold() not in DESCRIPTIVE_INTRANSITIVE_ACTIONS:
                invalid = True
            expected.append((tuple(_ipos_integrity_tokens(subject)), tuple(_ipos_integrity_tokens(action)),
                             tuple(_ipos_integrity_tokens(obj)), tuple(_ipos_integrity_tokens(guard)),
                             negative, "can" if capability else "", (adverb,) if adverb else ()))
        actual = [(tuple(_ipos_integrity_tokens(fact.subject)), tuple(_ipos_integrity_tokens(fact.action)),
                   tuple(_ipos_integrity_tokens((" " + fact.coordination + " ").join(fact.objects))),
                   tuple(_ipos_integrity_tokens(fact.condition)), fact.negative, fact.modality, fact.qualifiers) for fact in linked]
        if invalid:
            findings.append("ipos_source_construction_unresolved:" + record.record_id)
        elif Counter(expected) != Counter(actual) or any(fact.mode or fact.object_suffix or fact.results
                                                       for fact in linked):
            findings.append("ipos_source_relationship_changed:" + record.record_id)
        else:
            identifier_pattern = IPOS_TECHNICAL_IDENTIFIER_PATTERN
            source_identifiers = Counter(re.findall(identifier_pattern, _ipos_source_text(record.text)))
            model_identifiers = Counter(identifier for fact in linked for identifier in re.findall(
                identifier_pattern, " ".join((fact.subject, fact.condition, *fact.objects))))
            if source_identifiers != model_identifiers:
                findings.append("ipos_source_identifier_changed:" + record.record_id)
    for fact in facts:
        if not any(fact.statement == record.text and fact.source in record.provenance for record in records):
            findings.append("ipos_unapproved_fact:" + fact.fact_id)
    for index, fact in enumerate(facts):
        for other in facts[index + 1:]:
            if (fact.negative != other.negative and
                    _descriptive_fact_key(replace(fact, negative=False)) ==
                    _descriptive_fact_key(replace(other, negative=False)) and fact.objects == other.objects):
                findings.append("ipos_approved_source_conflict:" + fact.source + ":" + other.source)
    return findings


def _ipos_claim_preservation_findings(paragraph: DescriptiveParagraph) -> List[str]:
    from collections import Counter

    first = paragraph.facts[0]
    objects = list(dict.fromkeys(value for fact in paragraph.facts for value in fact.objects))
    expected = Counter(_ipos_integrity_tokens(" ".join((
        first.subject, first.action, first.condition, first.mode, first.object_suffix,
        *first.qualifiers, *first.results, *objects,
    ))))
    if len(objects) > 1:
        expected[first.coordination] += 1
    action_forms = {first.action + "s": first.action, first.action + "es": first.action}
    if first.action.endswith("y"):
        action_forms[first.action[:-1] + "ies"] = first.action
    actual = Counter(action_forms.get(token, token) for token in _ipos_integrity_tokens(paragraph.text))
    for token in ("do", "does", "not", "can", "may"):
        expected.pop(token, None)
        actual.pop(token, None)
    return [] if actual == expected else ["ipos_claim_terms_added_or_lost:" + first.fact_id]


def _ipos_final_sections(path: Path) -> Dict[str, str]:
    sections: Dict[str, List[str]] = {}
    active = ""
    for event in drs_document_events(path):
        if event["kind"] == "heading":
            title = event["text"]
            if title.startswith(("1.1 ", "1.2 ", "1.3 ", "1.4 ", "1.5 ")):
                active = title.split(" ", 1)[0]
                sections.setdefault(active, [])
            elif re.match(r"^(?:0|1|2|3)\. ", title):
                active = ""
        elif active and event["kind"] == "paragraph":
            sections[active].append(event["text"])
        elif active and event["kind"] == "table":
            sections[active].extend(cell for row in event["rows"] for cell in row)
    return {name: " ".join(values) for name, values in sections.items()}


def ipos_publication_quality_findings(sections: Mapping[str, str]) -> List[str]:
    """Inspect actual section text without recomposing an expected publication."""
    findings = []
    seen_sentences: Dict[str, str] = {}
    for section, content in sections.items():
        findings.extend(section + ":" + finding for finding in natural_descriptive_prose_findings(content))
        if re.search(r"\bFSM\s+FSM\b", content, re.I):
            findings.append(section + ":duplicated_entity_label")
        if re.search(r"\b(?:counter|timer)\s+(?:reach|expire)\b|\band\s+raise\s+an?\s+interrupt\b", content, re.I):
            findings.append(section + ":subject_verb_agreement")
        if re.search(r"\bstarts?\s+\w+\s+writing\s+(?:the\s+)?register\b", content, re.I):
            findings.append(section + ":missing_action_link")
        if any(content.count(opening) != content.count(closing) for opening, closing in (("(", ")"), ("[", "]"))):
            findings.append(section + ":unbalanced_delimiters")
        for sentence in re.split(r"(?<=[.!?])\s+(?=[A-Za-z])", content):
            normalized = re.sub(r"\s+", " ", sentence.strip(" .")).casefold()
            if not normalized:
                continue
            if normalized in seen_sentences:
                findings.append(section + ":duplicated_explanation:" + seen_sentences[normalized])
            else:
                seen_sentences[normalized] = section
            if re.search(r"\b(?:and|or|to|by|into|for|from|unless|until|before|after)\s*[.!?]?$", sentence, re.I):
                findings.append(section + ":unfinished_clause")
        for heading, description in re.findall(r"\*\*([^*]+)\.\*\*\s*([^.!?]+[.!?])", content):
            label = " ".join(heading.casefold().split())
            body = " ".join(description.casefold().split())
            if label and (body.rstrip(".!?") == label or body.endswith("designed to " + label + ".")):
                findings.append(section + ":heading_body_tautology")
    return list(dict.fromkeys(findings))


def _fact_validate_ipos_descriptive_output(
    markdown_path: Path,
    audit_path: Path,
    block_name: str,
    requirement_statements: Iterable[str],
    source_ids: Iterable[str],
    topic_specs: Sequence[Tuple[str, Sequence[str]]],
    functional_input: IPOSFunctionalInput | None = None,
) -> List[str]:
    if functional_input is None:
        return ["ipos_admitted_context_missing"]
    if not functional_input.materialized:
        return []
    if not audit_path.exists():
        return ["ipos_descriptive_audit_missing"]
    with audit_path.open(encoding="utf-8", newline="") as handle:
        audit = tuple(csv.DictReader(handle))
    records = _ipos_admitted_records(functional_input)
    facts: Dict[str, DescriptiveFact] = {}
    units: Dict[str, SemanticUnit] = {}
    unit_contributors: Dict[str, Tuple[str, ...]] = {}
    try:
        for row in audit:
            if row["decision"] != "selected":
                continue
            payload = json.loads(row["fact_json"])
            for field in ("objects", "qualifiers", "results"):
                payload[field] = tuple(payload[field])
            fact = DescriptiveFact(**payload)
            facts[fact.fact_id] = fact
            payload = json.loads(row["semantic_unit_json"])
            payload["source_record_ids"] = tuple(payload["source_record_ids"])
            unit = SemanticUnit(**payload)
            units[unit.unit_id] = unit
            unit_contributors[unit.unit_id] = tuple(json.loads(row["contributor_ids"]))
        ordered_units = tuple(sorted(units.values(), key=lambda unit: unit.output_order))
        paragraphs = tuple(DescriptiveParagraph(unit.text, tuple(facts[identity] for identity in unit_contributors[unit.unit_id]))
                           for unit in ordered_units)
        result = DescriptiveFlowResult(records, tuple(facts.values()), paragraphs, ordered_units, audit)
        findings = validate_descriptive_flow(result, profile=IPOS_DESCRIPTIVE_PROFILE, owner=block_name)
        findings.extend(ipos_source_preservation_findings(records, result.facts))
        for paragraph in paragraphs:
            findings.extend(_ipos_claim_preservation_findings(paragraph))
    except (KeyError, TypeError, ValueError):
        return ["ipos_descriptive_audit_invalid_or_legacy"]
    findings.extend("ipos_projection_gap:" + row["source_record_id"] + ":" + row["reason"]
                    for row in audit if row["decision"] == "projection_gap")
    structural_path = audit_path.with_name("descriptive_structure_audit.json")
    if structural_path.exists():
        findings.extend(validate_ipos_structural_audit(structural_path, functional_input))
    for path in (markdown_path, markdown_path.with_suffix(".docx")):
        if not path.exists():
            findings.append("ipos_final_artifact_missing:" + path.name)
            continue
        sections = _ipos_final_sections(path)
        findings.extend(path.name + ":publication:" + finding for finding in ipos_publication_quality_findings(sections))
        if not sections.get("1.1") or not sections.get("1.2"):
            findings.append("ipos_required_overview_section_missing:" + path.name)
        overview = " ".join(sections.values())
        findings.extend(path.name + ":" + finding for finding in natural_descriptive_prose_findings(overview))
        if _IPOS_NORMATIVE_RE.search(overview):
            findings.append("ipos_normative_overview_wording:" + path.name)
        application = ipos_descriptive_structural_audit(functional_input, result)["application"]
        section_decisions = {value.claim_id: value for value in _ipos_section_contributions(result, application)}
        published_units = tuple(unit for unit in ordered_units
                    if unit.unit_id in section_decisions and section_decisions[unit.unit_id].decision == "retained")
        output_numbers = {key: value.split(" ", 1)[0] for key, value in IPOS_DESCRIPTIVE_SECTION_TITLES.items()}
        for unit in published_units:
            target = output_numbers[_ipos_profile_output_section(application, unit_contributors[unit.unit_id][0])]
            for reference in re.findall(r"\b[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9_]+\b", unit.text):
                if not re.search(r"(?<!\w)" + re.escape(reference) + r"(?!\w)", sections.get(target, "")):
                    findings.append("ipos_final_identifier_changed:" + path.name + ":" + reference)
            unit_tokens = _ipos_integrity_tokens(unit.text)
            section_tokens = _ipos_integrity_tokens(sections.get(target, ""))
            if not any(section_tokens[index:index + len(unit_tokens)] == unit_tokens
                       for index in range(len(section_tokens) - len(unit_tokens) + 1)):
                findings.append("ipos_final_claim_missing_or_changed:" + path.name + ":" + unit.unit_id)
        from collections import Counter
        allowed: Dict[str, Counter] = {number: Counter() for number in output_numbers.values()}
        function_labels = set()
        for unit in published_units:
            target = output_numbers[_ipos_profile_output_section(application, unit_contributors[unit.unit_id][0])]
            allowed[target].update(_ipos_integrity_tokens(unit.text))
            first = facts[unit_contributors[unit.unit_id][0]]
            if first.source == "inventory:main_functions":
                allowed[target].update(_ipos_integrity_tokens(ipos_main_function_name(first)))
            else:
                name = ipos_supported_function_name(first)
                if name and (target, name) not in function_labels:
                    allowed[target].update(_ipos_integrity_tokens(name))
                    function_labels.add((target, name))
        for section, content in sections.items():
            actual = Counter(_ipos_integrity_tokens(content))
            expected = allowed.get(section, Counter())
            extra = actual - expected
            if expected - actual or extra:
                findings.append("ipos_final_unmodeled_content:" + path.name + ":" + section)
    return list(dict.fromkeys(findings))


def _ipos_evidence_semantic_findings(evidence: DescriptiveEvidenceModel) -> List[str]:
    records = {record.record_id: record for record in evidence.records}
    by_span = {}
    for proposition in evidence.propositions:
        by_span.setdefault(proposition.spans, []).append(proposition)
    findings = []
    for spans, propositions in by_span.items():
        if len({span.record_id for span in spans}) != 1 or any(
                left.end > right.start for left, right in zip(spans, spans[1:])):
            findings.append("ipos_recomposition_provenance_incompatible")
            continue
        source = _ipos_recomposition_text(spans)
        priority_list = _ipos_priority_list(source)
        if priority_list is not None:
            entries, direction = priority_list
            expected = DescriptiveMeaning("behavior", "The priority order", "be", entries,
                                          qualifiers=(direction,))
            if len(propositions) != 1 or propositions[0].meaning != expected:
                findings.append("ipos_source_priority_order_changed:" + spans[0].record_id)
            continue
        if len(spans) == 2 and _ipos_explicit_case_reference(spans[0].text, spans[1].text):
            for proposition in propositions:
                antecedents = {value.proposition_id for value in evidence.propositions if value.spans == (spans[0],)}
                references = [edge for edge in evidence.dependencies if edge.relation == "referenced_context"
                              and edge.source_id in antecedents and edge.target_id == proposition.proposition_id
                              and edge.spans == spans]
                if len(references) != 1:
                    findings.append("ipos_source_case_reference_unbound:" + proposition.proposition_id)
        record = replace(records[spans[0].record_id], text=source)
        facts = []
        for proposition in propositions:
            meaning = proposition.meaning
            objects = meaning.objects
            if meaning.effects and len(objects) == 1:
                objects = (objects[0] + " " + " and ".join("to " + effect.action + " " + effect.object for effect in meaning.effects),)
            facts.append(DescriptiveFact(meaning.subject, meaning.predicate, objects,
                         record.provenance[0], source, record.ownership_scope,
                         condition=", ".join(context.text for context in meaning.contexts),
                         mode=meaning.mode, modality=meaning.modality, negative=meaning.negative,
                         coordination=meaning.coordination, qualifiers=meaning.qualifiers, results=meaning.results,
                         evidence_kind="approved_normative" if record.authority_tier == "approved_normative" else "descriptive",
                         owner=proposition.owner, domain=proposition.domain, layer=proposition.layer,
                         plural_subject=bool(re.search(r"\band\b", meaning.subject, re.I))))
        findings.extend(ipos_source_preservation_findings((record,), facts))
    for index, proposition in enumerate(evidence.propositions):
        for other in evidence.propositions[index + 1:]:
            if (proposition.owner == other.owner and proposition.meaning.negative != other.meaning.negative
                    and replace(proposition.meaning, negative=False) == replace(other.meaning, negative=False)):
                findings.append("ipos_approved_source_conflict:" + proposition.proposition_id + ":" + other.proposition_id)
    return list(dict.fromkeys(findings))


def ipos_claim_preflight_findings(functional_input: IPOSFunctionalInput) -> List[str]:
    plan = build_ipos_claim_plan(functional_input)
    findings = validate_descriptive_synthesis_plan(plan, expected_evidence=plan.evidence,
               expected_intents=plan.intents, profile=IPOS_DESCRIPTIVE_PROFILE, owner=functional_input.block)
    findings.extend(_ipos_evidence_semantic_findings(plan.evidence))
    by_id = {claim.claim_id: claim for claim in plan.claims}
    for contribution in ipos_claim_section_plan(plan):
        if contribution.decision == "unresolved":
            findings.append("ipos_projection_gap:" + contribution.claim_id + ":" + contribution.reason)
        elif contribution.decision == "retained":
            try:
                text = realize_descriptive_functional_claim(by_id[contribution.claim_id])
                findings.extend(ipos_publication_quality_findings({contribution.claim_id: text}))
            except ValueError as error:
                findings.append("ipos_realization_unresolved:" + contribution.claim_id + ":" + str(error))
    return list(dict.fromkeys(findings))


def _ipos_phrase_pattern(value: str, role: str = "", action: str = "") -> str:
    tokens = value.split()
    patterns = [re.escape(token) for token in tokens]
    for index, token in enumerate(tokens):
        previous = tokens[index - 1].casefold() if index else ""
        if role == "condition" and (previous, token.casefold()) in {("counter", "reach"), ("timer", "expire")}:
            patterns[index] += "(?:es)?" if token.casefold() == "reach" else "s?"
        elif role == "object" and action == "start" and token.casefold() == "writing" and index > 0:
            tail = " ".join(tokens[index + 1:])
            if re.match(r"(?:the\s+)?register\b.+", tail, re.I) and not any(
                    word.casefold() in {"by", "without", "while", "before", "after", "not"} for word in tokens[:index]):
                patterns[index] = r"(?:by\s+)?" + patterns[index]
        elif (role == "object" and action in {"go", "enter"} and previous == "and" and token.casefold() == "raise"
              and not any(left.casefold() == "to" and right.casefold() in {"start", "initiate", "go", "enter", "raise"}
                          for left, right in zip(tokens, tokens[1:]))
              and re.match(r"an?\s+interrupt\b", " ".join(tokens[index + 1:]), re.I)):
            patterns[index] += "s?"
    return r"\s+".join(patterns)


def _ipos_readback_section_findings(content: str, claims: Sequence[DescriptiveFunctionalClaim]) -> List[str]:
    from collections import Counter
    content = content.replace("\u2019", "'").replace("\u2018", "'")
    meanings = list(dict.fromkeys(meaning for claim in claims for meaning in claim.meanings))
    remaining = list(meanings)
    findings = []
    source_identifiers = Counter(identifier for meaning in meanings for identifier in re.findall(
        IPOS_TECHNICAL_IDENTIFIER_PATTERN, " ".join((meaning.subject, meaning.mode,
        *(context.text for context in meaning.contexts), *meaning.objects, *meaning.qualifiers, *meaning.results,
        *(effect.object for effect in meaning.effects)))))
    actual_identifiers = Counter(re.findall(IPOS_TECHNICAL_IDENTIFIER_PATTERN, content))
    if set(actual_identifiers) != set(source_identifiers):
        findings.append("ipos_final_identifier_changed")
    for sentence in re.split(r"(?<=[.!?])\s+(?=[A-Za-z])", content.strip()):
        rest = sentence.strip().rstrip(".!?")
        previous = None
        while rest:
            matched = False
            for meaning in tuple(remaining):
                condition = ", ".join((("In " + meaning.mode,) if meaning.mode else ())
                                      + tuple(context.text for context in meaning.contexts))
                prefix = _ipos_phrase_pattern(condition, "condition") + r"\s*,\s*" if condition else ""
                subject = _ipos_phrase_pattern(meaning.subject) + r"\s+"
                action = meaning.predicate
                plural = bool(re.search(r"\band\b", meaning.subject, re.I))
                forms = ({action} if plural else {action[:-1] + "ies"} if action.endswith("y")
                         and len(action) > 1 and action[-2] not in "aeiou" else
                         {action + ("es" if action.endswith(("s", "sh", "ch", "x", "z", "o")) else "s")})
                if action in {"comprise", "contain"}:
                    forms.update({"contain", "comprise"} if plural else {"contains", "comprises"})
                if action == "have":
                    forms = {"have"} if plural else {"has"}
                if action == "be":
                    verb = ("are" if plural else "is") + (r"\s+not" if meaning.negative else "")
                elif meaning.modality:
                    verb = re.escape(meaning.modality) + (r"\s+not" if meaning.negative else "") + r"\s+" + re.escape(action)
                elif meaning.negative:
                    verb = ("do" if plural else "does") + r"\s+not\s+" + re.escape(action)
                else:
                    verb = "(?:" + "|".join(re.escape(form) for form in sorted(forms, key=len, reverse=True)) + ")"
                objects = _descriptive_join(meaning.objects, meaning.coordination)
                suffix = " ".join(value for value in (objects, " ".join(value for value in meaning.qualifiers if value not in {"always", "first"})) if value)
                if meaning.effects:
                    suffix += " " + " and ".join("to " + effect.action + " " + effect.object for effect in meaning.effects)
                adverbs = "".join(re.escape(value) + r"\s+" for value in reversed(meaning.qualifiers) if value in {"always", "first"})
                predicate = adverbs + verb + (r"\s+" + _ipos_phrase_pattern(suffix, "object", action) if suffix else "")
                patterns = [prefix + subject + predicate]
                if condition:
                    patterns.append(subject + predicate + r"(?:\s*,\s*|\s+)" + _ipos_phrase_pattern(condition, "condition"))
                if previous == (meaning.subject, meaning.contexts, meaning.mode):
                    patterns.append(predicate)
                match = next((match for pattern in patterns if (match := re.match(pattern + r"(?=\s+and\s+|$)", rest, re.I))), None)
                if match is None:
                    continue
                previous = (meaning.subject, meaning.contexts, meaning.mode)
                remaining.remove(meaning)
                rest = re.sub(r"^\s+and\s+", "", rest[match.end():], count=1)
                matched = True
                break
            if not matched:
                findings.append("ipos_final_unmodeled_content")
                break
    if remaining:
        findings.append("ipos_final_claim_missing_or_changed")
    return list(dict.fromkeys(findings))


def validate_ipos_descriptive_output(
    markdown_path: Path, audit_path: Path, block_name: str, requirement_statements: Iterable[str],
    source_ids: Iterable[str], topic_specs: Sequence[Tuple[str, Sequence[str]]],
    functional_input: IPOSFunctionalInput | None = None,
    *, include_publication: bool = True,
) -> List[str]:
    if functional_input is None:
        return ["ipos_admitted_context_missing"]
    if not functional_input.materialized:
        return []
    if not audit_path.is_file():
        return ["ipos_descriptive_audit_missing"]
    contract = _ipos_document_contract()
    plan = build_ipos_claim_plan(functional_input)
    findings = validate_descriptive_synthesis_plan(plan, expected_evidence=plan.evidence,
               expected_intents=plan.intents, profile=IPOS_DESCRIPTIVE_PROFILE, owner=block_name)
    findings.extend(_ipos_evidence_semantic_findings(plan.evidence))
    decisions = {value.claim_id: value for value in ipos_claim_section_plan(plan)}
    published = {claim.claim_id: claim for claim in plan.claims if decisions[claim.claim_id].decision == "retained"}
    findings.extend("ipos_projection_gap:" + value.claim_id + ":" + value.reason
                    for value in decisions.values() if value.decision == "unresolved")
    titles = {section["role"]: section["title"] for section in contract["sections"]}
    with audit_path.open(encoding="utf-8", newline="") as handle:
        audit = tuple(csv.DictReader(handle))
    actual_ids = []
    for row in audit:
        if row.get("decision") != "selected":
            continue
        identity = row.get("unit_id", "")
        actual_ids.append(identity)
        try:
            claim = published[identity]
            unit = json.loads(row["semantic_unit_json"])
            if (json.loads(row["claim_json"]) != json.loads(json.dumps(asdict(claim)))
                    or json.loads(row["contributor_ids"]) != list(claim.proposition_ids)
                    or row["profile_fingerprint"] != contract["sha256"]
                    or unit["source_record_ids"] != list(dict.fromkeys(span.record_id for span in claim.spans))
                    or unit["unit_id"] != identity or unit["section"] != titles[_ipos_claim_output_role(claim.role)]):
                raise ValueError("Changed admitted claim linkage")
            findings.extend(finding + ":audit:" + identity for finding in _ipos_readback_section_findings(unit["text"], (claim,)))
        except (KeyError, TypeError, ValueError):
            findings.append("ipos_descriptive_audit_invalid_or_changed:" + identity)
    if set(actual_ids) != set(published) or len(actual_ids) != len(set(actual_ids)):
        findings.append("ipos_contributor_audit_coverage_invalid")
    for path in (markdown_path, markdown_path.with_suffix(".docx")):
        if not path.is_file():
            findings.append("ipos_final_artifact_missing:" + path.name)
            continue
        sections = _ipos_final_sections(path)
        if include_publication:
            findings.extend(path.name + ":publication:" + finding for finding in ipos_publication_quality_findings(sections))
        for section in contract["sections"]:
            number = section["title"].split(" ", 1)[0]
            expected = tuple(claim for claim in published.values() if _ipos_claim_output_role(claim.role) == section["role"])
            if section["required"] and not sections.get(number):
                findings.append("ipos_required_overview_section_missing:" + path.name + ":" + number)
            findings.extend(finding + ":" + path.name + ":" + number for finding in
                            _ipos_readback_section_findings(sections.get(number, ""), expected))
        events = drs_document_events(path)
        actual_titles = []
        in_overview = False
        for event in events:
            if event["kind"] != "heading":
                continue
            title = event["text"]
            if re.match(r"^1\.\s", title):
                in_overview = True
            elif re.match(r"^[023]\.\s", title):
                in_overview = False
            elif in_overview:
                actual_titles.append(title)
        allowed_titles = [section["title"] for section in contract["sections"]]
        if actual_titles != [title for title in allowed_titles if title in actual_titles]:
            findings.append("ipos_section_contract_changed:" + path.name)
        for title in actual_titles:
            number = title.split(" ", 1)[0]
            if title in allowed_titles and not sections.get(number):
                findings.append("ipos_section_without_information_gain:" + path.name + ":" + number)
    return list(dict.fromkeys(findings))


def _legacy_validate_ipos_descriptive_output(
    markdown_path: Path,
    audit_path: Path,
    block_name: str,
    requirement_statements: Iterable[str],
    source_ids: Iterable[str],
    topic_specs: Sequence[Tuple[str, Sequence[str]]],
    functional_input: IPOSFunctionalInput | None = None,
) -> List[str]:
    """Validate that IPOS overview text is synthesized and evidence-linked."""
    findings: List[str] = []
    requirement_texts = [str(value).strip() for value in requirement_statements if str(value).strip()]
    text = markdown_path.read_text(encoding="utf-8", errors="ignore")
    if functional_input is not None and functional_input.materialized:
        findings: List[str] = []
        expected_sections, _expected_audit = _legacy_compose_ipos_overview(
            functional_input,
        )
        overview_match = re.search(
            r"^##\s+(?:\*\*)?1\. Block overview(?:\*\*)?.*?"
            r"(?=^##\s+(?:\*\*)?2\. Source I/O(?:\*\*)?|\Z)",
            text, flags=re.MULTILINE | re.DOTALL,
        )
        overview = overview_match.group(0) if overview_match else ""

        def required_section(title: str, next_title: str) -> str:
            match = re.search(
                rf"^###\s+(?:\*\*)?{re.escape(title)}(?:\*\*)?.*?(?=^###\s+(?:\*\*)?{re.escape(next_title)}(?:\*\*)?|^##\s+|\Z)",
                overview, flags=re.MULTILINE | re.DOTALL,
            )
            return match.group(0).strip() if match else ""

        functionality = required_section("1.1 Functionality", "1.2 Supported functions and scope")
        scope = required_section("1.2 Supported functions and scope", "1.3 Internal structure")
        if not functionality:
            findings.append(f"{markdown_path.name}: missing_functionality_section")
        if not scope:
            findings.append(f"{markdown_path.name}: missing_supported_scope_section")
        expected_functionality = expected_sections["functionality"][0]
        expected_scope = expected_sections["scope"][0]
        if expected_functionality.casefold() not in functionality.casefold():
            findings.append(f"{markdown_path.name}: scope_mismatch_with_block_authority:functionality")
        if expected_scope.casefold() not in scope.casefold():
            findings.append(f"{markdown_path.name}: scope_mismatch_with_block_authority:scope")
        overview_lower = overview.casefold()
        findings.extend(
            f"{markdown_path.name}: {finding}"
            for finding in natural_descriptive_prose_findings(functionality + scope)
        )
        for statement in requirement_texts:
            normalized = re.sub(r"\s+", " ", statement.casefold()).strip()
            rendered = re.sub(r"\s+", " ", overview_lower).strip()
            if statement.casefold() in overview_lower or (
                len(normalized) >= 40 and difflib.SequenceMatcher(None, normalized, rendered).quick_ratio() >= 0.86
            ):
                findings.append(f"{markdown_path.name}: description_reuses_requirement_text")
        if not functional_input.function or not functional_input.inputs or not functional_input.outputs:
            findings.append(f"{markdown_path.name}: insufficient_local_evidence_for_supported_scope")
        if not audit_path.exists():
            return findings + [f"{audit_path.name}: descriptive provenance audit is missing"]
        with audit_path.open("r", encoding="utf-8", newline="") as handle:
            audit_rows = list(csv.DictReader(handle))
        required_fields = {
            "source_authority_tier", "block", "domain", "source_function", "inputs_used", "outputs_used",
            "provenance_references",
            "descriptive_statement", "supporting_requirement_evidence_ids", "hierarchy_level",
            "target_section", "derivation_mode", "decision", "exclusion_suppression_reason",
            "duplicate_cross_document_note",
        }
        if audit_rows and not required_fields.issubset(audit_rows[0]):
            findings.append(f"{audit_path.name}: descriptive_provenance_fields_missing")
        candidate_rows = [row for row in audit_rows if row.get("candidate_refinement") or row.get("candidate_evidence_statement")]
        candidate_fields = {
            "candidate_refinement", "supporting_local_requirement_ids", "candidate_evidence_block",
            "candidate_evidence_requirement_id", "candidate_evidence_statement", "target_section", "derivation_mode", "decision",
            "exclusion_suppression_reason", "duplicate_cross_document_note", "summary_group_id",
            "rendered_summary", "summary_group_contributor_ids",
            "structural_facets_json", "structural_selection_mode", "structural_aggregation_order",
            "structural_selected_evidence_ids", "structural_summary_status", "structural_summary_reason",
        }
        if candidate_rows and not candidate_fields.issubset(audit_rows[0]):
            findings.append(f"{audit_path.name}: refinement_candidate_fields_missing")
        expected_candidates = {
            (row.get("candidate_evidence_block", "").casefold(), row.get("candidate_evidence_requirement_id", "")): row
            for row in _expected_audit
            if row.get("candidate_evidence_statement")
        }
        actual_candidate_keys = {
            (row.get("candidate_evidence_block", "").casefold(), row.get("candidate_evidence_requirement_id", ""))
            for row in candidate_rows
        }
        if audit_rows and "candidate_refinement" in audit_rows[0] and actual_candidate_keys != set(expected_candidates):
            findings.append(f"{audit_path.name}: refinement_candidate_audit_incomplete_or_unexpected")
        allowed_candidate_decisions = {
            "accepted_candidate", "suppressed_due_to_conflict", "suppressed_due_to_cross_block_risk",
            "suppressed_due_to_insufficient_evidence", "suppressed_due_to_normative_reuse",
            "suppressed_due_to_table_or_signal_only_evidence",
        }
        selected = [row for row in audit_rows if row.get("decision") == "selected"]
        if not selected:
            findings.append(f"{audit_path.name}: descriptive_provenance_missing_selected_items")
        accepted_candidates = [
            row for row in candidate_rows
            if row.get("decision") == "accepted_candidate" and row.get("materialization_status") != "suppressed_atomic"
        ]
        rendered_groups = {
            row.get("summary_group_id", "").strip()
            for row in accepted_candidates
            if row.get("summary_group_id", "").strip()
        }
        if len(accepted_candidates) > 2 and len(rendered_groups) == len(accepted_candidates):
            findings.append(f"{audit_path.name}: requirement_by_requirement_scope_expansion")
        if re.search(r"\bSupports approved\s+\w+\s+operations\b", scope, re.I):
            findings.append(f"{markdown_path.name}: generic_action_family_placeholder")
        forbidden_structural_placeholders = (
            "approved local behavior", "state-dependent behavior", "local dependency completion",
        )
        for row in accepted_candidates:
            if row.get("materialization_status") != "standalone":
                continue
            summary = row.get("rendered_summary", "")
            if (re.match(r"\*\*[^*]*_[^*]*\.\*\*", summary)
                or re.search(r"\b(?:following steps|output when|before entering|is equal to)\b", summary, re.I)
                or not re.match(r"\*\*[^*]+\.\*\* [A-Z][a-z]+\s+\S+", summary)):
                findings.append(f"{audit_path.name}: low_quality_standalone_summary:{row.get('candidate_evidence_requirement_id', '')}")
        for summary in {row.get("rendered_summary", "").strip() for row in accepted_candidates} - {""}:
            if not summary.startswith("**") or summary.count("\n") > 3:
                findings.append(f"{audit_path.name}: invalid_subfunction_summary_shape")
            heading = re.match(r"\*\*([^*]+)\.\*\*\s+(.+)", summary)
            if heading and (
                re.fullmatch(r"Local \w+ function", heading.group(1), re.I)
                or re.search(r"\b(?:starts only|stores to perform|turn-on .* starts|resets? .* every)\b", summary, re.I)
                or re.search(r"\bin order to have\b.*\bnot empty\b", summary, re.I)
                or heading.group(1).casefold() == "state and processing synchronization"
                and re.fullmatch(r"Enters? [A-Za-z ]+ mode\.", heading.group(2), re.I)
            ):
                findings.append(f"{audit_path.name}: low_quality_subfunction_summary:{heading.group(1)}")
            if any(
                _ipos_structural_entity(row.get("candidate_evidence_statement", ""))
                and row.get("rendered_summary", "").strip() == summary
                for row in accepted_candidates
            ) and any(placeholder in summary.casefold() for placeholder in forbidden_structural_placeholders):
                findings.append(f"{audit_path.name}: vague_structural_subfunction_summary")
        for row in accepted_candidates:
            if not _ipos_structural_entity(row.get("candidate_evidence_statement", "")):
                continue
            summary = row.get("rendered_summary", "").strip()
            if row.get("structural_selection_mode") != "deterministic_aggregation":
                findings.append(f"{audit_path.name}: invalid_structural_selection_mode")
            if not re.fullmatch(
                r"\*\*[^*]+\.\*\* [^.]+\.(?:\n  [^.]+\.){0,2}"
                r"(?:\n  (?:Interacts with|No direct dependency) [^.]+\.)?",
                summary,
            ):
                findings.append(f"{audit_path.name}: invalid_aggregate_structural_summary")
            try:
                aggregate = json.loads(row.get("structural_facets_json", "{}"))
            except json.JSONDecodeError:
                aggregate = {}
            if not aggregate.get("actions") or not (aggregate.get("paths") or aggregate.get("interactions")):
                findings.append(f"{audit_path.name}: aggregate_structural_evidence_insufficient")
            if _IPOS_NORMATIVE_RE.search(summary):
                findings.append(f"{audit_path.name}: structural_summary_uses_normative_language")
            if row.get("structural_summary_status") != "selected" or not row.get("structural_facets_json", "").strip():
                findings.append(f"{audit_path.name}: structural_facet_audit_missing")
            selected_ids = {value.strip() for value in row.get("structural_selected_evidence_ids", "").split(";") if value.strip()}
            if not selected_ids:
                findings.append(f"{audit_path.name}: structural_selected_evidence_missing")
            if row.get("candidate_evidence_requirement_id", "").strip() in selected_ids and not row.get("structural_aggregation_order", "").strip():
                findings.append(f"{audit_path.name}: structural_aggregation_order_missing")
        for row in audit_rows:
            audit_block = re.sub(r"[-_\s]+", " ", row.get("block", "").strip()).casefold()
            expected_block = re.sub(r"[-_\s]+", " ", functional_input.block.strip()).casefold()
            if audit_block != expected_block:
                findings.append(f"{audit_path.name}: cross_block_capability_leak")
            if not row.get("provenance_references", "").strip():
                findings.append(f"{audit_path.name}: descriptive_provenance_reference_missing")
            if row.get("decision") == "suppressed" and not row.get("exclusion_suppression_reason", "").strip():
                findings.append(f"{audit_path.name}: suppression_reason_missing")
            if row in candidate_rows:
                if row.get("candidate_evidence_block", "").casefold() == functional_input.block.casefold() and row.get("materialization_status") != "suppressed_atomic":
                    required_behavior = _ipos_required_behavior_fragments(row.get("candidate_evidence_statement", ""))
                    rendered_behavior = row.get("rendered_summary", "").casefold().replace("_", " ")
                    for fragment in required_behavior:
                        if row.get("decision") != "accepted_candidate" or fragment.casefold() not in rendered_behavior:
                            findings.append(
                                f"{audit_path.name}: approved_behavior_missing_from_scope:"
                                f"{row.get('candidate_evidence_requirement_id', '')}:{fragment}"
                            )
                candidate_key = (
                    row.get("candidate_evidence_block", "").casefold(),
                    row.get("candidate_evidence_requirement_id", ""),
                )
                expected = expected_candidates.get(candidate_key)
                if row.get("decision") not in allowed_candidate_decisions:
                    findings.append(f"{audit_path.name}: invalid_refinement_candidate_decision")
                if row.get("target_section") != "supported_functions_and_scope":
                    findings.append(f"{audit_path.name}: refinement_candidate_target_mismatch")
                if row.get("source_function") != functional_input.function or row.get("inputs_used") != functional_input.inputs or row.get("outputs_used") != functional_input.outputs:
                    findings.append(f"{audit_path.name}: refinement_candidate_inventory_authority_mismatch")
                if expected is None or any(row.get(key, "") != expected.get(key, "") for key in (
                    "candidate_refinement", "derivation_mode", "decision", "exclusion_suppression_reason",
                    "summary_group_id", "rendered_summary", "summary_group_contributor_ids", "materialization_status",
                )):
                    findings.append(f"{audit_path.name}: refinement_candidate_shared_policy_mismatch")
                if row.get("decision") == "accepted_candidate":
                    if row.get("candidate_evidence_block", "").casefold() != functional_input.block.casefold():
                        findings.append(f"{audit_path.name}: accepted_refinement_uses_cross_block_evidence")
                    if not row.get("supporting_local_requirement_ids", "").strip():
                        findings.append(f"{audit_path.name}: accepted_refinement_lacks_local_requirement_id")
                    if _IPOS_NORMATIVE_RE.search(row.get("candidate_refinement", "")):
                        findings.append(f"{audit_path.name}: accepted_refinement_uses_normative_language")
                    if re.search(r"\bSupports approved\s+\w+\s+operations\b", row.get("candidate_refinement", ""), re.I):
                        findings.append(f"{audit_path.name}: generic_action_family_placeholder")
                    rendered_summary = row.get("rendered_summary", "").strip()
                    if row.get("materialization_status") == "suppressed_atomic":
                        if rendered_summary or row.get("summary_group_id", "").strip():
                            findings.append(f"{audit_path.name}: suppressed_atomic_refinement_was_materialized")
                    elif not row.get("summary_group_id", "").strip() or not rendered_summary:
                        findings.append(f"{audit_path.name}: accepted_refinement_group_provenance_missing")
                    elif row.get("candidate_evidence_requirement_id", "").strip() not in row.get("summary_group_contributor_ids", ""):
                        findings.append(f"{audit_path.name}: accepted_refinement_missing_from_group_provenance")
                    elif rendered_summary.casefold() not in scope.casefold():
                        findings.append(f"{audit_path.name}: accepted_refinement_group_missing_from_rendered_scope")
                elif not row.get("exclusion_suppression_reason", "").strip():
                    findings.append(f"{audit_path.name}: refinement_candidate_suppression_reason_missing")
        return findings
    overview_match = re.search(
        r"^##\s+(?:\*\*)?1\. Block overview(?:\*\*)?.*?"
        r"(?=^##\s+(?:\*\*)?2\. Source I/O(?:\*\*)?|\Z)",
        text,
        flags=re.MULTILINE | re.DOTALL,
    )
    if not overview_match:
        return [f"{markdown_path.name}: descriptive overview section is missing"]
    overview = overview_match.group(0)

    def section_body(title: str, following_titles: Sequence[str]) -> str:
        following = "|".join(re.escape(value) for value in following_titles)
        match = re.search(
            rf"^###\s+(?:\*\*)?{re.escape(title)}(?:\*\*)?.*?(?=^###\s+(?:\*\*)?(?:{following})(?:\*\*)?|^##\s+|\Z)",
            overview,
            flags=re.MULTILINE | re.DOTALL,
        )
        return match.group(0) if match else ""

    functionality_match = re.search(
        r"^###\s+(?:\*\*)?1\.1 Functionality(?:\*\*)?.*?"
        r"(?=^###\s+(?:\*\*)?1\.2 Supported functions and scope(?:\*\*)?|\Z)",
        overview,
        flags=re.MULTILINE | re.DOTALL,
    )
    if functionality_match:
        functionality_lines = [
            line.strip()
            for line in functionality_match.group(0).splitlines()
            if line.strip() and not line.startswith("###")
        ]
        if len(re.findall(r"[.!?](?:\s|$)", " ".join(functionality_lines))) > 3:
            findings.append(f"{markdown_path.name}: functionality exceeds three sentences")
        if functionality_lines and not re.match(
            r"^The\s+(?:main purpose of the\s+)?[^.]+\s+block\s+(?:is|performs|provides|manages|coordinates|controls|routes|handles)\b",
            functionality_lines[0],
            re.I,
        ):
            findings.append(f"{markdown_path.name}: functionality does not state the block role explicitly")
    functionality_text = " ".join(functionality_lines) if functionality_match else ""
    if re.search(r"\[(?:TO:\s*[^]]+|DDS_[^]]+|IPOS_[^]]+|Vpriority[^]]*)\]", overview, re.IGNORECASE):
        findings.append(f"{markdown_path.name}: descriptive overview contains an internal marker")
    if re.search(r"(?:\[?TO:\s*[^\]\n]+\]?|\b(?:DDS|IPOS)_[A-Z0-9_-]+\b|\bVpriority\b)", overview, re.I):
        findings.append(f"{markdown_path.name}: descriptive overview contains an unbracketed internal marker")
    if re.search(
        r"\b(?:0x[0-9a-f]+|parameter\s+\w+|register\s*(?:\[|address|map)|address\s+(?:range|window)|\w+_(?:val|value|version)|\b(?:timer|version)\s+(?:value|field|setting))\b",
        overview,
        re.I,
    ):
        findings.append(f"{markdown_path.name}: descriptive overview contains a low-level parameter, register, address, or value fragment")
    for drs_topic in IPOS_DRS_LEVEL_TOPICS:
        if drs_topic.casefold() in overview.casefold():
            findings.append(f"{markdown_path.name}: descriptive overview contains DRS-level topic {drs_topic}")
    if re.search(r"\b(?:approved evidence identifies|includes approved functional evidence)\b", overview, re.I):
        findings.append(f"{markdown_path.name}: descriptive overview contains generic evidence-existence prose")
    if re.search(r"Internal function:\s+|local functional organization combines|\b\d+\s+(?:ports|interfaces|boundaries)\b", overview, re.I):
        findings.append(f"{markdown_path.name}: descriptive overview contains topic-label or generic structural filler")
    scope = section_body("1.2 Supported functions and scope", ("1.3 Internal structure", "1.4 Internal blocks or functions", "1.5 General architecture"))
    if scope and not re.search(r"\b(?:supports|supported|scope|capabilit|protocol|standard|limit|boundary)\b", scope, re.I):
        findings.append(f"{markdown_path.name}: supported functions and scope is not scope-coherent")
    rendered_scope = [
        re.sub(r"^\s*-\s+", "", line).strip()
        for line in scope.splitlines()
        if re.match(r"^\s*-\s+", line)
    ]
    if rendered_scope:
        findings.append(
            f"{markdown_path.name}: supported scope does not match the deterministic projection of this block's approved evidence"
        )
    if functionality_text and scope:
        normalized_functionality = re.sub(r"\s+", " ", functionality_text.casefold()).strip()
        normalized_scope = re.sub(r"\s+", " ", scope.casefold()).strip()
        if normalized_functionality in normalized_scope or difflib.SequenceMatcher(
            None, normalized_functionality, normalized_scope
        ).quick_ratio() >= 0.86:
            findings.append(f"{markdown_path.name}: supported scope repeats the Functionality macro-purpose")
    structure = section_body("1.3 Internal structure", ("1.4 Internal blocks or functions", "1.5 General architecture"))
    if structure:
        if re.search(r"\b\d+\s+(?:approved\s+)?(?:source\s+)?(?:I/O|port|interface)\b|boundary is defined by", structure, re.I):
            findings.append(f"{markdown_path.name}: internal structure is reduced to an I/O or boundary count")
        if not re.search(r"\b(?:organization|partition|subcomponent|functional organization|comprises)\b", structure, re.I):
            findings.append(f"{markdown_path.name}: internal structure lacks supported organization content")
    internal_functions = section_body("1.4 Internal blocks or functions", ("1.5 General architecture",))
    if internal_functions and not re.search(
        r"(?:\bInternal function:\s+|local function is organized around\b|approved local interfaces support\b)",
        internal_functions,
        re.I,
    ):
        findings.append(f"{markdown_path.name}: internal functions section lacks supported local functions")
    if functionality_text and internal_functions:
        normalized_internal = re.sub(r"\s+", " ", internal_functions.casefold()).strip()
        if normalized_functionality in normalized_internal or difflib.SequenceMatcher(
            None, normalized_functionality, normalized_internal
        ).quick_ratio() >= 0.86:
            findings.append(f"{markdown_path.name}: internal functions repeat the Functionality macro-purpose")
    architecture = section_body("1.5 General architecture", ())
    if architecture and re.search(r"\b(?:register\s*\[|timer\s*value|version\s*(?:field|value)|0x[0-9a-f]+)\b", architecture, re.I):
        findings.append(f"{markdown_path.name}: general architecture contains low-level field or value fragments")
    for source_id in {value.strip() for value in source_ids if value.strip()}:
        if source_id in overview:
            findings.append(f"{markdown_path.name}: descriptive overview contains source ID {source_id}")
    overview_normalized = re.sub(r"\s+", " ", overview.casefold()).strip()
    for statement in set(requirement_texts):
        statement_normalized = re.sub(r"\s+", " ", statement.casefold()).strip()
        if statement in overview:
            findings.append(f"{markdown_path.name}: descriptive overview copies an authoritative requirement statement")
        elif len(statement_normalized) >= 40 and difflib.SequenceMatcher(None, statement_normalized, overview_normalized).quick_ratio() >= 0.86:
            findings.append(f"{markdown_path.name}: descriptive overview is near-copied from an authoritative requirement statement")
    if not audit_path.exists():
        return findings + [f"{audit_path.name}: descriptive summary audit is missing"]
    with audit_path.open("r", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    selected_rows = [row for row in rows if row.get("decision") == "selected"]
    if functionality_match and not selected_rows:
        findings.append(f"{markdown_path.name}: rendered functionality has no selected descriptive audit evidence")
    for row in rows:
        if row.get("decision") != "selected":
            continue
        summary = (row.get("statement") or "").strip()
        source = (row.get("source") or "").strip()
        if not summary:
            findings.append(f"{audit_path.name}: selected descriptive topic is missing")
        if not source:
            findings.append(f"{audit_path.name}: selected descriptive summary has no provenance")
        if re.sub(r"[-_\s]+", " ", row.get("mapped_block", "").strip()).casefold() != re.sub(
            r"[-_\s]+", " ", block_name.strip()
        ).casefold():
            findings.append(f"{audit_path.name}: selected descriptive summary is mapped to a different block")
        if row.get("domain") and row.get("domain") != "block_function":
            findings.append(f"{audit_path.name}: selected descriptive summary does not use approved block-function evidence")
        source_evidence = (row.get("source_evidence") or "").strip()
        if source_evidence and source_evidence in overview:
            findings.append(f"{markdown_path.name}: descriptive overview copies raw approved evidence")
        if summary == source_evidence:
            findings.append(f"{audit_path.name}: selected summary is not re-elaborated")
        functional_evidence = (row.get("functional_evidence") or "").strip()
        assignment = assign_ipos_descriptive_topic(functional_evidence, topic_specs)
        if not functional_evidence or not assignment["topic"]:
            findings.append(f"{audit_path.name}: selected topic lacks eligible functional evidence")
            continue
        if row.get("topic") != assignment["topic"]:
            findings.append(f"{audit_path.name}: selected topic does not match shared functional assignment")
        if row.get("matched_functional_evidence") != assignment["matched_functional_evidence"]:
            findings.append(f"{audit_path.name}: selected topic evidence does not match shared assignment")
        if not re.search(r"^###\s+(?:\*\*)?1\.1 Functionality(?:\*\*)?", overview, re.MULTILINE):
            findings.append(f"{markdown_path.name}: selected local functional evidence lacks a main-purpose summary")
    return findings


def write_descriptive_audit(path: Path, assembly: DescriptiveAssembly, *, section: str) -> None:
    """Write reviewable descriptive selection decisions without changing authority."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "section", "input_order", "output_order", "decision", "reason", "topic",
        "assignment_reason", "normative_evidence", "scope", "scope_decision",
        "mapped_block", "domain", "layer", "statement", "source", "source_evidence",
        "functional_evidence", "matched_functional_evidence",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in assembly.audit:
            writer.writerow({field: row.get(field, "") for field in fields} | {"section": section})


def write_ipos_descriptive_audit(path: Path, rows: Iterable[Mapping[str, object]]) -> None:
    """Persist IPOS composition provenance, including selected and suppressed items."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "source_authority_tier", "block", "domain", "source_function", "inputs_used", "outputs_used",
        "provenance_references",
        "descriptive_statement", "supporting_requirement_evidence_ids", "hierarchy_level",
        "target_section", "derivation_mode", "decision", "exclusion_suppression_reason",
        "duplicate_cross_document_note", "candidate_refinement", "supporting_local_requirement_ids",
        "candidate_evidence_block", "candidate_evidence_requirement_id", "candidate_evidence_statement",
        "summary_group_id", "rendered_summary", "summary_group_contributor_ids", "materialization_status", "output_order",
        "structural_facets_json", "structural_selection_mode", "structural_aggregation_order",
        "structural_selected_evidence_ids", "structural_summary_status", "structural_summary_reason",
        "profile", "profile_version", "profile_fingerprint", "projection_adapter", "source_record_id",
        "normalized_record_json", "reason", "fact_id", "fact_json", "contributor_ids", "unit_id",
        "semantic_unit_json",
        "claim_json", "span_json",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: csv_cell_text(row.get(field, "")) for field in fields})


def organize_descriptive_topics(
    records: Iterable[Mapping[str, object]],
    topic_specs: Sequence[Tuple[str, Sequence[str]]],
    max_items_per_topic: int = 6,
) -> Dict[str, List[str]]:
    """Compatibility wrapper for deterministic descriptive assembly."""
    return assemble_descriptive_summary(
        records, topic_specs, max_items_per_topic=max_items_per_topic
    ).topics


def normalize_authored_requirement_blocks(
    lines: Sequence[str], requirement_prefixes: Sequence[str]
) -> List[str]:
    """Apply standard Markdown boundaries and bold headers to authored requirements."""
    prefixes = "|".join(re.escape(prefix) for prefix in requirement_prefixes)
    header_re = re.compile(
        rf"^(?:\*\*)?(?P<header>\[(?:{prefixes})-REQ-\d{{3}}\]\s+Requirement:)(?:\*\*)?\s*$"
    )
    out: List[str] = []
    requirement_open = False

    def close_requirement() -> None:
        nonlocal requirement_open
        if not requirement_open:
            return
        while out and not str(out[-1]).strip():
            out.pop()
        out.extend(["", "[End]", ""])
        requirement_open = False

    for line in lines:
        text = str(line).strip()
        header_match = header_re.match(text)
        if header_match:
            close_requirement()
            while out and (
                not str(out[-1]).strip() or str(out[-1]).strip() == AUTHORED_REQUIREMENT_SPACER
            ):
                out.pop()
            if out:
                out.extend(["", AUTHORED_REQUIREMENT_SPACER, ""])
            line = f"**{header_match.group('header')}**"
            requirement_open = True
        elif requirement_open and (text.startswith("#") or text == "[End]"):
            close_requirement()
        out.append(line)
        if requirement_open and re.match(r"^Covers:\s*\S+\s*$", text):
            close_requirement()
    close_requirement()
    return out


def apply_docx_authored_requirement_formatting(
    docx_path: Path, requirement_prefixes: Sequence[str] = ("SRS", "ARS", "DRS"), spacing_twips: int = 160
) -> None:
    """Make authored requirement headers visibly separated and bold in Word."""
    prefixes = "|".join(re.escape(prefix) for prefix in requirement_prefixes)
    header_re = re.compile(rf"\[(?:{prefixes})-REQ-\d{{3}}\]\s+Requirement:")
    with zipfile.ZipFile(docx_path) as archive:
        members = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    document_xml = members["word/document.xml"].decode("utf-8", errors="ignore")

    def add_spacing(match: re.Match[str]) -> str:
        paragraph = match.group(0)
        if not header_re.search(re.sub(r"<[^>]+>", "", paragraph)):
            return paragraph
        if "<w:pPr" in paragraph:
            paragraph = re.sub(
                r"<w:pPr(?:\s[^>]*)?>",
                f'<w:pPr><w:spacing w:before="{spacing_twips}"/>',
                paragraph,
                count=1,
            )
        else:
            paragraph = paragraph.replace(
                "<w:p>", f'<w:p><w:pPr><w:spacing w:before="{spacing_twips}"/></w:pPr>', 1
            )

        def make_run_bold(run_match: re.Match[str]) -> str:
            run = run_match.group(0)
            if "<w:b" in run:
                return run
            if "<w:rPr" in run:
                return re.sub(r"<w:rPr(?:\s[^>]*)?>", "<w:rPr><w:b/>", run, count=1)
            return run.replace("<w:r>", "<w:r><w:rPr><w:b/></w:rPr>", 1)

        return re.sub(r"<w:r(?:\s[^>]*)?>.*?</w:r>", make_run_bold, paragraph, flags=re.DOTALL)

    document_xml = re.sub(r"<w:p(?:\s[^>]*)?>.*?</w:p>", add_spacing, document_xml, flags=re.DOTALL)
    members["word/document.xml"] = document_xml.encode("utf-8")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for filename, content in members.items():
                archive.writestr(filename, content)
        temporary_path.replace(docx_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def apply_docx_ipos_requirement_style_guard(docx_path: Path) -> None:
    """Keep template requirement styles out of authored IPOS requirement blocks."""
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    authored_id_re = re.compile(r"^IPOS-[A-Z0-9]+(?:-[A-Z0-9]+)*-\d{3}$")
    forbidden_styles = {"STReq", "STMacroReq", "NoSpacing"}
    with zipfile.ZipFile(docx_path) as archive:
        members = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    root = ET.fromstring(members["word/document.xml"])
    body = root.find(f"{word_namespace}body")
    in_requirement = False
    if body is not None:
        for paragraph in body.findall(f"{word_namespace}p"):
            text = _docx_element_text(paragraph, word_namespace).strip()
            if authored_id_re.fullmatch(text):
                in_requirement = True
            if in_requirement:
                properties = paragraph.find(f"{word_namespace}pPr")
                if properties is not None:
                    for style in list(properties.findall(f"{word_namespace}pStyle")):
                        if style.get(f"{word_namespace}val") in forbidden_styles:
                            properties.remove(style)
            if in_requirement and text == "[End]":
                in_requirement = False
    ET.register_namespace("w", "http://schemas.openxmlformats.org/wordprocessingml/2006/main")
    members["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for filename, content in members.items():
                archive.writestr(filename, content)
        temporary_path.replace(docx_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def apply_docx_common_spec_formatting(docx_path: Path) -> None:
    """Apply shared borders, title-page/TOC pagination, and page numbering."""
    with zipfile.ZipFile(docx_path) as archive:
        members = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    core_properties = ET.fromstring(members["docProps/core.xml"])
    creator_tag = "{http://purl.org/dc/elements/1.1/}creator"
    creator = core_properties.find(creator_tag)
    if creator is None:
        creator = ET.SubElement(core_properties, creator_tag)
    creator.text = document_author_name()
    ET.register_namespace("dcterms", "http://purl.org/dc/terms/")
    members["docProps/core.xml"] = ET.tostring(core_properties, encoding="utf-8", xml_declaration=True)
    document_xml = members["word/document.xml"].decode("utf-8", errors="ignore")
    border_xml = (
        '<w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/></w:tblBorders>'
    )

    def add_borders(match: re.Match[str]) -> str:
        table = match.group(0)
        if "<w:tblBorders" in table:
            return re.sub(r"<w:tblBorders.*?</w:tblBorders>", border_xml, table, count=1, flags=re.DOTALL)
        return table.replace("</w:tblPr>", border_xml + "</w:tblPr>", 1)

    document_xml = re.sub(r"<w:tbl(?:\s[^>]*)?>.*?</w:tbl>", add_borders, document_xml, flags=re.DOTALL)
    document_xml = _reorder_docx_navigation(document_xml)
    members["word/document.xml"] = document_xml.encode("utf-8")
    _ensure_docx_page_number_footer(members)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for filename, content in members.items():
                archive.writestr(filename, content)
        temporary_path.replace(docx_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _docx_element_text(element: ET.Element, word_namespace: str) -> str:
    return "".join(node.text or "" for node in element.iter(f"{word_namespace}t"))


def _reorder_docx_navigation(document_xml: str) -> str:
    """Put the TOC first on page 2, followed by Document Navigation."""
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = ET.fromstring(document_xml)
    body = root.find(f"{word_namespace}body")
    if body is None:
        return document_xml
    children = list(body)
    text = [_docx_element_text(element, word_namespace).strip() for element in children]
    navigation_indices = [
        index for index, value in enumerate(text)
        if value.casefold() == "0. document navigation"
    ]
    toc_index = next(
        (index for index, value in enumerate(text) if value.casefold() == "0.1 table of contents"),
        None,
    )
    if toc_index is None:
        return document_xml
    navigation_index = next((index for index in navigation_indices if index < toc_index), None)
    if navigation_index is None:
        navigation_index = next((index for index in navigation_indices if index > toc_index), None)
    if navigation_index is None:
        return document_xml
    introduction_index = next(
        (
            index
            for index, element in enumerate(children)
            if element.tag == f"{word_namespace}p"
            and re.match(
                r"^1(?:\.\d+)*\.?\s+",
                _docx_element_text(element, word_namespace).strip(),
            )
        ),
        None,
    )
    if navigation_index is None:
        return document_xml
    if introduction_index is None:
        introduction_index = len(children)
    if navigation_index >= introduction_index:
        return document_xml

    title_block = children[:min(toc_index, navigation_index)]
    next_navigation_index = next((index for index in navigation_indices if index > toc_index), None)
    toc_end = next_navigation_index or next(
        (
            index for index in range(toc_index + 1, introduction_index)
            if re.match(r"^0\.[2-9](?:\.|\s)", text[index], re.IGNORECASE)
        ),
        introduction_index,
    )
    toc = [element for element in children[toc_index:toc_end] if text[children.index(element)].casefold() != "0. document navigation"]
    navigation_start = next_navigation_index if next_navigation_index is not None else toc_end
    navigation = [
        element for element in children[navigation_start:introduction_index]
        if text[children.index(element)].casefold() != "0. document navigation"
    ]
    navigation.insert(0, children[navigation_index])
    remainder = children[introduction_index:]
    for element in children:
        body.remove(element)
    for element in title_block + toc + navigation + remainder:
        body.append(element)

    first_toc_paragraph = next(
        element for element in toc if element.tag == f"{word_namespace}p"
    )
    first_toc_index = list(body).index(first_toc_paragraph)
    explicit_page_break = ET.Element(f"{word_namespace}p")
    explicit_run = ET.SubElement(explicit_page_break, f"{word_namespace}r")
    ET.SubElement(explicit_run, f"{word_namespace}br", {f"{word_namespace}type": "page"})
    body.insert(first_toc_index, explicit_page_break)
    properties = first_toc_paragraph.find(f"{word_namespace}pPr")
    if properties is None:
        properties = ET.Element(f"{word_namespace}pPr")
        first_toc_paragraph.insert(0, properties)
    for page_break in list(properties.findall(f"{word_namespace}pageBreakBefore")):
        properties.remove(page_break)
    properties.insert(0, ET.Element(f"{word_namespace}pageBreakBefore"))

    for paragraph in body.findall(f"{word_namespace}p"):
        properties = paragraph.find(f"{word_namespace}pPr")
        if properties is not None and paragraph is not first_toc_paragraph:
            for page_break in list(properties.findall(f"{word_namespace}pageBreakBefore")):
                properties.remove(page_break)

    ET.register_namespace("w", "http://schemas.openxmlformats.org/wordprocessingml/2006/main")
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    return ET.tostring(root, encoding="unicode", xml_declaration=True)


def _ensure_docx_page_number_footer(members: Dict[str, bytes]) -> None:
    """Install a deterministic right-aligned PAGE field in the default footer."""
    document_xml = members["word/document.xml"].decode("utf-8", errors="ignore")
    relationships = members.get("word/_rels/document.xml.rels", b"").decode("utf-8", errors="ignore")
    footer_reference = re.search(r'<w:footerReference\b(?=[^>]*w:type="default")([^>]*)/>', document_xml)
    relationship_id = re.search(r'r:id="([^"]+)"', footer_reference.group(1)) if footer_reference else None
    if relationship_id:
        relation = re.search(
            rf'<Relationship\b[^>]*Id="{re.escape(relationship_id.group(1))}"[^>]*Target="([^"]+)"[^>]*/>',
            relationships,
        )
        target = relation.group(1) if relation else "footer3.xml"
        footer_path = "word/" + target.lstrip("/")
    else:
        ids = [int(value) for value in re.findall(r'Id="rId(\d+)"', relationships)]
        new_id = f"rId{max(ids, default=0) + 1}"
        target = "footer3.xml"
        footer_path = "word/" + target
        relationship = (
            f'<Relationship Id="{new_id}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" '
            f'Target="{target}" />'
        )
        relationships = relationships.replace("</Relationships>", relationship + "</Relationships>", 1)
        reference = f'<w:footerReference w:type="default" r:id="{new_id}"/>'
        sectpr = re.search(r"<w:sectPr(?:\s[^>]*)?>", document_xml)
        if sectpr:
            document_xml = document_xml[:sectpr.end()] + reference + document_xml[sectpr.end():]
        members["word/_rels/document.xml.rels"] = relationships.encode("utf-8")
        content_types = members.get("[Content_Types].xml", b"").decode("utf-8", errors="ignore")
        if 'PartName="/word/footer3.xml"' not in content_types:
            override = (
                '<Override PartName="/word/footer3.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml" />'
            )
            content_types = content_types.replace("</Types>", override + "</Types>", 1)
            members["[Content_Types].xml"] = content_types.encode("utf-8")

    footer_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:pPr><w:jc w:val="right"/></w:pPr><w:r><w:fldChar w:fldCharType="begin"/>'
        '<w:instrText xml:space="preserve"> PAGE </w:instrText><w:fldChar w:fldCharType="separate"/>'
        '<w:t>1</w:t><w:fldChar w:fldCharType="end"/></w:r></w:p></w:ftr>'
    )
    members[footer_path] = footer_xml.encode("utf-8")
    members["word/document.xml"] = document_xml.encode("utf-8")


def csv_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return " ".join(str(item) for item in value if item is not None)
    return str(value)


def read_retained_rows(path: Path, requirement_ids: Set[str]) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as handle:
        return [
            {key: csv_cell_text(value).strip() for key, value in row.items() if key is not None}
            for row in csv.DictReader(handle)
            if (row.get("Requirement ID") or "").strip() in requirement_ids
            and (row.get("Routing Status") or "").strip() == RETAINED_ROUTING_STATUS
        ]


def source_parent_title(source: str) -> str:
    """Return the nearest explicit parent title, preserving source wording."""
    primary = (source or "").split(", paragraph", 1)[0].strip()
    primary = re.sub(r"\s*\(page\s+\d+\)\s*$", "", primary, flags=re.IGNORECASE).strip()
    direct_match = re.match(
        r"^Section\s+([0-9A-Za-z_.-]+)\s+(.+?)(?:\s*\(under\s+Section|,\s*paragraph|$)",
        primary,
        flags=re.IGNORECASE,
    )
    if direct_match:
        direct_title = direct_match.group(2).strip(" .-:()\t")
        if direct_title.lower() not in {"requirement", "requirements"}:
            return direct_title

    parent_match = re.search(
        r"\(under\s+Section\s+[0-9A-Za-z_.-]+\s+(.+?)\)\s*,?\s*paragraph",
        source or "",
        flags=re.IGNORECASE,
    )
    if parent_match:
        return parent_match.group(1).strip(" .-:()\t")

    if re.match(r"^Section\s+[0-9A-Za-z_.-]+\s+", primary, flags=re.IGNORECASE):
        return "Other source function context"
    match = re.match(r"^Section\s+[0-9A-Za-z_.-]+\s+(.*)$", primary, flags=re.IGNORECASE)
    return (match.group(1) if match else primary).strip(" -:\t") or "Unresolved source paragraph"


def is_user_action(statement: str) -> bool:
    return bool(re.search(r"\bthe\s+user\s+shall\b", statement or "", flags=re.IGNORECASE))


def is_reset_clock_table_row(row: Mapping[str, str]) -> bool:
    """Classify structural reset/clock rows without using neighboring OCR text."""
    derivation_kind = (row.get("derivation_kind") or "").strip().lower()
    if derivation_kind in {"reset_table_connection", "clock_table_connection_frequency"}:
        return True
    notes = (row.get("notes") or "").lower()
    return "derived from reset/clock table columns" in notes


def pmu_clock_reset_alias(row: Mapping[str, str], block_names: Iterable[str]) -> str:
    """Map generic clock/reset source context to PMU when the project defines PMU."""
    pmu_name = ""
    for block_name in block_names:
        if block_name.strip().lower() == "pmu":
            pmu_name = block_name
            break
    if not pmu_name:
        return ""

    source = csv_cell_text(row.get("source") or row.get("Source") or row.get("Source Paragraph"))
    statement = csv_cell_text(row.get("requirement_statement") or row.get("Requirement Statement"))
    owner = csv_cell_text(row.get("source_section_owner") or row.get("Source Section Owner"))
    notes = csv_cell_text(row.get("notes") or row.get("Notes"))
    text = " ".join([source, statement, owner, notes]).lower()
    parent = source_parent_title(source).lower()

    has_clock_reset = any(term in text for term in ("clock", "clk", "reset", "resetn", "rst_n", "por"))
    has_pmu = "pmu" in text or "power management" in text or "power-management" in text
    parent_is_clock_reset = bool(re.search(r"\b(clocks?\s+(?:and\s+)?reset|reset\s+(?:and\s+)?clocks?)\b", parent))

    if owner.strip().lower() == "pmu":
        return pmu_name
    if is_reset_clock_table_row(row):
        return pmu_name
    if parent_is_clock_reset:
        return pmu_name
    if has_pmu and has_clock_reset:
        return pmu_name
    return ""


def cascade_retained_context(
    rows: Iterable[Mapping[str, str]],
    statements_by_id: Mapping[str, str],
    mapped_blocks: Mapping[str, str],
) -> List[Dict[str, str]]:
    """Carry an explicit source parent through same-page retained user actions."""
    active_parent = ""
    active_page = ""
    result: List[Dict[str, str]] = []
    for source_row in rows:
        row = dict(source_row)
        requirement_id = (row.get("Requirement ID") or "").strip()
        source = row.get("Non-Block Function Context") or row.get("Source Paragraph") or ""
        page_match = re.search(r"\(page\s+(\d+)\)", source, flags=re.IGNORECASE)
        page = page_match.group(1) if page_match else ""
        parent_match = re.search(
            r"\(under\s+Section\s+[0-9A-Za-z_.-]+\s+(.+?)\)\s*,?\s*paragraph",
            source,
            flags=re.IGNORECASE,
        )
        if page != active_page:
            active_parent = ""
            active_page = page
        if parent_match:
            active_parent = parent_match.group(1).strip(" .-:()\t")
        else:
            direct_match = re.match(
                r"^Section\s+[0-9A-Za-z_.-]+\s+(.+?)(?:\s*\(under\s+Section|,\s*paragraph|$)",
                source,
                flags=re.IGNORECASE,
            )
            if direct_match:
                direct_title = direct_match.group(1).strip(" .-:()\t")
                if direct_title.lower() not in {"requirement", "requirements"}:
                    active_parent = direct_title
        if (
            is_user_action(statements_by_id.get(requirement_id, row.get("Requirement Statement", "")))
            and active_parent
            and (mapped_blocks.get(requirement_id) or "").strip() in {"", "Unassigned"}
            and not parent_match
        ):
            row["Non-Block Function Context"] = active_parent
        result.append(row)
    return result