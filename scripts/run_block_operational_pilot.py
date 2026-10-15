"""Exercise an existing approved block in an isolated real-data workspace."""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys

from openpyxl import load_workbook

from approved_snapshot_resolver import resolve_complete_snapshot_id
from validate_downstream_coherence import resolve_downstream_contract
from workflow_routing import ipos_derivative_residuals
from ipos_blocker_review import blocker_rows, read_blocker_review, review_paths, write_blocker_review


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inventory(root: Path, folders: tuple[str, ...]) -> dict:
    return {path.relative_to(root).as_posix(): digest(path)
            for folder in folders for path in sorted((root / folder).rglob("*"))
            if path.is_file() and "__pycache__" not in path.parts}


def run_step(workspace: Path, output: Path, name: str, arguments: list[str]) -> dict:
    result = subprocess.run([sys.executable, "-B", *arguments], cwd=workspace,
                            capture_output=True, text=True, encoding="utf-8", errors="replace")
    log = output / (name + ".log")
    log.write_text(result.stdout + "\n" + result.stderr, encoding="utf-8")
    print(name + ": exit=" + str(result.returncode), flush=True)
    return {"command": arguments, "exit_code": result.returncode, "log": str(log)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot-id", required=True)
    parser.add_argument("--block", required=True)
    parser.add_argument("--kind", choices=("digital", "analog"), default="digital")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    source = Path(__file__).resolve().parents[1]
    output = Path(args.output).resolve()
    if output.exists():
        raise ValueError("Pilot output already exists; preserve prior evidence and choose a new path")
    project_id = json.loads((source / "config/project_context.json").read_text(encoding="utf-8"))["project_name"]
    snapshot_id = resolve_complete_snapshot_id(source, project_id=project_id, snapshot_id=args.snapshot_id)
    folders = ("scripts", "config", "templates", "data/canonical", "artifacts/stage1_requirements",
               "artifacts/stage1_specs", "artifacts/stage2_mirco_arc", "artifacts/stage3_srs",
               "artifacts/stage4_ars", "artifacts/stage5_drs", "artifacts/stage6_digital_ipos",
               "artifacts/stage7_analog_ipos", "artifacts/traceability_reports",
               "artifacts/orchestrator/ipos_blocker_reviews")
    if any(output == source / folder or source / folder in output.parents for folder in folders):
        raise ValueError("Pilot output must not be inside a copied input tree")
    before = inventory(source, folders)
    workspace = output / "workspace"
    output.mkdir(parents=True)
    for folder in folders:
        if (source / folder).exists():
            shutil.copytree(source / folder, workspace / folder, ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    for manifest_path in (workspace / "artifacts/stage2_mirco_arc").glob("*architecture_model_manifest.json"):
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        address = Path(manifest["model_path"])
        if address.is_absolute():
            relative = address.relative_to(source)
            manifest["model_path"] = relative.as_posix()
        manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    contract = resolve_downstream_contract(workspace, snapshot_id)
    functional = contract.ipos_descriptive_input(workspace, block=args.block, domain=args.kind)
    archived_ids = list(functional.descriptive_omitted_claim_ids)
    report = {"snapshot_id": snapshot_id, "block": args.block, "kind": args.kind,
              "source_root": str(source), "workspace": str(workspace), "authority": False,
              "source_fact_approval": False, "steps": [], "review_cases": [],
              "archived_omitted_claim_ids": archived_ids, "status": "IN_PROGRESS"}

    def save() -> None:
        changed = [path for path, value in before.items() if not (source / path).exists() or digest(source / path) != value]
        report["source_input_changes"] = changed
        (output / "report.json").write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        if changed:
            raise RuntimeError("Original source inputs changed during isolated pilot")

    try:
        case_decisions = (("archived", ""), ("clarification", "user clarified"),
                          ("source-update", "need source/spec update - clarification"))
        for case, decision in case_decisions:
            if decision:
                functional = contract.ipos_descriptive_input(workspace, block=args.block, domain=args.kind)
                gap = next(value for value in ipos_derivative_residuals(functional) if value["kind"] == "projection_gap")
                row = blocker_rows(args.block, ["ipos_projection_gap:" + ",".join(gap["source_record_ids"]) + ":" + gap["reason"]], functional)[0]
                row["Flow recommendation"] = "user clarified"
                write_blocker_review(workspace, snapshot_id, args.block, [row])
                workbook_path, _manifest, _companion = review_paths(workspace, args.block)
                workbook = load_workbook(workbook_path)
                workbook.active.cell(2, 10).value = decision
                workbook.active.cell(2, 11).value = "Operational pilot routing only: existing approved evidence unchanged; no new source fact or approval."
                workbook.save(workbook_path)
                workbook.close()
            step = run_step(workspace, output, case + "-ipos", ["scripts/generate_ipos_specs.py", "--kind", args.kind,
                            "--snapshot-id", snapshot_id, "--block", args.block])
            report["steps"].append(step)
            current = contract.ipos_descriptive_input(workspace, block=args.block, domain=args.kind)
            reused = set(archived_ids).issubset(current.descriptive_omitted_claim_ids)
            review = read_blocker_review(workspace, snapshot_id=snapshot_id, block=args.block, require_saved=False) if decision else {}
            residuals = ipos_derivative_residuals(current)
            report["review_cases"].append({"case": case, "decision": decision, "generation_exit": step["exit_code"],
                "archived_omissions_preserved": reused, "omitted_claim_count": len(current.descriptive_omitted_claim_ids),
                "review": review, "residual_count": len(residuals),
                "residual_classes": dict(Counter(row["kind"] for row in residuals)),
                "residuals": residuals})
            save()
            if step["exit_code"] or not reused:
                report["status"] = "REVIEW_CYCLE_FAILED"
                return 1
        phase = run_step(workspace, output, "final-phase", ["scripts/run_final_sysml_phase.py", "--snapshot-id", snapshot_id])
        report["steps"].append(phase)
        readiness_path = workspace / "artifacts/validation/final_readiness.json"
        report["readiness"] = json.loads(readiness_path.read_text(encoding="utf-8")) if readiness_path.exists() else None
        report["status"] = "PASS" if phase["exit_code"] == 0 else "FINAL_PHASE_FAILED"
        return 0 if phase["exit_code"] == 0 else 1
    finally:
        save()
        print("Pilot report: " + str(output / "report.json"), flush=True)


if __name__ == "__main__":
    raise SystemExit(main())