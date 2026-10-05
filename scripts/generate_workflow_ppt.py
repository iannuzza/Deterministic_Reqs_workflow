from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def _workflow_title(repo_root):
    context_path = repo_root / "config" / "project_context.json"
    try:
        context = json.loads(context_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        context = {}
    return str(context.get("project_name") or "Engineering") + " Workflow Overview"


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs, title, bullets, level_map=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    level_map = level_map or {}
    for i, item in enumerate(bullets):
        p = body.add_paragraph() if i > 0 else body.paragraphs[0]
        p.text = item
        p.level = level_map.get(i, 0)
        p.font.size = Pt(20 if p.level == 0 else 16)


def add_stage_table_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Workflow Stages (Main Steps)"

    rows = 8
    cols = 4
    left = Inches(0.4)
    top = Inches(1.3)
    width = Inches(12.5)
    height = Inches(5.5)

    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(2.0)
    table.columns[2].width = Inches(5.4)
    table.columns[3].width = Inches(4.1)

    headers = ["Stage", "Gate", "Main objective", "Primary runner"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h

    data = [
        ("0", "Gate 0", "Ontology baseline and semantic blockers", "scripts/run_stage0_gate0.py"),
        ("1", "Gate 1", "Requirements extraction + taxonomy + RAG checks", "scripts/run_stage1_requirements_gate1.py"),
        ("2", "Gate 2", "Formal specs from Stage 1 requirements", "scripts/run_stage1_specs_gate2.py"),
        ("2a", "Micro-Arch Gate", "Micro-architecture synthesis and crosscheck", "scripts/run_stage2_micro_arc_gate.py"),
        ("3", "Stage 3 Gate", "SRS generation + markdown/latex checks", "scripts/run_stage3_srs_gate.py"),
        ("4", "Stage 4 Gate", "ARS generation + crosscheck", "scripts/run_stage4_ars_gate.py"),
        ("5", "Stage 5 Gate", "DRS generation + crosscheck", "scripts/run_stage5_drs_gate.py"),
    ]

    for r, row in enumerate(data, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    # Light formatting for readability
    for r in range(rows):
        for c in range(cols):
            tf = table.cell(r, c).text_frame
            for p in tf.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
            if r == 0:
                table.cell(r, c).fill.solid()
                table.cell(r, c).fill.fore_color.rgb = RGBColor(220, 230, 241)


def main():
    repo_root = Path(__file__).resolve().parents[1]
    out_path = repo_root / "docs" / "workflow_main_steps.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        _workflow_title(repo_root),
        "Stage-gated process: main steps, gate logic, and execution commands",
    )

    add_bullets_slide(
        prs,
        "What This Workflow Solves",
        [
            "Converts source specifications into traceable SRS/ARS/DRS artifacts",
            "Uses deterministic stage runners and explicit gate validators",
            "Enforces artifact-driven flow after Stage 1 source read",
            "Adds crosscheck evidence at each stage before progression",
        ],
    )

    add_stage_table_slide(prs)

    add_bullets_slide(
        prs,
        "Gate Logic",
        [
            "Each stage must generate required artifacts and pass crosschecks",
            "If a gate fails, loop on the same stage until blockers are resolved",
            "Canonical execution order: 0 -> 1 -> 2 -> 2a -> 3 -> 4 -> 5",
            "Stage 6 architecture comparison is standalone after Stage 5",
        ],
    )

    add_bullets_slide(
        prs,
        "Critical Policy Rules",
        [
            "Stage 1 is the only stage allowed to read source spec directly",
            "Stage 2+ must consume artifacts only (guarded by stage independence checks)",
            "SRS/ARS/DRS entries use one-to-one Covers mapping",
            "Requirement statements use normative wording and unique IDs",
        ],
    )

    add_bullets_slide(
        prs,
        "Main Execution Commands",
        [
            "python scripts/run_stage0_gate0.py",
            "python scripts/run_stage1_requirements_gate1.py",
            "python scripts/run_stage1_specs_gate2.py",
            "python scripts/run_stage2_micro_arc_gate.py",
            "python scripts/run_stage3_srs_gate.py",
            "python scripts/run_stage4_ars_gate.py",
            "python scripts/run_stage5_drs_gate.py",
            "python scripts/workflow_cli.py arch-compare --project-to-compare <project_name>",
        ],
    )

    add_bullets_slide(
        prs,
        "Primary Outputs",
        [
            "Stage 0: ontology, glossary, semantic issues, stage report",
            "Stage 1: requirements summary, taxonomy crosscheck, coverage checks",
            "Stage 2/2a: specs and micro-architecture traceability artifacts",
            "Stage 3/4/5: system/analog/digital requirement specs + trace matrices",
            "Stage 6: cross-project architecture comparison report set",
        ],
    )

    add_bullets_slide(
        prs,
        "Operational Pattern",
        [
            "Run script-first: generator -> crosscheck -> gate",
            "Use stable loops for normal runs; bounded retry loops for recovery",
            "Update docs/templates only when policy or behavior changes",
            "Keep artifacts and docs synchronized after approved fixes",
        ],
    )

    add_bullets_slide(
        prs,
        "Next-Step Usage",
        [
            "Use this deck as kickoff material for stage reviews",
            "Track PASS/FAIL evidence directly from orchestrator reports",
            "Extend with project-specific metrics from latest run outputs",
        ],
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    print(str(out_path))


if __name__ == "__main__":
    main()
