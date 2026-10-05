from pathlib import Path
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


def _workflow_title(repo_root):
	context_path = repo_root / "config" / "project_context.json"
	try:
		context = json.loads(context_path.read_text(encoding="utf-8"))
	except (OSError, json.JSONDecodeError):
		context = {}
	return str(context.get("project_name") or "Engineering") + " Workflow Architecture"


def _set_title(slide, text):
	if slide.shapes.title is not None:
		slide.shapes.title.text = text
		return
	title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.7))
	tf = title_box.text_frame
	tf.clear()
	p = tf.paragraphs[0]
	p.text = text
	p.font.size = Pt(30)
	p.font.bold = True


def _set_subtitle_or_body(slide, lines):
	body = None
	# Prefer subtitle/body placeholders when available.
	for shape in slide.placeholders:
		if shape.is_placeholder and shape.placeholder_format.idx in (1, 2):
			body = shape
			break
	if body is None:
		body = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(12.0), Inches(5.6))

	tf = body.text_frame
	tf.clear()
	for i, line in enumerate(lines):
		p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
		p.text = line
		p.level = 0
		p.font.size = Pt(20 if i == 0 else 18)


def _add_title_slide(prs, workflow_title):
	slide = prs.slides.add_slide(prs.slide_layouts[0])
	_set_title(slide, workflow_title)
	_set_subtitle_or_body(
		slide,
		[
			"Template-based compact view",
			"Agents pipeline + Orchestrator + Crosscheckers + CLI + Skills",
		],
	)


def _add_main_steps_slide(prs):
	slide = prs.slides.add_slide(prs.slide_layouts[1])
	_set_title(slide, "Main Steps (Stage 0 to Stage 6)")
	_set_subtitle_or_body(
		slide,
		[
			"Stage 0: Ontology baseline and Gate 0",
			"Stage 1: Requirements extraction + taxonomy + strict RAG/evidence checks and Gate 1",
			"Stage 2: Specs formalization and Gate 2",
			"Stage 2a: Micro-architecture synthesis + crosscheck (Decision: go) + gate",
			"Stage 3: SRS generation + crosschecks + Stage 3 gate",
			"Stage 4: ARS generation + crosschecks + Stage 4 gate",
			"Stage 5: DRS generation + crosschecks + Stage 5 gate",
			"Stage 6: Cross-project architecture comparison (standalone after Stage 5)",
			"Fail-fast rule: unresolved hard checks stop the pipeline before the next stage",
		],
	)


def _add_compact_diagram_slide(prs):
	slide = prs.slides.add_slide(prs.slide_layouts[5])
	_set_title(slide, "Compact Workflow Diagram")

	# Top strip: Skill and CLI.
	skill = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(0.5),
		Inches(1.1),
		Inches(2.3),
		Inches(0.8),
	)
	skill.text = "Skill Layer\n(Chronicle/Project Rules)"
	skill.fill.solid()
	skill.fill.fore_color.rgb = RGBColor(225, 236, 255)

	cli = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(3.1),
		Inches(1.1),
		Inches(2.0),
		Inches(0.8),
	)
	cli.text = "CLI\nworkflow_cli.py"
	cli.fill.solid()
	cli.fill.fore_color.rgb = RGBColor(225, 244, 225)

	orch = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(5.5),
		Inches(1.0),
		Inches(3.2),
		Inches(1.0),
	)
	orch.text = "Orchestrator\nrun_stage*_gate.py"
	orch.fill.solid()
	orch.fill.fore_color.rgb = RGBColor(255, 239, 213)

	# Agents pipeline row.
	agent_boxes = [
		("Stage 0\nOntology Agent", 0.5),
		("Stage 1\nReq Agent", 2.2),
		("Stage 2/2a\nSpecs+MicroArc", 3.9),
		("Stage 3\nSRS Agent", 5.6),
		("Stage 4\nARS Agent", 7.3),
		("Stage 5\nDRS Agent", 9.0),
		("Stage 6\nCompare Agent", 10.7),
	]
	agent_shapes = []
	for label, x in agent_boxes:
		shp = slide.shapes.add_shape(
			MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
			Inches(x),
			Inches(2.7),
			Inches(1.45),
			Inches(0.95),
		)
		shp.text = label
		shp.fill.solid()
		shp.fill.fore_color.rgb = RGBColor(242, 242, 242)
		agent_shapes.append(shp)

	cross = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(3.9),
		Inches(4.3),
		Inches(5.0),
		Inches(1.0),
	)
	cross.text = "Crosscheckers + Validators\n(run_*_crosscheck_agent.py + validate_stage*_gate.py)"
	cross.fill.solid()
	cross.fill.fore_color.rgb = RGBColor(255, 228, 225)

	outputs = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(9.3),
		Inches(4.3),
		Inches(3.5),
		Inches(1.0),
	)
	outputs.text = "Artifacts + Reports\n(stage_*_result.md, specs, traces)"
	outputs.fill.solid()
	outputs.fill.fore_color.rgb = RGBColor(230, 230, 250)

	# Simple connectors with lines.
	def line(x1, y1, x2, y2):
		ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
		ln.line.width = Pt(1.5)
		ln.line.color.rgb = RGBColor(80, 80, 80)

	line(2.8, 1.5, 5.5, 1.5)  # Skill -> Orchestrator
	line(5.1, 1.5, 5.5, 1.5)  # CLI -> Orchestrator
	line(7.1, 2.0, 7.1, 2.7)  # Orchestrator down

	for idx in range(len(agent_shapes) - 1):
		ax = 0.5 + 1.45 + idx * 1.7
		line(ax, 3.15, ax + 0.25, 3.15)

	line(7.1, 3.65, 7.1, 4.3)  # pipeline -> crosscheckers
	line(8.9, 4.8, 9.3, 4.8)   # crosscheckers -> outputs

	note = slide.shapes.add_textbox(Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.0))
	nt = note.text_frame
	nt.clear()
	p = nt.paragraphs[0]
	p.text = "Flow: Skills and CLI guide the orchestrator, which executes stage agents; crosscheckers and validators gate progress before artifacts are approved."
	p.font.size = Pt(14)


def _add_executive_architecture_slide(prs):
	slide = prs.slides.add_slide(prs.slide_layouts[5])
	_set_title(slide, "Executive Architecture View")

	# Layer 1: user interaction channels.
	user_box = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(0.7),
		Inches(1.1),
		Inches(3.2),
		Inches(1.0),
	)
	user_box.text = "User Interaction\nChat + Task Requests"
	user_box.fill.solid()
	user_box.fill.fore_color.rgb = RGBColor(214, 234, 248)

	skill_box = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(4.3),
		Inches(1.1),
		Inches(3.0),
		Inches(1.0),
	)
	skill_box.text = "Skill Layer\nRules + Prompt Patterns"
	skill_box.fill.solid()
	skill_box.fill.fore_color.rgb = RGBColor(232, 245, 233)

	cli_box = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(7.7),
		Inches(1.1),
		Inches(2.9),
		Inches(1.0),
	)
	cli_box.text = "CLI Layer\nworkflow_cli.py"
	cli_box.fill.solid()
	cli_box.fill.fore_color.rgb = RGBColor(255, 243, 224)

	# Layer 2: orchestration core.
	orch_box = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(2.6),
		Inches(2.7),
		Inches(6.0),
		Inches(1.1),
	)
	orch_box.text = "Orchestrator Core\nrun_stage*_gate.py + gate loop control"
	orch_box.fill.solid()
	orch_box.fill.fore_color.rgb = RGBColor(252, 228, 236)

	# Layer 3: pipeline and quality controls.
	agent_lane = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(0.8),
		Inches(4.2),
		Inches(6.3),
		Inches(1.2),
	)
	agent_lane.text = "Agents Pipeline\nStage 0 -> 1 -> 2 -> 2a -> 3 -> 4 -> 5 -> 6"
	agent_lane.fill.solid()
	agent_lane.fill.fore_color.rgb = RGBColor(245, 245, 245)

	qc_lane = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(7.4),
		Inches(4.2),
		Inches(5.0),
		Inches(1.2),
	)
	qc_lane.text = "Crosscheckers + Validators\nquality gates before accept"
	qc_lane.fill.solid()
	qc_lane.fill.fore_color.rgb = RGBColor(255, 235, 238)

	# Layer 4: outputs.
	output_box = slide.shapes.add_shape(
		MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
		Inches(3.3),
		Inches(5.9),
		Inches(6.0),
		Inches(0.9),
	)
	output_box.text = "Business Outputs: Specs, Traceability, Stage Reports, Comparison Pack"
	output_box.fill.solid()
	output_box.fill.fore_color.rgb = RGBColor(237, 231, 246)

	def line(x1, y1, x2, y2):
		ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
		ln.line.width = Pt(2)
		ln.line.color.rgb = RGBColor(88, 88, 88)

	line(3.9, 1.6, 5.6, 1.6)
	line(7.3, 1.6, 8.6, 1.6)
	line(5.8, 2.1, 5.8, 2.7)
	line(5.8, 3.8, 3.9, 4.2)
	line(6.8, 3.8, 9.1, 4.2)
	line(6.0, 5.4, 6.0, 5.9)


def _add_gui_description_slide(prs):
	slide = prs.slides.add_slide(prs.slide_layouts[5])
	_set_title(slide, "Workflow GUI Description")

	left = Inches(0.6)
	top = Inches(1.15)
	width = Inches(12.0)
	height = Inches(4.9)
	table = slide.shapes.add_table(8, 2, left, top, width, height).table
	table.columns[0].width = Inches(3.4)
	table.columns[1].width = Inches(8.6)

	table.cell(0, 0).text = "GUI area"
	table.cell(0, 1).text = "Description"

	rows = [
		("Workflow Diagram", "Displays Stage 0 to Stage 6 flow; the active stage is highlighted during a run and the diagram can scroll horizontally."),
		("CLI Runs", "Runs deterministic workflow commands and shows live process output, return codes, and stop-at-stage behavior."),
		("Requirement List", "Shows extracted requirement records with horizontal scrolling for long statements and source evidence."),
		("Live Chat Console", "Accepts normal questions and local workflow commands; the input area can be resized vertically."),
		("AI backend", "OpenAI is the default selection. Auto priority is OpenAI, Azure, then Local; the masked API-key field supports session setup."),
		("Local commands", "/help, /run, /status, /validate, /taxonomy, /analyze log, /stop, and /clear run locally without an LLM."),
		("Read-only analysis", "/analyze log and pasted error messages can be analyzed without allowing code fixes."),
	]

	for r, (area, description) in enumerate(rows, start=1):
		table.cell(r, 0).text = area
		table.cell(r, 1).text = description

	for r in range(8):
		for c in range(2):
			tf = table.cell(r, c).text_frame
			for p in tf.paragraphs:
				for run in p.runs:
					run.font.size = Pt(13 if r == 0 else 11)
			if r == 0:
				table.cell(r, c).fill.solid()
				table.cell(r, c).fill.fore_color.rgb = RGBColor(220, 230, 241)

	note = slide.shapes.add_textbox(Inches(0.7), Inches(6.25), Inches(12.0), Inches(0.5))
	nt = note.text_frame
	nt.clear()
	p = nt.paragraphs[0]
	p.text = "The GUI is an operator console: local commands remain available without AI credentials, while normal questions use the selected AI backend."
	p.font.size = Pt(12)


def _add_cli_commands_slide(prs):
	slide = prs.slides.add_slide(prs.slide_layouts[5])
	_set_title(slide, "workflow_cli.py Command Map")

	left = Inches(0.6)
	top = Inches(1.2)
	width = Inches(12.0)
	height = Inches(4.9)
	table = slide.shapes.add_table(8, 2, left, top, width, height).table
	table.columns[0].width = Inches(4.6)
	table.columns[1].width = Inches(7.4)

	table.cell(0, 0).text = "Command"
	table.cell(0, 1).text = "When to use"

	rows = [
		("python scripts/workflow_cli.py --help", "Show command reference and options"),
		("python scripts/workflow_cli.py status", "Inspect stage map and current runner status"),
		("python scripts/workflow_cli.py run --from-stage 0 --to-stage 5", "Run core pipeline in order"),
		("python scripts/workflow_cli.py run --stage 2", "Execute a single stage quickly"),
		("python scripts/workflow_cli.py retry --stage 3 --max-loops 2", "Bounded recovery for failed stages"),
		("python scripts/workflow_cli.py validate --all", "Validate all gates; stop on missing evidence or hard-check failures"),
		("python scripts/workflow_cli.py arch-compare --project-to-compare <project_name>", "Run Stage 6 architecture comparison"),
	]

	for r, (cmd, purpose) in enumerate(rows, start=1):
		table.cell(r, 0).text = cmd
		table.cell(r, 1).text = purpose

	for r in range(8):
		for c in range(2):
			tf = table.cell(r, c).text_frame
			for p in tf.paragraphs:
				for run in p.runs:
					run.font.size = Pt(13 if r == 0 else 12)
			if r == 0:
				table.cell(r, c).fill.solid()
				table.cell(r, c).fill.fore_color.rgb = RGBColor(220, 230, 241)

	note = slide.shapes.add_textbox(Inches(0.7), Inches(6.25), Inches(12.0), Inches(0.5))
	nt = note.text_frame
	nt.clear()
	p = nt.paragraphs[0]
	p.text = "Tip: for predictable runs, use --max-loops-per-stage 1 in normal mode and retry only failing stages."
	p.font.size = Pt(12)


def _add_execution_slide(prs):
	slide = prs.slides.add_slide(prs.slide_layouts[1])
	_set_title(slide, "Execution Shortlist")
	_set_subtitle_or_body(
		slide,
		[
			"python scripts/run_stage0_gate0.py",
			"python scripts/run_stage1_requirements_gate1.py",
			"python scripts/run_stage1_specs_gate2.py",
			"python scripts/run_stage2_micro_arc_gate.py",
			"python scripts/run_stage3_srs_gate.py",
			"python scripts/run_stage4_ars_gate.py",
			"python scripts/run_stage5_drs_gate.py",
			"python scripts/workflow_cli.py arch-compare --project-to-compare <project_name>",
		],
	)


def main():
	repo_root = Path(__file__).resolve().parents[1]
	template = repo_root / "templates" / "Presentation_template.pptx"
	out_path = repo_root / "docs" / "workflow_main_steps_template_v2.pptx"

	prs = Presentation(str(template))
	_add_title_slide(prs, _workflow_title(repo_root))
	_add_main_steps_slide(prs)
	_add_compact_diagram_slide(prs)
	_add_executive_architecture_slide(prs)
	_add_gui_description_slide(prs)
	_add_cli_commands_slide(prs)
	_add_execution_slide(prs)

	out_path.parent.mkdir(parents=True, exist_ok=True)
	prs.save(out_path)
	print(str(out_path))


if __name__ == "__main__":
	main()
