"""Generate the canonical STBIO workflow documentation presentation."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import shutil
import tempfile
import xml.etree.ElementTree as ElementTree
import zipfile

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None


NAVY = (23, 50, 77)
BLUE = (44, 112, 168)
TEAL = (26, 142, 138)
GOLD = (218, 153, 33)
PALE_BLUE = (225, 238, 248)
PALE_TEAL = (221, 242, 239)
PALE_GOLD = (255, 244, 215)
PALE_GRAY = (241, 244, 247)
WHITE = (255, 255, 255)

DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _replace_slide_text(xml: bytes, values: list[str]) -> bytes:
    root = ElementTree.fromstring(xml)
    text_nodes = root.findall(f".//{{{DRAWING_NS}}}t")
    for index, node in enumerate(text_nodes):
        node.text = values[index] if index < len(values) else ""
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)


def generate_without_python_pptx(root: Path, output: Path) -> None:
    """Generate a compatible deck from the checked-in v3 base when pip is unavailable."""
    base = root / "docs" / "workflow_main_steps_template_v3.pptx"
    if not base.exists():
        raise FileNotFoundError(f"Offline presentation base is missing: {base}")
    content = {
        1: ["STBIO AI", "Workflow Main Steps", "Canonical stage-gated requirements engineering workflow", "September 2026"],
        2: ["Agenda", "Authority model", "Source baseline", "Review and merge", "Approved snapshot", "Engineering specifications", "Traceability", "Operator workflow", "Key commands"],
        4: ["Canonical Workflow", "From source evidence to snapshot-backed engineering deliverables"],
        5: ["Outcome", "Source-faithful, reviewable, deterministic, and traceable engineering documentation"],
        6: ["Workflow Documentation", "Authoritative state is canonical SQLite. CSV, JSON, Markdown, XLSX, SysML, and retrieval indexes are derived artifacts."],
        7: ["STBIO AI Workflow Main Steps", "Source fidelity | explicit approval | immutable snapshots | traceable specifications"],
        8: [
            "Canonical Main Steps", "Step", "Primary activity", "Decision point",
            "S0", "Primary source baseline", "Approved bootstrap",
            "S1", "Ingest, OCR, parse, extract requirements", "Gate 1",
            "S2", "Classify and create mapping preview", "Gate 2",
            "S2A-S2C", "Supplementary review, impact analysis, merge approval", "Reviewer approval",
            "S2D-S2E", "Canonical merge and ontology/index refresh", "Canonical revision",
            "S2F-S2G", "Architecture Map Review and mapping approval", "Evidence hashes + reviewer",
            "S2H", "Immutable approved mapping snapshot", "Authoritative consumer boundary",
            "S3-S5", "Generate parallel SRS, ARS, and DRS", "Stage gates",
            "S6-S7", "Generate Digital and Analog IPOS", "Snapshot provenance",
        ],
        9: [
            "Stage-Gated Workflow", "S0", "Source\nBaseline", "Gate", "0", "S1", "Ingest +\nExtract", "S2", "Classify +\nPreview",
            "S2A-S2C", "Review +\nApproval", "S2D-S2E", "Merge +\nRefresh", "S2F-S2G", "Map Review +\nApproval",
            "S2H", "Immutable\nSnapshot", "S3-S5", "SRS / ARS / DRS", "S6-S7", "Digital / Analog IPOS",
            "Every deterministic stage records evidence and stops on its gate failure. Human approval boundaries validate decisions, identities, timestamps, and evidence hashes before authoritative downstream work.",
        ],
        10: [
            "Review, Merge, And Snapshot Controls",
            "S2A: reviewer evaluates supplementary source candidates and evidence.",
            "S2B-S2C: impact analysis identifies affected outputs; an explicit decision authorizes the merge.",
            "S2D-S2H: canonical revisions refresh context, Architecture Map Review completes, and an immutable approved snapshot is frozen.",
        ],
        11: [
            "System Traceability And Coverage", "Stage 1 Source IDs", "Complete ID catalog + location provenance", "Descriptions", "Scoped prose; provenance audit; no Covers",
            "SRS / ARS / DRS", "Parallel derivatives from approved snapshot", "IPOS", "Approved parent Covers + direct source origin", "direct", "direct", "direct",
            "IPOS keeps approved Covers hierarchy and direct Stage 1 source origin. Direct source-to-IPOS requires approved lineage and no meaningful intermediate; descriptions use provenance, not Covers.",
        ],
        12: [
            "Core Workflow Rules",
            "Stage 1 is the only direct source reader. Descriptive sections may reuse its scoped source-backed evidence with provenance; they create no normative Covers links.",
            "SRS, ARS, and DRS are parallel derivatives from the approved snapshot. Covers records only approved requirement lineage, not descriptive dependencies.",
            "IPOS retains approved parent Covers and direct Stage 1 source origin; direct source-to-IPOS requires approved lineage with no meaningful intermediate.",
            "Canonical SQLite stores revisions, approvals, snapshots, and audit records. SRS / ARS / DRS are parallel snapshot derivatives; IPOS keeps approved Covers hierarchy and direct Stage 1 source origin. Direct source-to-IPOS requires approved lineage without a meaningful intermediate.",
            "Approved vocabulary, taxonomy, architecture profile, mapping, and snapshot identity are required for authoritative generation, reports, SysML, and GUI actions.",
            "Execution is deterministic and local-only: one runner and one validator per selected stage, forward-only, with immediate stop on failure.",
            "Rollback preserves history, supersedes prior state, and requires a newly approved snapshot rather than deleting evidence.",
        ],
        13: [
            "Operator Actions And Key Commands",
            "Run python scripts/workflow_gui.py for the review, approval, artifact, and System Traceability workspace.",
            "Run python scripts/workflow_cli.py status to inspect workflow state and locks; use the CLI only for supported deterministic stage execution.",
            "Run python scripts/run_stage1_requirements_gate1.py for extraction and Gate 1, then python scripts/run_stage1_specs_gate2.py for mapping-preview generation and Gate 2.",
            "Use Architecture Map Review to confirm every mapping decision, reviewer identity, required evidence hashes, and ambiguity disposition before approval.",
            "Select an approved immutable snapshot before producing SRS, DRS, ARS, Digital IPOS, Analog IPOS, SysML, or authoritative reports.",
        ],
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx", dir=output.parent) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(base, "r") as source, zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as destination:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if entry.filename.startswith("ppt/slides/slide") and entry.filename.endswith(".xml"):
                    try:
                        slide_number = int(entry.filename.removeprefix("ppt/slides/slide").removesuffix(".xml"))
                    except ValueError:
                        slide_number = 0
                    if slide_number in content:
                        payload = _replace_slide_text(payload, content[slide_number])
                destination.writestr(entry, payload)
        shutil.move(temporary_path, output)
    finally:
        temporary_path.unlink(missing_ok=True)


def project_name(root: Path) -> str:
    try:
        context = json.loads((root / "config" / "project_context.json").read_text(encoding="utf-8"))
        return str(context.get("project_name") or "STBIO AI")
    except (OSError, json.JSONDecodeError):
        return "STBIO AI"


def set_fill(shape, color: tuple[int, int, int]) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*color)


def add_text(slide, text: str, x: float, y: float, width: float, height: float, size: int = 18,
             color: tuple[int, int, int] = NAVY, bold: bool = False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align is not None:
        paragraph.alignment = align
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(*color)
    return box


def add_title(slide, title: str, subtitle: str = "") -> None:
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.78))
    set_fill(band, NAVY)
    band.line.fill.background()
    add_text(slide, title, 0.55, 0.12, 9.8, 0.42, 28, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, 10.0, 0.16, 2.75, 0.3, 11, PALE_BLUE, False, PP_ALIGN.RIGHT)


def add_footer(slide, number: int) -> None:
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(7.05), Inches(12.2), Inches(0.02))
    set_fill(line, PALE_BLUE)
    line.line.fill.background()
    add_text(slide, "STBIO AI | Deterministic, stage-gated engineering workflow", 0.55, 7.12, 8.5, 0.2, 9, BLUE)
    add_text(slide, str(number), 12.15, 7.12, 0.6, 0.2, 9, BLUE, False, PP_ALIGN.RIGHT)


def add_bullets(slide, title: str, bullets: list[str], number: int, subtitle: str = "") -> None:
    add_title(slide, title, subtitle)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.15), Inches(11.6), Inches(5.65))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = RGBColor(*NAVY)
        paragraph.space_after = Pt(16)
        paragraph.level = 0
    add_footer(slide, number)


def add_card(slide, heading: str, body: str, x: float, y: float, width: float, height: float,
             fill: tuple[int, int, int] = PALE_BLUE, accent: tuple[int, int, int] = BLUE):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
    set_fill(card, fill)
    card.line.color.rgb = RGBColor(*accent)
    card.line.width = Pt(1.1)
    add_text(slide, heading, x + 0.16, y + 0.12, width - 0.32, 0.34, 15, accent, True)
    add_text(slide, body, x + 0.16, y + 0.53, width - 0.32, height - 0.64, 11, NAVY)
    return card


def add_connector(slide, source, target) -> None:
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        source.left + source.width,
        source.top + source.height // 2,
        target.left,
        target.top + target.height // 2,
    )
    line.line.color.rgb = RGBColor(*BLUE)
    line.line.width = Pt(1.4)


def add_cover(slide, name: str) -> None:
    background = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    set_fill(background, NAVY)
    background.line.fill.background()
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(5.8), Inches(13.333), Inches(0.22))
    set_fill(stripe, TEAL)
    stripe.line.fill.background()
    add_text(slide, name, 0.8, 1.35, 11.7, 0.65, 18, (152, 211, 207), True)
    add_text(slide, "Workflow Main Steps", 0.76, 2.05, 11.7, 1.0, 40, WHITE, True)
    add_text(slide, "Canonical stage-gated requirements engineering workflow", 0.8, 3.15, 9.5, 0.45, 20, PALE_BLUE)
    add_text(slide, "Source fidelity | Explicit approvals | Immutable snapshots | Traceable specifications", 0.8, 4.25, 10.6, 0.32, 14, (152, 211, 207))
    add_text(slide, "September 2026", 0.8, 6.4, 3.0, 0.3, 12, PALE_BLUE)


def add_stage_map(slide, number: int) -> None:
    add_title(slide, "Canonical Workflow At A Glance", "16 governed transitions")
    stages = [
        ("S0", "Primary\nBaseline", PALE_TEAL), ("S1", "Ingest +\nExtract", PALE_BLUE),
        ("S2", "Classify +\nPreview", PALE_BLUE), ("S2A", "Staging\nReview", PALE_GOLD),
        ("S2B", "Merge\nImpact", PALE_GOLD), ("S2C", "Merge\nApproval", PALE_GOLD),
        ("S2D", "Canonical\nMerge", PALE_TEAL), ("S2E", "Refresh\nIndex", PALE_TEAL),
        ("S2F", "Map\nReview", PALE_GOLD), ("S2G", "Map\nApproval", PALE_GOLD),
        ("S2H", "Approved\nSnapshot", PALE_TEAL), ("S3", "SRS", PALE_BLUE),
        ("S4", "ARS", PALE_BLUE), ("S5", "DRS", PALE_BLUE), ("S6", "Digital\nIPOS", PALE_TEAL),
        ("S7", "Analog\nIPOS", PALE_TEAL),
    ]
    previous = None
    for index, (stage, label, color) in enumerate(stages):
        row, column = divmod(index, 8)
        x, y = 0.58 + column * 1.58, 1.55 + row * 2.0
        shape = add_card(slide, stage, label, x, y, 1.25, 1.08, color, TEAL if color == PALE_TEAL else GOLD if color == PALE_GOLD else BLUE)
        if previous is not None and index % 8:
            add_connector(slide, previous, shape)
        previous = shape
    add_text(slide, "Yellow: human review or approval boundary     Green: canonical-state transition     Blue: deterministic generation", 0.78, 5.35, 11.8, 0.35, 14, NAVY)
    add_text(slide, "Each transition records artifacts and evidence. Authoritative downstream work begins only from the approved immutable snapshot.", 0.78, 5.8, 11.8, 0.5, 17, NAVY)
    add_footer(slide, number)


def add_governance(slide, number: int) -> None:
    add_title(slide, "Authority Model", "What is authoritative and what is derived")
    add_card(slide, "Canonical Authority", "data/canonical/canonical_store.sqlite\nRequirement revisions, reviews, approvals, snapshots, audit events", 0.75, 1.35, 3.8, 2.25, PALE_TEAL, TEAL)
    add_card(slide, "Approval Context", "Approved vocabulary, taxonomy, architecture profile, mapping decision, and immutable snapshot identity", 4.78, 1.35, 3.8, 2.25, PALE_GOLD, GOLD)
    add_card(slide, "Derived Outputs", "Retrieval/index SQLite, CSV, JSON, Markdown, XLSX, SysML, reports, and GUI views are replaceable outputs", 8.8, 1.35, 3.8, 2.25, PALE_BLUE, BLUE)
    add_text(slide, "Authoritative outputs require an approved snapshot selector. Derived files carry provenance but never become the source of truth.", 0.85, 4.35, 11.5, 0.6, 21, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "Audit and rollback preserve history: a rollback supersedes prior state and requires a new approved snapshot; it never deletes evidence.", 1.0, 5.3, 11.2, 0.55, 17, NAVY, False, PP_ALIGN.CENTER)
    add_footer(slide, number)


def add_stage_detail(slide, number: int, title: str, kicker: str, cards: list[tuple[str, str, tuple[int, int, int], tuple[int, int, int]]]) -> None:
    add_title(slide, title, kicker)
    for index, (heading, body, fill, accent) in enumerate(cards):
        row, column = divmod(index, 3)
        add_card(slide, heading, body, 0.7 + column * 4.15, 1.32 + row * 2.35, 3.7, 1.85, fill, accent)
    add_footer(slide, number)


def add_traceability(slide, number: int) -> None:
    add_title(slide, "Traceability And Coverage", "Separate descriptive provenance from normative lineage")
    add_text(slide, "DESCRIPTIVE EVIDENCE", 0.72, 1.12, 4.0, 0.28, 13, TEAL, True)
    source = add_card(slide, "Stage 1 source catalog", "Source-backed descriptions, IDs, and location provenance", 0.72, 1.48, 3.45, 1.05, PALE_BLUE, BLUE)
    prose = add_card(slide, "Descriptive sections", "Scoped, supported prose; no new requirements or normative Covers", 4.92, 1.48, 3.45, 1.05, PALE_TEAL, TEAL)
    audit = add_card(slide, "Provenance audit", "Retain the Stage 1/source reference separately from requirement hierarchy", 9.12, 1.48, 3.45, 1.05, PALE_GOLD, GOLD)
    add_connector(slide, source, prose)
    add_connector(slide, prose, audit)

    add_text(slide, "NORMATIVE REQUIREMENT LINEAGE", 0.72, 3.02, 5.0, 0.28, 13, BLUE, True)
    snapshot = add_card(slide, "Approved snapshot", "IDs, allocation, ownership, and approved lineage", 0.72, 3.38, 3.45, 1.05, PALE_TEAL, TEAL)
    levels = add_card(slide, "SRS / ARS / DRS", "Parallel semantic derivations; execution order is not authority", 4.92, 3.38, 3.45, 1.05, PALE_BLUE, BLUE)
    ipos = add_card(slide, "IPOS requirements", "Covers the approved parent when present; traceability also keeps direct Stage 1 origin", 9.12, 3.38, 3.45, 1.05, PALE_TEAL, TEAL)
    add_connector(slide, snapshot, levels)
    add_connector(slide, levels, ipos)

    add_text(slide, "Direct source-to-IPOS is permitted only for explicitly approved lineage with no meaningful intermediate requirement; do not invent a parent.", 0.82, 5.05, 11.7, 0.42, 16, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "Coverage joins use exact approved IDs. Descriptive source links remain provenance, not requirement coverage.", 0.82, 5.72, 11.7, 0.38, 15, NAVY, False, PP_ALIGN.CENTER)
    add_footer(slide, number)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the v4 workflow documentation presentation.")
    parser.add_argument("--output-path", type=Path, default=None, help="Output .pptx path.")
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    output = args.output_path or root / "docs" / "workflow_main_steps_template_v4.pptx"
    if not output.is_absolute():
        output = root / output
    if Presentation is None:
        generate_without_python_pptx(root, output)
        print(output)
        return
    template = root / "templates" / "Presentation_template.pptx"
    presentation = Presentation(str(template))
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    add_cover(presentation.slides.add_slide(presentation.slide_layouts[6]), project_name(root))
    add_stage_map(presentation.slides.add_slide(presentation.slide_layouts[6]), 2)
    add_governance(presentation.slides.add_slide(presentation.slide_layouts[6]), 3)
    add_stage_detail(presentation.slides.add_slide(presentation.slide_layouts[6]), 4, "S0-S2: Establish A Reviewable Baseline", "Source evidence becomes a governed draft", [
        ("S0 Primary Baseline", "Lock the approved primary source identity, metadata, tag policy, and bootstrap context.", PALE_TEAL, TEAL),
        ("S1 Ingestion", "OCR, parse, preserve source IDs and provenance; construct the integrated requirements corpus.", PALE_BLUE, BLUE),
        ("S2 Classification", "Create deterministic staging objects, classification evidence, specs baseline, and mapping preview.", PALE_BLUE, BLUE),
        ("Gate discipline", "Stage 1 is the only direct source reader. Stage 2+ consume artifacts only; independence guard enforces it.", PALE_GOLD, GOLD),
        ("Source fidelity", "Stable IDs, verbatim source statements, provenance, duplicate detection, and recovery review prevent silent loss.", PALE_GOLD, GOLD),
        ("Review package", "Editable mapping preview, profile draft, approval request, summaries, traceability, and ambiguity dispositions.", PALE_TEAL, TEAL),
    ])
    add_stage_detail(presentation.slides.add_slide(presentation.slide_layouts[6]), 5, "S2A-S2D: Review And Canonical Merge", "Supplementary sources are additive, not a replacement", [
        ("S2A Staging Review", "Review supplementary source candidates and their evidence before accepting any additions.", PALE_GOLD, GOLD),
        ("S2B Merge Impact", "Identify changed requirements, affected mappings, snapshots, specifications, SysML, and reports.", PALE_GOLD, GOLD),
        ("S2C Merge Approval", "Explicit reviewer decision authorizes the non-destructive merge into canonical history.", PALE_GOLD, GOLD),
        ("S2D Canonical Merge", "Create revisions in canonical SQLite. Merge by source_req_id and retain historical audit evidence.", PALE_TEAL, TEAL),
        ("No silent reuse", "Changed corpus evidence invalidates current mapping and snapshot; prior approvals remain historical only.", PALE_GOLD, GOLD),
        ("Incremental refresh", "The dependent workflow is recomputed from the updated baseline without destructive restart.", PALE_TEAL, TEAL),
    ])
    add_stage_detail(presentation.slides.add_slide(presentation.slide_layouts[6]), 6, "S2E-S2H: Approved Architecture Context", "The boundary before authoritative engineering outputs", [
        ("S2E Refresh", "Refresh ontology, vocabulary candidates, taxonomy context, and rebuildable retrieval indexes.", PALE_TEAL, TEAL),
        ("S2F Architecture Map Review", "Reviewer confirms or changes every requirement-to-block mapping, classification, decision, and notes.", PALE_GOLD, GOLD),
        ("S2G Mapping Approval", "Accept only approved or reassigned rows; require matching evidence hashes, reviewer identity, and resolved critical/major ambiguities.", PALE_GOLD, GOLD),
        ("S2H Snapshot", "Freeze immutable approved mapping plus canonical revision, profile, vocabulary, taxonomy, and retrieval metadata.", PALE_TEAL, TEAL),
        ("Mapping evidence", "Stage 2A validates ontology links, role evidence, functions, properties, relationships, and traceability.", PALE_BLUE, BLUE),
        ("Consumer boundary", "Downstream generators resolve an explicit approved snapshot, never a mutable artifact file.", PALE_BLUE, BLUE),
    ])
    add_stage_detail(presentation.slides.add_slide(presentation.slide_layouts[6]), 7, "Stages 3-7: Produce Engineering Deliverables", "All generation is deterministic and snapshot-backed", [
        ("Stage 3 SRS", "Generate system requirements and scoped descriptive sections from the approved snapshot; preserve Stage 1 source provenance.", PALE_BLUE, BLUE),
        ("Stage 4 ARS", "Parallel analog derivation from the approved snapshot and its scoped requirements.", PALE_BLUE, BLUE),
        ("Stage 5 DRS", "Parallel digital/integration derivation; Stage 1 descriptions retain provenance, not Covers.", PALE_BLUE, BLUE),
        ("Stage 6 Digital IPOS", "Block-local requirements keep approved parent Covers and direct Stage 1 source origin in traceability.", PALE_TEAL, TEAL),
        ("Stage 7 Analog IPOS", "Block-local requirements keep approved parent Covers and direct Stage 1 source origin in traceability.", PALE_TEAL, TEAL),
        ("SysML and reports", "Generate derived SysML, traceability reports, XLSX, and documentation with snapshot provenance.", PALE_TEAL, TEAL),
    ])
    add_traceability(presentation.slides.add_slide(presentation.slide_layouts[6]), 8)
    add_bullets(presentation.slides.add_slide(presentation.slide_layouts[6]), "Gate And Approval Controls", [
        "A gate validates the outputs of its preceding deterministic stage and stops advancement on failure.",
        "A human approval validates a review decision, identity, timestamp, evidence hashes, and required dispositions.",
        "No retry loops or backwards transitions are implicit. Correct the blocker, then rerun the affected stage or allowed forward range.",
        "S2B-S2H are authority boundaries coordinated by canonical services and GUI controls, not ordinary CLI generation stages.",
    ], 9, "Fail closed; retain evidence")
    add_bullets(presentation.slides.add_slide(presentation.slide_layouts[6]), "Deterministic Execution", [
        "workflow_cli.py runs a selected executable stage once, then its validator once, in forward order.",
        "A nonzero runner or validator result stops the selected range immediately.",
        "The legacy contiguous CLI path remains available for executable stages; review-only steps route to GUI/service actions.",
        "Local scripts and artifacts implement runtime behavior. Agents define specialized design-time responsibilities; standard execution makes no LLM, network, or cloud calls.",
    ], 10, "Transparent, repeatable, local-only")
    add_bullets(presentation.slides.add_slide(presentation.slide_layouts[6]), "Operator Workflow", [
        "Run source baseline, extraction, and classification; inspect Gate 0, Gate 1, and Gate 2 evidence.",
        "Use the GUI review workspaces for supplementary sources and Architecture Map Review. Save review decisions explicitly.",
        "Approve the mapping only after all required rows, hashes, reviewer identity, and ambiguity dispositions are valid.",
        "Create or select an approved snapshot before generating specifications, IPOS outputs, SysML, or authoritative reports.",
        "Use the System Traceability GUI view and generated reports to inspect coverage, direct links, and gaps.",
    ], 11, "Human decisions remain explicit")
    add_bullets(presentation.slides.add_slide(presentation.slide_layouts[6]), "Key Commands", [
        "python scripts/workflow_gui.py     Desktop review, approval, artifact, and traceability workspace",
        "python scripts/workflow_cli.py status     Report runner/validator state and lock status",
        "python scripts/run_stage1_requirements_gate1.py     Source extraction and Gate 1 flow",
        "python scripts/run_stage1_specs_gate2.py     Specs, mapping preview, and Gate 2 flow",
        "python scripts/run_stage2_micro_arc_gate.py     Validated architecture mapping flow",
        "python scripts/workflow_cli.py validate --all     Validate the available gate set",
    ], 12, "Run from the STBIO_AI repository root")
    add_bullets(presentation.slides.add_slide(presentation.slide_layouts[6]), "Workflow Outcomes", [
        "Source-faithful requirements retain stable identity and evidence provenance from ingestion through deliverables.",
        "Architecture ownership is reviewed, hash-bound, and frozen before it is used to author specifications.",
        "Canonical data, approvals, snapshots, and audit history make each authoritative result reproducible.",
        "Derived artifacts remain inspectable and regenerable without creating competing workflow authority.",
        "The result is a disciplined handoff from source document to traceable system, detailed, analog, and IP specifications.",
    ], 13, "Engineering evidence you can inspect and rerun")

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    print(output)


if __name__ == "__main__":
    main()