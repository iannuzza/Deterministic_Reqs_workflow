import subprocess
import sys
import threading
import time
import csv
import hashlib
from io import StringIO
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from tkinter import font as tkfont
from datetime import datetime
import re
from collections import Counter
from typing import Dict, List, Optional
import os
import json
import shutil
import ctypes
import tempfile
from urllib.parse import unquote
from urllib import request, error

from canonical_store import connect
from approved_snapshot_resolver import resolve_complete_authoritative_input, ApprovedResolverError
from allocation_ledger import read_csv
from mapping_review_sync import sync_mapping_workbook_to_csv
from generate_stage2_specs import _write_mapping_workbook
from ingest_source_spec import _sync_saved_review_workbook
from pre_freeze_coverage import validate_pre_freeze_coverage
from validate_downstream_coherence import (
    approved_generated_document_dependency_graph,
    approved_primary_source_name,
    approved_traceability_hierarchy,
    created_ipos_block_directories,
)
from workflow_routing import runtime_user_name
from workflow_cli import STAGE_ORDER

from s0_source_bootstrap import BootstrapAssessment, approve_candidate, assess_source, persist_candidate

try:
    import pymupdf as fitz
except ImportError:
    fitz = None

try:
    from PIL import Image, ImageTk
except ImportError:
    Image = None
    ImageTk = None

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import PatternFill
except ImportError:
    Workbook = None
    load_workbook = None
    PatternFill = None


REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = REPO_ROOT / "scripts"
PYTHON_EXE = sys.executable
LOG_DIR = REPO_ROOT / "logs"
STARTUP_LOG = LOG_DIR / "workflow_gui_startup.log"
CHAT_LOG = LOG_DIR / "workflow_gui_chat.log"
LOCAL_MODELS_PATH = REPO_ROOT / "config" / "local_models.json"
PROJECT_CONTEXT_PATH = REPO_ROOT / "config" / "project_context.json"
GUI_STATE_PATH = REPO_ROOT / "config" / "gui_state.json"
WORK_AREA_ROOTS = (
    "artifacts",
    "config",
    "docs",
    "local_memory",
    "logs",
    "scripts",
    "specs",
    "templates",
)
TEXT_FILE_SUFFIXES = {".csv", ".json", ".log", ".md", ".mmd", ".py", ".sysml", ".tex", ".txt"}
DOCX_FILE_SUFFIX = ".docx"
XLSX_FILE_SUFFIX = ".xlsx"
MAX_VIEWER_BYTES = 2 * 1024 * 1024
COVERAGE_REPORT_PATH = REPO_ROOT / "artifacts/stage5_drs/requirements_hierarchy_coverage_report.md"
TRACEABILITY_REPORT_DIR = REPO_ROOT / "artifacts/traceability_reports"
AUTHORITY_SCAN_PATH = REPO_ROOT / "artifacts/validation/authority_consistency_scan.json"
DOWNSTREAM_VALIDATION_REPORT_PATH = REPO_ROOT / "artifacts/validation/downstream_coherence_report.json"
SYSML_ROOT = REPO_ROOT / "artifacts/stage2_mirco_arc"
SYSML_BLOCKS_ROOT = SYSML_ROOT / "sysml" / "blocks"
STAGE_RUNNERS = {
    "0": "run_stage0_gate0.py",
    "1": "run_stage1_requirements_gate1.py",
    "2": "run_stage1_specs_gate2.py",
    "2a": "run_stage2_micro_arc_gate.py",
    "3": "run_stage3_srs_gate.py",
    "4": "run_stage4_ars_gate.py",
    "5": "run_stage5_drs_gate.py",
}

STAGE_KEYS = ["0", "1", "2", "2a", "3", "4", "5"]
STAGE_RESULT_PATHS = {
    "0": REPO_ROOT / "artifacts/orchestrator/stage_00_result.md",
    "1": REPO_ROOT / "artifacts/orchestrator/stage_01_result.md",
    "2": REPO_ROOT / "artifacts/orchestrator/stage_02_result.md",
    "2a": REPO_ROOT / "artifacts/orchestrator/stage_02_micro_arch_result.md",
}
SINGLE_STAGE_OPTIONS = (
    ("0", "S0 - Primary Source Baseline"),
    ("1", "S1 - Source Ingestion / OCR / Req-ID Extraction"),
    ("2", "S2 - Deterministic Classification / Specs"),
    ("2a", "S2A - Staging Review / Micro-Architecture"),
    ("2f", "S2F - Architecture Map Review"),
    ("3", "Stage 3 - SRS"),
    ("4", "Stage 4 - ARS"),
    ("5", "Stage 5 - DRS"),
    ("6", "Stage 6 - Digital IPOS"),
    ("7", "Stage 7 - Analog IPOS"),
    ("regenerate_workbook", "Regenerate Review Workbook from CSV"),
    ("arch-compare", "Optional Architecture Comparison"),
)
SINGLE_STAGE_LABELS = tuple(label for _key, label in SINGLE_STAGE_OPTIONS)
SINGLE_STAGE_LABEL_TO_KEY = {label: key for key, label in SINGLE_STAGE_OPTIONS}
RANGE_STAGE_OPTIONS = tuple(option for option in SINGLE_STAGE_OPTIONS if option[0] in {*STAGE_KEYS, "6", "7"})
RANGE_STAGE_LABELS = tuple(label for _key, label in RANGE_STAGE_OPTIONS)
RANGE_STAGE_LABEL_TO_KEY = {label: key for key, label in RANGE_STAGE_OPTIONS}


WORKFLOW_DIAGRAM_NODES = [
    {"id": "s0", "label": "S0 Primary Source\nBootstrap + Approval", "kind": "main", "x": 40, "y": 70},
    {"id": "s1", "label": "S1 OCR / RAG\nIDs + Source Evidence", "kind": "main", "x": 170, "y": 70},
    {"id": "s2", "label": "S2 Specs +\nMapping Preview", "kind": "main", "x": 300, "y": 70},
    {"id": "s2a", "label": "S2A", "kind": "main", "x": 430, "y": 70},
    {"id": "s2b", "label": "S2B", "kind": "main", "x": 560, "y": 70},
    {"id": "s2c", "label": "Mapping Review\nApproval Gate", "kind": "main", "x": 690, "y": 70},
    {"id": "s2d", "label": "Immutable\nSnapshot", "kind": "main", "x": 820, "y": 70},
    {"id": "s2e", "label": "Snapshot-bound\nGeneration", "kind": "main", "x": 950, "y": 70},
    {"id": "stage3", "label": "Stage 3\nSRS", "kind": "main", "x": 1080, "y": 70},
    {"id": "stage4", "label": "Stage 4\nARS", "kind": "main", "x": 1210, "y": 170},
    {"id": "stage5", "label": "Stage 5\nDRS", "kind": "main", "x": 1210, "y": 270},
    {"id": "stage6", "label": "Stage 6\nDigital IPOS", "kind": "main", "x": 1340, "y": 170},
    {"id": "stage7", "label": "Stage 7\nAnalog IPOS", "kind": "main", "x": 1340, "y": 270},
    {"id": "arch_compare", "label": "Optional\nArchitecture Comparison", "kind": "main", "x": 1470, "y": 170},
    {"id": "reports", "label": "Snapshot XLSX\nSRS / ARS / DRS", "kind": "main", "x": 1210, "y": 370},
    {"id": "sysml", "label": "Snapshot SysML\nFull Block Coverage", "kind": "main", "x": 1340, "y": 370},
    {"id": "supp_import", "label": "Supplementary Source\nReview + Import", "kind": "supplementary", "x": 330, "y": 480},
    {"id": "supp_ingest", "label": "Supplementary Ingestion /\nOCR / Parsing / Req-ID Extraction", "kind": "supplementary", "x": 560, "y": 480},
    {"id": "supp_review", "label": "Supplementary Deterministic\nClassification / Review", "kind": "supplementary", "x": 820, "y": 480},
    {"id": "supp_merge", "label": "Supplementary\nMerge Decision", "kind": "supplementary", "x": 1060, "y": 480},
]

WORKFLOW_DIAGRAM_EDGES = [
    ("s0", "s1", "main"), ("s1", "s2", "main"), ("s2", "s2a", "main"),
    ("s2a", "s2b", "main"), ("s2b", "s2c", "main"), ("s2c", "s2d", "main"),
    ("s2d", "s2e", "main"), ("s2e", "stage3", "main"),
    ("stage3", "stage4", "main"), ("stage4", "stage5", "main"),
    ("stage5", "stage6", "main"), ("stage6", "stage7", "main"),
    ("stage5", "arch_compare", "main"),
    ("s2e", "reports", "main"), ("s2e", "sysml", "main"),
    ("s2", "s1", "loopback"), ("s2c", "s2a", "loopback"), ("s2d", "s2c", "loopback"),
    ("supp_import", "supp_ingest", "supplementary"),
    ("supp_ingest", "supp_review", "supplementary"),
    ("supp_review", "supp_merge", "supplementary"),
    ("supp_merge", "s2a", "supplementary"),
    ("supp_merge", "supp_ingest", "loopback"),
    ("supp_merge", "supp_import", "loopback"),
]

WORKFLOW_DIAGRAM_NOTES = (
    "Descriptions: scoped descriptive prose may use Stage 1 source-backed evidence and provenance; it creates no normative Covers links.",
    "Requirements: SRS / ARS / DRS derive in parallel from one approved snapshot; runner order does not confer semantic authority.",
    "IPOS: retain approved Covers hierarchy and direct Stage 1 source origin in traceability. Direct source-to-IPOS requires approved lineage and no meaningful intermediate.",
)

WORKFLOW_NODE_MEANINGS = {
    "s2b": "S2B Merge Impact Analysis: evaluate affected requirements, mappings, profiles, snapshots, and regeneration scope before approval.",
    "s2c": "S2C Mapping Review and Approval Gate: resolve candidate ownership and require explicit review decisions before snapshotting.",
    "s2d": "S2D Immutable Snapshot: freeze the approved canonical revisions, mappings, profiles, and retrieval context for downstream generation.",
    "s2e": "S2E Downstream Generation: consume the approved immutable snapshot for specifications, reports, and SysML outputs.",
    "reports": "Snapshot XLSX Reports: generate derived SRS, ARS, or DRS reports from an approved immutable snapshot.",
    "sysml": "Snapshot SysML: generate the architecture model with full approved requirement coverage for represented blocks.",
    "arch_compare": "Optional Architecture Comparison: run the deterministic comparison against another project after Stage 5.",
    "stage6": "Stage 6 Digital IPOS: generate and validate digital implementation-oriented requirements.",
    "6": "Stage 6 Digital IPOS: generate and validate digital implementation-oriented requirements.",
    "stage7": "Stage 7 Analog IPOS: generate and validate analog implementation-oriented requirements.",
}

WORKFLOW_CLI_MENU_KEYS = frozenset(STAGE_ORDER) | {"arch-compare"}
WORKFLOW_NODE_STAGE_KEYS = {
    "s0": "0", "s1": "1", "s2": "2", "s2a": "2a",
    "stage3": "3", "stage4": "4", "stage5": "5",
    "stage6": "6", "stage7": "7", "arch_compare": "arch-compare",
}
SELECTABLE_WORKFLOW_NODE_IDS = frozenset(
    node_id for node_id, stage_key in WORKFLOW_NODE_STAGE_KEYS.items()
    if stage_key in WORKFLOW_CLI_MENU_KEYS
)
SELECTABLE_WORKFLOW_NODE_FILL = "#d8dadd"


def compact_workflow_scrollregion(bounds, margin: float = 12) -> tuple[float, float, float, float]:
    """Add a small margin around the actual compact diagram content bounds."""
    if not bounds:
        return (0, 0, 0, 0)
    left, top, right, bottom = bounds
    return (left - margin, top - margin, right + margin, bottom + margin)


def compact_workflow_node_y(node, zoom: float = 1.0) -> float:
    """Preserve declared main-flow rows while keeping the compact graph short."""
    return (node["y"] * 0.70 if node["kind"] == "main" else 330) * zoom


def responsive_layout_scale(screen_width: int, screen_height: int) -> float:
    """Scale GUI dimensions proportionally to the available screen resolution."""
    width_scale = screen_width / 1440
    height_scale = screen_height / 900
    return max(0.72, min(1.25, width_scale, height_scale))


def scale_window_geometry(geometry: str, factor: float) -> str:
    """Scale saved window dimensions while preserving its last screen position."""
    match = re.match(r"^(\d+)x(\d+)([+-]\d+)?([+-]\d+)?$", geometry)
    if not match:
        return geometry
    width = max(320, round(int(match.group(1)) * factor))
    height = max(240, round(int(match.group(2)) * factor))
    position = "".join(value or "" for value in match.groups()[2:])
    return f"{width}x{height}{position}"


class WorkflowDiagramRenderer:
    """Reusable scrollable Canvas renderer for the complete workflow graph."""

    COLORS = {
        "main": {"fill": "#eaf0f6", "outline": "#4a6075", "line": "#40566d"},
        "supplementary": {"fill": "#fce4d6", "outline": "#b35c21", "line": "#c26b2d"},
        "loopback": {"fill": "#fff3cd", "outline": "#8a5a12", "line": "#8a5a12"},
        "selectable": {"fill": SELECTABLE_WORKFLOW_NODE_FILL, "outline": "#64686c", "line": "#64686c"},
        "active": {"fill": "#9be59b", "outline": "#2c8a2c", "line": "#2c8a2c"},
    }

    def __init__(self, parent: ttk.Frame, nodes=None, edges=None, on_node_click=None, selectable_nodes=None) -> None:
        self.parent = parent
        self.nodes = nodes or WORKFLOW_DIAGRAM_NODES
        self.edges = edges or WORKFLOW_DIAGRAM_EDGES
        self.on_node_click = on_node_click
        self.selectable_nodes = set(selectable_nodes or ())
        self.node_items = {}
        self.active_stage = None
        self.canvas = tk.Canvas(parent, background="#ffffff", highlightthickness=0, scrollregion=(0, 0, 1500, 620))
        self.vertical_scroll = ttk.Scrollbar(parent, orient=tk.VERTICAL, command=self.canvas.yview)
        self.horizontal_scroll = ttk.Scrollbar(parent, orient=tk.HORIZONTAL, command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=self.vertical_scroll.set, xscrollcommand=self.horizontal_scroll.set)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        self.vertical_scroll.grid(row=0, column=1, sticky="ns")
        self.horizontal_scroll.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        parent.grid_rowconfigure(0, weight=1)
        parent.grid_columnconfigure(0, weight=1)
        self.canvas.bind("<MouseWheel>", self._on_vertical_wheel, add="+")
        self.canvas.bind("<Shift-MouseWheel>", self._on_horizontal_wheel, add="+")
        self._render()
        self._build_legend()

    def refresh(self) -> None:
        """Redraw the graph from the shared node and edge definitions."""
        self.canvas.delete("all")
        self.node_items = {}
        self._render()

    def refresh(self) -> None:
        """Redraw the graph from the shared node/edge definitions."""
        self.canvas.delete("all")
        self.node_items = {}
        self._render()

    def set_active_stage(self, stage: Optional[str]) -> None:
        self.active_stage = stage
        active_style = self.COLORS["active"]
        for node_id, (rect, _label) in self.node_items.items():
            node_stage = WORKFLOW_NODE_STAGE_KEYS.get(node_id)
            if stage is not None and node_stage == stage:
                self.canvas.itemconfig(rect, fill=active_style["fill"], outline=active_style["outline"], width=3)
            else:
                style_key = "selectable" if node_id in SELECTABLE_WORKFLOW_NODE_IDS else next(
                    (node["kind"] for node in self.nodes if node["id"] == node_id), "main"
                )
                style = self.COLORS[style_key]
                self.canvas.itemconfig(rect, fill=style["fill"], outline=style["outline"], width=2)

    def _on_vertical_wheel(self, event):
        self.canvas.yview_scroll(-1 if event.delta > 0 else 1, "units")

    def _on_horizontal_wheel(self, event):
        self.canvas.xview_scroll(-1 if event.delta > 0 else 1, "units")

    @staticmethod
    def _node_size(node):
        return (112, 80) if node["kind"] == "main" else (180, 60)

    def _center(self, node):
        width, height = self._node_size(node)
        return node["x"] + width / 2, node["y"] + height / 2

    def _loopback_points(self, source, target, lane_index):
        source_x, source_y = self._center(source)
        target_x, target_y = self._center(target)
        lane_y = min(source["y"], target["y"]) - 30 - (lane_index * 22)
        return (source_x, source["y"], source_x, lane_y, target_x, lane_y, target_x, target["y"])

    def _render(self) -> None:
        positions = {node["id"]: node for node in self.nodes}
        loopback_index = 0
        for source_id, target_id, edge_kind in self.edges:
            source = positions[source_id]
            target = positions[target_id]
            sx, sy = self._center(source)
            tx, ty = self._center(target)
            if edge_kind == "loopback":
                points = self._loopback_points(source, target, loopback_index)
                loopback_index += 1
            else:
                points = (sx, sy, tx, ty)
            color = self.COLORS[edge_kind]["line"]
            line_options = {"fill": color, "width": 2, "arrow": tk.LAST}
            if edge_kind == "loopback":
                line_options["dash"] = (6, 4)
            self.canvas.create_line(*points, **line_options)
            if edge_kind == "loopback":
                lane_y = points[3]
                self.canvas.create_text(
                    (points[2] + points[4]) / 2,
                    lane_y - 8,
                    text=f"loopback {source_id.upper()} -> {target_id.upper()}",
                    font=("Segoe UI", 8),
                    fill=self.COLORS["loopback"]["line"],
                    anchor=tk.CENTER,
                )
        for node in self.nodes:
            style_key = "selectable" if node["id"] in SELECTABLE_WORKFLOW_NODE_IDS else node["kind"]
            style = self.COLORS[style_key]
            width, height = self._node_size(node)
            rect = self.canvas.create_rectangle(
                node["x"], node["y"], node["x"] + width, node["y"] + height,
                fill=style["fill"], outline=style["outline"], width=2,
                tags=("workflow-node", node["id"]),
            )
            label = self.canvas.create_text(
                node["x"] + width / 2, node["y"] + height / 2, text=node["label"],
                width=width - 12,
                font=("Segoe UI", 9, "bold"), fill="#1f2d3d", justify=tk.CENTER,
                tags=("workflow-node", node["id"]),
            )
            self.node_items[node["id"]] = (rect, label)
            if self.on_node_click and node["id"] in self.selectable_nodes:
                self.canvas.tag_bind(node["id"], "<Button-1>", lambda _event, value=node["id"]: self.on_node_click(value))
        for index, note in enumerate(WORKFLOW_DIAGRAM_NOTES):
            self.canvas.create_text(
                24,
                570 + index * 18,
                text=note,
                width=1580,
                anchor=tk.NW,
                font=("Segoe UI", 9),
                fill="#2f3b46",
            )
        self.set_active_stage(self.active_stage)

    def _build_legend(self) -> None:
        legend = ttk.Frame(self.parent)
        bounds = self.canvas.bbox("all")
        if bounds:
            margin = 12
            self.canvas.configure(
                scrollregion=(bounds[0] - margin, bounds[1] - margin, bounds[2] + margin, bounds[3] + margin)
            )
        legend.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        ttk.Label(legend, text="Legend:").pack(side=tk.LEFT, padx=(4, 10))
        for label, kind in (
            ("CLI-selectable", "selectable"),
            ("Automatic flow", "main"),
            ("Supplementary flow", "supplementary"),
            ("Loopback / rework", "loopback"),
        ):
            style = self.COLORS[kind]
            swatch = tk.Canvas(legend, width=24, height=14, highlightthickness=0)
            swatch.create_rectangle(2, 2, 22, 12, fill=style["fill"], outline=style["outline"], width=2)
            swatch.pack(side=tk.LEFT, padx=(4, 3))
            ttk.Label(legend, text=label).pack(side=tk.LEFT, padx=(0, 12))


def build_gui_review_context(repo_root: Path) -> dict[str, str]:
    """Return read-only workflow context for GUI display and deterministic tests."""
    context: dict[str, str] = {
        "revision": "unknown",
        "snapshot_id": "none",
        "impact_status": "unknown",
        "vocabulary_taxonomy": "unknown",
        "primary_source": "not configured",
        "source_context": "primary",
        "primary_tagging": "unknown",
        "approved_req_prefix": "not recorded",
        "scanner": "not run",
    }
    context_path = repo_root / "config" / "project_context.json"
    try:
        project_context = json.loads(context_path.read_text(encoding="utf-8")) if context_path.exists() else {}
        context["primary_source"] = str(project_context.get("source_spec_path") or "not configured")
        id_rules = project_context.get("requirement_id_rules", {})
        if isinstance(id_rules, dict):
            context["primary_tagging"] = "yes" if bool(id_rules.get("tagged_source_mode")) else "no"
            patterns = id_rules.get("source_req_id_patterns") or id_rules.get("table_req_id_patterns") or []
            if isinstance(patterns, list) and patterns:
                pattern = str(patterns[0]).strip()
                prefix = re.sub(r"\\d.*$", "", pattern).rstrip("_-")
                if prefix:
                    context["approved_req_prefix"] = prefix
    except (OSError, ValueError, TypeError):
        project_context = {}
    profile_path = repo_root / "config" / "stage2_mirco_arc_profile.json"
    try:
        profile = json.loads(profile_path.read_text(encoding="utf-8")) if profile_path.exists() else {}
        approval = profile.get("approval", {})
        context["revision"] = str(approval.get("approved_at") or "profile revision not recorded")
    except (OSError, ValueError, TypeError):
        pass
    try:
        project_id = str(project_context.get("project_name") or repo_root.name)
        resolved, selection = resolve_complete_authoritative_input(repo_root, "gui", project_id=project_id)
        context["snapshot_id"] = resolved.snapshot_id
        context["impact_status"] = "approved"
        context["vocabulary_taxonomy"] = f"snapshot_hash={resolved.snapshot_hash}"
        manifest_path = repo_root / "artifacts/traceability_reports/requirement_allocation_materialization.json"
        ledger_path = repo_root / "artifacts/traceability_reports/requirement_allocation_ledger.csv"
        if not manifest_path.exists() or not ledger_path.exists():
            raise ApprovedResolverError("allocation ledger materialization is missing")
        ledger_rows = read_csv(ledger_path)
        if any(row.get("snapshot_id") != resolved.snapshot_id for row in ledger_rows):
            raise ApprovedResolverError("allocation ledger snapshot does not match selected snapshot")
        context["allocation_rows"] = str(len(ledger_rows))
        connection = connect(repo_root)
        source = connection.execute("SELECT profile_payload_json FROM profile_revisions WHERE profile_kind = 'source_baseline' AND approval_state = 'approved' ORDER BY approved_at DESC LIMIT 1").fetchone()
        connection.close()
        if source:
            payload = json.loads(source[0])
            if context["primary_tagging"] == "unknown":
                context["primary_tagging"] = "yes" if payload.get("status") == "approved" else "no"
            if context["approved_req_prefix"] == "not recorded":
                context["approved_req_prefix"] = str(payload.get("approved_prefix") or payload.get("proposed_prefix") or "not recorded")
    except (OSError, ValueError, TypeError, KeyError, ApprovedResolverError):
        context["snapshot_id"] = "unavailable"
        context["impact_status"] = "blocked: no complete approved snapshot/materialization"
    if (repo_root / "artifacts/stage1_requirements/integrated_requirements.csv").exists():
        context["source_context"] = "primary + approved supplementary baseline"
    scan_path = repo_root / "artifacts/validation/authority_consistency_scan.json"
    try:
        scan = json.loads(scan_path.read_text(encoding="utf-8"))
        context["scanner"] = f"{scan.get('finding_count', 0)} diagnostic findings"
    except (OSError, ValueError, TypeError):
        pass
    return context


def format_gui_review_context(context: dict[str, str]) -> str:
    return (
        f"Revision: {context['revision']} | Snapshot: {context['snapshot_id']} | Impact: {context['impact_status']}\n"
        f"Primary: {context['primary_source']} | Tagged: {context['primary_tagging']} | Req-ID prefix: {context['approved_req_prefix']}\n"
        f"Authority scanner: {context['scanner']} (diagnostic only)"
    )


class WorkflowGUI:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Requirement AI workflow")
        self.layout_scale = responsive_layout_scale(root.winfo_screenwidth(), root.winfo_screenheight())
        try:
            tk_scaling = float(root.tk.call("tk", "scaling"))
            root.tk.call("tk", "scaling", tk_scaling * self.layout_scale)
        except (tk.TclError, TypeError, ValueError):
            pass
        self.gui_state = self._load_gui_state()
        self.verbose_trace_enabled = bool(self.gui_state.get("verbose_trace", True))
        saved_layout = self.gui_state.get("last_user_layout")
        saved_geometry = saved_layout.get("geometry") if isinstance(saved_layout, dict) else None
        default_width = round(1180 * self.layout_scale)
        default_height = round(860 * self.layout_scale)
        saved_scale = (
            saved_layout.get("layout_scale", self.gui_state.get("layout_scale", 1.0))
            if isinstance(saved_layout, dict)
            else self.gui_state.get("layout_scale", 1.0)
        )
        try:
            geometry_scale = self.layout_scale / max(0.1, float(saved_scale))
        except (TypeError, ValueError):
            geometry_scale = self.layout_scale
        stored_geometry = saved_geometry or self.gui_state.get("geometry")
        self.root.geometry(
            scale_window_geometry(stored_geometry, geometry_scale)
            if stored_geometry
            else f"{default_width}x{default_height}"
        )
        self.root.minsize(round(980 * self.layout_scale), round(700 * self.layout_scale))
        self.root.resizable(True, True)
        self.root.protocol("WM_DELETE_WINDOW", self._close_window)
        self._save_state_after_id = None
        self._gui_layout_ready = False
        self.is_running = False
        self.current_process = None
        self.stage_rects = {}
        self.stage_labels = {}
        self.active_stage = None
        self.hover_stage = None
        self.stage_failure_reasons = {}
        self.compact_workflow_zoom = 1.0
        self.ai_running = False
        self.chat_history = []
        self.ai_request_count = 0
        self.ai_input_tokens = 0
        self.ai_output_tokens = 0
        self.ai_context_limit = int(os.getenv("AI_CONTEXT_WINDOW_TOKENS", "8192"))
        self.ai_last_backend = "-"
        self.ai_last_model = "-"
        self.ai_last_cost = None
        self.local_models = []
        self.help_popup = None
        self.help_popup_label = None
        self.help_popup_owner = None
        self.help_popup_owner_widget = None
        self._active_popup = None
        self._active_popup_kind = None
        self._install_hover_popup_controller()
        self.pdf_reader_window = None
        self.pdf_reader_document = None
        self.pdf_reader_path = None
        self.pdf_reader_page = 0
        self.pdf_reader_zoom = 1.0
        self.pdf_reader_photo = None
        self.pdf_reader_find_var = None
        self.pdf_reader_find_status_var = None
        self.pdf_reader_find_page = None
        self.workspace_root = REPO_ROOT.parent
        self.stage6_dialog = None
        self.table_review_dialog = None
        self.architecture_approval_dialog = None
        self.coverage_selected_link = None
        self.traceability_selected_edge = None
        self.traceability_selected_block = None
        self.sysml_selected_connection = None
        self.sysml_selected_block = None
        self.sysml_source_zoom = 1.0
        self.stage_meanings = {
            "supp": (
                    "Supplementary Source Review: stage, review, and approve added requirements before the integrated baseline is refreshed."
            ),
            "0": "Stage 0: Ontology baseline and semantic blocker analysis.",
            "1": "Stage 1: Requirements extraction, taxonomy update, and coverage checks.",
            "2": (
                "Stage 2: Formal specification generation and Architecture Map preview. "
                "Complete Architecture Map Review before running Stage 2A."
            ),
            "map": (
                "Architecture Map Review: Review every mapping-preview row. Only approved or "
                "reassigned rows, with current evidence and review hashes, allow Stage 2A to run."
            ),
            "2a": (
                "Stage 2a: Micro-architecture synthesis and architecture crosscheck. "
                "Requires current Architecture Map Review approval."
            ),
            "3": "Stage 3: SRS generation with markdown/latex and gate checks.",
            "4": "Stage 4: ARS generation, mapping checks, and gate validation.",
            "5": "Stage 5: DRS generation, traceability checks, and gate validation.",
            "6": "Stage 6: Generate and validate Digital IPOS from the approved snapshot.",
        }

        self._build_ui()
        self.root.after(250, self._restore_gui_layout)

    @staticmethod
    def _gui_pane_names() -> tuple[str, ...]:
        return (
            "main_vertical_pane",
            "top_horizontal_pane",
            "work_pane",
            "visual_pane",
            "chat_vertical_pane",
            "sysml_vertical_pane",
            "sysml_horizontal_pane",
            "traceability_pane",
        )

    @staticmethod
    def _is_sash_pane(pane) -> bool:
        """Return whether a registered layout widget supports sash operations."""
        return callable(getattr(pane, "sashpos", None)) and callable(getattr(pane, "panes", None))

    @staticmethod
    def _load_gui_state() -> dict:
        try:
            if GUI_STATE_PATH.exists():
                state = json.loads(GUI_STATE_PATH.read_text(encoding="utf-8"))
                return state if isinstance(state, dict) else {}
        except (OSError, ValueError, TypeError):
            pass
        return {}

    def _restore_gui_layout(self) -> None:
        panes = {
            name: getattr(self, name, None) for name in self._gui_pane_names()
        }
        defaults = {
                "main_vertical_pane": [300],
                "top_horizontal_pane": [420],
            "work_pane": [300],
            "visual_pane": [500],
            "chat_vertical_pane": [420],
            "sysml_vertical_pane": [500],
            "sysml_horizontal_pane": [300],
            "traceability_pane": [180],
        }
        legacy_names = {
            "main_vertical_pane": "main_vertical",
            "visual_pane": "visual_horizontal",
            "sysml_vertical_pane": "sysml_vertical",
            "traceability_pane": "traceability_vertical",
        }
        saved_layout = self.gui_state.get("last_user_layout")
        if not isinstance(saved_layout, dict):
            saved_layout = self.gui_state
        saved = saved_layout.get("sashes", {})
        saved_scale = saved_layout.get("layout_scale", self.gui_state.get("layout_scale", 1.0))
        try:
            sash_scale = self.layout_scale / max(0.1, float(saved_scale))
        except (TypeError, ValueError):
            sash_scale = self.layout_scale
        for name, pane in panes.items():
            if not self._is_sash_pane(pane):
                continue
            positions = saved.get(name, saved.get(legacy_names.get(name, ""), defaults[name]))
            positions = positions if isinstance(positions, list) else [positions]
            try:
                for index, position in enumerate(positions):
                    position = round(int(position) * sash_scale)
                    if self.gui_state.get("layout_saved_by_user") or self.gui_state.get("last_user_layout"):
                        pane.sashpos(index, max(0, position))
                    elif position > 20:
                        pane.sashpos(index, position)
                    elif index < len(defaults[name]):
                        pane.sashpos(index, defaults[name][index])
            except (AttributeError, tk.TclError, TypeError, ValueError):
                pass
        self.root.after(500, self._restore_collapsed_panes)
        self.root.after(750, self._reapply_saved_gui_sashes)
        if GUI_STATE_PATH.exists():
            try:
                saved_at = datetime.fromtimestamp(GUI_STATE_PATH.stat().st_mtime).astimezone().isoformat(timespec="seconds")
                self.root.after(
                    0,
                    self._append_log_ui,
                    f"[{datetime.now().astimezone().isoformat(timespec='seconds')}] [LAYOUT] "
                    f"Restored last saved GUI layout from {GUI_STATE_PATH.as_posix()} "
                    f"(file_time={saved_at}, geometry={self.root.geometry()})",
                )
            except OSError:
                pass

    def _reapply_saved_gui_sashes(self) -> None:
        """Apply user-saved sash positions after Tk has laid out all panes."""
        if not self.gui_state.get("layout_saved_by_user") and not self.gui_state.get("last_user_layout"):
            return
        saved_layout = self.gui_state.get("last_user_layout")
        if not isinstance(saved_layout, dict):
            saved_layout = self.gui_state
        saved = saved_layout.get("sashes", {})
        saved_scale = saved_layout.get("layout_scale", self.gui_state.get("layout_scale", 1.0))
        try:
            sash_scale = self.layout_scale / max(0.1, float(saved_scale))
        except (TypeError, ValueError):
            sash_scale = self.layout_scale
        self.root.update_idletasks()
        applied = {}
        for name in self._gui_pane_names():
            pane = getattr(self, name, None)
            positions = saved.get(name)
            if not self._is_sash_pane(pane) or not isinstance(positions, list):
                continue
            try:
                for index, position in enumerate(positions):
                    pane.sashpos(index, max(0, round(int(position) * sash_scale)))
                applied[name] = [pane.sashpos(index) for index in range(max(0, len(pane.panes()) - 1))]
            except (AttributeError, tk.TclError, TypeError, ValueError):
                continue
        self.root.after(
            0,
            self._append_log_ui,
            f"[{datetime.now().astimezone().isoformat(timespec='seconds')}] [LAYOUT] "
            f"Reapplied saved pane layout from {GUI_STATE_PATH.as_posix()} sashes={applied}",
        )

    def _restore_collapsed_panes(self) -> None:
        """Recover any persisted pane whose sash was restored at an unusable edge."""
        if self.gui_state.get("layout_saved_by_user"):
            return
        for pane_name, default_position in (
            ("main_vertical_pane", 300),
            ("top_horizontal_pane", 420),
            ("work_pane", 300),
            ("visual_pane", 500),
            ("chat_vertical_pane", 420),
            ("sysml_vertical_pane", 500),
            ("sysml_horizontal_pane", 300),
            ("traceability_pane", 180),
        ):
            pane = getattr(self, pane_name, None)
            if not self._is_sash_pane(pane):
                continue
            try:
                if len(pane.panes()) <= 1:
                    continue
                sash = pane.sashpos(0)
                extent = pane.winfo_width() if pane.cget("orient") == tk.HORIZONTAL else pane.winfo_height()
                if extent > 0:
                    minimum = 120 if pane_name in {"main_vertical_pane", "top_horizontal_pane"} else 80
                    fallback = min(default_position, max(minimum, extent - minimum))
                    if sash < minimum or sash > extent - minimum:
                        pane.sashpos(0, fallback)
                elif sash <= 20:
                    pane.sashpos(0, default_position)
            except (AttributeError, tk.TclError):
                continue
    def _restore_collapsed_visual_panes(self) -> None:
        """Compatibility wrapper for tab-change callers."""
        self._restore_collapsed_panes()

    def _restore_sysml_layout_when_selected(self, _event=None) -> None:
        selected_tab = self.work_tabs.select()
        if selected_tab == str(self.system_traceability_tab):
            self.root.after_idle(self._open_system_traceability_window)
        visible_tabs = {
            str(self.work_browser_tab),
            str(self.sysml_architecture_tab),
            str(self.system_traceability_tab),
        }
        if selected_tab not in visible_tabs:
            return
        self.root.after_idle(self._restore_collapsed_visual_panes)
        self.root.after(250, self._restore_collapsed_visual_panes)

    def _open_system_traceability_window(self) -> None:
        """Open the expanded hierarchy view when System Traceability is selected."""
        existing = getattr(self, "traceability_window", None)
        if existing is not None and existing.winfo_exists():
            existing.lift()
            existing.focus_set()
            self._refresh_traceability_window()
            return

        window = tk.Toplevel(self.root)
        window.title("System Traceability Hierarchy")
        self._set_responsive_geometry(window, 1180, 760, 820, 520)
        window.transient(self.root)
        self.traceability_window = window
        self.traceability_zoom = 1.0

        header = ttk.Frame(window, padding=(12, 10, 12, 6))
        header.pack(fill=tk.X)
        ttk.Label(
            header,
            text="System Traceability Hierarchy",
            font=("Segoe UI", 15, "bold"),
        ).pack(side=tk.LEFT)
        ttk.Button(
            header,
            text="Refresh",
            command=self._refresh_traceability_view,
            width=10,
        ).pack(side=tk.RIGHT, padx=(6, 0))
        self.traceability_maximize_button = ttk.Button(
            header,
            text="Maximize",
            command=self._toggle_traceability_window_size,
            width=10,
        )
        self.traceability_maximize_button.pack(side=tk.RIGHT)
        self._bind_button_help(
            self.traceability_maximize_button,
            "Toggle the System Traceability window between normal and maximized size.",
        )
        close_button = ttk.Button(header, text="Close", command=window.destroy, width=8)
        close_button.pack(side=tk.RIGHT, padx=(6, 0))
        self._bind_button_help(close_button, "Close the System Traceability window.")
        self.traceability_window_status_var = tk.StringVar(value="Loading traceability hierarchy...")
        ttk.Label(
            window,
            textvariable=self.traceability_window_status_var,
            wraplength=1120,
            padding=(12, 0, 12, 8),
        ).pack(fill=tk.X)

        visual_pane = ttk.Panedwindow(window, orient=tk.VERTICAL)
        visual_pane.pack(fill=tk.BOTH, expand=True, padx=12, pady=(0, 12))

        graph_frame = ttk.Frame(visual_pane)
        frame = ttk.Frame(visual_pane)
        visual_pane.add(graph_frame, weight=3)
        visual_pane.add(frame, weight=2)

        self.traceability_graph_canvas = tk.Canvas(graph_frame, background="#f8fafc", highlightthickness=0)
        graph_y = ttk.Scrollbar(graph_frame, orient=tk.VERTICAL, command=self.traceability_graph_canvas.yview)
        graph_x = ttk.Scrollbar(graph_frame, orient=tk.HORIZONTAL, command=self.traceability_graph_canvas.xview)
        self.traceability_graph_canvas.configure(yscrollcommand=graph_y.set, xscrollcommand=graph_x.set)
        self.traceability_graph_canvas.grid(row=0, column=0, sticky="nsew")
        graph_y.grid(row=0, column=1, sticky="ns")
        graph_x.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        graph_frame.grid_rowconfigure(0, weight=1)
        graph_frame.grid_columnconfigure(0, weight=1)
        self.traceability_graph_canvas.bind(
            "<MouseWheel>",
            lambda event: self.traceability_graph_canvas.yview_scroll(-1 if event.delta > 0 else 1, "units"),
        )
        self.traceability_graph_canvas.bind(
            "<Shift-MouseWheel>",
            lambda event: self.traceability_graph_canvas.xview_scroll(-1 if event.delta > 0 else 1, "units"),
        )
        self.traceability_graph_canvas.bind("<Control-MouseWheel>", self._zoom_traceability_graph_wheel, add="+")
        self.traceability_graph_canvas.bind("<Button-1>", self._on_traceability_canvas_click)

        columns = ("upstream", "coverage", "percentage", "artifact")
        self.traceability_tree = ttk.Treeview(frame, columns=columns, show="tree headings")
        self.traceability_tree.heading("#0", text="Specification hierarchy")
        self.traceability_tree.heading("upstream", text="Upstream specification / source ledger")
        self.traceability_tree.heading("coverage", text="Coverage")
        self.traceability_tree.heading("percentage", text="%")
        self.traceability_tree.heading("artifact", text="Artifact")
        self.traceability_tree.column("#0", width=350, minwidth=240, stretch=True)
        self.traceability_tree.column("upstream", width=360, minwidth=220, stretch=True)
        self.traceability_tree.column("coverage", width=120, minwidth=95, anchor=tk.CENTER)
        self.traceability_tree.column("percentage", width=70, minwidth=60, anchor=tk.CENTER)
        self.traceability_tree.column("artifact", width=330, minwidth=180, stretch=True)
        self.traceability_tree.tag_configure("traceability_even", background="#f7fafc")
        self.traceability_tree.tag_configure("traceability_odd", background="#ffffff")
        tree_y = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=self.traceability_tree.yview)
        tree_x = ttk.Scrollbar(frame, orient=tk.HORIZONTAL, command=self.traceability_tree.xview)
        self.traceability_tree.configure(yscrollcommand=tree_y.set, xscrollcommand=tree_x.set)
        self.traceability_tree.bind("<<TreeviewSelect>>", self._on_traceability_tree_select)
        self.traceability_tree.grid(row=0, column=0, sticky="nsew")
        tree_y.grid(row=0, column=1, sticky="ns")
        tree_x.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)
        window.bind("<Configure>", self._on_traceability_window_resize, add="+")
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        self._refresh_traceability_window()

    @staticmethod
    def _traceability_matrix_ids(path: Path) -> set[str]:
        if not path.exists():
            return set()
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                return {
                    (row.get("source_req_id") or row.get("id") or row.get("req_id") or "").strip()
                    for row in csv.DictReader(handle)
                    if (row.get("source_req_id") or row.get("id") or row.get("req_id") or "").strip()
                }
        except (OSError, UnicodeError):
            return set()

    def _traceability_source_ids(self) -> set[str]:
        ledger_path = TRACEABILITY_REPORT_DIR / "requirement_allocation_ledger.csv"
        return self._traceability_matrix_ids(ledger_path)

    def _traceability_supplementary_sources(self) -> dict[str, set[str]]:
        path = REPO_ROOT / "artifacts/stage1_specs/architecture_mapping_preview.csv"
        sources: dict[str, set[str]] = {}
        if not path.exists():
            return sources
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                for row in csv.DictReader(handle):
                    source = (row.get("supplementary_source_spec") or "").strip()
                    req_id = (row.get("requirement_id") or row.get("canonical_requirement_id") or "").strip()
                    if source and req_id and Path(source.replace("\\", "/")).name.casefold() != approved_primary_source_name(REPO_ROOT):
                        sources.setdefault(Path(source.replace("\\", "/")).name, set()).add(req_id)
        except (OSError, UnicodeError):
            return {}
        return sources

    def _traceability_document_nodes(self) -> list[dict[str, object]]:
        return list(approved_traceability_hierarchy(REPO_ROOT).get("children") or [])
        source_ids = self._traceability_source_ids()
        total = len(source_ids)
        documents: list[dict[str, object]] = []

        def add_document(label: str, kind: str, path: Path, dependencies: str) -> None:
            ids = self._traceability_matrix_ids(path)
            covered = len(ids & source_ids)
            documents.append({
                "label": label,
                "kind": kind,
                "covered": covered,
                "total": total,
                "dependencies": dependencies,
                "artifact": path.relative_to(REPO_ROOT).as_posix() if path.is_relative_to(REPO_ROOT) else str(path),
                "children": [],
            })

        add_document("SRS - System Requirements Specification", "spec", REPO_ROOT / "artifacts/stage3_srs/srs_traceability_matrix.csv", "Primary + supplementary sources")
        add_document("ARS - Analog Requirements Specification", "spec", REPO_ROOT / "artifacts/stage4_ars/ars_traceability_matrix.csv", "SRS; analog source scope")
        add_document("DRS - Digital Requirements Specification", "spec", REPO_ROOT / "artifacts/stage5_drs/drs_traceability_matrix.csv", "SRS; primary + supplementary digital scope")

        for kind, directory, prefix, dependency in (
            ("Digital IPOS", REPO_ROOT / "artifacts/stage6_digital_ipos", "Digital IPOS", "DRS; approved block allocation"),
            ("Analog IPOS", REPO_ROOT / "artifacts/stage7_analog_ipos", "Analog IPOS", "ARS; approved block allocation"),
        ):
            matrix = directory / "ipos_traceability_matrix.csv"
            blocks_directory = directory / "blocks"
            created_block_dirs = {
                name.casefold(): path
                for name, path in created_ipos_block_directories(REPO_ROOT, kind).items()
            }

            def block_slug(value: str) -> str:
                return re.sub(r"[^a-z0-9]+", "", value.casefold())

            ids_by_block: dict[str, set[str]] = {}
            if matrix.exists():
                try:
                    with matrix.open("r", encoding="utf-8-sig", newline="") as handle:
                        for row in csv.DictReader(handle):
                            block = (row.get("owning_block") or row.get("mapped_block") or "Unassigned").strip() or "Unassigned"
                            req_id = (row.get("source_req_id") or row.get("id") or "").strip()
                            if req_id:
                                ids_by_block.setdefault(block, set()).add(req_id)
                except (OSError, UnicodeError):
                    ids_by_block = {}
            created_ids_by_block = {
                block: ids
                for block, ids in ids_by_block.items()
                if any(block_slug(block) == block_slug(created) for created in created_block_dirs)
            }
            children = []
            created_ids: set[str] = set()
            for created_name, created_path in sorted(created_block_dirs.items()):
                matching_ids = next(
                    (ids for block, ids in created_ids_by_block.items() if block_slug(block) == block_slug(created_name)),
                    set(),
                )
                created_ids.update(matching_ids)
                children.append({
                    "label": f"{prefix}: {created_path.name}",
                    "kind": "block",
                    "covered": len(matching_ids & source_ids),
                    "total": total,
                    "dependencies": dependency,
                    "artifact": created_path.relative_to(REPO_ROOT).as_posix(),
                    "children": [],
                })
            if children:
                documents.append({
                    "label": kind,
                    "kind": "group",
                    "covered": len(created_ids & source_ids),
                    "total": total,
                    "dependencies": dependency,
                    "artifact": blocks_directory.relative_to(REPO_ROOT).as_posix(),
                    "children": children,
                })
        return documents

    def _zoom_traceability_graph_wheel(self, event):
        """Zoom the traceability graph while preserving ordinary scroll behavior."""
        current = float(getattr(self, "traceability_zoom", 1.0))
        self.traceability_zoom = max(0.65, min(2.25, current * (1.15 if event.delta > 0 else 1 / 1.15)))
        self._refresh_traceability_window()
        return "break"

    def _refresh_traceability_window(self) -> None:
        tree = getattr(self, "traceability_tree", None)
        window = getattr(self, "traceability_window", None)
        if tree is None or window is None or not window.winfo_exists():
            return
        graph = approved_generated_document_dependency_graph(REPO_ROOT)
        tree.delete(*tree.get_children())
        for index, node in enumerate(graph.get("nodes") or []):
            if not isinstance(node, dict):
                continue
            node_id = str(node.get("id") or "")
            metrics = list(node.get("upstream_coverage_metrics") or [])
            if node.get("kind") in {"source_spec", "supplementary_spec"} and not metrics:
                source_role = "Primary spec" if node.get("kind") == "source_spec" else "Supplementary spec"
                upstream_text = f"Source ledger: {source_role}"
                coverage_text = ""
                percentage_text = ""
            else:
                upstream_text = " | ".join(str(value.get("upstream") or "") for value in metrics)
                if not upstream_text:
                    upstream_text = "No approved upstream edge"
                coverage_text = " | ".join(
                    f"{int(value.get('covered') or 0)} / {int(value.get('total') or 0)}"
                    for value in metrics
                )
                percentage_text = " | ".join(
                    f"{float(value.get('percentage') or 0.0):.1f}%"
                    for value in metrics
                )
            tree.insert(
                "",
                tk.END,
                iid=node_id,
                text=str(node.get("label") or node_id),
                values=(
                    upstream_text,
                    coverage_text,
                    percentage_text,
                    str(node.get("artifact") or ""),
                ),
                tags=("traceability_even" if index % 2 == 0 else "traceability_odd",),
                open=True,
            )
        selected = self.traceability_selected_block
        if selected and tree.exists(selected):
            tree.selection_set(selected)
            tree.focus(selected)
            tree.see(selected)
        self._draw_generated_dependency_graph(self.traceability_graph_canvas)
        self.traceability_window_status_var.set(
            f"Central payload: {len(graph.get('nodes') or [])} nodes, "
            f"{sum(1 for edge in graph.get('edges') or [] if int((edge.get('coverage') or {}).get('covered') or 0) > 0)} visible approved RM edges, "
            f"{len(graph.get('source_spec_edges') or [])} source-to-document and "
            f"{len(graph.get('source_spec_relationship_edges') or [])} source-to-source coverage links. "
            "RM edges use approved Covers; source-spec links show reciprocal source and target ID coverage."
            " Source-ledger coverage is source IDs represented by a node divided by unique IDs in the approved allocation ledger."
        )

    def _on_traceability_window_resize(self, _event=None) -> None:
        """Redraw the central graph after user-driven window or pane resizing."""
        if getattr(self, "_traceability_resize_after_id", None):
            self.root.after_cancel(self._traceability_resize_after_id)
        self._traceability_resize_after_id = self.root.after_idle(self._refresh_traceability_layout)

    def _refresh_traceability_layout(self) -> None:
        self._traceability_resize_after_id = None
        canvas = getattr(self, "traceability_graph_canvas", None)
        if canvas is not None and canvas.winfo_exists():
            self._draw_generated_dependency_graph(canvas)
        tree = getattr(self, "traceability_tree", None)
        if tree is not None and tree.winfo_exists():
            tree.update_idletasks()

    def _toggle_traceability_window_size(self) -> None:
        window = getattr(self, "traceability_window", None)
        if window is None or not window.winfo_exists():
            return
        is_zoomed = window.state() == "zoomed"
        window.state("normal" if is_zoomed else "zoomed")
        self.traceability_maximize_button.configure(text="Maximize" if is_zoomed else "Restore")

    def _refresh_traceability_view(self) -> None:
        """Reload the approved dependency graph and clear stale selection state."""
        self.traceability_selected_edge = None
        self.traceability_selected_block = None
        self._refresh_traceability_window()
        canvas = getattr(self, "traceability_graph_canvas", None)
        if canvas is not None and canvas.winfo_exists():
            canvas.update_idletasks()

    def _draw_traceability_graph(
        self,
        source_name: str,
        source_ids: set[str],
        supplementary: dict[str, set[str]],
        documents: list[dict[str, object]],
    ) -> None:
        canvas = getattr(self, "traceability_graph_canvas", None)
        if canvas is None:
            return
        self._draw_generated_dependency_graph(canvas)
        return

    def _on_traceability_tree_select(self, _event=None) -> None:
        selected = self.traceability_tree.selection()
        if not selected:
            return
        node_id = selected[0]
        if node_id == self.traceability_selected_block:
            return
        self.traceability_selected_block = node_id
        self.traceability_selected_edge = None
        self._draw_generated_dependency_graph(self.traceability_graph_canvas)

    def _on_traceability_canvas_click(self, event) -> str:
        canvas = self.traceability_graph_canvas
        current = canvas.find_overlapping(event.x, event.y, event.x, event.y)
        if not current:
            current = canvas.find_withtag("current")
        tags = canvas.gettags(current[0]) if current else ()
        edge_id = next((tag for tag in tags if tag.startswith("traceability_edge:")), None)
        target_id = next((tag for tag in tags if tag.startswith("traceability_target:")), None)
        block_id = next((tag for tag in tags if tag.startswith("traceability_node:")), None)
        if edge_id:
            self.traceability_selected_edge = edge_id.split(":", 1)[1]
            self.traceability_selected_block = target_id.split(":", 1)[1] if target_id else None
        elif block_id:
            self.traceability_selected_block = block_id.split(":", 1)[1]
            self.traceability_selected_edge = None
        else:
            self.traceability_selected_edge = None
            self.traceability_selected_block = None
        self._refresh_traceability_window()
        return "break"

    def _clear_traceability_selection(self, _event=None) -> str:
        self.traceability_selected_edge = None
        self.traceability_selected_block = None
        self._refresh_traceability_window()
        return "break"

    def _select_traceability_edge(self, event) -> str:
        canvas = self.traceability_graph_canvas
        current = canvas.find_withtag("current")
        if not current:
            return "break"
        edge_id = next((tag for tag in canvas.gettags(current[0]) if tag.startswith("traceability_edge:")), None)
        if edge_id:
            self.traceability_selected_edge = edge_id.split(":", 1)[1]
            self.traceability_selected_block = None
            self._refresh_traceability_window()
        return "break"

    def _select_traceability_block(self, event) -> str:
        canvas = self.traceability_graph_canvas
        current = canvas.find_withtag("current")
        if not current:
            return "break"
        block_id = next((tag for tag in canvas.gettags(current[0]) if tag.startswith("traceability_node:")), None)
        if block_id:
            self.traceability_selected_block = block_id.split(":", 1)[1]
            self.traceability_selected_edge = None
            self._refresh_traceability_window()
        return "break"

    def _draw_generated_dependency_graph(self, canvas: tk.Canvas) -> None:
        """Render only central generated-document dependency edges."""
        canvas.delete("all")
        graph = approved_generated_document_dependency_graph(REPO_ROOT)
        node_list = list(graph.get("nodes") or [])
        edge_list = [
            edge for edge in graph.get("edges") or []
            if int(((edge.get("coverage") or {}).get("covered")) or 0) > 0
            and int(((edge.get("coverage") or {}).get("total")) or 0) > 0
        ]
        source_spec_edges = [
            edge for edge in graph.get("source_spec_edges") or []
            if int(((edge.get("source_coverage") or {}).get("covered")) or 0) > 0
            and int(((edge.get("target_coverage") or {}).get("covered")) or 0) > 0
        ]
        source_spec_relationship_edges = [
            edge for edge in graph.get("source_spec_relationship_edges") or []
            if int(((edge.get("source_coverage") or {}).get("covered")) or 0) > 0
            and int(((edge.get("target_coverage") or {}).get("covered")) or 0) > 0
        ]
        if not node_list:
            canvas.create_text(24, 24, anchor="nw", text="No approved generated-document dependencies are available.", fill="#52606d")
            canvas.configure(scrollregion=(0, 0, 700, 100))
            return

        zoom = float(getattr(self, "traceability_zoom", 1.0))
        body_font = tkfont.Font(family="Segoe UI", size=max(8, int(9 * zoom)))
        layout_edges = [*edge_list, *source_spec_edges, *source_spec_relationship_edges]
        incoming_nodes = {
            str(edge.get("target") or "")
            for edge in layout_edges
            if isinstance(edge, dict)
        }
        depth: dict[str, int] = {
            str(node.get("id")): 0
            for node in node_list
            if str(node.get("id") or "") not in incoming_nodes
        }
        changed = True
        while changed:
            changed = False
            for edge in layout_edges:
                source = str(edge["source"])
                target = str(edge["target"])
                if source in depth and (target not in depth or depth[target] < depth[source] + 1):
                    depth[target] = depth[source] + 1
                    changed = True
        for node in node_list:
            depth.setdefault(str(node["id"]), 0)
        for node in node_list:
            if node.get("kind") != "supplementary_spec":
                continue
            node_id = str(node.get("id") or "")
            targets = [
                depth.get(str(edge.get("target")), 0)
                for edge in layout_edges
                if str(edge.get("source")) == node_id
            ]
            if targets:
                depth[node_id] = max(0, min(targets) - 1)
        widths: dict[str, float] = {}
        heights: dict[str, float] = {}
        for node in node_list:
            label = str(node.get("label") or "")
            if node.get("kind") in {"source_spec", "supplementary_spec"}:
                source_role = "Primary spec" if node.get("kind") == "source_spec" else "Supplementary spec"
                lines = [label, f"Source ledger: {source_role}"]
                if node.get("kind") == "supplementary_spec":
                    lines.extend(str(value) for value in node.get("upstream_coverage") or [])
            else:
                lines = [label]
                lines.extend(
                    "Upstream coverage: " + str(value)
                    for value in node.get("upstream_coverage") or []
                )
            longest = max((body_font.measure(line) for line in lines), default=180)
            widths[str(node["id"])] = max(220 * zoom, min(430 * zoom, longest + 28 * zoom))
            heights[str(node["id"])] = max(70 * zoom, (len(lines) * 18 + 24) * zoom)

        positions: dict[str, tuple[float, float]] = {}
        column_gap = 90 * zoom
        row_gap = 58 * zoom
        column_widths = {
            current_depth: max(
                (widths[str(node["id"])] for node in node_list if depth[str(node["id"])] == current_depth),
                default=260 * zoom,
            )
            for current_depth in sorted(set(depth.values()))
        }
        column_x: dict[int, float] = {}
        current_x = 40 * zoom
        for current_depth in sorted(column_widths):
            column_x[current_depth] = current_x
            current_x += column_widths[current_depth] + column_gap
        for current_depth in sorted(set(depth.values())):
            current_y = 36 * zoom
            for node in sorted(
                (item for item in node_list if depth[str(item["id"])] == current_depth),
                key=lambda item: str(item.get("id") or "").casefold(),
            ):
                node_id = str(node["id"])
                positions[node_id] = (column_x[current_depth], current_y)
                current_y += heights[node_id] + row_gap

        def edge_is_selected(edge_id: str, source: str, target: str) -> bool:
            if self.traceability_selected_edge == edge_id:
                return True
            selected = self.traceability_selected_block
            return bool(selected and (source == selected or target == selected))

        for index, edge in enumerate(source_spec_edges):
            source = str(edge.get("source") or "")
            target = str(edge.get("target") or "")
            if source not in positions or target not in positions:
                continue
            sx, sy = positions[source]
            tx, ty = positions[target]
            edge_id = f"source-spec:{source}->{target}:{index}"
            selected = edge_is_selected(edge_id, source, target)
            source_coverage = edge.get("source_coverage") or {}
            target_coverage = edge.get("target_coverage") or {}
            canvas.create_line(
                sx + widths[source], sy + heights[source] / 2,
                tx, ty + heights[target] / 2,
                fill="#005fbd" if selected else "#c47a25",
                width=max(2, int(3 * zoom)) if selected else max(1, int(2 * zoom)),
                dash=(5, 3),
                arrow=tk.LAST,
                tags=("traceability_source_edge", f"traceability_edge:{edge_id}", f"traceability_target:{target}"),
            )
        for index, edge in enumerate(source_spec_relationship_edges):
            source = str(edge.get("source") or "")
            target = str(edge.get("target") or "")
            if source not in positions or target not in positions:
                continue
            sx, sy = positions[source]
            tx, ty = positions[target]
            edge_id = f"source-spec-link:{source}->{target}:{index}"
            selected = edge_is_selected(edge_id, source, target)
            canvas.create_line(
                sx + widths[source], sy + heights[source] / 2,
                tx, ty + heights[target] / 2,
                fill="#005fbd" if selected else "#c47a25",
                width=max(2, int(3 * zoom)) if selected else max(1, int(2 * zoom)),
                dash=(2, 3),
                arrow=tk.LAST,
                tags=("traceability_source_relationship_edge", f"traceability_edge:{edge_id}", f"traceability_target:{target}"),
            )
        for index, edge in enumerate(edge_list):
            source = str(edge["source"])
            target = str(edge["target"])
            if source not in positions or target not in positions:
                continue
            sx, sy = positions[source]
            tx, ty = positions[target]
            edge_id = f"{source}->{target}:{index}"
            selected = edge_is_selected(edge_id, source, target)
            canvas.create_line(
                sx + widths[source], sy + heights[source] / 2,
                tx, ty + heights[target] / 2,
                fill="#005fbd" if selected else "#c47a25",
                width=max(3, int(4 * zoom)) if selected else max(2, int(3 * zoom)),
                arrow=tk.LAST,
                tags=("traceability_edge", f"traceability_edge:{edge_id}", f"traceability_target:{target}"),
            )
        for node in node_list:
            node_id = str(node["id"])
            x, y = positions[node_id]
            if node.get("kind") in {"source_spec", "supplementary_spec"}:
                source_role = "Primary spec" if node.get("kind") == "source_spec" else "Supplementary spec"
                lines = [str(node.get("label") or ""), f"Source ledger: {source_role}"]
                if node.get("kind") == "supplementary_spec":
                    lines.extend(str(value) for value in node.get("upstream_coverage") or [])
            else:
                lines = [str(node.get("label") or "")]
                lines.extend("Upstream coverage: " + str(value) for value in node.get("upstream_coverage") or [])
            selected = self.traceability_selected_block == node_id
            tags = ("traceability_node", f"traceability_node:{node_id}")
            canvas.create_rectangle(
                x, y, x + widths[node_id], y + heights[node_id],
                fill="#cfe8ff" if selected else "#e8f1dc" if node.get("kind") == "ipos_block" else "#f8eadf" if node.get("kind") == "supplementary_spec" else "#e1edf7",
                outline="#1976d2" if selected else "#486581",
                width=max(2, int(3 * zoom)) if selected else max(1, int(2 * zoom)),
                tags=tags,
            )
            canvas.create_text(x + widths[node_id] / 2, y + heights[node_id] / 2, text="\n".join(lines), font=body_font, fill="#1f2d3d", justify=tk.CENTER, width=widths[node_id] - 16 * zoom, tags=tags)
        max_x = max((x + widths[node_id] for node_id, (x, _y) in positions.items()), default=0) + 50 * zoom
        max_y = max((y + heights[node_id] for node_id, (_x, y) in positions.items()), default=0) + 50 * zoom
        canvas.configure(scrollregion=(0, 0, max_x, max_y))

    @staticmethod
    def _traceability_percentage(node: dict[str, object]) -> str:
        covered = int(node.get("covered") or 0)
        total = int(node.get("total") or 0)
        return f"{covered / total * 100.0:.1f}%" if total else "0.0%"

    def _draw_traceability_graph_legacy(
        self,
        source_name: str,
        source_ids: set[str],
        supplementary: dict[str, set[str]],
        documents: list[dict[str, object]],
    ) -> None:
        """Compatibility body retained for historical generated GUI state."""
        canvas = getattr(self, "traceability_graph_canvas", None)
        if canvas is None:
            return
        canvas.delete("all")
        zoom = float(getattr(self, "traceability_zoom", 1.0))
        column_gap = 70 * zoom
        row_gap = 42 * zoom
        positions: dict[str, tuple[int, int]] = {}
        nodes: dict[str, dict[str, object]] = {}
        hierarchy_edges: list[tuple[str, str]] = []
        coverage_edges: list[tuple[str, str, str]] = []
        order_by_depth: dict[int, int] = {}

        def node_dimensions(node: dict[str, object]) -> tuple[float, float]:
            label = str(node.get("label") or "")
            dependency = str(node.get("dependency") or "")
            longest_line = max((len(line) for line in f"{label}\n{dependency}".splitlines()), default=20)
            width = max(220 * zoom, min(430 * zoom, longest_line * 6.2 * zoom + 30))
            wrap_chars = max(24, int(width / (6.2 * zoom)))
            wrapped_lines = sum(max(1, (len(line) + wrap_chars - 1) // wrap_chars) for line in f"{label}\n{dependency}".splitlines())
            height = max(82 * zoom, wrapped_lines * 15 * zoom + 28 * zoom)
            return width, height

        def add_node(key: str, node: dict[str, object], depth: int) -> None:
            order_by_depth[depth] = order_by_depth.get(depth, 0) + 1
            node["_depth"] = depth
            node["_width"], node["_height"] = node_dimensions(node)
            nodes[key] = node

        primary_key = "source_primary"
        primary_node = {
            "label": f"Primary source\n{source_name}",
            "covered": len(source_ids),
            "total": len(source_ids),
            "dependency": "approved Stage 1 source",
            "fill": "#d9e8f5",
            "coverage_edges": [],
        }
        add_node(primary_key, primary_node, 0)
        supplementary_keys: dict[str, str] = {}
        for index, (name, ids) in enumerate(sorted(supplementary.items(), key=lambda item: item[0].casefold())):
            key = f"source_supplementary_{index}"
            supplementary_keys[name] = key
            add_node(key, {
                "label": f"Supplementary source\n{name}",
                "covered": len(ids),
                "total": len(source_ids),
                "dependency": "approved mapping provenance",
                "fill": "#f8e5d2",
                "coverage_edges": [],
            }, 0)

        generated_key = "generated_hierarchy"
        add_node(generated_key, {
            "label": "Requirement Management / Design hierarchy",
            "covered": 0,
            "total": len(source_ids),
            "dependency": "central approved hierarchy contract",
            "fill": "#e7edf3",
            "coverage_edges": [],
        }, 1)
        hierarchy_edges.append((generated_key, primary_key))
        hierarchy_edges.extend((generated_key, key) for key in supplementary_keys.values())

        design_root = approved_traceability_hierarchy(REPO_ROOT)

        def visit(node: dict[str, object], parent: str, depth: int, ordinal: str) -> None:
            key = f"hierarchy_{ordinal}"
            ids = set(node.get("coverage_ids") or [])
            edges = list(node.get("coverage_edges") or [])
            add_node(key, {
                "label": node.get("label", ""),
                "covered": len(ids & source_ids),
                "total": len(source_ids),
                "dependency": "\n".join(edges[:3]) or "approved hierarchy parent",
                "fill": "#e8f1dc" if node.get("kind") == "block" else "#e1edf7",
                "coverage_edges": edges,
            }, depth)
            hierarchy_edges.append((parent, key))
            for source_id in sorted(ids):
                target_key = primary_key
                for name, source_ids_for_spec in supplementary.items():
                    if source_id in source_ids_for_spec:
                        target_key = supplementary_keys.get(name, primary_key)
                        break
                coverage_edges.append((key, target_key, source_id))
            for child_index, child in enumerate(node.get("children") or []):
                visit(child, key, depth + 1, f"{ordinal}_{child_index}")

        visit(design_root, generated_key, 2, "root")

        depth_widths: dict[int, float] = {}
        depth_nodes: dict[int, list[str]] = {}
        for key, node in nodes.items():
            depth = int(node["_depth"])
            depth_widths[depth] = max(depth_widths.get(depth, 0.0), float(node["_width"]))
            depth_nodes.setdefault(depth, []).append(key)
        depth_x: dict[int, float] = {}
        current_x = 40 * zoom
        for depth in sorted(depth_widths):
            depth_x[depth] = current_x
            current_x += depth_widths[depth] + column_gap
        for depth, keys in depth_nodes.items():
            current_y = 40 * zoom
            for key in keys:
                positions[key] = (depth_x[depth], current_y)
                current_y += float(nodes[key]["_height"]) + row_gap

        def coverage(node: dict[str, object]) -> str:
            covered = int(node.get("covered") or 0)
            total = int(node.get("total") or 0)
            percentage = covered / total * 100.0 if total else 0.0
            return f"{covered} / {total} ({percentage:.1f}%)"

        for parent, child in hierarchy_edges:
            if parent not in positions or child not in positions:
                continue
            px, py = positions[parent]
            cx, cy = positions[child]
            parent_width = float(nodes[parent]["_width"])
            parent_height = float(nodes[parent]["_height"])
            child_width = float(nodes[child]["_width"])
            canvas.create_line(px + parent_width / 2, py + parent_height, cx + child_width / 2, cy, fill="#718096", width=max(1, int(2 * zoom)), arrow=tk.LAST)
        for child, parent, source_id in coverage_edges:
            if child not in positions or parent not in positions:
                continue
            cx, cy = positions[child]
            px, py = positions[parent]
            child_width = float(nodes[child]["_width"])
            child_height = float(nodes[child]["_height"])
            parent_width = float(nodes[parent]["_width"])
            parent_height = float(nodes[parent]["_height"])
            canvas.create_line(cx, cy + child_height / 2, px + parent_width, py + parent_height / 2, fill="#b7791f", width=max(1, int(zoom)), dash=(4, 2), arrow=tk.LAST)
            canvas.create_text((cx + px + parent_width) / 2, (cy + py + parent_height) / 2, text="", fill="#975a16", font=("Segoe UI", max(7, int(7 * zoom))))
        for key, node in nodes.items():
            x, y = positions[key]
            width = float(node["_width"])
            height = float(node["_height"])
            canvas.create_rectangle(x, y, x + width, y + height, fill=str(node["fill"]), outline="#486581", width=max(1, int(2 * zoom)))
            canvas.create_text(x + width / 2, y + 18 * zoom, text=str(node["label"]), font=("Segoe UI", max(8, int(10 * zoom)), "bold"), fill="#1f2d3d", width=width - 16 * zoom)
            canvas.create_text(x + width / 2, y + height / 2 + 10 * zoom, text=f"Coverage: {coverage(node)}\n{node['dependency']}", font=("Segoe UI", max(7, int(8 * zoom))), fill="#334e68", width=width - 16 * zoom)
        max_y = max((y + float(nodes[key]["_height"]) for key, (_x, y) in positions.items()), default=0) + 40 * zoom
        max_x = max((x + float(nodes[key]["_width"]) for key, (x, _y) in positions.items()), default=0) + 40 * zoom
        canvas.configure(scrollregion=(0, 0, max_x, max_y))
        return

        # Legacy flat graph code retained below only as unreachable compatibility text.
        canvas.delete("all")
        node_width = 250
        node_height = 72
        column_gap = 55
        row_gap = 34
        columns = 4
        canvas_width = columns * node_width + (columns - 1) * column_gap + 80
        positions: dict[str, tuple[int, int]] = {}
        nodes: dict[str, dict[str, object]] = {}

        def add_node(key: str, label: str, covered: int, total: int, dependency: str, x: int, y: int, fill: str) -> None:
            nodes[key] = {
                "label": label,
                "covered": covered,
                "total": total,
                "dependency": dependency,
                "fill": fill,
            }
            positions[key] = (x, y)

        def coverage(node: dict[str, object]) -> str:
            covered = int(node.get("covered") or 0)
            total = int(node.get("total") or 0)
            percentage = covered / total * 100.0 if total else 0.0
            return f"{covered} / {total} ({percentage:.1f}%)"

        add_node("primary", f"Primary source\n{source_name}", len(source_ids), len(source_ids), "Stage 1 OCR -> allocation ledger", 40, 24, "#d9e8f5")
        add_node("supplementary", f"Supplementary sources\n{len(supplementary)} source specification(s)", len(set().union(*supplementary.values())) if supplementary else 0, len(source_ids), "Merged through Stage 2 review", 345, 24, "#f8e5d2")

        generated_y = 170
        generated_key = "generated"
        add_node(generated_key, "Generated specification hierarchy", 0, len(source_ids), "Approved snapshot", 650, generated_y, "#e7edf3")
        spec_nodes: dict[str, str] = {}
        spec_documents = [document for document in documents if document.get("kind") == "spec"]
        for index, document in enumerate(spec_documents):
            key = f"spec_{index}"
            spec_nodes[str(document.get("label"))] = key
            add_node(
                key,
                str(document.get("label")),
                int(document.get("covered") or 0),
                int(document.get("total") or 0),
                str(document.get("dependencies") or ""),
                40 + index * (node_width + column_gap),
                315,
                "#dcefd8" if "DRS" in str(document.get("label")) else "#e1edf7",
            )

        ipos_y = 480
        ipos_index = 0
        for document in documents:
            if document.get("kind") != "group":
                continue
            group_key = f"ipos_{ipos_index}"
            ipos_index += 1
            add_node(
                group_key,
                str(document.get("label")),
                int(document.get("covered") or 0),
                int(document.get("total") or 0),
                str(document.get("dependencies") or ""),
                40 + (len(spec_documents) + ipos_index - 1) * (node_width + column_gap),
                ipos_y,
                "#e8f1dc" if "Digital" in str(document.get("label")) else "#f6e6d8",
            )
            nodes[group_key]["children"] = []
            for child_index, child in enumerate(document.get("children") or []):
                child_key = f"{group_key}_{child_index}"
                child_y = ipos_y + node_height + row_gap + child_index * (node_height + row_gap)
                add_node(
                    child_key,
                    str(child.get("label")),
                    int(child.get("covered") or 0),
                    int(child.get("total") or 0),
                    str(child.get("dependencies") or ""),
                    40 + (len(spec_documents) + ipos_index - 1) * (node_width + column_gap),
                    child_y,
                    "#f0f5e9",
                )
                nodes[group_key]["children"].append(child_key)

        edges = [("primary", "generated"), ("supplementary", "generated")]
        edges.extend((generated_key, key) for key in spec_nodes.values())
        for key in list(nodes):
            if key.startswith("ipos_") and "_" not in key[5:]:
                edges.append((generated_key, key))
                if "Digital" in str(nodes[key]["label"]):
                    drs_key = next((value for label, value in spec_nodes.items() if "DRS" in label), None)
                    if drs_key:
                        edges.append((drs_key, key))
                if "Analog" in str(nodes[key]["label"]):
                    ars_key = next((value for label, value in spec_nodes.items() if "ARS" in label), None)
                    if ars_key:
                        edges.append((ars_key, key))
        for parent, child_keys in ((key, value.get("children", [])) for key, value in nodes.items()):
            edges.extend((parent, child) for child in child_keys)

        for parent, child in edges:
            if parent not in positions or child not in positions:
                continue
            px, py = positions[parent]
            cx, cy = positions[child]
            canvas.create_line(
                px + node_width / 2,
                py + node_height,
                cx + node_width / 2,
                cy,
                fill="#718096",
                width=2,
                arrow=tk.LAST,
            )

        for key, node in nodes.items():
            x, y = positions[key]
            canvas.create_rectangle(x, y, x + node_width, y + node_height, fill=str(node["fill"]), outline="#486581", width=2)
            canvas.create_text(
                x + node_width / 2,
                y + 22,
                text=str(node["label"]),
                font=("Segoe UI", 10, "bold"),
                fill="#1f2d3d",
                width=node_width - 16,
            )
            canvas.create_text(
                x + node_width / 2,
                y + 50,
                text=f"Coverage: {coverage(node)}\n{node['dependency']}",
                font=("Segoe UI", 8),
                fill="#334e68",
                width=node_width - 16,
            )
        max_y = max((y for _x, y in positions.values()), default=0) + node_height + 40
        canvas.configure(scrollregion=(0, 0, canvas_width, max_y))

    @staticmethod
    def _insert_traceability_node(tree: ttk.Treeview, parent: str, node: dict[str, object], *, open_node: bool = False) -> str:
        total = int(node.get("total") or 0)
        covered = int(node.get("covered") or 0)
        percentage = covered / total * 100.0 if total else 0.0
        item = tree.insert(
            parent,
            tk.END,
            text=str(node.get("label") or ""),
            values=(f"{covered} / {total}", f"{percentage:.1f}%", str(node.get("dependencies") or ""), str(node.get("artifact") or "")),
            open=open_node,
        )
        for child in node.get("children") or []:
            WorkflowGUI._insert_traceability_node(tree, item, child, open_node=False)
        return item

    def _save_gui_state(self, user_initiated: bool = False) -> None:
        if not self._gui_layout_ready:
            return
        self._save_state_after_id = None
        self.root.update_idletasks()
        default_sashes = {
            "main_vertical_pane": [300],
            "top_horizontal_pane": [280],
            "work_pane": [300],
            "visual_pane": [500],
            "chat_vertical_pane": [420],
            "sysml_vertical_pane": [500],
            "sysml_horizontal_pane": [300],
            "traceability_pane": [180],
        }
        current_layout = {
            "geometry": self.root.geometry(),
            "layout_scale": self.layout_scale,
            "sashes": {},
        }
        for name in self._gui_pane_names():
            pane = getattr(self, name, None)
            if self._is_sash_pane(pane):
                try:
                    saved_positions = []
                    for index in range(max(0, len(pane.panes()) - 1)):
                        position = pane.sashpos(index)
                        saved_positions.append(max(0, position))
                    current_layout["sashes"][name] = saved_positions
                except (AttributeError, tk.TclError):
                    pass
        state = {
            "geometry": current_layout["geometry"],
            "layout_scale": self.layout_scale,
            "verbose_trace": self.verbose_trace_enabled,
            "layout_saved_by_user": bool(self.gui_state.get("layout_saved_by_user", False)),
            "sashes": current_layout["sashes"],
        }
        if user_initiated:
            state["layout_saved_by_user"] = True
            state["last_user_layout"] = current_layout
        elif state["layout_saved_by_user"]:
            saved_layout = self.gui_state.get("last_user_layout")
            if isinstance(saved_layout, dict):
                state["last_user_layout"] = saved_layout
            else:
                state["last_user_layout"] = {
                    "geometry": state["geometry"],
                    "sashes": state["sashes"],
                }
        try:
            GUI_STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
            payload = json.dumps(state, indent=2) + "\n"
            GUI_STATE_PATH.write_text(payload, encoding="utf-8")
            self.gui_state = state
            saved_at = datetime.now().astimezone().isoformat(timespec="seconds")
            self.root.after(
                0,
                self._append_log_ui,
                f"[{saved_at}] [LAYOUT] Saved GUI layout to {GUI_STATE_PATH.as_posix()} "
                f"geometry={state['geometry']} sashes={state['sashes']}",
            )
        except OSError as exc:
            self.append_log(f"[WARN] Could not save GUI layout: {exc}")

    def _save_gui_layout_now(self) -> None:
        self.root.update_idletasks()
        self.gui_state["layout_saved_by_user"] = True
        self._save_gui_state(user_initiated=True)
        self.root.after(
            0,
            self._append_log_ui,
            f"[{datetime.now().astimezone().isoformat(timespec='seconds')}] [LAYOUT] "
            "Save Window Layout completed; persisted geometry and pane settings.",
        )

    def _close_window(self) -> None:
        if self.pdf_reader_document is not None:
            self._close_pdf_reader()
        self._save_gui_state()
        self.root.destroy()

    def _bind_button_help(self, button: ttk.Button, help_text: str) -> None:
        button.bind("<Enter>", lambda event, text=help_text, owner=button: self._show_hover_help(event, text, owner))
        button.bind("<Leave>", lambda _event: self._hide_hover_help())

    def _install_hover_popup_controller(self) -> None:
        """Install runtime-only cleanup for the centralized popup controller."""
        self.root.bind_all("<Motion>", self._dismiss_hover_help_outside_owner, add="+")
        self.root.bind_all("<Leave>", self._dismiss_hover_help_outside_owner, add="+")
        self.root.bind_all("<ButtonPress>", self._close_hover_popup_on_button, add="+")
        self.root.bind("<FocusOut>", self._close_hover_popup_on_focus, add="+")

    def _clear_popup_if_current(self, popup) -> None:
        if popup is not self._active_popup:
            return
        self._active_popup = None
        self._active_popup_kind = None
        if popup is self.help_popup:
            self.help_popup = None
            self.help_popup_label = None
            self.help_popup_owner = None
            self.help_popup_owner_widget = None

    def _close_popup(self) -> None:
        """Close the one active transient popup without touching persisted GUI state."""
        popup = self._active_popup
        self._active_popup = None
        self._active_popup_kind = None
        if popup is None:
            return
        try:
            popup.grab_release()
        except (AttributeError, tk.TclError):
            pass
        try:
            popup.destroy()
        except (AttributeError, tk.TclError):
            pass
        if popup is self.help_popup:
            self.help_popup = None
            self.help_popup_label = None
            self.help_popup_owner = None
            self.help_popup_owner_widget = None

    def _activate_popup(self, popup, kind: str) -> None:
        self._close_popup()
        self._active_popup = popup
        self._active_popup_kind = kind
        popup.bind(
            "<Destroy>",
            lambda _event, owner=popup: self._clear_popup_if_current(owner),
            add="+",
        )
        if kind == "menu":
            popup.bind(
                "<Unmap>",
                lambda _event, owner=popup: self._unmap_popup_if_current(owner),
                add="+",
            )

    def _unmap_popup_if_current(self, popup) -> None:
        if popup is not self._active_popup or self._active_popup_kind != "menu":
            return
        self._active_popup = None
        self._active_popup_kind = None
        try:
            popup.destroy()
        except (AttributeError, tk.TclError):
            pass

    def _close_hover_popup_on_button(self, _event=None) -> None:
        if self._active_popup_kind == "hover":
            self._close_popup()

    def _close_hover_popup_on_focus(self, _event=None) -> None:
        if self._active_popup_kind == "hover":
            self._close_popup()

    def _show_popup_menu(self, menu: tk.Menu, event) -> str:
        """Post one transient menu and replace any prior popup first."""
        self._activate_popup(menu, "menu")
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            try:
                menu.grab_release()
            except tk.TclError:
                pass
        return "break"

    def _show_hover_help(self, event, help_text: str, owner=None) -> None:
        self._close_popup()

        popup = tk.Toplevel(self.root)
        popup.wm_overrideredirect(True)
        popup.attributes("-topmost", True)

        label = tk.Label(
            popup,
            text=help_text,
            padx=10,
            pady=6,
            relief="solid",
            borderwidth=1,
            background="#fff9e8",
            foreground="#2f2f2f",
        )
        label.pack()

        popup.update_idletasks()
        margin = 8
        anchor_x = event.widget.winfo_rootx()
        anchor_y = event.widget.winfo_rooty()
        preferred_x = anchor_x + 10
        preferred_y = anchor_y + event.widget.winfo_height() + 4
        popup_width = popup.winfo_reqwidth()
        popup_height = popup.winfo_reqheight()
        work_left, work_top, work_right, work_bottom = self._monitor_work_area(anchor_x, anchor_y)

        x = min(max(preferred_x, work_left + margin), max(work_left + margin, work_right - popup_width - margin))
        if preferred_y + popup_height + margin <= work_bottom:
            y = preferred_y
        else:
            y = anchor_y - popup_height - 4
        y = min(max(y, work_top + margin), max(work_top + margin, work_bottom - popup_height - margin))
        popup.geometry(f"+{x}+{y}")

        self._activate_popup(popup, "hover")
        self.help_popup = popup
        self.help_popup_label = label
        self.help_popup_owner = owner or event.widget
        self.help_popup_owner_widget = event.widget

    def _dismiss_hover_help_outside_owner(self, _event=None) -> None:
        """Dismiss a button help popup as soon as the pointer leaves its button."""
        if self.help_popup is None or self.help_popup_owner is None:
            return
        try:
            owner = self.help_popup_owner
            owner_widget = self.help_popup_owner_widget
            if isinstance(owner, tuple):
                pointer_x = self.root.winfo_pointerx()
                pointer_y = self.root.winfo_pointery()
                widget_left = owner_widget.winfo_rootx()
                widget_top = owner_widget.winfo_rooty()
                widget_right = widget_left + owner_widget.winfo_width()
                widget_bottom = widget_top + owner_widget.winfo_height()
                if not (widget_left <= pointer_x < widget_right and widget_top <= pointer_y < widget_bottom):
                    self._hide_hover_help()
                    return
                current_items = owner_widget.find_withtag("current")
                current_tags = set()
                for item_id in current_items:
                    current_tags.update(owner_widget.gettags(item_id))
                expected_tag = owner[1] if owner[0] == "workflow-node" else f"stage-{owner[1]}"
                if expected_tag not in current_tags:
                    self._hide_hover_help()
                return
            pointer_x = self.root.winfo_pointerx()
            pointer_y = self.root.winfo_pointery()
            owner_left = owner.winfo_rootx()
            owner_top = owner.winfo_rooty()
            owner_right = owner_left + owner.winfo_width()
            owner_bottom = owner_top + owner.winfo_height()
            if not (owner_left <= pointer_x < owner_right and owner_top <= pointer_y < owner_bottom):
                self._hide_hover_help()
        except (AttributeError, tk.TclError):
            self._hide_hover_help()

    @staticmethod
    def _monitor_work_area(x: int, y: int) -> tuple[int, int, int, int]:
        """Return the work area of the monitor containing a screen coordinate."""
        virtual_left = 0
        virtual_top = 0
        virtual_right = 0
        virtual_bottom = 0
        try:
            user32 = ctypes.windll.user32
            class Point(ctypes.Structure):
                _fields_ = [("x", ctypes.c_long), ("y", ctypes.c_long)]

            class Rect(ctypes.Structure):
                _fields_ = [
                    ("left", ctypes.c_long),
                    ("top", ctypes.c_long),
                    ("right", ctypes.c_long),
                    ("bottom", ctypes.c_long),
                ]

            class MonitorInfo(ctypes.Structure):
                _fields_ = [("cbSize", ctypes.c_uint), ("rcMonitor", Rect), ("rcWork", Rect), ("dwFlags", ctypes.c_uint)]

            point = Point(int(x), int(y))
            monitor = user32.MonitorFromPoint(point, 2)
            if monitor:
                info = MonitorInfo()
                info.cbSize = ctypes.sizeof(MonitorInfo)
                if user32.GetMonitorInfoW(monitor, ctypes.byref(info)):
                    return (info.rcWork.left, info.rcWork.top, info.rcWork.right, info.rcWork.bottom)
            virtual_left = user32.GetSystemMetrics(76)
            virtual_top = user32.GetSystemMetrics(77)
            virtual_right = virtual_left + user32.GetSystemMetrics(78)
            virtual_bottom = virtual_top + user32.GetSystemMetrics(79)
        except (AttributeError, OSError):
            pass

        if virtual_right <= virtual_left or virtual_bottom <= virtual_top:
            virtual_left = 0
            virtual_top = 0
            virtual_right = 1920
            virtual_bottom = 1080
        return (virtual_left, virtual_top, virtual_right, virtual_bottom)

    def _set_responsive_geometry(
        self,
        window: tk.Toplevel,
        width: int,
        height: int,
        min_width: Optional[int] = None,
        min_height: Optional[int] = None,
    ) -> None:
        """Apply screen-scaled, screen-bounded geometry to a secondary window."""
        scale = getattr(self, "layout_scale", 1.0)
        screen_width = window.winfo_screenwidth()
        screen_height = window.winfo_screenheight()
        target_width = min(round(width * scale), max(320, round(screen_width * 0.94)))
        target_height = min(round(height * scale), max(240, round(screen_height * 0.90)))
        window.geometry(f"{target_width}x{target_height}")
        if min_width is not None or min_height is not None:
            window.minsize(
                min(round((min_width if min_width is not None else width) * scale), target_width),
                min(round((min_height if min_height is not None else height) * scale), target_height),
            )

    def _hide_hover_help(self) -> None:
        if self.help_popup is self._active_popup or self._active_popup_kind == "hover":
            self._close_popup()

    def _build_ui(self) -> None:
        header = ttk.Frame(self.root, padding=12)
        header.pack(fill=tk.X)

        ttk.Label(
            header,
            text="Requirement AI workflow",
            font=("Segoe UI", 18, "bold"),
        ).pack(anchor=tk.W)

        ttk.Label(
            header,
            text="Run pipeline commands through workflow_cli.py and inspect extracted requirements.",
        ).pack(anchor=tk.W, pady=(4, 0))

        save_layout_btn = ttk.Button(
            header,
            text="Save Window Layout",
            command=self._save_gui_layout_now,
            width=18,
        )
        save_layout_btn.pack(side=tk.RIGHT, anchor=tk.N)
        self._bind_button_help(
            save_layout_btn,
            "Save the current window size, position, and pane settings for the next session.",
        )

        self.ai_usage_var = tk.StringVar()

        self.main_vertical_pane = ttk.Panedwindow(self.root, orient=tk.VERTICAL)
        self.main_vertical_pane.pack(fill=tk.BOTH, expand=True, padx=12, pady=(2, 8))

        top = ttk.Frame(self.main_vertical_pane, padding=(0, 0, 0, 0))
        self.main_vertical_pane.add(top, weight=1)

        self.top_horizontal_pane = ttk.Panedwindow(top, orient=tk.HORIZONTAL)
        top_horizontal_pane = self.top_horizontal_pane
        top_horizontal_pane.pack(fill=tk.BOTH, expand=True)

        cli_group = ttk.LabelFrame(top_horizontal_pane, text="workflow_cli Runs", padding=6)
        top_horizontal_pane.add(cli_group, weight=3)

        cli_canvas = tk.Canvas(cli_group, highlightthickness=0, height=94, xscrollincrement=20)
        cli_canvas.pack(fill=tk.BOTH, expand=True)
        cli_hscroll = tk.Scrollbar(
            cli_group,
            orient=tk.HORIZONTAL,
            command=cli_canvas.xview,
            width=18,
            relief=tk.SUNKEN,
            bd=1,
            highlightthickness=1,
            bg="#607d8b",
            troughcolor="#d9e2ec",
            activebackground="#2f6f9f",
        )
        cli_hscroll.pack(fill=tk.X, pady=(4, 0))
        cli_canvas.configure(xscrollcommand=cli_hscroll.set)

        cli_content = ttk.Frame(cli_canvas)
        cli_canvas_window = cli_canvas.create_window((0, 0), window=cli_content, anchor="nw")

        def update_cli_scrollregion(_event=None) -> None:
            cli_canvas.configure(scrollregion=cli_canvas.bbox("all"))

        def update_cli_window_height(event) -> None:
            cli_canvas.itemconfigure(cli_canvas_window, height=event.height)

        cli_content.bind("<Configure>", update_cli_scrollregion)
        cli_canvas.bind("<Configure>", update_cli_window_height)
        cli_canvas.bind(
            "<Shift-MouseWheel>",
            lambda event: cli_canvas.xview_scroll(-1 if event.delta > 0 else 1, "units"),
            add="+",
        )

        row = ttk.Frame(cli_content)
        row.pack(fill=tk.X)

        ttk.Label(row, text="Single stage:").pack(side=tk.LEFT, padx=(0, 8))

        self.stage_var = tk.StringVar(value="2")
        self.stage_combo = ttk.Combobox(
            row,
            textvariable=self.stage_var,
            values=SINGLE_STAGE_LABELS,
            state="readonly",
            width=42,
        )
        self.stage_var.set(SINGLE_STAGE_LABELS[2])
        self.stage_combo.pack(side=tk.LEFT)

        run_selected = ttk.Button(
            row,
            text="Run Selected Stage",
            command=self.run_selected_stage_cli,
        )
        run_selected.pack(side=tk.LEFT, padx=(10, 0))
        self._bind_button_help(run_selected, "Run workflow_cli for one selected stage (deterministic script call).")

        range_row = ttk.Frame(cli_content)
        range_row.pack(fill=tk.X, pady=(6, 0))

        ttk.Label(range_row, text="Range:").pack(side=tk.LEFT, padx=(0, 8))

        self.from_stage_var = tk.StringVar(value="0")
        self.from_stage_combo = ttk.Combobox(
            range_row,
            textvariable=self.from_stage_var,
            values=RANGE_STAGE_LABELS,
            state="readonly",
            width=42,
        )
        self.from_stage_var.set(RANGE_STAGE_LABELS[0])
        self.from_stage_combo.pack(side=tk.LEFT)

        ttk.Label(range_row, text="to").pack(side=tk.LEFT, padx=(8, 8))

        self.to_stage_var = tk.StringVar(value="2")
        self.to_stage_combo = ttk.Combobox(
            range_row,
            textvariable=self.to_stage_var,
            values=RANGE_STAGE_LABELS,
            state="readonly",
            width=42,
        )
        self.to_stage_var.set(RANGE_STAGE_LABELS[2])
        self.to_stage_combo.pack(side=tk.LEFT)

        run_range_btn = ttk.Button(
            range_row,
            text="Run Stage Range",
            command=self.run_stage_range_cli,
        )
        run_range_btn.pack(side=tk.LEFT, padx=(10, 0))
        self._bind_button_help(run_range_btn, "Run workflow_cli from start stage to end stage (inclusive).")

        cli_actions = ttk.Frame(cli_content)
        cli_actions.pack(fill=tk.X, pady=(6, 0))

        validate_btn = ttk.Button(
            cli_actions,
            text="Validate All Gates",
            command=self.run_validate_all,
            width=18,
        )
        validate_btn.pack(side=tk.LEFT, padx=(0, 8))
        self._bind_button_help(validate_btn, "Validate all stage gates using workflow_cli.")

        status_btn = ttk.Button(
            cli_actions,
            text="Show Status",
            command=self.run_status,
            width=14,
        )
        status_btn.pack(side=tk.LEFT, padx=(0, 8))
        self._bind_button_help(status_btn, "Show current workflow status and recent artifacts.")

        taxonomy_btn = ttk.Button(
            cli_actions,
            text="Taxonomy Update",
            command=self.run_taxonomy_update,
            width=16,
        )
        taxonomy_btn.pack(side=tk.LEFT)
        self._bind_button_help(taxonomy_btn, "Run deterministic taxonomy crosscheck/update from Stage 1 OCR index.")

        source_spec_btn = ttk.Button(
            cli_actions,
            text="Source Specification...",
            command=self._show_source_spec_dialog,
            width=22,
        )
        source_spec_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(
            source_spec_btn,
            "Choose a primary source or stage a supplementary source for review.",
        )

        cli_secondary_actions = ttk.Frame(cli_content)
        cli_secondary_actions.pack(fill=tk.X, pady=(5, 0))

        cli_secondary_row1 = ttk.Frame(cli_secondary_actions)
        cli_secondary_row1.pack(fill=tk.X)
        cli_secondary_row2 = ttk.Frame(cli_secondary_actions)
        cli_secondary_row2.pack(fill=tk.X, pady=(5, 0))
        cli_secondary_row3 = ttk.Frame(cli_secondary_actions)
        cli_secondary_row3.pack(fill=tk.X, pady=(5, 0))

        architecture_map_btn = ttk.Button(
            cli_secondary_row1,
            text="Architecture Map Review",
            command=self._show_architecture_map_review_dialog,
            width=24,
        )
        architecture_map_btn.pack(side=tk.LEFT)
        self._bind_button_help(
            architecture_map_btn,
            "Open the architecture mapping CSV/Markdown preview and current review decision status.",
        )

        review_context_btn = ttk.Button(
            cli_secondary_row1,
            text="Review Context",
            command=self._show_review_context,
            width=16,
        )
        review_context_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(review_context_btn, "Show revision, snapshot, impact, source, vocabulary, taxonomy, and diagnostic status.")

        refresh_context_btn = ttk.Button(
            cli_secondary_row1,
            text="Refresh Context",
            command=self._refresh_review_context,
            width=18,
        )
        refresh_context_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(
            refresh_context_btn,
            "Reload revision, approved snapshot, impact, allocation, source, profile, and diagnostic status.",
        )

        scanner_btn = ttk.Button(
            cli_secondary_row1,
            text="Authority Scan",
            command=self.run_authority_scan,
            width=16,
        )
        scanner_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(scanner_btn, "Run the diagnostic authority consistency scan and open its JSON result from the Work Area.")

        table_row_review_btn = ttk.Button(
            cli_secondary_row2,
            text="Table Row Review",
            command=self._open_table_row_review_csv,
            width=20,
        )
        table_row_review_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(
            table_row_review_btn,
            "Open the Stage 1 table-row review CSV in Excel at any time.",
        )

        drs_req_only_btn = ttk.Button(
            cli_secondary_row2,
            text="S5 after S2a (Req-only)",
            command=self.run_stage5_after_stage2a,
            width=22,
        )
        drs_req_only_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(
            drs_req_only_btn,
            "Run Stage 5 directly after Stage 2A using Stage 1 + Stage 2A artifacts only.",
        )

        sysml_menu = tk.Menu(self.root, tearoff=False)
        sysml_menu.add_command(label="Generate from approved snapshot", command=self.regenerate_sysml)
        sysml_menu.add_command(label="Open SysML Architecture", command=lambda: self.work_tabs.select(self.sysml_architecture_tab))
        regenerate_sysml_btn = ttk.Menubutton(
            cli_secondary_row2,
            text="SysML...",
            menu=sysml_menu,
            width=18,
        )
        regenerate_sysml_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(
            regenerate_sysml_btn,
            "Generate the hierarchical SysML model from an approved snapshot or open the SysML Architecture view.",
        )

        report_btn = ttk.Button(
            cli_secondary_row3,
            text="Traceability Reports...",
            command=self._show_snapshot_report_menu,
            width=20,
        )
        report_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(report_btn, "Generate hierarchy-level XLSX reports or one inclusive, filterable workbook covering imported and generated requirements.")

        conflict_btn = ttk.Button(
            cli_secondary_row3,
            text="Conflict Review",
            command=self._open_conflict_review,
            width=16,
        )
        conflict_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(conflict_btn, "Open the latest derived same-req-ID conflict review report.")

        downstream_validation_btn = ttk.Button(
            cli_secondary_row3,
            text="Validate Downstream",
            command=self.run_validate_downstream,
            width=20,
        )
        downstream_validation_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(
            downstream_validation_btn,
            "Run the local downstream retrieval, preservation, mapping, hierarchy, and rendering validation. Recommendations require approval and are never applied automatically.",
        )

        work_area_group = ttk.LabelFrame(top_horizontal_pane, text="Work Area", padding=6)
        top_horizontal_pane.add(work_area_group, weight=4)

        work_head = ttk.Frame(work_area_group)
        work_head.pack(fill=tk.X, pady=(0, 3))
        ttk.Label(work_head, text="Select a repository file to open it read-only.").pack(side=tk.LEFT)
        refresh_work_btn = ttk.Button(work_head, text="Refresh", command=self._refresh_all_derived_views, width=10)
        refresh_work_btn.pack(side=tk.RIGHT)
        self._bind_button_help(refresh_work_btn, "Refresh repository files and reload System Traceability from the latest generated artifacts.")

        self.work_tabs = ttk.Notebook(work_area_group)
        self.work_tabs.pack(fill=tk.BOTH, expand=True)
        self.work_browser_tab = ttk.Frame(self.work_tabs)
        self.system_traceability_tab = ttk.Frame(self.work_tabs, padding=8)
        self.sysml_architecture_tab = ttk.Frame(self.work_tabs, padding=8)
        self.work_tabs.add(self.work_browser_tab, text="Artifacts")
        self.work_tabs.add(self.system_traceability_tab, text="System Traceability")
        self.work_tabs.add(self.sysml_architecture_tab, text="SysML Architecture")
        self.work_tabs.bind("<<NotebookTabChanged>>", self._restore_sysml_layout_when_selected, add="+")

        self.work_pane = ttk.Panedwindow(self.work_browser_tab, orient=tk.HORIZONTAL)
        work_pane = self.work_pane
        work_pane.pack(fill=tk.BOTH, expand=True)
        navigation_frame = ttk.Frame(work_pane)
        viewer_frame = ttk.Frame(work_pane)
        work_pane.add(navigation_frame, weight=1)
        work_pane.add(viewer_frame, weight=3)

        self.work_tree = ttk.Treeview(navigation_frame, show="tree", selectmode="browse")
        work_tree_y = ttk.Scrollbar(navigation_frame, orient=tk.VERTICAL, command=self.work_tree.yview)
        work_tree_x_frame = tk.Frame(navigation_frame, background="#d9e2ec", height=30)
        work_tree_x_frame.configure(height=26)
        work_tree_x_frame.grid_propagate(False)
        work_tree_x = tk.Scrollbar(
            work_tree_x_frame,
            orient=tk.HORIZONTAL,
            command=self.work_tree.xview,
            width=18,
            relief=tk.SUNKEN,
            bd=1,
            highlightthickness=1,
            bg="#607d8b",
            troughcolor="#d9e2ec",
            activebackground="#2f6f9f",
        )
        self.work_tree.configure(yscrollcommand=work_tree_y.set, xscrollcommand=work_tree_x.set)
        self.work_tree.grid(row=0, column=0, sticky="nsew")
        work_tree_y.grid(row=0, column=1, sticky="ns")
        work_tree_x_frame.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(4, 0))
        work_tree_x.pack(fill=tk.X, expand=True, padx=1, pady=2)
        navigation_frame.grid_rowconfigure(0, weight=1)
        navigation_frame.grid_rowconfigure(1, weight=0, minsize=30)
        navigation_frame.grid_columnconfigure(0, weight=1)
        self.work_tree.bind("<<TreeviewSelect>>", self._on_work_area_select)
        self.work_tree.bind("<Button-3>", self._show_work_tree_context_menu)
        self.work_tree.bind(
            "<Shift-MouseWheel>",
            lambda event: self.work_tree.xview_scroll(-1 if event.delta > 0 else 1, "units"),
            add="+",
        )

        self.work_file_var = tk.StringVar(value="No file selected")
        ttk.Label(viewer_frame, textvariable=self.work_file_var, anchor="w").pack(fill=tk.X, pady=(0, 4))

        self.work_find_var = tk.StringVar()
        self.work_find_status_var = tk.StringVar(value="")
        self.work_find_frame = ttk.Frame(viewer_frame)
        ttk.Label(self.work_find_frame, text="Find:").pack(side=tk.LEFT, padx=(0, 4))
        self.work_find_entry = ttk.Entry(self.work_find_frame, textvariable=self.work_find_var, width=32)
        self.work_find_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.work_find_entry.bind("<Return>", lambda _e: self._find_next_in_work_view())
        self.work_find_entry.bind("<Shift-Return>", lambda _e: self._find_previous_in_work_view())
        self.work_find_entry.bind("<Escape>", lambda _e: self._close_work_find())
        self.work_find_entry.bind("<KeyRelease>", self._on_work_find_changed)
        ttk.Button(self.work_find_frame, text="Prev", command=self._find_previous_in_work_view, width=7).pack(side=tk.LEFT, padx=(6, 0))
        ttk.Button(self.work_find_frame, text="Next", command=self._find_next_in_work_view, width=7).pack(side=tk.LEFT, padx=(4, 0))
        ttk.Button(self.work_find_frame, text="Close", command=self._close_work_find, width=7).pack(side=tk.LEFT, padx=(4, 0))
        ttk.Label(self.work_find_frame, textvariable=self.work_find_status_var, width=14, anchor="e").pack(side=tk.LEFT, padx=(8, 0))

        self.work_text_frame = ttk.Frame(viewer_frame)
        self.work_text_frame.pack(fill=tk.BOTH, expand=True)
        self.work_viewer = tk.Text(self.work_text_frame, wrap=tk.WORD, font=("Consolas", 11))
        self.work_viewer.tag_configure("comment", foreground="#6a737d")
        self.work_viewer.tag_configure("string", foreground="#a31515")
        self.work_viewer.tag_configure("keyword", foreground="#0000cc", font=("Consolas", 11, "bold"))
        self.work_viewer.tag_configure("number", foreground="#098658")
        self.work_viewer.tag_configure("heading", foreground="#7a3e00", font=("Consolas", 11, "bold"))
        self.work_viewer.tag_configure("link", foreground="#0563c1", underline=True)
        self.work_viewer.tag_configure("delimiter", foreground="#8a6d3b")
        self.work_viewer.tag_configure("find_hit", background="#fff2a8")
        self.work_viewer.tag_configure("find_current", background="#ffbf47")
        viewer_y = ttk.Scrollbar(self.work_text_frame, orient=tk.VERTICAL, command=self.work_viewer.yview)
        self.work_text_x_frame = ttk.Frame(viewer_frame, padding=(0, 4, 0, 0))
        self.work_text_x_frame.configure(height=26)
        self.work_text_x_frame.pack_propagate(False)
        viewer_x = tk.Scrollbar(
            self.work_text_x_frame,
            orient=tk.HORIZONTAL,
            command=self.work_viewer.xview,
            width=18,
            relief=tk.SUNKEN,
            bd=1,
            highlightthickness=1,
            bg="#607d8b",
            troughcolor="#d9e2ec",
            activebackground="#2f6f9f",
        )
        self.work_viewer.configure(yscrollcommand=viewer_y.set, xscrollcommand=viewer_x.set)
        self.work_viewer.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self._bind_copy_support(self.work_viewer)
        self.work_viewer.bind("<KeyPress>", self._block_rendered_text_edit, add="+")
        viewer_y.pack(side=tk.RIGHT, fill=tk.Y)
        self.work_text_x_frame.pack(fill=tk.X)
        viewer_x.pack(fill=tk.X, expand=True)
        self.work_viewer.bind(
            "<Shift-MouseWheel>",
            lambda event: self.work_viewer.xview_scroll(-1 if event.delta > 0 else 1, "units"),
            add="+",
        )

        self.work_csv_frame = ttk.Frame(viewer_frame)
        self.work_csv_table = ttk.Treeview(self.work_csv_frame, show="headings", selectmode="browse")
        csv_y = ttk.Scrollbar(self.work_csv_frame, orient=tk.VERTICAL, command=self.work_csv_table.yview)
        csv_x = tk.Scrollbar(
            self.work_csv_frame,
            orient=tk.HORIZONTAL,
            command=self.work_csv_table.xview,
            width=18,
            relief=tk.SUNKEN,
            bd=1,
            highlightthickness=1,
            bg="#607d8b",
            troughcolor="#d9e2ec",
            activebackground="#2f6f9f",
        )
        self.work_csv_table.configure(yscrollcommand=csv_y.set, xscrollcommand=csv_x.set)
        self.work_csv_table.grid(row=0, column=0, sticky="nsew")
        csv_y.grid(row=0, column=1, sticky="ns")
        csv_x.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(4, 0))
        self.work_csv_frame.grid_rowconfigure(0, weight=1)
        self.work_csv_frame.grid_columnconfigure(0, weight=1)
        self.work_csv_table.tag_configure("find_hit", background="#fff2a8")
        self.work_csv_table.tag_configure("find_current", background="#ffbf47")
        self.work_csv_table.bind("<Button-3>", self._show_csv_context_menu)
        self.work_csv_table.bind("<Double-1>", self._sort_csv_by_heading, add="+")
        self.work_csv_table.bind(
            "<Shift-MouseWheel>",
            lambda event: self.work_csv_table.xview_scroll(-1 if event.delta > 0 else 1, "units"),
            add="+",
        )
        self.current_work_path = None
        self.current_csv_rows = []
        self.current_csv_headers = []
        self.csv_sort_directions = {}
        self.work_find_results = []
        self.work_find_index = -1
        self.root.bind_all("<Control-f>", self._open_work_find, add="+")
        self.root.bind_all("<Control-F>", self._open_work_find, add="+")

        self.refresh_work_area()
        self._build_system_traceability_tab()
        self._build_sysml_architecture_tab()
        # The compact workflow remains in the lower visual panel.

        visuals = ttk.Frame(self.main_vertical_pane, padding=(0, 0, 0, 0))
        self.main_vertical_pane.add(visuals, weight=5)

        split_hint = ttk.Label(
            visuals,
            text="Drag the center divider to enlarge Workflow Diagram (left) or Live Chat Console (right).",
        )
        split_hint.pack(fill=tk.X, pady=(0, 6))

        self.visual_pane = ttk.Panedwindow(visuals, orient=tk.HORIZONTAL)
        self.visual_pane.pack(fill=tk.BOTH, expand=True)

        left_panel = ttk.Frame(self.visual_pane)
        right_panel = ttk.Frame(self.visual_pane)
        self.visual_pane.add(left_panel, weight=3)
        self.visual_pane.add(right_panel, weight=2)

        diagram_group = ttk.LabelFrame(left_panel, text="Compact Workflow Diagram", padding=8)
        diagram_group.pack(fill=tk.BOTH, expand=True, padx=(0, 4))

        self.diagram_canvas = tk.Canvas(
            diagram_group,
            height=210,
            bg="#ffffff",
            highlightthickness=0,
            xscrollincrement=20,
        )
        self.diagram_canvas.pack(fill=tk.BOTH, expand=True)
        self.diagram_vscroll = tk.Scrollbar(
            diagram_group,
            orient=tk.VERTICAL,
            command=self.diagram_canvas.yview,
            width=18,
            relief=tk.SUNKEN,
            bd=1,
            highlightthickness=1,
            bg="#607d8b",
            troughcolor="#d9e2ec",
            activebackground="#2f6f9f",
        )
        self.diagram_vscroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.diagram_canvas.bind("<Configure>", self._on_diagram_resize)
        self.diagram_canvas.bind(
            "<MouseWheel>",
            lambda event: self.diagram_canvas.yview_scroll(-1 if event.delta > 0 else 1, "units"),
            add="+",
        )
        self.diagram_canvas.bind(
            "<Shift-MouseWheel>",
            lambda event: self.diagram_canvas.xview_scroll(-1 if event.delta > 0 else 1, "units"),
            add="+",
        )
        self.diagram_canvas.bind("<Control-MouseWheel>", self._zoom_compact_workflow_wheel, add="+")

        self.diagram_hscroll_frame = ttk.Frame(diagram_group, padding=(0, 4, 0, 0))
        self.diagram_hscroll_frame.configure(height=26)
        self.diagram_hscroll_frame.pack_propagate(False)

        self.diagram_hscroll = tk.Scrollbar(
            self.diagram_hscroll_frame,
            orient=tk.HORIZONTAL,
            command=self.diagram_canvas.xview,
            width=18,
            relief=tk.SUNKEN,
            bd=1,
            highlightthickness=1,
            bg="#607d8b",
            troughcolor="#d9e2ec",
            activebackground="#2f6f9f",
        )
        self.diagram_canvas.configure(
            xscrollcommand=self.diagram_hscroll.set,
            yscrollcommand=self.diagram_vscroll.set,
        )

        self.diagram_hint_var = tk.StringVar(value="Hover a stage to see its meaning. Click to run.")
        ttk.Label(diagram_group, textvariable=self.diagram_hint_var, wraplength=560).pack(fill=tk.X, pady=(6, 0))
        self.diagram_hscroll_frame.pack(fill=tk.X)
        self.diagram_hscroll.pack(fill=tk.X, expand=True)

        self._draw_workflow_diagram()

        chat_group = ttk.LabelFrame(right_panel, text="Live Chat Console", padding=8)
        chat_group.pack(fill=tk.BOTH, expand=True, padx=(4, 0))

        chat_mode_row = ttk.Frame(chat_group)
        chat_mode_row.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(chat_mode_row, text="Mode:").pack(side=tk.LEFT)
        self.chat_mode_var = tk.StringVar(value="Auto")
        self.chat_mode_combo = ttk.Combobox(
            chat_mode_row,
            textvariable=self.chat_mode_var,
            values=["Auto", "Command", "AI"],
            state="readonly",
            width=10,
        )
        self.chat_mode_combo.pack(side=tk.LEFT, padx=(6, 10))

        ttk.Label(chat_mode_row, text="Backend:").pack(side=tk.LEFT)
        self.ai_backend_var = tk.StringVar(value="OpenAI")
        self.ai_backend_combo = ttk.Combobox(
            chat_mode_row,
            textvariable=self.ai_backend_var,
            values=["Auto", "Local", "Azure", "OpenAI"],
            state="readonly",
            width=10,
        )
        self.ai_backend_combo.pack(side=tk.LEFT, padx=(6, 10))
        self.ai_backend_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_ai_controls())

        chat_api_row = ttk.Frame(chat_group)
        chat_api_row.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(chat_api_row, text="OpenAI API key:").pack(side=tk.LEFT)
        self.openai_api_key_var = tk.StringVar(value=os.getenv("OPENAI_API_KEY", ""))
        self.openai_api_key_entry = ttk.Entry(chat_api_row, textvariable=self.openai_api_key_var, show="*", width=42)
        self.openai_api_key_entry.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)

        ttk.Label(chat_mode_row, text="Local model:").pack(side=tk.LEFT)
        self.local_model_var = tk.StringVar(value="")
        self.local_model_combo = ttk.Combobox(
            chat_mode_row,
            textvariable=self.local_model_var,
            values=[],
            state="disabled",
            width=22,
        )
        self.local_model_combo.pack(side=tk.LEFT, padx=(6, 6), fill=tk.X, expand=True)

        refresh_local_btn = ttk.Button(chat_mode_row, text="Refresh Local", command=self.refresh_local_models, width=12)
        refresh_local_btn.pack(side=tk.LEFT)
        self._bind_button_help(refresh_local_btn, "Refresh available local model list from env and config/local_models.json.")

        self.chat_vertical_pane = ttk.Panedwindow(chat_group, orient=tk.VERTICAL)
        chat_vertical_pane = self.chat_vertical_pane
        chat_vertical_pane.pack(fill=tk.BOTH, expand=True)

        chat_text_frame = ttk.Frame(chat_vertical_pane)
        chat_input_row = ttk.Frame(chat_vertical_pane, padding=(0, 6, 0, 0))
        chat_input_row.configure(height=38)
        chat_input_row.pack_propagate(False)
        chat_vertical_pane.add(chat_text_frame, weight=4)
        chat_vertical_pane.add(chat_input_row, weight=0)

        self.chat_console = tk.Text(chat_text_frame, wrap=tk.NONE, height=11, font=("Consolas", 11))
        self.chat_console.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        chat_y_scroll = ttk.Scrollbar(chat_text_frame, orient=tk.VERTICAL, command=self.chat_console.yview)
        chat_y_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        chat_x_scroll = ttk.Scrollbar(chat_text_frame, orient=tk.HORIZONTAL, command=self.chat_console.xview)
        chat_x_scroll.pack(fill=tk.X, pady=(4, 0))

        self.chat_console.config(yscrollcommand=chat_y_scroll.set, xscrollcommand=chat_x_scroll.set)
        self.chat_console.insert(
            tk.END,
            "GUI commands: /help, /run <stage>, run stage <start> to <end>, /status, "
            "/validate, /taxonomy, /analyze log, /stop, /clear\n",
        )
        self.chat_console.insert(tk.END, "Type /help for examples and AI mode details.\n")
        self._load_chat_log()

        ttk.Label(chat_input_row, text="Chat:").pack(side=tk.LEFT, padx=(0, 6))
        self.chat_input_var = tk.StringVar()
        self.chat_entry = ttk.Entry(chat_input_row, textvariable=self.chat_input_var)
        self.chat_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.chat_entry.bind("<Return>", self._on_chat_enter)

        send_btn = ttk.Button(chat_input_row, text="Send", command=self.send_chat_message, width=10)
        send_btn.pack(side=tk.LEFT, padx=(8, 0))
        self._bind_button_help(send_btn, "Send chat input to command handler or selected AI mode.")

        self.refresh_local_models()

        execution_panel = ttk.Frame(self.main_vertical_pane, padding=(0, 0, 0, 0))
        self.main_vertical_pane.add(execution_panel, weight=4)

        controls = ttk.Frame(execution_panel, padding=(0, 0, 0, 8))
        controls.pack(fill=tk.X)

        self.status_var = tk.StringVar(value="Idle")
        self.review_context_var = tk.StringVar(value=format_gui_review_context(build_gui_review_context(REPO_ROOT)))
        ttk.Label(controls, textvariable=self.status_var).pack(side=tk.LEFT)
        self.review_context_viewer = tk.Text(
            controls,
            height=2,
            width=92,
            wrap=tk.WORD,
            relief=tk.FLAT,
            borderwidth=0,
            highlightthickness=0,
            background=self.root.cget("background"),
            font=("Segoe UI", 9),
        )
        self.review_context_viewer.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(14, 0))
        self._bind_copy_support(self.review_context_viewer)
        self.review_context_viewer.bind("<KeyPress>", self._block_rendered_text_edit, add="+")
        self.review_context_var.trace_add("write", lambda *_args: self._refresh_review_context_viewer())
        self._refresh_review_context_viewer()

        self.verbose_trace_var = tk.BooleanVar(value=self.verbose_trace_enabled)
        self.verbose_trace_btn = ttk.Button(
            controls,
            text=self._verbose_trace_button_text(),
            command=self._toggle_verbose_trace,
            width=18,
        )
        self.verbose_trace_btn.pack(side=tk.LEFT, padx=(14, 0))
        self._bind_button_help(
            self.verbose_trace_btn,
            "Toggle detailed GUI activity entries in the Execution Log. Enabled by default.",
        )

        clear_log_btn = ttk.Button(
            controls,
            text="Clear Log",
            command=self.clear_log,
            width=10,
        )
        clear_log_btn.pack(side=tk.RIGHT, padx=(8, 0))
        self._bind_button_help(clear_log_btn, "Clear the execution log panel.")

        stop_btn = ttk.Button(
            controls,
            text="Stop Current",
            command=self.stop_current,
            width=12,
        )
        stop_btn.pack(side=tk.RIGHT)
        self._bind_button_help(stop_btn, "Terminate the currently running subprocess.")

        execution_content = ttk.Frame(execution_panel)
        execution_content.pack(fill=tk.BOTH, expand=True)

        log_frame = ttk.LabelFrame(execution_content, text="Execution Log", padding=8)
        log_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        usage_panel = ttk.LabelFrame(execution_content, text="AI Usage", padding=(8, 4), width=240)
        usage_panel.pack(side=tk.RIGHT, fill=tk.Y, padx=(8, 0))
        usage_panel.pack_propagate(False)
        ttk.Label(
            usage_panel,
            textvariable=self.ai_usage_var,
            anchor="nw",
            justify=tk.LEFT,
            wraplength=214,
        ).pack(fill=tk.BOTH, expand=True)

        log_text_frame = ttk.Frame(log_frame)
        log_text_frame.pack(fill=tk.BOTH, expand=True)

        self.log_font = tkfont.Font(family="Consolas", size=11)
        self.log = tk.Text(log_text_frame, wrap=tk.NONE, font=self.log_font)
        self.log.tag_configure("ipos_activity", background="#e8f5e9", foreground="#1b5e20")
        self.log.tag_configure("ipos_run_boundary", background="#c8e6c9", foreground="#14532d")
        self.log.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        log_y_scroll = ttk.Scrollbar(log_text_frame, orient=tk.VERTICAL, command=self.log.yview)
        log_y_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        log_x_scroll = ttk.Scrollbar(log_frame, orient=tk.HORIZONTAL, command=self.log.xview)
        log_x_scroll.pack(fill=tk.X, pady=(4, 0))

        self.log.config(yscrollcommand=log_y_scroll.set, xscrollcommand=log_x_scroll.set)
        self.log.bind("<Control-MouseWheel>", self._zoom_execution_log_wheel, add="+")
        self.log.bind("<Shift-MouseWheel>", lambda event: self.log.xview_scroll(-1 if event.delta > 0 else 1, "units"), add="+")

        self._refresh_ai_usage_panel()

    def _build_workflow_diagram_tab(self) -> None:
        """Build the full workflow graph in a dedicated work-area tab."""
        header = ttk.Frame(self.workflow_diagram_tab)
        header.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(
            header,
            text="Complete workflow: main stages, supplementary review, rework paths, and loopbacks.",
        ).pack(side=tk.LEFT)
        refresh_btn = ttk.Button(header, text="Refresh Diagram", command=self._refresh_workflow_diagram)
        refresh_btn.pack(side=tk.RIGHT)
        self._bind_button_help(refresh_btn, "Redraw the shared workflow definition after stage-model updates.")
        diagram_frame = ttk.Frame(self.workflow_diagram_tab)
        diagram_frame.pack(fill=tk.BOTH, expand=True)
        self.workflow_diagram_renderer = WorkflowDiagramRenderer(
            diagram_frame,
            on_node_click=self._on_full_workflow_node_click,
            selectable_nodes=SELECTABLE_WORKFLOW_NODE_IDS,
        )

    def _refresh_workflow_diagram(self) -> None:
        renderer = getattr(self, "workflow_diagram_renderer", None)
        if renderer is not None:
            renderer.refresh()
        self._draw_workflow_diagram()
        self.append_log("[GUI] Workflow diagram refreshed from the shared graph definition.")

    def _zoom_compact_workflow_wheel(self, event):
        direction = 0.1 if event.delta > 0 else -0.1
        self.compact_workflow_zoom = min(2.0, max(0.6, round(self.compact_workflow_zoom + direction, 2)))
        self._draw_workflow_diagram()
        return "break"

    def _on_full_workflow_node_click(self, node_id: str) -> None:
        """Route clickable graph nodes to the existing GUI actions."""
        if node_id == "supp_import":
            self._show_source_spec_dialog()
            return
        stage = WORKFLOW_NODE_STAGE_KEYS.get(node_id)
        if stage:
            self.run_stage_cli(stage)

    def ensure_visible(self) -> None:
        # Restore the saved layout and clamp it to the work area of its current monitor.
        self.root.update_idletasks()
        geometry_match = re.match(r"^(\d+)x(\d+)(?:([+-]\d+)([+-]\d+))?$", self.root.geometry())
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        if geometry_match:
            width, height = int(geometry_match.group(1)), int(geometry_match.group(2))
            x = int(geometry_match.group(3)) if geometry_match.group(3) else max(0, (screen_width - width) // 2)
            y = int(geometry_match.group(4)) if geometry_match.group(4) else max(0, (screen_height - height) // 3)
        else:
            width = round(1180 * getattr(self, "layout_scale", 1.0))
            height = round(860 * getattr(self, "layout_scale", 1.0))
            x = max(0, (screen_width - width) // 2)
            y = max(0, (screen_height - height) // 3)

        work_left, work_top, work_right, work_bottom = self._monitor_work_area(x + width // 2, y + height // 2)
        available_width = max(320, work_right - work_left)
        available_height = max(240, work_bottom - work_top)
        width = min(width, available_width)
        height = min(height, available_height)
        x = min(max(x, work_left), work_right - width)
        y = min(max(y, work_top), work_bottom - height)
        self.root.minsize(min(round(980 * self.layout_scale), width), min(round(700 * self.layout_scale), height))
        self.root.geometry(f"{width}x{height}+{x}+{y}")
        self.root.state("normal")
        self.root.deiconify()
        self._gui_layout_ready = True
        self.root.lift()
        self.root.attributes("-topmost", True)
        self.root.after(350, lambda: self.root.attributes("-topmost", False))
        self.root.focus_force()

    def append_log(self, text: str) -> None:
        if text.startswith("[GUI]") and "[TRACE]" not in text and not self.verbose_trace_enabled:
            return
        if text.startswith("[GUI]"):
            text = f"[{datetime.now():%H:%M:%S}] {text}"
        self.root.after(0, self._append_log_ui, text)

    def _verbose_trace_button_text(self) -> str:
        return f"Verbose Trace: {'ON' if self.verbose_trace_enabled else 'OFF'}"

    def _toggle_verbose_trace(self) -> None:
        self.verbose_trace_enabled = not self.verbose_trace_enabled
        self.verbose_trace_var.set(self.verbose_trace_enabled)
        self.verbose_trace_btn.configure(text=self._verbose_trace_button_text())
        state = "enabled" if self.verbose_trace_enabled else "disabled"
        self.root.after(0, self._append_log_ui, f"[TRACE] Verbose GUI trace {state}.")

    def _append_log_ui(self, text: str) -> None:
        visible_text = text.lstrip()
        if visible_text.startswith("[IPOS_ACTIVITY]"):
            tags = ("ipos_activity",)
        elif visible_text.startswith(("=== START Stage 6 Digital IPOS", "=== END Stage 6 Digital IPOS")):
            tags = ("ipos_run_boundary",)
        else:
            tags = ()
        self.log.insert(tk.END, text + "\n", tags)
        self.log.see(tk.END)

    def append_chat(self, text: str) -> None:
        try:
            LOG_DIR.mkdir(parents=True, exist_ok=True)
            with CHAT_LOG.open("a", encoding="utf-8") as chat_log:
                chat_log.write(f"[{datetime.now().astimezone().isoformat(timespec='seconds')}] {text}\n")
        except OSError as exc:
            self.append_log(f"[WARN] Could not write local chat log: {exc}")
        self.root.after(0, self._append_chat_ui, text)

    def _load_chat_log(self) -> None:
        try:
            if not CHAT_LOG.exists():
                return
            lines = CHAT_LOG.read_text(encoding="utf-8", errors="replace").splitlines()
            if lines:
                self.chat_console.insert(tk.END, "\nLocal chat log (latest entries):\n")
                self.chat_console.insert(tk.END, "\n".join(lines[-200:]) + "\n")
                self.chat_console.see(tk.END)
        except OSError as exc:
            self.chat_console.insert(tk.END, f"GUI: Could not load local chat log: {exc}\n")

    def _append_chat_ui(self, text: str) -> None:
        self.chat_console.insert(tk.END, text + "\n")
        self.chat_console.see(tk.END)

    def _set_status(self, text: str) -> None:
        self.root.after(0, lambda: self.status_var.set(text))

    def _refresh_review_context(self) -> None:
        text = format_gui_review_context(build_gui_review_context(REPO_ROOT))
        self.root.after(0, lambda: self.review_context_var.set(text))

    def _refresh_review_context_viewer(self) -> None:
        viewer = getattr(self, "review_context_viewer", None)
        if viewer is None:
            return
        viewer.delete("1.0", tk.END)
        viewer.insert("1.0", self.review_context_var.get())
        viewer.see("1.0")

    def _show_review_context(self) -> None:
        context = build_gui_review_context(REPO_ROOT)
        self.review_context_var.set(format_gui_review_context(context))
        dialog = tk.Toplevel(self.root)
        dialog.title("Workflow Review Context")
        dialog.transient(self.root)
        self._set_responsive_geometry(dialog, 760, 260)
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frame, text="Read-only workflow review context", font=("Segoe UI", 13, "bold")).pack(anchor=tk.W)
        context_viewer = tk.Text(
            frame,
            height=3,
            width=86,
            wrap=tk.WORD,
            relief=tk.FLAT,
            borderwidth=0,
            highlightthickness=0,
            background=dialog.cget("background"),
            font=("Segoe UI", 9),
        )
        context_viewer.pack(fill=tk.X, pady=(12, 16))
        context_viewer.insert("1.0", format_gui_review_context(context))
        self._bind_copy_support(context_viewer)
        context_viewer.bind("<KeyPress>", self._block_rendered_text_edit, add="+")
        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X)
        ttk.Button(buttons, text="Open Authority Scan", command=lambda: self._on_work_area_select_path(AUTHORITY_SCAN_PATH)).pack(side=tk.LEFT)
        ttk.Button(buttons, text="Close", command=dialog.destroy).pack(side=tk.RIGHT)

    def run_authority_scan(self) -> None:
        AUTHORITY_SCAN_PATH.parent.mkdir(parents=True, exist_ok=True)
        command = [PYTHON_EXE, str(SCRIPTS_DIR / "authority_consistency_scan.py"), "--output", str(AUTHORITY_SCAN_PATH)]
        self._run_command_async(command, "Authority consistency scan", on_success=self._refresh_review_context)

    def _refresh_ai_usage_panel(self) -> None:
        context_tokens = self.ai_input_tokens + self.ai_output_tokens
        context_percent = (context_tokens / self.ai_context_limit * 100) if self.ai_context_limit else 0.0
        cost_text = f"${self.ai_last_cost:.6f}" if self.ai_last_cost is not None else "Unavailable"
        self.ai_usage_var.set(
            f"Requests: {self.ai_request_count}    "
            f"Last: {self.ai_last_backend}/{self.ai_last_model}\n"
            f"Session tokens: {self.ai_input_tokens:,} in + {self.ai_output_tokens:,} out    "
            f"Context: {context_tokens:,}/{self.ai_context_limit:,} ({context_percent:.1f}%)    "
            f"Cost: {cost_text}"
        )

    def _update_ai_usage(self, backend: str, model: str, usage: Optional[dict], user_msg: str, assistant_text: str) -> None:
        usage = usage if isinstance(usage, dict) else {}
        input_tokens = usage.get("prompt_tokens", usage.get("input_tokens"))
        output_tokens = usage.get("completion_tokens", usage.get("output_tokens"))
        if not isinstance(input_tokens, int):
            input_tokens = max(1, len(user_msg) // 4)
        if not isinstance(output_tokens, int):
            output_tokens = max(1, len(assistant_text) // 4)

        self.ai_request_count += 1
        self.ai_input_tokens += input_tokens
        self.ai_output_tokens += output_tokens
        self.ai_last_backend = backend
        self.ai_last_model = model

        input_rate = os.getenv("AI_INPUT_USD_PER_MILLION_TOKENS")
        output_rate = os.getenv("AI_OUTPUT_USD_PER_MILLION_TOKENS")
        if input_rate and output_rate:
            try:
                self.ai_last_cost = (input_tokens * float(input_rate) + output_tokens * float(output_rate)) / 1_000_000
            except ValueError:
                self.ai_last_cost = None
        else:
            self.ai_last_cost = None
        self.root.after(0, self._refresh_ai_usage_panel)

    def _on_chat_enter(self, _event=None):
        self.send_chat_message()
        return "break"

    def send_chat_message(self) -> None:
        msg = self.chat_input_var.get().strip()
        if not msg:
            return
        self.chat_input_var.set("")
        self.append_chat(f"You: {msg}")

        if self._handle_local_command(msg):
            return

        mode = self.chat_mode_var.get().strip()
        if mode == "Command":
            self.append_chat("GUI: Command mode active. Use /help or switch mode to Auto/AI for model responses.")
            return

        self._run_ai_chat_async(msg)

    def _handle_local_command(self, msg: str) -> bool:
        # Lightweight command-style chat for quick operations.
        if msg.lower() in ("/help", "help"):
            self.append_chat("GUI: Commands: /run <stage>, run stage <start> to <end>, /status, /validate, /taxonomy, /analyze log, /stop, /clear")
            self.append_chat("GUI: Examples: /run 2, run stage 1 to 3, /run stage 2a to 5")
            self.append_chat("GUI: Log analysis: /analyze log or /analyze errors analyzes the current Execution Log read-only.")
            self.append_chat("GUI: Pasted errors: paste them in chat and ask 'analyze these errors; do not suggest code fixes'.")
            self.append_chat("GUI: AI modes: Auto (recommended), Command, AI.")
            return True
        if msg.lower() in (
            "/analyze log",
            "analyze log",
            "analyze run log",
            "/analyze errors",
            "analyze errors",
            "analyze possible errors",
        ):
            self._analyze_run_log()
            return True
        if msg.lower() in ("/status", "status"):
            self.run_status()
            return True
        if msg.lower() in ("/validate", "validate"):
            self.run_validate_all()
            return True
        if msg.lower() in ("/taxonomy", "taxonomy"):
            self.run_taxonomy_update()
            return True
        if msg.lower() in ("/stop", "stop"):
            self.stop_current()
            return True
        if msg.lower() in ("/clear", "clear"):
            self.chat_console.delete("1.0", tk.END)
            return True

        m = re.match(r"^/run\s+(0|1|2|2a|3|4|5)$", msg.lower())
        if m:
            stage = m.group(1)
            self.run_stage_script(stage)
            return True

        r = re.match(r"^/?run\s+stage\s+(0|1|2|2a|3|4|5)\s+to\s+(0|1|2|2a|3|4|5)$", msg.lower())
        if r:
            start_stage = r.group(1)
            end_stage = r.group(2)
            self.run_stage_range(start_stage, end_stage)
            return True
        return False

    def _analyze_run_log(self) -> None:
        log_text = self.log.get("1.0", tk.END).strip()
        if not log_text:
            self.append_chat("GUI: The execution log is empty; run a stage before requesting analysis.")
            return

        log_text = log_text[-16000:]
        analysis_request = (
            "Analyze the following Requirement AI workflow execution log in read-only mode. "
            "Identify error and warning messages, group related failures, explain likely causes, "
            "and list verification steps. Do not modify files, generate patches, or provide code fixes. "
            "If the evidence is insufficient, say what information is missing.\n\n"
            "EXECUTION LOG:\n"
            f"{log_text}"
        )
        self.append_chat("GUI: Sending the current run log for read-only error analysis.")
        self._run_ai_chat_async(analysis_request)

    def _update_ai_controls(self) -> None:
        backend = self.ai_backend_var.get().strip()
        local_enabled = backend in ("Auto", "Local") and len(self.local_models) > 0
        self.local_model_combo.configure(state="readonly" if local_enabled else "disabled")

    def refresh_local_models(self) -> None:
        models = []

        env_models = os.getenv("LOCAL_LLM_MODELS", "").strip()
        if env_models:
            for m in env_models.split(","):
                mm = m.strip()
                if mm:
                    models.append(mm)

        if LOCAL_MODELS_PATH.exists():
            try:
                data = json.loads(LOCAL_MODELS_PATH.read_text(encoding="utf-8"))
                if isinstance(data, list):
                    for item in data:
                        if isinstance(item, str) and item.strip():
                            models.append(item.strip())
                elif isinstance(data, dict):
                    # Accept {'models': ['model-a', ...]} shape.
                    for item in data.get("models", []):
                        if isinstance(item, str) and item.strip():
                            models.append(item.strip())
            except Exception as exc:
                self.append_chat(f"GUI: Could not parse local models config: {exc}")

        # De-duplicate while preserving order.
        seen = set()
        dedup = []
        for m in models:
            if m not in seen:
                seen.add(m)
                dedup.append(m)

        self.local_models = dedup
        if self.local_models:
            self.local_model_combo.configure(values=self.local_models)
            if self.local_model_var.get().strip() not in self.local_models:
                self.local_model_var.set(self.local_models[0])
            self.append_chat(f"GUI: Local models detected: {', '.join(self.local_models)}")
        else:
            self.local_model_combo.configure(values=[])
            self.local_model_var.set("")
            self.append_chat("GUI: No local models detected yet. Add LOCAL_LLM_MODELS or config/local_models.json.")
        self._update_ai_controls()

    def _run_ai_chat_async(self, user_msg: str) -> None:
        if self.ai_running:
            self.append_chat("GUI: AI request already in progress. Please wait.")
            return

        def worker():
            self.ai_running = True
            self._set_status("Running: AI chat")
            try:
                backend, model, assistant_text, usage = self._query_ai_backend(user_msg)
                self.chat_history.append({"role": "user", "content": user_msg})
                self.chat_history.append({"role": "assistant", "content": assistant_text})
                self._update_ai_usage(backend, model, usage, user_msg, assistant_text)
                self.append_chat(f"AI [{backend}:{model}]: {assistant_text}")
            except Exception as exc:
                self.append_chat(f"AI ERROR: {exc}")
            finally:
                self.ai_running = False
                self._set_status("Idle")

        threading.Thread(target=worker, daemon=True).start()

    def _query_ai_backend(self, user_msg: str):
        backend = self._resolve_backend()
        if backend == "Local":
            return self._chat_local(user_msg)
        if backend == "Azure":
            return self._chat_azure(user_msg)
        return self._chat_openai(user_msg)

    def _resolve_backend(self) -> str:
        requested = self.ai_backend_var.get().strip()
        if requested in ("Local", "Azure", "OpenAI"):
            return requested

        # Auto backend selection priority: OpenAI -> Azure -> Local.
        if self._openai_available():
            return "OpenAI"
        if self._azure_available():
            return "Azure"
        if self._local_available():
            return "Local"
        raise RuntimeError(
            "No AI backend configured. Set local endpoint env vars, or Azure/OpenAI env vars."
        )

    def _local_available(self) -> bool:
        # A configured local model implies the bundled Ollama-compatible default.
        return bool(os.getenv("LOCAL_LLM_URL")) or bool(self.local_models)

    def _azure_available(self) -> bool:
        return bool(
            os.getenv("AZURE_OPENAI_ENDPOINT")
            and os.getenv("AZURE_OPENAI_API_KEY")
            and os.getenv("AZURE_OPENAI_DEPLOYMENT")
        )

    def _openai_available(self) -> bool:
        return bool(os.getenv("OPENAI_API_KEY"))

    def _build_messages(self, user_msg: str):
        context = ""
        if PROJECT_CONTEXT_PATH.exists():
            try:
                project_context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8"))
                context = (
                    "Current GUI project context: "
                    f"source spec={project_context.get('source_spec_path', 'unknown')}; "
                    f"requirements index={project_context.get('ocr_index_path', 'unknown')}; "
                    f"repository root={REPO_ROOT}."
                )
            except (OSError, json.JSONDecodeError):
                context = f"Current GUI repository root={REPO_ROOT}."
        system_msg = {
            "role": "system",
            "content": (
                "You are the AI assistant inside the Requirement AI workflow GUI. "
                "Answer the user's question using the supplied conversation and project context. "
                "Be concise and actionable. You may explain logs, requirements, and workflow behavior, "
                "but do not modify files or claim to have run commands. "
                f"{context}"
            ),
        }
        recent = self.chat_history[-8:] if len(self.chat_history) > 8 else self.chat_history[:]
        return [system_msg] + recent + [{"role": "user", "content": user_msg}]

    def _http_post_json(self, url: str, payload: dict, headers: dict):
        req = request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json", **headers},
            method="POST",
        )
        try:
            with request.urlopen(req, timeout=90) as resp:
                text = resp.read().decode("utf-8")
                return json.loads(text)
        except error.HTTPError as exc:
            body = exc.read().decode("utf-8", errors="ignore")
            raise RuntimeError(f"HTTP {exc.code}: {body}")
        except error.URLError as exc:
            raise RuntimeError(
                f"AI endpoint is unreachable at {url}. Start Ollama or set LOCAL_LLM_URL. ({exc.reason})"
            )

    def _chat_openai(self, user_msg: str):
        api_key = self.openai_api_key_var.get().strip() or os.getenv("OPENAI_API_KEY", "").strip()
        base_url = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
        model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
        if not api_key:
            raise RuntimeError("OPENAI_API_KEY is not set.")

        payload = {
            "model": model,
            "messages": self._build_messages(user_msg),
            "temperature": 0.2,
        }
        data = self._http_post_json(
            f"{base_url}/chat/completions",
            payload,
            {"Authorization": f"Bearer {api_key}"},
        )
        text = data.get("choices", [{}])[0].get("message", {}).get("content", "")
        return "OpenAI", model, (text or "(empty response)").strip(), data.get("usage")

    def _chat_azure(self, user_msg: str):
        endpoint = os.getenv("AZURE_OPENAI_ENDPOINT", "").rstrip("/")
        key = os.getenv("AZURE_OPENAI_API_KEY", "").strip()
        deployment = os.getenv("AZURE_OPENAI_DEPLOYMENT", "").strip()
        api_version = os.getenv("AZURE_OPENAI_API_VERSION", "2024-10-21")
        if not endpoint or not key or not deployment:
            raise RuntimeError("Azure OpenAI vars missing. Need ENDPOINT, API_KEY, DEPLOYMENT.")

        payload = {
            "messages": self._build_messages(user_msg),
            "temperature": 0.2,
        }
        url = f"{endpoint}/openai/deployments/{deployment}/chat/completions?api-version={api_version}"
        data = self._http_post_json(url, payload, {"api-key": key})
        text = data.get("choices", [{}])[0].get("message", {}).get("content", "")
        return "Azure", deployment, (text or "(empty response)").strip(), data.get("usage")

    def _chat_local(self, user_msg: str):
        base_url = os.getenv("LOCAL_LLM_URL", "http://127.0.0.1:11434/v1").strip().rstrip("/")
        if not base_url:
            raise RuntimeError("Local AI endpoint is not configured. Set LOCAL_LLM_URL or start Ollama.")

        model = self.local_model_var.get().strip() or os.getenv("LOCAL_LLM_MODEL", "").strip()
        if not model and self.local_models:
            model = self.local_models[0]
        if not model:
            raise RuntimeError("No local model selected. Configure LOCAL_LLM_MODELS or config/local_models.json.")

        api_key = os.getenv("LOCAL_LLM_API_KEY", "").strip()
        headers = {}
        if api_key:
            headers["Authorization"] = f"Bearer {api_key}"

        payload = {
            "model": model,
            "messages": self._build_messages(user_msg),
            "temperature": 0.2,
        }
        data = self._http_post_json(f"{base_url}/chat/completions", payload, headers)
        text = data.get("choices", [{}])[0].get("message", {}).get("content", "")
        return "Local", model, (text or "(empty response)").strip(), data.get("usage")

    def _draw_workflow_diagram(self) -> None:
        self._draw_compact_workflow_graph()
        return

    def _draw_compact_workflow_graph(self) -> None:
        """Render the same complete graph as the dedicated workflow tab at compact scale."""
        c = self.diagram_canvas
        current_x = c.xview()[0] if c.xview() else 0.0
        c.delete("all")
        self.stage_rects = {}
        self.stage_labels = {}
        zoom = self.compact_workflow_zoom
        node_positions = {}
        stage_by_node = {
            "s0": "0", "s1": "1", "s2": "2", "s2a": "2a",
            "stage3": "3", "stage4": "4", "stage5": "5",
            "stage6": "6", "stage7": "7",
        }
        for node in WORKFLOW_DIAGRAM_NODES:
            x = (node["x"] + 10) * zoom
            y = compact_workflow_node_y(node, zoom)
            width = (112 if node["kind"] == "main" else 150) * zoom
            height = 64 * zoom
            node_positions[node["id"]] = (x, y, x + width, y + height)
        loopback_index = 0
        for source_id, target_id, edge_kind in WORKFLOW_DIAGRAM_EDGES:
            sx1, sy1, sx2, sy2 = node_positions[source_id]
            tx1, ty1, tx2, ty2 = node_positions[target_id]
            color = {"main": "#40566d", "supplementary": "#c26b2d", "loopback": "#8a5a12"}[edge_kind]
            if edge_kind == "loopback":
                lane_y = min(sy1, ty1) - (12 + (loopback_index * 10)) * zoom
                points = (
                    (sx1 + sx2) / 2, sy1,
                    (sx1 + sx2) / 2, lane_y,
                    (tx1 + tx2) / 2, lane_y,
                    (tx1 + tx2) / 2, ty1,
                )
                loopback_index += 1
            else:
                points = ((sx1 + sx2) / 2, (sy1 + sy2) / 2, (tx1 + tx2) / 2, (ty1 + ty2) / 2)
            line_options = {"fill": color, "width": 2, "arrow": tk.LAST}
            if edge_kind == "loopback":
                line_options["dash"] = (5, 3)
            c.create_line(*points, **line_options)
            if edge_kind == "loopback":
                c.create_text(
                    ((sx1 + sx2) / 2 + (tx1 + tx2) / 2) / 2,
                    lane_y - (6 * zoom),
                    text=f"{source_id.upper()} -> {target_id.upper()}",
                    font=("Segoe UI", max(6, int(7 * zoom))),
                    fill="#8a5a12",
                    anchor=tk.CENTER,
                )
        for node in WORKFLOW_DIAGRAM_NODES:
            x1, y1, x2, y2 = node_positions[node["id"]]
            style_key = "selectable" if node["id"] in SELECTABLE_WORKFLOW_NODE_IDS else node["kind"]
            style = WorkflowDiagramRenderer.COLORS[style_key]
            fill = style["fill"]
            outline = style["outline"]
            stage = stage_by_node.get(node["id"])
            tag = f"stage-{stage}" if stage else ("stage-supp" if node["id"] == "supp_import" else None)
            tags = ("workflow-node", node["id"], "stage-node", tag) if tag else ("workflow-node", node["id"])
            rect = c.create_rectangle(x1, y1, x2, y2, fill=fill, outline=outline, width=2, tags=tags)
            label = c.create_text(
                (x1 + x2) / 2,
                (y1 + y2) / 2,
                text=node["label"],
                width=max(1, x2 - x1 - 12 * zoom),
                font=("Segoe UI", 8, "bold"),
                justify=tk.CENTER,
                tags=tags,
            )
            if stage:
                self.stage_rects[stage] = rect
                self.stage_labels[stage] = label
                self._bind_stage_click(tag, stage)
            elif node["id"] == "supp_import":
                self.stage_rects["supp"] = rect
                self.stage_labels["supp"] = label
                self._bind_stage_click(tag, "supp")
            elif node["id"] in WORKFLOW_NODE_MEANINGS:
                node_tag = node["id"]
                if node_tag in SELECTABLE_WORKFLOW_NODE_IDS:
                    self.diagram_canvas.tag_bind(
                        node_tag,
                        "<Button-1>",
                        lambda _event, selected_node=node_tag: self._on_full_workflow_node_click(selected_node),
                    )
                self.diagram_canvas.tag_bind(
                    node_tag,
                    "<Enter>",
                    lambda event, node_id=node_tag: self._on_workflow_node_hover(event, node_id),
                )
                self.diagram_canvas.tag_bind(
                    node_tag,
                    "<Leave>",
                    lambda event, node_id=node_tag: self._on_workflow_node_leave(event, node_id),
                )
        for index, note in enumerate(WORKFLOW_DIAGRAM_NOTES):
            c.create_text(
                16 * zoom,
                (410 + index * 18) * zoom,
                anchor="w",
                text=note,
                width=1580 * zoom,
                font=("Segoe UI", max(7, int(8 * zoom))),
                fill="#2f3b46",
            )
        c.configure(scrollregion=compact_workflow_scrollregion(c.bbox("all"), margin=16 * zoom))
        c.xview_moveto(current_x)
        self._refresh_stage_styles()

    def _on_diagram_resize(self, _event) -> None:
        self._draw_workflow_diagram()

    def _bind_stage_click(self, tag: str, stage: str) -> None:
        self.diagram_canvas.tag_bind(tag, "<Button-1>", lambda _e, s=stage: self._on_stage_node_click(s))
        if stage == "5":
            self.diagram_canvas.tag_bind(tag, "<Button-3>", self._show_s5_context_menu)
        self.diagram_canvas.tag_bind(tag, "<Enter>", lambda event, s=stage: self._on_stage_hover_enter(event, s))
        self.diagram_canvas.tag_bind(tag, "<Leave>", lambda event, s=stage: self._on_stage_hover_leave(s, event.widget))

    def _show_s5_context_menu(self, event) -> str:
        menu = tk.Menu(self.root, tearoff=False)
        menu.add_command(label="S5 flow run", command=lambda: self.run_stage_cli("5"))
        menu.add_command(label="S5 after S2a run", command=self.run_stage5_after_stage2a)
        return self._show_popup_menu(menu, event)

    def _on_stage_hover_enter(self, event, stage: str) -> None:
        self.hover_stage = stage
        self.diagram_canvas.config(cursor="hand2")
        failure_reason = self.stage_failure_reasons.get(stage)
        if failure_reason:
            hover_text = f"Stage {stage.upper()} failed: {failure_reason}"
        else:
            hover_text = self.stage_meanings.get(stage, "")
        self.diagram_hint_var.set(hover_text)
        self._show_hover_help(event, hover_text, ("stage", stage))
        self._refresh_stage_styles()

    def _on_stage_hover_leave(self, stage: str, owner=None) -> None:
        if self.hover_stage == stage:
            self.hover_stage = None
        self.diagram_canvas.config(cursor="")
        self.diagram_hint_var.set("Hover a stage to see its meaning. Click to run.")
        self._refresh_stage_styles()

    def _on_workflow_node_hover(self, event, node_id: str) -> None:
        self.hover_stage = node_id
        self.diagram_canvas.config(cursor="hand2")
        hover_text = WORKFLOW_NODE_MEANINGS.get(node_id, "")
        self.diagram_hint_var.set(hover_text)
        self._show_hover_help(event, hover_text, ("workflow-node", node_id))

    def _on_workflow_node_leave(self, event, node_id: str) -> None:
        if self.hover_stage == node_id:
            self.hover_stage = None
        self.diagram_canvas.config(cursor="")
        self.diagram_hint_var.set("Hover a stage to see its meaning. Click to run.")

    def _on_stage_node_click(self, stage: str) -> None:
        if stage == "supp":
            self._hide_hover_help()
            self.append_chat("GUI: Diagram click -> Supplementary Source Review")
            self._show_source_spec_dialog()
            return
        if stage == "map":
            self.append_chat("GUI: Diagram click -> Architecture Map Review")
            self._show_architecture_map_review_dialog()
            return
        self.append_chat(f"GUI: Diagram click -> Stage {stage}")
        if stage == "arch-compare":
            self.run_stage6_compare()
            return
        self.run_stage_cli(stage)

    def set_active_stage(self, stage: Optional[str]) -> None:
        self.root.after(0, self._set_active_stage_ui, stage)

    def _set_active_stage_ui(self, stage: Optional[str]) -> None:
        self.active_stage = stage
        renderer = getattr(self, "workflow_diagram_renderer", None)
        if renderer is not None:
            renderer.set_active_stage(stage)
        self._refresh_stage_styles()

    def _refresh_stage_styles(self) -> None:
        for key, rect in self.stage_rects.items():
            if key in self.stage_failure_reasons:
                self.diagram_canvas.itemconfig(rect, fill="#f8b4b4", outline="#b3261e", width=3)
                self.diagram_canvas.itemconfig(self.stage_labels[key], fill="#7f1510")
            elif key == self.active_stage:
                active_style = WorkflowDiagramRenderer.COLORS["active"]
                self.diagram_canvas.itemconfig(rect, fill=active_style["fill"], outline=active_style["outline"], width=3)
                self.diagram_canvas.itemconfig(self.stage_labels[key], fill="#1f2d3d")
            elif key == self.hover_stage:
                self.diagram_canvas.itemconfig(rect, fill="#ffe599", outline="#c27c0e", width=3)
                self.diagram_canvas.itemconfig(self.stage_labels[key], fill="#1f2d3d")
            else:
                if key == "supp":
                    style = WorkflowDiagramRenderer.COLORS["supplementary"]
                else:
                    node_id = next(
                        (node_id for node_id, stage_key in WORKFLOW_NODE_STAGE_KEYS.items() if stage_key == key),
                        None,
                    )
                    style_key = "selectable" if node_id in SELECTABLE_WORKFLOW_NODE_IDS else "main"
                    style = WorkflowDiagramRenderer.COLORS[style_key]
                self.diagram_canvas.itemconfig(rect, fill=style["fill"], outline=style["outline"], width=2)
                self.diagram_canvas.itemconfig(self.stage_labels[key], fill="#1f2d3d")

    def clear_stage_failure(self, stage: Optional[str]) -> None:
        if stage:
            self.root.after(0, self._clear_stage_failure_ui, stage.lower())

    def _clear_stage_failure_ui(self, stage: str) -> None:
        self.stage_failure_reasons.pop(stage, None)
        self._refresh_stage_styles()

    def set_stage_failure(self, stage: Optional[str], reason: str) -> None:
        if stage:
            self.root.after(0, self._set_stage_failure_ui, stage.lower(), reason)

    def _set_stage_failure_ui(self, stage: str, reason: str) -> None:
        self.stage_failure_reasons[stage] = reason
        self.diagram_hint_var.set(f"Stage {stage.upper()} failed: {reason}")
        self._refresh_stage_styles()

    @staticmethod
    def _failure_reason(output_lines) -> str:
        meaningful_lines = []
        for line in output_lines:
            text = line.strip()
            lower = text.lower()
            if not text or "stop at stage" in lower or "workflow run:" in lower:
                continue
            if re.search(r"\b(error|exception|traceback|fail(?:ed|ure)?|missing|duplicate|unsupported|invalid)\b", lower):
                meaningful_lines.append(text)
        if meaningful_lines:
            return meaningful_lines[-1][:240]
        return "Command exited with a non-zero status; inspect the execution log."

    def _show_downstream_block_dialog(self, output_lines, cmd) -> None:
        """Present downstream validation failures without exposing internal diagnostics."""
        snapshot_id = "unavailable"
        command_text = " ".join(str(item) for item in cmd)
        snapshot_match = re.search(r"--snapshot-id\s+([^\s]+)", command_text)
        if snapshot_match:
            snapshot_id = snapshot_match.group(1)
        result = None
        try:
            from validate_downstream_coherence import validate

            result = validate(REPO_ROOT, snapshot_id if snapshot_id != "unavailable" else None)
            snapshot_id = str(result.get("selected_snapshot_id") or snapshot_id)
            DOWNSTREAM_VALIDATION_REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
            DOWNSTREAM_VALIDATION_REPORT_PATH.write_text(
                json.dumps(result, indent=2, sort_keys=True) + "\n",
                encoding="utf-8",
            )
        except Exception as exc:
            result = {"findings": [f"Validation report unavailable: {exc}"]}

        categories = Counter()
        for finding in result.get("findings", []):
            category = str(finding).split(":", 1)[0].strip()
            if category:
                categories[category] += 1
        if categories:
            problems = "\n".join(f"- {name} ({count})" for name, count in sorted(categories.items()))
        else:
            problems = "- Validation findings require review"
        report_path = DOWNSTREAM_VALIDATION_REPORT_PATH

        dialog = tk.Toplevel(self.root)
        dialog.title("Downstream execution blocked")
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.columnconfigure(0, weight=1)
        body = ttk.Frame(dialog, padding=16)
        body.grid(row=0, column=0, sticky="nsew")
        body.columnconfigure(0, weight=1)
        ttk.Label(body, text="The selected approved snapshot cannot continue to Stage 3–7 because some generated downstream outputs are not valid yet.", wraplength=620, justify="left").grid(row=0, column=0, sticky="w", pady=(0, 12))
        ttk.Label(body, text="Snapshot:", font=("Segoe UI", 10, "bold")).grid(row=1, column=0, sticky="w")
        ttk.Label(body, text=snapshot_id, wraplength=620).grid(row=2, column=0, sticky="w", pady=(0, 10))
        ttk.Label(body, text="Problems found:", font=("Segoe UI", 10, "bold")).grid(row=3, column=0, sticky="w")
        ttk.Label(body, text=problems, justify="left", anchor="w").grid(row=4, column=0, sticky="w", pady=(0, 10))
        ttk.Label(body, text="What you can do now:", font=("Segoe UI", 10, "bold")).grid(row=5, column=0, sticky="w")
        ttk.Label(body, text="- Do not run Stage 3–7 yet\n- Do not run SysML yet\n- Review the validation report\n- If available, use Repair downstream outputs and then try again\n- Otherwise, send the report to the workflow maintainer", justify="left", anchor="w").grid(row=6, column=0, sticky="w", pady=(0, 10))
        ttk.Label(body, text="Report:", font=("Segoe UI", 10, "bold")).grid(row=7, column=0, sticky="w")
        ttk.Label(body, text=report_path.relative_to(REPO_ROOT).as_posix(), wraplength=620).grid(row=8, column=0, sticky="w", pady=(0, 14))

        buttons = ttk.Frame(body)
        buttons.grid(row=9, column=0, sticky="e")

        def open_report() -> None:
            try:
                os.startfile(str(report_path))
            except OSError as exc:
                messagebox.showerror("Open validation report", str(exc), parent=dialog)

        def copy_report_path() -> None:
            self.root.clipboard_clear()
            self.root.clipboard_append(str(report_path))
            self.root.update()

        ttk.Button(buttons, text="Open validation report", command=open_report).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(buttons, text="Copy report path", command=copy_report_path).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(buttons, text="Close", command=dialog.destroy).pack(side=tk.LEFT)
        dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
        self._set_responsive_geometry(dialog, 820, 640, 680, 420)

    def _append_retrieval_evidence(self, output_lines) -> None:
        """Show workflow retrieval preflight and benchmark evidence in the log."""
        benchmark_path = REPO_ROOT / "artifacts/rag/retrieval_benchmark_report.json"
        try:
            if PROJECT_CONTEXT_PATH.exists():
                project_context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8"))
                configured_report = project_context.get("retrieval_benchmark_report_path")
                if configured_report:
                    benchmark_path = Path(str(configured_report))
                    if not benchmark_path.is_absolute():
                        benchmark_path = REPO_ROOT / benchmark_path
        except (OSError, ValueError, TypeError):
            pass
        report_path = REPO_ROOT / "docs/semantic-retrieval-benchmark-report.md"
        test_lines = []
        for line in output_lines:
            text = line.rstrip("\n")
            if re.search(r"Ran \d+ tests? in |^OK$|^FAILED \(", text.strip()):
                test_lines.append(text.strip())
        if not test_lines and not benchmark_path.exists():
            return

        self.append_log("--- Retrieval Test Evidence ---")
        if test_lines:
            for line in test_lines:
                self.append_log(line)
        else:
            self.append_log("Test result: see workflow_cli_runs.jsonl")

        if benchmark_path.exists():
            try:
                payload = json.loads(benchmark_path.read_text(encoding="utf-8"))
                selection = payload.get("selection", {})
                self.append_log(f"Selected mode: {selection.get('selected_mode', 'hybrid')}")
                self.append_log("Benchmark Results")
                self.append_log("Mode | Recall@1 | Recall@5 | Recall@10 | MRR | Mean ms | p95 ms | Status")
                for row in payload.get("results", []):
                    if "recall_at_1" not in row:
                        self.append_log(f"{row.get('mode', '')} | unavailable | {row.get('error', '')}")
                        continue
                    self.append_log(
                        f"{row.get('mode')} | {row.get('recall_at_1', 0):.3f} | "
                        f"{row.get('recall_at_5', 0):.3f} | {row.get('recall_at_10', 0):.3f} | "
                        f"{row.get('mrr', 0):.3f} | {row.get('mean_latency_ms', 0):.2f} | "
                        f"{row.get('p95_latency_ms', 0):.2f} | {row.get('status', '')}"
                    )
            except (OSError, ValueError, TypeError) as exc:
                self.append_log(f"Benchmark Results unavailable: {exc}")

        if report_path.exists():
            try:
                report_lines = report_path.read_text(encoding="utf-8").splitlines()
                for heading in ("## Benchmark Results", "### Focused Retrieval Test Evidence"):
                    try:
                        start = report_lines.index(heading)
                    except ValueError:
                        continue
                    self.append_log(heading)
                    for line in report_lines[start + 1 :]:
                        if line.startswith("## ") or line.startswith("### "):
                            break
                        if line.startswith("|") or line.startswith("Command:") or line.startswith("Recorded result:") or line.startswith("```") or re.match(r"^(Ran \d+ tests?|OK$)", line):
                            self.append_log(line)
            except OSError as exc:
                self.append_log(f"Retrieval Test Evidence unavailable: {exc}")
        self.append_log("--- End Retrieval Test Evidence ---")

    def _append_workflow_manifest_summary(self) -> None:
        """Show the resolved pipeline and called workflow files in the GUI log."""
        manifest_path = REPO_ROOT / "artifacts/orchestrator/workflow_cli_runs.jsonl"
        self.append_log("--- Workflow Execution Summary ---")
        self.append_log("Pipeline used: workflow_cli.py -> Stage 0 retrieval tests/benchmark when Stage 0 is selected -> stage runner -> gate validator")
        if not manifest_path.exists():
            self.append_log("Called run files: manifest unavailable")
            self.append_log("--- End Workflow Execution Summary ---")
            return
        try:
            records = []
            for line in manifest_path.read_text(encoding="utf-8", errors="ignore").splitlines():
                if line.strip():
                    try:
                        records.append(json.loads(line))
                    except json.JSONDecodeError:
                        continue
            for record in records[-8:]:
                kind = record.get("kind", "")
                script = record.get("script", "")
                status = "PASS" if record.get("exit_code") == 0 else f"FAIL ({record.get('exit_code')})"
                if kind == "retrieval_tests":
                    self.append_log(f"Preliminary retrieval tests: {status} ({len(record.get('modules', []))} modules)")
                elif kind == "retrieval_benchmark":
                    self.append_log(f"Project benchmark: {status}; report={record.get('report', '')}; selection={record.get('selection', '')}")
                elif script:
                    self.append_log(f"Called {kind}: {script} -> {status}")
        except OSError as exc:
            self.append_log(f"Called run files unavailable: {exc}")

        try:
            context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8"))
            rag_db = Path(str(context.get("rag_db_path", "artifacts/rag/rag_index.sqlite")))
            rag_manifest = Path(str(context.get("rag_manifest_path", "artifacts/rag/index_manifest.json")))
            if not rag_db.is_absolute():
                rag_db = REPO_ROOT / rag_db
            if not rag_manifest.is_absolute():
                rag_manifest = REPO_ROOT / rag_manifest
            self.append_log(f"RAG creation artifact: {rag_db} ({'present' if rag_db.exists() else 'missing'})")
            self.append_log(f"RAG manifest: {rag_manifest} ({'present' if rag_manifest.exists() else 'missing'})")
        except (OSError, ValueError, TypeError) as exc:
            self.append_log(f"RAG artifact summary unavailable: {exc}")
        self.append_log("--- End Workflow Execution Summary ---")

    def clear_log(self) -> None:
        self.log.delete("1.0", tk.END)

    def _zoom_execution_log(self, factor: float) -> str:
        if not hasattr(self, "log_font"):
            return "break"
        self.execution_log_zoom = min(2.5, max(0.6, self.execution_log_zoom * factor))
        self.log_font.configure(size=max(7, round(11 * self.execution_log_zoom)))
        self.execution_log_zoom_var.set(f"Log zoom: {round(self.execution_log_zoom * 100)}%")
        return "break"

    def _zoom_execution_log_wheel(self, event) -> str:
        return self._zoom_execution_log(1.15 if event.delta > 0 else 1 / 1.15)

    def _reset_execution_log_zoom(self) -> str:
        if not hasattr(self, "log_font"):
            return "break"
        self.execution_log_zoom = 1.0
        self.log_font.configure(size=11)
        self.execution_log_zoom_var.set("Log zoom: 100%")
        return "break"

    def stop_current(self) -> None:
        if self.current_process is None:
            messagebox.showinfo("Info", "No process is currently running.")
            return
        try:
            self.current_process.terminate()
            self.append_log("[INFO] Sent terminate signal to current process.")
        except Exception as exc:
            self.append_log(f"[ERROR] Failed to terminate process: {exc}")

    def _run_command_async(
        self,
        cmd,
        label: str,
        stage: Optional[str] = None,
        approval_after_success: bool = False,
        on_success=None,
    ) -> None:
        if self.is_running:
            messagebox.showwarning("Busy", "Another command is already running.")
            return

        def worker():
            self.is_running = True
            self._set_status(f"Running: {label}")
            self.clear_stage_failure(stage)
            self.set_active_stage(stage)
            is_stage6_run = stage == "6" or label == "Stage 6 Digital IPOS"
            started_at = datetime.now().astimezone() if is_stage6_run else None
            start_suffix = f" | start={started_at.isoformat(timespec='seconds')}" if started_at else ""
            self.append_log(f"\n=== START {label}{start_suffix} ===")
            self.append_log("Command: " + " ".join(str(c) for c in cmd))
            self.append_chat(f"GUI: Running {label}")
            output_lines = []
            last_stage = stage
            success_callback = None
            code = None

            try:
                self.current_process = subprocess.Popen(
                    cmd,
                    cwd=str(REPO_ROOT),
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    bufsize=1,
                )

                for line in self.current_process.stdout:
                    output_lines.append(line)
                    stage_match = re.search(r"\bStage\s+(0|1|2a|3|4|5|6)\b", line, re.IGNORECASE)
                    if stage_match:
                        last_stage = stage_match.group(1).lower()
                        self.clear_stage_failure(last_stage)
                        self.set_active_stage(last_stage)
                    self.append_log(line.rstrip("\n"))

                code = self.current_process.wait()
                if not is_stage6_run:
                    self.append_log(f"=== END {label} (exit={code}) ===\n")
                if code == 0 and stage == "0" and "workflow_cli.py" in " ".join(str(item) for item in cmd):
                    self._append_retrieval_evidence(output_lines)
                    self._append_workflow_manifest_summary()
                self.append_chat(f"GUI: Completed {label} with exit={code}")
                if code == 0 and approval_after_success:
                    self.root.after(0, self._show_architecture_profile_approval_dialog)
                if code == 0 and on_success is not None:
                    success_callback = on_success
                if code == 0 and stage in {"3", "4", "5", "6", "7"}:
                    prior_callback = success_callback
                    success_callback = lambda prior_callback=prior_callback: self._refresh_traceability_after_generation(prior_callback)
                if code != 0:
                    stopped_stage = re.search(r"\bSTOP at stage\s+(0|1|2a|3|4|5|6)\b", "".join(output_lines), re.IGNORECASE)
                    failed_stage = stopped_stage.group(1).lower() if stopped_stage else last_stage
                    failure_reason = self._failure_reason(output_lines)
                    if failed_stage:
                        self.set_stage_failure(failed_stage, failure_reason)
                        self.append_chat(f"GUI: Stage {failed_stage.upper()} failed: {failure_reason}")
                    else:
                        self.append_chat(f"GUI: {label} failed: {failure_reason}")
                    if "downstream snapshot coherence blocked" in "".join(output_lines).lower() or "block_downstream" in "".join(output_lines).lower():
                        self.root.after(0, self._show_downstream_block_dialog, output_lines, cmd)
                    if "manual table-row review required" in "".join(output_lines).lower():
                        self.root.after(0, self._show_table_row_manual_review_dialog)
                self._set_status("Idle")
            except Exception as exc:
                self.append_log(f"[ERROR] {exc}")
                self.append_chat(f"GUI: ERROR while running {label}: {exc}")
                self.set_stage_failure(last_stage, str(exc))
                self._set_status("Idle")
            finally:
                if is_stage6_run:
                    finished_at = datetime.now().astimezone()
                    elapsed = finished_at - started_at
                    elapsed_text = str(elapsed).split(".", 1)[0]
                    exit_text = str(code) if code is not None else "exception"
                    self.append_log(
                        f"=== END {label} | finish={finished_at.isoformat(timespec='seconds')} "
                        f"| elapsed={elapsed_text} | exit={exit_text} ==="
                    )
                self.set_active_stage(None)
                self.current_process = None
                self.is_running = False
                if success_callback is not None:
                    self.root.after(0, success_callback)

        threading.Thread(target=worker, daemon=True).start()

    def _refresh_traceability_after_generation(self, prior_callback=None) -> None:
        """Refresh derived hierarchy reports before updating the GUI after generation."""
        report_command = [
            PYTHON_EXE,
            str(SCRIPTS_DIR / "generate_snapshot_reports.py"),
            "--use-latest-approved",
            "--report",
            "all",
        ]
        self._run_command_async(
            report_command,
            "Refresh snapshot traceability reports",
            on_success=lambda: self._run_command_async(
                [PYTHON_EXE, str(SCRIPTS_DIR / "traceability_hierarchy_coverage.py")],
                "Refresh hierarchy coverage report",
                on_success=prior_callback,
            ),
        )

    def run_stage_script(self, stage: str) -> None:
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "run", "--stage", stage]
        self._run_command_async(
            cmd,
            f"CLI run stage {stage}",
            stage=stage,
            approval_after_success=stage == "2",
            on_success=self._show_architecture_map_review_dialog if stage == "2a" else None,
        )

    def _refresh_architecture_after_supplementary_merge(self) -> None:
        """Incrementally recompute dependent artifacts after integrated-baseline update."""
        self.run_stage_script_after(
            "0",
            lambda: self.run_stage_script_after(
                "1",
                lambda: self.run_stage_script_after(
                    "2",
                    self._show_architecture_profile_approval_dialog,
                ),
            ),
        )

    def run_stage_script_after(self, stage: str, on_success) -> None:
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "run", "--stage", stage]
        self._run_command_async(
            cmd,
            f"CLI run stage {stage}",
            stage=stage,
            approval_after_success=False,
            on_success=on_success,
        )

    def _show_source_spec_dialog(self) -> None:
        self._hide_hover_help()
        dialog = tk.Toplevel(self.root)
        dialog.title("Source Specification")
        dialog.transient(self.root)
        dialog.bind("<Destroy>", lambda _event: self._hide_hover_help(), add="+")
        dialog.resizable(False, False)
        dialog.grab_set()
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        mode = tk.StringVar(value="primary")
        selected_path = tk.StringVar()
        revision = tk.StringVar(value="")
        revision_date = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))

        ttk.Label(frame, text="Source mode:").grid(row=0, column=0, sticky=tk.W, pady=(0, 8))
        ttk.Combobox(
            frame,
            textvariable=mode,
            values=("primary", "supplementary"),
            state="readonly",
            width=18,
        ).grid(row=0, column=1, sticky=tk.W, pady=(0, 8))
        ttk.Label(frame, text="PDF file:").grid(row=1, column=0, sticky=tk.W, pady=4)
        ttk.Entry(frame, textvariable=selected_path, width=56).grid(row=1, column=1, sticky=tk.W, pady=4)

        def choose_pdf() -> None:
            self.append_log("[GUI] Source Specification: Browse button clicked; opening PDF chooser.")
            path = filedialog.askopenfilename(
                parent=dialog,
                title="Choose source specification",
                filetypes=(("Specifications", "*.pdf *.html *.htm *.docx"), ("PDF files", "*.pdf"), ("HTML files", "*.html *.htm"), ("DOCX files", "*.docx")),
            )
            if path:
                selected_path.set(path)
                self.append_log(f"[GUI] Source Specification: selected PDF {Path(path).name}.")
            else:
                self.append_log("[GUI] Source Specification: PDF chooser cancelled.")

        ttk.Button(frame, text="Browse...", command=choose_pdf, width=11).grid(row=1, column=2, padx=(8, 0), pady=4)
        ttk.Label(frame, text="Source revision:").grid(row=2, column=0, sticky=tk.W, pady=4)
        ttk.Entry(frame, textvariable=revision, width=24).grid(row=2, column=1, sticky=tk.W, pady=4)
        ttk.Label(frame, text="Revision date:").grid(row=3, column=0, sticky=tk.W, pady=4)
        ttk.Entry(frame, textvariable=revision_date, width=24).grid(row=3, column=1, sticky=tk.W, pady=4)
        ttk.Label(
            frame,
            text="Primary selection updates this project's configured source. Supplementary selection stages a review package and does not modify the tracked corpus.",
            wraplength=620,
            justify=tk.LEFT,
        ).grid(row=4, column=0, columnspan=3, sticky=tk.W, pady=(10, 14))

        def submit() -> None:
            self.append_log(f"[GUI] Source Specification: Continue clicked; mode={mode.get()}.")
            source = Path(selected_path.get()).resolve()
            if not source.exists() or source.suffix.lower() not in {".pdf", ".html", ".htm", ".docx"}:
                self.append_log(f"[GUI] Source Specification: rejected source path {source}.")
                messagebox.showerror("Source Specification", "Choose an existing PDF, HTML, or DOCX source file.", parent=dialog)
                return
            if mode.get() == "primary":
                context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8")) if PROJECT_CONTEXT_PATH.exists() else {}
                try:
                    assessment = assess_source(source, context)
                except Exception as exc:
                    messagebox.showerror("Primary Source Bootstrap", str(exc), parent=dialog)
                    return
                evidence_preview = "\n".join(str(item.get("text", "")) for item in assessment.evidence[:5]) or "No ID evidence detected."
                proposal = (
                    f"Detected status: {assessment.status}\n"
                    f"Proposed prefix: {assessment.proposed_prefix or '(none)'}\n"
                    f"Reason: {assessment.reason}\n\nEvidence:\n{evidence_preview}"
                )
                if not messagebox.askyesno("Primary Source Bootstrap Review", proposal + "\n\nApprove this proposal?", parent=dialog):
                    self.append_log("[GUI] Primary Source Bootstrap: proposal rejected.")
                    return
                prefix = assessment.proposed_prefix
                if messagebox.askyesno("Edit Req-ID Prefix", "Edit the proposed req-ID prefix before approval?", parent=dialog):
                    edited = simpledialog.askstring("Req-ID Prefix", "Approved base req-ID prefix:", initialvalue=prefix, parent=dialog)
                    if edited is None:
                        return
                    prefix = edited.strip()
                    assessment = assessment.__class__(assessment.source_path, assessment.status, prefix, assessment.tagged_count, assessment.table_id_count, assessment.unbracketed_id_count, assessment.evidence, assessment.reason)
                project_id = str(context.get("project_name") or REPO_ROOT.name)
                profile_id = persist_candidate(REPO_ROOT, project_id=project_id, assessment=assessment, actor=runtime_user_name())
                approve_candidate(REPO_ROOT, profile_id=profile_id, actor=runtime_user_name(), approved_prefix=prefix)
                try:
                    context["source_spec_path"] = source.relative_to(REPO_ROOT).as_posix()
                except ValueError:
                    context["source_spec_path"] = str(source)
                context["source_bootstrap_profile_id"] = profile_id
                context["source_tag_status"] = assessment.status
                context["source_req_id_prefix"] = prefix
                context["source_bootstrap_approved_by"] = "gui"
                context["source_bootstrap_approved_at"] = datetime.now().isoformat(timespec="seconds")
                PROJECT_CONTEXT_PATH.write_text(json.dumps(context, indent=2) + "\n", encoding="utf-8")
                self.append_log(f"[GUI] Source Specification: configured primary source {source.name}; awaiting Stage 1.")
                self.append_chat(f"GUI: Primary source selected: {source.name}. Run Stage 1 to extract it.")
                messagebox.showinfo("Source Specification", "Primary source configured. Run Stage 1 to extract and validate it.", parent=dialog)
                dialog.destroy()
                return
            if not revision.get().strip():
                self.append_log("[GUI] Source Specification: supplementary source rejected because revision is empty.")
                messagebox.showerror("Source Specification", "Enter a source revision identifier.", parent=dialog)
                return
            self.append_log(
                f"[GUI] Source Specification: staging supplementary source {source.name} "
                f"(revision={revision.get().strip()}); launching ingestion command."
            )
            dialog.destroy()
            command = [
                PYTHON_EXE,
                str(SCRIPTS_DIR / "ingest_source_spec.py"),
                "stage",
                "--source", str(source),
                "--revision", revision.get().strip(),
                "--revision-date", revision_date.get().strip(),
            ]
            self._run_command_async(
                command,
                f"Stage supplementary source {source.name}",
                on_success=lambda: self._show_source_ingestion_review_dialog(source.name),
            )

        def reopen_last_saved_review() -> None:
            self.append_log("[GUI] Source Specification: Reopen Last Saved Review clicked.")
            ingestion_root = REPO_ROOT / "artifacts" / "source_ingestion"
            candidates = []
            for metadata_path in ingestion_root.glob("*/source_metadata.json"):
                try:
                    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    continue
                workbook_path = metadata_path.parent / "comparison_results.xlsx"
                if workbook_path.exists():
                    candidates.append((workbook_path.stat().st_mtime, metadata.get("source_spec", workbook_path.name)))
            if not candidates:
                messagebox.showinfo("Source Specification", "No saved supplementary review workbook was found.", parent=dialog)
                return
            source_name = max(candidates, key=lambda item: item[0])[1]
            dialog.destroy()
            self._show_source_ingestion_review_dialog(source_name)

        buttons = ttk.Frame(frame)
        buttons.grid(row=5, column=0, columnspan=3, sticky=tk.E)
        ttk.Button(buttons, text="Reopen Last Saved Review", command=reopen_last_saved_review, width=24).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(buttons, text="Cancel", command=dialog.destroy, width=11).pack(side=tk.RIGHT)
        ttk.Button(buttons, text="Continue", command=submit, width=11).pack(side=tk.RIGHT, padx=(0, 8))

    def _show_source_ingestion_review_dialog(self, source_name: str) -> None:
        self._hide_hover_help()
        ingestion_root = REPO_ROOT / "artifacts" / "source_ingestion"
        candidates = []
        for metadata_path in ingestion_root.glob("*/source_metadata.json"):
            try:
                metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            except Exception:
                continue
            if metadata.get("source_spec") == source_name and (metadata_path.parent / "comparison_results.xlsx").exists():
                candidates.append(metadata_path.parent)
        if not candidates:
            self.append_chat(f"GUI: No review workbook found for staged source {source_name}.")
            return
        ingestion_dir = max(candidates, key=lambda path: path.stat().st_mtime)
        review_csv = ingestion_dir / "comparison_results.csv"
        review_workbook = review_csv.with_suffix(".xlsx")
        rows = []
        try:
            workbook_state = self._close_saved_open_excel_workbook(review_workbook)
            if workbook_state == "same_name_other_path":
                raise RuntimeError(
                    f"Excel has another workbook named {review_workbook.name} open from a different folder. "
                    "Close that workbook before reopening this review."
                )
            if workbook_state == "save_failed":
                raise RuntimeError(
                    f"The review workbook has unsaved changes: {review_workbook.name}. "
                    "Save it in Excel, close it, and reopen the review."
                )
            _sync_saved_review_workbook(ingestion_dir)
            with review_csv.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
        except (OSError, RuntimeError, ValueError) as exc:
            messagebox.showerror("Source Ingestion Review", str(exc), parent=self.root)
            return
        counts = Counter(row.get("req_class") or row.get("classification", "unknown") for row in rows)
        self._open_mapping_workbook(review_workbook, self.root)
        dialog = tk.Toplevel(self.root)
        dialog.title("Supplementary Source Approval")
        dialog.transient(self.root)
        dialog.bind("<Destroy>", lambda _event: self._hide_hover_help(), add="+")
        self._set_responsive_geometry(dialog, 640, 440)
        dialog.resizable(True, False)
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frame, text=f"Review {source_name}", font=("Segoe UI", 13, "bold")).pack(anchor=tk.W)
        ttk.Label(
            frame,
            text=(
                f"New: {counts['new']}   Duplication: {counts['duplication']}   Refines: {counts['refines']}\n"
                f"Conflict: {counts['conflict']}   Removed: {counts['removed']}\n\n"
                "Edit only the review fields exposed by the workbook menus. The tool sets review_decision automatically to reassigned when a mapping value changes. "
                "Accept All updates the currently opened review workbook immediately and keeps this menu open for crosscheck. "
                "Save Excel before importing or merging."
            ),
            wraplength=580,
            justify=tk.LEFT,
        ).pack(anchor=tk.W, pady=(8, 14))

        def import_workbook() -> List[Dict[str, str]]:
            if load_workbook is None:
                raise RuntimeError("openpyxl is required to import the source review workbook.")
            if PatternFill is None:
                raise RuntimeError("openpyxl styles are required to preserve review-decision colors.")
            workbook_state = self._close_saved_open_excel_workbook(review_workbook)
            if workbook_state == "same_name_other_path":
                raise RuntimeError(
                    f"Excel has another workbook named {review_workbook.name} open from a different folder. "
                    "Close that workbook before importing this review."
                )
            if workbook_state == "save_failed":
                raise RuntimeError(
                    f"The review workbook has unsaved changes: {review_workbook.name}. "
                    "Save it in Excel, close it, and click the import or merge button again."
                )
            self.append_log(f"[GUI] Reading saved review workbook before CSV synchronization: {review_workbook}")
            with review_csv.open("r", encoding="utf-8-sig", newline="") as handle:
                fieldnames = csv.DictReader(handle).fieldnames or []
            workbook = load_workbook(review_workbook, data_only=False)
            try:
                sheet = workbook["Source Ingestion Review"]
                values = list(sheet.iter_rows(values_only=True))
                headers = [str(value or "").strip() for value in values[0]] if values else []
                if headers != fieldnames:
                    raise ValueError("Review workbook columns do not match the staged CSV.")
                imported = [dict(zip(fieldnames, ["" if value is None else str(value).strip() for value in row])) for row in values[1:] if any(value is not None and str(value).strip() for value in row)]
                decisions = {row.get("review_decision", "").strip().lower() for row in imported}
                if not imported or not decisions.issubset({"approved", "rejected", "pending", "pending_review", "needs_clarification"}):
                    raise ValueError(
                        "Every review row must have review_decision set to approved, rejected, "
                        "pending, pending_review, or needs_clarification."
                    )
                classifications = {row.get("classification", "").strip().lower() for row in imported}
                if not classifications.issubset({"analog", "digital", "system", "block"}):
                    raise ValueError("Every review row must have classification set to Analog, Digital, System, or Block.")
                decision_column = headers.index("review_decision") + 1
                approved_fill = PatternFill("solid", fgColor="E2F0D9")
                pending_fill = PatternFill("solid", fgColor="FCE4D6")
                for row_index, imported_row in enumerate(imported, start=2):
                    sheet.cell(row=row_index, column=decision_column).fill = (
                        approved_fill if imported_row["review_decision"].casefold() == "approved" else pending_fill
                    )
                with review_csv.open("w", encoding="utf-8", newline="") as handle:
                    writer = csv.DictWriter(handle, fieldnames=fieldnames)
                    writer.writeheader()
                    writer.writerows(imported)
                workbook.save(review_workbook)
                return imported
            finally:
                workbook.close()

        def refresh_review_colors() -> None:
            self.append_log("[GUI] Supplementary Source Approval: Refresh Review Colors clicked.")
            try:
                workbook_state = self._close_saved_open_excel_workbook(review_workbook)
                if workbook_state == "same_name_other_path":
                    raise RuntimeError(
                        f"Excel has another workbook named {review_workbook.name} open from a different folder."
                    )
                if workbook_state == "save_failed":
                    raise RuntimeError(
                        f"Save {review_workbook.name} in Excel, close it, then refresh the colors again."
                    )
                self.append_log("[GUI] Supplementary Source Approval: saved workbook checked; no workbook write performed.")
                self.append_chat("GUI: Checked supplementary review colors from the user-saved Excel workbook.")
                messagebox.showinfo(
                    "Refresh Review Colors",
                    "The saved workbook was checked. No automatic workbook save was performed.",
                    parent=dialog,
                )
            except PermissionError:
                self.append_log("[GUI] Supplementary Source Approval: refresh blocked because the workbook is open in Excel.")
                messagebox.showerror(
                    "Refresh Review Colors",
                    "Close the review workbook in Excel, then refresh the colors.",
                    parent=dialog,
                )
            except Exception as exc:
                self.append_log(f"[GUI] Supplementary Source Approval: refresh failed: {exc}")
                messagebox.showerror("Refresh Review Colors", str(exc), parent=dialog)

        def merge_and_run_stage0() -> None:
            self.append_log("[GUI] Supplementary Source Approval: Update Integrated Source Baseline and Recompute Dependents clicked.")
            try:
                imported = import_workbook()
                self.append_log(f"[GUI] Supplementary Source Approval: imported {len(imported)} workbook rows.")
            except Exception as exc:
                self.append_log(f"[GUI] Supplementary Source Approval: workbook import failed: {exc}")
                messagebox.showerror("Supplementary Source Approval", str(exc), parent=dialog)
                return
            approved_count = sum(row.get("review_decision", "").strip().lower() == "approved" for row in imported)
            self.append_log(
                f"[GUI] Supplementary Source Approval: {approved_count} approved rows found; "
                "waiting for merge confirmation."
            )
            if not messagebox.askyesno(
                "Update Integrated Source Baseline",
                f"Add {approved_count} approved requirements to the integrated source baseline and recompute dependent artifacts?",
                parent=dialog,
            ):
                self.append_log("[GUI] Supplementary Source Approval: merge cancelled by user.")
                return
            approved_staged_ids = [row.get("staged_id", "").strip() for row in imported if row.get("review_decision", "").strip().lower() == "approved" and row.get("staged_id", "").strip()]
            if approved_staged_ids and not messagebox.askyesno(
                "Approve Same-ID Conflict Winners",
                "The merge will use approved staged text as the winner for any same-ID text conflicts and mark affected mappings or snapshots impacted. Continue?",
                parent=dialog,
            ):
                self.append_log("[GUI] Supplementary Source Approval: conflict winner approval cancelled by user.")
                return
            self.append_log(
                f"[GUI] Supplementary Source Approval: merge confirmed; closing review menu and launching {source_name} merge."
            )
            dialog.destroy()
            command = [PYTHON_EXE, str(SCRIPTS_DIR / "ingest_source_spec.py"), "merge", "--ingestion-dir", str(ingestion_dir)]
            for staged_id in approved_staged_ids:
                command.extend(["--conflict-winner-id", staged_id])
            self._run_command_async(
                command,
                f"Update integrated source baseline from supplementary source {source_name}",
                on_success=self._refresh_architecture_after_supplementary_merge,
            )

        def accept_all_and_run_stage0() -> None:
            self.append_log("[GUI] Supplementary Source Approval: Accept All Requirement IDs clicked.")
            if not messagebox.askyesno(
                "Accept All Requirement IDs",
                f"Set review_decision=approved for all {len(rows)} staged requirement IDs?\n\n"
                "First save and close the current comparison_results.xlsx in Excel. "
                "The GUI will then update that same file, synchronize the CSV, and reopen it.",
                parent=dialog,
            ):
                self.append_log("[GUI] Supplementary Source Approval: Accept All cancelled by user.")
                return
            if load_workbook is None:
                messagebox.showerror("Accept All Requirement IDs", "openpyxl is required to update the review workbook.", parent=dialog)
                return
            try:
                self.append_log(f"[GUI] Supplementary Source Approval: writing approved to {len(rows)} workbook rows.")
                workbook_state = self._close_saved_open_excel_workbook(review_workbook)
                if workbook_state == "same_name_other_path":
                    raise RuntimeError(
                        f"Excel has another workbook named {review_workbook.name} open from a different folder. "
                        "Close that workbook before using Accept All Requirement IDs."
                    )
                if workbook_state == "save_failed":
                    raise RuntimeError(
                        f"Save and close {review_workbook.name} in Excel first, then click Accept All Requirement IDs again."
                    )
                workbook = load_workbook(review_workbook)
                try:
                    sheet = workbook["Source Ingestion Review"]
                    headers = [str(cell.value or "").strip() for cell in sheet[1]]
                    decision_column = headers.index("review_decision") + 1
                    approved_fill = PatternFill("solid", fgColor="E2F0D9") if PatternFill else None
                    for row_index in range(2, sheet.max_row + 1):
                        sheet.cell(row=row_index, column=decision_column).value = "approved"
                        if approved_fill is not None:
                            sheet.cell(row=row_index, column=decision_column).fill = approved_fill
                    workbook.save(review_workbook)
                finally:
                    workbook.close()
                self.append_log(
                    "[GUI] Supplementary Source Approval: Accept All saved and reopened the workbook; "
                    "CSV synchronization is deferred until explicit Merge."
                )
            except PermissionError:
                self.append_log("[GUI] Supplementary Source Approval: Accept All blocked because the workbook is locked.")
                messagebox.showerror(
                    "Accept All Requirement IDs",
                    "Save and close the current comparison_results.xlsx in Excel, then click Accept All Requirement IDs again.",
                    parent=dialog,
                )
                return
            except Exception as exc:
                self.append_log(f"[GUI] Supplementary Source Approval: Accept All failed: {exc}")
                messagebox.showerror("Accept All Requirement IDs", str(exc), parent=dialog)
                return
            self.append_log(f"[GUI] Supplementary Source Approval: reopening saved workbook {review_workbook.name}.")
            self._open_mapping_workbook(review_workbook, dialog)
            self.append_log("[GUI] Supplementary Source Approval: Accept All completed; review menu remains open.")
            self.append_chat(f"GUI: Approved {len(rows)} supplementary source requirement IDs in the review workbook.")
            messagebox.showinfo(
                "Accept All Requirement IDs",
                "All review_decision cells are now approved and green.\n\n"
                "The saved workbook was reopened for crosscheck. CSV synchronization is deferred until you click "
                "Update the integrated source baseline and recompute dependent artifacts.",
                parent=dialog,
            )

        def open_review_workbook() -> None:
            self.append_log("[GUI] Supplementary Source Approval: Open Review Workbook in Excel clicked.")
            try:
                workbook_state = self._close_saved_open_excel_workbook(review_workbook)
                if workbook_state == "same_name_other_path":
                    raise RuntimeError(
                        f"Excel has another workbook named {review_workbook.name} open from a different folder. "
                        "Close that workbook before reopening this review."
                    )
                if workbook_state == "save_failed":
                    raise RuntimeError(
                        f"The review workbook has unsaved changes: {review_workbook.name}. "
                        "Save it in Excel, close it, and click Open Review Workbook in Excel again."
                    )
                _sync_saved_review_workbook(ingestion_dir)
                self.append_log("[GUI] Supplementary Source Approval: synchronized CSV and reapplied review colors before opening Excel.")
                self._open_mapping_workbook(review_workbook, dialog)
            except (OSError, RuntimeError, ValueError) as exc:
                self.append_log(f"[GUI] Supplementary Source Approval: workbook reopen blocked: {exc}")
                messagebox.showerror("Open Review Workbook", str(exc), parent=dialog)

        def close_review_dialog() -> None:
            self.append_log("[GUI] Supplementary Source Approval: Close clicked; review menu closed.")
            dialog.destroy()

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X)
        ttk.Button(
            buttons,
            text="Accept All Requirement IDs",
            command=accept_all_and_run_stage0,
        ).pack(fill=tk.X)
        ttk.Button(
            buttons,
            text="Update Integrated Source Baseline and Recompute Dependents",
            command=merge_and_run_stage0,
        ).pack(fill=tk.X, pady=(8, 0))
        ttk.Button(
            buttons,
            text="Open Review Workbook in Excel",
            command=open_review_workbook,
        ).pack(side=tk.LEFT, pady=(8, 0))
        ttk.Button(
            buttons,
            text="Refresh Review Colors",
            command=refresh_review_colors,
        ).pack(side=tk.LEFT, padx=(8, 0), pady=(8, 0))
        ttk.Button(buttons, text="Close", command=close_review_dialog).pack(side=tk.RIGHT, pady=(8, 0))

    def _show_table_row_manual_review_dialog(self) -> None:
        review_csv = REPO_ROOT / "artifacts/stage1_requirements/table_row_review.csv"
        review_request = REPO_ROOT / "artifacts/stage1_requirements/table_row_review_request.md"
        if self.table_review_dialog is not None and self.table_review_dialog.winfo_exists():
            self.table_review_dialog.lift()
            self.table_review_dialog.focus_force()
            return

        dialog = tk.Toplevel(self.root)
        dialog.title("Stage 1 Table Row Review Required")
        dialog.transient(self.root)
        dialog.resizable(False, False)
        self.table_review_dialog = dialog
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(
            frame,
            text="Stage 1 found ambiguous OCR table rows.",
            font=("Segoe UI", 13, "bold"),
        ).pack(anchor=tk.W)
        ttk.Label(
            frame,
            text=(
                "Open the review CSV in Excel, inspect or correct the reconstructed rows, "
                "then explicitly accept the review. Stage 1 must be rerun after acceptance."
            ),
            wraplength=560,
            justify=tk.LEFT,
        ).pack(anchor=tk.W, pady=(8, 12))

        def open_review_csv() -> None:
            if not review_csv.exists():
                messagebox.showerror("Table Row Review", "The table review CSV is missing.", parent=dialog)
                return
            try:
                subprocess.Popen(["cmd", "/c", "start", "", "excel.exe", str(review_csv)])
            except OSError:
                try:
                    os.startfile(str(review_csv))
                except OSError as exc:
                    messagebox.showerror("Table Row Review", str(exc), parent=dialog)

        def view_request() -> None:
            if review_request.exists():
                self._on_work_area_select_path(review_request)
            else:
                messagebox.showerror("Table Row Review", "The review request is missing.", parent=dialog)

        def accept_review() -> None:
            if not messagebox.askyesno(
                "Accept Table Row Review",
                "Have you reviewed and approved the ambiguous table rows in Excel?\n\n"
                "Stage 1 will remain stopped until it is rerun and passes validation.",
                parent=dialog,
            ):
                return
            self.append_chat("GUI: User approved the table-row manual review; rerun Stage 1 required.")
            self.append_log("GUI: User approved the table-row manual review; Stage 1 rerun required.")
            messagebox.showinfo(
                "Table Row Review Approved",
                "Review approved. Rerun Stage 1 to validate the reviewed table rows.",
                parent=dialog,
            )
            dialog.destroy()
            self.table_review_dialog = None

        def close_dialog() -> None:
            dialog.destroy()
            self.table_review_dialog = None

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X)
        ttk.Button(buttons, text="Open CSV in Excel", command=open_review_csv, width=21).pack(side=tk.LEFT)
        ttk.Button(buttons, text="View Review Request", command=view_request, width=20).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Button(buttons, text="Accept Review", command=accept_review, width=16).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Button(buttons, text="Close", command=close_dialog, width=10).pack(side=tk.RIGHT)

    def _open_table_row_review_csv(self) -> None:
        review_csv = REPO_ROOT / "artifacts/stage1_requirements/table_row_review.csv"
        if not review_csv.exists():
            messagebox.showinfo(
                "Table Row Review",
                "The table-row review CSV is not generated yet. Run Stage 1 first.",
                parent=self.root,
            )
            return
        try:
            subprocess.Popen(["cmd", "/c", "start", "", "excel.exe", str(review_csv)])
        except OSError:
            try:
                os.startfile(str(review_csv))
            except OSError as exc:
                messagebox.showerror("Table Row Review", str(exc), parent=self.root)

    @staticmethod
    def _resolve_gui_stage(stage: str) -> str:
        value = stage.strip()
        return SINGLE_STAGE_LABEL_TO_KEY.get(value, RANGE_STAGE_LABEL_TO_KEY.get(value, value)).lower()

    def run_stage_cli(self, stage: str) -> None:
        stage = self._resolve_gui_stage(stage)
        if stage == "3":
            self._prepare_stage3_run()
            return
        if stage in {"2b", "2c", "2g", "2h"}:
            self.append_chat(f"GUI: {stage.upper()} is an approval-gated workflow action; opening the approval controls.")
            self._show_architecture_profile_approval_dialog()
            return
        if stage == "2f":
            self.append_chat("GUI: S2F is Architecture Map Review; opening the review controls.")
            self._show_architecture_map_review_dialog()
            return
        if stage == "2d":
            messagebox.showinfo(
                "S2D Canonical Merge",
                "S2D is completed by the approval-gated merge action after impact review. "
                "Use the supplementary review or architecture approval controls to execute it.",
                parent=self.root,
            )
            return
        if stage == "2e":
            self.run_taxonomy_update()
            return
        if stage == "6":
            selector = self._snapshot_selector_args()
            if selector is None:
                return
            self._run_command_async(
                [PYTHON_EXE, str(SCRIPTS_DIR / "run_stage6_digital_ipos_gate.py"), *selector, "--regenerate-downstream"],
                "Stage 6 Digital IPOS",
                stage="6",
            )
            return
        if stage == "7":
            selector = self._snapshot_selector_args()
            if selector is None:
                return
            self._run_command_async(
                [PYTHON_EXE, str(SCRIPTS_DIR / "run_stage7_analog_ipos_gate.py"), *selector, "--regenerate-downstream"],
                "Stage 7 Analog IPOS",
            )
            return
        if stage == "arch-compare":
            self.run_stage6_compare()
            return
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "run", "--stage", stage]
        if stage.strip().lower() in {"3", "4", "5"}:
            selector = self._snapshot_selector_args()
            if selector is None:
                return
            cmd.extend(selector)
        self._run_command_async(
            cmd,
            f"CLI run stage {stage}",
            stage=stage,
            approval_after_success=stage == "2",
        )

    @staticmethod
    def _stage_result_passes(path: Path) -> bool:
        if not path.exists():
            return False
        try:
            return bool(re.search(r"^- Status:\s*PASS\s*$", path.read_text(encoding="utf-8"), re.MULTILINE | re.IGNORECASE))
        except OSError:
            return False

    def _first_stale_stage3_prerequisite(self) -> Optional[str]:
        previous_path = None
        for stage in ("0", "1", "2", "2a"):
            result_path = STAGE_RESULT_PATHS[stage]
            if not self._stage_result_passes(result_path):
                return stage
            if previous_path is not None and result_path.stat().st_mtime < previous_path.stat().st_mtime:
                return stage
            previous_path = result_path
        return None

    @staticmethod
    def _architecture_mapping_is_approved() -> bool:
        profile_path = REPO_ROOT / "config/stage2_mirco_arc_profile.json"
        try:
            profile = json.loads(profile_path.read_text(encoding="utf-8"))
            return str(profile.get("approval", {}).get("status") or "").lower() == "approved"
        except (OSError, ValueError, TypeError):
            return False

    @staticmethod
    def _latest_approved_snapshot_id() -> Optional[str]:
        try:
            context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8")) if PROJECT_CONTEXT_PATH.exists() else {}
            project_id = str(context.get("project_name") or REPO_ROOT.name)
            connection = connect(REPO_ROOT)
            try:
                row = connection.execute(
                    "SELECT snapshot_id FROM snapshots WHERE project_id = ? AND status = 'approved' ORDER BY approved_at DESC, snapshot_id DESC LIMIT 1",
                    (project_id,),
                ).fetchone()
            finally:
                connection.close()
            return str(row["snapshot_id"]) if row else None
        except Exception:
            return None

    @staticmethod
    def _sysml_review_is_current(snapshot_id: str) -> bool:
        manifest_path = REPO_ROOT / "artifacts/stage2_mirco_arc/stbio_architecture_model_manifest.json"
        review_path = REPO_ROOT / "artifacts/stage2_mirco_arc/stbio_architecture_model_review.md"
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            review = review_path.read_text(encoding="utf-8")
            return manifest.get("snapshot_id") == snapshot_id and "- Status: approved" in review and review_path.stat().st_mtime >= manifest_path.stat().st_mtime
        except (OSError, ValueError, TypeError):
            return False

    def _run_stage3_from_latest_snapshot(self) -> None:
        self._run_command_async(
            [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "run", "--stage", "3", "--use-latest-approved"],
            "CLI run stage 3",
            stage="3",
        )

    def _generate_sysml_then_run_stage3(self, snapshot_id: str) -> None:
        self._run_command_async(
            [PYTHON_EXE, str(SCRIPTS_DIR / "generate_architecture_sysml.py"), "--snapshot-id", snapshot_id],
            "Stage 3 SysML structural generation",
            on_success=lambda: self._run_command_async(
                [PYTHON_EXE, str(SCRIPTS_DIR / "review_architecture_sysml.py"), "--scope", "stage3-structural"],
                "Stage 3 SysML structural review",
                on_success=self._run_stage3_from_latest_snapshot,
            ),
        )

    def _show_stage3_snapshot_required(self) -> None:
        """Request only the remaining snapshot approval after a valid mapping review."""
        if not messagebox.askyesno(
            "Stage 3 Snapshot Required",
            "Architecture Map Review is already approved. Stage 3 needs an immutable approved snapshot before it can generate SysML and SRS. Create it now?",
            parent=self.root,
        ):
            return
        reviewer = simpledialog.askstring("Freeze Approved Snapshot", "Reviewer name or ID:", parent=self.root)
        if not reviewer or not reviewer.strip():
            return
        self._run_command_async(
            [PYTHON_EXE, str(SCRIPTS_DIR / "freeze_stage2b_snapshot.py"), "--reviewer", reviewer.strip()],
            "Freeze approved mapping snapshot for Stage 3",
            on_success=self._prepare_stage3_run,
        )

    def _run_stale_stage3_prefix(self, first_stale_stage: str) -> None:
        sequence = [stage for stage in ("0", "1", "2") if STAGE_KEYS.index(stage) >= STAGE_KEYS.index(first_stale_stage)]

        def run_next(index: int) -> None:
            if index >= len(sequence):
                self.append_chat("GUI: Stage 2 refresh completed. Architecture mapping approval is required before Stage 2A and Stage 3.")
                self._show_architecture_profile_approval_dialog()
                return
            stage = sequence[index]
            self.run_stage_script_after(stage, lambda: run_next(index + 1))

        run_next(0)

    def _prepare_stage3_run(self) -> None:
        stale_stage = self._first_stale_stage3_prerequisite()
        if stale_stage in {"0", "1", "2"}:
            self.append_chat(f"GUI: Stage 3 preflight found stale prerequisite S{stale_stage.upper()}; refreshing through S2.")
            self._run_stale_stage3_prefix(stale_stage)
            return
        if stale_stage == "2a":
            if not self._architecture_mapping_is_approved():
                self.append_chat("GUI: Stage 3 preflight requires Architecture Map Review approval before rerunning Stage 2A.")
                self._show_architecture_profile_approval_dialog()
                return
            self.append_chat("GUI: Stage 3 preflight is refreshing stale Stage 2A artifacts.")
            self._run_command_async(
                [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "run", "--stage", "2a"],
                "CLI refresh stage 2A for Stage 3",
                stage="2a",
                on_success=self._prepare_stage3_run,
            )
            return
        if not self._architecture_mapping_is_approved():
            self.append_chat("GUI: Stage 3 preflight requires current Architecture Map Review approval.")
            self._show_architecture_profile_approval_dialog()
            return
        snapshot_id = self._latest_approved_snapshot_id()
        if snapshot_id is None:
            self.append_chat("GUI: Architecture Map Review is current; Stage 3 requires only an approved immutable snapshot.")
            self._show_stage3_snapshot_required()
            return
        if not self._sysml_review_is_current(snapshot_id):
            self.append_chat("GUI: Stage 3 preflight is generating and reviewing SysML for the current approved snapshot.")
            self._generate_sysml_then_run_stage3(snapshot_id)
            return
        self._run_stage3_from_latest_snapshot()

    def run_selected_stage_cli(self) -> None:
        stage = self._resolve_gui_stage(self.stage_var.get())
        if stage == "regenerate_workbook":
            self.regenerate_review_workbook()
            return
        self.run_stage_cli(stage)

    def regenerate_review_workbook(self) -> None:
        """Rebuild only the review XLSX from the authoritative managed CSV."""
        mapping_csv = REPO_ROOT / "artifacts/stage1_specs/architecture_mapping_preview.csv"
        mapping_workbook = mapping_csv.with_suffix(".xlsx")
        self.append_log(
            f"[GUI][EXCEL] Regeneration requested: csv={mapping_csv.resolve()} "
            f"xlsx={mapping_workbook.resolve()}"
        )
        if not mapping_csv.exists():
            messagebox.showerror(
                "Regenerate Review Workbook",
                "The mapping CSV is missing. Run Stage 2 first to create the review package.",
                parent=self.root,
            )
            return
        if not messagebox.askyesno(
            "Regenerate Review Workbook",
            (
                "Regenerate the Architecture Mapping Review workbook from the current managed CSV?\n\n"
                "Only the XLSX will be rebuilt. Any Excel changes not already saved to the CSV will be lost."
            ),
            parent=self.root,
        ):
            return

        workbook_state = self._close_saved_open_excel_workbook(mapping_workbook)
        self.append_log(f"[GUI][EXCEL] Close/save result: {workbook_state or 'not_open'}")
        if workbook_state == "same_name_other_path":
            messagebox.showerror(
                "Regenerate Review Workbook",
                f"Excel has another workbook named {mapping_workbook.name} open from a different folder.",
                parent=self.root,
            )
            return
        if workbook_state == "save_failed":
            messagebox.showwarning(
                "Save and Close Workbook",
                f"Save and close {mapping_workbook.name} in Excel before regenerating it.",
                parent=self.root,
            )
            return

        try:
            if mapping_workbook.exists():
                self.append_log("[GUI][EXCEL] Syncing saved workbook values into managed CSV.")
                self._sync_mapping_workbook_to_csv(mapping_workbook, mapping_csv)
            self.append_log("[GUI][EXCEL] Generating review workbook from managed CSV.")
            regenerated_path = _write_mapping_workbook(mapping_csv)
            self.append_log(f"[GUI] Review workbook regenerated from CSV: {regenerated_path}")
            self.append_log(f"[GUI][EXCEL] Reopening generated workbook: {regenerated_path.resolve()}")
            self._open_mapping_workbook(regenerated_path, self.root)
            messagebox.showinfo(
                "Regenerate Review Workbook",
                "The review workbook was regenerated from the current CSV and reopened in Excel.",
                parent=self.root,
            )
        except Exception as exc:
            self.append_log(f"[GUI] Review workbook regeneration failed: {exc}")
            messagebox.showerror("Regenerate Review Workbook", str(exc), parent=self.root)

    def run_stage_range_cli(self) -> None:
        start_stage = self._resolve_gui_stage(self.from_stage_var.get())
        end_stage = self._resolve_gui_stage(self.to_stage_var.get())

        self.run_stage_range(start_stage, end_stage)

    def run_stage_range(self, start_stage: str, end_stage: str) -> None:
        start_stage = self._resolve_gui_stage(start_stage)
        end_stage = self._resolve_gui_stage(end_stage)
        range_order = [*STAGE_KEYS, "6", "7"]
        if start_stage not in range_order or end_stage not in range_order:
            messagebox.showinfo(
                "Unsupported Range",
                "Ranges support S0 through Stage 5 plus Digital IPOS and Analog IPOS. "
                "Approval-only workflow steps and optional architecture comparison run through their dedicated controls.",
                parent=self.root,
            )
            return

        try:
            start_idx = range_order.index(start_stage)
            end_idx = range_order.index(end_stage)
        except ValueError:
            messagebox.showerror("Invalid stage", "Please select valid start and end stages.")
            return

        if start_idx > end_idx:
            messagebox.showwarning(
                "Invalid range",
                "Start stage must be before or equal to end stage.",
            )
            return

        approval_stage_index = range_order.index("2")
        stage2a_index = range_order.index("2a")
        if start_idx <= approval_stage_index and end_idx >= stage2a_index:
            messagebox.showinfo(
                "Architecture Profile Approval Required",
                (
                    "Stage 2 generates the architecture-profile draft and opens the approval request.\n\n"
                    "Run through Stage 2 first, review and approve the profile, then run Stage 2A "
                    "and any required downstream range."
                ),
            )
            return

        selected_stages = range_order[start_idx : end_idx + 1]
        if any(stage in {"3", "4", "5", "6", "7"} for stage in selected_stages):
            selector = self._snapshot_selector_args()
            if selector is None:
                return
        else:
            selector = []

        ipos_stages = [stage for stage in selected_stages if stage in {"6", "7"}]

        def run_ipos(index: int = 0) -> None:
            if index >= len(ipos_stages):
                return
            stage = ipos_stages[index]
            script = "run_stage6_digital_ipos_gate.py" if stage == "6" else "run_stage7_analog_ipos_gate.py"
            label = "Stage 6 Digital IPOS" if stage == "6" else "Stage 7 Analog IPOS"
            self._run_command_async(
                [PYTHON_EXE, str(SCRIPTS_DIR / script), *selector],
                label,
                stage=stage,
                on_success=lambda: run_ipos(index + 1),
            )

        core_stages = [stage for stage in selected_stages if stage in STAGE_KEYS]
        if not core_stages:
            run_ipos()
            return
        cmd = [
            PYTHON_EXE,
            str(SCRIPTS_DIR / "workflow_cli.py"),
            "run",
            "--from-stage",
            core_stages[0],
            "--to-stage",
            core_stages[-1],
            *selector,
        ]
        self._run_command_async(
            cmd,
            f"CLI run range {start_stage} -> {end_stage}",
            stage=core_stages[0],
            approval_after_success=core_stages[-1] == "2",
            on_success=run_ipos if ipos_stages else None,
        )

    def _show_architecture_map_review_dialog(self) -> None:
        preview_dir = REPO_ROOT / "artifacts/stage1_specs"
        mapping_csv = preview_dir / "architecture_mapping_preview.csv"
        mapping_markdown = preview_dir / "architecture_mapping_preview.md"
        approved_profile = REPO_ROOT / "config/stage2_mirco_arc_profile.json"

        decisions = Counter()
        total_rows = 0
        csv_problem = ""
        if mapping_csv.exists():
            try:
                with mapping_csv.open("r", encoding="utf-8-sig", newline="") as handle:
                    for row in csv.DictReader(handle):
                        total_rows += 1
                        decision = (row.get("review_decision") or "blank").strip().lower() or "blank"
                        decisions[decision] += 1
            except Exception as exc:
                csv_problem = str(exc)

        gate_status = "not recorded"
        reviewer = ""
        if approved_profile.exists():
            try:
                approval = json.loads(approved_profile.read_text(encoding="utf-8")).get("approval", {})
                gate_status = str(approval.get("status") or "not recorded")
                reviewer = str(approval.get("approved_by") or "").strip()
            except Exception as exc:
                gate_status = f"unreadable ({exc})"

        dialog = tk.Toplevel(self.root)
        dialog.title("Architecture Map Review")
        dialog.transient(self.root)
        dialog.resizable(False, False)
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frame, text="Architecture Mapping Review", font=("Segoe UI", 13, "bold")).pack(anchor=tk.W)
        if not mapping_csv.exists():
            review_text = "Mapping CSV: not generated. Run Stage 2 to create the review package."
        elif csv_problem:
            review_text = f"Mapping CSV: could not read review decisions ({csv_problem})"
        else:
            decision_text = ", ".join(
                f"{decision}: {count}" for decision, count in sorted(decisions.items())
            ) or "no decisions"
            review_text = f"CSV review decisions ({total_rows} rows): {decision_text}"
        ttk.Label(frame, text=review_text, wraplength=590, justify=tk.LEFT).pack(anchor=tk.W, pady=(8, 2))

        gate_text = f"Stage 2A gate decision: {gate_status}"
        if reviewer:
            gate_text += f" (reviewer: {reviewer})"
        ttk.Label(frame, text=gate_text, wraplength=590, justify=tk.LEFT).pack(anchor=tk.W, pady=(0, 14))

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X)
        ttk.Button(buttons, text="Close", command=dialog.destroy, width=12).pack(side=tk.RIGHT)
        ttk.Button(
            buttons,
            text="View Markdown",
            command=lambda: self._on_work_area_select_path(mapping_markdown),
            width=16,
        ).pack(side=tk.LEFT)
        ttk.Button(
            buttons,
            text="View CSV",
            command=lambda: self._on_work_area_select_path(mapping_csv),
            width=12,
        ).pack(side=tk.LEFT, padx=(8, 0))

        def open_csv_in_excel() -> None:
            if not mapping_csv.exists():
                messagebox.showerror("Open Mapping CSV", "Mapping preview CSV is missing.", parent=dialog)
                return
            try:
                subprocess.Popen(["cmd", "/c", "start", "", "excel.exe", str(mapping_csv)])
            except OSError:
                try:
                    os.startfile(str(mapping_csv))
                except OSError as exc:
                    messagebox.showerror("Open Mapping CSV", str(exc), parent=dialog)

        ttk.Button(buttons, text="Open CSV in Excel", command=open_csv_in_excel, width=20).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Button(
            buttons,
            text="Approval / Snapshot Controls",
            command=self._show_architecture_profile_approval_dialog,
            width=28,
        ).pack(side=tk.LEFT, padx=(8, 0))

    def _open_mapping_workbook(self, workbook_path: Path, parent=None) -> None:
        self.append_log(f"[GUI][EXCEL] Opening review workbook: {workbook_path.resolve()}")
        if not workbook_path.exists():
            self.append_log(f"[GUI] Review workbook missing: {workbook_path}")
            messagebox.showerror("Open Mapping Workbook", "Mapping review workbook is missing.", parent=parent)
            return
        existing_state = self._activate_existing_excel_workbook(workbook_path)
        self.append_log(f"[GUI][EXCEL] Existing workbook state: {existing_state or 'not_found'}")
        if existing_state == "same_path":
            self.append_log(f"[GUI][EXCEL] Exact workbook already open and activated: {workbook_path.name}")
            return
        if existing_state == "same_path_readonly":
            self.append_log(
                f"[GUI] Review workbook is open read-only: {workbook_path.name}"
            )
            messagebox.showwarning(
                "Review Workbook Is Read-Only",
                f"{workbook_path.name} is already open in Excel as read-only. "
                "Close the other Excel workbook or instance, then open the review workbook again.",
                parent=parent,
            )
            return
        if existing_state == "same_name_other_path":
            self.append_log(
                f"[GUI] Excel already has another workbook named {workbook_path.name} open; "
                "did not open a duplicate."
            )
            messagebox.showwarning(
                "Review Workbook Already Open",
                f"Excel already has another workbook named {workbook_path.name} open from a different folder. "
                "Close that workbook before opening the current review workbook.",
                parent=parent,
            )
            return
        try:
            os.startfile(str(workbook_path))
            self.append_log(f"[GUI][EXCEL] Launch requested: {workbook_path.name}")
            self._confirm_mapping_workbook_open(workbook_path, parent=parent)
        except OSError as exc:
            self.append_log(f"[GUI] Review workbook launch failed: {exc}")
            messagebox.showerror("Open Mapping Workbook", str(exc), parent=parent)

    def _confirm_mapping_workbook_open(self, workbook_path: Path, parent=None, attempt: int = 0) -> None:
        """Confirm that Excel exposes the exact writable review workbook."""
        state = self._activate_existing_excel_workbook(workbook_path)
        self.append_log(
            f"[GUI][EXCEL] Attach check {attempt + 1}/20: path={workbook_path.resolve()} "
            f"state={state or 'not_found'}"
        )
        if state == "same_path":
            self.append_log(f"[GUI][EXCEL] Writable exact workbook confirmed: {workbook_path.name}")
            return
        if state == "same_name_other_path":
            self.append_log(
                f"[GUI] Another workbook named {workbook_path.name} is open; workbook not attached."
            )
            return
        if attempt < 20:
            if state == "same_path_readonly":
                self.append_log(
                    f"[GUI] Review workbook is temporarily read-only; retrying attach "
                    f"({attempt + 1}/20): {workbook_path.name}"
                )
            self.root.after(
                250,
                self._confirm_mapping_workbook_open,
                workbook_path,
                parent,
                attempt + 1,
            )
            return
        if state == "same_path_readonly":
            self.append_log(
                f"[GUI] Review workbook remained read-only: {workbook_path.name}"
            )
            messagebox.showwarning(
                "Review Workbook Is Read-Only",
                f"{workbook_path.name} remained open in Excel as read-only for 5 seconds. "
                "Close the other Excel workbook or instance, then open the review workbook again.",
                parent=parent,
            )
            return
        self.append_log(
            f"[GUI] Could not confirm writable exact review workbook: {workbook_path.name}"
        )

    def _ensure_architecture_block_reference_workbook(self) -> Optional[Path]:
        """Create the read-only block-name reference workbook for architecture review."""
        if Workbook is None:
            self.append_log("[GUI] Block reference workbook unavailable: openpyxl is not installed")
            return None
        block_defs: dict = {}
        inventory_path = REPO_ROOT / "artifacts/stage2_mirco_arc/block_inventory.csv"
        if inventory_path.exists():
            try:
                with inventory_path.open("r", encoding="utf-8-sig", newline="") as handle:
                    for row in csv.DictReader(handle):
                        block_name = (row.get("Block") or "").strip()
                        if block_name and block_name.casefold() != "unassigned":
                            block_defs[block_name] = {
                                "function": row.get("Function") or "",
                                "inputs": row.get("Inputs") or "",
                                "outputs": row.get("Outputs") or "",
                            }
            except OSError as exc:
                self.append_log(f"[GUI] Block inventory read failed: {exc}")
        if not block_defs:
            profile_paths = [
                REPO_ROOT / "config/stage2_mirco_arc_profile.json",
                REPO_ROOT / "artifacts/stage1_specs/architecture_profile_draft.json",
            ]
            for path in profile_paths:
                try:
                    candidate = json.loads(path.read_text(encoding="utf-8"))
                except (OSError, ValueError):
                    continue
                if isinstance(candidate, dict) and isinstance(candidate.get("block_defs"), dict):
                    block_defs = {
                        name: metadata
                        for name, metadata in candidate["block_defs"].items()
                        if name.casefold() != "unassigned"
                    }
                    break
        if not isinstance(block_defs, dict) or not block_defs:
            self.append_log("[GUI] Block reference workbook unavailable: no block_defs found")
            return None

        interface_owner_names = set()
        interface_catalog = REPO_ROOT / "artifacts/stage2_mirco_arc/interface_catalog.csv"
        if interface_catalog.exists():
            try:
                with interface_catalog.open("r", encoding="utf-8-sig", newline="") as handle:
                    catalog_owners = {
                        (row.get("Owner") or "").strip()
                        for row in csv.DictReader(handle)
                        if (row.get("Owner") or "").strip()
                    }
                    interface_owner_names = {
                        owner.casefold()
                        for owner in catalog_owners
                        if re.search(r"\b(?:interface|port|pin|endpoint)\b", owner, re.IGNORECASE)
                    }
            except OSError as exc:
                self.append_log(f"[GUI] Interface catalog read failed: {exc}")

        real_block_defs = {}
        for name, metadata in block_defs.items():
            if name.casefold() in interface_owner_names:
                continue
            function = str((metadata or {}).get("function") or "").casefold() if isinstance(metadata, dict) else ""
            if "convert" in function and ("adc" in function or "analog" in function):
                continue
            real_block_defs[name] = metadata
        block_defs = real_block_defs

        output_path = REPO_ROOT / "artifacts/stage1_specs/approved_block_list.xlsx"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Approved Block List"
        sheet.append(["Block name", "Classification", "Function", "Inputs", "Outputs", "Use in approved_block"])
        for block_name in sorted(block_defs, key=str.casefold):
            if block_name == "Unassigned":
                continue
            metadata = block_defs.get(block_name) or {}
            if not isinstance(metadata, dict):
                metadata = {}
            sheet.append([
                block_name,
                str(metadata.get("classification") or "Digital/System/Analog review").strip(),
                str(metadata.get("function") or "").strip(),
                str(metadata.get("inputs") or "").strip(),
                str(metadata.get("outputs") or "").strip(),
                "Copy the exact Block name into architecture_mapping_preview.csv",
            ])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in {"A": 24, "B": 24, "C": 72, "D": 42, "E": 42, "F": 58}.items():
            sheet.column_dimensions[column].width = width
        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True)
        workbook.save(output_path)
        self.append_log(f"[GUI] Block reference workbook generated: {output_path}")
        return output_path

    def _activate_existing_excel_workbook(self, workbook_path: Path) -> Optional[str]:
        """Return the state of an open workbook without launching a duplicate Excel window."""
        script = """Option Explicit
On Error Resume Next
Dim target, targetName, excel, book, candidate, fso, duplicateName
target = LCase(CStr(WScript.Arguments(0)))
Set fso = CreateObject("Scripting.FileSystemObject")
targetName = LCase(fso.GetFileName(target))
Set excel = GetObject(, "Excel.Application")
If Err.Number <> 0 Then WScript.Quit 2
Err.Clear
For Each candidate In excel.Workbooks
        If LCase(CStr(candidate.FullName)) = target Then
            If CBool(candidate.ReadOnly) Then WScript.Quit 7
            candidate.Activate
            WScript.Quit 0
    End If
    If LCase(CStr(candidate.Name)) = targetName Then duplicateName = True
Next
If duplicateName Then WScript.Quit 6
WScript.Quit 2
"""
        script_path = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", suffix=".vbs", delete=False, encoding="utf-8") as handle:
                handle.write(script)
                script_path = Path(handle.name)
            result = subprocess.run(
                ["cscript.exe", "//NoLogo", str(script_path), str(workbook_path.resolve())],
                cwd=str(REPO_ROOT),
                capture_output=True,
                text=True,
                timeout=10,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            self.append_log(f"[GUI] Could not inspect open Excel workbooks: {exc}")
            return None
        finally:
            if script_path is not None:
                try:
                    script_path.unlink()
                except OSError:
                    pass
        if result.returncode == 0:
            return "same_path"
        if result.returncode == 7:
            return "same_path_readonly"
        if result.returncode == 6:
            return "same_name_other_path"
        return None

    def _close_saved_open_excel_workbook(self, workbook_path: Path) -> Optional[str]:
        """Close the exact open workbook only after the user has saved it."""
        self.append_log(f"[GUI][EXCEL] Checking saved state before close: {workbook_path.resolve()}")
        script = """Option Explicit
On Error Resume Next
Dim target, targetName, excel, book, candidate, fso, duplicateName
target = LCase(CStr(WScript.Arguments(0)))
Set fso = CreateObject("Scripting.FileSystemObject")
targetName = LCase(fso.GetFileName(target))
Set excel = GetObject(, "Excel.Application")
If Err.Number <> 0 Then WScript.Quit 2
Err.Clear
For Each candidate In excel.Workbooks
    If LCase(CStr(candidate.FullName)) = target Then
        Set book = candidate
        Exit For
    End If
    If LCase(CStr(candidate.Name)) = targetName Then duplicateName = True
Next
If book Is Nothing Then
    If duplicateName Then WScript.Quit 6
    WScript.Quit 2
End If
If book.Saved = False Then WScript.Quit 5
book.Close False
If Err.Number <> 0 Then WScript.Quit 5
WScript.Quit 0
"""
        script_path = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", suffix=".vbs", delete=False, encoding="utf-8") as handle:
                handle.write(script)
                script_path = Path(handle.name)
            result = subprocess.run(
                ["cscript.exe", "//NoLogo", str(script_path), str(workbook_path.resolve())],
                cwd=str(REPO_ROOT),
                capture_output=True,
                text=True,
                timeout=20,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            self.append_log(f"[GUI] Could not inspect/close open review workbook: {exc}")
            return "save_failed"
        finally:
            if script_path is not None:
                try:
                    script_path.unlink()
                except OSError:
                    pass
        if result.returncode == 0:
            self.append_log(f"[GUI][EXCEL] Saved workbook closed: {workbook_path.name}")
            return "saved_closed"
        if result.returncode == 6:
            return "same_name_other_path"
        if result.returncode == 5:
            return "save_failed"
        return None

    def _approve_open_excel_workbook(self, workbook_path: Path) -> bool:
        """Update an already-open Excel workbook without requiring Excel to be closed."""
        script = """Option Explicit
On Error Resume Next
Dim target, excel, book, candidate, sheet, used, decisionColumn, column, row, cell, firstRow, lastRow, decisionRange
target = WScript.Arguments(0)
Set excel = GetObject(, "Excel.Application")
If Err.Number <> 0 Then WScript.Quit 2
Err.Clear
For Each candidate In excel.Workbooks
    If LCase(CStr(candidate.FullName)) = LCase(CStr(target)) Then
        Set book = candidate
        Exit For
    End If
Next
If book Is Nothing Then WScript.Quit 3
Set sheet = book.Worksheets.Item("Source Ingestion Review")
If Err.Number <> 0 Then WScript.Quit 4
Set used = sheet.UsedRange
decisionColumn = 0
For column = 1 To used.Columns.Count
    If LCase(CStr(sheet.Cells.Item(used.Row, used.Column + column - 1).Value2)) = "review_decision" Then
        decisionColumn = used.Column + column - 1
        Exit For
    End If
Next
If decisionColumn = 0 Then WScript.Quit 4
firstRow = used.Row + 1
lastRow = used.Row + used.Rows.Count - 1
Set decisionRange = sheet.Range(sheet.Cells.Item(firstRow, decisionColumn), sheet.Cells.Item(lastRow, decisionColumn))
decisionRange.Value2 = "approved"
decisionRange.Interior.Color = 14872793
For row = firstRow To lastRow
    If LCase(CStr(sheet.Cells.Item(row, decisionColumn).Value2)) <> "approved" Then WScript.Quit 5
Next
book.Save
If Err.Number <> 0 Then WScript.Quit 5
Err.Clear
If book.Saved = False Then WScript.Quit 5
book.Close False
If Err.Number <> 0 Then WScript.Quit 5
WScript.Quit 0
"""
        script_path = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", suffix=".vbs", delete=False, encoding="utf-8") as handle:
                handle.write(script)
                script_path = Path(handle.name)
            result = subprocess.run(
                ["cscript.exe", "//NoLogo", str(script_path), str(workbook_path.resolve())],
                cwd=str(REPO_ROOT),
                capture_output=True,
                text=True,
                timeout=20,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            self.append_log(f"[GUI] Open Excel refresh unavailable; using file update: {exc}")
            return False
        finally:
            if script_path is not None:
                try:
                    script_path.unlink()
                except OSError:
                    pass
        if result.returncode == 0:
            self.append_log(f"[GUI] Open Excel workbook refreshed green: {workbook_path.name}")
            return True
        if result.returncode == 2:
            self.append_log("[GUI] No active Excel instance found; using file update.")
        elif result.returncode == 3:
            self.append_log(f"[GUI] Excel is open, but the review workbook is not open: {workbook_path.name}")
        elif result.returncode == 5:
            self.append_log(f"[GUI] Excel found the review workbook but could not save it: {workbook_path.name}")
        else:
            detail = (result.stderr or result.stdout).strip()
            self.append_log(f"[GUI] Open Excel refresh skipped (code={result.returncode}): {detail}")
        return False

    def _sync_mapping_workbook_to_csv(self, workbook_path: Path, csv_path: Path) -> None:
        self.append_log(
            f"[GUI][EXCEL] XLSX->CSV sync start: workbook={workbook_path.resolve()} "
            f"csv={csv_path.resolve()}"
        )
        workbook_state = self._close_saved_open_excel_workbook(workbook_path)
        self.append_log(f"[GUI][EXCEL] XLSX->CSV close state: {workbook_state or 'not_open'}")
        if workbook_state == "same_name_other_path":
            raise RuntimeError(f"Excel has another workbook named {workbook_path.name} open from a different folder.")
        if workbook_state == "save_failed":
            raise RuntimeError(f"Save and close {workbook_path.name} in Excel before continuing.")
        count = sync_mapping_workbook_to_csv(workbook_path, csv_path)
        self.append_log(f"[GUI][EXCEL] XLSX->CSV sync complete: rows={count} workbook={workbook_path.name}")

    def _show_architecture_profile_approval_dialog(self) -> None:
        approval_request = REPO_ROOT / "artifacts/stage1_specs/architecture_profile_approval_request.md"
        profile_draft = REPO_ROOT / "artifacts/stage1_specs/architecture_profile_draft.json"
        mapping_preview = REPO_ROOT / "artifacts/stage1_specs/architecture_mapping_preview.csv"
        mapping_workbook = mapping_preview.with_suffix(".xlsx")
        approved_profile = REPO_ROOT / "config/stage2_mirco_arc_profile.json"
        if not approval_request.exists() or not profile_draft.exists() or not mapping_workbook.exists():
            self.append_chat("GUI: Architecture approval artifacts are missing after Stage 2.")
            return

        self._on_work_area_select_path(approval_request)
        self._open_mapping_workbook(mapping_workbook)
        block_reference_workbook = self._ensure_architecture_block_reference_workbook()
        if block_reference_workbook is not None:
            self._open_mapping_workbook(block_reference_workbook)
        if self.architecture_approval_dialog is not None and self.architecture_approval_dialog.winfo_exists():
            self.architecture_approval_dialog.lift()
            self.architecture_approval_dialog.focus_force()
            return

        try:
            draft_data = json.loads(profile_draft.read_text(encoding="utf-8"))
        except Exception:
            draft_data = {}
        approval = draft_data.get("approval", {}) if isinstance(draft_data, dict) else {}
        hashes = approval.get("evidence_hashes", {}) if isinstance(approval, dict) else {}
        dispositions = draft_data.get("ambiguity_dispositions", []) if isinstance(draft_data, dict) else []
        required_dispositions = [
            item for item in dispositions
            if isinstance(item, dict) and item.get("required_for_approval") == "yes"
        ]
        draft_hash = ""
        review_package = draft_data.get("review_package", {}) if isinstance(draft_data, dict) else {}
        if isinstance(review_package, dict):
            draft_hash = str(review_package.get("draft_profile_sha256") or "")
        freshness_text = "Evidence hashes available" if hashes and draft_hash else "Evidence hash metadata incomplete"

        dialog = tk.Toplevel(self.root)
        dialog.title("Architecture Profile Approval Required")
        dialog.transient(self.root)
        dialog.resizable(False, False)
        self.architecture_approval_dialog = dialog

        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(
            frame,
            text="Stage 2 completed. User approval is required before Stage 2A requirement mapping.",
            font=("Segoe UI", 13, "bold"),
            wraplength=560,
            justify=tk.LEFT,
        ).pack(anchor=tk.W)
        review_help = tk.Frame(frame, background="#fff3cd", highlightbackground="#d6b655", highlightthickness=1)
        review_help.pack(fill=tk.X, pady=(10, 8))
        tk.Label(
            review_help,
            text="1. Review and edit the workbook, then save it.  2. Enable the gate after every row is reviewed.",
            background="#fff3cd",
            foreground="#554000",
            font=("Segoe UI", 11, "bold"),
            justify=tk.LEFT,
            padx=10,
            pady=6,
            wraplength=540,
        ).pack(anchor=tk.W)
        choices_help = tk.Frame(frame, background="#e8f4fd", highlightbackground="#6ea8ce", highlightthickness=1)
        choices_help.pack(fill=tk.X, pady=(0, 10))
        tk.Label(
            choices_help,
            text=(
                "Workbook choices\n"
                "- review_decision: leave it pending while editing; a valid edited row becomes approved automatically, and the field remains editable.\n"
                "- approved_classification: Analog | Digital | System\n"
                "- approved_block: keep the candidate block when correct; otherwise replace it with the exact existing block name from the block inventory. For top-level ownership, choose System, Digital, or Analog.\n"
                "- Complete the allocation fields in the workbook. Do not invent a new block here.\n\n"
                "Color guide: red fields are incomplete or invalid; green means the row is explicitly reviewed and currently valid. "
                "After editing, save and close Excel, then select Enable Gate from Workbook.\n\n"
                "Supplementary source approval means the requirement entered the integrated corpus. "
                "This architecture review separately approves its block ownership and classification."
            ),
            background="#e8f4fd",
            foreground="#123b55",
            justify=tk.LEFT,
            padx=10,
            pady=6,
            wraplength=540,
        ).pack(anchor=tk.W)
        ttk.Label(
            frame,
            text=(
                f"Approval freshness: {freshness_text}\n"
                f"Critical/major ambiguity dispositions requiring approval: {len(required_dispositions)}\n"
                "Text, table, and figure extraction completeness is not enforced by this approval gate."
            ),
            wraplength=560,
            justify=tk.LEFT,
        ).pack(anchor=tk.W, pady=(0, 14))

        def close_dialog() -> None:
            if dialog.winfo_exists():
                dialog.destroy()
            self.architecture_approval_dialog = None

        def open_draft() -> None:
            self._on_work_area_select_path(profile_draft)

        def open_mapping_preview() -> None:
            self._on_work_area_select_path(mapping_preview)

        def open_mapping_preview_in_excel() -> None:
            self._open_mapping_workbook(mapping_workbook, dialog)

        def open_block_reference_in_excel() -> None:
            reference_path = self._ensure_architecture_block_reference_workbook()
            if reference_path is not None:
                self._open_mapping_workbook(reference_path, dialog)

        def run_stage2a_after_approval() -> None:
            try:
                self._sync_mapping_workbook_to_csv(mapping_workbook, mapping_preview)
            except Exception as exc:
                self._open_mapping_workbook(mapping_workbook, dialog)
                messagebox.showerror("Stage 2A Synchronization Required", str(exc), parent=dialog)
                return
            try:
                current_profile = json.loads(approved_profile.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                current_profile = {}
            current_approval = current_profile.get("approval", {}) if isinstance(current_profile, dict) else {}
            if not isinstance(current_approval, dict) or current_approval.get("status") != "approved":
                self._open_mapping_workbook(mapping_workbook, dialog)
                self.append_log(
                    "[GUI] Stage 2A approval is still required; reopened the architecture mapping workbook."
                )
                messagebox.showwarning(
                    "Stage 2A Approval Required",
                    "Review the reopened Architecture Map workbook, save and close it, "
                    "then click Enable Gate from Workbook before running Stage 2A.",
                    parent=dialog,
                )
                return
            self.append_chat("GUI: Stage 2A started; Architecture Profile Approval remains open for Stage 2B snapshot freezing.")
            self.run_stage_cli("2a")

        def freeze_stage2b_snapshot() -> None:
            try:
                self._sync_mapping_workbook_to_csv(mapping_workbook, mapping_preview)
                project_context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8"))
                validate_pre_freeze_coverage(
                    REPO_ROOT,
                    project_id=str(project_context.get("project_name") or REPO_ROOT.name),
                )
            except Exception as exc:
                self._open_mapping_workbook(mapping_workbook, dialog)
                messagebox.showerror("Stage 2B Pre-Freeze Validation Failed", str(exc), parent=dialog)
                return
            reviewer = simpledialog.askstring("Freeze Stage 2B Snapshot", "Reviewer name or ID:", parent=dialog)
            if not reviewer:
                return
            command = [
                PYTHON_EXE,
                str(SCRIPTS_DIR / "freeze_stage2b_snapshot.py"),
                "--reviewer",
                reviewer.strip(),
            ]
            close_dialog()
            self._run_command_async(
                command,
                "Freeze Stage 2B approved mapping snapshot",
                on_success=lambda: self.append_chat("GUI: Stage 2B approved mapping snapshot frozen."),
            )

        def approve_all_mapping_rows() -> None:
            self._open_mapping_workbook(mapping_workbook, dialog)
            messagebox.showinfo(
                "Review Workbook",
                "Edit and save the Architecture Map workbook. The tool will synchronize it to the managed CSV when you enable the gate.",
                parent=dialog,
            )

        def approve_reviewed_mapping_csv() -> None:
            if not mapping_preview.exists() or not mapping_workbook.exists():
                messagebox.showerror("Enable Stage 2A Gate", "Mapping review CSV or workbook is missing.", parent=dialog)
                return
            reviewer = simpledialog.askstring("Enable Stage 2A Gate", "Reviewer name or ID:", parent=dialog)
            if not reviewer:
                return
            try:
                self._sync_mapping_workbook_to_csv(mapping_workbook, mapping_preview)
                csv_findings = self._validate_mapping_preview_csv(mapping_preview)
                if csv_findings:
                    messagebox.showerror(
                        "Enable Stage 2A Gate",
                        "Fix the mapping preview CSV before enabling Stage 2A:\n\n" + "\n".join(csv_findings[:12]),
                        parent=dialog,
                    )
                    return
                digest = hashlib.sha256(mapping_preview.read_bytes()).hexdigest()
                profile = json.loads(approved_profile.read_text(encoding="utf-8"))
                draft = json.loads(profile_draft.read_text(encoding="utf-8"))
                approval = profile.setdefault("approval", {})
                draft_approval = draft.get("approval", {}) if isinstance(draft, dict) else {}
                review_package = draft.get("review_package", {}) if isinstance(draft, dict) else {}
                approval["status"] = "approved"
                approval["approved_by"] = reviewer.strip()
                approval["approved_at"] = datetime.now().isoformat(timespec="seconds")
                approval["evidence_hashes"] = draft_approval.get("evidence_hashes", {}) if isinstance(draft_approval, dict) else {}
                approval["reviewed_draft_profile_sha256"] = str(review_package.get("draft_profile_sha256") or "") if isinstance(review_package, dict) else ""
                approval["reviewed_mapping_preview_csv_sha256"] = digest
                if "ambiguity_dispositions" not in profile and isinstance(draft, dict):
                    profile["ambiguity_dispositions"] = draft.get("ambiguity_dispositions", [])
                approved_profile.write_text(json.dumps(profile, indent=2) + "\n", encoding="utf-8")
                self.append_chat("GUI: Stage 2A approval gate enabled from reviewed mapping workbook.")
                messagebox.showinfo("Enable Stage 2A Gate", "Reviewed mapping workbook imported and CSV hash recorded in the approved profile.", parent=dialog)
                self._on_work_area_select_path(approved_profile)
            except Exception as exc:
                messagebox.showerror("Enable Stage 2A Gate", str(exc), parent=dialog)

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X)
        ttk.Button(buttons, text="Close", command=close_dialog, width=12).pack(side=tk.RIGHT)
        ttk.Button(buttons, text="Run Stage 2A", command=run_stage2a_after_approval, width=14).pack(side=tk.RIGHT, padx=(0, 8))
        ttk.Button(buttons, text="Freeze Stage 2B Snapshot", command=freeze_stage2b_snapshot, width=23).pack(side=tk.RIGHT, padx=(0, 8))
        ttk.Button(buttons, text="Review Workbook", command=approve_all_mapping_rows, width=18).pack(side=tk.LEFT)
        tk.Button(
            buttons,
            text="Open Architecture Map in Excel",
            command=open_mapping_preview_in_excel,
            width=28,
            background="#ef8c00",
            foreground="white",
            activebackground="#c76f00",
            activeforeground="white",
            relief=tk.RAISED,
        ).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Button(
            buttons,
            text="Open Block List in Excel",
            command=open_block_reference_in_excel,
            width=24,
        ).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Button(buttons, text="View Mapping CSV", command=open_mapping_preview, width=18).pack(side=tk.LEFT, padx=(8, 0))
        tk.Button(
            buttons,
            text="Enable Gate from Workbook",
            command=approve_reviewed_mapping_csv,
            width=24,
            background="#2e7d32",
            foreground="white",
            activebackground="#1b5e20",
            activeforeground="white",
            relief=tk.RAISED,
        ).pack(side=tk.LEFT)
        ttk.Button(buttons, text="Open Draft", command=open_draft, width=12).pack(side=tk.LEFT, padx=(8, 0))
        dialog.protocol("WM_DELETE_WINDOW", close_dialog)

    def _validate_mapping_preview_csv(self, path: Path) -> list[str]:
        allowed = {"approved", "reassigned", "rejected", "needs_clarification"}
        approved_decisions = {"approved", "reassigned"}
        allowed_classifications = {"analog", "digital", "system"}
        allocation_targets = {
            "system_level": "srs",
            "top_digital_architecture": "drs",
            "top_analog_architecture": "ars",
            "block_local_digital": "digital ipos",
            "block_local_analog": "analog ipos",
            "descriptive_only": "none",
        }
        findings: list[str] = []
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            required = {
                "candidate_block",
                "requirement_id",
                "approved_classification",
                "review_decision",
                "approved_block",
                "allocation_class",
                "owning_target",
                "allocation_rationale",
                "lineage_mode",
            }
            missing = sorted(required - set(reader.fieldnames or []))
            if missing:
                return [f"missing required column(s): {', '.join(missing)}"]
            for index, row in enumerate(reader, start=2):
                req_id = (row.get("requirement_id") or "").strip()
                decision = (row.get("review_decision") or "").strip().lower()
                approved_block = (row.get("approved_block") or "").strip()
                allocation_class = (row.get("allocation_class") or "").strip().casefold()
                owning_target = (row.get("owning_target") or "").strip().casefold()
                allocation_rationale = (row.get("allocation_rationale") or "").strip()
                lineage_mode = (row.get("lineage_mode") or "").strip()
                candidate_block = (row.get("candidate_block") or "").strip()
                generated_preview = (row.get("generated_block_paragraph_preview") or "").strip().lower()
                approved_classification = (row.get("approved_classification") or "").strip().lower()
                label = req_id or f"row {index}"
                unresolved = (
                    candidate_block.lower() == "unresolved source paragraph"
                    or "unresolved source paragraph" in generated_preview
                )
                if unresolved and (decision in approved_decisions or approved_block.lower() == "unresolved source paragraph"):
                    findings.append(
                        f"{label}: unresolved source paragraph requires user block mapping before Stage 2A approval"
                    )
                if approved_classification not in allowed_classifications:
                    findings.append(f"{label}: approved_classification must be Analog, Digital, or System")
                if allocation_class not in allocation_targets:
                    findings.append(f"{label}: allocation_class is missing or invalid")
                elif owning_target != allocation_targets[allocation_class]:
                    findings.append(
                        f"{label}: owning_target={owning_target or '<empty>'} does not match "
                        f"allocation_class={allocation_class}"
                    )
                if not allocation_rationale or allocation_rationale.startswith("REVIEW_REQUIRED:"):
                    findings.append(f"{label}: allocation_rationale requires review")
                if not lineage_mode:
                    findings.append(f"{label}: lineage_mode is required")
                if decision not in allowed:
                    findings.append(f"{label}: review_decision must be one of {', '.join(sorted(allowed))}")
                elif decision not in approved_decisions:
                    findings.append(f"{label}: review_decision={decision} blocks approval")
                elif decision == "reassigned" and not approved_block:
                    findings.append(f"{label}: reassigned rows require approved_block")
                elif decision == "approved" and not approved_block:
                    findings.append(f"{label}: approved rows require approved_block")
                elif approved_block.casefold() in {"system", "digital", "analog"} and approved_block.casefold() != approved_classification:
                    findings.append(
                        f"{label}: top-level approved_block={approved_block} must match "
                        f"approved_classification={approved_classification.title()}"
                    )
        return findings

    def run_stage6_compare(self) -> None:
        if self.stage6_dialog is not None and self.stage6_dialog.winfo_exists():
            self.stage6_dialog.lift()
            self.stage6_dialog.focus_force()
            return

        dialog = tk.Toplevel(self.root)
        dialog.title("Optional Architecture Comparison")
        dialog.transient(self.root)
        self._set_responsive_geometry(dialog, 620, 220)
        dialog.resizable(True, False)
        dialog.grab_set()
        self.stage6_dialog = dialog

        dialog_content = ttk.Frame(dialog)
        dialog_content.pack(fill=tk.BOTH, expand=True)

        dialog_canvas = tk.Canvas(dialog_content, highlightthickness=0)
        dialog_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        dialog_hscroll = ttk.Scrollbar(
            dialog,
            orient=tk.HORIZONTAL,
            command=dialog_canvas.xview,
        )
        dialog_hscroll.pack(fill=tk.X, padx=12, pady=(0, 12))
        dialog_canvas.configure(xscrollcommand=dialog_hscroll.set)

        frame = ttk.Frame(dialog_canvas, padding=12)
        dialog_canvas.create_window((0, 0), window=frame, anchor="nw")

        def update_dialog_scrollregion(_event=None) -> None:
            dialog_canvas.configure(scrollregion=dialog_canvas.bbox("all"))

        frame.bind("<Configure>", update_dialog_scrollregion)

        ttk.Label(
            frame,
            text="Select two projects for a deterministic architecture comparison.",
        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))

        current_project = REPO_ROOT.name
        project_names = self._list_workspace_projects()
        if current_project not in project_names:
            project_names.insert(0, current_project)

        ttk.Label(frame, text="Base project:").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
        base_var = tk.StringVar(value=current_project)
        base_combo = ttk.Combobox(frame, textvariable=base_var, values=project_names, state="readonly", width=36)
        base_combo.grid(row=1, column=1, sticky="w", pady=4)

        ttk.Label(frame, text="Project to compare:").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=4)
        default_compare = next((p for p in project_names if p != current_project), current_project)
        compare_var = tk.StringVar(value=default_compare)
        compare_combo = ttk.Combobox(frame, textvariable=compare_var, values=project_names, state="readonly", width=36)
        compare_combo.grid(row=2, column=1, sticky="w", pady=4)

        hint = (
            "Execution uses existing scripts/workflow_cli.py only. No LLM calls are made.\n"
            "Results are written under the selected base project artifacts/comparison."
        )
        ttk.Label(frame, text=hint, wraplength=450).grid(row=3, column=0, columnspan=2, sticky="w", pady=(8, 10))

        btn_row = ttk.Frame(frame)
        btn_row.grid(row=4, column=0, columnspan=2, sticky="e")

        def close_dialog() -> None:
            if dialog.winfo_exists():
                dialog.destroy()
            self.stage6_dialog = None

        def run_compare() -> None:
            base_project = base_var.get().strip()
            compare_project = compare_var.get().strip()
            if not base_project or not compare_project:
                messagebox.showwarning("Architecture Comparison", "Please select both projects.", parent=dialog)
                return
            if base_project == compare_project:
                messagebox.showwarning("Architecture Comparison", "Please choose two different projects.", parent=dialog)
                return

            base_root = self.workspace_root / base_project
            cli_path = base_root / "scripts" / "workflow_cli.py"
            if not cli_path.exists():
                messagebox.showerror(
                    "Architecture Comparison",
                    f"Missing workflow CLI for base project: {cli_path}",
                    parent=dialog,
                )
                return

            close_dialog()
            cmd = [
                PYTHON_EXE,
                str(cli_path),
                "arch-compare",
                "--project-to-compare",
                compare_project,
                "--workspace-root",
                str(self.workspace_root),
            ]
            self._run_command_async(cmd, f"Architecture comparison ({base_project} vs {compare_project})")

        ttk.Button(btn_row, text="Cancel", command=close_dialog, width=12).pack(side=tk.RIGHT)
        ttk.Button(btn_row, text="Run Compare", command=run_compare, width=14).pack(side=tk.RIGHT, padx=(0, 8))

        dialog.protocol("WM_DELETE_WINDOW", close_dialog)

    def _list_workspace_projects(self) -> list[str]:
        names = []
        try:
            for child in sorted(self.workspace_root.iterdir(), key=lambda p: p.name.lower()):
                if not child.is_dir():
                    continue
                if (child / "scripts" / "workflow_cli.py").exists():
                    names.append(child.name)
        except Exception:
            pass
        return names

    def run_validate_all(self) -> None:
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "validate", "--all"]
        self._run_command_async(cmd, "CLI validate --all")

    def run_validate_downstream(self) -> None:
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "validate-downstream"]
        self._run_command_async(cmd, "CLI validate-downstream")

    def _is_safe_work_path(self, path: Path) -> Optional[Path]:
        try:
            resolved = path.resolve()
            resolved.relative_to(REPO_ROOT.resolve())
            return resolved
        except (OSError, ValueError):
            return None

    def _work_area_path(self, item_id: str) -> Optional[Path]:
        relative_path = self.work_tree.item(item_id, "values")
        if not relative_path:
            return None
        return self._is_safe_work_path(REPO_ROOT / relative_path[0])

    def _insert_work_tree_dir(self, parent_id: str, directory: Path) -> None:
        try:
            children = sorted(directory.iterdir(), key=lambda item: (not item.is_dir(), item.name.lower()))
        except OSError:
            return

        for child in children:
            if child.name.startswith(".") or child.name in {"__pycache__", ".pytest_cache"}:
                continue
            relative_path = child.relative_to(REPO_ROOT).as_posix()
            if child.is_dir():
                node = self.work_tree.insert(parent_id, tk.END, text=child.name, values=(relative_path,), open=False)
                self._insert_work_tree_dir(node, child)
            elif child.is_file():
                if child.suffix.lower() == ".sysml" and not self._sysml_block_file_is_concrete(child):
                    continue
                self.work_tree.insert(parent_id, tk.END, text=child.name, values=(relative_path,))

    def refresh_work_area(self) -> None:
        if not hasattr(self, "work_tree"):
            return

        open_path = self.current_work_path

        for item in self.work_tree.get_children():
            self.work_tree.delete(item)

        for root_name in WORK_AREA_ROOTS:
            directory = REPO_ROOT / root_name
            if directory.is_dir():
                node = self.work_tree.insert("", tk.END, text=root_name, values=(root_name,), open=False)
                self._insert_work_tree_dir(node, directory)

        for root_file in sorted(REPO_ROOT.iterdir(), key=lambda item: item.name.lower()):
            if root_file.is_file() and not root_file.name.startswith("."):
                self.work_tree.insert("", tk.END, text=root_file.name, values=(root_file.name,))

        if open_path is not None and open_path.exists() and open_path.is_file():
            if open_path.suffix.lower() == ".pdf":
                self.work_file_var.set(
                    f"PDF selected; use the Work Area selection to open in Adobe Acrobat Reader: "
                    f"{open_path.relative_to(REPO_ROOT).as_posix()}"
                )
            else:
                self._on_work_area_select_path(open_path)
            return

        self._set_work_view("Select a file from the Work Area navigation panel.", "No file selected")

    def _refresh_work_area_and_traceability(self) -> None:
        self.refresh_work_area()
        if hasattr(self, "coverage_canvas"):
            self._refresh_system_traceability()
        if hasattr(self, "sysml_tree"):
            self._refresh_sysml_architecture()
        if hasattr(self, "diagram_canvas"):
            self._refresh_workflow_diagram()

    def _refresh_all_derived_views(self) -> None:
        """Refresh reports first, then reload traceability, SysML, and workflow diagrams."""
        self._refresh_traceability_after_generation(self._refresh_work_area_and_traceability)

    def _build_sysml_architecture_tab(self) -> None:
        controls = ttk.Frame(self.sysml_architecture_tab)
        controls.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(controls, text="Generated SysML hierarchy and block models", font=("Segoe UI", 12, "bold")).pack(side=tk.LEFT)
        self.sysml_graph_zoom = 1.0
        self.sysml_graph_zoom_var = tk.StringVar(value="Zoom: 100%")
        ttk.Label(controls, textvariable=self.sysml_graph_zoom_var).pack(side=tk.RIGHT, padx=(8, 0))
        reset_zoom_btn = ttk.Button(controls, text="Reset Zoom", command=self._reset_sysml_graph_zoom, width=12)
        reset_zoom_btn.pack(side=tk.RIGHT, padx=(6, 0))
        zoom_out_btn = ttk.Button(controls, text="-", command=lambda: self._zoom_sysml_graph(1 / 1.15), width=3)
        zoom_out_btn.pack(side=tk.RIGHT)
        zoom_in_btn = ttk.Button(controls, text="+", command=lambda: self._zoom_sysml_graph(1.15), width=3)
        zoom_in_btn.pack(side=tk.RIGHT, padx=(6, 0))
        self._bind_button_help(reset_zoom_btn, "Return the SysML graph to 100 percent.")
        self._bind_button_help(zoom_out_btn, "Zoom out the SysML graph.")
        self._bind_button_help(zoom_in_btn, "Zoom in the SysML graph.")
        refresh_btn = ttk.Button(controls, text="Refresh", command=self._refresh_sysml_architecture, width=10)
        refresh_btn.pack(side=tk.RIGHT)
        self._bind_button_help(refresh_btn, "Reload the generated SysML top, DigitalSubsystem, shared definitions, and block files.")
        hierarchy_window_btn = ttk.Button(controls, text="Hierarchy Window", command=self._open_sysml_hierarchy_window, width=17)
        hierarchy_window_btn.pack(side=tk.RIGHT, padx=(6, 0))
        self._bind_button_help(hierarchy_window_btn, "Open the graphical SysML hierarchy navigator in a separate window.")

        self.sysml_vertical_pane = ttk.Panedwindow(self.sysml_architecture_tab, orient=tk.HORIZONTAL)
        self.sysml_vertical_pane.pack(fill=tk.BOTH, expand=True)
        vertical_pane = self.sysml_vertical_pane
        graph_frame = ttk.LabelFrame(vertical_pane, text="SysML Block Hierarchy", padding=4)
        self.sysml_graph_canvas = tk.Canvas(graph_frame, height=250, bg="#f7f9fb", highlightthickness=1, highlightbackground="#c8d2dc")
        graph_y = ttk.Scrollbar(graph_frame, orient=tk.VERTICAL, command=self.sysml_graph_canvas.yview)
        graph_x = ttk.Scrollbar(graph_frame, orient=tk.HORIZONTAL, command=self.sysml_graph_canvas.xview)
        self.sysml_graph_canvas.configure(yscrollcommand=graph_y.set, xscrollcommand=graph_x.set)
        self.sysml_graph_canvas.grid(row=0, column=0, sticky="nsew")
        graph_y.grid(row=0, column=1, sticky="ns")
        graph_x.grid(row=1, column=0, sticky="ew", pady=(2, 6))
        graph_frame.grid_rowconfigure(0, weight=1)
        graph_frame.grid_columnconfigure(0, weight=1)
        self.sysml_graph_canvas.bind("<MouseWheel>", lambda event: self.sysml_graph_canvas.yview_scroll(-1 if event.delta > 0 else 1, "units"), add="+")
        self.sysml_graph_canvas.bind("<Shift-MouseWheel>", lambda event: self.sysml_graph_canvas.xview_scroll(-1 if event.delta > 0 else 1, "units"), add="+")
        self.sysml_graph_canvas.bind("<Control-MouseWheel>", self._zoom_sysml_graph_wheel, add="+")
        self.sysml_horizontal_pane = ttk.Panedwindow(vertical_pane, orient=tk.HORIZONTAL)
        pane = self.sysml_horizontal_pane
        tree_frame = ttk.LabelFrame(pane, text="SysML Source Files", padding=4)
        text_frame = ttk.LabelFrame(pane, text="Selected Source", padding=4)
        pane.add(tree_frame, weight=1)
        pane.add(text_frame, weight=3)
        vertical_pane.add(graph_frame, weight=2)
        vertical_pane.add(pane, weight=3)

        navigation_pane = ttk.Panedwindow(tree_frame, orient=tk.VERTICAL)
        navigation_pane.pack(fill=tk.BOTH, expand=True)
        hierarchy_frame = ttk.LabelFrame(navigation_pane, text="Hierarchy Navigation", padding=3)
        source_tree_frame = ttk.LabelFrame(navigation_pane, text="Source Files", padding=3)
        navigation_pane.add(hierarchy_frame, weight=2)
        navigation_pane.add(source_tree_frame, weight=3)

        self.sysml_hierarchy_tree = ttk.Treeview(hierarchy_frame, show="tree", selectmode="browse")
        hierarchy_scroll = ttk.Scrollbar(hierarchy_frame, orient=tk.VERTICAL, command=self.sysml_hierarchy_tree.yview)
        self.sysml_hierarchy_tree.configure(yscrollcommand=hierarchy_scroll.set)
        self.sysml_hierarchy_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        hierarchy_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.sysml_hierarchy_tree.bind("<<TreeviewSelect>>", self._on_sysml_hierarchy_select)

        self.sysml_tree = ttk.Treeview(tree_frame, show="tree", selectmode="browse")
        tree_scroll = ttk.Scrollbar(source_tree_frame, orient=tk.VERTICAL, command=self.sysml_tree.yview)
        self.sysml_tree.configure(yscrollcommand=tree_scroll.set)
        self.sysml_tree.pack(in_=source_tree_frame, side=tk.LEFT, fill=tk.BOTH, expand=True)
        tree_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.sysml_tree.bind("<<TreeviewSelect>>", self._on_sysml_select)

        self.sysml_file_var = tk.StringVar(value="No SysML file selected")
        source_controls = ttk.Frame(text_frame)
        source_controls.pack(fill=tk.X, pady=(0, 4))
        ttk.Label(source_controls, text="File:").pack(side=tk.LEFT)
        self.sysml_file_choice_var = tk.StringVar(value="")
        self.sysml_file_choice = ttk.Combobox(
            source_controls,
            textvariable=self.sysml_file_choice_var,
            state="readonly",
            width=42,
        )
        self.sysml_file_choice.pack(side=tk.LEFT, padx=(4, 8), fill=tk.X, expand=True)
        self.sysml_file_choice.bind("<<ComboboxSelected>>", self._on_sysml_file_choice)
        ttk.Label(source_controls, textvariable=self.sysml_file_var, anchor="w").pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.sysml_source_zoom_var = tk.StringVar(value="Source: 100%")
        ttk.Label(source_controls, textvariable=self.sysml_source_zoom_var).pack(side=tk.RIGHT, padx=(8, 0))
        source_reset_btn = ttk.Button(source_controls, text="Reset", command=self._reset_sysml_source_zoom, width=7)
        source_reset_btn.pack(side=tk.RIGHT, padx=(6, 0))
        source_zoom_out_btn = ttk.Button(source_controls, text="-", command=lambda: self._zoom_sysml_source(1 / 1.15), width=3)
        source_zoom_out_btn.pack(side=tk.RIGHT)
        source_zoom_in_btn = ttk.Button(source_controls, text="+", command=lambda: self._zoom_sysml_source(1.15), width=3)
        source_zoom_in_btn.pack(side=tk.RIGHT, padx=(6, 0))
        self._bind_button_help(source_reset_btn, "Reset selected SysML source text size.")
        self._bind_button_help(source_zoom_out_btn, "Zoom out selected SysML source text.")
        self._bind_button_help(source_zoom_in_btn, "Zoom in selected SysML source text.")
        text_area = ttk.Frame(text_frame)
        text_area.pack(fill=tk.BOTH, expand=True)
        self.sysml_source_font = tkfont.Font(self.root, family="Consolas", size=10)
        self.sysml_viewer = tk.Text(text_area, wrap=tk.NONE, font=self.sysml_source_font)
        self.sysml_viewer.tag_configure("find_hit", background="#fff2a8")
        self.sysml_viewer.tag_configure("find_current", background="#ffbf47")
        text_scroll_y = ttk.Scrollbar(text_area, orient=tk.VERTICAL, command=self.sysml_viewer.yview)
        text_scroll_x = ttk.Scrollbar(text_frame, orient=tk.HORIZONTAL, command=self.sysml_viewer.xview)
        self.sysml_viewer.configure(yscrollcommand=text_scroll_y.set, xscrollcommand=text_scroll_x.set)
        self.sysml_viewer.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self._bind_copy_support(self.sysml_viewer)
        self.sysml_viewer.bind("<KeyPress>", self._block_rendered_text_edit, add="+")
        self.sysml_viewer.bind("<Control-f>", lambda _event: self._open_sysml_find(self.sysml_viewer), add="+")
        self.sysml_viewer.bind("<Control-F>", lambda _event: self._open_sysml_find(self.sysml_viewer), add="+")
        text_scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        text_scroll_x.pack(fill=tk.X, pady=(4, 0))
        self.sysml_viewer.bind("<Control-MouseWheel>", self._zoom_sysml_source_wheel, add="+")
        self._refresh_sysml_architecture()

    def _zoom_sysml_source(self, factor: float) -> str:
        new_zoom = min(2.5, max(0.6, self.sysml_source_zoom * factor))
        self.sysml_source_zoom = new_zoom
        self.sysml_source_font.configure(size=max(8, round(10 * new_zoom)))
        self.sysml_source_zoom_var.set(f"Source: {round(new_zoom * 100)}%")
        return "break"

    def _zoom_sysml_source_wheel(self, event) -> str:
        return self._zoom_sysml_source(1.15 if event.delta > 0 else 1 / 1.15)

    def _reset_sysml_source_zoom(self) -> None:
        self.sysml_source_zoom = 1.0
        self.sysml_source_font.configure(size=10)
        self.sysml_source_zoom_var.set("Source: 100%")

    def _zoom_sysml_graph(self, factor: float, origin=None) -> str:
        if not hasattr(self, "sysml_graph_canvas"):
            return "break"
        new_zoom = min(2.5, max(0.5, self.sysml_graph_zoom * factor))
        factor = new_zoom / self.sysml_graph_zoom
        self.sysml_graph_zoom = new_zoom
        canvas = self.sysml_graph_canvas
        if origin is None:
            origin = (canvas.winfo_width() / 2, canvas.winfo_height() / 2)
        canvas.scale("all", origin[0], origin[1], factor, factor)
        left, top, right, bottom = canvas.bbox("all") or (0, 0, 600, 100)
        canvas.configure(scrollregion=(min(0, left), min(0, top), max(600, right), max(100, bottom)))
        self.sysml_graph_zoom_var.set(f"Zoom: {round(self.sysml_graph_zoom * 100)}%")
        return "break"

    def _zoom_sysml_graph_wheel(self, event) -> str:
        factor = 1.15 if event.delta > 0 else 1 / 1.15
        return self._zoom_sysml_graph(factor, (event.x, event.y))

    def _open_sysml_hierarchy_window(self) -> None:
        existing = getattr(self, "sysml_hierarchy_window", None)
        if existing is not None and existing.winfo_exists():
            existing.deiconify()
            existing.lift()
            existing.focus_force()
            self._draw_sysml_hierarchy_window()
            return

        window = tk.Toplevel(self.root)
        window.title("SysML Hierarchy Navigation")
        self._set_responsive_geometry(window, 1100, 700, 620, 420)
        self.sysml_hierarchy_window = window
        header = ttk.Frame(window, padding=(8, 8, 8, 4))
        header.pack(fill=tk.X)
        ttk.Label(header, text="Graphical SysML Hierarchy Navigation", font=("Segoe UI", 12, "bold")).pack(side=tk.LEFT)
        self.sysml_hierarchy_zoom = getattr(self, "sysml_hierarchy_zoom", 1.0)
        self.sysml_hierarchy_zoom_var = tk.StringVar(value=f"Zoom: {round(self.sysml_hierarchy_zoom * 100)}%")
        ttk.Label(header, textvariable=self.sysml_hierarchy_zoom_var).pack(side=tk.RIGHT, padx=(8, 0))
        reset_hierarchy_zoom_btn = ttk.Button(header, text="Reset Zoom", command=self._reset_sysml_hierarchy_zoom, width=12)
        reset_hierarchy_zoom_btn.pack(side=tk.RIGHT, padx=(6, 0))
        hierarchy_zoom_out_btn = ttk.Button(header, text="-", command=lambda: self._zoom_sysml_hierarchy(1 / 1.15), width=3)
        hierarchy_zoom_out_btn.pack(side=tk.RIGHT)
        hierarchy_zoom_in_btn = ttk.Button(header, text="+", command=lambda: self._zoom_sysml_hierarchy(1.15), width=3)
        hierarchy_zoom_in_btn.pack(side=tk.RIGHT, padx=(6, 0))
        self._bind_button_help(reset_hierarchy_zoom_btn, "Return the hierarchy navigator to 100 percent.")
        self._bind_button_help(hierarchy_zoom_out_btn, "Zoom out the hierarchy navigator.")
        self._bind_button_help(hierarchy_zoom_in_btn, "Zoom in the hierarchy navigator.")
        ttk.Button(header, text="Refresh", command=self._refresh_sysml_hierarchy_window, width=10).pack(side=tk.RIGHT)
        source_selector = ttk.Frame(window, padding=(8, 0, 8, 4))
        source_selector.pack(fill=tk.X)
        ttk.Label(source_selector, text="SysML source:").pack(side=tk.LEFT)
        self.sysml_hierarchy_source_var = tk.StringVar(value="")
        self.sysml_hierarchy_source_paths = {}
        self.sysml_hierarchy_source_choice = ttk.Combobox(
            source_selector,
            textvariable=self.sysml_hierarchy_source_var,
            state="readonly",
            width=72,
        )
        self.sysml_hierarchy_source_choice.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)
        self.sysml_hierarchy_source_choice.bind("<<ComboboxSelected>>", self._on_sysml_hierarchy_source_choice)
        ttk.Button(
            source_selector,
            text="Find (Ctrl+F)",
            command=lambda: self._open_sysml_find(self.sysml_hierarchy_details),
            width=15,
        ).pack(side=tk.RIGHT, padx=(6, 0))
        self._bind_button_help(source_selector, "Select any generated .sysml file to inspect it in this hierarchy window.")
        content_pane = ttk.Panedwindow(window, orient=tk.VERTICAL)
        content_pane.pack(fill=tk.BOTH, expand=True, padx=8, pady=(4, 8))
        canvas_frame = ttk.Frame(content_pane, padding=(0, 0, 0, 4))
        content_pane.add(canvas_frame, weight=4)
        canvas = tk.Canvas(canvas_frame, background="#f7f9fb", highlightthickness=1, highlightbackground="#c8d2dc")
        scroll_y = ttk.Scrollbar(canvas_frame, orient=tk.VERTICAL, command=canvas.yview)
        scroll_x = ttk.Scrollbar(canvas_frame, orient=tk.HORIZONTAL, command=canvas.xview)
        canvas.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        scroll_y.grid(row=0, column=1, sticky="ns")
        scroll_x.grid(row=1, column=0, sticky="ew")
        canvas_frame.grid_rowconfigure(0, weight=1)
        canvas_frame.grid_columnconfigure(0, weight=1)
        canvas.bind("<MouseWheel>", lambda event: canvas.yview_scroll(-1 if event.delta > 0 else 1, "units"), add="+")
        canvas.bind("<Shift-MouseWheel>", lambda event: canvas.xview_scroll(-1 if event.delta > 0 else 1, "units"), add="+")
        canvas.bind("<Control-MouseWheel>", self._zoom_sysml_hierarchy_wheel, add="+")
        self.sysml_hierarchy_window_canvas = canvas
        details_frame = ttk.LabelFrame(content_pane, text="Selected Block Port Map and Full SysML Source", padding=6)
        content_pane.add(details_frame, weight=3)
        port_frame = ttk.Frame(details_frame)
        port_frame.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 8))
        port_controls = ttk.Frame(port_frame)
        port_controls.pack(fill=tk.X, pady=(0, 4))
        self.sysml_port_zoom = getattr(self, "sysml_port_zoom", 1.0)
        self.sysml_port_zoom_var = tk.StringVar(value=f"Ports: {round(self.sysml_port_zoom * 100)}%")
        ttk.Label(port_controls, textvariable=self.sysml_port_zoom_var).pack(side=tk.LEFT)
        self.sysml_hierarchy_port_canvas = tk.Canvas(port_frame, width=620, height=300, background="#ffffff", highlightthickness=1, highlightbackground="#c8d2dc")
        port_scroll = ttk.Scrollbar(port_frame, orient=tk.VERTICAL, command=self.sysml_hierarchy_port_canvas.yview)
        self.sysml_hierarchy_port_canvas.configure(yscrollcommand=port_scroll.set)
        self.sysml_hierarchy_port_canvas.pack(side=tk.LEFT, fill=tk.Y)
        port_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.sysml_hierarchy_port_canvas.bind("<Control-MouseWheel>", self._zoom_sysml_port_wheel, add="+")
        self.sysml_hierarchy_details = tk.Text(details_frame, height=14, wrap=tk.WORD, state=tk.DISABLED)
        self.sysml_hierarchy_details.tag_configure("find_hit", background="#fff2a8")
        self.sysml_hierarchy_details.tag_configure("find_current", background="#ffbf47")
        details_scroll = ttk.Scrollbar(details_frame, orient=tk.VERTICAL, command=self.sysml_hierarchy_details.yview)
        self.sysml_hierarchy_details.configure(yscrollcommand=details_scroll.set)
        self.sysml_hierarchy_details.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        details_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self._bind_copy_support(self.sysml_hierarchy_details)
        self.sysml_hierarchy_details.bind(
            "<Control-f>",
            lambda _event: self._open_sysml_find(self.sysml_hierarchy_details),
            add="+",
        )
        self.sysml_hierarchy_details.bind(
            "<Control-F>",
            lambda _event: self._open_sysml_find(self.sysml_hierarchy_details),
            add="+",
        )

        def close_window() -> None:
            if getattr(self, "sysml_find_target", None) is getattr(self, "sysml_hierarchy_details", None):
                self._close_sysml_find()
            self.sysml_hierarchy_window_canvas = None
            self.sysml_hierarchy_window = None
            window.destroy()

        window.protocol("WM_DELETE_WINDOW", close_window)
        self._draw_sysml_hierarchy_window()

    def _open_sysml_find(self, target: tk.Text) -> str:
        existing = getattr(self, "sysml_find_window", None)
        if existing is not None and existing.winfo_exists():
            self.sysml_find_target = target
            existing.deiconify()
            existing.lift()
            existing.focus_force()
            self.sysml_find_entry.selection_range(0, tk.END)
            return "break"

        window = tk.Toplevel(target.winfo_toplevel())
        window.title("Find in SysML Text")
        window.transient(target.winfo_toplevel())
        window.resizable(False, False)
        self.sysml_find_window = window
        self.sysml_find_target = target
        self.sysml_find_results = []
        self.sysml_find_index = -1

        frame = ttk.Frame(window, padding=8)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frame, text="Find:").pack(side=tk.LEFT, padx=(0, 4))
        self.sysml_find_var = tk.StringVar()
        self.sysml_find_entry = ttk.Entry(frame, textvariable=self.sysml_find_var, width=32)
        self.sysml_find_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.sysml_find_status_var = tk.StringVar(value="")
        self.sysml_find_entry.bind("<Return>", self._find_next_in_sysml_text)
        self.sysml_find_entry.bind("<Shift-Return>", self._find_previous_in_sysml_text)
        self.sysml_find_entry.bind("<Escape>", self._close_sysml_find)
        self.sysml_find_entry.bind("<KeyRelease>", self._refresh_sysml_find_results)
        ttk.Button(frame, text="Prev", command=self._find_previous_in_sysml_text, width=7).pack(side=tk.LEFT, padx=(6, 0))
        ttk.Button(frame, text="Next", command=self._find_next_in_sysml_text, width=7).pack(side=tk.LEFT, padx=(4, 0))
        ttk.Button(frame, text="Close", command=self._close_sysml_find, width=7).pack(side=tk.LEFT, padx=(4, 0))
        ttk.Label(frame, textvariable=self.sysml_find_status_var, width=12, anchor="e").pack(side=tk.LEFT, padx=(8, 0))
        window.protocol("WM_DELETE_WINDOW", self._close_sysml_find)
        self.sysml_find_entry.focus_set()
        self._refresh_sysml_find_results()
        return "break"

    def _clear_sysml_find_highlights(self) -> None:
        target = getattr(self, "sysml_find_target", None)
        if target is not None and target.winfo_exists():
            target.tag_remove("find_hit", "1.0", tk.END)
            target.tag_remove("find_current", "1.0", tk.END)
        self.sysml_find_results = []
        self.sysml_find_index = -1

    def _refresh_sysml_find_results(self, _event=None) -> None:
        target = getattr(self, "sysml_find_target", None)
        if target is None or not target.winfo_exists():
            return
        self._clear_sysml_find_highlights()
        query = self.sysml_find_var.get()
        if not query:
            self.sysml_find_status_var.set("")
            return
        count_var = tk.IntVar()
        start = "1.0"
        while True:
            index = target.search(query, start, nocase=True, stopindex=tk.END, count=count_var)
            if not index or count_var.get() <= 0:
                break
            end = f"{index} + {count_var.get()} chars"
            target.tag_add("find_hit", index, end)
            self.sysml_find_results.append((index, end))
            start = end
        if self.sysml_find_results:
            self.sysml_find_index = 0
            self._show_current_sysml_find_result()
        else:
            self.sysml_find_status_var.set("0 matches")

    def _show_current_sysml_find_result(self) -> None:
        if not self.sysml_find_results:
            self.sysml_find_status_var.set("0 matches")
            return
        target = self.sysml_find_target
        target.tag_remove("find_current", "1.0", tk.END)
        start, end = self.sysml_find_results[self.sysml_find_index]
        target.tag_add("find_current", start, end)
        target.see(start)
        self.sysml_find_status_var.set(f"{self.sysml_find_index + 1}/{len(self.sysml_find_results)}")

    def _find_next_in_sysml_text(self, _event=None) -> str:
        if not getattr(self, "sysml_find_results", None):
            self._refresh_sysml_find_results()
        if self.sysml_find_results:
            self.sysml_find_index = (self.sysml_find_index + 1) % len(self.sysml_find_results)
            self._show_current_sysml_find_result()
        return "break"

    def _find_previous_in_sysml_text(self, _event=None) -> str:
        if not getattr(self, "sysml_find_results", None):
            self._refresh_sysml_find_results()
        if self.sysml_find_results:
            self.sysml_find_index = (self.sysml_find_index - 1) % len(self.sysml_find_results)
            self._show_current_sysml_find_result()
        return "break"

    def _close_sysml_find(self, _event=None) -> str:
        self._clear_sysml_find_highlights()
        window = getattr(self, "sysml_find_window", None)
        self.sysml_find_window = None
        if window is not None and window.winfo_exists():
            window.destroy()
        return "break"

    def _zoom_sysml_hierarchy(self, factor: float) -> str:
        self.sysml_hierarchy_zoom = min(2.5, max(0.5, self.sysml_hierarchy_zoom * factor))
        if hasattr(self, "sysml_hierarchy_zoom_var"):
            self.sysml_hierarchy_zoom_var.set(f"Zoom: {round(self.sysml_hierarchy_zoom * 100)}%")
        self._draw_sysml_hierarchy_window()
        return "break"

    def _zoom_sysml_hierarchy_wheel(self, event) -> str:
        return self._zoom_sysml_hierarchy(1.15 if event.delta > 0 else 1 / 1.15)

    def _reset_sysml_hierarchy_zoom(self) -> None:
        self.sysml_hierarchy_zoom = 1.0
        if hasattr(self, "sysml_hierarchy_zoom_var"):
            self.sysml_hierarchy_zoom_var.set("Zoom: 100%")
        self._draw_sysml_hierarchy_window()

    def _zoom_sysml_port(self, factor: float) -> str:
        self.sysml_port_zoom = min(2.5, max(0.6, getattr(self, "sysml_port_zoom", 1.0) * factor))
        if hasattr(self, "sysml_port_zoom_var"):
            self.sysml_port_zoom_var.set(f"Ports: {round(self.sysml_port_zoom * 100)}%")
        selected = getattr(self, "sysml_hierarchy_selected_path", None)
        if selected is not None:
            self._draw_sysml_port_map(selected)
        return "break"

    def _zoom_sysml_port_wheel(self, event) -> str:
        return self._zoom_sysml_port(1.15 if event.delta > 0 else 1 / 1.15)

    def _reset_sysml_port_zoom(self) -> None:
        self.sysml_port_zoom = 1.0
        if hasattr(self, "sysml_port_zoom_var"):
            self.sysml_port_zoom_var.set("Ports: 100%")
        selected = getattr(self, "sysml_hierarchy_selected_path", None)
        if selected is not None:
            self._draw_sysml_port_map(selected)

    def _on_sysml_hierarchy_source_choice(self, _event=None) -> None:
        selected = self.sysml_hierarchy_source_var.get().strip()
        path = self.sysml_hierarchy_source_paths.get(selected)
        if path is None:
            return
        self.sysml_hierarchy_selected_path = path
        self._show_sysml_block_details(path)

    def _draw_sysml_hierarchy_window(self) -> None:
        canvas = getattr(self, "sysml_hierarchy_window_canvas", None)
        if canvas is None or not canvas.winfo_exists():
            return
        canvas.delete("all")
        top_path = SYSML_ROOT / "sysml" / "STBIOSystem.sysml"
        digital_path = SYSML_ROOT / "sysml" / "DigitalSubsystem.sysml"
        source_paths = self._concrete_sysml_block_paths()
        if hasattr(self, "sysml_hierarchy_source_choice"):
            self.sysml_hierarchy_source_paths = {
                path.relative_to(REPO_ROOT).as_posix(): path for path in source_paths
            }
            self.sysml_hierarchy_source_choice.configure(values=list(self.sysml_hierarchy_source_paths))
            selected_path = getattr(self, "sysml_hierarchy_selected_path", None)
            if selected_path in source_paths:
                self.sysml_hierarchy_source_var.set(selected_path.relative_to(REPO_ROOT).as_posix())
        if not source_paths:
            canvas.create_text(24, 24, anchor="nw", text="No generated SysML files were found.", fill="#52606d", font=("Segoe UI", 11))
            canvas.configure(scrollregion=(0, 0, 700, 120))
            return

        zoom = getattr(self, "sysml_hierarchy_zoom", 1.0)
        node_width, node_height = 220 * zoom, 48 * zoom
        nodes = []
        if top_path.exists():
            nodes.append(("STBIOSystem", top_path, 440 * zoom, 30 * zoom, "root"))
        if digital_path.exists():
            nodes.append(("DigitalSubsystem", digital_path, 150 * zoom, 150 * zoom, "subsystem"))
        analog_path = SYSML_BLOCKS_ROOT / "adc.sysml"
        if analog_path.exists():
            nodes.append(("AnalogSubsystem", analog_path, 730 * zoom, 150 * zoom, "subsystem"))
        block_paths = list(source_paths)
        hierarchy_ports = {}
        for path in block_paths + ([analog_path] if analog_path.exists() else []):
            source = path.read_text(encoding="utf-8", errors="replace")
            hierarchy_ports[path] = re.findall(
                r"\b(?:(in|out|inout)\s+)?port\s+([A-Za-z_][A-Za-z0-9_]*)\s*:",
                source,
            )
        block_width = 280 * zoom
        block_heights = {
            path: max(74 * zoom, (42 + max(0, len(hierarchy_ports.get(path, []))) * 14) * zoom)
            for path in block_paths
        }
        block_columns = 3
        block_x_step = 315 * zoom
        block_y = 285 * zoom
        for row_start in range(0, len(block_paths), block_columns):
            row_paths = block_paths[row_start:row_start + block_columns]
            for column, path in enumerate(row_paths):
                nodes.append((path.stem, path, 25 * zoom + column * block_x_step, block_y, "block"))
            block_y += max(block_heights[path] for path in row_paths) + 34 * zoom
        positions = {path: (x, y) for _label, path, x, y, _kind in nodes}
        node_sizes = {
            path: (block_width, block_heights.get(path, 58 * zoom))
            for _label, path, _x, _y, _kind in nodes
        }
        node_sizes.update({path: (220 * zoom, 48 * zoom) for _label, path, _x, _y, kind in nodes if kind != "block"})
        root_path = top_path if top_path.exists() else source_paths[0]
        for label, path, x, y, kind in nodes:
            if path == root_path:
                continue
            parent_path = digital_path if kind == "block" and digital_path.exists() else root_path
            if parent_path in positions:
                px, py = positions[parent_path]
                parent_width, parent_height = node_sizes.get(parent_path, (node_width, node_height))
                child_width, _child_height = node_sizes.get(path, (node_width, node_height))
                canvas.create_line(px + parent_width / 2, py + parent_height, x + child_width / 2, y, fill="#78909c", width=2, arrow=tk.LAST)
        for label, path, x, y, kind in nodes:
            fill = {"root": "#dceef7", "subsystem": "#fff0d6", "block": "#ffffff"}[kind]
            outline = {"root": "#236477", "subsystem": "#a66a00", "block": "#8493a1"}[kind]
            current_width, current_height = node_sizes.get(path, (node_width, node_height))
            canvas.create_rectangle(x, y, x + current_width, y + current_height, fill=fill, outline=outline, width=2, tags=("sysml_hierarchy_node", path.as_posix()))
            canvas.create_text(x + current_width / 2, y + 16 * zoom, text=label, fill="#263238", font=("Segoe UI", max(7, round(10 * zoom)), "bold" if kind != "block" else "normal"), tags=("sysml_hierarchy_node", path.as_posix()))
            if kind == "block":
                canvas.create_line(x + 8 * zoom, y + 30 * zoom, x + current_width - 8 * zoom, y + 30 * zoom, fill="#c8d2dc", width=1)
                ports = hierarchy_ports.get(path, [])
                input_ports = [name for direction, name in ports if direction in {"in", "inout"} or not direction]
                output_ports = [name for direction, name in ports if direction in {"out", "inout"}]
                rows = max(len(input_ports), len(output_ports))
                for index in range(rows):
                    port_y = y + (44 + index * 14) * zoom
                    if index < len(input_ports):
                        canvas.create_line(x + 2 * zoom, port_y, x + 12 * zoom, port_y, fill="#1976d2", width=max(1, round(zoom)))
                        canvas.create_text(x + 16 * zoom, port_y, text=input_ports[index], anchor=tk.W, width=max(70, int(current_width / 2 - 28 * zoom)), fill="#245174", font=("Segoe UI", max(6, round(7 * zoom))), tags=("sysml_hierarchy_node", path.as_posix()))
                    if index < len(output_ports):
                        canvas.create_line(x + current_width - 12 * zoom, port_y, x + current_width - 2 * zoom, port_y, fill="#c75b12", width=max(1, round(zoom)))
                        canvas.create_text(x + current_width - 16 * zoom, port_y, text=output_ports[index], anchor=tk.E, width=max(70, int(current_width / 2 - 28 * zoom)), fill="#7a3b12", font=("Segoe UI", max(6, round(7 * zoom))), tags=("sysml_hierarchy_node", path.as_posix()))
        canvas.configure(
            scrollregion=(
                0,
                0,
                max(980 * zoom, (25 + min(block_columns, len(block_paths)) * 315) * zoom),
                max(390 * zoom, block_y + 40 * zoom),
            )
        )

        def open_node(event) -> None:
            current = canvas.find_withtag("current")
            if not current:
                return
            path_text = next((tag for tag in canvas.gettags(current[0]) if tag.endswith(".sysml")), "")
            if path_text:
                self.work_tabs.select(self.sysml_architecture_tab)
                selected_path = Path(path_text)
                self._select_sysml_file(selected_path)
                self._show_sysml_block_details(selected_path)

        canvas.tag_bind("sysml_hierarchy_node", "<Button-1>", open_node)
        canvas.tag_bind("sysml_hierarchy_node", "<Double-1>", open_node)

    def _show_sysml_block_details(self, path: Path) -> None:
        details = getattr(self, "sysml_hierarchy_details", None)
        if details is None or not details.winfo_exists():
            return
        if not path.exists():
            text = f"File not found: {path}"
        else:
            source = path.read_text(encoding="utf-8", errors="replace")
            text = f"File: {path.relative_to(REPO_ROOT).as_posix()}\n"
            text += "Full selected SysML source:\n\n"
            text += source
        self.sysml_hierarchy_selected_path = path
        if hasattr(self, "sysml_hierarchy_source_var"):
            try:
                self.sysml_hierarchy_source_var.set(path.relative_to(REPO_ROOT).as_posix())
            except ValueError:
                self.sysml_hierarchy_source_var.set(path.name)
        self._draw_sysml_port_map(path)
        details.configure(state=tk.NORMAL)
        details.delete("1.0", tk.END)
        details.insert("1.0", text)
        details.configure(state=tk.DISABLED)

    @staticmethod
    def _sysml_attribute_value(source: str, attribute_name: str) -> str:
        match = re.search(
            rf'\battribute\s+{re.escape(attribute_name)}\b\s*(?::\s*[A-Za-z_][A-Za-z0-9_]*)?\s*=\s*(?:"((?:\\.|[^"\\])*)"|([^;\s]+))',
            source,
        )
        if not match:
            return ""
        if match.group(2) is not None:
            return match.group(2)
        try:
            return json.loads(f'"{match.group(1)}"')
        except (TypeError, ValueError, json.JSONDecodeError):
            return match.group(1).replace('\\"', '"').replace("\\\\", "\\")

    def _draw_sysml_port_map(self, path: Path) -> None:
        canvas = getattr(self, "sysml_hierarchy_port_canvas", None)
        if canvas is None or not canvas.winfo_exists():
            return
        canvas.delete("all")
        if not path.exists():
            canvas.create_text(140, 75, text="No block selected", anchor=tk.CENTER, fill="#52606d")
            return
        source = path.read_text(encoding="utf-8", errors="replace")
        definition = next(iter(re.findall(r"\bpart\s+def\s+([A-Za-z_][A-Za-z0-9_]*)\s*\{", source)), path.stem)
        ports = re.findall(
            r"\b(?:(in|out|inout)\s+)?port\s+([A-Za-z_][A-Za-z0-9_]*)\s*:\s*([A-Za-z_][A-Za-z0-9_]*)\s*\{",
            source,
        )
        input_ports = [port for port in ports if port[0] in {"in", "inout"} or not port[0]]
        output_ports = [port for port in ports if port[0] in {"out", "inout"}]
        zoom = getattr(self, "sysml_port_zoom", 1.0)
        port_rows = max(len(input_ports), len(output_ports))
        block_x1, block_y1 = 270 * zoom, 48 * zoom
        block_x2, block_y2 = 410 * zoom, (110 + port_rows * 22) * zoom
        label_font = ("Segoe UI", max(8, round(10 * zoom)))
        label_width = max(180, round(245 * zoom))
        row_step = 22 * zoom
        canvas.create_rectangle(block_x1, block_y1, block_x2, block_y2, fill="#eaf0f6", outline="#236477", width=2)
        canvas.create_text((block_x1 + block_x2) / 2, block_y1 + (block_y2 - block_y1) / 2, text=definition, fill="#263238", font=("Segoe UI", max(9, round(11 * zoom)), "bold"), width=max(120, round(150 * zoom)))
        for index, (_direction, name, _port_type) in enumerate(input_ports):
            y = 68 * zoom + index * row_step
            canvas.create_text(block_x1 - 14 * zoom, y, text=name, anchor=tk.E, width=label_width, fill="#263238", font=label_font)
            canvas.create_line(block_x1 - 10 * zoom, y, block_x1, y, fill="#1976d2", width=max(1, round(1.5 * zoom)))
            canvas.create_oval(block_x1 - 13 * zoom, y - 4 * zoom, block_x1 - 5 * zoom, y + 4 * zoom, fill="#1976d2", outline="")
        for index, (_direction, name, _port_type) in enumerate(output_ports):
            y = 68 * zoom + index * row_step
            canvas.create_line(block_x2, y, block_x2 + 10 * zoom, y, fill="#c75b12", width=max(1, round(1.5 * zoom)))
            canvas.create_oval(block_x2 + 5 * zoom, y - 4 * zoom, block_x2 + 13 * zoom, y + 4 * zoom, fill="#c75b12", outline="")
            canvas.create_text(block_x2 + 14 * zoom, y, text=name, anchor=tk.W, width=label_width, fill="#263238", font=label_font)
        canvas.configure(scrollregion=(0, 0, max(620, round((block_x2 + label_width + 24 * zoom))), max(300, round(block_y2 + 40 * zoom))))

    def _reset_sysml_graph_zoom(self) -> None:
        if not hasattr(self, "sysml_graph_canvas"):
            return
        self.sysml_graph_zoom = 1.0
        self._refresh_sysml_architecture()
        self.sysml_graph_zoom_var.set("Zoom: 100%")

    def _draw_sysml_architecture(self) -> None:
        if not hasattr(self, "sysml_graph_canvas"):
            return
        canvas = self.sysml_graph_canvas
        canvas.delete("all")
        top_path = SYSML_ROOT / "sysml" / "STBIOSystem.sysml"
        digital_path = SYSML_ROOT / "sysml" / "DigitalSubsystem.sysml"
        if not top_path.exists() or not digital_path.exists():
            canvas.create_text(20, 20, anchor="nw", text="Generate the SysML architecture to display the graph.", fill="#52606d", font=("Segoe UI", 11))
            canvas.configure(scrollregion=(0, 0, 600, 100))
            return
        source_paths = [top_path, digital_path, SYSML_ROOT / "sysml" / "shared_definitions.sysml"]
        source_paths.extend(self._concrete_sysml_block_paths())
        source_text = {
            path: path.read_text(encoding="utf-8", errors="replace")
            for path in source_paths
            if path.exists()
        }
        top_text = source_text[top_path]
        digital_text = source_text[digital_path]
        self.sysml_connection_details = {}

        node_width, node_height = 170, 42
        positions = {"STBIOSystem": (520, 18), "topAnalog": (150, 92), "topDigital": (890, 92)}
        block_def_paths = {
            block_type: path
            for path in source_paths
            for block_type in re.findall(
                r"\bpart\s+def\s+([A-Za-z_][A-Za-z0-9_]*)\b",
                path.read_text(encoding="utf-8", errors="replace"),
            )
        }
        blocks = [
            (usage, block_type)
            for usage, block_type in re.findall(
                r"\bpart\s+(?!def\b)(\w+)\s*:\s*([A-Za-z_][A-Za-z0-9_]*)\s*;",
                digital_text,
            )
            if block_type in block_def_paths and self._sysml_block_file_is_concrete(block_def_paths[block_type])
        ]
        if not blocks:
            canvas.create_text(
                20,
                20,
                anchor="nw",
                text="No SysML part declarations were found in DigitalSubsystem.sysml.",
                fill="#52606d",
                font=("Segoe UI", 11),
            )
            canvas.configure(scrollregion=(0, 0, 600, 100))
            return
        for index, (usage, block_type) in enumerate(blocks):
            positions[f"topDigital.{usage}"] = (650 + (index % 4) * 190, 175 + (index // 4) * 72)
        analog_match = re.search(r"part def AnalogSubsystem\s*\{(.*?)\n\s*\}", top_text, flags=re.S)
        if analog_match:
            for usage, block_type in re.findall(r"\s+part\s+(\w+)\s*:\s*(\w+)\s*;", analog_match.group(1)):
                positions[f"topAnalog.analogSubsystem.{usage}"] = (70, 175)

        def node_label(path: str) -> str:
            return path.rsplit(".", 1)[-1]

        connection_pattern = re.compile(
            r"connection\s+([A-Za-z0-9_]+).*?connect\s+([A-Za-z0-9_.]+)\s+to\s+([A-Za-z0-9_.]+);",
            flags=re.S,
        )
        metadata_pattern = re.compile(
            r"doc\s+[A-Za-z0-9_]+\s+/\*\s*signal=([^|]*)\|\s*trigger=([^|]*)\|\s*notes=([^|]*)\|\s*sourceReqIds=([^|]*)[^*]*\*/",
            flags=re.S,
        )
        metadata_matches = list(metadata_pattern.finditer(top_text))
        for connection_match in connection_pattern.finditer(top_text):
            connection_id, source, target = connection_match.groups()
            metadata = next(
                (match for match in metadata_matches if match.start() >= connection_match.end()),
                None,
            )
            signal, trigger, notes, source_ids = metadata.groups() if metadata else ("", "", "", "")
            if source in positions and target in positions:
                sx, sy = positions[source]
                tx, ty = positions[target]
                self.sysml_connection_details[connection_id] = {
                    "source": source,
                    "target": target,
                    "signal": signal.strip(),
                    "trigger": trigger.strip(),
                    "notes": notes.strip(),
                    "source_ids": source_ids.strip() or "None",
                }
                canvas.create_line(
                    sx + node_width / 2,
                    sy + node_height / 2,
                    tx + node_width / 2,
                    ty + node_height / 2,
                    fill="#1976d2" if connection_id == self.sysml_selected_connection else "#78909c",
                    arrow=tk.LAST,
                    width=3 if connection_id == self.sysml_selected_connection else 1.5,
                    tags=("sysml_connection", connection_id),
                )

        for path, (x, y) in positions.items():
            is_selected = path == self.sysml_selected_block
            fill = "#cfe8ff" if is_selected else "#dceef7" if path == "STBIOSystem" else "#e9f2e8" if path == "analogSubsystem" else "#fff0d6" if path == "digitalSubsystem" else "#ffffff"
            outline = "#1976d2" if is_selected else "#236477" if path == "STBIOSystem" else "#527a50" if path == "analogSubsystem" else "#a66a00" if path == "digitalSubsystem" else "#8493a1"
            canvas.create_rectangle(x, y, x + node_width, y + node_height, fill=fill, outline=outline, width=3 if is_selected else 2, tags=("sysml_node", path))
            canvas.create_text(x + node_width / 2, y + node_height / 2, text=node_label(path), fill="#0b4f8a" if is_selected else "#263238", font=("Segoe UI", 10, "bold" if is_selected or path in {"STBIOSystem", "analogSubsystem", "digitalSubsystem"} else "normal"), tags=("sysml_node", path))
        canvas.configure(scrollregion=(0, 0, 1450, max(260, 250 + ((len(blocks) + 3) // 4) * 72)))

        def open_node(event) -> None:
            current = canvas.find_withtag("current")
            if not current:
                return
            tags = canvas.gettags(current[0])
            path = next((tag for tag in tags if "." in tag and tag.startswith("digitalSubsystem.")), None)
            if path is None:
                return
            usage = path.rsplit(".", 1)[-1]
            block_type = next((candidate_type for candidate_usage, candidate_type in blocks if candidate_usage == usage), None)
            if block_type is None:
                return
            self.sysml_selected_block = path
            self._draw_sysml_architecture()
            block_file = next(
                (candidate for candidate in self._concrete_sysml_block_paths()
                 if re.search(
                     rf"\bpart\s+def\s+{re.escape(block_type)}\b",
                     candidate.read_text(encoding="utf-8", errors="replace"),
                 )),
                None,
            )
            if block_file is not None:
                self.work_tabs.select(self.sysml_architecture_tab)
                self._select_sysml_file(block_file)

        canvas.tag_bind("sysml_node", "<Button-1>", open_node)
        canvas.tag_bind("sysml_node", "<Double-1>", open_node)
        canvas.tag_bind("sysml_connection", "<Button-1>", self._show_sysml_connection_popup)

    def _show_sysml_connection_popup(self, event) -> str:
        canvas = self.sysml_graph_canvas
        current = canvas.find_withtag("current")
        if not current:
            return "break"
        connection_id = next((tag for tag in canvas.gettags(current[0]) if tag.startswith("connection_")), None)
        details = self.sysml_connection_details.get(connection_id)
        if not details:
            return "break"
        self.sysml_selected_connection = connection_id
        self._draw_sysml_architecture()
        menu = tk.Menu(self.root, tearoff=False)
        menu.add_command(label=f"{connection_id}: {details['source']} -> {details['target']}", state=tk.DISABLED)
        menu.add_separator()
        menu.add_command(label=f"Signal: {details['signal'] or 'None'}", state=tk.DISABLED)
        menu.add_command(label=f"Trigger: {details['trigger'] or 'None'}", state=tk.DISABLED)
        menu.add_command(label=f"Notes: {details['notes'] or 'None'}", state=tk.DISABLED)
        menu.add_command(label=f"Source IDs: {details['source_ids']}", state=tk.DISABLED)
        return self._show_popup_menu(menu, event)

    def _select_sysml_file(self, path: Path) -> None:
        try:
            self.sysml_file_choice_var.set(path.relative_to(REPO_ROOT).as_posix())
        except ValueError:
            self.sysml_file_choice_var.set(path.name)

        def find_item(parent: str):
            for item in self.sysml_tree.get_children(parent):
                values = self.sysml_tree.item(item, "values")
                if values and Path(values[0]) == path:
                    return item
                found = find_item(item)
                if found:
                    return found
            return None

        item = find_item("")
        if item:
            parent = self.sysml_tree.parent(item)
            while parent:
                self.sysml_tree.item(parent, open=True)
                parent = self.sysml_tree.parent(parent)
            self.sysml_tree.selection_set(item)
            self.sysml_tree.see(item)
            self._on_sysml_select()
            return
        if path.is_file():
            self.work_tabs.select(self.sysml_architecture_tab)
            try:
                content = path.read_text(encoding="utf-8", errors="replace")
                label = path.relative_to(REPO_ROOT).as_posix()
            except (OSError, ValueError) as exc:
                content = f"Unable to read SysML file: {exc}"
                label = path.name
            self.sysml_file_var.set(label)
            self.sysml_viewer.configure(state=tk.NORMAL)
            self.sysml_viewer.delete("1.0", tk.END)
            self.sysml_viewer.insert("1.0", content)
            self.sysml_viewer.configure(state=tk.NORMAL)

    def _on_sysml_file_choice(self, _event=None) -> None:
        selected = self.sysml_file_choice_var.get().strip()
        path = self.sysml_file_choice_paths.get(selected)
        if path is not None:
            self._select_sysml_file(path)

    def _on_sysml_hierarchy_select(self, _event=None) -> None:
        selection = self.sysml_hierarchy_tree.selection()
        if not selection:
            return
        values = self.sysml_hierarchy_tree.item(selection[0], "values")
        if values and Path(values[0]).is_file():
            self._select_sysml_file(Path(values[0]))

    @staticmethod
    def _sysml_concrete_block_names() -> set[str]:
        inventory_path = REPO_ROOT / "artifacts/stage2_mirco_arc/block_inventory.csv"
        if not inventory_path.exists():
            return set()
        try:
            with inventory_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            interface_catalog_path = REPO_ROOT / "artifacts/stage2_mirco_arc/interface_catalog.csv"
            interface_owners = set()
            if interface_catalog_path.exists():
                with interface_catalog_path.open("r", encoding="utf-8-sig", newline="") as handle:
                    interface_owners = {
                        (row.get("Owner") or "").strip().casefold()
                        for row in csv.DictReader(handle)
                        if (row.get("Owner") or "").strip()
                        and re.search(r"\b(?:interface|port|pin|endpoint)\b", (row.get("Owner") or ""), re.IGNORECASE)
                    }
            return {
                name
                for row in rows
                for name in [(row.get("Block") or "").strip()]
                if name
                and name.casefold() != "unassigned"
                and name.casefold() not in interface_owners
                and "described by stage 2 interaction evidence" not in (row.get("Function") or "").casefold()
                and (row.get("Entity kind") or "concrete_block").strip().lower() == "concrete_block"
            }
        except OSError:
            return set()

    @classmethod
    def _sysml_block_file_is_concrete(cls, path: Path) -> bool:
        concrete_names = cls._sysml_concrete_block_names()
        if not concrete_names or path.parent.resolve() != SYSML_BLOCKS_ROOT.resolve():
            return False
        try:
            source = path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            return False
        display_match = re.search(r'attribute\s+displayName\s*:\s*String\s*=\s*"([^"]+)"', source)
        return bool(display_match and display_match.group(1).strip() in concrete_names)

    @classmethod
    def _concrete_sysml_block_paths(cls) -> list[Path]:
        if not SYSML_BLOCKS_ROOT.is_dir():
            return []
        created_ipos_names = {
            re.sub(r"[^a-z0-9]", "", name.casefold())
            for stage in ("Digital IPOS", "Analog IPOS")
            for name in created_ipos_block_directories(REPO_ROOT, stage)
        }
        return [
            path
            for path in sorted(SYSML_BLOCKS_ROOT.glob("*.sysml"), key=lambda item: item.name.lower())
            if cls._sysml_block_file_is_concrete(path)
            and re.sub(r"[^a-z0-9]", "", path.stem.casefold()) in created_ipos_names
        ]

    def _refresh_sysml_hierarchy_window(self) -> None:
        self._refresh_sysml_architecture()
        self._draw_sysml_hierarchy_window()

    def _populate_sysml_hierarchy(self, source_paths: list[Path]) -> None:
        for item in self.sysml_hierarchy_tree.get_children():
            self.sysml_hierarchy_tree.delete(item)
        path_by_name = {path.name: path for path in source_paths}
        top_path = path_by_name.get("STBIOSystem.sysml") or (source_paths[0] if source_paths else None)
        if top_path is None:
            return
        root = self.sysml_hierarchy_tree.insert("", tk.END, text=top_path.stem, values=(top_path.as_posix(),), open=True)
        top_text = top_path.read_text(encoding="utf-8", errors="replace")
        subsystem_nodes = {}
        for usage, block_type in re.findall(r"\bpart\s+(\w+)\s*:\s*([A-Za-z_][A-Za-z0-9_]*)\s*;", top_text):
            match = next((path for path in source_paths if re.search(rf"\bpart\s+def\s+{re.escape(block_type)}\b", path.read_text(encoding="utf-8", errors="replace"))), None)
            parent = root
            node = self.sysml_hierarchy_tree.insert(parent, tk.END, text=usage, values=((match or top_path).as_posix(),), open=True)
            subsystem_nodes[usage] = node
        for path in source_paths:
            if path == top_path:
                continue
            if path.name in {"shared_definitions.sysml", "DigitalSubsystem.sysml"}:
                self.sysml_hierarchy_tree.insert(root, tk.END, text=path.stem, values=(path.as_posix(),))
        digital_node = subsystem_nodes.get("topDigital")
        if digital_node is not None:
            digital_path = path_by_name.get("DigitalSubsystem.sysml")
            if digital_path is not None:
                digital_text = digital_path.read_text(encoding="utf-8", errors="replace")
                concrete_names = self._sysml_concrete_block_names()
                for usage, block_type in re.findall(r"\bpart\s+(?!def\b)(\w+)\s*:\s*([A-Za-z_][A-Za-z0-9_]*)\s*;", digital_text):
                    block_path = next((path for path in source_paths if re.search(rf"\bpart\s+def\s+{re.escape(block_type)}\b", path.read_text(encoding="utf-8", errors="replace"))), digital_path)
                    if concrete_names:
                        display_match = re.search(r'attribute\s+displayName\s*:\s*String\s*=\s*"([^"]+)"', block_path.read_text(encoding="utf-8", errors="replace"))
                        if not display_match or display_match.group(1).strip() not in concrete_names:
                            continue
                    self.sysml_hierarchy_tree.insert(digital_node, tk.END, text=usage, values=(block_path.as_posix(),))
        analog_node = subsystem_nodes.get("topAnalog")
        if analog_node is not None:
            analog_match = re.search(r"\bpart\s+def\s+AnalogSubsystem\s*\{(.*?)\n\s*\}", top_text, flags=re.S)
            if analog_match:
                for usage, block_type in re.findall(r"\bpart\s+(?!def\b)(\w+)\s*:\s*([A-Za-z_][A-Za-z0-9_]*)\s*;", analog_match.group(1)):
                    block_path = next(
                        (path for path in source_paths if re.search(rf"\bpart\s+def\s+{re.escape(block_type)}\b", path.read_text(encoding="utf-8", errors="replace"))),
                        top_path,
                    )
                    self.sysml_hierarchy_tree.insert(analog_node, tk.END, text=usage, values=(block_path.as_posix(),))

    def _refresh_sysml_architecture(self) -> None:
        if not hasattr(self, "sysml_tree"):
            return
        for item in self.sysml_tree.get_children():
            self.sysml_tree.delete(item)
        top_path = SYSML_ROOT / "sysml" / "STBIOSystem.sysml"
        model_paths = [
            path
            for path in [top_path, SYSML_ROOT / "sysml" / "DigitalSubsystem.sysml", SYSML_ROOT / "sysml" / "shared_definitions.sysml"]
            if path.exists()
        ]
        concrete_paths = self._concrete_sysml_block_paths()
        self.sysml_file_choice_paths = {
            path.relative_to(REPO_ROOT).as_posix(): path for path in concrete_paths
        }
        self.sysml_file_choice.configure(values=list(self.sysml_file_choice_paths))
        self._populate_sysml_hierarchy(model_paths + concrete_paths)
        concrete_node = self.sysml_tree.insert("", tk.END, text="Concrete SysML blocks", open=True)
        for path in concrete_paths:
            self.sysml_tree.insert(concrete_node, tk.END, text=path.stem, values=(path.as_posix(),))
        if concrete_paths:
            self.sysml_tree.selection_set(self.sysml_tree.get_children(concrete_node)[0])
            self.sysml_tree.focus(self.sysml_tree.get_children(concrete_node)[0])
        self._on_sysml_select()
        self._draw_sysml_architecture()
        if getattr(self, "sysml_hierarchy_window", None) is not None:
            self._draw_sysml_hierarchy_window()

    def _on_sysml_select(self, _event=None) -> None:
        selection = self.sysml_tree.selection()
        if not selection:
            return
        values = self.sysml_tree.item(selection[0], "values")
        if not values:
            return
        path = Path(values[0])
        if not path.is_file():
            return
        try:
            content = path.read_text(encoding="utf-8", errors="replace")
            label = path.relative_to(REPO_ROOT).as_posix()
        except (OSError, ValueError) as exc:
            content = f"Unable to read SysML file: {exc}"
            label = path.name
        self.sysml_file_var.set(label)
        self.sysml_viewer.configure(state=tk.NORMAL)
        self.sysml_viewer.delete("1.0", tk.END)
        self.sysml_viewer.insert("1.0", content)
        self.sysml_viewer.configure(state=tk.NORMAL)

    def _build_system_traceability_tab(self) -> None:
        self.traceability_pane = ttk.Frame(self.system_traceability_tab, padding=18)
        self.traceability_pane.pack(fill=tk.BOTH, expand=True)
        ttk.Label(
            self.traceability_pane,
            text="System Traceability",
            font=("Segoe UI", 16, "bold"),
        ).pack(anchor=tk.W, pady=(40, 8))
        ttk.Label(
            self.traceability_pane,
            text="The complete vertical dependency graph and hierarchy open in a dedicated window.",
            wraplength=720,
        ).pack(anchor=tk.W, pady=(0, 14))
        open_hierarchy_btn = ttk.Button(
            self.traceability_pane,
            text="Open Traceability Hierarchy",
            command=self._open_system_traceability_window,
            width=30,
        )
        open_hierarchy_btn.pack(anchor=tk.W)
        self._bind_button_help(open_hierarchy_btn, "Open the dedicated graphical hierarchy with dependency links and coverage metrics.")
        controls = ttk.Frame(self.traceability_pane)
        controls.pack(anchor=tk.W, pady=(12, 0))
        open_report_btn = ttk.Button(controls, text="Open Report", command=self._open_coverage_report, width=12)
        open_report_btn.pack(side=tk.LEFT)
        self._bind_button_help(open_report_btn, "Open the generated hierarchy coverage report in the Artifacts tab.")
        refresh_coverage_btn = ttk.Button(controls, text="Refresh Artifacts", command=self._crosscheck_and_refresh_system_traceability, width=16)
        refresh_coverage_btn.pack(side=tk.LEFT, padx=(6, 0))
        self._bind_button_help(refresh_coverage_btn, "Regenerate the source-baselined traceability artifacts.")
        self.coverage_status_var = tk.StringVar(value="Loading coverage report...")
        ttk.Label(self.traceability_pane, textvariable=self.coverage_status_var, wraplength=720).pack(anchor=tk.W, pady=(16, 0))
        self._refresh_system_traceability()

    def _read_coverage_report(self) -> dict[str, str]:
        values = {
            "source_document": "Unknown source document",
            "source_count": "-",
            "SRS": "Unavailable",
            "ARS": "Unavailable",
            "DRS": "Unavailable",
            "Source -> SRS": "Unavailable",
            "SRS -> ARS": "Unavailable",
            "SRS -> DRS": "Unavailable",
            "DRS -> Digital IPOS": "Unavailable",
            "ARS -> Analog IPOS": "Unavailable",
            "Source specification -> Digital IPOS": "Unavailable",
            "Source specification -> Analog IPOS": "Unavailable",
            "Source specification -> SRS -> DRS -> Digital IPOS": "Unavailable",
            "Source specification -> SRS -> ARS -> Analog IPOS": "Unavailable",
            "allocation_summary": "Unavailable",
            "allocation_findings": "",
            "error": "",
        }
        if not COVERAGE_REPORT_PATH.exists():
            values["error"] = "Run Stage 5 to generate the hierarchy coverage report."
            return values

        try:
            report_text = COVERAGE_REPORT_PATH.read_text(encoding="utf-8")
        except OSError as exc:
            values["error"] = f"Unable to read coverage report: {exc}"
            return values

        source_match = re.search(r"^- Source requirements:\s*(\d+)\s*$", report_text, flags=re.M)
        if source_match:
            values["source_count"] = source_match.group(1)
        source_document_match = re.search(r"^- Source document:\s*`?([^`\n]+)`?", report_text, flags=re.M)
        if source_document_match:
            values["source_document"] = source_document_match.group(1).strip()
        for label in (
            "SRS", "ARS", "DRS", "Digital IPOS", "Analog IPOS",
            "Source -> SRS", "SRS -> ARS", "SRS -> DRS",
            "DRS -> Digital IPOS", "ARS -> Analog IPOS",
            "Source specification -> Digital IPOS", "Source specification -> Analog IPOS",
            "Source specification -> SRS -> DRS -> Digital IPOS",
            "Source specification -> SRS -> ARS -> Analog IPOS",
        ):
            match = re.search(rf"^\|\s*{re.escape(label)}\s*\|\s*(.*?)\s*\|\s*$", report_text, flags=re.M)
            if match:
                values[label] = match.group(1)
        try:
            edge_statistics = approved_generated_document_dependency_graph(REPO_ROOT).get("edge_statistics", [])
            edge_lookup = {
                (str(item.get("upstream_node")), str(item.get("derived_node"))): item
                for item in edge_statistics
                if isinstance(item, dict)
            }

            def edge_value(upstream: str, derived: str) -> str:
                statistic = edge_lookup.get((upstream, derived))
                if not statistic:
                    return "Unavailable"
                covered = len(statistic.get("covered_unique_upstream_req_ids") or [])
                total = len(statistic.get("total_unique_upstream_req_ids") or [])
                percentage = float(statistic.get("percentage") or 0.0)
                return f"{covered} / {total} ({percentage:.1f}%)"

            values["SRS -> ARS"] = edge_value("SRS", "ARS")
            values["SRS -> DRS"] = edge_value("SRS", "DRS")
            values["DRS -> Digital IPOS"] = edge_value("DRS", "Digital IPOS")
            values["ARS -> Analog IPOS"] = edge_value("ARS", "Analog IPOS")
            values["rm_edge_statistics"] = edge_statistics
        except Exception as exc:
            values["error"] = f"Unable to read central RM edge statistics: {exc}"
        crosscheck_path = REPO_ROOT / "artifacts/traceability_reports/allocation_crosscheck_report.json"
        if crosscheck_path.exists():
            try:
                crosscheck = json.loads(crosscheck_path.read_text(encoding="utf-8"))
                values["allocation_summary"] = (
                    f"{crosscheck.get('row_count', 0)} ledger rows; "
                    f"statuses={crosscheck.get('by_coverage_status', {})}; "
                    f"lineage={crosscheck.get('by_lineage_mode', {})}"
                )
                values["allocation_findings"] = "; ".join(crosscheck.get("findings", []))
            except (OSError, json.JSONDecodeError):
                values["allocation_summary"] = "Allocation crosscheck unavailable"
        return values

    def _refresh_system_traceability(self) -> None:
        self.coverage_values = self._read_coverage_report()
        error_text = self.coverage_values["error"]
        self.coverage_status_var.set(
            error_text or (
                "Coverage is calculated from the centralized allocation ledger. "
                + self.coverage_values.get("allocation_summary", "")
                + (" Findings: " + self.coverage_values.get("allocation_findings", "") if self.coverage_values.get("allocation_findings") else "")
            )
        )
        self._draw_system_traceability()

    def _crosscheck_and_refresh_system_traceability(self) -> None:
        """Regenerate the hierarchy report through the DRS crosscheck before rendering it."""
        selector = self._snapshot_selector_args()
        if selector is None:
            return
        self._run_command_async(
            [
                PYTHON_EXE,
                str(SCRIPTS_DIR / "run_drs_crosscheck_agent.py"),
                *selector,
            ],
            "Refresh System Traceability",
            on_success=self._refresh_work_area_and_traceability,
        )

    def _zoom_system_traceability(self, event) -> str:
        factor = 1.1 if event.delta > 0 else 1 / 1.1
        self.coverage_zoom = min(1.8, max(0.6, self.coverage_zoom * factor))
        self.coverage_zoom_var.set(f"Zoom: {round(self.coverage_zoom * 100)}%")
        self._draw_system_traceability()
        return "break"

    def _open_coverage_report(self) -> None:
        if not COVERAGE_REPORT_PATH.exists():
            messagebox.showinfo("System Traceability", "Run Stage 5 to generate the hierarchy coverage report.")
            return
        self.work_tabs.select(self.work_browser_tab)
        self._on_work_area_select_path(COVERAGE_REPORT_PATH)

    @staticmethod
    def _coverage_link_color(value: str) -> str:
        if "not applicable" in value.lower() or "unavailable" in value.lower():
            return "#7a8793"
        match = re.search(r"\(([0-9.]+)%\)", value)
        if not match:
            return "#7a8793"
        percentage = float(match.group(1))
        if percentage >= 95.0:
            return "#2c8a2c"
        if percentage >= 80.0:
            return "#c27c0e"
        return "#b3261e"

    def _draw_system_traceability(self) -> None:
        if not hasattr(self, "coverage_canvas"):
            return
        canvas = self.coverage_canvas
        current_x = canvas.xview()[0] if canvas.xview() else 0.0
        current_y = canvas.yview()[0] if canvas.yview() else 0.0
        canvas.delete("all")
        values = getattr(self, "coverage_values", self._read_coverage_report())
        zoom = getattr(self, "coverage_zoom", 1.0)
        width = max(int(max(canvas.winfo_width(), 920) * zoom), 600)
        height = max(int(max(canvas.winfo_height(), 420) * zoom), 280)
        node_width = min(142, max(112, (width - 100) // 5))
        node_height = max(56, int(78 * zoom))
        source_x = 16
        srs_x = int(width * 0.28)
        child_x = width - node_width - 16
        srs_y = int(height * 0.42)
        ars_y = 18
        drs_y = height - node_height - 18
        ipos_x = int(width * 0.57)

        nodes = {
            "Source Spec": (source_x, srs_y),
            "SRS": (srs_x, srs_y),
            "ARS": (child_x, ars_y),
            "DRS": (child_x, drs_y),
            "Analog IPOS": (ipos_x, ars_y),
            "Digital IPOS": (ipos_x, drs_y),
            "SysML": (srs_x, 18),
        }
        generated_paths = {
            "Analog IPOS": REPO_ROOT / "artifacts/stage7_analog_ipos",
            "Digital IPOS": REPO_ROOT / "artifacts/stage6_digital_ipos",
            "SysML": SYSML_ROOT / "sysml" / "STBIOSystem.sysml",
        }
        def node_metric(document: str) -> str:
            value = values[document]
            if "not applicable" in value.lower():
                return "Not applicable\nfor this source spec"
            return f"Source Spec coverage\n{value}"

        def link_metric(link: str) -> str:
            value = values[link]
            return "N/A" if "not applicable" in value.lower() else value

        def generated_metric(document: str) -> str:
            path = generated_paths[document]
            exists = path.exists() if path.suffix else path.is_dir()
            return "Generated" if exists else "Not generated"

        labels = {
            "Source Spec": f"Source Spec\n{values['source_document']}\n{values['source_count']} source requirements",
            "SRS": f"SRS\n{node_metric('SRS')}",
            "ARS": f"ARS\n{node_metric('ARS')}",
            "DRS": f"DRS\n{node_metric('DRS')}",
            "Analog IPOS": f"Analog IPOS\n{generated_metric('Analog IPOS')}",
            "Digital IPOS": f"Digital IPOS\n{generated_metric('Digital IPOS')}",
            "SysML": f"SysML Architecture\n{generated_metric('SysML')}",
        }
        fills = {"Source Spec": "#eaf0f6", "SRS": "#d9edf7", "ARS": "#f1f3f5", "DRS": "#d9ead3", "Analog IPOS": "#f8eadf", "Digital IPOS": "#e8f1dc", "SysML": "#eee4f7"}

        def connect(source: str, target: str, metric: str, label_offset: int = -14) -> None:
            link_name = f"{source} -> {target}"
            start_x = nodes[source][0] + node_width
            start_y = nodes[source][1] + node_height / 2
            end_x = nodes[target][0]
            end_y = nodes[target][1] + node_height / 2
            color = self._coverage_link_color(metric)
            selected = self.coverage_selected_link == link_name
            line_id = canvas.create_line(
                start_x + 3,
                start_y,
                end_x - 8,
                end_y,
                fill="#1f6feb" if selected else color,
                width=5 if selected else 3,
                arrow=tk.LAST,
                tags=("coverage_link", f"coverage_link_{source}_{target}"),
            )
            canvas.create_text(
                (start_x + end_x) / 2,
                (start_y + end_y) / 2 + label_offset,
                text=metric,
                font=("Segoe UI", max(8, round(10 * zoom)), "bold"),
                fill="#1f6feb" if selected else color,
                tags=("coverage_link_label", f"coverage_link_label_{source}_{target}"),
            )
            canvas.tag_bind(line_id, "<Button-1>", self._show_coverage_link_popup)
            canvas.tag_bind(line_id, "<Enter>", lambda _event: canvas.configure(cursor="hand2"))
            canvas.tag_bind(line_id, "<Leave>", lambda _event: canvas.configure(cursor=""))

        connect("Source Spec", "SRS", link_metric("Source -> SRS"))
        connect("SRS", "ARS", link_metric("SRS -> ARS"), -16)
        connect("SRS", "DRS", link_metric("SRS -> DRS"), 16)
        connect("ARS", "Analog IPOS", link_metric("ARS -> Analog IPOS"), -10)
        connect("DRS", "Digital IPOS", link_metric("DRS -> Digital IPOS"), 10)
        connect("SRS", "SysML", generated_metric("SysML"), -16)

        for name, (x, y) in nodes.items():
            canvas.create_rectangle(x, y, x + node_width, y + node_height, fill=fills[name], outline="#4a6075", width=2)
            canvas.create_text(
                x + node_width / 2,
                y + node_height / 2,
                text=labels[name],
                font=("Segoe UI", max(8, round(9 * zoom)), "bold"),
                fill="#1f2d3d",
                width=node_width - 12,
            )
            canvas.configure(scrollregion=(0, 0, width, height))
            canvas.xview_moveto(current_x)
            canvas.yview_moveto(current_y)

    def _show_coverage_link_popup(self, event) -> str:
        """Highlight and describe a clicked System Traceability connection."""
        canvas = self.coverage_canvas
        current = canvas.find_withtag("current")
        if not current:
            return "break"
        tags = canvas.gettags(current[0])
        link_tag = next((tag for tag in tags if tag.startswith("coverage_link_") and not tag.startswith("coverage_link_label_")), None)
        if not link_tag:
            return "break"
        link_name = link_tag.replace("coverage_link_", "", 1).replace("_", " -> ", 1)
        # The rendered tag uses one separator between source and target; recover the known link names.
        link_name = next((name for name in ("Source Spec -> SRS", "SRS -> ARS", "SRS -> DRS", "ARS -> Analog IPOS", "DRS -> Digital IPOS", "SRS -> SysML") if link_tag == f"coverage_link_{name.replace(' -> ', '_')}") , link_name)
        self.coverage_selected_link = link_name
        self._draw_system_traceability()
        metric_key = {
            "Source Spec -> SRS": "Source -> SRS",
            "SRS -> ARS": "SRS -> ARS",
            "SRS -> DRS": "SRS -> DRS",
            "ARS -> Analog IPOS": "ARS -> Analog IPOS",
            "DRS -> Digital IPOS": "DRS -> Digital IPOS",
        }.get(link_name, "")
        metric = self.coverage_values.get(metric_key, "Unavailable")
        messagebox.showinfo(
            "System Traceability Connection",
            f"{link_name}\n\nCoverage: {metric}\nSource specification: {self.coverage_values.get('source_document', 'Unknown')}",
            parent=self.root,
        )
        return "break"

    def _on_work_area_select_path(self, path: Path) -> None:
        if path.suffix.lower() == ".pdf":
            self._open_pdf_in_acrobat(path)
            return
        if path.suffix.lower() == DOCX_FILE_SUFFIX:
            self._open_docx_in_word(path)
            return
        if path.suffix.lower() == XLSX_FILE_SUFFIX:
            self._open_xlsx_in_excel(path)
            return
        if path.suffix.lower() not in TEXT_FILE_SUFFIXES:
            self._set_work_view(
                "This file type is not displayed as text in the read-only viewer.",
                f"Unsupported file: {path.relative_to(REPO_ROOT).as_posix()}",
            )
            return
        try:
            if path.stat().st_size > MAX_VIEWER_BYTES:
                self._set_work_view(
                    f"File exceeds the {MAX_VIEWER_BYTES // (1024 * 1024)} MB viewer limit.",
                    f"File too large: {path.relative_to(REPO_ROOT).as_posix()}",
                )
                return
            content = path.read_text(encoding="utf-8", errors="replace")
            if path.suffix.lower() == ".csv":
                self._set_csv_view(content, path)
            else:
                self._set_work_view(content, path.relative_to(REPO_ROOT).as_posix())
                self.current_work_path = path
        except OSError as exc:
            self._set_work_view(str(exc), f"Unable to read: {path.relative_to(REPO_ROOT).as_posix()}")

    def _open_docx_in_word(self, path: Path) -> None:
        label = path.relative_to(REPO_ROOT).as_posix()
        try:
            if not path.exists():
                raise FileNotFoundError(path)
            if hasattr(os, "startfile"):
                os.startfile(str(path))
            else:
                raise OSError("Microsoft Word file association is unavailable")
            self.current_work_path = path
            self.work_file_var.set(f"Opened in Microsoft Word: {label}")
        except OSError as exc:
            messagebox.showerror("Open DOCX", f"Unable to open in Microsoft Word:\n{exc}", parent=self.root)

    def _open_pdf_in_acrobat(self, path: Path) -> None:
        label = path.relative_to(REPO_ROOT).as_posix()
        try:
            if not path.exists():
                raise FileNotFoundError(path)
            program_files = os.environ.get("ProgramFiles", r"C:\Program Files")
            program_files_x86 = os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")
            local_app_data = os.environ.get("LOCALAPPDATA", "")
            candidates = [
                shutil.which("AcroRd32.exe"),
                str(Path(program_files) / "Adobe" / "Acrobat Reader DC" / "Reader" / "AcroRd32.exe"),
                str(Path(program_files_x86) / "Adobe" / "Acrobat Reader DC" / "Reader" / "AcroRd32.exe"),
                str(Path(program_files) / "Adobe" / "Acrobat Classic" / "Acrobat" / "Acrobat.exe"),
                str(Path(program_files_x86) / "Adobe" / "Acrobat Classic" / "Acrobat" / "Acrobat.exe"),
            ]
            if local_app_data:
                candidates.append(str(Path(local_app_data) / "Adobe" / "Acrobat Reader DC" / "Reader" / "AcroRd32.exe"))
            executable = next((candidate for candidate in candidates if candidate and Path(candidate).is_file()), None)
            if executable is None:
                raise OSError("Adobe Acrobat Reader (AcroRd32.exe) was not found")
            subprocess.Popen([executable, str(path)])
            self.current_work_path = path
            self.work_file_var.set(f"Opened in Adobe Acrobat Reader: {label}")
        except OSError as exc:
            messagebox.showerror("Open PDF", f"Unable to open in Adobe Acrobat Reader:\n{exc}", parent=self.root)

    def _open_xlsx_in_excel(self, path: Path) -> None:
        label = path.relative_to(REPO_ROOT).as_posix()
        try:
            if not path.exists():
                raise FileNotFoundError(path)
            if hasattr(os, "startfile"):
                os.startfile(str(path))
            else:
                raise OSError("Microsoft Excel file association is unavailable")
            self.current_work_path = path
            self.work_file_var.set(f"Opened in Microsoft Excel: {label}")
        except OSError as exc:
            messagebox.showerror("Open XLSX", f"Unable to open in Microsoft Excel:\n{exc}", parent=self.root)

    def _open_pdf_in_work_viewer(self, path: Path) -> None:
        label = path.relative_to(REPO_ROOT).as_posix()
        if fitz is None or Image is None or ImageTk is None:
            self._set_work_view("PDF reader dependencies are unavailable. Install PyMuPDF and Pillow.", f"PDF unavailable: {label}")
            return
        try:
            if self.pdf_reader_window is not None and self.pdf_reader_window.winfo_exists():
                self.pdf_reader_window.destroy()
            if self.pdf_reader_document is not None:
                self.pdf_reader_document.close()
            self.pdf_reader_document = fitz.open(str(path))
            self.pdf_reader_path = path
            self.pdf_reader_page = 0
            self.pdf_reader_zoom = 1.0
            self.pdf_reader_find_page = None
            if self.pdf_reader_find_var is not None:
                self.pdf_reader_find_var.set("")
            if self.pdf_reader_find_status_var is not None:
                self.pdf_reader_find_status_var.set("")
            self.current_work_path = path
            self._build_pdf_reader_window(label)
        except (OSError, RuntimeError, ValueError) as exc:
            self._set_work_view(f"Unable to read PDF: {exc}", f"Unable to read: {label}")

    def _build_pdf_reader_window(self, label: str) -> None:
        window = tk.Toplevel(self.root)
        self.pdf_reader_window = window
        window.title(f"PDF Reader - {Path(label).name}")
        self._set_responsive_geometry(window, 1100, 800, 640, 480)
        window.transient(self.root)

        controls = ttk.Frame(window, padding=6)
        controls.pack(fill=tk.X)
        self.pdf_reader_page_var = tk.StringVar()
        ttk.Button(controls, text="Previous", command=lambda: self._show_pdf_page(-1), width=10).pack(side=tk.LEFT)
        ttk.Button(controls, text="Next", command=lambda: self._show_pdf_page(1), width=10).pack(side=tk.LEFT, padx=(6, 0))
        ttk.Label(controls, textvariable=self.pdf_reader_page_var).pack(side=tk.LEFT, padx=(12, 0))
        self.pdf_reader_find_var = tk.StringVar()
        self.pdf_reader_find_status_var = tk.StringVar(value="")
        ttk.Label(controls, text="Find:").pack(side=tk.LEFT, padx=(18, 4))
        find_entry = ttk.Entry(controls, textvariable=self.pdf_reader_find_var, width=28)
        find_entry.pack(side=tk.LEFT)
        find_entry.bind("<Return>", lambda _event: self._find_pdf_text(1))
        ttk.Button(controls, text="Find Next", command=lambda: self._find_pdf_text(1), width=10).pack(side=tk.LEFT, padx=(4, 0))
        ttk.Button(controls, text="Find Previous", command=lambda: self._find_pdf_text(-1), width=12).pack(side=tk.LEFT, padx=(4, 0))
        ttk.Label(controls, textvariable=self.pdf_reader_find_status_var).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Button(controls, text="-", command=lambda: self._zoom_pdf_page(1 / 1.15), width=3).pack(side=tk.RIGHT)
        ttk.Button(controls, text="+", command=lambda: self._zoom_pdf_page(1.15), width=3).pack(side=tk.RIGHT, padx=(6, 0))
        self.pdf_reader_zoom_var = tk.StringVar(value="100%")
        ttk.Label(controls, textvariable=self.pdf_reader_zoom_var).pack(side=tk.RIGHT, padx=(8, 6))

        frame = ttk.Frame(window)
        frame.pack(fill=tk.BOTH, expand=True, padx=6, pady=(0, 6))
        canvas = tk.Canvas(frame, background="#4b5563", highlightthickness=0)
        scroll_y = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=canvas.yview)
        scroll_x = ttk.Scrollbar(frame, orient=tk.HORIZONTAL, command=canvas.xview)
        canvas.configure(xscrollcommand=scroll_x.set, yscrollcommand=scroll_y.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        scroll_y.grid(row=0, column=1, sticky="ns")
        scroll_x.grid(row=1, column=0, sticky="ew")
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)
        self.pdf_reader_canvas = canvas
        canvas.bind("<Control-MouseWheel>", lambda event: self._zoom_pdf_page(1.15 if event.delta > 0 else 1 / 1.15), add="+")
        window.bind("<Control-f>", lambda _event: (find_entry.focus_set(), "break")[1])
        window.protocol("WM_DELETE_WINDOW", self._close_pdf_reader)
        find_entry.focus_set()
        self._render_pdf_page()

    def _find_pdf_text(self, direction: int = 1) -> str:
        document = self.pdf_reader_document
        query = self.pdf_reader_find_var.get().strip() if self.pdf_reader_find_var is not None else ""
        if document is None or not query:
            if self.pdf_reader_find_status_var is not None:
                self.pdf_reader_find_status_var.set("Enter text")
            return "break"

        start_page = self.pdf_reader_page
        if self.pdf_reader_find_page == start_page:
            start_page = (start_page + direction) % document.page_count
        page_indexes = [
            (start_page + direction * offset) % document.page_count
            for offset in range(document.page_count)
        ]
        query_folded = query.casefold()
        for page_index in page_indexes:
            page_text = document.load_page(page_index).get_text("text")
            if query_folded in page_text.casefold():
                self.pdf_reader_find_page = page_index
                self.pdf_reader_page = page_index
                self._render_pdf_page()
                if self.pdf_reader_find_status_var is not None:
                    self.pdf_reader_find_status_var.set(f"Page {page_index + 1}")
                return "break"

        self.pdf_reader_find_page = None
        if self.pdf_reader_find_status_var is not None:
            self.pdf_reader_find_status_var.set("Not found")
        return "break"

    def _render_pdf_page(self) -> None:
        document = self.pdf_reader_document
        if document is None or not document.page_count or not hasattr(self, "pdf_reader_canvas"):
            return
        page = document.load_page(self.pdf_reader_page)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(self.pdf_reader_zoom, self.pdf_reader_zoom), alpha=False)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        self.pdf_reader_photo = ImageTk.PhotoImage(image)
        canvas = self.pdf_reader_canvas
        canvas.delete("all")
        canvas.create_image(12, 12, anchor=tk.NW, image=self.pdf_reader_photo)
        canvas.configure(scrollregion=(0, 0, pixmap.width + 24, pixmap.height + 24))
        self.pdf_reader_page_var.set(f"Page {self.pdf_reader_page + 1} of {document.page_count}")
        self.pdf_reader_zoom_var.set(f"{round(self.pdf_reader_zoom * 100)}%")

    def _show_pdf_page(self, offset: int) -> None:
        document = self.pdf_reader_document
        if document is None:
            return
        self.pdf_reader_page = min(max(0, self.pdf_reader_page + offset), document.page_count - 1)
        self._render_pdf_page()

    def _zoom_pdf_page(self, factor: float) -> str:
        self.pdf_reader_zoom = min(3.0, max(0.5, self.pdf_reader_zoom * factor))
        self._render_pdf_page()
        return "break"

    def _close_pdf_reader(self) -> None:
        window = self.pdf_reader_window
        if self.pdf_reader_document is not None:
            self.pdf_reader_document.close()
        self.pdf_reader_document = None
        self.pdf_reader_window = None
        self.pdf_reader_path = None
        self.pdf_reader_find_page = None
        if window is not None and window.winfo_exists():
            window.destroy()

    def _set_work_view(self, text: str, label: str) -> None:
        self.current_work_path = None
        self.current_csv_rows = []
        self.current_csv_headers = []
        self._clear_work_find_highlights()
        self.work_csv_frame.pack_forget()
        self.work_text_frame.pack(fill=tk.BOTH, expand=True)
        self.work_text_x_frame.pack(fill=tk.X)
        self.work_file_var.set(label)
        self.work_viewer.configure(state=tk.NORMAL)
        self.work_viewer.delete("1.0", tk.END)
        self.work_viewer.insert("1.0", text)
        self._highlight_work_view(text, Path(label).suffix.lower())
        self.work_viewer.configure(state=tk.NORMAL)
        self.work_viewer.see("1.0")
        self._refresh_work_find_results()

    def _bind_copy_support(self, widget: tk.Text) -> None:
        widget.bind("<Control-c>", lambda _event, target=widget: self._copy_text_selection(target), add="+")
        widget.bind("<Button-3>", lambda event, target=widget: self._show_text_copy_menu(event, target), add="+")

    def _copy_text_selection(self, widget: tk.Text) -> str:
        try:
            selected = widget.get(tk.SEL_FIRST, tk.SEL_LAST)
        except tk.TclError:
            return "break"
        self.root.clipboard_clear()
        self.root.clipboard_append(selected)
        self.root.update()
        return "break"

    @staticmethod
    def _block_rendered_text_edit(event) -> Optional[str]:
        """Keep rendered text selectable without allowing keyboard edits."""
        control_pressed = bool(event.state & 0x0004)
        if control_pressed:
            return None if event.keysym.lower() in {"c", "a", "f"} else "break"
        if event.char or event.keysym in {"BackSpace", "Delete", "Return", "Tab"}:
            return "break"
        return None

    def _show_text_copy_menu(self, event, widget: tk.Text) -> str:
        menu = tk.Menu(widget, tearoff=False)
        menu.add_command(label="Copy", command=lambda: self._copy_text_selection(widget))
        return self._show_popup_menu(menu, event)

    def _open_work_find(self, _event=None) -> str:
        if not self.work_find_frame.winfo_ismapped():
            before_widget = self.work_csv_frame if self.work_csv_frame.winfo_ismapped() else self.work_text_frame
            self.work_find_frame.pack(fill=tk.X, pady=(0, 4), before=before_widget)
        self.work_find_entry.focus_set()
        self.work_find_entry.selection_range(0, tk.END)
        self._refresh_work_find_results()
        return "break"

    def _close_work_find(self, _event=None) -> str:
        self.work_find_frame.pack_forget()
        self.work_find_status_var.set("")
        self._clear_work_find_highlights()
        return "break"

    def _on_work_find_changed(self, event=None) -> None:
        if event is not None and event.keysym in {"Return", "Escape", "Shift_L", "Shift_R", "Control_L", "Control_R"}:
            return
        self._refresh_work_find_results()

    def _clear_work_find_highlights(self) -> None:
        self.work_viewer.tag_remove("find_hit", "1.0", tk.END)
        self.work_viewer.tag_remove("find_current", "1.0", tk.END)
        if hasattr(self, "work_csv_table"):
            for index, item_id in enumerate(self.work_csv_table.get_children(), start=1):
                self.work_csv_table.item(item_id, tags=("even" if index % 2 == 0 else "odd",))
        self.work_find_results = []
        self.work_find_index = -1

    def _is_csv_work_view_active(self) -> bool:
        return bool(self.current_work_path and self.current_work_path.suffix.lower() == ".csv" and self.work_csv_frame.winfo_ismapped())

    def _refresh_work_find_results(self) -> None:
        query = self.work_find_var.get()
        self._clear_work_find_highlights()
        if not query:
            self.work_find_status_var.set("")
            return
        if self._is_csv_work_view_active():
            self._refresh_csv_find_results(query)
        else:
            self._refresh_text_find_results(query)

    def _refresh_text_find_results(self, query: str) -> None:
        count_var = tk.IntVar()
        start = "1.0"
        while True:
            index = self.work_viewer.search(query, start, nocase=True, stopindex=tk.END, count=count_var)
            if not index:
                break
            length = count_var.get()
            if length <= 0:
                break
            end = f"{index} + {length} chars"
            self.work_viewer.tag_add("find_hit", index, end)
            self.work_find_results.append((index, end))
            start = end

        if self.work_find_results:
            self.work_find_index = 0
            self._show_current_text_find_result()
        else:
            self.work_find_status_var.set("0 matches")

    def _refresh_csv_find_results(self, query: str) -> None:
        needle = query.lower()
        for item_id in self.work_csv_table.get_children():
            values = [str(value) for value in self.work_csv_table.item(item_id, "values")]
            if any(needle in value.lower() for value in values):
                self.work_find_results.append(item_id)
                self.work_csv_table.item(item_id, tags=("find_hit",))

        if self.work_find_results:
            self.work_find_index = 0
            self._show_current_csv_find_result()
        else:
            self.work_find_status_var.set("0 matches")

    def _show_current_text_find_result(self) -> None:
        if not self.work_find_results:
            self.work_find_status_var.set("0 matches")
            return
        self.work_viewer.tag_remove("find_current", "1.0", tk.END)
        start, end = self.work_find_results[self.work_find_index]
        self.work_viewer.tag_add("find_current", start, end)
        self.work_viewer.see(start)
        self.work_find_status_var.set(f"{self.work_find_index + 1}/{len(self.work_find_results)}")

    def _show_current_csv_find_result(self) -> None:
        if not self.work_find_results:
            self.work_find_status_var.set("0 matches")
            return
        for item_id in self.work_csv_table.get_children():
            if item_id in self.work_find_results:
                self.work_csv_table.item(item_id, tags=("find_hit",))
        item_id = self.work_find_results[self.work_find_index]
        self.work_csv_table.item(item_id, tags=("find_current",))
        self.work_csv_table.selection_set(item_id)
        self.work_csv_table.focus(item_id)
        self.work_csv_table.see(item_id)
        self.work_find_status_var.set(f"{self.work_find_index + 1}/{len(self.work_find_results)}")

    def _find_next_in_work_view(self, _event=None) -> str:
        if not self.work_find_frame.winfo_ismapped():
            return self._open_work_find()
        if not self.work_find_results:
            self._refresh_work_find_results()
        if self.work_find_results:
            self.work_find_index = (self.work_find_index + 1) % len(self.work_find_results)
            if self._is_csv_work_view_active():
                self._show_current_csv_find_result()
            else:
                self._show_current_text_find_result()
        return "break"

    def _find_previous_in_work_view(self, _event=None) -> str:
        if not self.work_find_frame.winfo_ismapped():
            return self._open_work_find()
        if not self.work_find_results:
            self._refresh_work_find_results()
        if self.work_find_results:
            self.work_find_index = (self.work_find_index - 1) % len(self.work_find_results)
            if self._is_csv_work_view_active():
                self._show_current_csv_find_result()
            else:
                self._show_current_text_find_result()
        return "break"

    def _markdown_anchor_slug(self, heading: str) -> str:
        slug = re.sub(r"\s*\{#[A-Za-z0-9_.:-]+\}\s*$", "", heading.strip()).lower()
        slug = re.sub(r"[`*_~\[\]()]", "", slug)
        slug = re.sub(r"[^a-z0-9\s-]", "", slug)
        slug = re.sub(r"\s+", "-", slug)
        slug = re.sub(r"-+", "-", slug).strip("-")
        return slug

    def _markdown_anchor_target(self, href: str) -> str:
        target = unquote(href[1:]).strip().lower()
        target = re.sub(r"^#", "", target)
        target = re.sub(r"\s+", "-", target)
        return target

    def _jump_to_work_anchor(self, line_index: str) -> None:
        self.work_viewer.see(line_index)
        self.work_viewer.tag_remove("anchor_hit", "1.0", tk.END)
        line_end = f"{line_index} lineend"
        self.work_viewer.tag_add("anchor_hit", line_index, line_end)
        self.root.after(1200, lambda: self.work_viewer.tag_remove("anchor_hit", "1.0", tk.END))

    def _set_csv_view(self, content: str, path: Path) -> None:
        self.current_work_path = path
        self.csv_sort_directions = {}
        self._clear_work_find_highlights()
        self.work_text_frame.pack_forget()
        self.work_text_x_frame.pack_forget()
        if self.work_find_frame.winfo_ismapped():
            self.work_find_frame.pack_forget()
            self.work_find_frame.pack(fill=tk.X, pady=(0, 4), before=self.work_csv_frame)
        self.work_csv_frame.pack(fill=tk.BOTH, expand=True)

        for item in self.work_csv_table.get_children():
            self.work_csv_table.delete(item)

        rows = list(csv.reader(StringIO(content)))
        self.current_csv_rows = rows
        if not rows:
            self.current_csv_headers = []
            self.work_csv_table.configure(columns=())
            self.work_file_var.set(f"CSV: {path.relative_to(REPO_ROOT).as_posix()} (0 rows)")
            return

        headers = [header.strip() or f"Column {index + 1}" for index, header in enumerate(rows[0])]
        max_columns = max(len(row) for row in rows)
        while len(headers) < max_columns:
            headers.append(f"Column {len(headers) + 1}")
        columns = [f"col_{index}" for index in range(max_columns)]
        self.current_csv_headers = headers
        self.work_csv_table.configure(columns=columns)

        for column_id, header in zip(columns, headers):
            sample_values = [header]
            sample_values.extend(row[columns.index(column_id)] for row in rows[1:80] if columns.index(column_id) < len(row))
            width = min(max(max(len(value) for value in sample_values) * 8 + 24, 80), 360)
            self.work_csv_table.heading(column_id, text=header)
            self.work_csv_table.column(column_id, width=width, minwidth=60, stretch=False, anchor="w")

        for index, row in enumerate(rows[1:], start=1):
            values = row + [""] * (max_columns - len(row))
            self.work_csv_table.insert("", tk.END, values=values, tags=("even" if index % 2 == 0 else "odd",))

        self.work_csv_table.tag_configure("even", background="#f7fbff")
        self.work_csv_table.tag_configure("odd", background="#ffffff")
        self.work_file_var.set(f"CSV: {path.relative_to(REPO_ROOT).as_posix()} ({max(len(rows) - 1, 0)} rows)")
        self._refresh_work_find_results()

    def _sort_csv_by_heading(self, event) -> str:
        if self.work_csv_table.identify_region(event.x, event.y) != "heading":
            return "break"

        column_id = self.work_csv_table.identify_column(event.x)
        try:
            column_index = int(column_id.lstrip("#")) - 1
        except ValueError:
            return "break"

        if column_index < 0 or column_index >= len(self.current_csv_headers):
            return "break"

        ascending = not self.csv_sort_directions.get(column_index, False)
        self.csv_sort_directions[column_index] = ascending
        header = self.current_csv_headers[column_index]
        data_rows = self.current_csv_rows[1:]
        data_rows.sort(
            key=lambda row: (row[column_index] if column_index < len(row) else "").casefold(),
            reverse=not ascending,
        )
        self.current_csv_rows = [self.current_csv_rows[0], *data_rows]

        for item_id in self.work_csv_table.get_children():
            self.work_csv_table.delete(item_id)
        for index, row in enumerate(data_rows, start=1):
            values = row + [""] * (len(self.current_csv_headers) - len(row))
            self.work_csv_table.insert(
                "",
                tk.END,
                values=values,
                tags=("even" if index % 2 == 0 else "odd",),
            )

        direction = "ascending" if ascending else "descending"
        self.work_file_var.set(
            f"CSV: {self.current_work_path.relative_to(REPO_ROOT).as_posix()} "
            f"({len(data_rows)} rows, sorted by {header} {direction})"
        )
        self._refresh_work_find_results()
        return "break"

    def _show_work_tree_context_menu(self, event) -> str:
        item_id = self.work_tree.identify_row(event.y)
        if item_id:
            self.work_tree.selection_set(item_id)
            self.work_tree.focus(item_id)
            self._on_work_area_select()
        path = self._work_area_path(item_id) if item_id else None
        if path is None or path.is_dir():
            return "break"
        menu = tk.Menu(self.root, tearoff=False)
        menu.add_command(label="Open read-only", command=self._on_work_area_select)
        if path.name == "requirements_summary.csv":
            menu.add_command(label="Show ID counts", command=self._show_requirement_id_counts)
        return self._show_popup_menu(menu, event)

    def _show_csv_context_menu(self, event) -> str:
        if self.current_work_path is None:
            return "break"
        row_id = self.work_csv_table.identify_row(event.y)
        if row_id:
            self.work_csv_table.selection_set(row_id)
            self.work_csv_table.focus(row_id)
        menu = tk.Menu(self.root, tearoff=False)
        menu.add_command(label="Show CSV size", command=self._show_csv_size)
        if self.current_work_path.name == "requirements_summary.csv":
            menu.add_command(label="Show ID counts", command=self._show_requirement_id_counts)
        return self._show_popup_menu(menu, event)

    def _show_csv_size(self) -> None:
        row_count = max(len(self.current_csv_rows) - 1, 0)
        column_count = len(self.current_csv_headers)
        messagebox.showinfo("CSV size", f"Rows: {row_count}\nColumns: {column_count}", parent=self.root)

    def _show_requirement_id_counts(self) -> None:
        if self.current_work_path is None or self.current_work_path.name != "requirements_summary.csv":
            messagebox.showinfo("ID counts", "Open requirements_summary.csv first.", parent=self.root)
            return
        if not self.current_csv_rows:
            try:
                content = self.current_work_path.read_text(encoding="utf-8", errors="replace")
                self.current_csv_rows = list(csv.reader(StringIO(content)))
                self.current_csv_headers = self.current_csv_rows[0] if self.current_csv_rows else []
            except OSError as exc:
                messagebox.showerror("ID counts", str(exc), parent=self.root)
                return

        headers = [header.strip() for header in self.current_csv_headers]
        header_lookup = {header.lower(): index for index, header in enumerate(headers)}
        id_index = header_lookup.get("id")
        source_id_index = header_lookup.get("source_req_id")
        data_rows = self.current_csv_rows[1:]

        def values_at(index: Optional[int]) -> list[str]:
            if index is None:
                return []
            return [row[index].strip() for row in data_rows if index < len(row) and row[index].strip()]

        ids = values_at(id_index)
        source_ids = values_at(source_id_index)
        duplicate_ids = sorted(value for value, count in Counter(ids).items() if count > 1)
        duplicate_source_ids = sorted(value for value, count in Counter(source_ids).items() if count > 1)
        message = (
            f"Requirement rows: {len(data_rows)}\n"
            f"id values: {len(ids)}\n"
            f"unique id values: {len(set(ids))}\n"
            f"duplicate id values: {len(duplicate_ids)}\n"
            f"source_req_id values: {len(source_ids)}\n"
            f"unique source_req_id values: {len(set(source_ids))}\n"
            f"duplicate source_req_id values: {len(duplicate_source_ids)}"
        )
        messagebox.showinfo("requirements_summary.csv ID counts", message, parent=self.root)

    def _highlight_work_view(self, text: str, suffix: str) -> None:
        for tag in list(self.work_viewer.tag_names()):
            if tag.startswith("internal_link_"):
                self.work_viewer.tag_delete(tag)
        for tag in ("comment", "string", "keyword", "number", "heading", "link", "delimiter", "anchor_hit"):
            self.work_viewer.tag_remove(tag, "1.0", tk.END)

        self.work_viewer.tag_configure("anchor_hit", background="#fff2a8")

        def add_matches(tag: str, pattern: str) -> None:
            for match in re.finditer(pattern, text, flags=re.MULTILINE):
                start = f"1.0 + {match.start()} chars"
                end = f"1.0 + {match.end()} chars"
                self.work_viewer.tag_add(tag, start, end)

        def add_markdown_internal_links() -> None:
            anchors = {}
            slug_counts = Counter()
            for match in re.finditer(r"^(#{1,6})\s+(.+?)\s*#*\s*$", text, flags=re.MULTILINE):
                line_index = f"1.0 + {match.start()} chars"
                explicit_anchor = re.search(r"\{#([A-Za-z0-9_.:-]+)\}\s*$", match.group(2))
                if explicit_anchor:
                    anchors[explicit_anchor.group(1).lower()] = line_index

                slug = self._markdown_anchor_slug(match.group(2))
                if slug:
                    slug_counts[slug] += 1
                    anchor = slug if slug_counts[slug] == 1 else f"{slug}-{slug_counts[slug] - 1}"
                    anchors.setdefault(anchor, line_index)

            for index, match in enumerate(re.finditer(r"\[[^\]]+\]\((#[^)]+)\)", text), start=1):
                anchor = self._markdown_anchor_target(match.group(1))
                if anchor not in anchors:
                    continue
                tag = f"internal_link_{index}"
                start = f"1.0 + {match.start()} chars"
                end = f"1.0 + {match.end()} chars"
                self.work_viewer.tag_add("link", start, end)
                self.work_viewer.tag_add(tag, start, end)
                self.work_viewer.tag_configure(tag, foreground="#0563c1", underline=True)
                self.work_viewer.tag_bind(tag, "<Button-1>", lambda _e, target=anchors[anchor]: self._jump_to_work_anchor(target))
                self.work_viewer.tag_bind(tag, "<Enter>", lambda _e: self.work_viewer.config(cursor="hand2"))
                self.work_viewer.tag_bind(tag, "<Leave>", lambda _e: self.work_viewer.config(cursor=""))

        if suffix in {".py", ".json", ".tex", ".mmd"}:
            add_matches("comment", r"#.*$" if suffix == ".py" else r"%.*$" if suffix == ".tex" else r"//.*$")
            add_matches("string", r"(?:\"(?:\\.|[^\"])*\"|'(?:\\.|[^'])*')")
            if suffix == ".py":
                keywords = r"\b(?:and|as|assert|class|def|elif|else|except|False|for|from|if|import|in|is|None|not|or|pass|raise|return|True|try|while|with|yield)\b"
            elif suffix == ".json":
                keywords = r"\b(?:true|false|null)\b"
            elif suffix == ".tex":
                keywords = r"\\[A-Za-z]+"
            else:
                keywords = r"\b(?:flowchart|graph|subgraph|style|classDef)\b"
            add_matches("keyword", keywords)
            add_matches("number", r"\b(?:0x[0-9A-Fa-f]+|\d+(?:\.\d+)?)\b")
        elif suffix == ".md":
            add_matches("heading", r"^#{1,6} .*$")
            add_matches("keyword", r"^```.*$|^> .*$|^[-*] .*$")
            add_matches("link", r"\[[^\]]+\]\([^\)]+\)")
            add_markdown_internal_links()
        elif suffix == ".csv":
            add_matches("heading", r"\A[^\n]*")
            add_matches("delimiter", r",|;|\t")
        elif suffix in {".txt", ".log"}:
            add_matches("comment", r"^\s*(?:#|//|\[DEBUG\]).*$")

    def _on_work_area_select(self, _event=None) -> None:
        selection = self.work_tree.selection()
        if not selection:
            return
        path = self._work_area_path(selection[0])
        if path is None:
            self._set_work_view("The selected path is outside the repository and cannot be opened.", "Unsafe path")
            return
        if path.is_dir():
            self.work_file_var.set(f"Folder: {path.relative_to(REPO_ROOT).as_posix()}")
            return
        if path.suffix.lower() == ".pdf":
            self._open_pdf_in_acrobat(path)
            return
        if path.suffix.lower() == DOCX_FILE_SUFFIX:
            self._open_docx_in_word(path)
            return
        if path.suffix.lower() == XLSX_FILE_SUFFIX:
            self._open_xlsx_in_excel(path)
            return
        if path.suffix.lower() not in TEXT_FILE_SUFFIXES:
            self._set_work_view(
                "This file type is not displayed as text in the read-only viewer.",
                f"Unsupported file: {path.relative_to(REPO_ROOT).as_posix()}",
            )
            return
        self._on_work_area_select_path(path)

    def run_status(self) -> None:
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "status"]
        self._run_command_async(cmd, "CLI status")

    def _snapshot_selector_args(self, parent=None) -> Optional[List[str]]:
        owner = parent or self.root
        try:
            context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8")) if PROJECT_CONTEXT_PATH.exists() else {}
            project_id = str(context.get("project_name") or REPO_ROOT.name)
            connection = connect(REPO_ROOT)
            try:
                rows = connection.execute(
                    """SELECT snapshot_id, approved_by, approved_at
                         FROM snapshots
                        WHERE project_id = ? AND status = 'approved'
                        ORDER BY approved_at DESC, snapshot_id DESC LIMIT 5""",
                    (project_id,),
                ).fetchall()
            finally:
                connection.close()
        except Exception as exc:
            messagebox.showerror("Approved Snapshot", f"Could not read approved snapshots: {exc}", parent=owner)
            return None
        if not rows:
            messagebox.showinfo("Approved Snapshot", "No approved snapshot is available. Create an immutable approved snapshot first.", parent=owner)
            return None

        dialog = tk.Toplevel(owner)
        dialog.title("Select Approved Snapshot")
        dialog.transient(owner)
        dialog.resizable(False, False)
        dialog.grab_set()
        self._activate_popup(dialog, "dialog")
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frame, text="Approved Snapshot", font=("Segoe UI", 13, "bold")).pack(anchor=tk.W)
        ttk.Label(frame, text="Choose Latest approved or one of the five most recently approved snapshots.", wraplength=540).pack(anchor=tk.W, pady=(4, 10))
        choices = ["Latest approved"]
        choice_to_selector = {"Latest approved": ["--use-latest-approved"]}
        for row in rows:
            snapshot_id = str(row["snapshot_id"])
            label = f"{snapshot_id} | {row['approved_at'] or 'unknown date'} | {row['approved_by'] or 'unknown reviewer'}"
            choices.append(label)
            choice_to_selector[label] = ["--snapshot-id", snapshot_id]
        selected = tk.StringVar(value=choices[0])
        combo = ttk.Combobox(frame, textvariable=selected, values=choices, state="readonly", width=78)
        combo.pack(fill=tk.X, pady=(0, 14))
        result: List[str] = []

        def choose() -> None:
            result.extend(choice_to_selector[selected.get()])
            dialog.destroy()

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X)
        ttk.Button(buttons, text="Cancel", command=dialog.destroy, width=12).pack(side=tk.RIGHT)
        ttk.Button(buttons, text="Use Snapshot", command=choose, width=15).pack(side=tk.RIGHT, padx=(0, 8))
        combo.focus_set()
        dialog.wait_window()
        return result or None

    def _show_snapshot_report_menu(self) -> None:
        try:
            context = json.loads(PROJECT_CONTEXT_PATH.read_text(encoding="utf-8")) if PROJECT_CONTEXT_PATH.exists() else {}
            project_id = str(context.get("project_name") or REPO_ROOT.name)
            connection = connect(REPO_ROOT)
            try:
                snapshot_rows = connection.execute(
                    "SELECT snapshot_id FROM snapshots WHERE project_id = ? AND status = 'approved' ORDER BY approved_at DESC, snapshot_id DESC",
                    (project_id,),
                ).fetchall()
            finally:
                connection.close()
        except Exception as exc:
            messagebox.showerror("Hierarchy Traceability Reports", f"Could not read approved snapshots: {exc}", parent=self.root)
            return
        if not snapshot_rows:
            messagebox.showinfo("Hierarchy Traceability Reports", "No approved snapshot is available. Complete mapping approval and create an immutable approved snapshot first.", parent=self.root)
            return

        dialog = tk.Toplevel(self.root)
        dialog.title("Hierarchy Traceability Reports")
        self._set_responsive_geometry(dialog, 620, 520, 620, 520)
        dialog.transient(self.root)
        dialog.grab_set()
        self._activate_popup(dialog, "dialog")
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frame, text="Generate Hierarchy Traceability XLSX", font=("Segoe UI", 13, "bold")).pack(anchor=tk.W)
        ttk.Label(frame, text="Choose a source-requirement catalog, a direct-upstream report, or all levels with imported and generated requirements.", wraplength=560).pack(anchor=tk.W, pady=(4, 12))

        report_var = tk.StringVar(value="all")
        report_options = (
            ("primary_source", "Primary source requirements (XLSX)"),
            ("integrated_sources", "Source spec integrated reqs (primary + supplementary XLSX)"),
            ("srs", "SRS vs Top Specification"),
            ("drs", "DRS vs SRS"),
            ("ars", "ARS vs SRS"),
            ("digital_ipos", "Digital IPOS vs DRS"),
            ("analog_ipos", "Analog IPOS vs ARS"),
            ("all", "All levels + inclusive workbook (imported and generated reqs)"),
        )
        for value, label in report_options:
            ttk.Radiobutton(frame, text=label, variable=report_var, value=value).pack(anchor=tk.W, pady=2)

        ttk.Separator(frame, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=12)
        ttk.Label(frame, text="Approved snapshot:").pack(anchor=tk.W)
        snapshot_choices = ["Latest approved"] + [str(row["snapshot_id"]) for row in snapshot_rows]
        snapshot_var = tk.StringVar(value=snapshot_choices[0])
        snapshot_combo = ttk.Combobox(frame, textvariable=snapshot_var, values=snapshot_choices, state="readonly", width=66)
        snapshot_combo.pack(fill=tk.X, pady=(4, 14))

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X, side=tk.BOTTOM)

        def generate() -> None:
            selected_snapshot = snapshot_var.get()
            selector = ["--use-latest-approved"] if selected_snapshot == "Latest approved" else ["--snapshot-id", selected_snapshot]
            report = report_var.get()
            dialog.destroy()
            command = [PYTHON_EXE, str(SCRIPTS_DIR / "generate_snapshot_reports.py"), *selector, "--report", report]

            def show_report() -> None:
                self._refresh_work_area_and_traceability()
                source_outputs = {
                    "primary_source": "primary_source_requirements.xlsx",
                    "integrated_sources": "source_spec_integrated_reqs.xlsx",
                }
                if report in source_outputs:
                    self._open_xlsx_in_excel(TRACEABILITY_REPORT_DIR / source_outputs[report])

            self._run_command_async(
                command,
                f"Generate {report.upper()} hierarchy traceability report",
                on_success=show_report,
            )

        ttk.Button(buttons, text="Cancel", command=dialog.destroy, width=12).pack(side=tk.RIGHT)
        ttk.Button(buttons, text="Generate XLSX", command=generate, width=16).pack(side=tk.RIGHT, padx=(0, 8))
        snapshot_combo.focus_set()

    def _open_conflict_review(self) -> None:
        reports = sorted((REPO_ROOT / "artifacts/canonical_workflow/conflicts").glob("*.json"))
        if not reports:
            messagebox.showinfo("Conflict Review", "No conflict review report has been generated.", parent=self.root)
            return
        self._on_work_area_select_path(reports[-1])

    def regenerate_sysml(self) -> None:
        selector = self._snapshot_selector_args()
        if selector is None:
            return
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "run_final_sysml_phase.py"), *selector]
        self._run_command_async(cmd, "Final SysML generation and central validation", on_success=self._refresh_generated_sysml)

    def _review_regenerated_sysml(self) -> None:
        self._run_command_async(
            [PYTHON_EXE, str(SCRIPTS_DIR / "review_architecture_sysml.py")],
            "SysML architecture model review",
            on_success=self._refresh_generated_sysml,
        )

    def _refresh_generated_sysml(self) -> None:
        self.refresh_work_area()
        self._refresh_sysml_architecture()
        self.work_tabs.select(self.sysml_architecture_tab)

    def run_taxonomy_update(self) -> None:
        if not self._show_taxonomy_update_guidance():
            self.append_chat("GUI: Taxonomy update canceled by user.")
            return
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "taxonomy-update"]
        self._run_command_async(cmd, "CLI taxonomy-update")

    def run_stage5_after_stage2a(self) -> None:
        confirm = messagebox.askyesno(
            "Run S5 after S2a",
            (
                "Run Stage 5 in requirements-only mode?\n\n"
                "This path skips Stage 3 (SRS) and Stage 4 (ARS) generation, and builds DRS "
                "from Stage 1 + Stage 2A artifacts only."
            ),
            parent=self.root,
        )
        if not confirm:
            return
        cmd = [PYTHON_EXE, str(SCRIPTS_DIR / "workflow_cli.py"), "drs-after-stage2a"]
        self._run_command_async(cmd, "CLI drs-after-stage2a", stage="5")

    def _show_taxonomy_update_guidance(self) -> bool:
        manual_file = str(PROJECT_CONTEXT_PATH)
        message = (
            "Before running taxonomy update, review/update this file manually:\n\n"
            f"- {manual_file}\n\n"
            "Fields to check:\n"
            "- source_spec_path (current source spec)\n"
            "- requirement_id_rules.source_req_id_patterns (expected ID family)\n\n"
            "Then run taxonomy update to regenerate taxonomy from Stage 1 OCR index.\n\n"
            "Do you want to continue now?"
        )
        return messagebox.askyesno("Taxonomy Update Guidance", message, parent=self.root)


def main() -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    with STARTUP_LOG.open("a", encoding="utf-8") as fh:
        fh.write(f"[{datetime.now().isoformat()}] Startup begin. Python={sys.executable}\n")

    try:
        root = tk.Tk()
        style = ttk.Style(root)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        for font_name in ("TkDefaultFont", "TkTextFont", "TkMenuFont", "TkHeadingFont", "TkCaptionFont"):
            named_font = tkfont.nametofont(font_name)
            named_font.configure(size=named_font.cget("size") + 2)
        style.configure("Treeview", font=("Segoe UI", 11), rowheight=28)
        style.configure("Treeview.Heading", font=("Segoe UI", 11, "bold"))

        app = WorkflowGUI(root)
        app.ensure_visible()
        app.append_log("Ready. Choose a stage button or run selected stage from dropdown.")
        app.append_chat("GUI: Ready. Type /help for quick chat commands.")

        with STARTUP_LOG.open("a", encoding="utf-8") as fh:
            fh.write(f"[{datetime.now().isoformat()}] GUI created; entering mainloop.\n")

        try:
            root.mainloop()
        except KeyboardInterrupt:
            print("GUI interrupted from the terminal.")
            root.destroy()

        with STARTUP_LOG.open("a", encoding="utf-8") as fh:
            fh.write(f"[{datetime.now().isoformat()}] Mainloop ended.\n")
    except Exception as exc:
        err = f"[{datetime.now().isoformat()}] GUI startup failed: {exc!r}\n"
        with STARTUP_LOG.open("a", encoding="utf-8") as fh:
            fh.write(err)
        print(err)
        raise


if __name__ == "__main__":
    main()
